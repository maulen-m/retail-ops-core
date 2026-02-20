#!/usr/bin/env python3
"""Prepare deterministic fixtures required by strict headless gate chain."""

from __future__ import annotations

import argparse
import json
import sqlite3
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook


def _date_window(as_of: date, *, days: int = 14) -> list[date]:
    start = as_of - timedelta(days=max(1, int(days)) - 1)
    return [start + timedelta(days=i) for i in range(max(1, int(days)))]


def _write_sales_workbook(path: Path, *, as_of: date, days: int = 14) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "SALES_KSP_CRM_1"
    ws.append(["Date", "Quantity", "Total_price", "Total_net_rev", "Sell_price_kzt"])
    for day in _date_window(as_of, days=days):
        ws.append([day.isoformat(), 1, 10000.0, 9500.0, 10000.0])
    wb.save(path)


def _write_inbound_workbook(path: Path, *, as_of: date) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "PO_part_id_Totals"
    ws.append(
        [
            "PO_part_id",
            "Status",
            "Actual_Arrival_date",
            "is_paid_BASE",
            "is_paid_DLV",
            "To_pay_BASE_KZT",
            "To_pay_DLV_KZT",
            "Est. Weight (kg)",
            "Total Bags",
            "Total Units",
        ]
    )
    ws.append(["PO-1.0", "Transit", as_of.isoformat(), "YES", "YES", 0.0, 0.0, 12.5, 5, 14])
    wb.save(path)


def _write_dim_sku_light_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "DIM_SKU_light_v5"
    ws.append(["SKU_key", "CNY", "Wt (kg)", "AvgPrc", "Active"])
    ws.append(["CL_FIX_SKU", 10.0, 1.0, 10000.0, 1])
    wb.save(path)


def _replace_with_symlink(link_path: Path, target_path: Path) -> None:
    if link_path.exists() or link_path.is_symlink():
        if link_path.is_dir() and not link_path.is_symlink():
            raise RuntimeError(f"refusing to replace directory with symlink: {link_path}")
        link_path.unlink()
    link_path.parent.mkdir(parents=True, exist_ok=True)
    link_path.symlink_to(target_path)


def _init_headless_db(path: Path, *, as_of: date, days: int = 14) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        path.unlink()

    conn = sqlite3.connect(str(path))
    try:
        conn.executescript(
            """
            CREATE TABLE dim_fx_rates (
                effective_date TEXT,
                cny_kzt REAL,
                usd_kzt REAL,
                dlv_rate_usd_kg REAL,
                source TEXT
            );
            CREATE TABLE dim_params (
                param_key TEXT,
                param_value REAL,
                product_type TEXT
            );
            CREATE TABLE dim_budget_caps (
                effective_date TEXT,
                active_flag INTEGER,
                global_monthly_cap_kzt REAL,
                per_draft_cap_kzt REAL
            );
            CREATE TABLE fact_runs (
                run_id INTEGER PRIMARY KEY AUTOINCREMENT,
                run_date TEXT,
                run_type TEXT,
                status TEXT,
                completed_at TEXT
            );
            CREATE TABLE fact_po_draft (
                draft_id TEXT,
                status TEXT,
                supplier_code TEXT,
                total_units REAL,
                total_cost_cny REAL,
                total_cost_kzt REAL,
                total_po_value_kzt REAL,
                total_order_qty REAL,
                skus_count INTEGER,
                roic_action_summary TEXT,
                guardrail_status TEXT,
                notes TEXT,
                created_at TEXT,
                updated_at TEXT
            );
            CREATE TABLE fact_po_draft_lines (
                line_id TEXT,
                draft_id TEXT,
                sku_key TEXT,
                sku_id TEXT,
                my_size TEXT,
                quantity REAL,
                unit_cost_cny REAL,
                roic_pct REAL
            );
            CREATE TABLE fact_po_approvals (
                approval_id TEXT,
                draft_id TEXT,
                sku_key TEXT,
                roic_action TEXT,
                approved_by TEXT,
                approved_at TEXT,
                decision TEXT,
                notes TEXT,
                po_value_kzt REAL,
                order_qty REAL
            );
            CREATE TABLE fact_po_execution (
                execution_id TEXT,
                draft_id TEXT,
                po_id TEXT,
                approval_id TEXT,
                planned_units REAL,
                executed_lines REAL,
                total_value_kzt REAL,
                notes TEXT,
                executed_by TEXT,
                executed_at TEXT,
                status TEXT
            );
            CREATE TABLE po_part (
                po_part_id TEXT PRIMARY KEY,
                po_id TEXT,
                status TEXT,
                est_weight_kg REAL,
                total_bags INTEGER,
                total_units INTEGER,
                is_paid_base INTEGER,
                is_paid_dlv INTEGER,
                to_pay_base_kzt REAL,
                to_pay_dlv_kzt REAL,
                base_cost_kzt REAL,
                est_delivery_kzt REAL
            );
            CREATE TABLE fact_orders_kaspi (
                order_id TEXT,
                internal_status TEXT,
                status_updated_at TEXT,
                sku_key TEXT,
                sku_id TEXT
            );
            CREATE TABLE fact_cashflow_events (
                account TEXT,
                ref_type TEXT,
                ref_id TEXT,
                amount_kzt REAL,
                event_date TEXT
            );
            CREATE TABLE fact_cashflow_daily (
                date TEXT PRIMARY KEY,
                inventory_on_delivery_close REAL,
                receivables_close REAL
            );
            CREATE TABLE fact_inventory_snapshot_size (
                snapshot_date TEXT,
                sku_key TEXT,
                current_stock REAL
            );
            CREATE TABLE dim_sku (
                sku_key TEXT PRIMARY KEY,
                base_cost_cny REAL,
                weight_kg REAL,
                cogs_kzt REAL,
                active_flag INTEGER
            );
            CREATE TABLE dim_kaspi_article_map (
                kaspi_article TEXT,
                store_code TEXT,
                sku_key TEXT,
                sku_id TEXT,
                active_flag INTEGER
            );
            CREATE TABLE dim_sku_weight_write_guard (
                guard_key TEXT PRIMARY KEY,
                allow_updates INTEGER NOT NULL DEFAULT 0,
                source TEXT,
                expires_at TEXT,
                updated_at TEXT
            );
            CREATE TABLE sales_fact_v2 (
                order_id TEXT,
                order_date TEXT,
                store_code TEXT,
                sku_key TEXT,
                sku_id TEXT,
                my_size TEXT,
                quantity REAL,
                net_rev REAL,
                cogs REAL,
                profit REAL,
                status TEXT,
                return_flag INTEGER
            );
            CREATE TABLE fact_sales (
                order_id TEXT,
                order_date TEXT,
                store_code TEXT,
                sku_key TEXT,
                sku_id TEXT,
                my_size TEXT,
                quantity REAL,
                line_net_rev REAL,
                cogs_line REAL,
                profit_line REAL
            );
            """
        )
        conn.execute(
            """
            INSERT INTO dim_fx_rates (effective_date, cny_kzt, usd_kzt, dlv_rate_usd_kg, source)
            VALUES (?, 75.0, 520.0, 2.0, 'CI_FIXTURE')
            """,
            (as_of.isoformat(),),
        )
        for key, value in (
            ("L_days", 21.0),
            ("R_days", 10.0),
            ("B_days", 14.0),
            ("z_factor", 1.65),
            ("TV_mix_floor", 0.23),
        ):
            conn.execute(
                "INSERT INTO dim_params (param_key, param_value, product_type) VALUES (?, ?, NULL)",
                (key, value),
            )
        conn.execute(
            """
            INSERT INTO dim_budget_caps (effective_date, active_flag, global_monthly_cap_kzt, per_draft_cap_kzt)
            VALUES (?, 1, 0, 0)
            """,
            (as_of.isoformat(),),
        )
        conn.execute(
            """
            INSERT INTO fact_runs (run_date, run_type, status, completed_at)
            VALUES (?, 'CI_FIXTURE', 'SUCCESS', ?)
            """,
            (as_of.isoformat(), f"{as_of.isoformat()} 00:00:00"),
        )
        conn.execute(
            """
            INSERT INTO po_part (
                po_part_id, po_id, status, est_weight_kg, total_bags, total_units,
                is_paid_base, is_paid_dlv, to_pay_base_kzt, to_pay_dlv_kzt,
                base_cost_kzt, est_delivery_kzt
            ) VALUES ('PO-1.0', 'PO-1', 'IN_TRANSIT', 12.5, 5, 14, 1, 1, 0, 0, 500000, 120000)
            """
        )
        conn.execute(
            """
            INSERT INTO fact_orders_kaspi (order_id, internal_status, status_updated_at, sku_key, sku_id)
            VALUES ('ORD-CI-1', 'COMPLETED', ?, 'CL_FIX_SKU', 'CL_FIX_SKU_M')
            """,
            (as_of.isoformat(),),
        )
        conn.execute(
            """
            INSERT INTO fact_cashflow_daily (date, inventory_on_delivery_close, receivables_close)
            VALUES (?, 0, 0)
            """,
            (as_of.isoformat(),),
        )
        conn.execute(
            """
            INSERT INTO fact_inventory_snapshot_size (snapshot_date, sku_key, current_stock)
            VALUES (?, 'CL_FIX_SKU', 100)
            """,
            (as_of.isoformat(),),
        )
        conn.execute(
            """
            INSERT INTO dim_sku (sku_key, base_cost_cny, weight_kg, cogs_kzt, active_flag)
            VALUES ('CL_FIX_SKU', 10.0, 1.0, 1790.0, 1)
            """
        )
        conn.execute(
            """
            INSERT INTO dim_sku_weight_write_guard
            (guard_key, allow_updates, source, expires_at, updated_at)
            VALUES ('dim_sku_weight_kg', 0, NULL, NULL, datetime('now'))
            """
        )
        conn.execute(
            """
            CREATE TRIGGER trg_block_dim_sku_weight_update
            BEFORE UPDATE OF weight_kg ON dim_sku
            FOR EACH ROW
            WHEN COALESCE(NEW.weight_kg, 0.0) <> COALESCE(OLD.weight_kg, 0.0)
             AND NOT EXISTS (
                SELECT 1
                FROM dim_sku_weight_write_guard
                WHERE guard_key = 'dim_sku_weight_kg'
                  AND allow_updates = 1
                  AND source = 'sync_dim_sku_from_dim_sku_light'
                  AND datetime(COALESCE(expires_at, '1970-01-01')) > datetime('now')
             )
            BEGIN
                SELECT RAISE(ABORT, 'dim_sku.weight_kg updates are blocked; use sync_dim_sku_from_dim_sku_light');
            END
            """
        )
        for idx, day in enumerate(_date_window(as_of, days=days), start=1):
            order_id = f"ORD-{day.isoformat()}"
            values = (
                order_id,
                day.isoformat(),
                "UNIVERSAL",
                "CL_FIX_SKU",
                "CL_FIX_SKU_M",
                "M",
                1.0,
                9500.0,
                1790.0,
                7710.0,
                "DELIVERED",
                0,
            )
            conn.execute(
                """
                INSERT INTO sales_fact_v2 (
                    order_id, order_date, store_code, sku_key, sku_id, my_size,
                    quantity, net_rev, cogs, profit, status, return_flag
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                values,
            )
            conn.execute(
                """
                INSERT INTO fact_sales (
                    order_id, order_date, store_code, sku_key, sku_id, my_size,
                    quantity, line_net_rev, cogs_line, profit_line
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                values[:7] + values[7:10],
            )
        conn.commit()
    finally:
        conn.close()


def _write_po_dashboard_payload(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "archived_pos": ["PO-1.0"],
        "pos": {"PO-1.0": {"po_id": "PO-1.0"}},
        "real_pos": [
            {
                "po_id": "PO-1.0",
                "weight_nom_kg": 12.5,
                "total_places": 5,
            }
        ],
        "sku_level": [{"sku_key": "CL_FIX_SKU", "cogs_unresolved_rows": 0, "profit_publishable": True}],
    }
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def _write_business_insides_snapshot(path: Path, *, as_of: date) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    snapshot_text = "\n".join(
        [
            f"# BUSINESS_INSIDES_{as_of.isoformat()}",
            "",
            "## Transparency",
            "- Unresolved COGS rows: `0`.",
            "- Unresolved SKU count: `0`.",
            "",
        ]
    )
    path.write_text(snapshot_text, encoding="utf-8")


def prepare_ci_headless_fixture(*, project_root: Path, as_of: date) -> dict[str, str]:
    root = project_root.resolve()
    fixture_dir = root / "config" / "anchors" / "fixtures"
    sales_target = fixture_dir / "SALES_KSP_CRM_V3.fixture.xlsx"
    inbound_target = fixture_dir / "INBOUND_CALENDAR_V10.002.fixture.xlsx"
    dim_sku_light_target = fixture_dir / "DIM_SKU_LIGHT_V5.fixture.xlsx"
    db_target = root / "db" / "app.db"
    crm_anchor = root / "config" / "anchors" / "SALES_KSP_CRM_LATEST.xlsx"
    inbound_anchor = root / "config" / "anchors" / "INBOUND_CALENDAR_LATEST.xlsx"
    dashboard_path = root / "exports" / "po_dashboard_data.json"
    business_insides_path = root / "config" / "business_insides" / f"BUSINESS_INSIDES_{as_of.isoformat()}.md"
    business_insides_snapshot = (
        root / "config" / "business_insides" / "snapshots" / f"BUSINESS_INSIDES_{as_of.isoformat()}.md"
    )

    _write_sales_workbook(sales_target, as_of=as_of)
    _write_inbound_workbook(inbound_target, as_of=as_of)
    _write_dim_sku_light_workbook(dim_sku_light_target)
    _init_headless_db(db_target, as_of=as_of)
    _write_po_dashboard_payload(dashboard_path)
    _write_business_insides_snapshot(business_insides_path, as_of=as_of)
    _write_business_insides_snapshot(business_insides_snapshot, as_of=as_of)
    _replace_with_symlink(crm_anchor, sales_target)
    _replace_with_symlink(inbound_anchor, inbound_target)

    return {
        "project_root": str(root),
        "fixture_dir": str(fixture_dir),
        "sales_target": str(sales_target),
        "inbound_target": str(inbound_target),
        "dim_sku_light_target": str(dim_sku_light_target),
        "db_target": str(db_target),
        "dashboard_path": str(dashboard_path),
        "business_insides_path": str(business_insides_path),
        "crm_anchor": str(crm_anchor),
        "inbound_anchor": str(inbound_anchor),
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Prepare headless fixtures for strict gate chain (anchors, db, and snapshots)"
    )
    parser.add_argument("--project-root", type=Path, default=Path(__file__).resolve().parent.parent)
    parser.add_argument("--as-of", type=str, default=date.today().isoformat())
    args = parser.parse_args()

    as_of = date.fromisoformat(args.as_of)
    result = prepare_ci_headless_fixture(project_root=args.project_root, as_of=as_of)
    for key, value in result.items():
        print(f"{key}={value}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
