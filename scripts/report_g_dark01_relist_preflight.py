#!/usr/bin/env python3
"""Build a no-write G-DARK-01 relist preflight from fresh Kaspi pricelists."""

from __future__ import annotations

import argparse
import csv
from datetime import datetime
import json
from pathlib import Path
import re
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
ALMATY_TZ = ZoneInfo("Asia/Almaty")

DEFAULT_OFFER_STATE_CSV = (
    PROJECT_ROOT
    / "exports"
    / "validation"
    / "g_dark01_relist_state"
    / "20260618_193323_0500"
    / "dark_relist_offer_state_rows.csv"
)
DEFAULT_OUTPUT_ROOT = PROJECT_ROOT / "exports" / "validation" / "g_dark01_relist_preflight"
TEMPLATE_COLUMNS = ["SKU", "model", "brand", "price", "PP1", "PP2", "PP3", "PP4", "PP5", "preorder"]
WAREHOUSE_COLUMNS = ["PP1", "PP2", "PP3", "PP4", "PP5"]
TARGET_STATUSES = {"MISSING_BUYABLE", "EXCLUDED_OR_ZERO_STOCK_BUYABLE"}
OUTPUT_COLUMNS = [
    "family_id",
    "sku_key",
    "size",
    "desired_state",
    "current_stock",
    "store",
    "source_state",
    "SKU",
    "price",
    "PP1",
    "PP2",
    "PP3",
    "PP4",
    "PP5",
    "preorder",
    "row_status",
]


def _now_almaty() -> str:
    return datetime.now(ALMATY_TZ).replace(microsecond=0).isoformat()


def _run_id(value: str) -> str:
    return value.replace("-", "").replace(":", "").replace("+", "_").replace("T", "_")


def _resolve_path(raw: str | Path) -> Path:
    path = Path(raw)
    if path.is_absolute():
        return path
    return PROJECT_ROOT / path


def _clean(raw: Any) -> str:
    text = str(raw or "").strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def _write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fieldnames})


def _load_pricelist(path: Path, *, store: str, source_state: str) -> list[dict[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Лист1" not in wb.sheetnames:
            raise ValueError(f"missing Лист1 in {path}")
        ws = wb["Лист1"]
        rows_iter = ws.iter_rows(values_only=True)
        headers = [_clean(value) for value in next(rows_iter)]
        index = {name: idx for idx, name in enumerate(headers)}
        missing = [column for column in TEMPLATE_COLUMNS if column not in index]
        if missing:
            raise ValueError(f"missing columns in {path}: {', '.join(missing)}")
        rows: list[dict[str, str]] = []
        for values in rows_iter:
            row = {column: _clean(values[index[column]] if index[column] < len(values) else "") for column in TEMPLATE_COLUMNS}
            if not row["SKU"]:
                continue
            row["store"] = store
            row["source_state"] = source_state
            rows.append(row)
        return rows
    finally:
        wb.close()


def _article_size(sku: str, sku_key: str) -> str:
    text = _clean(sku)
    prefix = f"{sku_key}_"
    if text.startswith(prefix):
        return text[len(prefix) :].split("_", 1)[0].upper()
    match = re.search(r"_(S|M|L|XL|2XL|3XL|4XL)(?:_|$)", text.upper())
    return match.group(1) if match else ""


def _target_rows(offer_state_csv: Path) -> list[dict[str, Any]]:
    targets: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    for row in _read_csv(offer_state_csv):
        status = _clean(row.get("status"))
        if status not in TARGET_STATUSES:
            continue
        sku_key = _clean(row.get("sku_key"))
        size = _clean(row.get("size")).upper()
        desired = "BUYABLE" if status == "MISSING_BUYABLE" else "NON_BUYABLE"
        key = (sku_key, size, desired)
        if key in seen:
            continue
        seen.add(key)
        targets.append(
            {
                "family_id": _clean(row.get("family_id")),
                "sku_key": sku_key,
                "size": size,
                "desired_state": desired,
                "current_stock": _clean(row.get("current_stock")),
                "source_status": status,
            }
        )
    return targets


def _mutation_update(row: dict[str, Any]) -> dict[str, str]:
    return {
        "SKU": _clean(row.get("SKU")),
        "PP1": "no",
        "PP2": "no",
        "PP3": "no",
        "PP4": "no",
        "PP5": "no",
        "preorder": "0",
        "note": "G-DARK-01 OD-033 zero-stock/excluded row off preflight",
    }


def build_relist_preflight(
    *,
    offer_state_csv: Path,
    store_files: dict[str, dict[str, Path]],
    output_root: Path,
    generated_at: str | None = None,
) -> dict[str, Any]:
    generated_at = generated_at or _now_almaty()
    out_dir = output_root / _run_id(generated_at)
    out_dir.mkdir(parents=True, exist_ok=True)

    targets = _target_rows(offer_state_csv)
    all_pricelist_rows: list[dict[str, str]] = []
    source_counts: dict[str, dict[str, int]] = {}
    for store, paths in sorted(store_files.items()):
        source_counts[store] = {}
        for state, path in sorted(paths.items()):
            rows = _load_pricelist(path, store=store, source_state=state)
            source_counts[store][state] = len(rows)
            all_pricelist_rows.extend(rows)

    target_evidence: list[dict[str, Any]] = []
    per_store_off_updates: dict[str, list[dict[str, str]]] = {}
    missing_platform_rows: list[str] = []
    archive_activation_rows: list[str] = []
    active_off_rows: list[str] = []

    for target in targets:
        matches: list[dict[str, Any]] = []
        for row in all_pricelist_rows:
            if _article_size(row["SKU"], target["sku_key"]) != target["size"]:
                continue
            if not row["SKU"].startswith(f"{target['sku_key']}_"):
                continue
            matches.append({**target, **row})
        if target["desired_state"] == "BUYABLE":
            active_matches = [row for row in matches if row["source_state"] == "ACTIVE"]
            archive_matches = [row for row in matches if row["source_state"] == "ARCHIVE"]
            if active_matches:
                for row in active_matches:
                    row["row_status"] = "OK_ALREADY_ACTIVE"
                    target_evidence.append(row)
            elif archive_matches:
                for row in archive_matches:
                    row["row_status"] = "ARCHIVE_ACTIVATION_CANDIDATE"
                    target_evidence.append(row)
                    archive_activation_rows.append(row["SKU"])
            else:
                missing = {
                    **target,
                    "store": "",
                    "source_state": "",
                    "SKU": "",
                    "price": "",
                    "PP1": "",
                    "PP2": "",
                    "PP3": "",
                    "PP4": "",
                    "PP5": "",
                    "preorder": "",
                    "row_status": "MISSING_FROM_ACTIVE_AND_ARCHIVE",
                }
                target_evidence.append(missing)
                missing_platform_rows.append(f"{target['sku_key']}:{target['size']}")
        else:
            active_matches = [row for row in matches if row["source_state"] == "ACTIVE"]
            if active_matches:
                for row in active_matches:
                    row["row_status"] = "ACTIVE_OFF_PATCH_CANDIDATE"
                    target_evidence.append(row)
                    active_off_rows.append(row["SKU"])
                    per_store_off_updates.setdefault(row["store"], []).append(_mutation_update(row))
            else:
                if matches:
                    for row in matches:
                        row["row_status"] = "OK_NOT_ACTIVE"
                        target_evidence.append(row)
                else:
                    target_evidence.append(
                        {
                            **target,
                            "store": "",
                            "source_state": "",
                            "SKU": "",
                            "price": "",
                            "PP1": "",
                            "PP2": "",
                            "PP3": "",
                            "PP4": "",
                            "PP5": "",
                            "preorder": "",
                            "row_status": "OK_NO_PLATFORM_ACTIVE_ROW",
                        }
                    )

    target_csv = out_dir / "dark_relist_preflight_target_rows.csv"
    _write_csv(target_csv, target_evidence, OUTPUT_COLUMNS)
    update_paths: dict[str, str] = {}
    for store, rows in sorted(per_store_off_updates.items()):
        update_path = out_dir / f"{store.lower().replace('-', '')}_off_updates.csv"
        _write_csv(update_path, rows, ["SKU", "PP1", "PP2", "PP3", "PP4", "PP5", "preorder", "note"])
        update_paths[store] = str(update_path)

    status = "BLOCKED_NO_WRITE" if missing_platform_rows else "READY_FOR_REVIEW"
    report = {
        "gate_id": "G-DARK-01",
        "status": status,
        "generated_at": generated_at,
        "offer_state_csv": str(offer_state_csv),
        "output_dir": str(out_dir),
        "source_counts": source_counts,
        "target_count": len(targets),
        "target_rows_csv": str(target_csv),
        "missing_platform_row_count": len(missing_platform_rows),
        "missing_platform_rows": sorted(missing_platform_rows),
        "archive_activation_candidate_count": len(set(archive_activation_rows)),
        "archive_activation_candidate_skus": sorted(set(archive_activation_rows)),
        "active_off_candidate_count": len(set(active_off_rows)),
        "active_off_candidate_skus": sorted(set(active_off_rows)),
        "off_updates_csv_by_store": update_paths,
        "safe_full_action_apply_allowed": False if missing_platform_rows else True,
        "production_db_written": False,
        "external_writes_performed": False,
        "notes": [
            "Preflight uses exact seller-article size token matching for the current G-DARK-01 artifact.",
            "Missing positive-stock marketplace rows require a separate offer-creation/catalog path or an owner-approved mapping change; no live write is applied here.",
        ],
    }
    json_path = out_dir / "dark_relist_preflight_report.json"
    md_path = out_dir / "dark_relist_preflight_report.md"
    report["json_path"] = str(json_path)
    report["md_path"] = str(md_path)
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    md_path.write_text(_render_markdown(report), encoding="utf-8")
    return report


def _render_markdown(report: dict[str, Any]) -> str:
    lines = [
        "# G-DARK-01 Relist Preflight",
        "",
        f"Status: {report['status']}",
        f"Generated: {report['generated_at']}",
        "",
        "## Summary",
        "",
        f"- target_count: {report['target_count']}",
        f"- missing_platform_row_count: {report['missing_platform_row_count']}",
        f"- archive_activation_candidate_count: {report['archive_activation_candidate_count']}",
        f"- active_off_candidate_count: {report['active_off_candidate_count']}",
        f"- safe_full_action_apply_allowed: {report['safe_full_action_apply_allowed']}",
        "",
        "## Missing Platform Rows",
        "",
    ]
    if report["missing_platform_rows"]:
        lines.extend(f"- {item}" for item in report["missing_platform_rows"])
    else:
        lines.append("- none")
    lines.extend(["", "## Candidate Off Rows", ""])
    if report["active_off_candidate_skus"]:
        lines.extend(f"- {item}" for item in report["active_off_candidate_skus"])
    else:
        lines.append("- none")
    return "\n".join(lines) + "\n"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--offer-state-csv", type=Path, default=DEFAULT_OFFER_STATE_CSV)
    parser.add_argument("--universal-active", type=Path, required=True)
    parser.add_argument("--universal-archive", type=Path, required=True)
    parser.add_argument("--storeb-active", type=Path, required=True)
    parser.add_argument("--storeb-archive", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)

    report = build_relist_preflight(
        offer_state_csv=_resolve_path(args.offer_state_csv),
        store_files={
            "UNIVERSAL": {
                "ACTIVE": _resolve_path(args.universal_active),
                "ARCHIVE": _resolve_path(args.universal_archive),
            },
            "STORE-B": {
                "ACTIVE": _resolve_path(args.storeb_active),
                "ARCHIVE": _resolve_path(args.storeb_archive),
            },
        },
        output_root=_resolve_path(args.output_root),
    )
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True))
    else:
        print(f"Status: {report['status']}")
        print(f"Report: {report['json_path']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
