#!/usr/bin/env python3
"""Hard-fail validator for Kaspi offer XLSM templates before ZIP/upload."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
import sqlite3
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DB = PROJECT_ROOT / "db" / "app.db"

MODEL_TOKEN_STOPWORDS = {
    "CL",
    "OC",
    "NK",
    "NEW",
    "CLO2",
    "CLO",
    "MEN",
    "WOMEN",
    "KO",
    "BLACK",
    "WHITE",
    "ORANGE",
    "BLUE",
    "RED",
    "GREEN",
    "GRAY",
}


def _norm(value: Any) -> str:
    return str(value or "").strip()


def _norm_token(value: Any) -> str:
    token = _norm(value).upper()
    token = re.sub(r"\s+", "", token)
    return token


def _table_exists(conn: sqlite3.Connection, table_name: str) -> bool:
    row = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",
        (table_name,),
    ).fetchone()
    return bool(row)


def _table_columns(conn: sqlite3.Connection, table_name: str) -> set[str]:
    rows = conn.execute(f"PRAGMA table_info({table_name})").fetchall()
    return {str(row[1]) for row in rows}


def _pick_column(columns: set[str], candidates: list[str]) -> str | None:
    for name in candidates:
        if name in columns:
            return name
    return None


def _parse_bool_hint(value: Any, hint: str) -> bool:
    return hint in _norm(value).lower()


def _load_values_dict(ws_values) -> dict[str, set[str]]:
    by_display: dict[str, set[str]] = {}
    max_col = int(ws_values.max_column or 0)
    max_row = int(ws_values.max_row or 0)
    for col_num in range(1, max_col + 1):
        display_name = _norm(ws_values.cell(1, col_num).value)
        if not display_name:
            continue
        allowed: set[str] = set()
        for row_num in range(2, max_row + 1):
            value = _norm(ws_values.cell(row_num, col_num).value)
            if value:
                allowed.add(value.lower())
        by_display[display_name] = allowed
    return by_display


def _token_variants(raw_token: str) -> set[str]:
    token = _norm_token(raw_token)
    if not token:
        return set()
    out = {token}
    out.add(token.replace("_", ""))
    out.add(token.replace("-", ""))
    parts = token.split("_")
    for trim in (1, 2):
        if len(parts) > trim:
            out.add("_".join(parts[:-trim]))
    return {x for x in out if x}


def _extract_model_token(raw: str) -> str | None:
    text = _norm_token(raw)
    if not text:
        return None
    matches = re.findall(r"[A-Z]+-?\d+[A-Z0-9-]*", text)
    for token in reversed(matches):
        cleaned = token.strip("-")
        if cleaned and cleaned not in MODEL_TOKEN_STOPWORDS:
            return cleaned
    return None


def _load_sku_catalog(conn: sqlite3.Connection) -> set[str]:
    if not _table_exists(conn, "dim_sku"):
        raise RuntimeError("dim_sku table missing")
    rows = conn.execute(
        """
        SELECT DISTINCT sku_key
        FROM dim_sku
        WHERE COALESCE(TRIM(sku_key), '') != ''
        """
    ).fetchall()
    return {_norm(row["sku_key"]) for row in rows}


def _load_token_to_sku(
    conn: sqlite3.Connection,
    *,
    store_code: str | None,
) -> dict[str, set[str]]:
    token_to_sku: dict[str, set[str]] = {}

    def add_pair(raw_token: str, sku_key: str) -> None:
        token = _norm_token(raw_token)
        sku = _norm(sku_key)
        if not token or not sku:
            return
        token_to_sku.setdefault(token, set()).add(sku)

    if _table_exists(conn, "dim_kaspi_article_map"):
        cols = _table_columns(conn, "dim_kaspi_article_map")
        token_col = _pick_column(cols, ["kaspi_article", "kaspi_offer_id", "offer_id", "article"])
        sku_col = _pick_column(cols, ["sku_key"])
        store_col = _pick_column(cols, ["store_code"])
        if token_col and sku_col:
            if store_code and store_col:
                rows = conn.execute(
                    f"""
                    SELECT {token_col} AS token_value, {sku_col} AS sku_key
                    FROM dim_kaspi_article_map
                    WHERE UPPER(COALESCE(TRIM({store_col}), '')) = ?
                      AND COALESCE(TRIM({token_col}), '') != ''
                      AND COALESCE(TRIM({sku_col}), '') != ''
                    """,
                    (store_code.upper(),),
                ).fetchall()
            else:
                rows = conn.execute(
                    f"""
                    SELECT {token_col} AS token_value, {sku_col} AS sku_key
                    FROM dim_kaspi_article_map
                    WHERE COALESCE(TRIM({token_col}), '') != ''
                      AND COALESCE(TRIM({sku_col}), '') != ''
                    """
                ).fetchall()
            for row in rows:
                add_pair(_norm(row["token_value"]), _norm(row["sku_key"]))

    if _table_exists(conn, "fact_offer_stock_mapper_current"):
        cols = _table_columns(conn, "fact_offer_stock_mapper_current")
        token_col = _pick_column(
            cols,
            [
                "kaspi_article",
                "offer_id",
                "kaspi_offer_id",
                "merchant_sku_key",
                "article",
            ],
        )
        sku_col = _pick_column(cols, ["sku_key"])
        store_col = _pick_column(cols, ["store_code"])
        method_col = _pick_column(cols, ["mapping_method"])
        if token_col and sku_col:
            where = [
                f"COALESCE(TRIM({token_col}), '') != ''",
                f"COALESCE(TRIM({sku_col}), '') != ''",
            ]
            params: list[Any] = []
            if store_code and store_col:
                where.append(f"UPPER(COALESCE(TRIM({store_col}), '')) = ?")
                params.append(store_code.upper())
            if method_col:
                where.append(f"COALESCE(TRIM({method_col}), '') != 'unresolved'")
            rows = conn.execute(
                f"""
                SELECT {token_col} AS token_value, {sku_col} AS sku_key
                FROM fact_offer_stock_mapper_current
                WHERE {" AND ".join(where)}
                """,
                tuple(params),
            ).fetchall()
            for row in rows:
                add_pair(_norm(row["token_value"]), _norm(row["sku_key"]))

    return token_to_sku


def _resolve_sku_candidates(
    token: str,
    token_to_sku: dict[str, set[str]],
) -> set[str]:
    candidates: set[str] = set()
    for variant in _token_variants(token):
        candidates.update(token_to_sku.get(variant, set()))
    return candidates


def validate_kaspi_offer_template(
    *,
    xlsm_path: Path,
    db_path: Path = DEFAULT_DB,
    category: str,
    store_code: str | None = None,
    expect_sku_key: str | None = None,
    expect_sku_prefix: str | None = None,
) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    if not xlsm_path.exists():
        raise FileNotFoundError(f"XLSM file not found: {xlsm_path}")

    category_norm = _norm(category).lower()
    allowed_categories = {"men-sport-suits", "men-thermal-underwear"}
    if category_norm not in allowed_categories:
        raise ValueError(f"Unsupported category: {category}")

    wb = load_workbook(xlsm_path, read_only=True, data_only=True)
    try:
        if "attributes" not in wb.sheetnames:
            raise RuntimeError("Template missing 'attributes' sheet")
        if "values" not in wb.sheetnames:
            raise RuntimeError("Template missing 'values' sheet")

        ws_attr = wb["attributes"]
        ws_values = wb["values"]

        max_col = int(ws_attr.max_column or 0)
        max_row = int(ws_attr.max_row or 0)

        columns: list[dict[str, Any]] = []
        for col_num in range(1, max_col + 1):
            col = {
                "col_num": col_num,
                "spec": _norm(ws_attr.cell(1, col_num).value),
                "machine_key": _norm(ws_attr.cell(2, col_num).value),
                "display_name": _norm(ws_attr.cell(3, col_num).value),
            }
            if not col["machine_key"] and not col["display_name"]:
                continue
            col["required"] = _parse_bool_hint(col["spec"], "обязательное поле")
            col["list_bound"] = _parse_bool_hint(col["spec"], "из списка")
            col["multi_value"] = _parse_bool_hint(col["spec"], "множество")
            columns.append(col)

        if not columns:
            raise RuntimeError("attributes sheet has no usable columns")

        by_machine = {col["machine_key"]: col for col in columns if col["machine_key"]}
        by_display = {col["display_name"]: col for col in columns if col["display_name"]}

        merchant_col = by_machine.get("merchant_sku") or by_display.get("Артикул")
        manufacturer_col = None
        for col in columns:
            mk = str(col["machine_key"] or "").lower()
            dn = str(col["display_name"] or "").lower()
            if "manufacturer code" in mk or dn == "артикул производителя":
                manufacturer_col = col
                break

        if merchant_col is None:
            raise RuntimeError("attributes sheet missing merchant_sku / Артикул column")
        if manufacturer_col is None:
            raise RuntimeError("attributes sheet missing manufacturer code / Артикул производителя column")

        value_dict = _load_values_dict(ws_values)

        with sqlite3.connect(str(db_path)) as conn:
            conn.row_factory = sqlite3.Row
            sku_catalog = _load_sku_catalog(conn)
            token_to_sku = _load_token_to_sku(conn, store_code=store_code)

        data_rows: list[dict[str, Any]] = []
        for row_num in range(4, max_row + 1):
            values = {col["col_num"]: _norm(ws_attr.cell(row_num, col["col_num"]).value) for col in columns}
            if not any(values.values()):
                continue
            data_rows.append({"row_num": row_num, "values": values})

        resolved_sku_keys: list[str] = []
        for row in data_rows:
            row_num = row["row_num"]
            values = row["values"]

            for col in columns:
                value = values[col["col_num"]]
                display = col["display_name"] or col["machine_key"] or f"col_{col['col_num']}"
                if col["required"] and not value:
                    errors.append(f"row {row_num}: missing required value in '{display}'")

                if value and col["list_bound"]:
                    if display not in value_dict:
                        errors.append(
                            f"row {row_num}: values dictionary missing for list-bound column '{display}'"
                        )
                    else:
                        allowed = value_dict[display]
                        if col["multi_value"] or "," in value:
                            parts = [x.strip() for x in value.split(",") if x.strip()]
                        else:
                            parts = [value]
                        for part in parts:
                            if part.lower() not in allowed:
                                errors.append(
                                    f"row {row_num}: value '{part}' in '{display}' "
                                    "not present in values dictionary"
                                )

            merchant = values[merchant_col["col_num"]]
            manufacturer = values[manufacturer_col["col_num"]]
            resolve_targets = [x for x in [merchant, manufacturer] if _norm(x)]
            resolved_per_token: dict[str, str] = {}

            for token in resolve_targets:
                candidates = _resolve_sku_candidates(token, token_to_sku)
                if not candidates:
                    errors.append(f"row {row_num}: unresolved SKU mapping for token '{token}'")
                    continue
                if len(candidates) > 1:
                    joined = ", ".join(sorted(candidates))
                    errors.append(f"row {row_num}: ambiguous SKU mapping for token '{token}' -> [{joined}]")
                    continue
                resolved_per_token[token] = next(iter(candidates))

            if resolved_per_token:
                unique_resolved = sorted(set(resolved_per_token.values()))
                if len(unique_resolved) > 1:
                    errors.append(
                        f"row {row_num}: merchant/manufacturer tokens resolve to different sku_keys: "
                        + ", ".join(unique_resolved)
                    )
                resolved = unique_resolved[0]
                resolved_sku_keys.append(resolved)

                if resolved not in sku_catalog:
                    errors.append(f"row {row_num}: resolved sku_key '{resolved}' not found in dim_sku")

                if expect_sku_key and resolved != expect_sku_key:
                    errors.append(
                        f"row {row_num}: resolved sku_key '{resolved}' != expected '{expect_sku_key}'"
                    )
                if expect_sku_prefix and not resolved.startswith(expect_sku_prefix):
                    errors.append(
                        f"row {row_num}: resolved sku_key '{resolved}' does not start with '{expect_sku_prefix}'"
                    )

                source_model = _extract_model_token(merchant or manufacturer)
                target_model = _extract_model_token(resolved)
                if source_model and target_model and source_model != target_model:
                    errors.append(
                        f"row {row_num}: model token mismatch source='{source_model}' target='{target_model}'"
                    )

        if not data_rows:
            warnings.append("No data rows found in attributes sheet (rows >= 4).")

        return {
            "ok": len(errors) == 0,
            "category": category_norm,
            "store_code": _norm(store_code).upper() if store_code else None,
            "xlsm_path": str(xlsm_path),
            "db_path": str(db_path),
            "row_count": len(data_rows),
            "resolved_sku_keys": sorted(set(resolved_sku_keys)),
            "error_count": len(errors),
            "warning_count": len(warnings),
            "errors": errors,
            "warnings": warnings,
        }
    finally:
        wb.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Validate Kaspi offer XLSM template and hard-fail on SKU misalignment",
    )
    parser.add_argument("--xlsm", type=Path, required=True, help="Path to Kaspi XLSM template")
    parser.add_argument("--db", type=Path, default=DEFAULT_DB, help="Path to app.db")
    parser.add_argument(
        "--category",
        choices=["men-sport-suits", "men-thermal-underwear"],
        required=True,
        help="Kaspi template category",
    )
    parser.add_argument("--store", type=str, default=None, help="Store code scope, e.g. ACMEWEAR")
    parser.add_argument("--expect-sku-key", type=str, default=None, help="Expected sku_key for all rows")
    parser.add_argument("--expect-sku-prefix", type=str, default=None, help="Expected sku_key prefix")
    parser.add_argument("--report-json", type=Path, default=None, help="Write full report as JSON")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    report = validate_kaspi_offer_template(
        xlsm_path=args.xlsm,
        db_path=args.db,
        category=args.category,
        store_code=args.store,
        expect_sku_key=args.expect_sku_key,
        expect_sku_prefix=args.expect_sku_prefix,
    )

    print(
        "kaspi_offer_template: "
        f"ok={report['ok']} rows={report['row_count']} "
        f"errors={report['error_count']} warnings={report['warning_count']}"
    )
    if report["resolved_sku_keys"]:
        print("resolved_sku_keys: " + ", ".join(report["resolved_sku_keys"]))

    for warning in report["warnings"]:
        print(f"WARNING: {warning}")
    for error in report["errors"]:
        print(f"ERROR: {error}")

    if args.report_json:
        args.report_json.parent.mkdir(parents=True, exist_ok=True)
        args.report_json.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"report_json: {args.report_json}")

    return 0 if report["ok"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
