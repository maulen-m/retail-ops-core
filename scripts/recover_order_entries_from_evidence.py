#!/usr/bin/env python3
"""Recover missing Kaspi order-entry rows from already-saved evidence.

The script is dry-run by default. Writes require both ``--apply`` and
``ENABLE_ORDER_ENTRY_RECOVERY_WRITE=1``. Applying to the repo production DB also
requires ``ENABLE_ORDER_ENTRY_RECOVERY_PROD_WRITE=1``; this starter does not
authorize that production gate.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import os
import re
import sqlite3
import sys
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DB = PROJECT_ROOT / "db" / "app.db"
DEFAULT_CURRENT_CRM = PROJECT_ROOT / "excel_ui" / "SALES_KSP_CRM_V3.xlsx"
DEFAULT_RESERVE_ARCHIVE = Path(
    "~/Documents/useful tables/Main crm spreadsheets/main tables/"
    "Purchase_orders/vibe_code_PO/Sales_archive/SALES_KSP_CRM_GPT_Sales_archive.xlsx"
)
DEFAULT_API_ENTRY_ROOTS = (
    PROJECT_ROOT / "exports/kaspi_archive_history_2024-06-06_to_2026-02-26_20260227_233318",
    Path(
        "~/Documents/useful tables/Main crm spreadsheets/main tables/"
        "Purchase_orders/vibe_code_PO/Sales_archive/"
        "kaspi_archive_history_2024-06-06_to_2026-02-27_20260228_193433_creationdate"
    ),
)
DEFAULT_FACT_ORDER_TARGET_STORES = ("STOREB", "ACMEWEAR", "UNIVERSAL")
DEFAULT_WEBUI_ARCHIVE_CSVS = (
    PROJECT_ROOT
    / "exports/webui_archive_full_parse_runs/"
    "webui_archive_full_parse_2024-06-06_to_2026-03-05_20260306_2211/"
    "final_merged/ArchiveOrders_WEBUI_MERGED_ALL_STORES.csv",
    PROJECT_ROOT
    / "exports/webui_archive_full_parse_runs/"
    "webui_archive_delta_2026-03-06_to_2026-03-19_20260320_live/"
    "final_merged/ArchiveOrders_WEBUI_MERGED_ALL_STORES.csv",
)
CURRENT_CRM_SHEET = "SALES_KSP_CRM_1"
RESERVE_ARCHIVE_SHEET = "Archive_sales"
WRITE_ENV_GATE = "ENABLE_ORDER_ENTRY_RECOVERY_WRITE"
PROD_WRITE_ENV_GATE = "ENABLE_ORDER_ENTRY_RECOVERY_PROD_WRITE"

DELIVERED_SALES_STATUSES = {"COMPLETED", "DELIVERED", "SOLD"}
STORE_ALIASES = {
    "": "UNIVERSAL",
    "UNIVERSAL": "UNIVERSAL",
    "ACMEWEAR": "ACMEWEAR",
    "ONLY FIT": "ACMEWEAR",
    "11KZ": "11KZ",
    "MELVIS": "MELVIS",
    "STOREB": "STOREB",
}
SAFE_WORKBOOK_RAW_KEYS = {
    "OrderID",
    "№ заказа",
    "STORE_NAME",
    "store_code",
    "Date",
    "Дата поступления заказа",
    "Дата изменения статуса",
    "Status",
    "Статус",
    "Quantity",
    "Количество",
    "SKU_ID_KSP",
    "Артикул",
    "SKU_ID",
    "SKU_key",
    "MY_SIZE",
    "PROBABLE_SIZE",
    "KASPI_OFFER_NAME",
    "Название товара в Kaspi Магазине",
    "Название в системе продавца",
    "Kaspi_name_source",
    "Kaspi_name_core",
    "Sell_price_kzt",
    "Total_price",
    "Сумма",
    "Delivery_fee_kzt",
    "Стоимость доставки для продавца",
    "Product_Type",
    "Категория",
    "Плановая дата передачи курьеру",
    "Склад передачи КД",
    "pack_id",
    "source_file",
    "source_file_sha256",
    "source_file_format",
    "source_row_number",
    "created_at",
    "status_change_at",
    "status_raw",
    "status_internal",
    "net_rev_kzt",
    "warehouse_code",
    "article",
    "kaspi_offer_name",
    "seller_system_name",
    "category",
    "delivery_fee_seller_kzt",
    "row_fingerprint",
    "window_since",
    "window_until",
}
PII_KEY_RE = re.compile(
    r"(phone|address|customer|покупател|адрес|телефон|отзыв|оценка|оформил)",
    re.IGNORECASE,
)
RAW_JSON_PII_RE = re.compile(
    r"(phone|address|customer_phone|pickup_or_delivery_address|адрес|телефон)",
    re.IGNORECASE,
)

if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.backup_db import backup_database  # noqa: E402


class RecoveryError(RuntimeError):
    """Raised for fail-closed recovery stoplines."""


@dataclass(frozen=True)
class TargetRow:
    sale_id: int
    order_id: str
    store_code: str
    order_date: str
    sku_key: str = ""
    sku_id: str = ""
    my_size: str = ""
    kaspi_offer_name: str = ""
    quantity: float | None = None
    sell_price_kzt: float | None = None

    @property
    def pair(self) -> tuple[str, str]:
        return (self.order_id, self.store_code)


@dataclass(frozen=True)
class EvidenceRow:
    source_name: str
    source_kind: str
    source_path: Path
    order_id: str
    store_code: str
    quantity: float
    business_date: str = ""
    offer_id: str = ""
    unit_price_kzt: float | None = None
    total_price_kzt: float | None = None
    entry_id: str = ""
    product_id: str = ""
    sku_id: str = ""
    sku_key: str = ""
    my_size: str = ""
    source_sheet: str = ""
    source_row_number: int | None = None
    source_line_number: int | None = None
    entry_index: int | None = None
    source_file_sha256: str = ""
    offer_name: str = ""
    seller_name: str = ""
    category_code: str = ""
    category_title: str = ""
    delivery_cost_kzt: float | None = None
    base_price_kzt: float | None = None
    entry_number: int | None = None
    raw_api_entry: dict[str, Any] | None = None
    item_payload: dict[str, Any] = field(default_factory=dict, compare=False)

    @property
    def pair(self) -> tuple[str, str]:
        return (self.order_id, self.store_code)

    @property
    def sufficient(self) -> bool:
        return bool(
            self.order_id
            and self.store_code
            and self.quantity
            and self.quantity > 0
            and (self.offer_id or self.product_id)
        )


@dataclass(frozen=True)
class SourceBundle:
    source_name: str
    confidence: str
    rows_by_pair: dict[tuple[str, str], list[EvidenceRow]]


@dataclass(frozen=True)
class Assignment:
    order_id: str
    store_code: str
    source_name: str
    confidence: str
    evidence_rows: list[EvidenceRow]
    target_rows: list[TargetRow]
    article_present_rows: int
    dim_article_mapped_rows: int
    direct_sku_rows: int

    @property
    def pair(self) -> tuple[str, str]:
        return (self.order_id, self.store_code)

    @property
    def target_row_count(self) -> int:
        return len(self.target_rows)


@dataclass(frozen=True)
class AcceptedNoEntryQuarantine:
    order_id: str
    store_code: str
    classification: str
    copied_temp_only: bool
    production_write_authorized: bool
    source_authority: str = ""
    notes: str = ""

    @property
    def pair(self) -> tuple[str, str]:
        return (self.order_id, self.store_code)


def norm(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    text = str(value).strip()
    if text.lower() in {"nan", "none", "null"}:
        return ""
    return text


def normalize_store(value: Any) -> str:
    text = norm(value).upper()
    return STORE_ALIASES.get(text, text or "UNIVERSAL")


def parse_date(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    text = norm(value)
    if not text:
        return ""
    if "T" in text:
        text = text.split("T", 1)[0]
    if len(text) >= 10 and text[4:5] == "-" and text[7:8] == "-":
        return text[:10]
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y/%m/%d", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return text[:10]


def coerce_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = norm(value).replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def coerce_int(value: Any) -> int | None:
    number = coerce_float(value)
    if number is None:
        return None
    return int(number)


def coerce_bool(value: Any) -> bool:
    text = norm(value).lower()
    return text in {"1", "true", "yes", "y", "approved"}


def _json_default(value: Any) -> Any:
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    if isinstance(value, Path):
        return str(value)
    return str(value)


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True, default=_json_default) + "\n",
        encoding="utf-8",
    )


def _hash_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _file_sha256(path: Path) -> str:
    if not path.exists():
        return ""
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _sqlite_integrity_check(path: Path) -> str:
    with sqlite3.connect(f"file:{path}?mode=ro", uri=True) as conn:
        row = conn.execute("PRAGMA integrity_check").fetchone()
    return str(row[0] if row else "")


def _sidecar_paths(db_path: Path) -> list[Path]:
    return [Path(f"{db_path}-wal"), Path(f"{db_path}-shm"), Path(f"{db_path}-journal")]


def _fail_on_sqlite_sidecars(db_path: Path) -> None:
    existing = [path for path in _sidecar_paths(db_path) if path.exists()]
    if existing:
        joined = ", ".join(str(path) for path in existing)
        raise RecoveryError(f"refusing production apply while SQLite sidecars exist: {joined}")


def load_accepted_no_entry_quarantine(path: Path | None) -> dict[tuple[str, str], AcceptedNoEntryQuarantine]:
    """Load an explicit copied-temp no-entry quarantine contract.

    This contract is intentionally narrow: it can only retain target rows that
    have no real item-entry evidence, and it must explicitly deny production
    write authority.
    """
    if path is None:
        return {}
    if not path.exists():
        raise RecoveryError(f"Accepted no-entry quarantine CSV missing: {path}")
    required = {
        "order_id",
        "store_code",
        "classification",
        "copied_temp_only",
        "production_write_authorized",
    }
    accepted: dict[tuple[str, str], AcceptedNoEntryQuarantine] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        if reader.fieldnames is None:
            raise RecoveryError(f"Accepted no-entry quarantine CSV has no header: {path}")
        missing = sorted(required - set(reader.fieldnames))
        if missing:
            raise RecoveryError(f"Accepted no-entry quarantine CSV missing columns: {missing}")
        for line_no, row in enumerate(reader, start=2):
            order_id = norm(row.get("order_id"))
            store_code = normalize_store(row.get("store_code"))
            classification = norm(row.get("classification")).upper()
            if not order_id:
                raise RecoveryError(f"Accepted no-entry quarantine row {line_no} missing order_id")
            if classification != "RETAINED_ORDER_ENTRY_QUARANTINE":
                raise RecoveryError(
                    "Accepted no-entry quarantine row "
                    f"{line_no} has unsupported classification={classification!r}"
                )
            item = AcceptedNoEntryQuarantine(
                order_id=order_id,
                store_code=store_code,
                classification=classification,
                copied_temp_only=coerce_bool(row.get("copied_temp_only")),
                production_write_authorized=coerce_bool(row.get("production_write_authorized")),
                source_authority=norm(row.get("source_authority")),
                notes=norm(row.get("notes")),
            )
            if not item.copied_temp_only:
                raise RecoveryError(
                    f"Accepted no-entry quarantine row {line_no} must set copied_temp_only=true"
                )
            if item.production_write_authorized:
                raise RecoveryError(
                    "Accepted no-entry quarantine row "
                    f"{line_no} must not authorize production writes"
                )
            if item.pair in accepted:
                raise RecoveryError(f"Duplicate accepted no-entry quarantine pair: {item.pair}")
            accepted[item.pair] = item
    return accepted


def _table_exists(conn: sqlite3.Connection, table: str) -> bool:
    return (
        conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,)).fetchone()
        is not None
    )


def _columns(conn: sqlite3.Connection, table: str) -> set[str]:
    if not _table_exists(conn, table):
        return set()
    return {row[1] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()}


def build_missing_targets(db_path: Path, *, as_of: str) -> list[TargetRow]:
    if not db_path.exists():
        raise RecoveryError(f"DB not found: {db_path}")
    uri = f"file:{db_path}?mode=ro"
    with sqlite3.connect(uri, uri=True) as conn:
        conn.row_factory = sqlite3.Row
        required = {"sales_fact_v2", "fact_order_entries_kaspi"}
        missing_tables = sorted(table for table in required if not _table_exists(conn, table))
        if missing_tables:
            raise RecoveryError(f"Missing required tables: {missing_tables}")
        sales_cols = _columns(conn, "sales_fact_v2")
        return_flag_clause = "AND COALESCE(s.return_flag, 0) = 0" if "return_flag" in sales_cols else ""
        as_of_clause = "AND date(s.order_date) <= date(?)" if as_of else ""
        params: tuple[Any, ...] = (as_of,) if as_of else ()
        status_sql = ",".join(repr(item) for item in sorted(DELIVERED_SALES_STATUSES))
        rows = conn.execute(
            f"""
            SELECT
                s.rowid AS row_id,
                COALESCE(s.sale_id, s.rowid) AS sale_id,
                s.order_id,
                date(s.order_date) AS order_date,
                UPPER(COALESCE(s.store_code, 'UNIVERSAL')) AS store_code,
                COALESCE(s.sku_key, '') AS sku_key,
                COALESCE(s.sku_id, '') AS sku_id,
                COALESCE(s.my_size, '') AS my_size,
                COALESCE(s.kaspi_offer_name, '') AS kaspi_offer_name,
                COALESCE(s.quantity, 0) AS quantity,
                s.sell_price_kzt AS sell_price_kzt
            FROM sales_fact_v2 s
            WHERE UPPER(COALESCE(s.status, '')) IN ({status_sql})
              {return_flag_clause}
              {as_of_clause}
              AND NOT EXISTS (
                  SELECT 1
                  FROM fact_order_entries_kaspi e
                  WHERE e.order_id = s.order_id
                    AND UPPER(COALESCE(e.store_code, 'UNIVERSAL'))
                        = UPPER(COALESCE(s.store_code, 'UNIVERSAL'))
              )
            ORDER BY date(s.order_date), store_code, s.order_id, s.rowid
            """,
            params,
        ).fetchall()
    return [
        TargetRow(
            sale_id=int(row["sale_id"]),
            order_id=norm(row["order_id"]),
            store_code=normalize_store(row["store_code"]),
            order_date=parse_date(row["order_date"]),
            sku_key=norm(row["sku_key"]),
            sku_id=norm(row["sku_id"]),
            my_size=norm(row["my_size"]),
            kaspi_offer_name=norm(row["kaspi_offer_name"]),
            quantity=coerce_float(row["quantity"]),
            sell_price_kzt=coerce_float(row["sell_price_kzt"]),
        )
        for row in rows
        if norm(row["order_id"])
    ]


def _fact_order_entry_required(row: sqlite3.Row, *, as_of: str) -> bool:
    """Mirror validate_order_entries_freshness without importing CLI-only code."""
    status_detail = norm(row["kaspi_status_detail"]).upper()
    internal_status = norm(row["internal_status"]).upper()
    kaspi_status = norm(row["kaspi_status"]).upper()

    if status_detail in {"CANCELLED", "RETURNED"}:
        return False
    if internal_status in {"CANCELLED", "RETURNED"}:
        return False

    pending_like = {"ACCEPTED_BY_MERCHANT", "APPROVED_BY_BANK", "NEW", "ASSEMBLY"}
    if status_detail in pending_like or internal_status in {"NEW", "ACCEPTED", "READY"} or kaspi_status in {"NEW", "ASSEMBLY"}:
        created_date = parse_date(row["created_at"])
        if created_date <= as_of:
            return False
    return True


def build_missing_fact_order_targets(
    db_path: Path,
    *,
    start_date: str,
    as_of: str,
    stores: tuple[str, ...],
    entry_required_only: bool = False,
) -> list[TargetRow]:
    if not db_path.exists():
        raise RecoveryError(f"DB not found: {db_path}")
    if not start_date:
        raise RecoveryError("start_date is required for fact_orders_kaspi target mode")
    if not stores:
        raise RecoveryError("at least one store is required for fact_orders_kaspi target mode")

    uri = f"file:{db_path}?mode=ro"
    with sqlite3.connect(uri, uri=True) as conn:
        conn.row_factory = sqlite3.Row
        required = {"fact_orders_kaspi", "fact_order_entries_kaspi"}
        missing_tables = sorted(table for table in required if not _table_exists(conn, table))
        if missing_tables:
            raise RecoveryError(f"Missing required tables: {missing_tables}")
        placeholders = ",".join("?" for _ in stores)
        rows = conn.execute(
            f"""
            SELECT
                COALESCE(o.id, o.rowid) AS sale_id,
                o.order_id,
                date(o.created_at) AS order_date,
                COALESCE(o.created_at, '') AS created_at,
                UPPER(COALESCE(o.store_code, 'UNIVERSAL')) AS store_code,
                COALESCE(o.sku_key, '') AS sku_key,
                COALESCE(o.sku_id, '') AS sku_id,
                COALESCE(o.my_size, '') AS my_size,
                COALESCE(o.kaspi_offer_name, '') AS kaspi_offer_name,
                COALESCE(o.quantity, 0) AS quantity,
                o.unit_price_kzt AS sell_price_kzt,
                COALESCE(o.kaspi_status_detail, '') AS kaspi_status_detail,
                COALESCE(o.internal_status, '') AS internal_status,
                COALESCE(o.kaspi_status, '') AS kaspi_status
            FROM fact_orders_kaspi o
            WHERE date(o.created_at) BETWEEN date(?) AND date(?)
              AND UPPER(COALESCE(o.store_code, 'UNIVERSAL')) IN ({placeholders})
              AND NOT EXISTS (
                  SELECT 1
                  FROM fact_order_entries_kaspi e
                  WHERE e.order_id = o.order_id
                    AND UPPER(COALESCE(e.store_code, 'UNIVERSAL'))
                        = UPPER(COALESCE(o.store_code, 'UNIVERSAL'))
              )
            ORDER BY date(o.created_at), store_code, o.order_id, o.rowid
            """,
            (start_date, as_of, *stores),
        ).fetchall()

    if entry_required_only:
        rows = [row for row in rows if _fact_order_entry_required(row, as_of=as_of)]

    return [
        TargetRow(
            sale_id=int(row["sale_id"]),
            order_id=norm(row["order_id"]),
            store_code=normalize_store(row["store_code"]),
            order_date=parse_date(row["order_date"]),
            sku_key=norm(row["sku_key"]),
            sku_id=norm(row["sku_id"]),
            my_size=norm(row["my_size"]),
            kaspi_offer_name=norm(row["kaspi_offer_name"]),
            quantity=coerce_float(row["quantity"]),
            sell_price_kzt=coerce_float(row["sell_price_kzt"]),
        )
        for row in rows
        if norm(row["order_id"])
    ]


def load_target_order_pairs(path: Path | None) -> set[tuple[str, str]]:
    if path is None:
        return set()
    if not path.exists():
        raise RecoveryError(f"Target order CSV missing: {path}")
    required = {"order_id", "store_code"}
    pairs: set[tuple[str, str]] = set()
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        if reader.fieldnames is None:
            raise RecoveryError(f"Target order CSV has no header: {path}")
        missing = sorted(required - set(reader.fieldnames))
        if missing:
            raise RecoveryError(f"Target order CSV missing columns: {missing}")
        for line_no, row in enumerate(reader, start=2):
            order_id = norm(row.get("order_id"))
            store_code = normalize_store(row.get("store_code"))
            if not order_id:
                raise RecoveryError(f"Target order CSV row {line_no} missing order_id")
            pair = (order_id, store_code)
            if pair in pairs:
                raise RecoveryError(f"Target order CSV duplicate order/store pair: {pair}")
            pairs.add(pair)
    if not pairs:
        raise RecoveryError(f"Target order CSV has no order rows: {path}")
    return pairs


def filter_targets_by_order_pairs(
    targets: list[TargetRow],
    *,
    requested_pairs: set[tuple[str, str]],
) -> tuple[list[TargetRow], dict[str, Any]]:
    if not requested_pairs:
        return targets, {
            "applied": False,
            "requested_order_store_pairs": 0,
            "matched_order_store_pairs": 0,
            "unmatched_order_store_pairs": [],
        }
    target_pairs = {target.pair for target in targets}
    unmatched_pairs = sorted(requested_pairs - target_pairs)
    if unmatched_pairs:
        raise RecoveryError(
            "Target order CSV contains pairs that are not in the current missing target set: "
            f"{[{'order_id': order_id, 'store_code': store} for order_id, store in unmatched_pairs]}"
        )
    filtered = [target for target in targets if target.pair in requested_pairs]
    matched_pairs = {target.pair for target in filtered}
    return filtered, {
        "applied": True,
        "requested_order_store_pairs": len(requested_pairs),
        "matched_order_store_pairs": len(matched_pairs),
        "unmatched_order_store_pairs": [],
    }


def load_article_map(db_path: Path) -> dict[tuple[str, str], dict[str, str]]:
    uri = f"file:{db_path}?mode=ro"
    with sqlite3.connect(uri, uri=True) as conn:
        conn.row_factory = sqlite3.Row
        if not _table_exists(conn, "dim_kaspi_article_map"):
            return {}
        rows = conn.execute(
            """
            SELECT
                UPPER(COALESCE(store_code, '')) AS store_code,
                UPPER(COALESCE(kaspi_article, '')) AS kaspi_article,
                MAX(COALESCE(sku_key, '')) AS sku_key,
                MAX(COALESCE(sku_id, '')) AS sku_id
            FROM dim_kaspi_article_map
            GROUP BY UPPER(COALESCE(store_code, '')), UPPER(COALESCE(kaspi_article, ''))
            """
        ).fetchall()
    out: dict[tuple[str, str], dict[str, str]] = {}
    for row in rows:
        article = norm(row["kaspi_article"]).upper()
        if not article:
            continue
        out[(normalize_store(row["store_code"]), article)] = {
            "sku_key": norm(row["sku_key"]),
            "sku_id": norm(row["sku_id"]),
        }
    return out


def _header_map(headers: list[Any]) -> dict[str, int]:
    out: dict[str, int] = {}
    for idx, header in enumerate(headers):
        name = norm(header)
        if name and name not in out:
            out[name] = idx
    return out


def _pick(row: tuple[Any, ...], header_map: dict[str, int], *names: str) -> Any:
    for name in names:
        idx = header_map.get(name)
        if idx is not None and idx < len(row):
            value = row[idx]
            if norm(value):
                return value
    return None


def _safe_payload(payload: dict[str, Any]) -> dict[str, Any]:
    safe: dict[str, Any] = {}
    for key, value in payload.items():
        key_text = norm(key)
        if not key_text:
            continue
        if key_text not in SAFE_WORKBOOK_RAW_KEYS:
            continue
        if PII_KEY_RE.search(key_text):
            continue
        safe[key_text] = value
    return safe


def redacted_workbook_raw_json(row: EvidenceRow, item_payload: dict[str, Any] | None = None) -> str:
    payload = _safe_payload(item_payload or row.item_payload)
    raw = {
        "recovery_source": {
            "source_name": row.source_name,
            "source_kind": row.source_kind,
            "source_path": str(row.source_path),
            "source_sheet": row.source_sheet or None,
            "source_row_number": row.source_row_number,
            "source_line_number": row.source_line_number,
            "entry_index": row.entry_index,
            "source_file_sha256": row.source_file_sha256 or None,
        },
        "item": payload,
    }
    text = json.dumps(raw, ensure_ascii=False, sort_keys=True, default=_json_default)
    if RAW_JSON_PII_RE.search(text):
        raise RecoveryError(f"PII-like key leaked into workbook-derived raw_json for {row.order_id}")
    return text


def read_workbook_evidence(
    path: Path,
    sheet_name: str,
    *,
    source_name: str,
    source_kind: str,
) -> dict[tuple[str, str], list[EvidenceRow]]:
    if not path.exists():
        raise RecoveryError(f"Workbook evidence source missing: {path}")
    file_sha = _file_sha256(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise RecoveryError(f"Workbook sheet missing: {path}::{sheet_name}")
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            headers = next(rows_iter)
        except StopIteration:
            return {}
        headers_by_name = _header_map(list(headers))
        out: dict[tuple[str, str], list[EvidenceRow]] = defaultdict(list)
        for row_no, row in enumerate(rows_iter, start=2):
            order_id = norm(_pick(row, headers_by_name, "OrderID", "№ заказа"))
            if not order_id:
                continue
            store = normalize_store(_pick(row, headers_by_name, "STORE_NAME", "store_code"))
            quantity = coerce_float(_pick(row, headers_by_name, "Quantity", "Количество"))
            business_date = parse_date(
                _pick(row, headers_by_name, "Date", "Дата изменения статуса", "Дата поступления заказа")
            )
            offer_id = norm(_pick(row, headers_by_name, "SKU_ID_KSP", "Артикул"))
            sku_id = norm(_pick(row, headers_by_name, "SKU_ID"))
            sku_key = norm(_pick(row, headers_by_name, "SKU_key"))
            my_size = norm(_pick(row, headers_by_name, "MY_SIZE"))
            offer_name = norm(
                _pick(row, headers_by_name, "KASPI_OFFER_NAME", "Название товара в Kaspi Магазине")
            )
            seller_name = norm(
                _pick(row, headers_by_name, "Kaspi_name_source", "Название в системе продавца")
            )
            unit_price = coerce_float(_pick(row, headers_by_name, "Sell_price_kzt", "Сумма"))
            total_price = coerce_float(_pick(row, headers_by_name, "Total_price", "Сумма"))
            item_payload = {
                header: row[idx]
                for header, idx in headers_by_name.items()
                if idx < len(row) and row[idx] is not None and norm(row[idx])
            }
            evidence = EvidenceRow(
                source_name=source_name,
                source_kind=source_kind,
                source_path=path,
                source_sheet=sheet_name,
                source_row_number=row_no,
                source_file_sha256=file_sha,
                order_id=order_id,
                store_code=store,
                quantity=quantity or 0.0,
                business_date=business_date,
                offer_id=offer_id,
                sku_id=sku_id,
                sku_key=sku_key,
                my_size=my_size,
                unit_price_kzt=unit_price,
                total_price_kzt=total_price,
                delivery_cost_kzt=coerce_float(
                    _pick(row, headers_by_name, "Delivery_fee_kzt", "Стоимость доставки для продавца")
                ),
                offer_name=offer_name,
                seller_name=seller_name,
                category_title=norm(_pick(row, headers_by_name, "Категория")),
                item_payload=item_payload,
            )
            out[evidence.pair].append(evidence)
        return dict(out)
    finally:
        wb.close()


def read_api_entry_evidence(root: Path) -> dict[tuple[str, str], list[EvidenceRow]]:
    out: dict[tuple[str, str], list[EvidenceRow]] = defaultdict(list)
    if not root.exists():
        raise RecoveryError(f"API entry evidence root missing: {root}")
    for path in sorted(root.glob("store_*/archive_order_entries_raw.jsonl")):
        store_from_path = path.parent.name.replace("store_", "")
        with path.open("r", encoding="utf-8") as fh:
            for line_no, line in enumerate(fh, start=1):
                text = line.strip()
                if not text:
                    continue
                payload = json.loads(text)
                order_id = norm(payload.get("order_code"))
                store = normalize_store(payload.get("store_code") or store_from_path)
                for entry_index, entry in enumerate(payload.get("entries") or []):
                    attrs = entry.get("attributes") or {}
                    relationships = entry.get("relationships") or {}
                    offer = attrs.get("offer") or {}
                    category = attrs.get("category") or {}
                    product_rel = (relationships.get("product") or {}).get("data") or {}
                    delivery_pos = (relationships.get("deliveryPointOfService") or {}).get("data") or {}
                    point_pos = (relationships.get("pointOfService") or {}).get("data") or {}
                    quantity = coerce_float(attrs.get("quantity")) or 0.0
                    evidence = EvidenceRow(
                        source_name="API_RAW_ORDER_ENTRIES",
                        source_kind="api_raw_order_entry",
                        source_path=path,
                        source_line_number=line_no,
                        entry_index=entry_index,
                        entry_id=norm(entry.get("id")),
                        order_id=order_id,
                        store_code=store,
                        quantity=quantity,
                        offer_id=norm(offer.get("code") or attrs.get("offerId")),
                        product_id=norm(product_rel.get("id")),
                        unit_price_kzt=coerce_float(attrs.get("basePrice") or attrs.get("price")),
                        total_price_kzt=coerce_float(attrs.get("totalPrice")),
                        delivery_cost_kzt=coerce_float(attrs.get("deliveryCost")),
                        base_price_kzt=coerce_float(attrs.get("basePrice")),
                        entry_number=coerce_int(attrs.get("entryNumber")),
                        category_code=norm(category.get("code")),
                        category_title=norm(category.get("title")),
                        offer_name=norm(offer.get("name")),
                        raw_api_entry={
                            **entry,
                            "_recovery_point_of_service_id": norm(point_pos.get("id")),
                            "_recovery_delivery_point_of_service_id": norm(delivery_pos.get("id")),
                        },
                    )
                    out[evidence.pair].append(evidence)
    return dict(out)


def read_webui_csv_evidence(paths: list[Path]) -> dict[tuple[str, str], list[EvidenceRow]]:
    out: dict[tuple[str, str], list[EvidenceRow]] = defaultdict(list)
    for path in paths:
        if not path.exists():
            raise RecoveryError(f"WebUI archive CSV missing: {path}")
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            reader = csv.DictReader(fh)
            for csv_row_no, row in enumerate(reader, start=2):
                order_id = norm(row.get("order_id") or row.get("№ заказа"))
                if not order_id:
                    continue
                source_row_number = coerce_int(row.get("source_row_number")) or csv_row_no
                store = normalize_store(row.get("store_code") or row.get("STORE_NAME"))
                quantity = coerce_float(row.get("quantity") or row.get("Количество")) or 0.0
                item_payload = {
                    key: value
                    for key, value in row.items()
                    if value is not None and norm(value) and key in SAFE_WORKBOOK_RAW_KEYS
                }
                evidence = EvidenceRow(
                    source_name="WEBUI_ARCHIVE_SOURCE_BACKUP",
                    source_kind="webui_archive_source_backup",
                    source_path=path,
                    source_row_number=source_row_number,
                    source_file_sha256=norm(row.get("source_file_sha256")),
                    order_id=order_id,
                    store_code=store,
                    quantity=quantity,
                    business_date=parse_date(
                        row.get("status_change_at")
                        or row.get("Date")
                        or row.get("Дата изменения статуса")
                        or row.get("Дата поступления заказа")
                    ),
                    offer_id=norm(row.get("article") or row.get("Артикул")),
                    total_price_kzt=coerce_float(row.get("net_rev_kzt") or row.get("Сумма")),
                    delivery_cost_kzt=coerce_float(
                        row.get("delivery_fee_seller_kzt") or row.get("Стоимость доставки для продавца")
                    ),
                    offer_name=norm(row.get("kaspi_offer_name") or row.get("Название товара в Kaspi Магазине")),
                    seller_name=norm(row.get("seller_system_name") or row.get("Название в системе продавца")),
                    category_title=norm(row.get("category") or row.get("Категория")),
                    item_payload=item_payload,
                )
                out[evidence.pair].append(evidence)
    return dict(out)


def merge_source_maps(*maps: dict[tuple[str, str], list[EvidenceRow]]) -> dict[tuple[str, str], list[EvidenceRow]]:
    out: dict[tuple[str, str], list[EvidenceRow]] = defaultdict(list)
    for source_map in maps:
        for pair, rows in source_map.items():
            out[pair].extend(rows)
    return dict(out)


def dedupe_evidence_rows(rows: list[EvidenceRow]) -> list[EvidenceRow]:
    """Project evidence to one current row per source-stable line identity.

    Real API entries retain their immutable entry IDs. Workbook evidence has no
    entry ID, so row numbers are provenance/ranking only: repeated lifecycle
    snapshots for the same order/store/public offer supersede each other rather
    than becoming separate order lines.
    """
    out: list[EvidenceRow] = []
    positions: dict[tuple[Any, ...], int] = {}

    def _position(row: EvidenceRow) -> tuple[str, int, int, int]:
        return (
            row.business_date or "",
            int(row.source_row_number or -1),
            int(row.source_line_number or -1),
            int(row.entry_index or -1),
        )

    def _material(row: EvidenceRow) -> tuple[Any, ...]:
        return (
            row.order_id,
            row.store_code,
            row.offer_id.upper(),
            row.product_id.upper(),
            float(row.quantity or 0.0),
            row.unit_price_kzt,
            row.total_price_kzt,
            row.sku_key,
            row.sku_id,
            row.my_size,
        )

    for row in rows:
        if row.entry_id:
            sig = (
                "API_ENTRY_ID",
                row.entry_id,
            )
        else:
            public_offer_identity = row.offer_id.upper() or row.product_id.upper()
            if not public_offer_identity:
                raise RecoveryError(
                    f"non-API evidence lacks public-offer identity for {row.order_id}"
                )
            sig = (
                "WORKBOOK_PUBLIC_OFFER",
                row.order_id,
                row.store_code,
                public_offer_identity,
            )
        existing_position = positions.get(sig)
        if existing_position is None:
            positions[sig] = len(out)
            out.append(row)
            continue
        if row.entry_id:
            if _material(row) != _material(out[existing_position]):
                raise RecoveryError(
                    f"conflicting payloads for immutable API entry_id {row.entry_id}"
                )
            continue
        current = out[existing_position]
        if _position(row) == _position(current) and _material(row) != _material(current):
            raise RecoveryError(
                "irreconcilable workbook snapshots share the same business-date/rank: "
                f"{row.order_id}:{row.store_code}:{public_offer_identity}"
            )
        if _position(row) >= _position(current):
            out[existing_position] = row
    return out


def assign_evidence_to_targets(
    targets: list[TargetRow],
    source_bundles: list[SourceBundle],
    *,
    article_map: dict[tuple[str, str], dict[str, str]],
) -> dict[tuple[str, str], Assignment]:
    targets_by_pair: dict[tuple[str, str], list[TargetRow]] = defaultdict(list)
    for target in targets:
        targets_by_pair[target.pair].append(target)

    assignments: dict[tuple[str, str], Assignment] = {}
    for pair, target_rows in sorted(targets_by_pair.items(), key=lambda item: item[0]):
        chosen_source = "UNRECOVERED_QUARANTINE"
        chosen_confidence = "NONE"
        chosen_rows: list[EvidenceRow] = []
        for bundle in source_bundles:
            evidence_rows = dedupe_evidence_rows([row for row in bundle.rows_by_pair.get(pair, []) if row.sufficient])
            if evidence_rows:
                chosen_source = bundle.source_name
                chosen_confidence = bundle.confidence
                chosen_rows = evidence_rows
                break
        article_present = 0
        dim_mapped = 0
        direct_sku = 0
        for row in chosen_rows:
            article = row.offer_id.upper()
            if row.sku_id or (row.sku_key and row.my_size):
                direct_sku += 1
            if article:
                article_present += 1
                if (row.store_code, article) in article_map or ("", article) in article_map:
                    dim_mapped += 1
        assignments[pair] = Assignment(
            order_id=pair[0],
            store_code=pair[1],
            source_name=chosen_source,
            confidence=chosen_confidence,
            evidence_rows=chosen_rows,
            target_rows=target_rows,
            article_present_rows=article_present,
            dim_article_mapped_rows=dim_mapped,
            direct_sku_rows=direct_sku,
        )
    return assignments


def make_non_api_entry_id(row: EvidenceRow) -> str:
    public_offer_identity = row.offer_id.upper() or row.product_id.upper()
    if not public_offer_identity:
        raise RecoveryError(
            f"non-API evidence lacks public-offer identity for {row.order_id}"
        )
    seed = {
        "identity_version": 2,
        "order_id": row.order_id,
        "store_code": row.store_code,
        "public_offer_identity": public_offer_identity,
    }
    return f"RECOV-WORKBOOK-{_hash_text(json.dumps(seed, sort_keys=True, default=_json_default))[:32]}"


def _raw_json_for_entry(row: EvidenceRow) -> str:
    source = {
        "source_name": row.source_name,
        "source_kind": row.source_kind,
        "source_path": str(row.source_path),
        "source_line_number": row.source_line_number,
        "entry_index": row.entry_index,
    }
    if row.source_kind == "api_raw_order_entry":
        return json.dumps(
            {"recovery_source": source, "entry": row.raw_api_entry or {}},
            ensure_ascii=False,
            sort_keys=True,
            default=_json_default,
        )
    return redacted_workbook_raw_json(row)


def evidence_to_entry(
    row: EvidenceRow,
    *,
    recovery_ts: str,
    article_mapped: bool,
) -> dict[str, Any]:
    entry_id = row.entry_id if row.source_kind == "api_raw_order_entry" and row.entry_id else make_non_api_entry_id(row)
    raw_json = _raw_json_for_entry(row)
    if row.source_kind != "api_raw_order_entry" and RAW_JSON_PII_RE.search(raw_json):
        raise RecoveryError(f"PII-like workbook raw_json leak for {row.order_id}")
    return {
        "entry_id": entry_id,
        "order_id": row.order_id,
        "store_code": row.store_code,
        "product_id": row.product_id or None,
        "offer_id": row.offer_id or None,
        "quantity": row.quantity,
        "unit_price_kzt": row.unit_price_kzt,
        "total_price_kzt": row.total_price_kzt,
        "raw_json": raw_json,
        "updated_at": recovery_ts,
        "unit_type": None,
        "min_allowed_weight": None,
        "weight_kg": None,
        "entry_number": row.entry_number,
        "category_code": row.category_code or None,
        "category_title": row.category_title or None,
        "delivery_cost_kzt": row.delivery_cost_kzt,
        "base_price_kzt": row.base_price_kzt,
        "point_of_service_id": None,
        "delivery_point_of_service_id": None,
        "sku_rebuild_mappable": bool(article_mapped or row.sku_id or (row.sku_key and row.my_size)),
        "recovery_source_name": row.source_name,
        "recovery_source_kind": row.source_kind,
    }


def _target_summary(targets: list[TargetRow]) -> dict[str, Any]:
    by_store: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"validator_rows": 0, "order_store_pairs": set(), "min_date": None, "max_date": None}
    )
    for target in targets:
        item = by_store[target.store_code]
        item["validator_rows"] += 1
        item["order_store_pairs"].add(target.pair)
        day = target.order_date
        item["min_date"] = day if item["min_date"] is None or day < item["min_date"] else item["min_date"]
        item["max_date"] = day if item["max_date"] is None or day > item["max_date"] else item["max_date"]
    return {
        "validator_rows": len(targets),
        "order_store_pairs": len({target.pair for target in targets}),
        "order_store_date_pairs": len({(target.order_id, target.store_code, target.order_date) for target in targets}),
        "min_date": min((target.order_date for target in targets), default=None),
        "max_date": max((target.order_date for target in targets), default=None),
        "by_store": {
            store: {
                "validator_rows": data["validator_rows"],
                "order_store_pairs": len(data["order_store_pairs"]),
                "min_date": data["min_date"],
                "max_date": data["max_date"],
            }
            for store, data in sorted(by_store.items())
        },
    }


def _assignment_summary(assignments: dict[tuple[str, str], Assignment]) -> dict[str, dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = {}
    for assignment in assignments.values():
        source = assignment.source_name
        if source not in grouped:
            grouped[source] = {
                "order_store_pairs": 0,
                "target_rows": 0,
                "evidence_rows": 0,
                "article_present_rows": 0,
                "dim_article_mapped_rows": 0,
                "direct_sku_rows": 0,
                "by_store": defaultdict(lambda: {"pairs": 0, "target_rows": 0, "evidence_rows": 0}),
            }
        item = grouped[source]
        item["order_store_pairs"] += 1
        item["target_rows"] += assignment.target_row_count
        item["evidence_rows"] += len(assignment.evidence_rows)
        item["article_present_rows"] += assignment.article_present_rows
        item["dim_article_mapped_rows"] += assignment.dim_article_mapped_rows
        item["direct_sku_rows"] += assignment.direct_sku_rows
        store_item = item["by_store"][assignment.store_code]
        store_item["pairs"] += 1
        store_item["target_rows"] += assignment.target_row_count
        store_item["evidence_rows"] += len(assignment.evidence_rows)
    for item in grouped.values():
        item["by_store"] = {store: dict(data) for store, data in sorted(item["by_store"].items())}
    return grouped


def _accepted_no_entry_quarantine_summary(
    assignments: dict[tuple[str, str], Assignment],
    accepted: dict[tuple[str, str], AcceptedNoEntryQuarantine],
) -> dict[str, Any]:
    quarantine_assignments = {
        pair: assignment
        for pair, assignment in assignments.items()
        if assignment.source_name == "UNRECOVERED_QUARANTINE"
    }
    quarantine_pairs = set(quarantine_assignments)
    accepted_pairs = set(accepted)
    matched_pairs = sorted(quarantine_pairs & accepted_pairs)
    missing_pairs = sorted(quarantine_pairs - accepted_pairs)
    unexpected_pairs = sorted(accepted_pairs - quarantine_pairs)
    accepted_target_rows = sum(quarantine_assignments[pair].target_row_count for pair in matched_pairs)
    quarantine_target_rows = sum(assignment.target_row_count for assignment in quarantine_assignments.values())
    by_store: dict[str, dict[str, int]] = defaultdict(lambda: {"pairs": 0, "target_rows": 0})
    for pair in matched_pairs:
        assignment = quarantine_assignments[pair]
        item = by_store[assignment.store_code]
        item["pairs"] += 1
        item["target_rows"] += assignment.target_row_count
    return {
        "contract_rows": len(accepted),
        "contract_pairs": len(accepted_pairs),
        "matched_pairs": len(matched_pairs),
        "accepted_target_rows": accepted_target_rows,
        "quarantine_target_rows": quarantine_target_rows,
        "missing_pairs": [{"order_id": order_id, "store_code": store} for order_id, store in missing_pairs],
        "unexpected_pairs": [{"order_id": order_id, "store_code": store} for order_id, store in unexpected_pairs],
        "all_unrecovered_targets_accepted": bool(
            quarantine_target_rows > 0
            and accepted
            and accepted_target_rows == quarantine_target_rows
            and not missing_pairs
            and not unexpected_pairs
        ),
        "by_store": {store: dict(data) for store, data in sorted(by_store.items())},
    }


def _candidate_entries(
    assignments: dict[tuple[str, str], Assignment],
    article_map: dict[tuple[str, str], dict[str, str]],
    *,
    recovery_ts: str,
) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    seen_entry_ids: set[str] = set()
    for assignment in assignments.values():
        if assignment.source_name == "UNRECOVERED_QUARANTINE":
            continue
        for row in assignment.evidence_rows:
            article_mapped = bool(
                row.offer_id
                and ((row.store_code, row.offer_id.upper()) in article_map or ("", row.offer_id.upper()) in article_map)
            )
            entry = evidence_to_entry(row, recovery_ts=recovery_ts, article_mapped=article_mapped)
            if entry["entry_id"] in seen_entry_ids:
                continue
            seen_entry_ids.add(entry["entry_id"])
            candidates.append(entry)
    return candidates


def _existing_entry_ids(conn: sqlite3.Connection, entry_ids: list[str]) -> set[str]:
    if not entry_ids:
        return set()
    existing: set[str] = set()
    chunk_size = 500
    for idx in range(0, len(entry_ids), chunk_size):
        chunk = entry_ids[idx : idx + chunk_size]
        placeholders = ",".join("?" for _ in chunk)
        for row in conn.execute(
            f"SELECT entry_id FROM fact_order_entries_kaspi WHERE entry_id IN ({placeholders})",
            chunk,
        ).fetchall():
            existing.add(str(row[0]))
    return existing


def _logical_entry_key(entry: dict[str, Any]) -> tuple[str, str, str]:
    public_offer_identity = str(entry.get("offer_id") or entry.get("product_id") or "").strip().upper()
    return (
        str(entry.get("order_id") or "").strip(),
        str(entry.get("store_code") or "").strip().upper(),
        public_offer_identity,
    )


def _legacy_logical_entry_collisions(
    conn: sqlite3.Connection,
    candidates: list[dict[str, Any]],
    exact_existing_ids: set[str],
) -> list[dict[str, Any]]:
    target_pairs = sorted(
        {
            (key[0], key[1])
            for key in (_logical_entry_key(candidate) for candidate in candidates)
            if key[0] and key[1]
        }
    )
    existing_by_key: dict[tuple[str, str, str], list[str]] = defaultdict(list)
    for order_id, store_code in target_pairs:
        rows = conn.execute(
            """
            SELECT entry_id, order_id, store_code, offer_id, product_id
            FROM fact_order_entries_kaspi
            WHERE CAST(order_id AS TEXT) = ? AND UPPER(COALESCE(store_code, '')) = ?
            """,
            (order_id, store_code),
        ).fetchall()
        for row in rows:
            payload = dict(row) if isinstance(row, sqlite3.Row) else {
                "entry_id": row[0],
                "order_id": row[1],
                "store_code": row[2],
                "offer_id": row[3],
                "product_id": row[4],
            }
            existing_by_key[_logical_entry_key(payload)].append(str(payload["entry_id"]))
    collisions: list[dict[str, Any]] = []
    for candidate in candidates:
        entry_id = str(candidate.get("entry_id") or "")
        existing_ids = existing_by_key.get(_logical_entry_key(candidate), [])
        legacy_ids = [value for value in existing_ids if value not in exact_existing_ids]
        if legacy_ids and entry_id not in existing_ids:
            collisions.append(
                {
                    "order_id": candidate.get("order_id"),
                    "store_code": candidate.get("store_code"),
                    "offer_id": candidate.get("offer_id"),
                    "candidate_entry_id": entry_id,
                    "existing_entry_ids": legacy_ids,
                }
            )
    return collisions


def _insert_entries(conn: sqlite3.Connection, entries: list[dict[str, Any]]) -> int:
    cols = _columns(conn, "fact_order_entries_kaspi")
    if not cols:
        raise RecoveryError("fact_order_entries_kaspi table missing")
    table_cols = [
        "entry_id",
        "order_id",
        "store_code",
        "product_id",
        "offer_id",
        "quantity",
        "unit_price_kzt",
        "total_price_kzt",
        "raw_json",
        "updated_at",
        "unit_type",
        "min_allowed_weight",
        "weight_kg",
        "entry_number",
        "category_code",
        "category_title",
        "delivery_cost_kzt",
        "base_price_kzt",
        "point_of_service_id",
        "delivery_point_of_service_id",
    ]
    insert_cols = [col for col in table_cols if col in cols]
    placeholders = ",".join("?" for _ in insert_cols)
    columns_sql = ",".join(insert_cols)
    inserted = 0
    for entry in entries:
        values = [entry.get(col) for col in insert_cols]
        cur = conn.execute(
            f"INSERT INTO fact_order_entries_kaspi ({columns_sql}) VALUES ({placeholders})",
            values,
        )
        inserted += int(cur.rowcount or 0)
    return inserted


def _is_production_db(db_path: Path) -> bool:
    try:
        resolved = db_path.resolve()
        production = DEFAULT_DB.resolve()
    except FileNotFoundError:
        resolved = db_path.absolute()
        production = DEFAULT_DB.absolute()
    return resolved == production


def _guard_apply_path(
    db_path: Path,
    *,
    expected_pre_sha256: str | None = None,
    backup_dir: Path | None = None,
) -> bool:
    if os.environ.get(WRITE_ENV_GATE) != "1":
        raise RecoveryError(f"{WRITE_ENV_GATE}=1 is required with --apply")
    production = _is_production_db(db_path)
    if production and os.environ.get(PROD_WRITE_ENV_GATE) != "1":
        raise RecoveryError(
            f"Refusing production db/app.db apply without {PROD_WRITE_ENV_GATE}=1"
        )
    if production and not expected_pre_sha256:
        raise RecoveryError("--expected-pre-sha256 is required for production apply")
    if production and backup_dir is None:
        raise RecoveryError("--backup-dir is required for production apply")
    if production:
        _fail_on_sqlite_sidecars(db_path)
    return production


def _write_preview_outputs(
    output_root: Path,
    *,
    summary: dict[str, Any],
    entries: list[dict[str, Any]],
    assignments: dict[tuple[str, str], Assignment],
    accepted_no_entry_quarantine: dict[tuple[str, str], AcceptedNoEntryQuarantine],
) -> None:
    output_root.mkdir(parents=True, exist_ok=True)
    _write_json(output_root / "summary.json", summary)
    with (output_root / "recovered_entries_preview.jsonl").open("w", encoding="utf-8") as fh:
        for entry in entries:
            preview = {
                key: value
                for key, value in entry.items()
                if key
                in {
                    "entry_id",
                    "order_id",
                    "store_code",
                    "product_id",
                    "offer_id",
                    "quantity",
                    "unit_price_kzt",
                    "total_price_kzt",
                    "updated_at",
                    "category_code",
                    "category_title",
                    "delivery_cost_kzt",
                    "base_price_kzt",
                    "sku_rebuild_mappable",
                    "recovery_source_name",
                    "recovery_source_kind",
                }
            }
            preview["raw_json_sha256"] = _hash_text(str(entry.get("raw_json") or ""))
            fh.write(json.dumps(preview, ensure_ascii=False, sort_keys=True, default=_json_default) + "\n")
    with (output_root / "quarantine_preview.jsonl").open("w", encoding="utf-8") as fh:
        for assignment in assignments.values():
            if assignment.source_name != "UNRECOVERED_QUARANTINE":
                continue
            for target in assignment.target_rows:
                fh.write(
                    json.dumps(
                        {
                            "sale_id": target.sale_id,
                            "order_id": target.order_id,
                            "store_code": target.store_code,
                            "order_date": target.order_date,
                            "sku_id": target.sku_id,
                            "sku_key": target.sku_key,
                            "target_quantity": target.quantity,
                            "reason": "no real item-entry evidence in approved source hierarchy",
                        },
                        ensure_ascii=False,
                        sort_keys=True,
                        default=_json_default,
                    )
                    + "\n"
                )
    with (output_root / "accepted_no_entry_quarantine_preview.jsonl").open("w", encoding="utf-8") as fh:
        for pair, accepted in sorted(accepted_no_entry_quarantine.items()):
            assignment = assignments.get(pair)
            fh.write(
                json.dumps(
                    {
                        "order_id": accepted.order_id,
                        "store_code": accepted.store_code,
                        "classification": accepted.classification,
                        "copied_temp_only": accepted.copied_temp_only,
                        "production_write_authorized": accepted.production_write_authorized,
                        "source_authority": accepted.source_authority,
                        "notes": accepted.notes,
                        "matched_unrecovered_quarantine": bool(
                            assignment and assignment.source_name == "UNRECOVERED_QUARANTINE"
                        ),
                        "target_rows": assignment.target_row_count if assignment else 0,
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                    default=_json_default,
                )
                + "\n"
            )


def _load_default_source_bundles(
    *,
    current_crm: Path,
    api_entry_roots: list[Path],
    webui_archive_csvs: list[Path],
    reserve_archive: Path,
) -> tuple[list[SourceBundle], dict[str, Any]]:
    current = read_workbook_evidence(
        current_crm,
        CURRENT_CRM_SHEET,
        source_name="CURRENT_CRM",
        source_kind="current_crm_workbook",
    )
    api_maps = [read_api_entry_evidence(root) for root in api_entry_roots]
    api = merge_source_maps(*api_maps)
    webui = read_webui_csv_evidence(webui_archive_csvs)
    reserve = read_workbook_evidence(
        reserve_archive,
        RESERVE_ARCHIVE_SHEET,
        source_name="RESERVE_ARCHIVE_WORKBOOK",
        source_kind="reserve_archive_workbook",
    )
    bundles = [
        SourceBundle("CURRENT_CRM", "HIGH", current),
        SourceBundle("API_RAW_ORDER_ENTRIES", "VERY_HIGH", api),
        SourceBundle("WEBUI_ARCHIVE_SOURCE_BACKUP", "HIGH", webui),
        SourceBundle("RESERVE_ARCHIVE_WORKBOOK", "MEDIUM", reserve),
    ]
    source_info = {
        "current_crm": str(current_crm),
        "api_entry_roots": [str(path) for path in api_entry_roots],
        "webui_archive_csvs": [str(path) for path in webui_archive_csvs],
        "reserve_archive": str(reserve_archive),
    }
    return bundles, source_info


def _load_api_only_source_bundles(api_entry_roots: list[Path]) -> tuple[list[SourceBundle], dict[str, Any]]:
    api_maps = [read_api_entry_evidence(root) for root in api_entry_roots]
    api = merge_source_maps(*api_maps)
    return [SourceBundle("API_RAW_ORDER_ENTRIES", "VERY_HIGH", api)], {
        "api_entry_roots": [str(path) for path in api_entry_roots],
        "workbook_sources_used": False,
        "webui_archive_sources_used": False,
        "source_mode": "api_raw_order_entries_only",
    }


def recover_order_entries(
    *,
    db_path: Path,
    as_of: str,
    output_root: Path,
    recovery_ts: str | None = None,
    source_bundles: list[SourceBundle] | None = None,
    current_crm: Path = DEFAULT_CURRENT_CRM,
    api_entry_roots: list[Path] | None = None,
    webui_archive_csvs: list[Path] | None = None,
    reserve_archive: Path = DEFAULT_RESERVE_ARCHIVE,
    target_source: str = "sales_fact_v2",
    start_date: str | None = None,
    stores: tuple[str, ...] = DEFAULT_FACT_ORDER_TARGET_STORES,
    entry_required_only: bool = False,
    target_order_csv: Path | None = None,
    accepted_no_entry_quarantine_csv: Path | None = None,
    expected_pre_sha256: str | None = None,
    backup_dir: Path | None = None,
    apply: bool = False,
    strict: bool = False,
) -> dict[str, Any]:
    if recovery_ts is None:
        recovery_ts = dt.datetime.now(dt.timezone.utc).astimezone().isoformat(timespec="seconds")
    elif as_of and parse_date(recovery_ts) > parse_date(as_of):
        raise RecoveryError(f"recovery_ts {recovery_ts} is after as_of {as_of}")
    normalized_target_source = target_source.strip().lower()
    if normalized_target_source == "sales_fact_v2":
        targets = build_missing_targets(db_path, as_of=as_of)
    elif normalized_target_source == "fact_orders_kaspi":
        targets = build_missing_fact_order_targets(
            db_path,
            start_date=start_date or "",
            as_of=as_of,
            stores=tuple(normalize_store(store) for store in stores),
            entry_required_only=entry_required_only,
        )
    else:
        raise RecoveryError(f"unsupported target_source: {target_source}")
    requested_target_pairs = load_target_order_pairs(target_order_csv)
    targets, target_filter_summary = filter_targets_by_order_pairs(
        targets,
        requested_pairs=requested_target_pairs,
    )
    accepted_no_entry_quarantine = load_accepted_no_entry_quarantine(accepted_no_entry_quarantine_csv)
    if accepted_no_entry_quarantine:
        try:
            accepted_db_is_production = db_path.resolve() == DEFAULT_DB.resolve()
        except FileNotFoundError:
            accepted_db_is_production = db_path.absolute() == DEFAULT_DB.absolute()
        if accepted_db_is_production:
            raise RecoveryError(
                "Accepted no-entry quarantine contract is copied-temp only and cannot target production db/app.db"
            )
    article_map = load_article_map(db_path)
    source_info: dict[str, Any] = {}
    if source_bundles is None:
        if normalized_target_source == "fact_orders_kaspi":
            source_bundles, source_info = _load_api_only_source_bundles(
                list(api_entry_roots or DEFAULT_API_ENTRY_ROOTS)
            )
        else:
            source_bundles, source_info = _load_default_source_bundles(
                current_crm=current_crm,
                api_entry_roots=list(api_entry_roots or DEFAULT_API_ENTRY_ROOTS),
                webui_archive_csvs=list(webui_archive_csvs or DEFAULT_WEBUI_ARCHIVE_CSVS),
                reserve_archive=reserve_archive,
            )
    assignments = assign_evidence_to_targets(targets, source_bundles, article_map=article_map)
    candidates = _candidate_entries(assignments, article_map, recovery_ts=recovery_ts)
    source_summary = _assignment_summary(assignments)
    quarantine = source_summary.get(
        "UNRECOVERED_QUARANTINE",
        {
            "order_store_pairs": 0,
            "target_rows": 0,
            "evidence_rows": 0,
            "article_present_rows": 0,
            "dim_article_mapped_rows": 0,
            "direct_sku_rows": 0,
            "by_store": {},
        },
    )
    accepted_quarantine_summary = _accepted_no_entry_quarantine_summary(
        assignments,
        accepted_no_entry_quarantine,
    )
    if accepted_no_entry_quarantine and accepted_quarantine_summary["unexpected_pairs"]:
        raise RecoveryError(
            "Accepted no-entry quarantine CSV contains pairs that are not unrecovered target rows: "
            f"{accepted_quarantine_summary['unexpected_pairs']}"
        )
    accepted_strict_pass = bool(accepted_quarantine_summary["all_unrecovered_targets_accepted"])
    strict_passed = quarantine["target_rows"] == 0 or accepted_strict_pass

    backup_path: Path | None = None
    backup_sha256 = ""
    backup_integrity = None
    pre_sha256 = _file_sha256(db_path)
    integrity_before = _sqlite_integrity_check(db_path)
    if apply:
        production_apply = _guard_apply_path(
            db_path,
            expected_pre_sha256=expected_pre_sha256,
            backup_dir=backup_dir,
        )
        if integrity_before.lower() != "ok":
            raise RecoveryError(f"pre-write integrity_check failed: {integrity_before}")
        if expected_pre_sha256 and pre_sha256 != expected_pre_sha256:
            raise RecoveryError(
                f"pre-write SHA mismatch: expected {expected_pre_sha256}, got {pre_sha256}"
            )
        if production_apply and backup_dir is not None:
            backup_path = backup_database(db_path, backup_dir, compress=False)
            backup_sha256 = _file_sha256(backup_path)
            backup_integrity = _sqlite_integrity_check(backup_path)
            if backup_integrity.lower() != "ok":
                raise RecoveryError(f"backup integrity_check failed: {backup_integrity}")
        conn_cm = sqlite3.connect(db_path)
    else:
        production_apply = False
        conn_cm = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    with conn_cm as conn:
        conn.row_factory = sqlite3.Row
        existing = _existing_entry_ids(conn, [entry["entry_id"] for entry in candidates])
        logical_collisions = _legacy_logical_entry_collisions(conn, candidates, existing)
        if logical_collisions:
            raise RecoveryError(
                "refusing recovered-entry insert over existing logical public-offer identities: "
                f"{logical_collisions[:5]}"
            )
        new_candidates = [entry for entry in candidates if entry["entry_id"] not in existing]
        inserted = 0
        if apply:
            inserted = _insert_entries(conn, new_candidates)
            conn.commit()

    post_sha256 = _file_sha256(db_path)
    integrity_after = _sqlite_integrity_check(db_path)
    mappable = sum(1 for entry in candidates if entry.get("sku_rebuild_mappable"))
    by_entry_source = Counter(str(entry.get("recovery_source_name")) for entry in candidates)
    production_db_target = db_path.resolve() == DEFAULT_DB.resolve()
    summary: dict[str, Any] = {
        "as_of": as_of,
        "recovery_ts": recovery_ts,
        "db_path": str(db_path),
        "output_root": str(output_root),
        "dry_run": not apply,
        "target_source": normalized_target_source,
        "target_start_date": start_date,
        "target_stores": list(stores),
        "entry_required_only": bool(entry_required_only),
        "target_filter": {
            **target_filter_summary,
            "target_order_csv": str(target_order_csv) if target_order_csv else None,
            "target_order_csv_sha256": _file_sha256(target_order_csv) if target_order_csv else "",
        },
        "target": _target_summary(targets),
        "source_hierarchy": [bundle.source_name for bundle in source_bundles],
        "sources": source_info,
        "recovery_by_source": source_summary,
        "quarantine": {
            "order_store_pairs": quarantine["order_store_pairs"],
            "target_rows": quarantine["target_rows"],
            "preview_path": str(output_root / "quarantine_preview.jsonl"),
        },
        "accepted_no_entry_quarantine": {
            **accepted_quarantine_summary,
            "contract_path": str(accepted_no_entry_quarantine_csv) if accepted_no_entry_quarantine_csv else None,
            "preview_path": str(output_root / "accepted_no_entry_quarantine_preview.jsonl"),
            "contract_sha256": _file_sha256(accepted_no_entry_quarantine_csv)
            if accepted_no_entry_quarantine_csv
            else "",
        },
        "entry_candidates": {
            "candidate_entry_rows": len(candidates),
            "new_candidate_entry_rows": len(new_candidates),
            "existing_entry_id_rows": len(existing),
            "api_entry_rows": by_entry_source.get("API_RAW_ORDER_ENTRIES", 0),
            "non_api_entry_rows": len(candidates) - by_entry_source.get("API_RAW_ORDER_ENTRIES", 0),
            "sku_rebuild_mappable_rows": mappable,
            "sku_rebuild_blocked_rows": len(candidates) - mappable,
            "by_recovery_source": dict(sorted(by_entry_source.items())),
        },
        "apply": {
            "applied": apply,
            "would_insert_entry_rows": len(new_candidates),
            "inserted_entry_rows": inserted,
            "skipped_existing_entry_rows": len(candidates) - len(new_candidates),
            "production_apply": production_apply,
        },
        "strict": {
            "requested": strict,
            "passed": strict_passed,
            "passed_by_accepted_no_entry_quarantine": bool(
                quarantine["target_rows"] and accepted_strict_pass
            ),
        },
        "production_db_modified": bool(apply and inserted > 0) if production_db_target else None,
        "pre_sha256": pre_sha256,
        "post_sha256": post_sha256,
        "expected_pre_sha256": expected_pre_sha256,
        "backup_path": str(backup_path) if backup_path else None,
        "backup_sha256": backup_sha256,
        "integrity_check": {
            "before": integrity_before,
            "backup": backup_integrity,
            "after": integrity_after,
        },
    }
    _write_preview_outputs(
        output_root,
        summary=summary,
        entries=candidates,
        assignments=assignments,
        accepted_no_entry_quarantine=accepted_no_entry_quarantine,
    )
    if strict and not strict_passed:
        raise RecoveryError(
            "Strict recovery failed: "
            f"unrecovered target rows={quarantine['target_rows']}, "
            f"accepted_no_entry_missing_pairs={accepted_quarantine_summary['missing_pairs']}"
        )
    return summary


def _parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    parser.add_argument("--as-of", required=True)
    parser.add_argument("--start-date", help="Start date for fact_orders_kaspi target mode.")
    parser.add_argument(
        "--target-source",
        choices=("sales_fact_v2", "fact_orders_kaspi"),
        default="sales_fact_v2",
        help="Target rows to recover. fact_orders_kaspi is for May 5-current header rows.",
    )
    parser.add_argument(
        "--stores",
        default=",".join(DEFAULT_FACT_ORDER_TARGET_STORES),
        help="Comma-separated stores for fact_orders_kaspi target mode.",
    )
    parser.add_argument(
        "--entry-required-only",
        action="store_true",
        help="In fact_orders_kaspi mode, target only rows required by the freshness validator.",
    )
    parser.add_argument(
        "--target-order-csv",
        type=Path,
        help=(
            "Optional narrow order/store CSV filter. The CSV must contain order_id and store_code, "
            "and every pair must already be present in the current missing target set."
        ),
    )
    parser.add_argument("--output-root", type=Path, required=True)
    parser.add_argument("--current-crm", type=Path, default=DEFAULT_CURRENT_CRM)
    parser.add_argument("--api-entry-root", type=Path, action="append", default=None)
    parser.add_argument("--webui-archive-csv", type=Path, action="append", default=None)
    parser.add_argument("--reserve-archive", type=Path, default=DEFAULT_RESERVE_ARCHIVE)
    parser.add_argument(
        "--accepted-no-entry-quarantine-csv",
        type=Path,
        help=(
            "Copied-temp-only CSV contract that explicitly retains unrecovered target rows "
            "without creating synthetic item entries."
        ),
    )
    parser.add_argument(
        "--recovery-ts",
        help="Timestamp to store in fact_order_entries_kaspi.updated_at; defaults to execution time.",
    )
    parser.add_argument("--expected-pre-sha256")
    parser.add_argument("--backup-dir", type=Path)
    parser.add_argument("--strict", action="store_true")
    parser.add_argument("--apply", action="store_true")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(argv or sys.argv[1:])
    stores = tuple(normalize_store(store) for store in str(args.stores).split(",") if store.strip())
    try:
        summary = recover_order_entries(
            db_path=args.db,
            as_of=args.as_of,
            output_root=args.output_root,
            target_source=args.target_source,
            start_date=args.start_date,
            stores=stores,
            entry_required_only=bool(args.entry_required_only),
            target_order_csv=args.target_order_csv,
            current_crm=args.current_crm,
            api_entry_roots=args.api_entry_root,
            webui_archive_csvs=args.webui_archive_csv,
            reserve_archive=args.reserve_archive,
            accepted_no_entry_quarantine_csv=args.accepted_no_entry_quarantine_csv,
            expected_pre_sha256=args.expected_pre_sha256,
            backup_dir=args.backup_dir,
            recovery_ts=args.recovery_ts,
            apply=args.apply,
            strict=args.strict,
        )
    except RecoveryError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    print(f"summary_json={args.output_root / 'summary.json'}")
    print(f"target_validator_rows={summary['target']['validator_rows']}")
    print(f"target_order_store_pairs={summary['target']['order_store_pairs']}")
    print(f"candidate_entry_rows={summary['entry_candidates']['candidate_entry_rows']}")
    print(f"would_insert_entry_rows={summary['apply']['would_insert_entry_rows']}")
    print(f"inserted_entry_rows={summary['apply']['inserted_entry_rows']}")
    print(f"quarantine_target_rows={summary['quarantine']['target_rows']}")
    print(f"sku_rebuild_blocked_rows={summary['entry_candidates']['sku_rebuild_blocked_rows']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
