#!/usr/bin/env python3
"""LaunchAgent entrypoint for Google Ops Board publish schedule."""

from __future__ import annotations

import argparse
import os
import hashlib
import json
import sqlite3
import subprocess
import sys
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

SCRIPT_PATH = PROJECT_ROOT / "scripts" / "sync_google_ops_board.py"
DB_CHECK_PATH = PROJECT_ROOT / "scripts" / "check_local_app_db.py"
EXPORT_API_ORDERS_PATH = PROJECT_ROOT / "scripts" / "export_api_orders.py"
SYNC_KASPI_ORDERS_PATH = PROJECT_ROOT / "scripts" / "sync_kaspi_orders.py"
APPLY_ORDER_ENTRY_SIDECAR_PATH = PROJECT_ROOT / "scripts" / "apply_kaspi_order_entry_sidecar.py"
ENRICH_ACTIVEORDERS_PATH = PROJECT_ROOT / "scripts" / "enrich_kaspi_orders_from_activeorders.py"
VALIDATE_ACTIVEORDERS_COLUMNS_PATH = PROJECT_ROOT / "scripts" / "validate_activeorders_columns.py"
ACTIVEORDERS_PATH = PROJECT_ROOT / "excel_ui" / "ActiveOrders" / "ActiveOrders.xlsx"
DEFAULT_SOURCE_SNAPSHOT_ROOT = PROJECT_ROOT / "exports" / "google_ops_board" / "source_snapshots"
ACTIVEORDERS_PLANNED_DATE_HEADER = "Плановая дата передачи курьеру"
DEFAULT_LOOKBACK_DAYS = 5
MORNING_SOURCE_REFRESH_HOUR = 7
MORNING_SOURCE_REFRESH_MINUTE = 0

from core.integrations.google_ops_board import (  # noqa: E402
    DEFAULT_CONTRACT_PATH,
    load_ops_board_contract,
    resolve_service_account_json,
    resolve_spreadsheet_id,
)
from core.stores.roster import load_sync_enabled_kaspi_store_codes  # noqa: E402
from scripts.google_ops_board_automation_common import (  # noqa: E402
    AUTOMATION_LOCK_HELD_ENV,
    GoogleOpsBoardAutomationLock,
    ensure_kaspi_api_call_ledger_env,
    now_almaty,
    record_lock_contention,
    reset_lock_contention,
    run_guarded,
    today_almaty,
)
from scripts.backup_db import backup_database, verify_backup  # noqa: E402
from scripts.run_google_ops_board_prewindow_health import ensure_prewindow_health  # noqa: E402

IDENTITY_SYNC_WRITE_ENV_GATE = "ENABLE_KASPI_WORKBOOK_MAP_SYNC"
ACTIVEORDERS_DB_WRITE_ENV_GATE = "ENABLE_KASPI_ACTIVEORDERS_DB_WRITE"
CURRENT_ORDER_ENTRY_SYNC_ENV_GATE = "ENABLE_KASPI_CURRENT_ORDER_ENTRY_SYNC"
KASPI_ENRICHMENT_WRITE_ENV_GATE = "ENABLE_KASPI_ENRICHMENT"
CURRENT_ORDER_STATUS_EVENT_SYNC_ENV_GATE = "ENABLE_KASPI_CURRENT_ORDER_STATUS_EVENT_SYNC"
ORDER_STATUS_EVENT_WRITE_ENV_GATE = "ENABLE_ORDER_STATUS_EVENT_WRITE"
LOCK_CONTENTION_EXIT_CODE = 75


def is_source_refresh_slot(now: datetime | None = None) -> bool:
    local_now = now or now_almaty()
    return local_now.hour == MORNING_SOURCE_REFRESH_HOUR and local_now.minute == MORNING_SOURCE_REFRESH_MINUTE


def inspect_activeorders_source(workbook_path: Path, *, target_date: date) -> dict[str, Any]:
    report: dict[str, Any] = {
        "path": str(workbook_path),
        "target_date": target_date.isoformat(),
        "target_label": target_date.strftime("%d.%m.%Y"),
        "exists": workbook_path.exists(),
        "mtime_date": "",
        "row_count": 0,
        "target_row_count": 0,
        "contains_target_date": False,
        "fresh": False,
        "planned_date_counts": {},
    }
    if not workbook_path.exists():
        return report

    stat = workbook_path.stat()
    report["mtime_date"] = datetime.fromtimestamp(stat.st_mtime, tz=now_almaty().tzinfo).date().isoformat()
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        worksheet = workbook.active
        headers = [str(cell) if cell is not None else "" for cell in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        if ACTIVEORDERS_PLANNED_DATE_HEADER not in headers:
            report["error"] = f"Missing required header: {ACTIVEORDERS_PLANNED_DATE_HEADER}"
            return report
        planned_date_idx = headers.index(ACTIVEORDERS_PLANNED_DATE_HEADER)
        target_label = report["target_label"]
        counts: Counter[str] = Counter()
        total_rows = 0
        target_rows = 0
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            total_rows += 1
            planned_date = str(row[planned_date_idx] or "").strip()
            if not planned_date:
                continue
            counts[planned_date] += 1
            if planned_date == target_label:
                target_rows += 1
        report["row_count"] = total_rows
        report["target_row_count"] = target_rows
        report["contains_target_date"] = target_rows > 0
        report["planned_date_counts"] = dict(counts)
        # A successful current-day refresh may legitimately contain zero new
        # target-date orders. Freshness is the source observation date, not a
        # business-volume assertion; carryover obligations are added downstream.
        report["fresh"] = report["mtime_date"] == target_date.isoformat()
        return report
    except Exception as exc:
        report["error"] = str(exc)
        return report


def source_snapshot_path(target_date: date, *, snapshot_root: Path = DEFAULT_SOURCE_SNAPSHOT_ROOT) -> Path:
    return Path(snapshot_root) / target_date.isoformat() / "source_snapshot.json"


def order_entry_sidecar_path(
    target_date: date,
    *,
    snapshot_root: Path = DEFAULT_SOURCE_SNAPSHOT_ROOT,
) -> Path:
    return Path(snapshot_root) / target_date.isoformat() / "order_entries_sidecar.json"


def order_entry_sidecar_report_path(
    target_date: date,
    *,
    snapshot_root: Path = DEFAULT_SOURCE_SNAPSHOT_ROOT,
) -> Path:
    return Path(snapshot_root) / target_date.isoformat() / "order_entries_sidecar_apply_report.json"


def build_file_fingerprint(path: Path) -> dict[str, Any]:
    target = Path(path).expanduser().resolve()
    if not target.exists():
        return {
            "path": str(target),
            "exists": False,
            "size": 0,
            "mtime_ns": 0,
            "sha256": "",
        }
    stat = target.stat()
    digest = hashlib.sha256()
    with target.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return {
        "path": str(target),
        "exists": True,
        "size": int(stat.st_size),
        "mtime_ns": int(stat.st_mtime_ns),
        "sha256": digest.hexdigest(),
    }


def write_source_snapshot(
    *,
    target_date: date,
    workbook_path: Path = ACTIVEORDERS_PATH,
    snapshot_root: Path = DEFAULT_SOURCE_SNAPSHOT_ROOT,
    refresh_slot: bool = False,
    include_entry_sidecar: bool = False,
    prewrite_backup_path: Path | None = None,
) -> dict[str, Any]:
    source_state = inspect_activeorders_source(workbook_path, target_date=target_date)
    if refresh_slot and source_state.get("contains_target_date"):
        source_state = dict(source_state)
        source_state["fresh"] = True
        source_state["freshness_basis"] = "refresh_slot_contains_target_date"
    path = source_snapshot_path(target_date, snapshot_root=snapshot_root)
    payload = {
        "target_date": target_date.isoformat(),
        "generated_at": now_almaty().isoformat(),
        "refresh_slot": bool(refresh_slot),
        "source_state": source_state,
        "source_fingerprint": build_file_fingerprint(workbook_path),
        "active_store_roster": sorted(load_sync_enabled_kaspi_store_codes()),
        "path": str(path),
    }
    if include_entry_sidecar:
        sidecar = order_entry_sidecar_path(target_date, snapshot_root=snapshot_root)
        receipt = order_entry_sidecar_report_path(target_date, snapshot_root=snapshot_root)
        if not sidecar.exists() or not receipt.exists():
            raise RuntimeError(
                "entry sidecar commit marker is incomplete: "
                f"sidecar_exists={sidecar.exists()} receipt_exists={receipt.exists()}"
            )
        try:
            receipt_payload = json.loads(receipt.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise RuntimeError(f"entry sidecar receipt is invalid: {receipt}") from exc
        if not receipt_payload.get("readback_complete"):
            raise RuntimeError("entry sidecar receipt does not prove complete readback")
        payload["order_entry_sidecar_fingerprint"] = build_file_fingerprint(sidecar)
        payload["order_entry_sidecar_receipt_fingerprint"] = build_file_fingerprint(receipt)
        payload["order_entry_sidecar_receipt"] = {
            key: receipt_payload.get(key)
            for key in (
                "payload_sha256",
                "store_count",
                "order_count",
                "entry_count",
                "preexisting_entry_count",
                "inserted_entry_count",
                "readback_complete",
                "integrity_after",
                "non_target_hash_after",
                "backup_path",
            )
        }
    if prewrite_backup_path is not None:
        if not prewrite_backup_path.exists():
            raise RuntimeError(f"prewrite DB backup is missing: {prewrite_backup_path}")
        payload["prewrite_db_backup_fingerprint"] = build_file_fingerprint(prewrite_backup_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return payload


def load_committed_source_snapshot(
    *,
    target_date: date,
    workbook_path: Path = ACTIVEORDERS_PATH,
    snapshot_root: Path = DEFAULT_SOURCE_SNAPSHOT_ROOT,
) -> dict[str, Any]:
    """Read back the refresh commit marker without weakening its evidence.

    A refresh may pin an order-entry sidecar, its apply receipt, and the
    pre-write database backup.  The publish cycle must preserve that stronger
    marker instead of replacing it with a generic workbook-only snapshot.
    """

    path = source_snapshot_path(target_date, snapshot_root=snapshot_root)
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise RuntimeError(f"committed source snapshot is missing or invalid: {path}") from exc
    if payload.get("target_date") != target_date.isoformat():
        raise RuntimeError("committed source snapshot target date does not match publish date")
    if payload.get("refresh_slot") is not True:
        raise RuntimeError("committed source snapshot is not a strong refresh commit marker")
    expected_fingerprint = build_file_fingerprint(workbook_path)
    actual_fingerprint = payload.get("source_fingerprint") or {}
    if not expected_fingerprint.get("exists") or (
        actual_fingerprint.get("sha256") != expected_fingerprint.get("sha256")
    ):
        raise RuntimeError("committed source snapshot no longer matches ActiveOrders")
    backup_fingerprint = payload.get("prewrite_db_backup_fingerprint") or {}
    backup_path_text = str(backup_fingerprint.get("path") or "").strip()
    if not backup_path_text:
        raise RuntimeError("committed source snapshot does not pin its prewrite DB backup")
    observed_backup = build_file_fingerprint(Path(backup_path_text))
    if not observed_backup.get("exists") or (
        observed_backup.get("sha256") != backup_fingerprint.get("sha256")
    ):
        raise RuntimeError("committed source snapshot prewrite DB backup fingerprint is invalid")
    sidecar_keys = {
        "order_entry_sidecar_fingerprint",
        "order_entry_sidecar_receipt_fingerprint",
        "order_entry_sidecar_receipt",
    }
    present_sidecar_keys = sidecar_keys & set(payload)
    if present_sidecar_keys and present_sidecar_keys != sidecar_keys:
        raise RuntimeError("committed source snapshot has partial order-entry sidecar evidence")
    return payload


def build_source_refresh_commands(
    *,
    target_date: date,
    workbook_path: Path,
    python_executable: str,
    lookback_days: int = DEFAULT_LOOKBACK_DAYS,
    enable_entry_sync: bool = False,
) -> list[list[str]]:
    since_date = (target_date - timedelta(days=lookback_days)).isoformat()
    sync_command = [
        python_executable,
        str(SYNC_KASPI_ORDERS_PATH),
        "--all",
        "--since",
        since_date,
        "-v",
    ]
    export_command = [
            python_executable,
            str(EXPORT_API_ORDERS_PATH),
            "--all-stores",
            "--require-complete",
            "--state",
            "KASPI_DELIVERY",
            "--days",
            str(lookback_days),
            "--include-overdue",
            "--no-archive",
            "--output",
            str(workbook_path),
            "--verbose",
        ]
    if enable_entry_sync:
        export_command.extend(
            [
                "--entry-sidecar-output",
                str(order_entry_sidecar_path(target_date)),
            ]
        )
    return [
        export_command,
        [
            python_executable,
            str(VALIDATE_ACTIVEORDERS_COLUMNS_PATH),
            str(workbook_path),
        ],
        sync_command,
        [
            python_executable,
            str(ENRICH_ACTIVEORDERS_PATH),
            "--apply",
            "--file",
            str(workbook_path),
            "--target-date",
            target_date.isoformat(),
        ],
    ]


def run_pinned_order_entry_sidecar_apply(
    *,
    target_date: date,
    env: dict[str, str],
    sidecar_path: Path | None = None,
    report_path: Path | None = None,
    prewrite_backup_path: Path | None = None,
) -> int:
    from scripts.apply_kaspi_order_entry_sidecar import (
        OrderEntrySidecarError,
        load_validated_sidecar,
    )

    sidecar = sidecar_path or order_entry_sidecar_path(target_date)
    report = report_path or order_entry_sidecar_report_path(target_date)
    try:
        envelope = json.loads(sidecar.read_text(encoding="utf-8"))
        pinned_hash = str(envelope.get("payload_sha256") or "").strip().lower()
        load_validated_sidecar(
            sidecar,
            expected_payload_sha256=pinned_hash,
            expected_target_date=target_date.isoformat(),
        )
    except (OSError, json.JSONDecodeError, AttributeError, OrderEntrySidecarError) as exc:
        print(f"ERROR: cannot pin exported order-entry sidecar: {exc}", file=sys.stderr)
        return 78
    command = [
        sys.executable,
        str(APPLY_ORDER_ENTRY_SIDECAR_PATH),
        "--sidecar",
        str(sidecar),
        "--expected-payload-sha256",
        pinned_hash,
        "--target-date",
        target_date.isoformat(),
        "--db-path",
        str(PROJECT_ROOT / "db" / "app.db"),
        "--report",
        str(report),
        "--apply",
    ]
    if prewrite_backup_path is not None:
        command.extend(["--prewrite-backup", str(prewrite_backup_path)])
    result = subprocess.run(command, cwd=str(PROJECT_ROOT), env=env)
    return int(result.returncode)


def create_current_order_refresh_backup(*, target_date: date) -> Path:
    db_path = PROJECT_ROOT / "db" / "app.db"
    backup_root = PROJECT_ROOT / "backups" / "kaspi_current_order_refresh" / target_date.isoformat()
    backup_path = backup_database(db_path, backup_root, compress=False)
    if not verify_backup(backup_path):
        raise RuntimeError(f"current-order refresh DB backup verification failed: {backup_path}")
    return backup_path


def run_source_refresh(*, target_date: date, env: dict[str, str]) -> int:
    if not EXPORT_API_ORDERS_PATH.exists():
        print(f"ERROR: missing ActiveOrders exporter: {EXPORT_API_ORDERS_PATH}", file=sys.stderr)
        return 78
    if not SYNC_KASPI_ORDERS_PATH.exists():
        print(f"ERROR: missing Kaspi order sync script: {SYNC_KASPI_ORDERS_PATH}", file=sys.stderr)
        return 78
    if not ENRICH_ACTIVEORDERS_PATH.exists():
        print(f"ERROR: missing ActiveOrders enrichment script: {ENRICH_ACTIVEORDERS_PATH}", file=sys.stderr)
        return 78
    if not VALIDATE_ACTIVEORDERS_COLUMNS_PATH.exists():
        print(f"ERROR: missing ActiveOrders validator: {VALIDATE_ACTIVEORDERS_COLUMNS_PATH}", file=sys.stderr)
        return 78

    entry_sync_enabled = str(env.get(CURRENT_ORDER_ENTRY_SYNC_ENV_GATE) or "").strip() == "1"
    if entry_sync_enabled and not APPLY_ORDER_ENTRY_SIDECAR_PATH.exists():
        print(
            f"ERROR: missing order-entry sidecar writer: {APPLY_ORDER_ENTRY_SIDECAR_PATH}",
            file=sys.stderr,
        )
        return 78

    lookback_days = int(str(env.get("KASPI_LOOKBACK_DAYS") or DEFAULT_LOOKBACK_DAYS).strip() or DEFAULT_LOOKBACK_DAYS)
    commands = build_source_refresh_commands(
        target_date=target_date,
        workbook_path=ACTIVEORDERS_PATH,
        python_executable=sys.executable,
        lookback_days=lookback_days,
        enable_entry_sync=entry_sync_enabled,
    )
    refresh_env = env.copy()
    refresh_env.setdefault(ACTIVEORDERS_DB_WRITE_ENV_GATE, "1")
    os.environ.setdefault(ACTIVEORDERS_DB_WRITE_ENV_GATE, refresh_env[ACTIVEORDERS_DB_WRITE_ENV_GATE])
    if entry_sync_enabled:
        refresh_env[KASPI_ENRICHMENT_WRITE_ENV_GATE] = "1"
    if str(env.get(CURRENT_ORDER_STATUS_EVENT_SYNC_ENV_GATE) or "").strip() == "1":
        refresh_env[ORDER_STATUS_EVENT_WRITE_ENV_GATE] = "1"
    else:
        refresh_env.pop(ORDER_STATUS_EVENT_WRITE_ENV_GATE, None)

    labels = (
        "Refreshing ActiveOrders export from live Kaspi API...",
        "Validating refreshed ActiveOrders workbook...",
        "Syncing DB order headers from Kaspi API...",
        "Enriching DB order identities from refreshed ActiveOrders workbook...",
    )
    prewrite_backup_path: Path | None = None
    for index, (label, command) in enumerate(zip(labels, commands)):
        if index == 2:
            print("Backing up DB before current-order refresh writes...")
            try:
                prewrite_backup_path = create_current_order_refresh_backup(target_date=target_date)
            except (OSError, sqlite3.Error, RuntimeError) as exc:
                print(f"ERROR: current-order refresh backup failed: {exc}", file=sys.stderr)
                return 78
            refresh_env["AB_CURRENT_ORDER_REFRESH_BACKUP_PATH"] = str(prewrite_backup_path)
        print(label)
        result = subprocess.run(command, cwd=str(PROJECT_ROOT), env=refresh_env)
        if result.returncode != 0:
            print(f"ERROR: source refresh step failed: {label}", file=sys.stderr)
            return int(result.returncode)
        if index == 2 and entry_sync_enabled:
            print("Applying hash-pinned order-entry sidecar without duplicate API fetches...")
            sidecar_result = run_pinned_order_entry_sidecar_apply(
                target_date=target_date,
                env=refresh_env,
                prewrite_backup_path=prewrite_backup_path,
            )
            if sidecar_result != 0:
                print("ERROR: source refresh step failed: order-entry sidecar apply", file=sys.stderr)
                return sidecar_result
    try:
        snapshot = write_source_snapshot(
            target_date=target_date,
            workbook_path=ACTIVEORDERS_PATH,
            refresh_slot=True,
            include_entry_sidecar=entry_sync_enabled,
            prewrite_backup_path=prewrite_backup_path,
        )
    except (OSError, RuntimeError, ValueError, json.JSONDecodeError) as exc:
        print(f"ERROR: source snapshot commit marker failed: {exc}", file=sys.stderr)
        return 78
    print(f"Source snapshot: {snapshot['path']}")
    return 0


def run_publish_cycle(
    *,
    env: dict[str, str],
    service_account_json: str,
    spreadsheet_id: str,
    force_source_refresh: bool = False,
) -> int:
    target_date = today_almaty()
    ensure_kaspi_api_call_ledger_env(env, target_date=target_date, project_root=PROJECT_ROOT)
    if env.get("KASPI_API_CALL_LEDGER_PATH"):
        os.environ.setdefault("KASPI_API_CALL_LEDGER_PATH", env["KASPI_API_CALL_LEDGER_PATH"])

    check_cmd = [
        sys.executable,
        str(DB_CHECK_PATH),
        "--db-path",
        str(PROJECT_ROOT / "db" / "app.db"),
    ]
    check = subprocess.run(check_cmd, cwd=str(PROJECT_ROOT), env=env)
    if check.returncode != 0:
        print("ERROR: local DB preflight failed; skipping Google Ops Board publish.", file=sys.stderr)
        return int(check.returncode)

    refresh_requested = bool(force_source_refresh or is_source_refresh_slot())
    if refresh_requested:
        reason = "forced shipping-source refresh" if force_source_refresh else "07:00 publish slot"
        print(f"{reason} detected; running full ActiveOrders -> DB source refresh before publish.")
        refresh_rc = run_source_refresh(target_date=target_date, env=env)
        if refresh_rc != 0:
            return int(refresh_rc)

    source_state = inspect_activeorders_source(ACTIVEORDERS_PATH, target_date=target_date)
    try:
        snapshot = load_committed_source_snapshot(
            target_date=target_date,
            workbook_path=ACTIVEORDERS_PATH,
        )
    except RuntimeError as exc:
        print(f"ERROR: source snapshot readback failed: {exc}", file=sys.stderr)
        return 78
    print(f"Source snapshot: {snapshot['path']}")
    if not source_state.get("fresh"):
        print(
            "ERROR: ActiveOrders source is stale for target date; skipping Google Ops Board publish. "
            f"Source: {source_state.get('path', '')} "
            f"mtime_date={source_state.get('mtime_date', '')} "
            f"target_row_count={source_state.get('target_row_count', 0)}",
            file=sys.stderr,
        )
        return 1

    health = ensure_prewindow_health(
        target_date=target_date,
        db_path=PROJECT_ROOT / "db" / "app.db",
        contract_path=DEFAULT_CONTRACT_PATH,
        service_account_json=Path(service_account_json),
        spreadsheet_id=spreadsheet_id,
        apply=True,
        reason="publish_scheduler",
        profile="publish",
    )
    if not health.get("ok"):
        print(
            f"ERROR: prewindow health gate is red; skipping Google Ops Board publish. "
            f"Report: {health.get('report_path', '')}",
            file=sys.stderr,
        )
        return 1

    cmd = [sys.executable, str(SCRIPT_PATH), "--apply"]
    cmd.extend(["--spreadsheet-id", spreadsheet_id])
    cmd.extend(["--service-account-json", service_account_json])
    result = subprocess.run(cmd, cwd=str(PROJECT_ROOT), env=env)
    if result.returncode != 0:
        return int(result.returncode)
    post_publish_health = ensure_prewindow_health(
        target_date=target_date,
        db_path=PROJECT_ROOT / "db" / "app.db",
        contract_path=DEFAULT_CONTRACT_PATH,
        service_account_json=Path(service_account_json),
        spreadsheet_id=spreadsheet_id,
        apply=True,
        reason="publish_scheduler_post_publish_readback",
        profile="publish",
        force=True,
        require_live_board_parity=True,
        emit_alerts=False,
    )
    if not post_publish_health.get("ok"):
        print(
            "ERROR: published Google Ops Board failed exact live readback; "
            f"report: {post_publish_health.get('report_path', '')}",
            file=sys.stderr,
        )
        return 1
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Publish the Google Ops Board from the canonical shipping source.")
    parser.add_argument(
        "--force-source-refresh",
        action="store_true",
        help="Run complete Kaspi source refresh and DB enrichment before publishing.",
    )
    args = parser.parse_args([] if argv is None else argv)

    if not SCRIPT_PATH.exists():
        print(f"ERROR: missing Google Ops Board publisher: {SCRIPT_PATH}", file=sys.stderr)
        return 78
    if not DB_CHECK_PATH.exists():
        print(f"ERROR: missing DB preflight script: {DB_CHECK_PATH}", file=sys.stderr)
        return 78

    env = os.environ.copy()
    env.setdefault("TERM", "dumb")
    env.setdefault("PYTHONUNBUFFERED", "1")
    env.setdefault(IDENTITY_SYNC_WRITE_ENV_GATE, "1")
    os.environ.setdefault(IDENTITY_SYNC_WRITE_ENV_GATE, env[IDENTITY_SYNC_WRITE_ENV_GATE])

    contract = load_ops_board_contract(DEFAULT_CONTRACT_PATH)
    try:
        service_account_path = resolve_service_account_json(contract=contract)
    except Exception:
        service_account_path = None
    service_account_json = str(service_account_path or "").strip()
    if not service_account_json or not Path(service_account_json).exists():
        print("ERROR: AB_GOOGLE_SERVICE_ACCOUNT_JSON is missing or does not exist.", file=sys.stderr)
        return 78
    spreadsheet_id = resolve_spreadsheet_id(
        str(env.get("AB_GOOGLE_OPS_BOARD_SPREADSHEET_ID") or "").strip() or None,
        contract=contract,
    )

    if str(os.environ.get(AUTOMATION_LOCK_HELD_ENV) or "").strip() == "1":
        env[AUTOMATION_LOCK_HELD_ENV] = "1"
        return run_publish_cycle(
            env=env,
            service_account_json=service_account_json,
            spreadsheet_id=spreadsheet_id,
            force_source_refresh=args.force_source_refresh,
        )

    previous_lock_env = os.environ.get(AUTOMATION_LOCK_HELD_ENV)
    lock_acquired = False
    try:
        with GoogleOpsBoardAutomationLock():
            lock_acquired = True
            try:
                reset_lock_contention("run_google_ops_board_publish_scheduler")
            except Exception as exc:
                print(f"WARNING: unable to reset lock-contention counter: {exc}", file=sys.stderr)
            os.environ[AUTOMATION_LOCK_HELD_ENV] = "1"
            env[AUTOMATION_LOCK_HELD_ENV] = "1"
            return run_publish_cycle(
                env=env,
                service_account_json=service_account_json,
                spreadsheet_id=spreadsheet_id,
                force_source_refresh=args.force_source_refresh,
            )
    except RuntimeError as exc:
        print(str(exc), file=sys.stderr)
        if not lock_acquired and not args.force_source_refresh:
            try:
                record_lock_contention("run_google_ops_board_publish_scheduler")
            except Exception as counter_exc:
                print(f"WARNING: unable to record lock contention: {counter_exc}", file=sys.stderr)
        # Quiet backstop publication is redundant while another canonical
        # Board owner holds the lock. A forced source refresh is not: its
        # caller must see temporary failure and retry or alert rather than
        # silently losing an import slot.
        return LOCK_CONTENTION_EXIT_CODE if args.force_source_refresh else 0
    finally:
        if previous_lock_env is None:
            os.environ.pop(AUTOMATION_LOCK_HELD_ENV, None)
        else:
            os.environ[AUTOMATION_LOCK_HELD_ENV] = previous_lock_env


if __name__ == "__main__":
    raise SystemExit(
        run_guarded(
            "run_google_ops_board_publish_scheduler",
            lambda: main(sys.argv[1:]),
        )
    )
