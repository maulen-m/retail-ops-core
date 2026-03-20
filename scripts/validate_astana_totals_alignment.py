#!/usr/bin/env python3
"""Validate 2.3.26 Astana totals sheet against authoritative workbook sheets."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import openpyxl

ASTANA_SHEET = "2.3.26_astana_totals"
TOTALS_SHEET = "PO_part_id_Totals"
DLV_SHEET = "dlv_payment_2.3.26"


def _to_float(value: Any) -> float:
    try:
        if value is None:
            return 0.0
        if isinstance(value, str) and not value.strip():
            return 0.0
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _is_valid_part_id(value: Any) -> bool:
    txt = str(value or "").strip().upper()
    if not txt:
        return False
    if any(token in txt for token in ("TOTAL", "PENDING", "UNPAID", "PAYMENT")):
        return False
    return txt.startswith("PO-") or txt.startswith("ARC-")


def _find_header_index(ws: openpyxl.worksheet.worksheet.Worksheet, names: list[str], max_cols: int = 80) -> dict[str, int]:
    header = {}
    for c in range(1, max_cols + 1):
        val = str(ws.cell(1, c).value or "").strip()
        if not val:
            continue
        for name in names:
            if val == name:
                header[name] = c
    return header


def _extract_totals_sheet(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, dict[str, float]]:
    cols = _find_header_index(
        ws,
        [
            "PO_part_id",
            "Total Units",
            "Total Bags",
            "Base_cost_CNY",
            "Actual_Weight_kg",
            "Paid_DLV_USD",
            "Paid_DLV_KZT",
        ],
    )
    required = {"PO_part_id", "Total Units", "Total Bags", "Base_cost_CNY", "Actual_Weight_kg", "Paid_DLV_USD", "Paid_DLV_KZT"}
    if not required.issubset(set(cols.keys())):
        missing = sorted(required - set(cols.keys()))
        raise RuntimeError(f"{TOTALS_SHEET} missing columns: {', '.join(missing)}")

    out: dict[str, dict[str, float]] = {}
    for r in range(2, ws.max_row + 1):
        part_id = str(ws.cell(r, cols["PO_part_id"]).value or "").strip()
        if not _is_valid_part_id(part_id):
            continue
        out[part_id] = {
            "units": _to_float(ws.cell(r, cols["Total Units"]).value),
            "bags": _to_float(ws.cell(r, cols["Total Bags"]).value),
            "cny": _to_float(ws.cell(r, cols["Base_cost_CNY"]).value),
            "actual_weight": _to_float(ws.cell(r, cols["Actual_Weight_kg"]).value),
            "dlv_usd": _to_float(ws.cell(r, cols["Paid_DLV_USD"]).value),
            "dlv_kzt": _to_float(ws.cell(r, cols["Paid_DLV_KZT"]).value),
        }
    return out


def _extract_astana_summary(ws: openpyxl.worksheet.worksheet.Worksheet) -> tuple[dict[str, dict[str, float]], dict[str, float]]:
    part_cols: dict[str, int] = {}
    for c in range(3, 40):
        part_id = str(ws.cell(3, c).value or "").strip()
        if _is_valid_part_id(part_id):
            part_cols[part_id] = c
    if not part_cols:
        raise RuntimeError(f"{ASTANA_SHEET}: failed to parse part columns from row 3")

    by_part: dict[str, dict[str, float]] = {}
    for part_id, col in sorted(part_cols.items()):
        by_part[part_id] = {
            "bags": _to_float(ws.cell(8, col).value),
            "units": _to_float(ws.cell(9, col).value),
            "actual_weight": _to_float(ws.cell(10, col).value),
            "dlv_usd": _to_float(ws.cell(11, col).value),
            "dlv_kzt": _to_float(ws.cell(12, col).value),
        }

    overall = {
        "bags": _to_float(ws.cell(8, 2).value),
        "units": _to_float(ws.cell(9, 2).value),
        "actual_weight": _to_float(ws.cell(10, 2).value),
        "dlv_usd": _to_float(ws.cell(11, 2).value),
        "dlv_kzt": _to_float(ws.cell(12, 2).value),
    }
    return by_part, overall


def _extract_astana_detail_cny(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, float]:
    out: dict[str, float] = {}
    for r in range(1, ws.max_row + 1):
        part_id = str(ws.cell(r, 1).value or "").strip()
        if not _is_valid_part_id(part_id):
            continue
        cny = _to_float(ws.cell(r, 6).value)
        if cny <= 0:
            continue
        out[part_id] = out.get(part_id, 0.0) + cny
    return out


def _extract_dlv_totals(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, float]:
    totals = {"actual_weight": 0.0, "dlv_usd": 0.0, "dlv_kzt": 0.0}
    for r in range(1, ws.max_row + 1):
        label = str(ws.cell(r, 1).value or "").strip()
        if label == "Total Actual Weight (kg)":
            totals["actual_weight"] = _to_float(ws.cell(r, 2).value)
        elif label == "TOTAL DELIVERY COST":
            totals["dlv_usd"] = _to_float(ws.cell(r, 2).value)
            totals["dlv_kzt"] = _to_float(ws.cell(r, 3).value)
    return totals


def _extract_per_bag_grand_weight(ws: openpyxl.worksheet.worksheet.Worksheet) -> float:
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip().upper() != "GRAND TOTAL":
            continue
        qty = _to_float(ws.cell(r, 5).value)
        weight = _to_float(ws.cell(r, 6).value)
        if qty > 0 and weight > 0:
            return weight
    return 0.0


def validate_astana_totals_alignment(
    *,
    workbook_path: Path,
    tol_qty: float = 0.0,
    tol_weight: float = 0.1,
    tol_money: float = 1.0,
) -> dict[str, Any]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    for required in (ASTANA_SHEET, TOTALS_SHEET, DLV_SHEET):
        if required not in wb.sheetnames:
            raise RuntimeError(f"Workbook missing required sheet: {required}")

    ws_astana = wb[ASTANA_SHEET]
    ws_totals = wb[TOTALS_SHEET]
    ws_dlv = wb[DLV_SHEET]

    totals = _extract_totals_sheet(ws_totals)
    astana_parts, astana_overall = _extract_astana_summary(ws_astana)
    astana_cny = _extract_astana_detail_cny(ws_astana)
    dlv_totals = _extract_dlv_totals(ws_dlv)
    per_bag_grand_weight = _extract_per_bag_grand_weight(ws_astana)

    errors: list[str] = []
    warnings: list[str] = []

    missing_in_totals = sorted(set(astana_parts) - set(totals))
    if missing_in_totals:
        errors.append(f"astana parts missing in totals sheet: {', '.join(missing_in_totals)}")

    for part_id in sorted(set(astana_parts) & set(totals)):
        a = astana_parts[part_id]
        t = totals[part_id]
        if abs(a["units"] - t["units"]) > tol_qty:
            errors.append(f"{part_id}: units astana={a['units']} totals={t['units']}")
        if abs(a["bags"] - t["bags"]) > tol_qty:
            errors.append(f"{part_id}: bags astana={a['bags']} totals={t['bags']}")
        if abs(a["actual_weight"] - t["actual_weight"]) > tol_weight:
            errors.append(
                f"{part_id}: actual_weight astana={a['actual_weight']} totals={t['actual_weight']}"
            )
        if abs(a["dlv_usd"] - t["dlv_usd"]) > tol_money:
            errors.append(f"{part_id}: dlv_usd astana={a['dlv_usd']} totals={t['dlv_usd']}")
        if abs(a["dlv_kzt"] - t["dlv_kzt"]) > tol_money:
            errors.append(f"{part_id}: dlv_kzt astana={a['dlv_kzt']} totals={t['dlv_kzt']}")
        if abs(astana_cny.get(part_id, 0.0) - t["cny"]) > tol_money:
            errors.append(
                f"{part_id}: cny detail astana={astana_cny.get(part_id, 0.0)} totals={t['cny']}"
            )

    sum_units = sum(v["units"] for v in astana_parts.values())
    sum_bags = sum(v["bags"] for v in astana_parts.values())
    sum_weight = sum(v["actual_weight"] for v in astana_parts.values())
    sum_dlv_usd = sum(v["dlv_usd"] for v in astana_parts.values())
    sum_dlv_kzt = sum(v["dlv_kzt"] for v in astana_parts.values())
    if abs(astana_overall["units"] - sum_units) > tol_qty:
        errors.append(f"astana overall units={astana_overall['units']} vs parts sum={sum_units}")
    if abs(astana_overall["bags"] - sum_bags) > tol_qty:
        errors.append(f"astana overall bags={astana_overall['bags']} vs parts sum={sum_bags}")
    if abs(astana_overall["actual_weight"] - sum_weight) > tol_weight:
        errors.append(
            f"astana overall actual_weight={astana_overall['actual_weight']} vs parts sum={sum_weight}"
        )
    if abs(astana_overall["dlv_usd"] - sum_dlv_usd) > tol_money:
        errors.append(f"astana overall dlv_usd={astana_overall['dlv_usd']} vs parts sum={sum_dlv_usd}")
    if abs(astana_overall["dlv_kzt"] - sum_dlv_kzt) > tol_money:
        errors.append(f"astana overall dlv_kzt={astana_overall['dlv_kzt']} vs parts sum={sum_dlv_kzt}")

    if abs(astana_overall["actual_weight"] - dlv_totals["actual_weight"]) > tol_weight:
        errors.append(
            f"overall actual_weight astana={astana_overall['actual_weight']} dlv_sheet={dlv_totals['actual_weight']}"
        )
    if abs(astana_overall["dlv_usd"] - dlv_totals["dlv_usd"]) > tol_money:
        errors.append(f"overall dlv_usd astana={astana_overall['dlv_usd']} dlv_sheet={dlv_totals['dlv_usd']}")
    if abs(astana_overall["dlv_kzt"] - dlv_totals["dlv_kzt"]) > tol_money:
        errors.append(f"overall dlv_kzt astana={astana_overall['dlv_kzt']} dlv_sheet={dlv_totals['dlv_kzt']}")

    if per_bag_grand_weight > 0 and abs(per_bag_grand_weight - astana_overall["actual_weight"]) > tol_weight:
        warnings.append(
            "per-bag grand weight is estimate-based and differs from summary actual weight: "
            f"per_bag={per_bag_grand_weight}, summary_actual={astana_overall['actual_weight']}"
        )

    return {
        "ok": len(errors) == 0,
        "workbook_path": str(workbook_path),
        "parts_checked": len(astana_parts),
        "errors": errors,
        "warnings": warnings,
        "summary": {
            "astana_overall": astana_overall,
            "dlv_totals": dlv_totals,
            "per_bag_grand_weight": per_bag_grand_weight,
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate 2.3.26_astana_totals alignment")
    parser.add_argument("--xlsx", type=Path, required=True, help="Inbound workbook path")
    parser.add_argument("--json", action="store_true", help="Print JSON payload")
    parser.add_argument("--tolerance-qty", type=float, default=0.0)
    parser.add_argument("--tolerance-weight", type=float, default=0.1)
    parser.add_argument("--tolerance-money", type=float, default=1.0)
    args = parser.parse_args()

    try:
        report = validate_astana_totals_alignment(
            workbook_path=args.xlsx.expanduser(),
            tol_qty=max(0.0, float(args.tolerance_qty)),
            tol_weight=max(0.0, float(args.tolerance_weight)),
            tol_money=max(0.0, float(args.tolerance_money)),
        )
    except Exception as exc:
        payload = {"ok": False, "error": str(exc)}
        if args.json:
            print(json.dumps(payload, ensure_ascii=False))
        else:
            print(f"ERROR: {exc}")
        return 1

    if args.json:
        print(json.dumps(report, ensure_ascii=False))
    else:
        print(
            f"astana_alignment: ok={report['ok']} parts_checked={report['parts_checked']} "
            f"errors={len(report['errors'])} warnings={len(report['warnings'])}"
        )
        for err in report["errors"]:
            print(f"ERROR {err}")
        for warn in report["warnings"]:
            print(f"WARN  {warn}")

    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
