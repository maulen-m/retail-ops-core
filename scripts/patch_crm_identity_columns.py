#!/usr/bin/env python3
"""
Patch historical CRM identity columns with canonical mapping:
- SKU_key
- Kaspi_name_core

Does not touch MY_SIZE.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List
import sys

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent.parent))
from scripts.import_orders_to_crm import (
    _derive_identity_from_raw_row,
    _load_article_identity_for_articles,
    _load_sku_meta_for_keys,
    compute_fixed_value_columns,
)


DEFAULT_WORKBOOK = Path("excel_ui/SALES_KSP_CRM_V3.xlsx")
DEFAULT_SHEET = "SALES_KSP_CRM_1"
DEFAULT_REPORT = Path("logs/crm_identity_patch_report.csv")
DEFAULT_BACKUP_DIR = Path("excel_ui/backups")


def _backup_file(target: Path, backup_dir: Path) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = backup_dir / f"CRM_backup_identity_patch_{ts}.xlsx"
    out.write_bytes(target.read_bytes())
    return out


def compute_identity_patch(
    raw_row: Dict[str, Any],
    article_identity_by_article: Dict[str, Dict[str, str]],
    valid_sku_keys: set[str] | None = None,
    sku_meta: Dict[str, Dict[str, Any]] | None = None,
    kaspi_core: Dict[str, str] | None = None,
) -> Dict[str, Any]:
    article = str(raw_row.get("Артикул") or "").strip().upper()
    if article.startswith("OF_SUIT-61_BLK_"):
        return {
            "sku_key": "CL_NEW-CLO2_MEN_SUIT-61_BLACK",
            "kaspi_name_core": "6в1_Черный_+Сумка",
        }

    identity = _derive_identity_from_raw_row(
        raw_row,
        article_identity_by_article=article_identity_by_article,
        valid_sku_keys=valid_sku_keys,
    )
    fixed = compute_fixed_value_columns(
        raw_row,
        sku_meta_by_key=sku_meta or {},
        kaspi_core_by_key=kaspi_core or {},
        article_identity_by_article=article_identity_by_article,
        valid_sku_keys=valid_sku_keys,
    )
    return {
        "sku_key": identity.get("sku_key") or fixed.get("SKU_key") or "",
        "kaspi_name_core": fixed.get("Kaspi_name_core") or "",
    }


def patch_workbook(
    workbook: Path,
    sheet: str,
    apply: bool,
    report: Path,
    backup_dir: Path,
    only_line61: bool = False,
) -> Dict[str, int]:
    df = pd.read_excel(workbook, sheet_name=sheet)
    cols = {str(c).strip(): c for c in df.columns}
    required = ["Артикул", "Название товара в Kaspi Магазине", "SKU_key", "Kaspi_name_core"]
    for name in required:
        if name not in cols:
            raise RuntimeError(f"Missing required column: {name}")

    rows: List[Dict[str, Any]] = []
    articles: List[str] = []
    for idx, rec in enumerate(df.to_dict(orient="records"), start=2):
        article = rec.get("Артикул")
        offer = rec.get("Название товара в Kaspi Магазине")
        if not article and not offer:
            continue
        raw = {
            "Артикул": article,
            "Название товара в Kaspi Магазине": offer,
            "SKU_key": rec.get("SKU_key"),
            "Product_Type": rec.get("Product_Type", ""),
        }
        rows.append({"row": idx, "raw": raw, "old_core": str(rec.get("Kaspi_name_core") or "").strip()})
        articles.append(str(article or "").strip())

    article_identity = _load_article_identity_for_articles(articles)
    seed_sku_keys = [str(x["raw"].get("SKU_key") or "").strip() for x in rows if str(x["raw"].get("SKU_key") or "").strip()]
    sku_meta, kaspi_core, valid_sku_keys = _load_sku_meta_for_keys(seed_sku_keys)

    updates: List[Dict[str, Any]] = []
    for item in rows:
        row_num = item["row"]
        raw = item["raw"]
        article = str(raw.get("Артикул") or "").strip().upper()
        old_sku = str(raw.get("SKU_key") or "").strip()
        old_core = item["old_core"]
        map_hit = article in article_identity
        line61 = article.startswith("OF_SUIT-61_BLK_")
        valid_current_sku = bool(old_sku and (not valid_sku_keys or old_sku in valid_sku_keys))

        patch = compute_identity_patch(
            raw,
            article_identity_by_article=article_identity,
            valid_sku_keys=valid_sku_keys,
            sku_meta=sku_meta,
            kaspi_core=kaspi_core,
        )
        new_sku = str(patch.get("sku_key") or "").strip()
        new_core = str(patch.get("kaspi_name_core") or "").strip()
        if only_line61 and not line61:
            continue
        if not only_line61:
            should_patch = (
                line61
                or map_hit
                or (old_sku and (not valid_current_sku))
                or (not old_core)
                or (new_sku and new_sku != old_sku)
                or (new_core and new_core != old_core and new_sku == old_sku)
            )
            if not should_patch:
                continue
        if (new_sku and new_sku != old_sku) or (new_core and new_core != old_core):
            updates.append(
                {
                    "row": row_num,
                    "old_sku_key": old_sku,
                    "new_sku_key": new_sku or old_sku,
                    "old_kaspi_name_core": old_core,
                    "new_kaspi_name_core": new_core or old_core,
                    "article": str(raw.get("Артикул") or ""),
                }
            )

    backup_path = None
    if apply and updates:
        backup_path = _backup_file(workbook, backup_dir)
        wb = load_workbook(workbook)
        try:
            sh = wb[sheet]
            headers = [sh.cell(row=1, column=c).value for c in range(1, sh.max_column + 1)]
            h2c = {str(h).strip(): i + 1 for i, h in enumerate(headers) if str(h or "").strip()}
            sku_key_col = h2c["SKU_key"]
            core_col = h2c["Kaspi_name_core"]
            for upd in updates:
                sh.cell(row=upd["row"], column=sku_key_col).value = upd["new_sku_key"]
                sh.cell(row=upd["row"], column=core_col).value = upd["new_kaspi_name_core"]
            wb.save(workbook)
        finally:
            wb.close()

    report.parent.mkdir(parents=True, exist_ok=True)
    with report.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["apply", int(apply)])
        w.writerow(["backup_path", str(backup_path) if backup_path else ""])
        w.writerow(["rows_scanned", len(rows)])
        w.writerow(["rows_updated", len(updates)])
        w.writerow([])
        w.writerow(["row", "article", "old_sku_key", "new_sku_key", "old_kaspi_name_core", "new_kaspi_name_core"])
        for upd in updates:
            w.writerow(
                [
                    upd["row"],
                    upd["article"],
                    upd["old_sku_key"],
                    upd["new_sku_key"],
                    upd["old_kaspi_name_core"],
                    upd["new_kaspi_name_core"],
                ]
            )

    return {"rows_scanned": len(rows), "rows_updated": len(updates), "backup_path": str(backup_path or "")}


def _build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Patch CRM SKU_key/Kaspi_name_core from canonical mapping")
    p.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    p.add_argument("--sheet", default=DEFAULT_SHEET)
    p.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    p.add_argument("--backup-dir", type=Path, default=DEFAULT_BACKUP_DIR)
    p.add_argument("--only-line61", action="store_true", help="Patch only OF_SUIT-61_BLK_* rows")
    p.add_argument("--apply", action="store_true")
    return p


def main() -> None:
    args = _build_parser().parse_args()
    stats = patch_workbook(
        workbook=args.workbook,
        sheet=args.sheet,
        apply=args.apply,
        report=args.report,
        backup_dir=args.backup_dir,
        only_line61=args.only_line61,
    )
    mode = "APPLY" if args.apply else "DRY RUN"
    print(f"CRM identity patch {mode}")
    print(f"  rows_scanned: {stats['rows_scanned']}")
    print(f"  rows_updated: {stats['rows_updated']}")
    if stats["backup_path"]:
        print(f"  backup: {stats['backup_path']}")
    print(f"  report: {args.report}")


if __name__ == "__main__":
    main()
