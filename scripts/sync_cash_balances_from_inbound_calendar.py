#!/usr/bin/env python3
"""Sync Cash_Balances sheet snapshot into bank_accounts history + snapshots."""

from __future__ import annotations

import argparse
from datetime import date, datetime
import os
from pathlib import Path
import sys
from typing import Any

import openpyxl
import yaml

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from core.db import DEFAULT_DB_PATH
from scripts import generate_bank_accounts_history_totals
from scripts import generate_bank_snapshot

DEFAULT_XLSX = (
    PROJECT_ROOT
    / "config"
    / "anchors"
    / "INBOUND_CALENDAR_LATEST.xlsx"
)
DEFAULT_HISTORY = PROJECT_ROOT / "config" / "bank_accounts_history.yaml"
DEFAULT_SNAPSHOT = PROJECT_ROOT / "config" / "bank_accounts.yaml"
DEFAULT_HISTORY_TOTALS = PROJECT_ROOT / "config" / "bank_accounts_history_totals.md"
SHEET_NAME = "Cash_Balances"
VALID_STORES = {"UNIVERSAL", "11KZ", "STOREB", "ACMEWEAR", "MELVIS"}


def _to_float(value: Any) -> float:
    try:
        if value is None:
            return 0.0
        if isinstance(value, str) and not value.strip():
            return 0.0
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _parse_sheet_ts(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value.replace(microsecond=0)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    raw = str(value or "").replace("\n", " ").strip()
    if not raw:
        raise ValueError("empty timestamp header")
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    raise ValueError(f"Unrecognized sheet timestamp: {value!r}")


def _to_as_of(dt: datetime) -> str:
    return f"{dt.strftime('%Y-%m-%d %H:%M:%S')} GMT+5"


def _read_snapshot(
    *,
    xlsx_path: Path,
    sheet_name: str,
    snapshot_ts: str | None = None,
) -> tuple[str, list[dict[str, Any]], dict[str, Any]]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    if sheet_name not in wb.sheetnames:
        raise RuntimeError(f"Workbook missing sheet: {sheet_name}")
    ws = wb[sheet_name]

    ts_cols: list[tuple[int, datetime, str]] = []
    for c in range(4, ws.max_column + 1):
        raw = ws.cell(4, c).value
        if raw is None or str(raw).strip() == "":
            continue
        try:
            dt = _parse_sheet_ts(raw)
        except ValueError:
            continue
        ts_cols.append((c, dt, str(raw).replace("\n", " ").strip()))
    if not ts_cols:
        raise RuntimeError(f"{sheet_name}: no snapshot timestamp columns found in row 4")

    selected_col = None
    selected_dt = None
    selected_label = None
    if snapshot_ts:
        target = snapshot_ts.strip()
        for c, dt, label in ts_cols:
            if label == target:
                selected_col, selected_dt, selected_label = c, dt, label
                break
        if selected_col is None:
            raise RuntimeError(f"Requested snapshot timestamp not found: {snapshot_ts}")
    else:
        selected_col, selected_dt, selected_label = max(ts_cols, key=lambda x: x[1])

    balances: list[dict[str, Any]] = []
    for r in range(5, min(ws.max_row, 240) + 1):
        store = str(ws.cell(r, 1).value or "").strip()
        account = str(ws.cell(r, 2).value or "").strip()
        currency = str(ws.cell(r, 3).value or "").strip().upper()
        if store not in VALID_STORES:
            continue
        if not store or not account or not currency:
            continue
        amount = _to_float(ws.cell(r, selected_col).value)
        balances.append(
            {
                "store": store,
                "account": account,
                "amount": amount,
                "currency": currency,
            }
        )

    if not balances:
        raise RuntimeError(f"{sheet_name}: no account rows parsed for selected snapshot")

    metadata = {
        "selected_column": selected_col,
        "selected_label": selected_label,
        "selected_dt": selected_dt.isoformat() if selected_dt else "",
        "account_rows": len(balances),
    }
    as_of = _to_as_of(selected_dt)
    return as_of, balances, metadata


def _load_history(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    payload = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    entries = payload.get("entries", [])
    if not isinstance(entries, list):
        raise RuntimeError(f"Invalid history structure at {path}")
    return entries


def _write_history(path: Path, entries: list[dict[str, Any]]) -> None:
    payload = {"entries": entries}
    text = yaml.safe_dump(payload, sort_keys=False, allow_unicode=False)
    if not text.startswith("entries:\n"):
        text = "entries:\n" + text
    path.write_text(text, encoding="utf-8")


def sync_cash_balances(
    *,
    xlsx_path: Path,
    history_path: Path,
    snapshot_path: Path,
    history_totals_path: Path,
    db_path: Path,
    snapshot_ts: str | None,
    apply: bool,
) -> dict[str, Any]:
    as_of, balances, metadata = _read_snapshot(
        xlsx_path=xlsx_path,
        sheet_name=SHEET_NAME,
        snapshot_ts=snapshot_ts,
    )
    entries = _load_history(history_path)

    new_entry = {
        "as_of": as_of,
        "source": "manual snapshot (Cash_Balances)",
        "balances": balances,
    }
    existing_idx = next((i for i, row in enumerate(entries) if str(row.get("as_of") or "") == as_of), None)
    action = "updated" if existing_idx is not None else "appended"
    if existing_idx is not None:
        entries[existing_idx] = new_entry
    else:
        entries.append(new_entry)

    entries_sorted = sorted(entries, key=lambda row: generate_bank_snapshot._parse_as_of(str(row.get("as_of") or "")))
    latest_effective = generate_bank_snapshot.get_effective_latest_entry(entries_sorted)
    fx_rates = generate_bank_snapshot.get_fx_rates(db_path)
    snapshot_content = generate_bank_snapshot.generate_snapshot(latest_effective, fx_rates)
    totals_content = generate_bank_accounts_history_totals.generate_history_totals_markdown(
        entries_sorted,
        fx_rates,
        str(history_path),
        mode="compact",
    )

    if apply:
        if os.environ.get("ENABLE_CASH_BALANCE_SYNC_WRITE") != "1":
            raise RuntimeError("ENABLE_CASH_BALANCE_SYNC_WRITE=1 is required to apply writes.")
        history_path.parent.mkdir(parents=True, exist_ok=True)
        snapshot_path.parent.mkdir(parents=True, exist_ok=True)
        history_totals_path.parent.mkdir(parents=True, exist_ok=True)
        _write_history(history_path, entries_sorted)
        snapshot_path.write_text(snapshot_content + ("\n" if not snapshot_content.endswith("\n") else ""), encoding="utf-8")
        history_totals_path.write_text(totals_content + ("\n" if not totals_content.endswith("\n") else ""), encoding="utf-8")

    return {
        "apply": bool(apply),
        "action": action,
        "as_of": as_of,
        "selected_label": metadata["selected_label"],
        "selected_column": metadata["selected_column"],
        "rows": metadata["account_rows"],
        "history_entries": len(entries_sorted),
        "snapshot_output": str(snapshot_path),
        "history_output": str(history_path),
        "history_totals_output": str(history_totals_path),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Sync Cash_Balances sheet into bank accounts history/snapshot")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX, help="Inbound workbook path")
    parser.add_argument("--history", type=Path, default=DEFAULT_HISTORY, help="bank_accounts_history.yaml path")
    parser.add_argument("--snapshot", type=Path, default=DEFAULT_SNAPSHOT, help="bank_accounts.yaml path")
    parser.add_argument(
        "--history-totals",
        type=Path,
        default=DEFAULT_HISTORY_TOTALS,
        help="bank_accounts_history_totals.md path",
    )
    parser.add_argument("--db", type=Path, default=DEFAULT_DB_PATH, help="DB path for FX rates")
    parser.add_argument(
        "--snapshot-ts",
        type=str,
        default=None,
        help="Timestamp label from row 4 (e.g. '02.03.2026 14:41:00'). Default: latest.",
    )
    parser.add_argument("--apply", action="store_true", help="Write outputs (requires env gate)")
    args = parser.parse_args()

    try:
        report = sync_cash_balances(
            xlsx_path=args.xlsx.expanduser(),
            history_path=args.history.expanduser(),
            snapshot_path=args.snapshot.expanduser(),
            history_totals_path=args.history_totals.expanduser(),
            db_path=args.db.expanduser(),
            snapshot_ts=args.snapshot_ts,
            apply=bool(args.apply),
        )
    except Exception as exc:
        print(f"ERROR: {exc}")
        return 1

    mode = "APPLY" if report["apply"] else "DRY-RUN"
    print(f"{mode} cash balance sync")
    print(f"as_of: {report['as_of']}")
    print(f"snapshot_column: {report['selected_column']} ({report['selected_label']})")
    print(f"rows: {report['rows']}")
    print(f"history_action: {report['action']}")
    print(f"history_entries: {report['history_entries']}")
    print(f"history_output: {report['history_output']}")
    print(f"snapshot_output: {report['snapshot_output']}")
    print(f"history_totals_output: {report['history_totals_output']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
