#!/usr/bin/env python3
"""Validate source freshness behind the LINE31 active-goal owner facts.

This verifier is intentionally read-only. It does not prove that final creative
assets are ready; it only prevents the active-goal audit from treating cached
owner-fact JSON as sufficient when the underlying workbook/source packet is
missing, stale, or contradictory.
"""

from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import sys
from typing import Any
from zoneinfo import ZoneInfo

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

ALMATY_TZ = ZoneInfo("Asia/Almaty")
DEFAULT_OWNER_FACTS_PATH = (
    PROJECT_ROOT
    / "exports/validation/line31_green_except_creative_owner_clarified_current_20260601_130108/"
    "CURRENT_OWNER_CLARIFICATION_FACTS.json"
)

REQUIRED_CASH_SHEETS = {
    "Cash_Balances",
    "base_payment_SHR_log",
    "SHR_Receipt_Ledger",
    "SHR_Pay_Lock_105952",
}
DATE_RE = re.compile(
    r"(?P<day>\d{1,2})[.](?P<month>\d{1,2})[.](?P<year>\d{4})"
    r"(?:[\s_\n]+(?P<hour>\d{1,2})[:_](?P<minute>\d{1,2})(?:[:_](?P<second>\d{1,2}))?)?"
)


def _read_json(path: Path) -> dict[str, Any]:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return {}


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _parse_cash_timestamp(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=ALMATY_TZ)
    if value is None:
        return None
    match = DATE_RE.search(str(value))
    if not match:
        return None
    parts = match.groupdict()
    return datetime(
        int(parts["year"]),
        int(parts["month"]),
        int(parts["day"]),
        int(parts.get("hour") or 0),
        int(parts.get("minute") or 0),
        int(parts.get("second") or 0),
        tzinfo=ALMATY_TZ,
    )


def _owner_timestamp(dt: datetime | None) -> str | None:
    if dt is None:
        return None
    return dt.astimezone(ALMATY_TZ).strftime("%Y-%m-%d %H:%M:%S GMT+5")


def _latest_cash_balance_timestamp(ws: Any) -> str | None:
    latest: datetime | None = None
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True):
        for value in row:
            parsed = _parse_cash_timestamp(value)
            if parsed and (latest is None or parsed > latest):
                latest = parsed
    return _owner_timestamp(latest)


def _rows(ws: Any) -> list[tuple[Any, ...]]:
    return [tuple(row) for row in ws.iter_rows(values_only=True)]


def _find_row(rows: list[tuple[Any, ...]], predicate: Any) -> tuple[Any, ...] | None:
    for row in rows:
        if predicate(row):
            return row
    return None


def _find_value_after_label(rows: list[tuple[Any, ...]], label: str) -> Any:
    for row in rows:
        for idx, value in enumerate(row):
            if str(value).strip() == label:
                for candidate in row[idx + 1 :]:
                    if candidate is not None:
                        return candidate
    return None


def _as_number(value: Any) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    if value is None:
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def _close_enough(value: Any, expected: Any) -> bool:
    left = _as_number(value)
    right = _as_number(expected)
    return left is not None and right is not None and abs(left - right) < 0.0001


def _validate_cash_and_shr(
    owner_facts: dict[str, Any],
    errors: list[str],
    warnings: list[str],
) -> dict[str, Any]:
    cash = owner_facts.get("cash", {})
    shr = owner_facts.get("shr", {})
    workbook_path = Path(str(cash.get("source_workbook", "")))
    result: dict[str, Any] = {
        "workbook_path": str(workbook_path),
        "exists": workbook_path.exists(),
        "expected_latest_timestamp": cash.get("current_timestamp"),
        "latest_cash_balance_timestamp": None,
        "sha256": None,
        "required_sheets_present": [],
        "required_sheets_missing": sorted(REQUIRED_CASH_SHEETS),
        "shr_payment_18_base_log_present": False,
        "shr_payment_18_ledger_present": False,
        "shr_actual_paid_base_cny": None,
        "shr_actual_remaining_cny": None,
    }
    if not workbook_path.exists():
        errors.append(f"cash workbook missing: {workbook_path}")
        return result

    result["sha256"] = _sha256(workbook_path)
    stat = workbook_path.stat()
    result["size_bytes"] = stat.st_size
    result["mtime_almaty"] = datetime.fromtimestamp(
        stat.st_mtime, ALMATY_TZ
    ).isoformat(timespec="seconds")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment guard
        errors.append(f"openpyxl unavailable for cash workbook freshness check: {exc}")
        return result

    wb = None
    try:
        wb = load_workbook(workbook_path, read_only=True, data_only=True)
        sheet_names = set(wb.sheetnames)
        missing = sorted(REQUIRED_CASH_SHEETS - sheet_names)
        present = sorted(REQUIRED_CASH_SHEETS & sheet_names)
        result["required_sheets_present"] = present
        result["required_sheets_missing"] = missing
        if missing:
            errors.append(f"cash workbook missing required sheets: {', '.join(missing)}")
            return result

        result["latest_cash_balance_timestamp"] = _latest_cash_balance_timestamp(
            wb["Cash_Balances"]
        )
        if result["latest_cash_balance_timestamp"] != cash.get("current_timestamp"):
            errors.append(
                "Cash_Balances latest timestamp mismatch: "
                f"{result['latest_cash_balance_timestamp']} != {cash.get('current_timestamp')}"
            )

        base_rows = _rows(wb["base_payment_SHR_log"])
        ledger_rows = _rows(wb["SHR_Receipt_Ledger"])
        result["shr_actual_paid_base_cny"] = _find_value_after_label(
            base_rows, "Actual paid (receipts + approval)"
        )
        result["shr_actual_remaining_cny"] = _find_value_after_label(
            base_rows, "Actual remaining CNY"
        )
        if not _close_enough(
            result["shr_actual_paid_base_cny"], shr.get("paid_base_total_cny")
        ):
            errors.append("SHR paid base total does not match owner facts")
        if not _close_enough(
            result["shr_actual_remaining_cny"], shr.get("remaining_base_payable_cny")
        ):
            errors.append("SHR remaining payable does not match owner facts")

        payment_18 = _find_row(
            base_rows,
            lambda row: len(row) >= 3
            and row[0] == 18
            and _close_enough(row[2], shr.get("po5_payment_18_cny")),
        )
        ledger_18 = _find_row(
            ledger_rows,
            lambda row: len(row) >= 9
            and row[0] == "PO5_018"
            and _close_enough(row[8], shr.get("po5_payment_18_cny")),
        )
        result["shr_payment_18_base_log_present"] = payment_18 is not None
        result["shr_payment_18_ledger_present"] = ledger_18 is not None
        if payment_18 is None:
            errors.append("SHR payment #18 7000 CNY is missing from base_payment_SHR_log")
        if ledger_18 is None:
            warnings.append(
                "SHR payment #18 owner-paid row is not present in SHR_Receipt_Ledger; "
                "final exchanger receipt may still be pending"
            )
    except Exception as exc:  # noqa: BLE001
        errors.append(f"cash workbook freshness read failed: {exc}")
    finally:
        if wb is not None:
            wb.close()
    return result


def _validate_line31_stock_packet(
    owner_facts: dict[str, Any],
    errors: list[str],
) -> dict[str, Any]:
    line31_stock = owner_facts.get("line31_stock", {})
    source_packet = Path(str(line31_stock.get("source_packet", "")))
    manifest_path = source_packet / "product_truth_yellow_to_apply_ready_manifest.json"
    validation_path = source_packet / "validation_summary.json"
    closeout_path = source_packet / "closeout.md"
    result: dict[str, Any] = {
        "source_packet": str(source_packet),
        "exists": source_packet.exists(),
        "manifest_path": str(manifest_path),
        "manifest_exists": manifest_path.exists(),
        "validation_summary_exists": validation_path.exists(),
        "closeout_exists": closeout_path.exists(),
        "manifest_gate": None,
        "validation_gate": None,
        "totals": {},
    }
    if not source_packet.exists():
        errors.append(f"LINE31 stock source packet missing: {source_packet}")
        return result
    manifest = _read_json(manifest_path)
    validation = _read_json(validation_path)
    if not manifest:
        errors.append(f"LINE31 stock manifest missing or invalid: {manifest_path}")
    if not validation:
        errors.append(f"LINE31 stock validation summary missing or invalid: {validation_path}")

    result["manifest_gate"] = manifest.get("gate")
    result["validation_gate"] = validation.get("gate")
    result["totals"] = manifest.get("totals", {})
    if manifest.get("gate") != "GREEN":
        errors.append("LINE31 stock manifest gate is not GREEN")
    if validation.get("gate") != "GREEN":
        errors.append("LINE31 stock validation gate is not GREEN")
    if closeout_path.exists() and "Gate: `GREEN`" not in closeout_path.read_text(
        encoding="utf-8", errors="replace"
    ):
        errors.append("LINE31 stock closeout does not contain Gate: `GREEN`")

    expected_totals = {
        "line31_physical_total_stock_qty": "physical_total_stock_qty",
        "line31_physical_sellable_stock_qty": "physical_sellable_stock_qty",
        "line31_physical_sellable_available_after_open_reserved_qty": (
            "physical_sellable_available_after_open_reserved_qty"
        ),
        "line31_physical_not_for_sale_reserve_qty": "physical_not_for_sale_reserve_qty",
        "line31_economic_final_sales_estimate_qty": "economic_final_sales_estimate_qty",
        "line31_open_reserved_on_delivery_exposure_qty": (
            "open_reserved_on_delivery_exposure_qty"
        ),
    }
    for manifest_key, facts_key in expected_totals.items():
        if not _close_enough(result["totals"].get(manifest_key), line31_stock.get(facts_key)):
            errors.append(f"LINE31 stock total mismatch for {manifest_key}")
    return result


def validate_owner_objective_source_freshness(
    owner_facts_path: Path = DEFAULT_OWNER_FACTS_PATH,
) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    owner_facts = _read_json(owner_facts_path)
    if not owner_facts:
        errors.append(f"owner facts JSON missing or invalid: {owner_facts_path}")
        owner_facts = {}

    cash = owner_facts.get("cash", {})
    shr = owner_facts.get("shr", {})
    line31_stock = owner_facts.get("line31_stock", {})
    protected_reserve_min_kzt = cash.get("protected_reserve_min_kzt")
    protected_reserve_ok = protected_reserve_min_kzt == 800000
    if not protected_reserve_ok:
        errors.append("protected reserve is not exactly 800000 KZT")
    if shr.get("po5_payment_18_cny") != 7000:
        errors.append("SHR payment #18 is not recorded as 7000 CNY")
    if shr.get("po5_payment_18_final_exchanger_receipt_pending") is not True:
        errors.append("SHR payment #18 receipt-pending flag is not true")
    if (
        line31_stock.get("basis")
        != "owner-approved exact rebuild from April leftovers plus PO1-A arrival"
    ):
        errors.append("LINE31 stock basis is not the owner-approved April+PO1-A rebuild")

    result = {
        "generated_at": datetime.now(ALMATY_TZ).isoformat(timespec="seconds"),
        "owner_facts_path": str(owner_facts_path),
        "owner_facts_exists": owner_facts_path.exists(),
        "protected_reserve": {
            "expected_min_kzt": 800000,
            "owner_fact_min_kzt": protected_reserve_min_kzt,
            "ok": protected_reserve_ok,
        },
        "cash_and_shr": _validate_cash_and_shr(owner_facts, errors, warnings),
        "line31_stock": _validate_line31_stock_packet(owner_facts, errors),
        "warnings": warnings,
        "errors": errors,
    }
    result["ok"] = not errors
    result["gate"] = "GREEN_SOURCE_FRESHNESS" if result["ok"] else "YELLOW_SOURCE_WEAK"
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--owner-facts", type=Path, default=DEFAULT_OWNER_FACTS_PATH)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)

    result = validate_owner_objective_source_freshness(args.owner_facts)
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True))
    else:
        print(f"Gate: {result['gate']}")
        for error in result["errors"]:
            print(f"ERROR: {error}")
        for warning in result["warnings"]:
            print(f"WARNING: {warning}")
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
