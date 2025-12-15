#!/usr/bin/env python3
"""
Refresh D_size_mix_reference.xlsx with recent sales data.

Closes the ~20% gap between stale anchor data and actual demand.
Uses 60-day rolling average for stability while capturing trends.

Usage:
    python3 scripts/refresh_anchor_demand.py [--lookback DAYS] [--dry-run]
"""

import sqlite3
import argparse
from pathlib import Path
from datetime import date, timedelta
from collections import defaultdict
import openpyxl

PROJECT_ROOT = Path(__file__).parent.parent
DB_PATH = PROJECT_ROOT / "db" / "app.db"
ANCHOR_FILE = PROJECT_ROOT / "excel" / "D_size_mix_reference.xlsx"
DEFAULT_LOOKBACK = 60  # Days


def get_demand_from_sales(conn, lookback_days: int) -> dict:
    """
    Calculate D_active and D_size from fact_sales_daily_size.

    Returns:
        {sku_key: {
            'd_active': float,      # Total daily demand
            'sizes': {size: d_size}, # Per-size daily demand
            'sigma': float,          # Volatility (D x 0.4)
            'days_with_sales': int   # Data quality indicator
        }}
    """
    cutoff = (date.today() - timedelta(days=lookback_days)).isoformat()

    # Get size-level aggregates
    rows = conn.execute("""
        SELECT
            sku_key,
            my_size,
            SUM(units) as total_units,
            COUNT(DISTINCT sale_date) as days_with_sales,
            AVG(units) as avg_daily
        FROM fact_sales_daily_size
        WHERE sale_date >= ?
          AND my_size IS NOT NULL
          AND my_size != ''
        GROUP BY sku_key, my_size
    """, (cutoff,)).fetchall()

    # Aggregate by SKU
    sku_data = defaultdict(lambda: {
        'sizes': {},
        'total_units': 0,
        'days_with_sales': 0
    })

    for row in rows:
        sku_key = row['sku_key']
        size = row['my_size']
        total = row['total_units'] or 0
        days = row['days_with_sales'] or 0

        d_size = total / lookback_days
        sku_data[sku_key]['sizes'][size] = round(d_size, 4)
        sku_data[sku_key]['total_units'] += total
        sku_data[sku_key]['days_with_sales'] = max(
            sku_data[sku_key]['days_with_sales'], days
        )

    # Build final result
    result = {}
    for sku_key, data in sku_data.items():
        d_active = data['total_units'] / lookback_days
        sigma = d_active * 0.4  # Standard volatility factor

        result[sku_key] = {
            'd_active': round(d_active, 4),
            'sizes': data['sizes'],
            'sigma': round(sigma, 4),
            'days_with_sales': data['days_with_sales']
        }

    return result


def update_anchor_file(demand_data: dict, anchor_path: Path, dry_run: bool = False) -> tuple:
    """
    Update anchor Excel file with fresh demand data.

    Returns:
        ({'updated': int, 'skipped': int, 'not_found': int}, changes_list)
    """
    wb = openpyxl.load_workbook(anchor_path)
    ws = wb.active

    # Build header map
    headers = {}
    for cell in ws[1]:
        if cell.value:
            headers[cell.value] = cell.column

    # Size column mapping
    size_col_map = {
        'S': 'S_D', 'M': 'M_D', 'L': 'L_D', 'XL': 'XL_D',
        '2XL': '2XL_D', '3XL': '3XL_D', '4XL': '4XL_D',
        '22': '22_D', '24': '24_D', '26': '26_D', '28': '28_D', '30': '30_D'
    }

    stats = {'updated': 0, 'skipped': 0, 'not_found': 0}
    changes = []

    for row in ws.iter_rows(min_row=2):
        sku_key = row[0].value
        if not sku_key:
            continue

        if sku_key not in demand_data:
            stats['not_found'] += 1
            continue

        data = demand_data[sku_key]
        row_num = row[0].row
        old_d = ws.cell(row=row_num, column=headers.get('D_active', 2)).value or 0

        # Handle non-numeric old values
        try:
            old_d = float(old_d)
        except (ValueError, TypeError):
            old_d = 0

        new_d = data['d_active']

        # Skip if no significant change (< 5%)
        if old_d > 0 and abs(new_d / old_d - 1) < 0.05:
            stats['skipped'] += 1
            continue

        change_pct = ((new_d / old_d - 1) * 100) if old_d > 0 else 999
        changes.append({
            'sku': sku_key,
            'old_d': old_d,
            'new_d': new_d,
            'change_pct': change_pct
        })

        if not dry_run:
            # Update D_active
            if 'D_active' in headers:
                ws.cell(row=row_num, column=headers['D_active'], value=data['d_active'])

            # Update size demands
            for size, col_name in size_col_map.items():
                if col_name in headers:
                    d_size = data['sizes'].get(size, 0)
                    ws.cell(row=row_num, column=headers[col_name], value=d_size)

            # Update sigma
            if 'sigma' in headers:
                ws.cell(row=row_num, column=headers['sigma'], value=data['sigma'])

        stats['updated'] += 1

    if not dry_run:
        wb.save(anchor_path)

    wb.close()

    return stats, changes


def main():
    parser = argparse.ArgumentParser(description='Refresh anchor demand data from sales')
    parser.add_argument('--lookback', type=int, default=DEFAULT_LOOKBACK,
                        help=f'Days of sales history to use (default: {DEFAULT_LOOKBACK})')
    parser.add_argument('--dry-run', action='store_true',
                        help='Show changes without applying them')
    args = parser.parse_args()

    print(f"{'[DRY RUN] ' if args.dry_run else ''}Refreshing anchor demand...")
    print(f"  Database: {DB_PATH}")
    print(f"  Anchor file: {ANCHOR_FILE}")
    print(f"  Lookback: {args.lookback} days")
    print()

    # Check files exist
    if not DB_PATH.exists():
        print(f"ERROR: Database not found: {DB_PATH}")
        return 1
    if not ANCHOR_FILE.exists():
        print(f"ERROR: Anchor file not found: {ANCHOR_FILE}")
        return 1

    # Get demand from sales
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row

    demand_data = get_demand_from_sales(conn, args.lookback)
    conn.close()

    print(f"Calculated demand for {len(demand_data)} SKUs from sales data")

    # Update anchor file
    stats, changes = update_anchor_file(demand_data, ANCHOR_FILE, args.dry_run)

    print(f"\nResults:")
    print(f"  Updated: {stats['updated']}")
    print(f"  Skipped (< 5% change): {stats['skipped']}")
    print(f"  Not in sales data: {stats['not_found']}")

    if changes:
        print(f"\nTop changes:")
        # Sort by absolute change
        sorted_changes = sorted(changes, key=lambda x: abs(x['change_pct']), reverse=True)[:10]
        print(f"  {'SKU':<40} {'Old D':>8} {'New D':>8} {'Change':>8}")
        print(f"  {'-'*40} {'-'*8} {'-'*8} {'-'*8}")
        for c in sorted_changes:
            print(f"  {c['sku']:<40} {c['old_d']:>8.2f} {c['new_d']:>8.2f} {c['change_pct']:>+7.1f}%")

    if args.dry_run:
        print(f"\n[DRY RUN] No changes saved. Remove --dry-run to apply.")
    else:
        print(f"\nAnchor file updated: {ANCHOR_FILE.name}")

    return 0


if __name__ == "__main__":
    exit(main())
