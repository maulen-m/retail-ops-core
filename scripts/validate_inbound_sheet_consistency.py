#!/usr/bin/env python3
"""Fail-closed inbound workbook sheet consistency validator."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd


INBOUNDS_SHEET = "Inbounds_sheet"
TOTALS_SHEET = "PO_part_id_Totals"
CARGO_PREFIX = "cargo_send_"
LINE61_SHORTAGE_CLASSIFICATION_ID = "PO_ACCEPTED_REAL_SHORTAGE_LINE61_2026_05_OWNER_CONFIRMED"
LINE61_SHORTAGE_BY_SIZE = {"XL": 7, "2XL": 5, "3XL": 6, "4XL": 5}
LINE61_SHORTAGE_CONTRACT_PATH = (
    "docs/contracts/mvos_source_contracts/PO_LINE61_REAL_SHORTAGE_PRODUCTION_SAFE_V1.md"
)
LINE61_ACCEPTED_SHORTAGE_RULES = (
    {
        "type": "cargo_vs_inbounds",
        "po_part_id": "PO-4.0",
        "sku_key": "CL_NEW-CLO2_MEN_SUIT-61_BLACK",
        "expected_qty": 92.0,
        "observed_qty": 115.0,
        "delta_qty": 23.0,
    },
    {
        "type": "totals_vs_inbounds",
        "po_part_id": "PO-4.0",
        "sku_key": "*PART_TOTAL*",
        "expected_qty": 1902.0,
        "observed_qty": 1925.0,
        "delta_qty": 23.0,
    },
)


def _is_valid_part_id(value: Any) -> bool:
    txt = _norm_part(value)
    if not txt:
        return False
    upper = txt.upper()
    if any(token in upper for token in ("TOTAL", "PENDING", "UNPAID", "PAYMENT")):
        return False
    return True


def _normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    mapping: dict[str, str] = {}
    for col in df.columns:
        raw = str(col or "").strip()
        key = raw.lower().replace(" ", "_")
        mapping[col] = key
    return df.rename(columns=mapping)


def _pick_column(columns: set[str], aliases: list[str]) -> str | None:
    for alias in aliases:
        key = alias.lower().replace(" ", "_")
        if key in columns:
            return key
    return None


def _to_float(value: Any) -> float:
    try:
        if value is None:
            return 0.0
        if isinstance(value, str) and not value.strip():
            return 0.0
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _norm_part(value: Any) -> str:
    return str(value or "").strip()


def _norm_sku(value: Any) -> str:
    return str(value or "").strip()


def _load_inbounds(path: Path) -> tuple[dict[tuple[str, str], float], dict[str, float], str]:
    df = pd.read_excel(path, sheet_name=INBOUNDS_SHEET, dtype=object)
    df = _normalize_columns(df)
    cols = set(df.columns)

    part_col = _pick_column(cols, ["PO_part_id", "po_part_id", "po_part"])
    sku_col = _pick_column(cols, ["SKU_key", "sku_key", "sku_id", "sku"])
    qty_col = _pick_column(cols, ["Qty", "qty", "quantity"])
    actual_col = _pick_column(cols, ["Actual_qty", "actual_qty", "actual_quantity"])

    if not part_col or not sku_col:
        raise RuntimeError("Inbounds_sheet missing required columns: PO_part_id/SKU_key")
    if not qty_col and not actual_col:
        raise RuntimeError("Inbounds_sheet missing quantity columns (Qty/Actual_qty)")

    use_actual = bool(actual_col)
    chosen_col = actual_col if use_actual else qty_col
    assert chosen_col is not None

    expected_by_key: dict[tuple[str, str], float] = {}
    expected_by_part: dict[str, float] = {}

    for _, row in df.iterrows():
        part_id = _norm_part(row.get(part_col))
        sku_key = _norm_sku(row.get(sku_col))
        if not _is_valid_part_id(part_id) or not sku_key:
            continue
        qty = _to_float(row.get(qty_col)) if qty_col else 0.0
        actual = _to_float(row.get(actual_col)) if actual_col else 0.0
        # Prefer actual quantity when available and positive.
        expected = actual if (actual_col and actual > 0) else qty
        key = (part_id, sku_key)
        expected_by_key[key] = expected_by_key.get(key, 0.0) + expected
        expected_by_part[part_id] = expected_by_part.get(part_id, 0.0) + expected

    return expected_by_key, expected_by_part, chosen_col


def _load_totals(path: Path) -> dict[str, float]:
    df = pd.read_excel(path, sheet_name=TOTALS_SHEET, dtype=object)
    df = _normalize_columns(df)
    cols = set(df.columns)

    part_col = _pick_column(cols, ["PO_part_id", "po_part_id", "po_part"])
    total_units_col = _pick_column(cols, ["Total_Units", "total_units", "total_unit", "total_qty"])
    if not part_col or not total_units_col:
        return {}

    totals: dict[str, float] = {}
    for _, row in df.iterrows():
        part_id = _norm_part(row.get(part_col))
        if not _is_valid_part_id(part_id):
            continue
        totals[part_id] = _to_float(row.get(total_units_col))
    return totals


def _load_cargo_sheets(path: Path) -> tuple[dict[tuple[str, str], float], dict[tuple[str, str], set[str]]]:
    observed: dict[tuple[str, str], float] = {}
    key_sources: dict[tuple[str, str], set[str]] = {}
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)

    def _add_row(part_id: str, sku_key: str, qty: float, sheet: str) -> None:
        if not _is_valid_part_id(part_id) or not sku_key or qty <= 0:
            return
        key = (part_id, sku_key)
        observed[key] = observed.get(key, 0.0) + qty
        key_sources.setdefault(key, set()).add(sheet)

    def _collect_tabular(sheet: str) -> bool:
        df = pd.read_excel(path, sheet_name=sheet, dtype=object)
        df = _normalize_columns(df)
        cols = set(df.columns)
        part_col = _pick_column(cols, ["PO_part_id", "po_part_id", "po_part"])
        sku_col = _pick_column(cols, ["SKU_key", "sku_key", "sku_id", "sku"])
        qty_col = _pick_column(cols, ["Qty", "qty", "quantity", "actual_qty"])
        if not part_col or not sku_col or not qty_col:
            return False

        found = False
        for _, row in df.iterrows():
            part_id = _norm_part(row.get(part_col))
            sku_key = _norm_sku(row.get(sku_col))
            qty = _to_float(row.get(qty_col))
            if not _is_valid_part_id(part_id) or not sku_key or qty <= 0:
                continue
            _add_row(part_id, sku_key, qty, sheet)
            found = True
        return found

    def _collect_styled(sheet: str) -> bool:
        ws = workbook[sheet]
        part_id_fallback = ""

        for r in range(1, 30):
            k = str(ws.cell(r, 1).value or "").strip().upper()
            if k != "PO_PART_ID":
                continue
            # row can be single-part (col B) or multi-part (cols C..)
            candidates = [ws.cell(r, c).value for c in range(2, 10)]
            for val in candidates:
                cand = _norm_part(val)
                if _is_valid_part_id(cand):
                    part_id_fallback = cand
                    break
            if part_id_fallback:
                break

        header_row = None
        header_cols: dict[str, int] = {}
        for r in range(1, 260):
            row_vals = [str(ws.cell(r, c).value or "").strip().lower() for c in range(1, 26)]
            if "sku_key" in row_vals and "qty" in row_vals:
                header_row = r
                for c, val in enumerate(row_vals, start=1):
                    if val:
                        header_cols[val] = c
                break

        if not header_row:
            return False

        sku_col = header_cols.get("sku_key")
        qty_col = header_cols.get("qty")
        part_col = header_cols.get("po_part_id")
        po_name_col = header_cols.get("po_name")
        if not sku_col or not qty_col:
            return False

        found = False
        for r in range(header_row + 1, min(ws.max_row, header_row + 1200) + 1):
            marker = str(ws.cell(r, 1).value or "").strip().upper()
            if found and (
                "TOTALS" in marker
                or "PER BAG" in marker
                or "COMBINED" in marker
                or marker.startswith("SKU_KEY")
            ):
                break

            sku_key = _norm_sku(ws.cell(r, sku_col).value)
            if not sku_key:
                continue
            upper = sku_key.upper()
            if any(token in upper for token in ("TOTAL", "SUBTOTAL", "PER BAG", "COMBINED")):
                continue
            qty = _to_float(ws.cell(r, qty_col).value)
            if qty <= 0:
                continue

            part_id = ""
            if part_col:
                part_id = _norm_part(ws.cell(r, part_col).value)
            if (not part_id or not _is_valid_part_id(part_id)) and po_name_col:
                part_id = _norm_part(ws.cell(r, po_name_col).value)
            if not part_id or not _is_valid_part_id(part_id):
                part_id = part_id_fallback
            if not part_id or not _is_valid_part_id(part_id):
                continue

            _add_row(part_id, sku_key, qty, sheet)
            found = True
        return found

    for sheet in workbook.sheetnames:
        if not sheet.lower().startswith(CARGO_PREFIX):
            continue
        if _collect_tabular(sheet):
            continue
        _collect_styled(sheet)

    return observed, key_sources


def _matches_rule(mismatch: dict[str, Any], rule: dict[str, Any]) -> bool:
    if mismatch.get("type") != rule["type"]:
        return False
    if mismatch.get("po_part_id") != rule["po_part_id"]:
        return False
    if mismatch.get("sku_key") != rule["sku_key"]:
        return False
    for qty_key in ("expected_qty", "observed_qty", "delta_qty"):
        if round(_to_float(mismatch.get(qty_key)), 2) != round(_to_float(rule[qty_key]), 2):
            return False
    return True


def _classify_mismatch(mismatch: dict[str, Any]) -> dict[str, Any]:
    for rule in LINE61_ACCEPTED_SHORTAGE_RULES:
        if not _matches_rule(mismatch, rule):
            continue
        classified = dict(mismatch)
        classified.update(
            {
                "classification": "accepted_real_shortage_production_safe",
                "classification_id": LINE61_SHORTAGE_CLASSIFICATION_ID,
                "classification_reason": "owner_confirmed_po4_line61_real_shortage",
                "contract_path": LINE61_SHORTAGE_CONTRACT_PATH,
                "ordered_cargo_units": 115.0,
                "actual_received_units": 92.0,
                "shortage_units": 23.0,
                "shortage_by_size": LINE61_SHORTAGE_BY_SIZE,
                "production_authority": True,
                "clears_inbound_sheet_consistency": True,
                "clears_po_money_gate": False,
            }
        )
        return classified
    return dict(mismatch)


def validate_inbound_sheet_consistency(
    *,
    workbook_path: Path,
    tolerance: float = 0.0,
    allow_accepted_shortages_for_copied_temp: bool = False,
) -> dict[str, Any]:
    expected_by_key, expected_by_part, quantity_column = _load_inbounds(workbook_path)
    totals_by_part = _load_totals(workbook_path)
    observed_by_key, key_sources = _load_cargo_sheets(workbook_path)

    mismatches: list[dict[str, Any]] = []

    common_keys = sorted(set(expected_by_key.keys()) & set(observed_by_key.keys()))
    if expected_by_key and not common_keys:
        mismatches.append(
            {
                "type": "cargo_parse_contract",
                "po_part_id": "*",
                "sku_key": "*",
                "expected_qty": float(sum(expected_by_key.values())),
                "observed_qty": 0.0,
                "delta_qty": float(-sum(expected_by_key.values())),
                "source_sheets": [],
            }
        )

    for part_id, sku_key in common_keys:
        expected = round(expected_by_key[(part_id, sku_key)], 2)
        observed = round(observed_by_key[(part_id, sku_key)], 2)
        delta = round(observed - expected, 2)
        if abs(delta) > tolerance:
            mismatches.append(
                {
                    "type": "cargo_vs_inbounds",
                    "po_part_id": part_id,
                    "sku_key": sku_key,
                    "expected_qty": expected,
                    "observed_qty": observed,
                    "delta_qty": delta,
                    "source_sheets": sorted(key_sources.get((part_id, sku_key), set())),
                }
            )

    for part_id, total_units in sorted(totals_by_part.items()):
        expected = round(expected_by_part.get(part_id, 0.0), 2)
        observed = round(_to_float(total_units), 2)
        delta = round(observed - expected, 2)
        if abs(delta) > tolerance:
            mismatches.append(
                {
                    "type": "totals_vs_inbounds",
                    "po_part_id": part_id,
                    "sku_key": "*PART_TOTAL*",
                    "expected_qty": expected,
                    "observed_qty": observed,
                    "delta_qty": delta,
                    "source_sheets": [TOTALS_SHEET],
                }
            )

    classified_mismatches = [_classify_mismatch(mismatch) for mismatch in mismatches]
    accepted_shortages = [
        mismatch
        for mismatch in classified_mismatches
        if mismatch.get("classification_id") == LINE61_SHORTAGE_CLASSIFICATION_ID
    ]
    unknown_mismatches = [
        mismatch
        for mismatch in classified_mismatches
        if mismatch.get("classification_id") != LINE61_SHORTAGE_CLASSIFICATION_ID
    ]

    display_quantity_column = "Actual_qty" if quantity_column == "actual_qty" else ("Qty" if quantity_column == "qty" else quantity_column)

    ok = len(unknown_mismatches) == 0
    copied_temp_allowance_applied = False
    if allow_accepted_shortages_for_copied_temp and classified_mismatches and not unknown_mismatches:
        ok = True
        copied_temp_allowance_applied = True

    return {
        "ok": ok,
        "workbook_path": str(workbook_path),
        "authoritative_sheet": INBOUNDS_SHEET,
        "authoritative_quantity_column": display_quantity_column,
        "checked_keys": len(common_keys),
        "mismatch_count": len(classified_mismatches),
        "unknown_mismatch_count": len(unknown_mismatches),
        "accepted_shortage_count": len(accepted_shortages),
        "copied_temp_accepted_shortage_allowance_applied": copied_temp_allowance_applied,
        "copied_temp_accepted_shortage_allowance_note": (
            "accepted Line61 shortage is production-safe shortage truth and clears no PO money gate"
            if copied_temp_allowance_applied
            else ""
        ),
        "production_safe_accepted_shortage_applied": bool(accepted_shortages and not unknown_mismatches),
        "mismatches": classified_mismatches,
        "unknown_mismatches": unknown_mismatches,
        "accepted_shortages": accepted_shortages,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate inbound workbook consistency across sheets")
    parser.add_argument("--xlsx", type=Path, required=True, help="Inbound workbook path")
    parser.add_argument("--tolerance", type=float, default=0.0, help="Absolute qty tolerance")
    parser.add_argument("--json", action="store_true", help="Print JSON report")
    parser.add_argument(
        "--allow-accepted-shortages-for-copied-temp",
        action="store_true",
        help=(
            "Exit 0 only when every mismatch is the exact owner-confirmed Line61 "
            "accepted real shortage. Kept for compatibility; the exact shortage "
            "is now production-safe inbound truth and still does not clear PO money gate."
        ),
    )
    args = parser.parse_args()

    try:
        report = validate_inbound_sheet_consistency(
            workbook_path=args.xlsx.expanduser(),
            tolerance=max(0.0, float(args.tolerance)),
            allow_accepted_shortages_for_copied_temp=args.allow_accepted_shortages_for_copied_temp,
        )
    except Exception as exc:
        payload = {"ok": False, "error": str(exc)}
        if args.json:
            print(json.dumps(payload, ensure_ascii=False))
        else:
            print(f"ERROR: {exc}")
        return 1

    if args.json:
        print(json.dumps(report, ensure_ascii=False))
    else:
        print(
            "inbound_sheet_consistency: "
            f"ok={report['ok']} checked_keys={report['checked_keys']} mismatch_count={report['mismatch_count']}"
        )
        for mismatch in report["mismatches"]:
            classification = (
                f" classification_id={mismatch['classification_id']}"
                if mismatch.get("classification_id")
                else ""
            )
            print(
                "MISMATCH "
                f"type={mismatch['type']} po_part_id={mismatch['po_part_id']} "
                f"sku_key={mismatch['sku_key']} expected={mismatch['expected_qty']} "
                f"observed={mismatch['observed_qty']}{classification}"
            )

    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
