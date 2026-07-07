#!/usr/bin/env python3
"""Validate Phase 1 direct feeder shadow rows against the CRM workbook lane."""

from __future__ import annotations

import argparse
import csv
import json
import shutil
import sqlite3
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DB = PROJECT_ROOT / "db" / "app.db"
DEFAULT_CRM = PROJECT_ROOT / "excel_ui" / "SALES_KSP_CRM_V3.xlsx"
DEFAULT_SHEET = "SALES_KSP_CRM_1"
DEFAULT_OUTPUT_ROOT = PROJECT_ROOT / "exports" / "decommission" / "crm_excel_shadow"

sys.path.insert(0, str(PROJECT_ROOT))

from core.ingest.sales_ingest import normalize_store_code, resolve_sales_identity_detail  # noqa: E402
from core.utils.sku_normalize import normalize_sku_key, normalize_size  # noqa: E402
from scripts.feed_fact_orders_kaspi_to_sales_fact_v2 import (  # noqa: E402
    SHADOW_COLUMNS,
    UNMAPPED_COLUMNS,
    _clean_text,
    _connect_readonly,
    _date_text,
    _parse_date,
    _table_exists,
    _to_float,
    _to_int,
)


BASELINE_COLUMNS = [
    "source_lane",
    "order_id",
    "order_date",
    "planned_shipment_date",
    "status_change_date",
    "store_code",
    "sku_key",
    "sku_id",
    "my_size",
    "kaspi_offer_name",
    "quantity",
    "sell_price_kzt",
    "delivery_fee",
    "net_rev",
    "status",
    "crm_status_raw",
    "return_flag",
    "logical_dedupe_key",
    "db_unique_key",
]

ROW_DIFF_COLUMNS = [
    "severity",
    "diff_type",
    "order_id",
    "store_code",
    "kaspi_offer_name",
    "sku_key",
    "sku_id",
    "my_size",
    "field",
    "crm_value",
    "direct_value",
    "match_key_type",
    "classification",
    "evidence",
    "crm_status_change_date",
    "direct_status_updated_date",
    "direct_planned_shipment_date",
    "direct_created_date",
    "direct_internal_status",
    "direct_kaspi_status",
    "direct_kaspi_status_detail",
]

DATE_BASIS_COLUMNS = [
    "order_id",
    "store_code",
    "kaspi_offer_name",
    "sku_key",
    "sku_id",
    "my_size",
    "crm_order_date",
    "direct_order_date",
    "explained_by",
    "evidence",
]


HEADER_ALIASES = {
    "order_id": ["orderid", "order_id", "№ заказа"],
    "order_date": ["date", "order_date", "дата поступления заказа"],
    "planned_shipment_date": ["planned_shipping_date", "плановая дата передачи курьеру"],
    "kaspi_offer_name": ["kaspi_offer_name", "название товара в kaspi магазине"],
    "sku_id": ["sku_id", "skuid"],
    "sku_key": ["sku_key", "skukey"],
    "my_size": ["my_size", "mysize", "size"],
    "quantity": ["quantity", "qty", "количество"],
    "sell_price_kzt": ["sell_price_kzt", "price", "цена", "сумма"],
    "total_price": ["total_price", "totalprice"],
    "store_name": ["store_name", "storename", "store"],
    "return_flag": ["return", "return_flag"],
    "status": ["status"],
    "crm_status_raw": ["статус"],
    "status_change_date": ["дата изменения статуса", "датаизменениястатуса", "status_change_date"],
    "net_rev": ["total_net_rev", "net_rev"],
    "delivery_fee": ["delivery_fee_kzt", "delivery_fee"],
    "delivery_fee_seller": ["delivery_fee_seller", "стоимость доставки для продавца"],
    "delivery_fee_buyer": ["delivery_fee_buyer", "стоимость доставки для покупателя"],
    "product_type": ["product_type"],
}


def _canonical_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _header_indexes(headers: list[Any]) -> dict[str, int]:
    normalized = [_canonical_header(header).replace("_", " ") for header in headers]
    indexes: dict[str, int] = {}
    for target, aliases in HEADER_ALIASES.items():
        alias_norm = {_canonical_header(alias).replace("_", " ") for alias in aliases}
        for idx, header in enumerate(normalized):
            if header in alias_norm:
                indexes[target] = idx
                break
    return indexes


def _cell(row: tuple[Any, ...], indexes: dict[str, int], key: str) -> Any:
    idx = indexes.get(key)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _truthy(value: Any) -> bool:
    text = str(value or "").strip().lower()
    return text in {"1", "true", "yes", "y", "returned", "return", "возврат"}


def _num_text(value: Any) -> str:
    number = _to_float(value)
    if number is None:
        return ""
    if float(number).is_integer():
        return str(int(number))
    return str(number)


def _decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value).replace(",", "."))
    except (InvalidOperation, ValueError):
        return None


def _numbers_equal(left: Any, right: Any) -> bool:
    return _decimal(left) == _decimal(right)


def _primary_key(row: dict[str, Any]) -> tuple[str, str, str, str, str]:
    return (
        str(row.get("order_id") or ""),
        str(row.get("store_code") or ""),
        str(row.get("kaspi_offer_name") or ""),
        str(row.get("sku_key") or ""),
        str(row.get("my_size") or ""),
    )


def _fallback_key(row: dict[str, Any]) -> tuple[str, str, str, str]:
    return (
        str(row.get("order_id") or ""),
        str(row.get("store_code") or ""),
        str(row.get("kaspi_offer_name") or ""),
        str(row.get("sku_id") or ""),
    )


def _key_text(key: tuple[Any, ...]) -> str:
    return "|".join(str(part or "") for part in key)


def _read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _resolve_workbook_identity(
    conn: sqlite3.Connection,
    *,
    sku_id: str | None,
    sku_key: str | None,
    my_size: str | None,
    product_type: str | None,
    kaspi_offer_name: str,
    store_code: str,
) -> tuple[str | None, str | None, str | None]:
    normalized_key = normalize_sku_key(sku_key) if sku_key else None
    normalized_size = normalize_size(my_size, product_type=product_type) if my_size else None
    resolution = resolve_sales_identity_detail(
        conn,
        sku_id,
        normalized_key,
        normalized_size,
        kaspi_offer_name,
        store_code,
    )
    return resolution.sku_key, resolution.sku_id, resolution.my_size


def load_crm_baseline_rows(
    conn: sqlite3.Connection,
    *,
    crm_file: Path,
    sheet_name: str,
    from_date: str,
    to_date: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    start = _parse_date(from_date)
    end = _parse_date(to_date)
    if not start or not end:
        raise ValueError("from-date and to-date must be YYYY-MM-DD")

    wb = load_workbook(crm_file, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"sheet not found: {sheet_name}")
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        headers = next(rows_iter)
        indexes = _header_indexes(list(headers))
        required = {"order_id", "order_date", "kaspi_offer_name", "quantity"}
        missing = sorted(required - set(indexes))
        if missing:
            raise ValueError(f"CRM workbook missing required headers: {missing}")

        baseline: list[dict[str, Any]] = []
        unmapped: list[dict[str, Any]] = []
        for row in rows_iter:
            order_date = _parse_date(_cell(row, indexes, "order_date"))
            if not order_date or order_date < start or order_date > end:
                continue
            order_id = _clean_text(_cell(row, indexes, "order_id")) or ""
            kaspi_offer_name = _clean_text(_cell(row, indexes, "kaspi_offer_name")) or ""
            if not order_id or not kaspi_offer_name:
                continue
            store_code = normalize_store_code(_clean_text(_cell(row, indexes, "store_name")) or "UNIVERSAL")
            sku_id = _clean_text(_cell(row, indexes, "sku_id"))
            sku_key = _clean_text(_cell(row, indexes, "sku_key"))
            my_size = _clean_text(_cell(row, indexes, "my_size"))
            product_type = _clean_text(_cell(row, indexes, "product_type"))
            resolved_key, resolved_id, resolved_size = _resolve_workbook_identity(
                conn,
                sku_id=sku_id,
                sku_key=sku_key,
                my_size=my_size,
                product_type=product_type,
                kaspi_offer_name=kaspi_offer_name,
                store_code=store_code,
            )
            if not resolved_key or not resolved_id or not resolved_size:
                unmapped.append(
                    {
                        "run_id": "",
                        "source_lane": "crm_baseline",
                        "reason": "unresolved_identity",
                        "source_row_id": "",
                        "order_id": order_id,
                        "store_code": store_code,
                        "line_identity_key": "",
                        "kaspi_offer_name": kaspi_offer_name,
                        "source_sku_key": sku_key or "",
                        "source_sku_id": sku_id or "",
                        "assigned_size": "",
                        "my_size": my_size or "",
                        "final_my_size": resolved_size or "",
                        "size_source": "CRM_MY_SIZE",
                        "created_date": "",
                        "planned_shipment_date": _date_text(_cell(row, indexes, "planned_shipment_date")),
                        "actual_shipment_date": "",
                        "status_updated_date": _date_text(_cell(row, indexes, "status_change_date")),
                    }
                )
                continue
            quantity = _to_int(_cell(row, indexes, "quantity"), 1) or 1
            sell_price = _to_float(_cell(row, indexes, "sell_price_kzt"))
            if sell_price is None:
                total_price = _to_float(_cell(row, indexes, "total_price"))
                if total_price is not None and quantity:
                    sell_price = total_price / quantity
            seller_fee = _to_float(_cell(row, indexes, "delivery_fee_seller"))
            legacy_fee = _to_float(_cell(row, indexes, "delivery_fee"))
            buyer_fee = _to_float(_cell(row, indexes, "delivery_fee_buyer"))
            delivery_fee = seller_fee if seller_fee is not None else legacy_fee
            if delivery_fee is None:
                delivery_fee = buyer_fee
            return_flag = 1 if _truthy(_cell(row, indexes, "return_flag")) else 0
            status = "RETURNED" if return_flag else "DELIVERED"
            crm_status_raw = _clean_text(_cell(row, indexes, "crm_status_raw")) or ""
            status_change_date = _date_text(_cell(row, indexes, "status_change_date"))
            out = {
                "source_lane": "crm_baseline",
                "order_id": order_id,
                "order_date": order_date.isoformat(),
                "planned_shipment_date": _date_text(_cell(row, indexes, "planned_shipment_date")),
                "status_change_date": status_change_date,
                "store_code": store_code,
                "sku_key": resolved_key,
                "sku_id": resolved_id,
                "my_size": resolved_size,
                "kaspi_offer_name": kaspi_offer_name,
                "quantity": quantity,
                "sell_price_kzt": _num_text(sell_price),
                "delivery_fee": _num_text(delivery_fee),
                "net_rev": _num_text(_cell(row, indexes, "net_rev")),
                "status": status,
                "crm_status_raw": crm_status_raw,
                "return_flag": return_flag,
            }
            out["logical_dedupe_key"] = _key_text(_primary_key(out))
            out["db_unique_key"] = _key_text(_fallback_key(out))
            baseline.append(out)
        baseline.sort(key=lambda r: (r["order_date"], r["order_id"], r["store_code"], r["kaspi_offer_name"], r["sku_id"]))
        return baseline, unmapped
    finally:
        wb.close()


def _sales_fact_count(conn: sqlite3.Connection, from_date: str, to_date: str) -> int:
    if not _table_exists(conn, "sales_fact_v2"):
        return 0
    row = conn.execute(
        "SELECT COUNT(*) AS c FROM sales_fact_v2 WHERE order_date BETWEEN ? AND ?",
        (from_date, to_date),
    ).fetchone()
    return int(row["c"] or 0)


def _load_fact_order_lookup(
    conn: sqlite3.Connection,
    order_ids: set[str],
) -> dict[str, list[dict[str, Any]]]:
    if not order_ids or not _table_exists(conn, "fact_orders_kaspi"):
        return {}
    columns = {
        row["name"] for row in conn.execute("PRAGMA table_info(fact_orders_kaspi)").fetchall()
    }
    wanted = [
        "order_id",
        "store_code",
        "kaspi_offer_name",
        "created_at",
        "planned_shipment_date",
        "actual_shipment_date",
        "courier_transmission_date",
        "status_updated_at",
        "internal_status",
        "kaspi_status",
        "kaspi_status_detail",
    ]
    select_exprs = []
    for column in wanted:
        if column in columns:
            select_exprs.append(f"{column} AS {column}")
        else:
            select_exprs.append(f"'' AS {column}")
    lookup: dict[str, list[dict[str, Any]]] = defaultdict(list)
    order_list = sorted(str(order_id) for order_id in order_ids if str(order_id))
    chunk = 500
    for start in range(0, len(order_list), chunk):
        batch = order_list[start : start + chunk]
        placeholders = ",".join("?" for _ in batch)
        rows = conn.execute(
            f"""
            SELECT {', '.join(select_exprs)}
            FROM fact_orders_kaspi
            WHERE order_id IN ({placeholders})
            ORDER BY order_id, store_code, kaspi_offer_name
            """,
            batch,
        ).fetchall()
        for row in rows:
            lookup[str(row["order_id"] or "")].append(
                {
                    "order_id": _clean_text(row["order_id"]) or "",
                    "store_code": normalize_store_code(_clean_text(row["store_code"]) or "UNIVERSAL"),
                    "kaspi_offer_name": _clean_text(row["kaspi_offer_name"]) or "",
                    "created_date": _date_text(row["created_at"]),
                    "planned_shipment_date": _date_text(row["planned_shipment_date"]),
                    "actual_shipment_date": _date_text(row["actual_shipment_date"]),
                    "courier_transmission_date": _date_text(row["courier_transmission_date"]),
                    "status_updated_date": _date_text(row["status_updated_at"]),
                    "internal_status": _clean_text(row["internal_status"]) or "",
                    "kaspi_status": _clean_text(row["kaspi_status"]) or "",
                    "kaspi_status_detail": _clean_text(row["kaspi_status_detail"]) or "",
                }
            )
    return lookup


def _index_rows(rows: list[dict[str, Any]], key_func) -> dict[tuple[Any, ...], list[dict[str, Any]]]:
    indexed: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
    for idx, row in enumerate(rows):
        row["_row_index"] = str(idx)
        indexed[key_func(row)].append(row)
    return indexed


def _row_identity(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "order_id": row.get("order_id", ""),
        "store_code": row.get("store_code", ""),
        "kaspi_offer_name": row.get("kaspi_offer_name", ""),
        "sku_key": row.get("sku_key", ""),
        "sku_id": row.get("sku_id", ""),
        "my_size": row.get("my_size", ""),
    }


def _diff_row(
    *,
    severity: str,
    diff_type: str,
    row: dict[str, Any],
    direct_row: dict[str, Any] | None = None,
    field: str = "",
    crm_value: Any = "",
    direct_value: Any = "",
    match_key_type: str = "",
    classification: str = "",
    evidence: str = "",
) -> dict[str, Any]:
    evidence_row = direct_row or row
    out = {
        "severity": severity,
        "diff_type": diff_type,
        **_row_identity(row),
        "field": field,
        "crm_value": crm_value,
        "direct_value": direct_value,
        "match_key_type": match_key_type,
        "classification": classification or diff_type,
        "evidence": evidence,
        "crm_status_change_date": row.get("status_change_date", ""),
        "direct_status_updated_date": evidence_row.get("status_updated_date", ""),
        "direct_planned_shipment_date": evidence_row.get("planned_shipment_date", ""),
        "direct_created_date": evidence_row.get("created_date", ""),
        "direct_internal_status": evidence_row.get("internal_status", ""),
        "direct_kaspi_status": evidence_row.get("kaspi_status", ""),
        "direct_kaspi_status_detail": evidence_row.get("kaspi_status_detail", ""),
    }
    return out


def _classify_date_basis_diff(crm_row: dict[str, Any], direct_row: dict[str, Any]) -> str | None:
    crm_date = _parse_date(crm_row.get("order_date"))
    direct_date = _parse_date(direct_row.get("order_date"))
    if crm_date == direct_date:
        return None
    explanations = []
    for column in (
        "planned_shipment_date",
        "created_date",
        "source_order_date",
        "actual_shipment_date",
        "courier_transmission_date",
        "status_updated_date",
    ):
        candidate = _parse_date(direct_row.get(column))
        if candidate and candidate == crm_date:
            explanations.append(f"crm_date_matches_{column}")
    if explanations:
        return ",".join(explanations)
    planned_date = _parse_date(direct_row.get("planned_shipment_date"))
    created_date = _parse_date(direct_row.get("created_date"))
    status_date = _parse_date(direct_row.get("status_updated_date"))
    if planned_date and direct_date == planned_date:
        return "crm_date_is_append_date_not_planned_shipment_date"
    if created_date and direct_date == created_date:
        return "crm_date_is_append_date_not_created_date"
    if status_date and direct_date == status_date:
        return "crm_date_is_append_date_not_status_updated_date"
    return None


def _date_basis_evidence(crm_row: dict[str, Any], direct_row: dict[str, Any]) -> str:
    return (
        f"crm_date={crm_row.get('order_date', '')}; "
        f"direct_order_date={direct_row.get('order_date', '')}; "
        f"planned={direct_row.get('planned_shipment_date', '')}; "
        f"created={direct_row.get('created_date', '')}; "
        f"status_updated={direct_row.get('status_updated_date', '')}"
    )


def _loose_key(row: dict[str, Any]) -> tuple[str, str, str]:
    return (
        str(row.get("order_id") or ""),
        str(row.get("store_code") or ""),
        str(row.get("kaspi_offer_name") or ""),
    )


def _identity_evidence(crm_row: dict[str, Any], direct_row: dict[str, Any]) -> str:
    return (
        f"crm_sku_key={crm_row.get('sku_key', '')}; crm_sku_id={crm_row.get('sku_id', '')}; "
        f"crm_size={crm_row.get('my_size', '')}; direct_sku_key={direct_row.get('sku_key', '')}; "
        f"direct_sku_id={direct_row.get('sku_id', '')}; direct_size={direct_row.get('my_size', '')}; "
        f"direct_size_source={direct_row.get('final_my_size_source', '')}"
    )


def _identity_bucket(crm_row: dict[str, Any], direct_row: dict[str, Any]) -> str:
    if (
        str(crm_row.get("my_size") or "") != str(direct_row.get("my_size") or "")
        and direct_row.get("final_my_size_source") == "assigned_size"
    ):
        return "direct_size_assignment_diff"
    return "identity_resolution_delta"


def _status_fresher_evidence(crm_row: dict[str, Any], direct_row: dict[str, Any]) -> str:
    return (
        f"crm_status={crm_row.get('status', '')}; crm_raw_status={crm_row.get('crm_status_raw', '')}; "
        f"crm_status_change_date={crm_row.get('status_change_date', '')}; "
        f"direct_status={direct_row.get('status', '')}; "
        f"direct_status_updated_date={direct_row.get('status_updated_date', '')}; "
        f"direct_internal_status={direct_row.get('internal_status', '')}; "
        f"direct_kaspi_status_detail={direct_row.get('kaspi_status_detail', '')}"
    )


def _classify_status_mismatch(crm_row: dict[str, Any], direct_row: dict[str, Any]) -> tuple[str, str] | None:
    crm_status_date = _parse_date(crm_row.get("status_change_date"))
    direct_status_date = _parse_date(direct_row.get("status_updated_date"))
    direct_status = str(direct_row.get("status") or "").upper()
    if direct_status in {"CANCELLED", "RETURNED"} and (
        crm_status_date is None or (direct_status_date is not None and direct_status_date >= crm_status_date)
    ):
        return "direct_status_fresher", _status_fresher_evidence(crm_row, direct_row)
    return None


def _classify_price_mismatch(crm_row: dict[str, Any], direct_row: dict[str, Any]) -> tuple[str, str] | None:
    quantity = _to_int(direct_row.get("quantity"), 0)
    crm_price = _to_float(crm_row.get("sell_price_kzt"))
    direct_price = _to_float(direct_row.get("sell_price_kzt"))
    if quantity <= 1 or crm_price is None or direct_price is None:
        return None
    raw_direct = _to_float(direct_row.get("raw_unit_price_kzt"))
    if _numbers_equal(crm_price, direct_price * quantity):
        return (
            "crm_line_total_price_for_qty_gt1",
            f"quantity={quantity}; crm_sell_price={crm_price}; direct_unit_price={direct_price}; "
            f"direct_raw_price={raw_direct}; direct_basis={direct_row.get('sell_price_basis', '')}",
        )
    if _numbers_equal(direct_price, crm_price * quantity):
        return (
            "direct_line_total_price_for_qty_gt1",
            f"quantity={quantity}; crm_unit_price={crm_price}; direct_sell_price={direct_price}; "
            f"direct_raw_price={raw_direct}; direct_basis={direct_row.get('sell_price_basis', '')}",
        )
    return None


def _source_evidence(source_rows: list[dict[str, Any]]) -> str:
    if not source_rows:
        return "fact_orders_kaspi rows=0"
    bits = []
    for row in source_rows[:3]:
        bits.append(
            "planned={planned_shipment_date}; created={created_date}; status_updated={status_updated_date}; "
            "internal={internal_status}; kaspi_detail={kaspi_status_detail}".format(**row)
        )
    if len(source_rows) > 3:
        bits.append(f"additional_rows={len(source_rows) - 3}")
    return " | ".join(bits)


def _classify_missing_direct(
    crm_row: dict[str, Any],
    *,
    source_lookup: dict[str, list[dict[str, Any]]] | None,
    from_date: str | None,
    to_date: str | None,
) -> tuple[str, str]:
    source_rows = (source_lookup or {}).get(str(crm_row.get("order_id") or ""), [])
    if not source_rows:
        return "crm_only_no_fact_order_source", "CRM row has no fact_orders_kaspi source row for order_id"
    planned_dates = [_parse_date(row.get("planned_shipment_date")) for row in source_rows]
    start = _parse_date(from_date) if from_date else None
    end = _parse_date(to_date) if to_date else None
    if start and end and planned_dates and all((not d) or d < start or d > end for d in planned_dates):
        return "crm_append_date_outside_direct_planned_window", _source_evidence(source_rows)
    return "crm_only_append_history_or_identity_delta", _source_evidence(source_rows)


def _classify_extra_direct(direct_row: dict[str, Any]) -> tuple[str, str]:
    status = str(direct_row.get("status") or "").upper()
    evidence = (
        f"planned={direct_row.get('planned_shipment_date', '')}; "
        f"created={direct_row.get('created_date', '')}; "
        f"status_updated={direct_row.get('status_updated_date', '')}; "
        f"status={direct_row.get('status', '')}; internal={direct_row.get('internal_status', '')}; "
        f"kaspi_detail={direct_row.get('kaspi_status_detail', '')}"
    )
    if status in {"CANCELLED", "RETURNED"}:
        return "direct_terminal_status_not_booked_by_crm", evidence
    return "direct_only_no_crm_append_row", evidence


def compare_rows(
    *,
    crm_rows: list[dict[str, Any]],
    direct_rows: list[dict[str, Any]],
    direct_unmapped: list[dict[str, Any]],
    crm_unmapped: list[dict[str, Any]],
    source_lookup: dict[str, list[dict[str, Any]]] | None = None,
    from_date: str | None = None,
    to_date: str | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    diffs: list[dict[str, Any]] = []
    date_basis_rows: list[dict[str, Any]] = []
    matched_direct_indexes: set[str] = set()
    unmatched_crm_rows: list[dict[str, Any]] = []
    matched = 0
    classified_identity_pairs = 0

    direct_primary = _index_rows(direct_rows, _primary_key)
    direct_fallback = _index_rows(direct_rows, _fallback_key)

    def _crm_match_priority(row: dict[str, Any]) -> tuple[int, str, str, str]:
        candidates = direct_primary.get(_primary_key(row), []) + direct_fallback.get(_fallback_key(row), [])
        priority = 3
        for candidate in candidates:
            same_date = _parse_date(row.get("order_date")) == _parse_date(candidate.get("order_date"))
            same_price = _numbers_equal(row.get("sell_price_kzt"), candidate.get("sell_price_kzt"))
            same_qty = _to_int(row.get("quantity"), 0) == _to_int(candidate.get("quantity"), 0)
            if same_date and same_price and same_qty:
                priority = min(priority, 0)
            elif same_date and same_price:
                priority = min(priority, 1)
            elif same_date:
                priority = min(priority, 2)
        return (
            priority,
            str(row.get("order_date") or ""),
            str(row.get("order_id") or ""),
            str(row.get("kaspi_offer_name") or ""),
        )

    for key, rows in direct_primary.items():
        if len(rows) > 1:
            diffs.append(
                _diff_row(
                        severity="RED",
                        diff_type="duplicate_direct_logical_key",
                        row=rows[0],
                        direct_row=rows[0],
                        field="logical_dedupe_key",
                        direct_value=_key_text(key),
                    )
            )
    for key, rows in direct_fallback.items():
        if len(rows) > 1:
            diffs.append(
                _diff_row(
                        severity="RED",
                        diff_type="duplicate_direct_db_unique_key",
                        row=rows[0],
                        direct_row=rows[0],
                        field="db_unique_key",
                        direct_value=_key_text(key),
                    )
            )

    for crm_row in sorted(crm_rows, key=_crm_match_priority):
        direct_row = None
        match_key_type = "primary"
        primary_matches = direct_primary.get(_primary_key(crm_row), [])
        for candidate in primary_matches:
            if candidate.get("_row_index") not in matched_direct_indexes:
                direct_row = candidate
                break
        if direct_row is None:
            match_key_type = "fallback"
            fallback_matches = direct_fallback.get(_fallback_key(crm_row), [])
            for candidate in fallback_matches:
                if candidate.get("_row_index") not in matched_direct_indexes:
                    direct_row = candidate
                    break
        if direct_row is None:
            unmatched_crm_rows.append(crm_row)
            continue

        matched += 1
        matched_direct_indexes.add(str(direct_row["_row_index"]))
        field_mismatches = []
        if _to_int(crm_row.get("quantity"), 0) != _to_int(direct_row.get("quantity"), 0):
            field_mismatches.append(("quantity", crm_row.get("quantity"), direct_row.get("quantity")))
        if not _numbers_equal(crm_row.get("sell_price_kzt"), direct_row.get("sell_price_kzt")):
            field_mismatches.append(("sell_price_kzt", crm_row.get("sell_price_kzt"), direct_row.get("sell_price_kzt")))
        if str(crm_row.get("status") or "").upper() != str(direct_row.get("status") or "").upper():
            field_mismatches.append(("status", crm_row.get("status"), direct_row.get("status")))
        if str(crm_row.get("my_size") or "") != str(direct_row.get("my_size") or ""):
            field_mismatches.append(("my_size", crm_row.get("my_size"), direct_row.get("my_size")))
        if str(crm_row.get("sku_key") or "") != str(direct_row.get("sku_key") or ""):
            field_mismatches.append(("sku_key", crm_row.get("sku_key"), direct_row.get("sku_key")))
        if str(crm_row.get("sku_id") or "") != str(direct_row.get("sku_id") or ""):
            field_mismatches.append(("sku_id", crm_row.get("sku_id"), direct_row.get("sku_id")))

        date_explanation = _classify_date_basis_diff(crm_row, direct_row)
        same_date = _parse_date(crm_row.get("order_date")) == _parse_date(direct_row.get("order_date"))
        if field_mismatches:
            red_field_mismatch = False
            for field, crm_value, direct_value in field_mismatches:
                bucket = None
                evidence = ""
                if field == "status":
                    classified = _classify_status_mismatch(crm_row, direct_row)
                    if classified:
                        bucket, evidence = classified
                elif field == "sell_price_kzt":
                    classified = _classify_price_mismatch(crm_row, direct_row)
                    if classified:
                        bucket, evidence = classified
                elif field in {"my_size", "sku_key", "sku_id"}:
                    bucket = _identity_bucket(crm_row, direct_row)
                    evidence = _identity_evidence(crm_row, direct_row)
                if bucket:
                    diffs.append(
                        _diff_row(
                            severity="INFO",
                            diff_type=bucket,
                            row=crm_row,
                            direct_row=direct_row,
                            field=field,
                            crm_value=crm_value,
                            direct_value=direct_value,
                            match_key_type=match_key_type,
                            evidence=evidence,
                        )
                    )
                    continue
                red_field_mismatch = True
                diffs.append(
                    _diff_row(
                        severity="RED",
                        diff_type="field_mismatch",
                        row=crm_row,
                        direct_row=direct_row,
                        field=field,
                        crm_value=crm_value,
                        direct_value=direct_value,
                        match_key_type=match_key_type,
                    )
                )
            if not same_date and red_field_mismatch:
                diffs.append(
                    _diff_row(
                        severity="RED",
                        diff_type="date_mismatch_with_field_mismatch",
                        row=crm_row,
                        direct_row=direct_row,
                        field="order_date",
                        crm_value=crm_row.get("order_date"),
                        direct_value=direct_row.get("order_date"),
                        match_key_type=match_key_type,
                    )
                )
            elif not same_date and date_explanation:
                date_basis_rows.append(
                    {
                        **_row_identity(crm_row),
                        "crm_order_date": crm_row.get("order_date"),
                        "direct_order_date": direct_row.get("order_date"),
                        "explained_by": date_explanation,
                        "evidence": _date_basis_evidence(crm_row, direct_row),
                    }
                )
        elif not same_date:
            if date_explanation:
                date_basis_rows.append(
                    {
                        **_row_identity(crm_row),
                        "crm_order_date": crm_row.get("order_date"),
                        "direct_order_date": direct_row.get("order_date"),
                        "explained_by": date_explanation,
                        "evidence": _date_basis_evidence(crm_row, direct_row),
                    }
                )
            else:
                diffs.append(
                    _diff_row(
                        severity="RED",
                        diff_type="unclassified_date_mismatch",
                        row=crm_row,
                        direct_row=direct_row,
                        field="order_date",
                        crm_value=crm_row.get("order_date"),
                        direct_value=direct_row.get("order_date"),
                        match_key_type=match_key_type,
                    )
                )

    unmatched_direct_rows = [
        row for row in direct_rows if str(row.get("_row_index")) not in matched_direct_indexes
    ]
    direct_loose: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in unmatched_direct_rows:
        direct_loose[_loose_key(row)].append(row)

    for crm_row in unmatched_crm_rows:
        direct_row = None
        for candidate in direct_loose.get(_loose_key(crm_row), []):
            if str(candidate.get("_row_index")) not in matched_direct_indexes:
                direct_row = candidate
                break
        if direct_row is not None:
            matched_direct_indexes.add(str(direct_row.get("_row_index")))
            classified_identity_pairs += 1
            bucket = _identity_bucket(crm_row, direct_row)
            diffs.append(
                _diff_row(
                    severity="INFO",
                    diff_type=bucket,
                    row=crm_row,
                    direct_row=direct_row,
                    field="identity",
                    crm_value=crm_row.get("logical_dedupe_key"),
                    direct_value=direct_row.get("logical_dedupe_key"),
                    match_key_type="loose_order_store_offer",
                    evidence=_identity_evidence(crm_row, direct_row),
                )
            )
            continue

        bucket, evidence = _classify_missing_direct(
            crm_row,
            source_lookup=source_lookup,
            from_date=from_date,
            to_date=to_date,
        )
        diffs.append(
            _diff_row(
                severity="INFO",
                diff_type=bucket,
                row=crm_row,
                field="key",
                crm_value=crm_row.get("logical_dedupe_key"),
                direct_value="",
                match_key_type="none",
                evidence=evidence,
            )
        )

    for direct_row in direct_rows:
        if str(direct_row.get("_row_index")) in matched_direct_indexes:
            continue
        bucket, evidence = _classify_extra_direct(direct_row)
        diffs.append(
            _diff_row(
                severity="INFO",
                diff_type=bucket,
                row=direct_row,
                direct_row=direct_row,
                field="key",
                crm_value="",
                direct_value=direct_row.get("logical_dedupe_key", ""),
                match_key_type="none",
                evidence=evidence,
            )
        )

    crm_unmapped_keys = {
        (row.get("order_id", ""), row.get("store_code", ""), row.get("kaspi_offer_name", ""))
        for row in crm_unmapped
    }
    crm_mapped_keys = {_loose_key(row) for row in crm_rows}
    for row in direct_unmapped:
        key = (row.get("order_id", ""), row.get("store_code", ""), row.get("kaspi_offer_name", ""))
        if key in crm_unmapped_keys:
            bucket = "unmapped_in_both_lanes"
        elif key in crm_mapped_keys:
            bucket = "direct_unmapped_for_crm_mapped_row"
        else:
            bucket = "direct_unmapped_source_not_in_crm_selector"
        diffs.append(
            _diff_row(
                severity="INFO",
                diff_type=bucket,
                row=row,
                direct_row=row,
                field="reason",
                crm_value="",
                direct_value=row.get("reason", ""),
                match_key_type="unmapped_key",
                evidence=(
                    f"reason={row.get('reason', '')}; planned={row.get('planned_shipment_date', '')}; "
                    f"created={row.get('created_date', '')}; status_updated={row.get('status_updated_date', '')}"
                ),
            )
        )

    severity_counts = Counter(row["severity"] for row in diffs)
    summary_counts = {
        "matched": matched,
        "classified_identity_pairs": classified_identity_pairs,
        "true_mismatches": severity_counts.get("RED", 0),
        "classified_diffs": len(diffs) - severity_counts.get("RED", 0),
        "date_basis_classified": len(date_basis_rows),
    }
    return diffs, date_basis_rows, summary_counts


def _load_direct_rows(path: Path) -> list[dict[str, Any]]:
    rows = _read_csv(path)
    normalized = []
    for row in rows:
        out = {column: row.get(column, "") for column in SHADOW_COLUMNS}
        if not out.get("logical_dedupe_key"):
            out["logical_dedupe_key"] = _key_text(_primary_key(out))
        if not out.get("db_unique_key"):
            out["db_unique_key"] = _key_text(_fallback_key(out))
        normalized.append(out)
    return normalized


def _write_summary_md(path: Path, summary: dict[str, Any]) -> None:
    lines = [
        "# Direct Sales Fact Parity Summary",
        "",
        f"Gate: {summary['verdict']}",
        "",
        f"- window: {summary['from_date']}..{summary['to_date']}",
        f"- baseline_rows: {summary['baseline_rows']}",
        f"- shadow_rows: {summary['shadow_rows']}",
        f"- sales_fact_v2_rows: {summary['sales_fact_v2_rows']}",
        f"- matched: {summary['matched']}",
        f"- classified_identity_pairs: {summary.get('classified_identity_pairs', 0)}",
        f"- date_basis_classified: {summary['date_basis_classified']}",
        f"- unmapped: {summary['unmapped']}",
        f"- missing_size: {summary['missing_size']}",
        f"- classified_diffs: {summary.get('classified_diffs', 0)}",
        f"- true_mismatches: {summary['true_mismatches']}",
        f"- unexplained_mismatches: {summary.get('unexplained_mismatches', summary['true_mismatches'])}",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def validate_parity(
    *,
    db_path: Path,
    crm_file: Path,
    crm_sheet: str,
    direct_shadow: Path,
    from_date: str,
    to_date: str,
    out_dir: Path,
) -> dict[str, Any]:
    out_dir.mkdir(parents=True, exist_ok=True)
    direct_rows = _load_direct_rows(direct_shadow)
    source_unmapped_path = direct_shadow.parent / "unmapped_rows.csv"
    direct_unmapped = _read_csv(source_unmapped_path)

    with _connect_readonly(db_path) as conn:
        crm_rows, crm_unmapped = load_crm_baseline_rows(
            conn,
            crm_file=crm_file,
            sheet_name=crm_sheet,
            from_date=from_date,
            to_date=to_date,
        )
        sales_fact_rows = _sales_fact_count(conn, from_date, to_date)
        source_lookup = _load_fact_order_lookup(
            conn,
            {str(row.get("order_id") or "") for row in crm_rows},
        )

    diffs, date_basis_rows, counts = compare_rows(
        crm_rows=crm_rows,
        direct_rows=direct_rows,
        direct_unmapped=direct_unmapped,
        crm_unmapped=crm_unmapped,
        source_lookup=source_lookup,
        from_date=from_date,
        to_date=to_date,
    )

    all_unmapped = []
    for row in crm_unmapped:
        out = {column: row.get(column, "") for column in UNMAPPED_COLUMNS}
        all_unmapped.append(out)
    for row in direct_unmapped:
        out = {column: row.get(column, "") for column in UNMAPPED_COLUMNS}
        all_unmapped.append(out)

    missing_size = sum(1 for row in all_unmapped if row.get("reason") == "missing_size")
    if counts["true_mismatches"] > 0 or not crm_rows or not direct_rows:
        verdict = "RED"
    elif diffs:
        verdict = "GREEN_WITH_CLASSIFIED_DIFFS"
    elif date_basis_rows:
        verdict = "GREEN_WITH_DATE_BASIS_DIFF"
    else:
        verdict = "GREEN"

    summary = {
        "verdict": verdict,
        "from_date": from_date,
        "to_date": to_date,
        "baseline_rows": len(crm_rows),
        "shadow_rows": len(direct_rows),
        "sales_fact_v2_rows": sales_fact_rows,
        "matched": counts["matched"],
        "classified_identity_pairs": counts.get("classified_identity_pairs", 0),
        "date_basis_classified": len(date_basis_rows),
        "unmapped": len(all_unmapped),
        "missing_size": missing_size,
        "true_mismatches": counts["true_mismatches"],
        "unexplained_mismatches": counts["true_mismatches"],
        "classified_diffs": counts.get("classified_diffs", 0),
        "diff_breakdown": dict(Counter(row["diff_type"] for row in diffs)),
        "severity_breakdown": dict(Counter(row["severity"] for row in diffs)),
        "date_basis_breakdown": dict(Counter(row["explained_by"] for row in date_basis_rows)),
    }

    target_shadow = out_dir / "direct_feeder_shadow_rows.csv"
    if direct_shadow.resolve() != target_shadow.resolve():
        shutil.copyfile(direct_shadow, target_shadow)
    else:
        _write_csv(target_shadow, direct_rows, SHADOW_COLUMNS)
    _write_csv(out_dir / "crm_baseline_rows.csv", crm_rows, BASELINE_COLUMNS)
    _write_csv(out_dir / "row_diff.csv", diffs, ROW_DIFF_COLUMNS)
    _write_csv(out_dir / "date_basis_diff.csv", date_basis_rows, DATE_BASIS_COLUMNS)
    _write_csv(out_dir / "unmapped_rows.csv", all_unmapped, UNMAPPED_COLUMNS)
    (out_dir / "summary.json").write_text(
        json.dumps(summary, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    _write_summary_md(out_dir / "summary.md", summary)
    return summary


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    parser.add_argument("--crm-file", type=Path, default=DEFAULT_CRM)
    parser.add_argument("--crm-sheet", default=DEFAULT_SHEET)
    parser.add_argument("--direct-shadow", type=Path, required=True)
    parser.add_argument("--from-date", required=True)
    parser.add_argument("--to-date", required=True)
    parser.add_argument("--out-dir", type=Path)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    out_dir = args.out_dir or DEFAULT_OUTPUT_ROOT / date.today().isoformat() / "manual_parity"
    summary = validate_parity(
        db_path=args.db,
        crm_file=args.crm_file,
        crm_sheet=args.crm_sheet,
        direct_shadow=args.direct_shadow,
        from_date=args.from_date,
        to_date=args.to_date,
        out_dir=out_dir,
    )
    print(json.dumps({**summary, "out_dir": str(out_dir)}, sort_keys=True))
    green_verdicts = {"GREEN", "GREEN_WITH_DATE_BASIS_DIFF", "GREEN_WITH_CLASSIFIED_DIFFS"}
    return 0 if summary["verdict"] in green_verdicts else 1


if __name__ == "__main__":
    raise SystemExit(main())
