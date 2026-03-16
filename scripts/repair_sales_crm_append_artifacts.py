#!/usr/bin/env python3
"""
Repair missing append formulas/styles and normalize conditional-formatting ranges in CRM workbook.

Default target:
  excel_ui/SALES_KSP_CRM_V3.xlsx
  sheet: SALES_KSP_CRM_1
  table: tb_SalesRaw
"""

from __future__ import annotations

import argparse
import os
import shutil
import sys
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Dict

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from core.paths import data_path
from scripts.import_orders_to_crm import (
    CRM_SALES_SHEET_NAME,
    CRM_SALES_TABLE_NAME,
    _find_template_row_for_append,
    _formula_template_columns,
    _has_formula_payload,
    _resolve_crm_template_path,
    _restore_conditional_formatting_from_template,
    _restore_formula_and_style_window_from_template,
    _resolve_table,
    _table_bounds,
    _verify_appended_rows_integrity,
)


def _backup_file(workbook_path: Path, backup_dir: Path) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = backup_dir / f"{workbook_path.stem}.append_repair_backup_{ts}{workbook_path.suffix}"
    shutil.copy2(workbook_path, backup_path)
    return backup_path


def _repair_append_window(
    workbook_path: Path,
    sheet_name: str,
    table_name: str,
    row_from: int | None,
    row_to: int | None,
    verbose: bool = False,
) -> tuple[Dict[str, Any], Any]:
    wb = load_workbook(filename=str(workbook_path), read_only=False, data_only=False)
    try:
        ws = wb[sheet_name]
        table = _resolve_table(ws, table_name)
        tbl_start_col, tbl_start_row, tbl_end_col, tbl_end_row = _table_bounds(table)
        data_start = tbl_start_row + 1

        header_values = [
            ws.cell(row=tbl_start_row, column=col).value
            for col in range(tbl_start_col, tbl_end_col + 1)
        ]
        header_to_col = {
            str(name).strip(): tbl_start_col + i
            for i, name in enumerate(header_values)
            if str(name or "").strip()
        }
        formula_cols = _formula_template_columns(header_to_col)
        start_row = row_from or data_start
        end_row = row_to or tbl_end_row
        start_row = max(start_row, data_start)
        end_row = min(end_row, tbl_end_row)
        if end_row < start_row:
            raise RuntimeError(
                f"Requested window {row_from}-{row_to} resolves outside table rows {data_start}-{tbl_end_row}."
            )

        template_wb = None
        template_ws = None
        template_end_row = start_row - 1
        if sheet_name == CRM_SALES_SHEET_NAME and table_name == CRM_SALES_TABLE_NAME:
            resolved_template = _resolve_crm_template_path()
            if resolved_template is not None and resolved_template.resolve() != workbook_path.resolve():
                template_wb = load_workbook(filename=str(resolved_template), read_only=False, data_only=False)
                if sheet_name in template_wb.sheetnames:
                    template_ws = template_wb[sheet_name]
                    template_table = _resolve_table(template_ws, table_name)
                    _, _, _, template_end_row = _table_bounds(template_table)

        source_ws = template_ws or ws
        template_row = _find_template_row_for_append(
            ws=source_ws,
            header_row=tbl_start_row,
            table_end_row=template_end_row,
            formula_cols=formula_cols,
        )
        if template_row is None:
            raise RuntimeError(
                f"Could not find a template row with formulas before repair window {start_row}-{end_row}."
            )

        restored_formulas = 0
        restored_styles = 0
        touched_rows = max(end_row - start_row + 1, 0)
        restored_styles, restored_formulas = _restore_formula_and_style_window_from_template(
            ws=ws,
            template_ws=source_ws,
            header_row=tbl_start_row,
            target_start_row=start_row,
            target_end_row=end_row,
            start_col=tbl_start_col,
            end_col=tbl_end_col,
            formula_cols=formula_cols,
            template_data_end_row=template_end_row,
        )

        if template_ws is not None:
            cf_blocks_updated = _restore_conditional_formatting_from_template(
                ws=ws,
                template_ws=template_ws,
                header_row=tbl_start_row,
                data_end_row=tbl_end_row,
                template_data_end_row=template_end_row,
                verbose=verbose,
            )
            template_wb.close()
        else:
            cf_blocks_updated = 0

        stats: Dict[str, Any] = {
            "table_ref": table.ref,
            "table_start_row": tbl_start_row,
            "table_end_row": tbl_end_row,
            "repair_start_row": start_row,
            "repair_end_row": end_row,
            "template_row": template_row,
            "rows_touched": touched_rows,
            "restored_formulas": restored_formulas,
            "restored_styles": restored_styles,
            "cf_blocks_updated": cf_blocks_updated,
        }
        return stats, wb
    except Exception:
        wb.close()
        raise


def main() -> None:
    parser = argparse.ArgumentParser(description="Repair CRM append artifacts (formulas/styles/CF ranges)")
    parser.add_argument(
        "--workbook",
        type=Path,
        default=data_path("excel_ui", "SALES_KSP_CRM_V3.xlsx"),
        help="Workbook to repair.",
    )
    parser.add_argument("--sheet", default="SALES_KSP_CRM_1", help="Worksheet name.")
    parser.add_argument("--table", default="tb_SalesRaw", help="Excel table name.")
    parser.add_argument("--row-from", type=int, default=None, help="Repair window start row (inclusive).")
    parser.add_argument("--row-to", type=int, default=None, help="Repair window end row (inclusive).")
    parser.add_argument("--apply", action="store_true", help="Apply in-place with backup.")
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Output workbook path for dry-run/preview mode.",
    )
    parser.add_argument(
        "--backup-dir",
        type=Path,
        default=data_path("excel_ui", "backups"),
        help="Backup directory for --apply mode.",
    )
    parser.add_argument(
        "--verify",
        action=argparse.BooleanOptionalAction,
        default=True,
        help="Run append-integrity verification on repaired row window (default: on).",
    )
    parser.add_argument("--verbose", action="store_true", help="Verbose output.")
    args = parser.parse_args()

    workbook_path = args.workbook.expanduser().resolve()
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    stats, wb = _repair_append_window(
        workbook_path=workbook_path,
        sheet_name=args.sheet,
        table_name=args.table,
        row_from=args.row_from,
        row_to=args.row_to,
        verbose=args.verbose,
    )

    if args.apply:
        backup_path = _backup_file(workbook_path, args.backup_dir)
        tmp_path = workbook_path.with_name(f"{workbook_path.stem}.append_repair.tmp{workbook_path.suffix}")
        wb.save(str(tmp_path))
        wb.close()
        os.replace(str(tmp_path), str(workbook_path))
        final_path = workbook_path
        print("Append repair APPLY")
        print(f"  backup: {backup_path}")
    else:
        output_path = args.output or workbook_path.with_name(
            f"{workbook_path.stem}.append_repair_preview{workbook_path.suffix}"
        )
        wb.save(str(output_path))
        wb.close()
        final_path = output_path
        print("Append repair DRY RUN")
        print(f"  output: {output_path}")

    if args.verify:
        _verify_appended_rows_integrity(
            workbook_path=final_path,
            sheet_name=args.sheet,
            table_name=args.table,
            start_row=int(stats["repair_start_row"]),
            end_row=int(stats["repair_end_row"]),
            verbose=args.verbose,
        )

    print(f"  workbook: {final_path}")
    for key in (
        "table_ref",
        "repair_start_row",
        "repair_end_row",
        "template_row",
        "rows_touched",
        "restored_formulas",
        "restored_styles",
        "cf_blocks_updated",
    ):
        print(f"  {key}: {stats[key]}")


if __name__ == "__main__":
    main()
