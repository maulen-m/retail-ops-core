#!/usr/bin/env python3
from __future__ import annotations

import argparse
import os
import shutil
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from scripts.validate_crm_workbook_integrity import validate_workbook_integrity

CANONICAL_LINE61_SKU_KEY = "CL_NEW-CLO2_MEN_SUIT-61_BLACK"
CANONICAL_LINE61_CORE = "6в1_Черный_+Сумка"


@dataclass
class BackfillStats:
    scanned_rows: int
    line61_rows: int
    updated_rows: int
    backup_path: Optional[Path]


def _resolve_table(ws, table_name: str) -> Table:
    if table_name in ws.tables:
        return ws.tables[table_name]
    if ws.tables:
        return next(iter(ws.tables.values()))
    raise RuntimeError(f"No Excel table found on sheet '{ws.title}'")


def _table_bounds(table: Table) -> Tuple[int, int, int, int]:
    start, end = table.ref.split(":")
    start_col, start_row = coordinate_from_string(start)
    end_col, end_row = coordinate_from_string(end)
    return (
        column_index_from_string(start_col),
        int(start_row),
        column_index_from_string(end_col),
        int(end_row),
    )


def _is_line61_row(sku_key: str, sku_id_ksp: str) -> bool:
    sku_key_u = str(sku_key or "").strip().upper()
    sku_id_u = str(sku_id_ksp or "").strip().upper()
    return (
        CANONICAL_LINE61_SKU_KEY.upper() == sku_key_u
        or "SUIT-61" in sku_key_u
        or sku_id_u.startswith("OF_SUIT-61_BLK_")
    )


def _create_backup(workbook: Path, backup_dir: Path) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = backup_dir / f"{workbook.stem}_line61_core_backup_{ts}{workbook.suffix}"
    shutil.copy2(workbook, backup)
    return backup


def backfill_line61_kaspi_core(
    workbook: Path,
    sheet_name: str = "SALES_KSP_CRM_1",
    table_name: str = "tb_SalesRaw",
    apply: bool = False,
    backup_dir: Optional[Path] = None,
) -> BackfillStats:
    workbook = Path(workbook)
    backup_dir = Path(backup_dir) if backup_dir else workbook.parent / "backups"

    wb = load_workbook(filename=str(workbook), read_only=False, data_only=False)
    temp_path: Optional[Path] = None
    try:
        ws = wb[sheet_name]
        table = _resolve_table(ws, table_name)
        start_col, start_row, end_col, end_row = _table_bounds(table)

        header_to_col = {}
        for col_num in range(start_col, end_col + 1):
            header = str(ws.cell(row=start_row, column=col_num).value or "").strip()
            if header:
                header_to_col[header] = col_num

        core_col = header_to_col.get("Kaspi_name_core")
        sku_key_col = header_to_col.get("SKU_key")
        sku_id_col = header_to_col.get("SKU_ID_KSP") or header_to_col.get("Артикул")
        if not core_col:
            raise RuntimeError("Column 'Kaspi_name_core' is required in CRM table.")
        if not sku_key_col and not sku_id_col:
            raise RuntimeError("Neither 'SKU_key' nor 'SKU_ID_KSP'/'Артикул' was found.")

        scanned_rows = max(0, end_row - start_row)
        line61_rows = 0
        rows_to_update: list[int] = []
        for row_num in range(start_row + 1, end_row + 1):
            sku_key = ws.cell(row=row_num, column=sku_key_col).value if sku_key_col else ""
            sku_id = ws.cell(row=row_num, column=sku_id_col).value if sku_id_col else ""
            if not _is_line61_row(str(sku_key or ""), str(sku_id or "")):
                continue
            line61_rows += 1
            core_val = str(ws.cell(row=row_num, column=core_col).value or "").strip()
            if core_val != CANONICAL_LINE61_CORE:
                rows_to_update.append(row_num)

        if not apply or not rows_to_update:
            return BackfillStats(
                scanned_rows=scanned_rows,
                line61_rows=line61_rows,
                updated_rows=len(rows_to_update),
                backup_path=None,
            )

        backup_path = _create_backup(workbook, backup_dir)
        for row_num in rows_to_update:
            ws.cell(row=row_num, column=core_col, value=CANONICAL_LINE61_CORE)

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        temp_path = workbook.parent / f".{workbook.stem}.line61_core_{ts}.tmp{workbook.suffix}"
        wb.save(str(temp_path))
        wb.close()

        integrity = validate_workbook_integrity(temp_path)
        if integrity.errors:
            details = "; ".join(integrity.errors[:3])
            raise RuntimeError(
                "Line61 backfill produced invalid workbook package. "
                f"Sample errors: {details}"
            )

        os.replace(str(temp_path), str(workbook))
        return BackfillStats(
            scanned_rows=scanned_rows,
            line61_rows=line61_rows,
            updated_rows=len(rows_to_update),
            backup_path=backup_path,
        )
    finally:
        try:
            wb.close()
        except Exception:
            pass
        if temp_path and temp_path.exists():
            temp_path.unlink(missing_ok=True)


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Backfill historical Line61 Kaspi_name_core to canonical value."
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=Path("excel_ui/SALES_KSP_CRM_V3.xlsx"),
        help="Path to CRM workbook.",
    )
    parser.add_argument("--sheet", default="SALES_KSP_CRM_1", help="Sheet name.")
    parser.add_argument("--table", default="tb_SalesRaw", help="Excel table name.")
    parser.add_argument(
        "--backup-dir",
        type=Path,
        default=Path("excel_ui/backups"),
        help="Backup directory for apply mode.",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Apply updates in-place. Without this flag script runs in dry-run mode.",
    )
    return parser


def main() -> int:
    args = _build_parser().parse_args()
    stats = backfill_line61_kaspi_core(
        workbook=args.workbook,
        sheet_name=args.sheet,
        table_name=args.table,
        apply=bool(args.apply),
        backup_dir=args.backup_dir,
    )
    mode = "APPLY" if args.apply else "DRY RUN"
    print(f"Line61 Kaspi_name_core backfill ({mode})")
    print(f"  scanned_rows: {stats.scanned_rows}")
    print(f"  line61_rows: {stats.line61_rows}")
    print(f"  updated_rows: {stats.updated_rows}")
    if stats.backup_path:
        print(f"  backup: {stats.backup_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
