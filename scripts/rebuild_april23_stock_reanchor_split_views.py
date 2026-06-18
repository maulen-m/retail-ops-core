#!/usr/bin/env python3
"""Build April 23 stock re-anchor split views without production writes."""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
import sqlite3
import sys
from typing import Any, Iterable

import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

try:  # noqa: SIM105 - optional only for generated workbook validation.
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


ANCHOR_XLSX = Path(
    "~/Docs/Oracle/Autonomous_business/2026-04-23/"
    "092937_TASK-000_decision-grade-stock-anchor-selection-rebuild-v2-with-db-tail/"
    "answer copy/stock_anchor_selection_and_rebuild_v2_2026-04-23.xlsx"
)
MANUAL_WEBUI_CSV = (
    PROJECT_ROOT
    / "exports/validation/webui_archive_autonomous_refresh_20260526/source_refresh_runs/"
    "webui_archive_autonomous_refresh_20260526_manual_import_20260301_to_20260526/"
    "final_merged/ArchiveOrders_WEBUI_MERGED_ALL_STORES.csv"
)
LIVE_WEBUI_CSV = (
    PROJECT_ROOT
    / "exports/validation/webui_archive_autonomous_refresh_20260526/full_parse_runs/"
    "webui_archive_autonomous_refresh_20260526_live_all_enabled_20260505_to_20260525__full_parse/"
    "final_merged/ArchiveOrders_WEBUI_MERGED_ALL_STORES.csv"
)
WEBUI_CLOSEOUT = Path(
    "~/Docs/Autonomous_business_agent_handoffs/2026-05-26_webui_archive_autonomous_refresh_workflow/"
    "WEBUI_ARCHIVE_AUTONOMOUS_REFRESH_CLI_AND_UI_RESEARCH_CLOSEOUT.md"
)
DB_PATH = PROJECT_ROOT / "db/app.db"
CRM_WORKBOOK = PROJECT_ROOT / "excel_ui/SALES_KSP_CRM_V3.xlsx"
GOOGLE_CLOSEOUT_DIR = (
    PROJECT_ROOT
    / "exports/google_ops_board/workflow_runs/2026-05-25/20260525_180737_2026-05-25_closeout"
)

TARGET_SKUS = {
    "CL_NEW-CLO2_MEN_SUIT-61_BLACK": "LINE61",
    "CL_OC_MEN_LINE51_WHITE": "LINE51",
}
LINE61_SKU = "CL_NEW-CLO2_MEN_SUIT-61_BLACK"
LINE51_SKU = "CL_OC_MEN_LINE51_WHITE"
VALID_SIZES = ("S", "M", "L", "XL", "2XL", "3XL", "4XL")
SIZE_SORT = {size: idx for idx, size in enumerate(VALID_SIZES, start=1)}
DELIVERED_STATUSES = {"DELIVERED", "COMPLETED", "ВЫДАН", "ЗАВЕРШЕН"}
RETURNED_STATUSES = {"RETURNED", "RETURN", "ВОЗВРАЩЕН", "ВОЗВРАТ"}
CANCELLED_STATUSES = {"CANCELLED", "CANCELED", "ОТМЕНЕН", "ОТМЕНЁН"}


class RebuildError(RuntimeError):
    """Raised when the split rebuild contract cannot be satisfied."""


@dataclass(frozen=True)
class TargetMapping:
    family: str
    sku_key: str
    my_size: str
    sku_id: str
    mapping_rule: str
    size_source: str
    child_bundle_flag: bool
    mapping_blocker: str


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def file_evidence(path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "path": str(path.resolve()),
        "exists": path.exists(),
        "size": int(stat.st_size),
        "mtime_epoch": int(stat.st_mtime),
        "mtime_iso": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
        "sha256": sha256_file(path),
    }


def sqlite_integrity(db_path: Path) -> str:
    conn = sqlite3.connect(str(db_path))
    try:
        return str(conn.execute("PRAGMA integrity_check;").fetchone()[0])
    finally:
        conn.close()


def protected_boundary_sample() -> dict[str, Any]:
    return {
        "sampled_at": datetime.now().isoformat(timespec="seconds"),
        "db_app": file_evidence(DB_PATH),
        "crm_workbook": file_evidence(CRM_WORKBOOK),
        "db_integrity_check": sqlite_integrity(DB_PATH),
    }


def protected_boundary_stable(before: dict[str, Any], after: dict[str, Any]) -> bool:
    checks = [
        before["db_app"]["sha256"] == after["db_app"]["sha256"],
        before["db_app"]["mtime_epoch"] == after["db_app"]["mtime_epoch"],
        before["crm_workbook"]["sha256"] == after["crm_workbook"]["sha256"],
        before["crm_workbook"]["mtime_epoch"] == after["crm_workbook"]["mtime_epoch"],
        after["db_integrity_check"] == "ok",
    ]
    return all(checks)


def parse_date(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if re.match(r"^\d{4}-\d{2}-\d{2}", text):
        try:
            return date.fromisoformat(text[:10]).isoformat()
        except ValueError:
            return ""
    parsed = pd.to_datetime(text, errors="coerce", dayfirst=True)
    if pd.isna(parsed):
        return ""
    return parsed.date().isoformat()


def parse_datetime_text(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if re.match(r"^\d{4}-\d{2}-\d{2}", text):
        parsed = pd.to_datetime(text, errors="coerce")
        if pd.isna(parsed):
            return ""
        return parsed.isoformat()
    parsed = pd.to_datetime(text, errors="coerce", dayfirst=True)
    if pd.isna(parsed):
        return ""
    return parsed.isoformat()


def to_float(value: Any, default: float = 0.0) -> float:
    text = str(value or "").strip().replace(" ", "").replace(",", ".")
    if not text:
        return default
    try:
        return float(text)
    except ValueError:
        return default


def normalize_status(value: Any) -> str:
    return str(value or "").strip().upper()


def normalize_store(value: Any) -> str:
    raw = str(value or "").strip().upper()
    aliases = {
        "ACMEWEAR": "ACMEWEAR",
        "ACMEWEAR.KZ": "ACMEWEAR",
        "ACMEWEAR ": "ACMEWEAR",
        "ACMEWEAR": "ACMEWEAR",
        "STORE-B": "STOREB",
        "M GROUP": "STOREB",
        "STOREB": "STOREB",
        "UNIVERSAL": "UNIVERSAL",
        "11KZ": "11KZ",
        "MELVIS": "MELVIS",
    }
    return aliases.get(raw, raw or "UNKNOWN")


def normalize_size(value: Any) -> str:
    text = str(value or "").strip().upper()
    return text if text in VALID_SIZES else ""


def extract_size_from_text(*values: Any) -> tuple[str, str]:
    for value in values:
        text = str(value or "").upper()
        if not text:
            continue
        normalized = re.sub(r"[^A-Z0-9]+", " ", text)
        tokens = normalized.split()
        for size in ("4XL", "3XL", "2XL", "XL", "L", "M", "S"):
            if size in tokens:
                return size, str(value)
    return "", ""


def extract_size_from_sku_id(value: Any) -> str:
    text = str(value or "").upper().strip()
    for size in ("4XL", "3XL", "2XL", "XL", "L", "M", "S"):
        if text.endswith(f"_{size}") or text.endswith(f"-{size}"):
            return size
    size, _ = extract_size_from_text(text)
    return size


def build_sku_id(sku_key: str, size: str) -> str:
    return f"{sku_key}_{size}" if sku_key and size else sku_key


def first_valid_size(record: dict[str, Any], fields: Iterable[str]) -> tuple[str, str]:
    for field in fields:
        value = record.get(field, "")
        size = normalize_size(value)
        if size:
            return size, field
        if field.lower().endswith("sku_id") or field == "sku_id":
            size = extract_size_from_sku_id(value)
            if size:
                return size, field
        size, source = extract_size_from_text(value)
        if size:
            return size, field if source else ""
    return "", ""


def map_target_family(record: dict[str, Any], *, prefer_size_fields: Iterable[str] = ()) -> TargetMapping | None:
    sku_key = str(record.get("sku_key") or record.get("SKU_key") or record.get("mapped_sku_key") or "").strip()
    sku_id = str(record.get("sku_id") or record.get("SKU_ID") or record.get("mapped_sku_id") or "").strip()
    article = str(record.get("article") or record.get("Артикул") or record.get("kaspi_article") or "").strip()
    offer = str(record.get("kaspi_offer_name") or record.get("KASPI_OFFER_NAME") or record.get("offer_name") or "").strip()
    seller = str(record.get("seller_system_name") or record.get("Название в системе продавца") or "").strip()
    core = str(record.get("Kaspi_name_core") or record.get("kaspi_name_core") or "").strip()

    upper_text = " ".join([sku_key, sku_id, article, offer, seller, core]).upper()
    mapping_rule = ""
    child_bundle_flag = False
    blocker = ""

    if sku_key in TARGET_SKUS:
        mapped_sku = sku_key
        mapping_rule = "sku_key_exact"
    elif sku_id.startswith(f"{LINE61_SKU}_"):
        mapped_sku = LINE61_SKU
        mapping_rule = "sku_id_prefix"
    elif sku_id.startswith(f"{LINE51_SKU}_"):
        mapped_sku = LINE51_SKU
        mapping_rule = "sku_id_prefix"
    elif any(token in upper_text for token in ["CL_NEW-CLO2_MEN_SUIT-61_BLACK", "OF_SUIT-61", "SUIT-61", "LINE61"]):
        mapped_sku = LINE61_SKU
        mapping_rule = "text_family_line61"
    elif (
        "CL_OC_MEN_LINE51_WHITE" in upper_text
        or "LINE51" in upper_text
        or sku_key.upper().startswith("LINE-")
        or article.upper().startswith("LINE-")
        or "БЕЛИ" in upper_text
    ):
        mapped_sku = LINE51_SKU
        mapping_rule = "text_family_line51"
        child_bundle_flag = sku_key.upper().startswith("LINE-") or article.upper().startswith("LINE-")
    else:
        return None

    size_fields = list(prefer_size_fields) + [
        "my_size",
        "MY_SIZE",
        "assigned_size",
        "final_size",
        "mapped_size",
        "sku_id",
        "SKU_ID",
        "article",
        "Артикул",
        "kaspi_offer_name",
        "KASPI_OFFER_NAME",
        "seller_system_name",
    ]
    size, size_source = first_valid_size(record, size_fields)
    if not size:
        blocker = "SIZE_UNRESOLVED"
        size = "UNKNOWN"

    if child_bundle_flag:
        blocker = "; ".join(filter(None, [blocker, "CHILD_BUNDLE_COMPONENT_STOCK_SPLIT_NOT_PROVEN"]))

    return TargetMapping(
        family=TARGET_SKUS[mapped_sku],
        sku_key=mapped_sku,
        my_size=size,
        sku_id=build_sku_id(mapped_sku, size),
        mapping_rule=mapping_rule,
        size_source=size_source,
        child_bundle_flag=child_bundle_flag,
        mapping_blocker=blocker,
    )


def load_anchor(anchor_xlsx: Path) -> pd.DataFrame:
    if not anchor_xlsx.exists():
        raise RebuildError(f"anchor workbook missing: {anchor_xlsx}")
    frame = pd.read_excel(anchor_xlsx, sheet_name="Current_Stock_Rebuild", header=5, dtype=str).fillna("")
    required = {"sku_key", "my_size", "estimated_current_stock", "cogs_kzt"}
    missing = required - set(frame.columns)
    if missing:
        raise RebuildError(f"anchor workbook missing columns: {sorted(missing)}")

    rows: list[dict[str, Any]] = []
    anchor_hash = sha256_file(anchor_xlsx)
    for _, row in frame.iterrows():
        sku_key = str(row.get("sku_key") or "").strip()
        if sku_key not in TARGET_SKUS:
            continue
        size = normalize_size(row.get("my_size"))
        if not size:
            raise RebuildError(f"anchor row has unresolved size for {sku_key}: {row.to_dict()}")
        rows.append(
            {
                "family": TARGET_SKUS[sku_key],
                "sku_key": sku_key,
                "my_size": size,
                "sku_id": str(row.get("stock_id") or build_sku_id(sku_key, size)),
                "april23_anchor_qty": to_float(row.get("estimated_current_stock")),
                "raw_anchor_qty": to_float(row.get("anchor_qty")),
                "cogs_kzt": to_float(row.get("cogs_kzt")),
                "anchor_date": "2026-04-23",
                "anchor_source": "Current_Stock_Rebuild.estimated_current_stock",
                "anchor_workbook_sha256": anchor_hash,
                "anchor_risk_flags": str(row.get("risk_flags") or "").strip(),
                "anchor_confidence_band": str(row.get("confidence_band") or "").strip(),
                "anchor_confidence_score": str(row.get("confidence_score") or "").strip(),
            }
        )
    out = pd.DataFrame(rows)
    if out.empty:
        raise RebuildError("anchor workbook has no target LINE61/LINE51 rows")
    return out.sort_values(["sku_key", "my_size"], key=lambda col: col.map(lambda x: SIZE_SORT.get(x, 99) if x in VALID_SIZES else x)).reset_index(drop=True)


def load_webui_source(path: Path, source_pack: str, source_priority: int) -> pd.DataFrame:
    if not path.exists():
        raise RebuildError(f"WebUI source missing: {path}")
    frame = pd.read_csv(path, dtype=str, keep_default_na=False).fillna("")
    frame["_source_pack"] = source_pack
    frame["_source_priority"] = source_priority
    frame["_source_path"] = str(path.resolve())
    frame["_source_sha256"] = sha256_file(path)
    return frame


def build_status_change_sales_ledger_from_frames(
    frames: list[pd.DataFrame],
    *,
    anchor_date: date,
    as_of: date,
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    missing_delivered_status_change = 0
    for frame in frames:
        for _, raw in frame.iterrows():
            record = raw.to_dict()
            mapping = map_target_family(
                record,
                prefer_size_fields=("article", "seller_system_name", "kaspi_offer_name"),
            )
            if mapping is None:
                continue
            status_internal = normalize_status(record.get("status_internal") or record.get("status_raw"))
            status_change_at = parse_date(record.get("status_change_at"))
            if status_internal in DELIVERED_STATUSES and not status_change_at:
                missing_delivered_status_change += 1
                continue
            if not status_change_at:
                continue
            event_date = date.fromisoformat(status_change_at)
            if event_date <= anchor_date or event_date > as_of:
                continue

            quantity = to_float(record.get("quantity"), 1.0)
            is_delivered = status_internal in DELIVERED_STATUSES
            is_returned = status_internal in RETURNED_STATUSES
            is_cancelled = status_internal in CANCELLED_STATUSES
            date_basis = "webui_status_change_at"
            rows.append(
                {
                    "ledger_row_id": "",
                    "source_pack": record.get("_source_pack", ""),
                    "source_priority": int(record.get("_source_priority", 99)),
                    "source_path": record.get("_source_path", ""),
                    "source_sha256": record.get("_source_sha256", ""),
                    "source_file": record.get("source_file", ""),
                    "source_file_sha256": record.get("source_file_sha256", ""),
                    "source_row_number": record.get("source_row_number", ""),
                    "store_code": normalize_store(record.get("store_code")),
                    "order_id": str(record.get("order_id") or "").strip(),
                    "order_intake_date": parse_date(record.get("created_at")),
                    "ship_date": "",
                    "sale_date": status_change_at if is_delivered else "",
                    "transaction_date": status_change_at if is_delivered else "",
                    "delivered_at": status_change_at if is_delivered else "",
                    "cancel_date": status_change_at if is_cancelled else "",
                    "return_date": status_change_at if is_returned else "",
                    "status_change_at": status_change_at,
                    "status_internal": status_internal,
                    "status_raw": record.get("status_raw", ""),
                    "article": record.get("article", ""),
                    "kaspi_offer_name": record.get("kaspi_offer_name", ""),
                    "seller_system_name": record.get("seller_system_name", ""),
                    "family": mapping.family,
                    "sku_key": mapping.sku_key,
                    "my_size": mapping.my_size,
                    "sku_id": mapping.sku_id,
                    "quantity": quantity,
                    "net_rev_kzt": to_float(record.get("net_rev_kzt")),
                    "economic_final_sales_depletion_qty": quantity if is_delivered else 0.0,
                    "physical_warehouse_depletion_qty": 0.0,
                    "accepted_return_qc_addback_qty": 0.0,
                    "accepted_physical_return_addback_qty": 0.0,
                    "date_basis": date_basis if is_delivered else "webui_status_change_at_non_sale",
                    "date_vocabulary_rule": (
                        "sale_date=WebUI status_change_at for delivered/completed; "
                        "created_at retained only as order_intake_date"
                    ),
                    "mapping_rule": mapping.mapping_rule,
                    "size_source": mapping.size_source,
                    "child_bundle_flag": mapping.child_bundle_flag,
                    "mapping_blocker": mapping.mapping_blocker,
                    "row_fingerprint": record.get("row_fingerprint", ""),
                    "coverage_note": coverage_note_for_source(str(record.get("_source_pack", "")), status_change_at),
                }
            )

    ledger = pd.DataFrame(rows)
    if ledger.empty:
        ledger = pd.DataFrame(columns=status_change_ledger_columns())
        ledger.attrs["missing_delivered_status_change"] = missing_delivered_status_change
        return ledger

    ledger["_dedup_key"] = ledger.apply(
        lambda row: "|".join(
            [
                str(row.get("store_code", "")),
                str(row.get("order_id", "")),
                str(row.get("sku_key", "")),
                str(row.get("my_size", "")),
                str(row.get("status_internal", "")),
                str(row.get("status_change_at", "")),
                str(row.get("quantity", "")),
                str(row.get("row_fingerprint", "")),
            ]
        ),
        axis=1,
    )
    ledger = ledger.sort_values(["source_priority", "store_code", "order_id"]).drop_duplicates("_dedup_key", keep="first")
    ledger = ledger.drop(columns=["_dedup_key"]).sort_values(
        ["status_change_at", "store_code", "order_id", "sku_key", "my_size"]
    )
    ledger = ledger.reset_index(drop=True)
    ledger["ledger_row_id"] = [f"SC{idx:06d}" for idx in range(1, len(ledger) + 1)]
    ledger.attrs["missing_delivered_status_change"] = missing_delivered_status_change
    return ledger[status_change_ledger_columns()]


def coverage_note_for_source(source_pack: str, event_date: str) -> str:
    if source_pack.startswith("live"):
        return "live_all_enabled_webui_2026-05-05_to_2026-05-25"
    if not event_date:
        return ""
    if "manual" in source_pack and ("2026-04-24" <= event_date <= "2026-05-04"):
        return "manual_webui_only_storeb_acmewear_universal_omits_store-d_store-c"
    if "manual" in source_pack and event_date == "2026-05-26":
        return "manual_webui_2026-05-26_storeb_acmewear_universal_omits_store-d_store-c"
    return "manual_webui_storeb_acmewear_universal"


def status_change_ledger_columns() -> list[str]:
    return [
        "ledger_row_id",
        "source_pack",
        "source_priority",
        "source_path",
        "source_sha256",
        "source_file",
        "source_file_sha256",
        "source_row_number",
        "store_code",
        "order_id",
        "order_intake_date",
        "ship_date",
        "sale_date",
        "transaction_date",
        "delivered_at",
        "cancel_date",
        "return_date",
        "status_change_at",
        "status_internal",
        "status_raw",
        "article",
        "kaspi_offer_name",
        "seller_system_name",
        "family",
        "sku_key",
        "my_size",
        "sku_id",
        "quantity",
        "net_rev_kzt",
        "economic_final_sales_depletion_qty",
        "physical_warehouse_depletion_qty",
        "accepted_return_qc_addback_qty",
        "accepted_physical_return_addback_qty",
        "date_basis",
        "date_vocabulary_rule",
        "mapping_rule",
        "size_source",
        "child_bundle_flag",
        "mapping_blocker",
        "row_fingerprint",
        "coverage_note",
    ]


def load_status_change_sales_ledger(manual_csv: Path, live_csv: Path, *, anchor_date: date, as_of: date) -> pd.DataFrame:
    frames = [
        load_webui_source(live_csv, "live_all_enabled_20260505_to_20260525", 0),
        load_webui_source(manual_csv, "manual_import_20260301_to_20260526", 1),
    ]
    return build_status_change_sales_ledger_from_frames(frames, anchor_date=anchor_date, as_of=as_of)


def build_shipped_deduction_ledger_from_db(
    db_path: Path,
    *,
    anchor_date: date,
    as_of: date,
    db_sha256: str,
) -> pd.DataFrame:
    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            """
            SELECT
              order_id,
              store_code,
              kaspi_offer_name,
              sku_key,
              sku_id,
              my_size,
              assigned_size,
              quantity,
              unit_price_kzt,
              created_at,
              planned_shipment_date,
              actual_shipment_date,
              courier_transmission_date,
              kaspi_status,
              internal_status,
              kaspi_status_detail,
              waybill_url,
              waybill_number,
              source,
              source_file,
              updated_at
            FROM fact_orders_kaspi
            WHERE date(coalesce(courier_transmission_date, actual_shipment_date)) > date(?)
              AND date(coalesce(courier_transmission_date, actual_shipment_date)) <= date(?)
            ORDER BY coalesce(courier_transmission_date, actual_shipment_date), store_code, order_id
            """,
            (anchor_date.isoformat(), as_of.isoformat()),
        ).fetchall()
    finally:
        conn.close()

    out_rows: list[dict[str, Any]] = []
    for row in rows:
        record = dict(row)
        mapping = map_target_family(
            record,
            prefer_size_fields=("assigned_size", "my_size", "sku_id", "kaspi_offer_name"),
        )
        if mapping is None:
            continue
        ship_dt = parse_datetime_text(record.get("courier_transmission_date")) or parse_datetime_text(
            record.get("actual_shipment_date")
        )
        ship_date = parse_date(ship_dt)
        if not ship_date:
            continue
        quantity = to_float(record.get("quantity"), 1.0)
        sku_suffix_size = extract_size_from_sku_id(record.get("sku_id"))
        size_mismatch = bool(sku_suffix_size and mapping.my_size != "UNKNOWN" and sku_suffix_size != mapping.my_size)
        mapping_blocker = mapping.mapping_blocker
        if size_mismatch:
            mapping_blocker = "; ".join(
                filter(None, [mapping_blocker, "SIZE_SOURCE_MISMATCH_SKU_ID_VS_ASSIGNED_SIZE"])
            )
        out_rows.append(
            {
                "ledger_row_id": "",
                "source_name": "fact_orders_kaspi_readonly_courier_transmission_date",
                "source_path": str(db_path.resolve()),
                "source_sha256": db_sha256,
                "store_code": normalize_store(record.get("store_code")),
                "order_id": str(record.get("order_id") or "").strip(),
                "order_intake_date": parse_date(record.get("created_at")),
                "ship_date": ship_date,
                "ship_datetime": ship_dt,
                "ship_date_source": "courier_transmission_date"
                if str(record.get("courier_transmission_date") or "").strip()
                else "actual_shipment_date",
                "planned_shipment_date": parse_date(record.get("planned_shipment_date")),
                "status_internal_at_ship_source": normalize_status(record.get("internal_status")),
                "kaspi_status": record.get("kaspi_status", ""),
                "kaspi_status_detail": record.get("kaspi_status_detail", ""),
                "waybill_number": record.get("waybill_number", ""),
                "waybill_url_present": bool(str(record.get("waybill_url") or "").strip()),
                "family": mapping.family,
                "sku_key": mapping.sku_key,
                "my_size": mapping.my_size,
                "sku_id": mapping.sku_id,
                "quantity": quantity,
                "physical_warehouse_depletion_qty": quantity,
                "date_basis": "ship_date",
                "date_vocabulary_rule": "ship_date deducts warehouse on-hand; order_intake_date is not used as sale truth",
                "kaspi_offer_name": record.get("kaspi_offer_name", ""),
                "mapping_rule": mapping.mapping_rule,
                "size_source": mapping.size_source,
                "child_bundle_flag": mapping.child_bundle_flag,
                "mapping_blocker": mapping_blocker,
            }
        )
    return pd.DataFrame(out_rows)


def build_shipped_deduction_ledger_from_closeout(closeout_dir: Path, *, db_shipped_keys: set[str]) -> pd.DataFrame:
    snapshot_path = closeout_dir / "salesraw_snapshot.json"
    expected_path = closeout_dir / "expected_closeout_orders.json"
    if not snapshot_path.exists():
        return pd.DataFrame()

    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    expected = json.loads(expected_path.read_text(encoding="utf-8")) if expected_path.exists() else {}
    generated_at = str(expected.get("generated_at") or snapshot.get("generated_at") or "")
    target_date = str(expected.get("target_date") or snapshot.get("target_date") or "2026-05-25")
    snapshot_sha = sha256_file(snapshot_path)
    rows: list[dict[str, Any]] = []
    for record in snapshot.get("rows", []):
        mapping = map_target_family(
            record,
            prefer_size_fields=("MY_SIZE", "PROBABLE_SIZE", "KASPI_OFFER_NAME", "SKU_key"),
        )
        if mapping is None:
            continue
        quantity = to_float(record.get("Quantity"), 1.0)
        store_code = normalize_store(record.get("STORE_NAME"))
        key = shipped_dedup_key(
            store_code=store_code,
            order_id=str(record.get("OrderID") or "").strip(),
            sku_key=mapping.sku_key,
            my_size=mapping.my_size,
            quantity=quantity,
        )
        if key in db_shipped_keys:
            continue
        rows.append(
            {
                "ledger_row_id": "",
                "source_name": "google_ops_board_closeout_owner_confirmed_shipped_2026_05_25",
                "source_path": str(snapshot_path.resolve()),
                "source_sha256": snapshot_sha,
                "store_code": store_code,
                "order_id": str(record.get("OrderID") or "").strip(),
                "order_intake_date": "",
                "ship_date": target_date,
                "ship_datetime": generated_at,
                "ship_date_source": "owner_confirmed_daily_closeout_all_orders_shipped",
                "planned_shipment_date": parse_date(record.get("Date")),
                "status_internal_at_ship_source": str(record.get("Status") or ""),
                "kaspi_status": "",
                "kaspi_status_detail": "",
                "waybill_number": "",
                "waybill_url_present": False,
                "family": mapping.family,
                "sku_key": mapping.sku_key,
                "my_size": mapping.my_size,
                "sku_id": mapping.sku_id,
                "quantity": quantity,
                "physical_warehouse_depletion_qty": quantity,
                "date_basis": "ship_date",
                "date_vocabulary_rule": "owner-confirmed closeout ship_date deducts warehouse on-hand only",
                "kaspi_offer_name": record.get("KASPI_OFFER_NAME", ""),
                "mapping_rule": mapping.mapping_rule,
                "size_source": mapping.size_source,
                "child_bundle_flag": mapping.child_bundle_flag,
                "mapping_blocker": mapping.mapping_blocker,
            }
        )
    return pd.DataFrame(rows)


def shipped_dedup_key(*, store_code: str, order_id: str, sku_key: str, my_size: str, quantity: float) -> str:
    return "|".join([store_code, order_id, sku_key, my_size, f"{quantity:g}"])


def load_shipped_deduction_ledger(db_path: Path, closeout_dir: Path, *, anchor_date: date, as_of: date) -> pd.DataFrame:
    db_sha = sha256_file(db_path)
    db_rows = build_shipped_deduction_ledger_from_db(db_path, anchor_date=anchor_date, as_of=as_of, db_sha256=db_sha)
    if db_rows.empty:
        db_rows = pd.DataFrame(columns=shipped_ledger_columns())
    db_keys = {
        shipped_dedup_key(
            store_code=str(row.store_code),
            order_id=str(row.order_id),
            sku_key=str(row.sku_key),
            my_size=str(row.my_size),
            quantity=float(row.quantity),
        )
        for row in db_rows.itertuples(index=False)
    }
    closeout_rows = build_shipped_deduction_ledger_from_closeout(closeout_dir, db_shipped_keys=db_keys)
    ledger = pd.concat([db_rows, closeout_rows], ignore_index=True)
    if ledger.empty:
        return pd.DataFrame(columns=shipped_ledger_columns())
    ledger["_dedup_key"] = ledger.apply(
        lambda row: shipped_dedup_key(
            store_code=str(row.get("store_code", "")),
            order_id=str(row.get("order_id", "")),
            sku_key=str(row.get("sku_key", "")),
            my_size=str(row.get("my_size", "")),
            quantity=float(row.get("quantity", 0) or 0),
        ),
        axis=1,
    )
    ledger = ledger.drop_duplicates("_dedup_key", keep="first").drop(columns=["_dedup_key"])
    ledger = ledger.sort_values(["ship_date", "store_code", "order_id", "sku_key", "my_size"]).reset_index(drop=True)
    ledger["ledger_row_id"] = [f"SH{idx:06d}" for idx in range(1, len(ledger) + 1)]
    return ledger[shipped_ledger_columns()]


def shipped_ledger_columns() -> list[str]:
    return [
        "ledger_row_id",
        "source_name",
        "source_path",
        "source_sha256",
        "store_code",
        "order_id",
        "order_intake_date",
        "ship_date",
        "ship_datetime",
        "ship_date_source",
        "planned_shipment_date",
        "status_internal_at_ship_source",
        "kaspi_status",
        "kaspi_status_detail",
        "waybill_number",
        "waybill_url_present",
        "family",
        "sku_key",
        "my_size",
        "sku_id",
        "quantity",
        "physical_warehouse_depletion_qty",
        "date_basis",
        "date_vocabulary_rule",
        "kaspi_offer_name",
        "mapping_rule",
        "size_source",
        "child_bundle_flag",
        "mapping_blocker",
    ]


def group_qty(frame: pd.DataFrame, qty_col: str) -> pd.DataFrame:
    if frame.empty:
        return pd.DataFrame(columns=["sku_key", "my_size", qty_col])
    return (
        frame.groupby(["sku_key", "my_size"], dropna=False)[qty_col]
        .sum()
        .reset_index()
    )


def concat_blockers(values: Iterable[Any]) -> str:
    blockers: list[str] = []
    for value in values:
        for part in str(value or "").split(";"):
            clean = part.strip()
            if clean and clean not in blockers:
                blockers.append(clean)
    return "; ".join(blockers)


def build_stock_views(
    anchor_df: pd.DataFrame,
    status_ledger: pd.DataFrame,
    shipped_ledger: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    anchor = anchor_df.copy()
    sales = group_qty(status_ledger, "economic_final_sales_depletion_qty").rename(
        columns={"economic_final_sales_depletion_qty": "completed_sales_depletion_qty"}
    )
    shipped = group_qty(shipped_ledger, "physical_warehouse_depletion_qty").rename(
        columns={"physical_warehouse_depletion_qty": "shipped_sent_depletion_qty"}
    )

    returns = (
        status_ledger[status_ledger["return_date"].astype(str).str.len() > 0]
        .groupby(["sku_key", "my_size"], dropna=False)["quantity"]
        .sum()
        .reset_index()
        .rename(columns={"quantity": "returned_status_change_qty_no_qc"})
        if not status_ledger.empty
        else pd.DataFrame(columns=["sku_key", "my_size", "returned_status_change_qty_no_qc"])
    )
    cancelled = (
        status_ledger[status_ledger["cancel_date"].astype(str).str.len() > 0]
        .groupby(["sku_key", "my_size"], dropna=False)["quantity"]
        .sum()
        .reset_index()
        .rename(columns={"quantity": "cancelled_status_change_qty"})
        if not status_ledger.empty
        else pd.DataFrame(columns=["sku_key", "my_size", "cancelled_status_change_qty"])
    )
    status_blockers = (
        status_ledger.groupby(["sku_key", "my_size"], dropna=False)["mapping_blocker"]
        .apply(concat_blockers)
        .reset_index()
        .rename(columns={"mapping_blocker": "status_mapping_blockers"})
        if not status_ledger.empty
        else pd.DataFrame(columns=["sku_key", "my_size", "status_mapping_blockers"])
    )
    shipped_blockers = (
        shipped_ledger.groupby(["sku_key", "my_size"], dropna=False)["mapping_blocker"]
        .apply(concat_blockers)
        .reset_index()
        .rename(columns={"mapping_blocker": "shipped_mapping_blockers"})
        if not shipped_ledger.empty
        else pd.DataFrame(columns=["sku_key", "my_size", "shipped_mapping_blockers"])
    )

    base = anchor.merge(sales, on=["sku_key", "my_size"], how="left")
    base = base.merge(shipped, on=["sku_key", "my_size"], how="left")
    base = base.merge(returns, on=["sku_key", "my_size"], how="left")
    base = base.merge(cancelled, on=["sku_key", "my_size"], how="left")
    base = base.merge(status_blockers, on=["sku_key", "my_size"], how="left")
    base = base.merge(shipped_blockers, on=["sku_key", "my_size"], how="left")
    for col in [
        "completed_sales_depletion_qty",
        "shipped_sent_depletion_qty",
        "returned_status_change_qty_no_qc",
        "cancelled_status_change_qty",
    ]:
        base[col] = base[col].fillna(0.0).astype(float)
    for col in ["status_mapping_blockers", "shipped_mapping_blockers"]:
        base[col] = base[col].fillna("")

    base["confirmed_inbound_receipts_after_anchor_qty"] = 0.0
    base["accepted_physical_returns_after_anchor_qty"] = 0.0
    base["accepted_return_qc_stock_after_anchor_qty"] = 0.0

    physical = base.copy()
    physical["stock_view"] = "PHYSICAL_WAREHOUSE_STOCK_ESTIMATE"
    physical["date_basis"] = "ship_date"
    physical["estimated_warehouse_on_hand_qty"] = (
        physical["april23_anchor_qty"]
        + physical["confirmed_inbound_receipts_after_anchor_qty"]
        + physical["accepted_physical_returns_after_anchor_qty"]
        - physical["shipped_sent_depletion_qty"]
    )
    physical["current_stock_value_kzt_review"] = physical["estimated_warehouse_on_hand_qty"] * physical["cogs_kzt"]
    physical["negative_stock_flag"] = physical["estimated_warehouse_on_hand_qty"] < 0
    physical["retained_blockers"] = physical.apply(
        lambda row: concat_blockers(
            [
                "NO_CONFIRMED_POST_ANCHOR_INBOUND_SOURCE",
                "NO_ACCEPTED_PHYSICAL_RETURN_SOURCE",
                "2026_05_25_SHIP_DATE_USES_OWNER_CLOSEOUT_WHEN_API_COURIER_TIMESTAMP_NOT_PRESENT",
                row.get("anchor_risk_flags", ""),
                row.get("shipped_mapping_blockers", ""),
            ]
        ),
        axis=1,
    )

    economic = base.copy()
    economic["stock_view"] = "ECONOMIC_FINAL_SALES_STOCK"
    economic["date_basis"] = "sale_date_webui_status_change_at"
    economic["estimated_final_sales_stock_qty"] = (
        economic["april23_anchor_qty"]
        + economic["confirmed_inbound_receipts_after_anchor_qty"]
        + economic["accepted_return_qc_stock_after_anchor_qty"]
        - economic["completed_sales_depletion_qty"]
    )
    economic["current_stock_value_kzt_review"] = economic["estimated_final_sales_stock_qty"] * economic["cogs_kzt"]
    economic["negative_stock_flag"] = economic["estimated_final_sales_stock_qty"] < 0
    economic["retained_blockers"] = economic.apply(
        lambda row: concat_blockers(
            [
                "NO_CONFIRMED_POST_ANCHOR_INBOUND_SOURCE",
                "NO_RETURN_QC_ACCEPTANCE_SOURCE",
                "MANUAL_WEBUI_PRE_2026_05_05_AND_2026_05_26_OMITS_11KZ_MELVIS",
                row.get("anchor_risk_flags", ""),
                row.get("status_mapping_blockers", ""),
            ]
        ),
        axis=1,
    )

    exposure = base.copy()
    exposure["stock_view"] = "INVENTORY_ON_DELIVERY_EXPOSURE"
    exposure["date_basis"] = "ship_date_minus_sale_date"
    exposure["on_delivery_exposure_qty"] = (
        exposure["shipped_sent_depletion_qty"] - exposure["completed_sales_depletion_qty"]
    ).clip(lower=0)
    exposure["missing_shipped_source_gap_qty"] = (
        exposure["completed_sales_depletion_qty"] - exposure["shipped_sent_depletion_qty"]
    ).clip(lower=0)
    exposure["returned_or_cancelled_no_qc_exposure_qty"] = (
        exposure["returned_status_change_qty_no_qc"] + exposure["cancelled_status_change_qty"]
    )
    exposure["on_delivery_exposure_value_kzt_review"] = exposure["on_delivery_exposure_qty"] * exposure["cogs_kzt"]
    exposure["retained_blockers"] = exposure.apply(
        lambda row: concat_blockers(
            [
                "NO_RETURN_QC_ACCEPTANCE_SOURCE" if row["returned_status_change_qty_no_qc"] else "",
                "MISSING_SHIPPED_SOURCE_FOR_SOME_DELIVERED_ROWS"
                if row["missing_shipped_source_gap_qty"]
                else "",
                row.get("status_mapping_blockers", ""),
                row.get("shipped_mapping_blockers", ""),
            ]
        ),
        axis=1,
    )

    physical_cols = [
        "stock_view",
        "family",
        "sku_key",
        "my_size",
        "sku_id",
        "date_basis",
        "april23_anchor_qty",
        "confirmed_inbound_receipts_after_anchor_qty",
        "accepted_physical_returns_after_anchor_qty",
        "shipped_sent_depletion_qty",
        "estimated_warehouse_on_hand_qty",
        "cogs_kzt",
        "current_stock_value_kzt_review",
        "negative_stock_flag",
        "retained_blockers",
        "anchor_source",
        "anchor_workbook_sha256",
    ]
    economic_cols = [
        "stock_view",
        "family",
        "sku_key",
        "my_size",
        "sku_id",
        "date_basis",
        "april23_anchor_qty",
        "confirmed_inbound_receipts_after_anchor_qty",
        "accepted_return_qc_stock_after_anchor_qty",
        "completed_sales_depletion_qty",
        "estimated_final_sales_stock_qty",
        "cogs_kzt",
        "current_stock_value_kzt_review",
        "negative_stock_flag",
        "retained_blockers",
        "anchor_source",
        "anchor_workbook_sha256",
    ]
    exposure_cols = [
        "stock_view",
        "family",
        "sku_key",
        "my_size",
        "sku_id",
        "date_basis",
        "shipped_sent_depletion_qty",
        "completed_sales_depletion_qty",
        "on_delivery_exposure_qty",
        "missing_shipped_source_gap_qty",
        "returned_status_change_qty_no_qc",
        "cancelled_status_change_qty",
        "returned_or_cancelled_no_qc_exposure_qty",
        "cogs_kzt",
        "on_delivery_exposure_value_kzt_review",
        "retained_blockers",
    ]
    return physical[physical_cols], economic[economic_cols], exposure[exposure_cols], base


def write_dataframe_csv(frame: pd.DataFrame, path: Path) -> None:
    frame.to_csv(path, index=False, encoding="utf-8")


def write_workbooks(
    *,
    current_stock_xlsx: Path,
    business_eval_xlsx: Path,
    physical: pd.DataFrame,
    exposure: pd.DataFrame,
    economic: pd.DataFrame,
    status_ledger: pd.DataFrame,
    shipped_ledger: pd.DataFrame,
    blockers: pd.DataFrame,
    sources: pd.DataFrame,
    summary: pd.DataFrame,
) -> list[str]:
    written: list[str] = []
    with pd.ExcelWriter(current_stock_xlsx, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        physical.to_excel(writer, sheet_name="Physical_Warehouse", index=False)
        exposure.to_excel(writer, sheet_name="On_Delivery_Exposure", index=False)
        economic.to_excel(writer, sheet_name="Economic_Final_Sales", index=False)
        blockers.to_excel(writer, sheet_name="Blockers", index=False)
        sources.to_excel(writer, sheet_name="Sources", index=False)
    written.append(str(current_stock_xlsx))

    with pd.ExcelWriter(business_eval_xlsx, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        physical.to_excel(writer, sheet_name="Physical_Warehouse", index=False)
        exposure.to_excel(writer, sheet_name="On_Delivery_Exposure", index=False)
        economic.to_excel(writer, sheet_name="Economic_Final_Sales", index=False)
        status_ledger.to_excel(writer, sheet_name="Status_Change_Ledger", index=False)
        shipped_ledger.to_excel(writer, sheet_name="Shipped_Deduction_Ledger", index=False)
        blockers.to_excel(writer, sheet_name="Blockers", index=False)
        sources.to_excel(writer, sheet_name="Sources", index=False)
    written.append(str(business_eval_xlsx))

    if load_workbook is not None:
        for workbook in [current_stock_xlsx, business_eval_xlsx]:
            wb = load_workbook(workbook, read_only=True, data_only=True)
            wb.close()
    return written


def make_summary(physical: pd.DataFrame, economic: pd.DataFrame, exposure: pd.DataFrame) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "metric": "physical_estimated_warehouse_on_hand_qty",
                "value": float(physical["estimated_warehouse_on_hand_qty"].sum()),
                "date_basis": "ship_date",
            },
            {
                "metric": "economic_estimated_final_sales_stock_qty",
                "value": float(economic["estimated_final_sales_stock_qty"].sum()),
                "date_basis": "sale_date_webui_status_change_at",
            },
            {
                "metric": "inventory_on_delivery_exposure_qty",
                "value": float(exposure["on_delivery_exposure_qty"].sum()),
                "date_basis": "ship_date_minus_sale_date",
            },
            {
                "metric": "missing_shipped_source_gap_qty",
                "value": float(exposure["missing_shipped_source_gap_qty"].sum()),
                "date_basis": "diagnostic",
            },
        ]
    )


def make_blockers(physical: pd.DataFrame, economic: pd.DataFrame, exposure: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for view_name, frame, blocker_col in [
        ("PHYSICAL_WAREHOUSE_STOCK_ESTIMATE", physical, "retained_blockers"),
        ("ECONOMIC_FINAL_SALES_STOCK", economic, "retained_blockers"),
        ("INVENTORY_ON_DELIVERY_EXPOSURE", exposure, "retained_blockers"),
    ]:
        for _, row in frame.iterrows():
            for blocker in str(row.get(blocker_col) or "").split(";"):
                clean = blocker.strip()
                if clean:
                    rows.append(
                        {
                            "stock_view": view_name,
                            "sku_key": row.get("sku_key", ""),
                            "my_size": row.get("my_size", ""),
                            "blocker": clean,
                        }
                    )
    out = pd.DataFrame(rows).drop_duplicates() if rows else pd.DataFrame(columns=["stock_view", "sku_key", "my_size", "blocker"])
    return out.sort_values(["stock_view", "sku_key", "my_size", "blocker"]).reset_index(drop=True)


def make_sources(paths: dict[str, Path]) -> pd.DataFrame:
    rows = []
    for name, path in paths.items():
        rows.append({"source_name": name, **file_evidence(path)})
    return pd.DataFrame(rows)


def write_date_vocabulary_lock(path: Path, blockers: pd.DataFrame) -> None:
    blocker_count = int(len(blockers))
    text = f"""# DATE_VOCABULARY_LOCK

Generated: {datetime.now().isoformat(timespec="seconds")}

## Canonical vocabulary
- `order_intake_date`: customer order/creation date, WebUI `Дата поступления заказа`, API `createdAt` / `creationDate`. Demand/intake only.
- `ship_date`: courier handoff, waybill, Telegram PDF shipped workflow, or actual shipped workflow date. Warehouse-on-hand depletion only.
- `sale_date` / `transaction_date` / `delivered_at`: WebUI `Дата изменения статуса` for `Выдан` / delivered / completed rows. COGS, cash, PnL, sales economics, and final-sales stock only.
- `cancel_date`: status-change date for cancelled rows. Not a positive sale.
- `return_date`: status-change date for returned rows. Not active sellable stock unless a return-QC/source rule accepts it.

## Locked stock views
- `PHYSICAL_WAREHOUSE_STOCK_ESTIMATE`: April 23 owner-approved stock anchor + confirmed inbound receipts + accepted physical returns - shipped/sent orders by `ship_date`.
- `ECONOMIC_FINAL_SALES_STOCK`: April 23 owner-approved stock anchor + confirmed inbound receipts + accepted return-QC stock - completed/bought-out sales by WebUI status-change `sale_date`.
- `INVENTORY_ON_DELIVERY_EXPOSURE`: shipped/sent deductions that have not settled as final completed sales, plus missing-shipped-source diagnostics.

## Fail-closed rules
- `order_intake_date` must not be used as `sale_date`, `transaction_date`, `delivered_at`, COGS, cash, PnL, or final-sales stock depletion.
- Cancelled rows set `cancel_date` only and never create positive sales.
- Returned rows set `return_date` only and do not add stock back without accepted return-QC/source evidence.
- If shipped-source, inbound-source, return-QC, or store/date coverage is incomplete, the correct gate is `YELLOW`.

## Audited implementation surfaces
- `scripts/rebuild_april23_stock_reanchor_split_views.py` builds this copied-temp split view.
- `scripts/export_sales_archive_statusdate_mapped.py` and `scripts/validate_sales_archive_statusdate_mapped.py` already enforce the status-change cutover for sales-archive economics.
- `docs/inventory/Sales_Data_Model_V16.md`, `docs/validation/WEBUI_ARCHIVE_SINGLE_TRUTH_CONTRACT.md`, and `docs/validation/WEBUI_ARCHIVE_AUTONOMOUS_REFRESH_WORKFLOW.md` carry the durable vocabulary lock.

Retained blocker rows in this packet: {blocker_count}
"""
    path.write_text(text, encoding="utf-8")


def write_closeout(
    path: Path,
    *,
    gate_color: str,
    summary: pd.DataFrame,
    blockers: pd.DataFrame,
    boundary_stable: bool,
    output_paths: dict[str, Path],
) -> None:
    summary_rows = "\n".join(
        f"- `{row.metric}` = `{row.value:g}` ({row.date_basis})"
        for row in summary.itertuples(index=False)
    )
    blocker_preview = blockers["blocker"].drop_duplicates().head(12).tolist() if not blockers.empty else []
    blocker_lines = "\n".join(f"- `{item}`" for item in blocker_preview) or "- none"
    outputs = "\n".join(f"- `{name}`: `{path.name}`" for name, path in output_paths.items())
    text = f"""# April 23 Stock Re-Anchor Split-View Closeout

Gate: **{gate_color}**

## Result
Both required stock views were produced as copied-temp/read-only outputs:

{summary_rows}

## Retained blockers
{blocker_lines}

## Protected surfaces
- `db/app.db` unchanged during packet generation: `{str(boundary_stable)}`
- `excel_ui/SALES_KSP_CRM_V3.xlsx` unchanged during packet generation: `{str(boundary_stable)}`
- No production DB/workbook/scheduler/source-pointer/external mutation was performed by this script.

## Output files
{outputs}

## Validation note
This generated closeout records packet generation. The orchestrator should append or cite
the focused pytest/docs-lint/protected-surface final sample after running external checks.
"""
    path.write_text(text, encoding="utf-8")


def write_manifest(
    path: Path,
    *,
    gate_color: str,
    inputs: dict[str, Path],
    outputs: dict[str, Path],
    boundary_before: dict[str, Any],
    boundary_after: dict[str, Any],
    boundary_stable: bool,
    summary: pd.DataFrame,
    blockers: pd.DataFrame,
) -> None:
    manifest = {
        "run_id": path.parent.name,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "gate_color": gate_color,
        "scope": "copied-temp/read-only review outputs only",
        "production_write_authorized": False,
        "artifact_tool_note": "@oai/artifact-tool unavailable in local runtime; openpyxl fallback used for generated review-only XLSX files",
        "date_vocabulary": {
            "order_intake_date": "WebUI Дата поступления заказа / API createdAt / creationDate; demand only",
            "ship_date": "courier handoff / waybill / shipped workflow; warehouse on-hand only",
            "sale_date_transaction_date_delivered_at": "WebUI Дата изменения статуса for delivered/completed rows; COGS/cash/PnL/final-sales stock",
            "cancel_date": "status-change date for cancelled rows; not positive sale",
            "return_date": "status-change date for returned rows; addback only with return-QC/source acceptance",
        },
        "inputs": {name: file_evidence(path) for name, path in inputs.items() if path.exists()},
        "outputs": {name: file_evidence(path) for name, path in outputs.items() if path.exists()},
        "protected_boundary": {
            "before_generation": boundary_before,
            "after_generation": boundary_after,
            "stable": boundary_stable,
        },
        "summary": summary.to_dict(orient="records"),
        "retained_blocker_count": int(len(blockers)),
        "retained_blockers": sorted(blockers["blocker"].drop_duplicates().tolist()) if not blockers.empty else [],
        "red_stopline_checks": {
            "order_intake_used_as_final_sale_truth": False,
            "protected_surface_drift_during_generation": not boundary_stable,
        },
    }
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def rebuild_split_views(
    *,
    anchor_xlsx: Path,
    manual_webui_csv: Path,
    live_webui_csv: Path,
    webui_closeout: Path,
    db_path: Path,
    google_closeout_dir: Path,
    output_root: Path,
    run_id: str,
    as_of: date,
) -> dict[str, Any]:
    output_dir = output_root / run_id
    output_dir.mkdir(parents=True, exist_ok=False)

    boundary_before = protected_boundary_sample()
    anchor_df = load_anchor(anchor_xlsx)
    status_ledger = load_status_change_sales_ledger(
        manual_webui_csv,
        live_webui_csv,
        anchor_date=date(2026, 4, 23),
        as_of=as_of,
    )
    shipped_ledger = load_shipped_deduction_ledger(
        db_path,
        google_closeout_dir,
        anchor_date=date(2026, 4, 23),
        as_of=as_of,
    )
    physical, economic, exposure, _base = build_stock_views(anchor_df, status_ledger, shipped_ledger)
    summary = make_summary(physical, economic, exposure)
    blockers = make_blockers(physical, economic, exposure)
    sources = make_sources(
        {
            "anchor_workbook": anchor_xlsx,
            "manual_webui_csv": manual_webui_csv,
            "live_webui_csv": live_webui_csv,
            "webui_closeout": webui_closeout,
            "db_app_readonly": db_path,
            "crm_workbook_protected_surface": CRM_WORKBOOK,
            "google_ops_closeout_salesraw_snapshot": google_closeout_dir / "salesraw_snapshot.json",
            "google_ops_closeout_expected_orders": google_closeout_dir / "expected_closeout_orders.json",
        }
    )

    outputs = {
        "date_vocabulary_lock": output_dir / "DATE_VOCABULARY_LOCK.md",
        "physical_warehouse_stock_estimate_by_sku_size": output_dir / "physical_warehouse_stock_estimate_by_sku_size.csv",
        "economic_final_sales_stock_by_sku_size": output_dir / "economic_final_sales_stock_by_sku_size.csv",
        "inventory_on_delivery_exposure_by_sku_size": output_dir / "inventory_on_delivery_exposure_by_sku_size.csv",
        "status_change_sales_ledger": output_dir / "status_change_sales_ledger.csv",
        "shipped_deduction_ledger": output_dir / "shipped_deduction_ledger.csv",
        "current_stock_valuation": output_dir / "current_stock_valuation.xlsx",
        "inventory_business_eval_tables": output_dir / "inventory_business_eval_tables.xlsx",
        "evidence_manifest": output_dir / "evidence_manifest.json",
        "closeout": output_dir / "closeout.md",
        "retained_blockers": output_dir / "retained_blockers.csv",
        "source_evidence": output_dir / "source_evidence.csv",
        "summary": output_dir / "summary.csv",
    }

    write_dataframe_csv(physical, outputs["physical_warehouse_stock_estimate_by_sku_size"])
    write_dataframe_csv(economic, outputs["economic_final_sales_stock_by_sku_size"])
    write_dataframe_csv(exposure, outputs["inventory_on_delivery_exposure_by_sku_size"])
    write_dataframe_csv(status_ledger, outputs["status_change_sales_ledger"])
    write_dataframe_csv(shipped_ledger, outputs["shipped_deduction_ledger"])
    write_dataframe_csv(blockers, outputs["retained_blockers"])
    write_dataframe_csv(sources, outputs["source_evidence"])
    write_dataframe_csv(summary, outputs["summary"])
    write_date_vocabulary_lock(outputs["date_vocabulary_lock"], blockers)
    write_workbooks(
        current_stock_xlsx=outputs["current_stock_valuation"],
        business_eval_xlsx=outputs["inventory_business_eval_tables"],
        physical=physical,
        exposure=exposure,
        economic=economic,
        status_ledger=status_ledger,
        shipped_ledger=shipped_ledger,
        blockers=blockers,
        sources=sources,
        summary=summary,
    )

    boundary_after = protected_boundary_sample()
    boundary_stable = protected_boundary_stable(boundary_before, boundary_after)
    gate_color = "RED" if not boundary_stable else ("YELLOW" if not blockers.empty else "GREEN")
    write_closeout(
        outputs["closeout"],
        gate_color=gate_color,
        summary=summary,
        blockers=blockers,
        boundary_stable=boundary_stable,
        output_paths=outputs,
    )
    write_manifest(
        outputs["evidence_manifest"],
        gate_color=gate_color,
        inputs={
            "anchor_workbook": anchor_xlsx,
            "manual_webui_csv": manual_webui_csv,
            "live_webui_csv": live_webui_csv,
            "webui_closeout": webui_closeout,
            "db_app_readonly": db_path,
            "google_closeout_salesraw_snapshot": google_closeout_dir / "salesraw_snapshot.json",
            "google_closeout_expected_orders": google_closeout_dir / "expected_closeout_orders.json",
        },
        outputs=outputs,
        boundary_before=boundary_before,
        boundary_after=boundary_after,
        boundary_stable=boundary_stable,
        summary=summary,
        blockers=blockers,
    )
    return {
        "output_dir": str(output_dir.resolve()),
        "gate_color": gate_color,
        "summary": summary.to_dict(orient="records"),
        "blocker_count": int(len(blockers)),
        "outputs": {name: str(path.resolve()) for name, path in outputs.items()},
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--anchor-xlsx", type=Path, default=ANCHOR_XLSX)
    parser.add_argument("--manual-webui-csv", type=Path, default=MANUAL_WEBUI_CSV)
    parser.add_argument("--live-webui-csv", type=Path, default=LIVE_WEBUI_CSV)
    parser.add_argument("--webui-closeout", type=Path, default=WEBUI_CLOSEOUT)
    parser.add_argument("--db-path", type=Path, default=DB_PATH)
    parser.add_argument("--google-closeout-dir", type=Path, default=GOOGLE_CLOSEOUT_DIR)
    parser.add_argument("--output-root", type=Path, default=PROJECT_ROOT / "exports/validation")
    parser.add_argument("--run-id", default="")
    parser.add_argument("--as-of", default="2026-05-26")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    run_id = args.run_id or f"april23_stock_reanchor_statusdate_rebuild_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    result = rebuild_split_views(
        anchor_xlsx=args.anchor_xlsx.expanduser().resolve(),
        manual_webui_csv=args.manual_webui_csv.expanduser().resolve(),
        live_webui_csv=args.live_webui_csv.expanduser().resolve(),
        webui_closeout=args.webui_closeout.expanduser().resolve(),
        db_path=args.db_path.expanduser().resolve(),
        google_closeout_dir=args.google_closeout_dir.expanduser().resolve(),
        output_root=args.output_root.expanduser().resolve(),
        run_id=run_id,
        as_of=date.fromisoformat(args.as_of),
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
