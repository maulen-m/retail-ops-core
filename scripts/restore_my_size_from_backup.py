#!/usr/bin/env python3
"""
Restore MY_SIZE values in CRM workbook from a trusted backup workbook.

Safety defaults:
- Dry-run by default
- Writes only when --apply is provided
- Creates timestamped backup before writing
"""
from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Tuple

import pandas as pd
from openpyxl import load_workbook


DEFAULT_SOURCE = Path("excel_ui/backups/CRM_backup_20260206_160719.xlsx")
DEFAULT_TARGET = Path("excel_ui/SALES_KSP_CRM_V3.xlsx")
DEFAULT_SHEET = "SALES_KSP_CRM_1"
DEFAULT_REPORT = Path("logs/my_size_restore_report.csv")
DEFAULT_BACKUP_DIR = Path("excel_ui/backups")

KEY_COLUMNS = [
    "№ заказа",
    "Date",
    "Артикул",
    "Название товара в Kaspi Магазине",
    "Количество",
]


def _norm_text(value: Any) -> str:
    text = str(value or "").strip()
    return " ".join(text.split())


def _norm_order_id(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    try:
        num = float(text)
        if num.is_integer():
            return str(int(num))
    except Exception:
        pass
    return text


def _norm_date(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return ""
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return text


def _norm_qty(value: Any) -> str:
    if value is None:
        return "0"
    text = str(value).strip()
    if not text:
        return "0"
    try:
        return str(int(float(text)))
    except Exception:
        return text


def build_row_key(row: Dict[str, Any]) -> str:
    return "|".join(
        [
            _norm_order_id(row.get("№ заказа")),
            _norm_date(row.get("Date")),
            _norm_text(row.get("Артикул")).lower(),
            _norm_text(row.get("Название товара в Kaspi Магазине")).lower(),
            _norm_qty(row.get("Количество")),
        ]
    )


def _valid_size(value: Any) -> bool:
    text = _norm_text(value).upper()
    if not text:
        return False
    valid = {
        "XS",
        "S",
        "M",
        "L",
        "XL",
        "2XL",
        "3XL",
        "4XL",
        "5XL",
        "22",
        "24",
        "26",
        "28",
        "30",
        "ONE_SIZE",
    }
    return text in valid


def compute_restore_updates(
    source_df: pd.DataFrame,
    target_df: pd.DataFrame,
    mode: str = "strict",
) -> Tuple[List[Dict[str, Any]], Dict[str, int]]:
    source_map: Dict[str, str] = {}
    for row in source_df.to_dict(orient="records"):
        key = build_row_key(row)
        if not key:
            continue
        source_map[key] = _norm_text(row.get("MY_SIZE"))

    updates: List[Dict[str, Any]] = []
    stats = {
        "source_rows": len(source_df),
        "target_rows": len(target_df),
        "matched_rows": 0,
        "updated_rows": 0,
        "unchanged_rows": 0,
        "unmatched_rows": 0,
    }

    for row_idx, row in enumerate(target_df.to_dict(orient="records"), start=2):
        key = build_row_key(row)
        if key not in source_map:
            stats["unmatched_rows"] += 1
            continue

        stats["matched_rows"] += 1
        old_size = _norm_text(row.get("MY_SIZE"))
        new_size = source_map[key]

        should_write = old_size != new_size
        if mode == "fill-missing":
            should_write = (not _valid_size(old_size)) and bool(new_size)

        if should_write:
            updates.append(
                {
                    "row_number": row_idx,
                    "row_key": key,
                    "old_my_size": old_size,
                    "new_my_size": new_size,
                    "order_id": _norm_order_id(row.get("№ заказа")),
                }
            )
            stats["updated_rows"] += 1
        else:
            stats["unchanged_rows"] += 1

    return updates, stats


def _backup_file(target: Path, backup_dir: Path) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = backup_dir / f"CRM_backup_my_size_restore_{ts}.xlsx"
    out.write_bytes(target.read_bytes())
    return out


def _load_sheet_df(path: Path, sheet: str) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet)


def _apply_updates_to_workbook(
    target_path: Path,
    sheet_name: str,
    updates: List[Dict[str, Any]],
) -> int:
    if not updates:
        return 0

    wb = load_workbook(target_path)
    try:
        sh = wb[sheet_name]
        headers = [sh.cell(row=1, column=c).value for c in range(1, sh.max_column + 1)]
        header_to_col = {str(h).strip(): i + 1 for i, h in enumerate(headers) if str(h or "").strip()}
        my_size_col = header_to_col.get("MY_SIZE")
        if not my_size_col:
            raise RuntimeError("MY_SIZE column not found in target sheet")

        for upd in updates:
            sh.cell(row=upd["row_number"], column=my_size_col).value = upd["new_my_size"] or ""

        wb.save(target_path)
    finally:
        wb.close()
    return len(updates)


def _write_report(path: Path, stats: Dict[str, int], updates: List[Dict[str, Any]], dry_run: bool, backup_path: Path | None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["metric", "value"])
        writer.writerow(["dry_run", int(dry_run)])
        writer.writerow(["backup_path", str(backup_path) if backup_path else ""])
        for k, v in stats.items():
            writer.writerow([k, v])

        writer.writerow([])
        writer.writerow(["row_number", "order_id", "old_my_size", "new_my_size", "row_key"])
        for u in updates:
            writer.writerow([u["row_number"], u["order_id"], u["old_my_size"], u["new_my_size"], u["row_key"]])


@dataclass
class RunResult:
    stats: Dict[str, int]
    updates: List[Dict[str, Any]]
    backup_path: Path | None


def run_restore(
    source_xlsx: Path,
    target_xlsx: Path,
    sheet_name: str,
    mode: str,
    apply: bool,
    report_path: Path,
    backup_dir: Path,
) -> RunResult:
    source_df = _load_sheet_df(source_xlsx, sheet_name)
    target_df = _load_sheet_df(target_xlsx, sheet_name)
    updates, stats = compute_restore_updates(source_df, target_df, mode=mode)

    backup_path = None
    if apply:
        backup_path = _backup_file(target_xlsx, backup_dir)
        _apply_updates_to_workbook(target_xlsx, sheet_name, updates)

    _write_report(report_path, stats, updates, dry_run=not apply, backup_path=backup_path)
    return RunResult(stats=stats, updates=updates, backup_path=backup_path)


def _build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Restore MY_SIZE values from backup workbook")
    p.add_argument("--source-xlsx", type=Path, default=DEFAULT_SOURCE)
    p.add_argument("--target-xlsx", type=Path, default=DEFAULT_TARGET)
    p.add_argument("--sheet", default=DEFAULT_SHEET)
    p.add_argument("--mode", choices=["strict", "fill-missing"], default="strict")
    p.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    p.add_argument("--backup-dir", type=Path, default=DEFAULT_BACKUP_DIR)
    p.add_argument("--apply", action="store_true", help="Apply writes to target workbook")
    return p


def main() -> None:
    args = _build_parser().parse_args()
    result = run_restore(
        source_xlsx=args.source_xlsx,
        target_xlsx=args.target_xlsx,
        sheet_name=args.sheet,
        mode=args.mode,
        apply=args.apply,
        report_path=args.report,
        backup_dir=args.backup_dir,
    )

    mode = "APPLY" if args.apply else "DRY RUN"
    print(f"MY_SIZE restore {mode}")
    print(f"  Source: {args.source_xlsx}")
    print(f"  Target: {args.target_xlsx}")
    print(f"  Matched: {result.stats['matched_rows']}")
    print(f"  Updated: {result.stats['updated_rows']}")
    print(f"  Report: {args.report}")
    if result.backup_path:
        print(f"  Backup: {result.backup_path}")


if __name__ == "__main__":
    main()
