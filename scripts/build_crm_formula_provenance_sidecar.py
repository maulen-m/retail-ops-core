#!/usr/bin/env python3
"""Build a PII-excluding CRM formula-provenance sidecar, read-only.

The source workbook is opened read-only in both formula and cached-value mode.
Only allowlisted operational fields and one-way row hashes are emitted; phone,
address, customer name, and other customer fields never enter the artifacts.
"""

from __future__ import annotations

import argparse
import json
import re
import sqlite3
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in __import__("sys").path:
    __import__("sys").path.insert(0, str(PROJECT_ROOT))

from core.calc.economics import calc_net_rev  # noqa: E402
from scripts.build_sales_formula_provenance_sidecar import (  # noqa: E402
    MONEY,
    ProvenanceError,
    _canonical_hash,
    _clean,
    _db_row_payload,
    _decimal,
    _decimal_text,
    _json_bytes,
    _jsonl_bytes,
    _load_jsonl,
    _money,
    _money_text,
    _parse_date,
    _policy_manifest,
    _selected_db_rows,
    _selected_view_rows,
    _sha256_bytes,
    _upper,
    _write_atomic,
    sha256_file,
)


BUILD_VERSION = "sales_formula_provenance_crm_v4"
DEFAULT_SHEET = "SALES_KSP_CRM_1"
CRM_SOURCE_FILE = "SALES_KSP_CRM_V3.xlsx"
OBSERVATION_POSITIVE_STATUSES = {"DELIVERED", "COMPLETED"}
OBSERVATION_SOURCES = {"API", "WEBUI"}
ORDER_POSITIVE_INTERNAL = "COMPLETED"
ORDER_POSITIVE_KASPI = "ARCHIVE"
ORDER_SOURCES = {"API", "KASPI_ARCHIVE", "ACTIVEORDERS_ENRICH"}
NEGATIVE_STAGES = {
    "CANCELLED",
    "RETURNED",
    "RETURN",
    "REFUND",
    "CANCELLED_AFTER_DELIVERY",
    "CANCELLED_DELIVERED",
}
RAW_LITERAL_HEADERS = (
    "№ заказа",
    "Дата поступления заказа",
    "Сумма",
    "Количество",
    "Стоимость доставки для продавца",
)
REQUIRED_HEADERS = (
    "Date",
    "STORE_NAME",
    "Quantity",
    "OrderID",
    "MY_SIZE",
    "SKU_key",
    "SKU_ID",
    "Sell_price_kzt",
    "Delivery_fee_kzt",
    "№ заказа",
    "Дата поступления заказа",
    "Артикул",
    "Сумма",
    "Статус",
    "Причина отмены",
    "Количество",
    "Стоимость доставки для продавца",
)


def _store_code(value: Any) -> str:
    token = _upper(value).replace("-", "").replace(" ", "")
    aliases = {
        "UNIVERSAL": "UNIVERSAL",
        "ACMEWEAR": "ACMEWEAR",
        "STOREB": "STOREB",
        "11KZ": "11KZ",
        "MELVIS": "MELVIS",
    }
    return aliases.get(token, token)


def _canonical_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return _decimal_text(Decimal(str(value)))
    if type(value).__module__.startswith("openpyxl.worksheet.formula"):
        return {
            "openpyxl_formula_type": type(value).__name__,
            "attributes": {
                str(key): _canonical_cell(item)
                for key, item in sorted(getattr(value, "__dict__", {}).items())
            },
        }
    return str(value)


def _is_formula_cell(value: Any) -> bool:
    return bool(
        (isinstance(value, str) and value.lstrip().startswith("="))
        or type(value).__module__.startswith("openpyxl.worksheet.formula")
    )


def _date_text(value: Any, *, epoch: Any = None) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float, Decimal)):
        converted = from_excel(value, epoch=epoch)
        if isinstance(converted, datetime):
            return converted.date().isoformat()
        if isinstance(converted, date):
            return converted.isoformat()
        raise ProvenanceError(f"invalid Excel source date: {value!r}")
    return _parse_date(value).isoformat()


def _optional_decimal_text(value: Any, *, field: str) -> str | None:
    if not _clean(value):
        return None
    return _decimal_text(_decimal(value, field=field))


def _timestamp(value: Any, *, field: str) -> datetime:
    text = _clean(value)
    if not text:
        raise ProvenanceError(f"missing lifecycle timestamp: {field}")
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ProvenanceError(f"invalid lifecycle timestamp {field}: {text!r}") from exc
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def _read_workbook_rows(path: Path, sheet: str) -> tuple[list[dict[str, Any]], list[str]]:
    formula_wb = load_workbook(path, read_only=True, data_only=False)
    cached_wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in formula_wb.sheetnames or sheet not in cached_wb.sheetnames:
            raise ProvenanceError(f"CRM sheet missing: {sheet}")
        formula_ws = formula_wb[sheet]
        cached_ws = cached_wb[sheet]
        formula_headers = [cell.value for cell in next(formula_ws.iter_rows(min_row=1, max_row=1))]
        cached_headers = [cell.value for cell in next(cached_ws.iter_rows(min_row=1, max_row=1))]
        if formula_headers != cached_headers:
            raise ProvenanceError("formula/cached workbook headers differ")
        headers = [_clean(value) for value in formula_headers]
        missing = [header for header in REQUIRED_HEADERS if header not in headers]
        if missing:
            raise ProvenanceError(f"CRM source missing required headers: {missing}")
        index = {header: headers.index(header) for header in REQUIRED_HEADERS}

        formula_iter = formula_ws.iter_rows(min_row=2, values_only=True)
        cached_iter = cached_ws.iter_rows(min_row=2, values_only=True)
        rows: list[dict[str, Any]] = []
        physical_row = 1
        while True:
            try:
                formula_values = next(formula_iter)
                formula_done = False
            except StopIteration:
                formula_values = ()
                formula_done = True
            try:
                cached_values = next(cached_iter)
                cached_done = False
            except StopIteration:
                cached_values = ()
                cached_done = True
            if formula_done and cached_done:
                break
            if formula_done != cached_done:
                raise ProvenanceError("formula/cached workbook row counts differ")
            physical_row += 1
            width = max(len(formula_values), len(cached_values), len(headers))
            formula_values = tuple(formula_values) + (None,) * (width - len(formula_values))
            cached_values = tuple(cached_values) + (None,) * (width - len(cached_values))

            def cached(header: str) -> Any:
                return cached_values[index[header]]

            def formula(header: str) -> Any:
                return formula_values[index[header]]

            derived_order_id = _clean(cached("OrderID"))
            raw_order_id = _clean(cached("№ заказа"))
            if not derived_order_id and not raw_order_id:
                continue
            formula_row_sha256 = _canonical_hash([_canonical_cell(value) for value in formula_values])
            safe = {
                "physical_row": physical_row,
                "formula_row_sha256": formula_row_sha256,
                "derived_order_id": derived_order_id,
                "raw_order_id": raw_order_id,
                "store_code": _store_code(cached("STORE_NAME")),
                "order_date": (
                    _date_text(cached("Date"), epoch=cached_wb.epoch)
                    if _clean(cached("Date"))
                    else ""
                ),
                "raw_order_received_date": (
                    _date_text(cached("Дата поступления заказа"), epoch=cached_wb.epoch)
                    if _clean(cached("Дата поступления заказа"))
                    else ""
                ),
                "sku_key": _clean(cached("SKU_key")),
                "sku_id": _clean(cached("SKU_ID")),
                "size": _upper(cached("MY_SIZE")),
                "derived_quantity": _optional_decimal_text(cached("Quantity"), field="crm.derived_quantity"),
                "raw_quantity": _optional_decimal_text(cached("Количество"), field="crm.raw_quantity"),
                "derived_sell_price_kzt": _optional_decimal_text(
                    cached("Sell_price_kzt"), field="crm.derived_sell_price"
                ),
                "legacy_delivery_fee_kzt": _optional_decimal_text(
                    cached("Delivery_fee_kzt"), field="crm.legacy_delivery_fee"
                ),
                "raw_gross_kzt": _optional_decimal_text(cached("Сумма"), field="crm.raw_gross"),
                "raw_seller_delivery_fee_kzt": _optional_decimal_text(
                    cached("Стоимость доставки для продавца"), field="crm.raw_seller_fee"
                ),
                "raw_article": _clean(cached("Артикул")),
                "raw_status": _clean(cached("Статус")),
                "raw_cancellation_reason": _clean(cached("Причина отмены")),
                "raw_formula_headers": [
                    header
                    for header in RAW_LITERAL_HEADERS
                    if _is_formula_cell(formula(header))
                ],
            }
            safe["allowlisted_row_sha256"] = _canonical_hash(safe)
            rows.append(safe)
    finally:
        formula_wb.close()
        cached_wb.close()
    return rows, headers


def _row_hash(row: sqlite3.Row | dict[str, Any], *, source_table: str) -> str:
    payload = dict(row)
    return _canonical_hash({"source_table": source_table, "row": payload})


def _lifecycle_evidence(
    conn: sqlite3.Connection,
    *,
    order_id: str,
    store_code: str,
) -> tuple[list[dict[str, Any]], list[str]]:
    reasons: list[str] = []
    positives: list[dict[str, Any]] = []
    positive_times: list[datetime] = []
    negative_times: list[datetime] = []

    observations = conn.execute(
        """
        SELECT id, order_id, store_code, status_internal, observed_at, source,
               ledger_run_id, source_detail, created_at
        FROM fact_order_status_observations
        WHERE order_id=?
          AND store_code=?
        ORDER BY datetime(observed_at) DESC, id DESC
        """,
        (order_id, store_code),
    ).fetchall()
    if observations:
        latest = observations[0]
        latest_status = _upper(latest["status_internal"])
        latest_source = _upper(latest["source"])
        if latest_status in OBSERVATION_POSITIVE_STATUSES and latest_source in OBSERVATION_SOURCES:
            observed_dt = _timestamp(
                latest["observed_at"], field="fact_order_status_observations.observed_at"
            )
            positives.append(
                {
                    "source_table": "fact_order_status_observations",
                    "row_id": int(latest["id"]),
                    "status_internal": _upper(latest["status_internal"]),
                    "source": _upper(latest["source"]),
                    "observed_at": _clean(latest["observed_at"]),
                    "ledger_run_id": _clean(latest["ledger_run_id"]),
                    "source_detail_sha256": _sha256_bytes(_clean(latest["source_detail"]).encode("utf-8")),
                    "source_row_sha256": _row_hash(latest, source_table="fact_order_status_observations"),
                }
            )
            positive_times.append(observed_dt)
        elif latest_status in NEGATIVE_STAGES:
            negative_times.append(
                _timestamp(
                    latest["observed_at"],
                    field="fact_order_status_observations.negative_observed_at",
                )
            )

    order_rows = conn.execute(
        """
        SELECT * FROM fact_orders_kaspi
        WHERE order_id=?
          AND store_code=?
        ORDER BY id
        """,
        (order_id, store_code),
    ).fetchall()
    if order_rows and all(
        _upper(row["internal_status"]) == ORDER_POSITIVE_INTERNAL
        and _upper(row["kaspi_status"]) == ORDER_POSITIVE_KASPI
        and _upper(row["source"]) in ORDER_SOURCES
        for row in order_rows
    ):
        safe_rows = []
        for row in order_rows:
            event_ts = next(
                (
                    _clean(row[column])
                    for column in ("status_updated_at", "updated_at", "imported_at", "created_at")
                    if _clean(row[column])
                ),
                "",
            )
            positive_times.append(
                _timestamp(event_ts, field="fact_orders_kaspi.positive_status_ts")
            )
            safe_rows.append(
                {
                    "row_id": int(row["id"]),
                    "internal_status": _upper(row["internal_status"]),
                    "kaspi_status": _upper(row["kaspi_status"]),
                    "source": _upper(row["source"]),
                    "status_ts": event_ts,
                    "source_row_sha256": _row_hash(row, source_table="fact_orders_kaspi"),
                }
            )
        positives.append({"source_table": "fact_orders_kaspi", "rows": safe_rows})

    for row in order_rows:
        if (
            _upper(row["internal_status"]) in NEGATIVE_STAGES
            or _upper(row["kaspi_status"]) in NEGATIVE_STAGES
        ):
            negative_ts = next(
                (
                    _clean(row[column])
                    for column in ("status_updated_at", "updated_at", "imported_at", "created_at")
                    if _clean(row[column])
                ),
                "",
            )
            negative_times.append(
                _timestamp(negative_ts, field="fact_orders_kaspi.negative_status_ts")
            )

    if not positives:
        reasons.append("NO_FIRST_PARTY_TERMINAL_POSITIVE")
        return [], reasons

    status_events = conn.execute(
        """
        SELECT event_id, store_code, order_id, stage_code, event_ts, source,
               raw_state, raw_status, source_status_change_at, source_run_id,
               flags_json, source_row_hash, idempotency_key, observed_at, created_at
        FROM order_status_event
        WHERE order_id=?
          AND store_code=?
        ORDER BY datetime(event_ts), event_id
        """,
        (order_id, store_code),
    ).fetchall()
    event_guard = []
    for row in status_events:
        event_dt = _timestamp(row["event_ts"], field="order_status_event.event_ts")
        source_row_hash = _clean(row["source_row_hash"])
        if (
            not _clean(row["source_run_id"])
            or not re.fullmatch(r"[0-9a-fA-F]{64}", source_row_hash)
        ):
            reasons.append("UNTRUSTED_STATUS_EVENT_GUARD")
            return [], reasons
        if _upper(row["stage_code"]) in NEGATIVE_STAGES:
            negative_times.append(event_dt)
        event_guard.append(
            {
                "event_id": int(row["event_id"]),
                "stage_code": _upper(row["stage_code"]),
                "event_ts": _clean(row["event_ts"]),
                "source": _clean(row["source"]),
                "source_status_change_at": _clean(row["source_status_change_at"]),
                "source_run_id": _clean(row["source_run_id"]),
                "source_row_hash": source_row_hash,
            }
        )
    latest_positive_dt = max(positive_times)
    contradiction = any(value >= latest_positive_dt for value in negative_times)
    if contradiction:
        reasons.append("TERMINAL_NEGATIVE_CONTRADICTION")
        return [], reasons
    positives.append(
        {
            "source_table": "order_status_event_negative_guard",
            "latest_positive_ts": latest_positive_dt.isoformat(),
            "rows": event_guard,
        }
    )
    return positives, reasons


def _workbook_match(row: dict[str, Any], db_row: dict[str, Any]) -> bool:
    order_id = _clean(db_row.get("order_id"))
    quantity = _decimal(db_row.get("quantity"), field="db.quantity")
    unit_price = _decimal(db_row.get("sell_price_kzt"), field="db.sell_price_kzt")
    if row["derived_order_id"] != order_id or row["raw_order_id"] != order_id:
        return False
    if row["raw_formula_headers"]:
        return False
    if row["store_code"] != _store_code(db_row.get("store_code")):
        return False
    if row["order_date"] != _clean(db_row.get("order_date"))[:10]:
        return False
    if row["sku_key"] != _clean(db_row.get("sku_key")):
        return False
    if row["sku_id"] != _clean(db_row.get("sku_id")):
        return False
    if row["size"] != _upper(db_row.get("my_size")):
        return False
    if row["raw_quantity"] is None or _decimal(row["raw_quantity"], field="crm.raw_quantity") != quantity:
        return False
    if row["derived_quantity"] is None or _decimal(
        row["derived_quantity"], field="crm.derived_quantity"
    ) != quantity:
        return False
    if row["derived_sell_price_kzt"] is None or _money(
        _decimal(row["derived_sell_price_kzt"], field="crm.derived_sell_price")
    ) != _money(unit_price):
        return False
    if row["raw_gross_kzt"] is None or abs(
        _money(_decimal(row["raw_gross_kzt"], field="crm.raw_gross"))
        - _money(unit_price * quantity)
    ) > MONEY:
        return False
    if row["raw_seller_delivery_fee_kzt"] is None:
        return False
    stored_delivery_fee = db_row.get("delivery_fee")
    if not _clean(stored_delivery_fee) or _money(
        _decimal(row["raw_seller_delivery_fee_kzt"], field="crm.raw_seller_fee")
    ) != _money(_decimal(stored_delivery_fee, field="db.delivery_fee")):
        return False
    if row["raw_cancellation_reason"]:
        return False
    return True


def _build_proof(
    *,
    db_row: dict[str, Any],
    workbook_row: dict[str, Any],
    lifecycle: list[dict[str, Any]],
    workbook_path: Path,
    workbook_sha256: str,
    db_sha256: str,
    view_sha256: str,
    policy_sha256: str,
    base_excluded_sha256: str,
    sheet: str,
) -> dict[str, Any]:
    quantity = _decimal(db_row.get("quantity"), field="db.quantity")
    unit_price = _decimal(db_row.get("sell_price_kzt"), field="db.sell_price_kzt")
    seller_fee_total = _money(
        _decimal(workbook_row["raw_seller_delivery_fee_kzt"], field="crm.raw_seller_fee")
    )
    seller_fee_unit = seller_fee_total / quantity
    order_date = date.fromisoformat(_clean(db_row.get("order_date"))[:10])
    canonical_line_net = Decimal(
        str(
            round(
                calc_net_rev(
                    float(unit_price),
                    delivery_fee=float(seller_fee_unit),
                    as_of_date=order_date,
                )
                * float(quantity),
                2,
            )
        )
    )
    db_payload = _db_row_payload(db_row)
    lifecycle_sha256 = _canonical_hash(lifecycle)
    proof_key = _sha256_bytes(
        (
            f"{BUILD_VERSION}\0{workbook_sha256}\0{sheet}\0{workbook_row['physical_row']}\0"
            f"{workbook_row['formula_row_sha256']}\0{workbook_row['allowlisted_row_sha256']}\0"
            f"{lifecycle_sha256}\0{_canonical_hash(db_payload)}\0{db_sha256}\0{view_sha256}\0"
            f"{policy_sha256}\0{base_excluded_sha256}"
        ).encode("utf-8")
    )
    stored_net = _money(_decimal(db_row.get("net_rev"), field="db.net_rev"))
    current_delivery = (
        _money(_decimal(db_row.get("delivery_fee"), field="db.delivery_fee"))
        if _clean(db_row.get("delivery_fee"))
        else None
    )
    safe_source_fields = {
        key: workbook_row[key]
        for key in (
            "derived_order_id",
            "raw_order_id",
            "store_code",
            "order_date",
            "raw_order_received_date",
            "sku_key",
            "sku_id",
            "size",
            "derived_quantity",
            "raw_quantity",
            "derived_sell_price_kzt",
            "raw_gross_kzt",
            "raw_seller_delivery_fee_kzt",
            "raw_article",
            "raw_status",
            "raw_cancellation_reason",
        )
    }
    return {
        "proof_key": proof_key,
        "source_kind": "CRM_XLSX_PHYSICAL_ROW",
        "source_workbook_path": str(workbook_path.resolve()),
        "source_workbook_sha256": workbook_sha256,
        "source_sheet": sheet,
        "source_physical_row": int(workbook_row["physical_row"]),
        "source_formula_row_sha256": workbook_row["formula_row_sha256"],
        "source_allowlisted_row_sha256": workbook_row["allowlisted_row_sha256"],
        "source_allowlisted_fields": safe_source_fields,
        "lifecycle_evidence": lifecycle,
        "lifecycle_evidence_sha256": lifecycle_sha256,
        "order_id": _clean(db_row.get("order_id")),
        "store_code": _store_code(db_row.get("store_code")),
        "sale_id": int(db_row["sale_id"]),
        "order_date": order_date.isoformat(),
        "sku_key": _clean(db_row.get("sku_key")),
        "sku_id": _clean(db_row.get("sku_id")),
        "size": _upper(db_row.get("my_size")),
        "quantity": _decimal_text(quantity),
        "gross_total_kzt": _money_text(unit_price * quantity),
        "unit_sell_price_kzt": _money_text(unit_price),
        "seller_delivery_fee_total_kzt": _money_text(seller_fee_total),
        "seller_delivery_fee_unit_kzt": _decimal_text(seller_fee_unit),
        "canonical_line_net_rev_kzt": _money_text(canonical_line_net),
        "stored_line_net_rev_kzt": _money_text(stored_net),
        "stored_delivery_fee_total_kzt": (
            _money_text(current_delivery) if current_delivery is not None else None
        ),
        "repair_required": stored_net != canonical_line_net or current_delivery != seller_fee_total,
        "db_row": db_payload,
        "db_row_sha256": _canonical_hash(db_payload),
        "copied_db_sha256": db_sha256,
        "selected_view_multiset_sha256": view_sha256,
        "economics_policy_sha256": policy_sha256,
        "base_excluded_sha256": base_excluded_sha256,
        "evidence_status": "SOURCE_PROVEN_REPAIR_CANDIDATE"
        if stored_net != canonical_line_net or current_delivery != seller_fee_total
        else "SOURCE_PROVEN_ALREADY_CANONICAL",
    }


def construct_sidecar(
    *,
    db_path: Path,
    workbook_path: Path,
    expected_db_sha256: str,
    expected_workbook_sha256: str,
    base_excluded_path: Path,
    base_excluded_sha256: str,
    since: date,
    until: date,
    sheet: str,
) -> dict[str, Any]:
    db_path = db_path.resolve()
    workbook_path = workbook_path.resolve()
    base_excluded_path = base_excluded_path.resolve()
    db_sha_before = sha256_file(db_path)
    workbook_sha_before = sha256_file(workbook_path)
    if db_sha_before != expected_db_sha256:
        raise ProvenanceError("copied DB SHA-256 mismatch")
    if workbook_sha_before != expected_workbook_sha256:
        raise ProvenanceError("CRM workbook SHA-256 mismatch")
    if sha256_file(base_excluded_path) != base_excluded_sha256:
        raise ProvenanceError("base excluded artifact SHA-256 mismatch")
    policy = _policy_manifest()

    conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA query_only=ON")
    try:
        integrity = str(conn.execute("PRAGMA integrity_check").fetchone()[0])
        if integrity.lower() != "ok":
            raise ProvenanceError(f"copied DB integrity failed: {integrity}")
        view_rows = _selected_view_rows(conn, since, until)
        view_sha256 = _canonical_hash(view_rows)
        db_rows, view_mismatch_orders = _selected_db_rows(conn, view_rows)
        db_by_order: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
        for row in db_rows:
            db_by_order[(_clean(row["order_id"]), _store_code(row["store_code"]))].append(row)

        workbook_rows, workbook_headers = _read_workbook_rows(workbook_path, sheet)
        workbook_by_order: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
        for row in workbook_rows:
            key = (row["derived_order_id"] or row["raw_order_id"], row["store_code"])
            workbook_by_order[key].append(row)

        base_excluded = _load_jsonl(base_excluded_path)
        base_keys = sorted(
            {(_clean(row.get("order_id")), _store_code(row.get("store_code"))) for row in base_excluded}
        )
        proofs: list[dict[str, Any]] = []
        excluded: list[dict[str, Any]] = []
        used_rows: list[int] = []
        for order_id, store_code in base_keys:
            key = (order_id, store_code)
            reasons: list[str] = []
            selected_lines = db_by_order.get(key, [])
            if key in view_mismatch_orders or len(selected_lines) != 1:
                reasons.append("SELECTED_DB_LINE_COUNT_NOT_ONE")
            elif _clean(selected_lines[0].get("source_file")) != CRM_SOURCE_FILE:
                reasons.append("SOURCE_FILE_NOT_CRM")
            matches = (
                [row for row in workbook_by_order.get(key, []) if _workbook_match(row, selected_lines[0])]
                if len(selected_lines) == 1
                else []
            )
            if len(selected_lines) == 1 and _clean(selected_lines[0].get("source_file")) == CRM_SOURCE_FILE:
                if not matches:
                    reasons.append("NO_EXACT_RAW_FEE_WORKBOOK_ROW")
                elif len(matches) > 1:
                    reasons.append("DUPLICATE_EXACT_WORKBOOK_ROWS")
            lifecycle: list[dict[str, Any]] = []
            if not reasons and len(matches) == 1:
                lifecycle, lifecycle_reasons = _lifecycle_evidence(
                    conn, order_id=order_id, store_code=store_code
                )
                reasons.extend(lifecycle_reasons)
            if reasons:
                excluded.append(
                    {
                        "order_id": order_id,
                        "store_code": store_code,
                        "reasons": sorted(set(reasons)),
                        "selected_db_line_count": len(selected_lines),
                        "workbook_exact_match_count": len(matches),
                    }
                )
                continue
            proof = _build_proof(
                db_row=selected_lines[0],
                workbook_row=matches[0],
                lifecycle=lifecycle,
                workbook_path=workbook_path,
                workbook_sha256=workbook_sha_before,
                db_sha256=db_sha_before,
                view_sha256=view_sha256,
                policy_sha256=policy["sha256"],
                base_excluded_sha256=base_excluded_sha256,
                sheet=sheet,
            )
            if not proof["repair_required"]:
                excluded.append(
                    {
                        "order_id": order_id,
                        "store_code": store_code,
                        "reasons": ["ALREADY_CANONICAL"],
                        "selected_db_line_count": 1,
                        "workbook_exact_match_count": 1,
                    }
                )
                continue
            proofs.append(proof)
            used_rows.append(int(matches[0]["physical_row"]))
    finally:
        conn.close()

    if len(used_rows) != len(set(used_rows)):
        raise ProvenanceError("CRM physical workbook row reused across proofs")
    proofs.sort(key=lambda row: row["proof_key"])
    excluded.sort(key=lambda row: (row["order_id"], row["store_code"], row["reasons"]))
    if len(proofs) + len(excluded) != len(base_keys):
        raise ProvenanceError("CRM proof/excluded partition mismatch")
    db_sha_after = sha256_file(db_path)
    workbook_sha_after = sha256_file(workbook_path)
    if db_sha_after != db_sha_before or workbook_sha_after != workbook_sha_before:
        raise ProvenanceError("read-only CRM sidecar input changed during construction")
    return {
        "proofs": proofs,
        "excluded": excluded,
        "metadata": {
            "schema": BUILD_VERSION,
            "since": since.isoformat(),
            "until": until.isoformat(),
            "db_path": str(db_path),
            "db_sha256": db_sha_before,
            "db_integrity": integrity,
            "workbook_path": str(workbook_path),
            "workbook_sha256": workbook_sha_before,
            "workbook_sheet": sheet,
            "workbook_headers_sha256": _canonical_hash(workbook_headers),
            "workbook_operational_row_count": len(workbook_rows),
            "base_excluded_path": str(base_excluded_path),
            "base_excluded_sha256": base_excluded_sha256,
            "base_excluded_count": len(base_keys),
            "policy": policy,
            "selected_view_line_count": len(view_rows),
            "selected_view_order_count": len(
                {(_clean(row["order_id"]), _store_code(row["store_code"])) for row in view_rows}
            ),
            "selected_view_multiset_sha256": view_sha256,
            "proof_count": len(proofs),
            "proof_order_count": len({(row["order_id"], row["store_code"]) for row in proofs}),
            "proof_physical_row_count": len(used_rows),
            "excluded_count": len(excluded),
            "exclusion_reason_counts": dict(
                sorted(Counter(reason for row in excluded for reason in row["reasons"]).items())
            ),
            "input_hashes_unchanged": True,
            "pii_fields_emitted": [],
        },
    }


def build_sidecar(**kwargs: Any) -> dict[str, Any]:
    output_dir = Path(kwargs.pop("output_dir")).resolve()
    constructed = construct_sidecar(**kwargs)
    proof_bytes = _jsonl_bytes(constructed["proofs"])
    excluded_bytes = _jsonl_bytes(constructed["excluded"])
    proof_path = output_dir / "formula_provenance.jsonl"
    excluded_path = output_dir / "excluded.jsonl"
    _write_atomic(proof_path, proof_bytes)
    _write_atomic(excluded_path, excluded_bytes)
    manifest = {
        **constructed["metadata"],
        "proof_path": str(proof_path),
        "proof_sha256": _sha256_bytes(proof_bytes),
        "excluded_path": str(excluded_path),
        "excluded_sha256": _sha256_bytes(excluded_bytes),
    }
    manifest["manifest_sha256"] = _canonical_hash(manifest)
    _write_atomic(output_dir / "manifest.json", _json_bytes(manifest, pretty=True))
    return manifest


def validate_manifest(manifest_path: Path) -> dict[str, Any]:
    manifest_path = manifest_path.resolve()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    unsigned = dict(manifest)
    expected_manifest_sha = _clean(unsigned.pop("manifest_sha256", None))
    errors: list[str] = []
    if expected_manifest_sha != _canonical_hash(unsigned):
        errors.append("MANIFEST_HASH_MISMATCH")
    proof_path = Path(_clean(manifest.get("proof_path")))
    excluded_path = Path(_clean(manifest.get("excluded_path")))
    if not proof_path.exists() or sha256_file(proof_path) != manifest.get("proof_sha256"):
        errors.append("PROOF_HASH_OR_PATH_MISMATCH")
    if not excluded_path.exists() or sha256_file(excluded_path) != manifest.get("excluded_sha256"):
        errors.append("EXCLUDED_HASH_OR_PATH_MISMATCH")
    try:
        constructed = construct_sidecar(
            db_path=Path(manifest["db_path"]),
            workbook_path=Path(manifest["workbook_path"]),
            expected_db_sha256=manifest["db_sha256"],
            expected_workbook_sha256=manifest["workbook_sha256"],
            base_excluded_path=Path(manifest["base_excluded_path"]),
            base_excluded_sha256=manifest["base_excluded_sha256"],
            since=date.fromisoformat(manifest["since"]),
            until=date.fromisoformat(manifest["until"]),
            sheet=manifest["workbook_sheet"],
        )
        expected_proof_bytes = _jsonl_bytes(constructed["proofs"])
        expected_excluded_bytes = _jsonl_bytes(constructed["excluded"])
        if proof_path.exists() and proof_path.read_bytes() != expected_proof_bytes:
            errors.append("PROOF_CONTENT_RECHECK_FAILED")
        if excluded_path.exists() and excluded_path.read_bytes() != expected_excluded_bytes:
            errors.append("EXCLUDED_CONTENT_RECHECK_FAILED")
        for key, value in constructed["metadata"].items():
            if manifest.get(key) != value:
                errors.append(f"MANIFEST_METADATA_RECHECK_FAILED:{key}")
    except (KeyError, OSError, ValueError, sqlite3.Error, ProvenanceError):
        errors.append("SOURCE_DB_RECONSTRUCTION_FAILED")
        constructed = {"proofs": [], "excluded": []}
    proofs = _load_jsonl(proof_path) if proof_path.exists() else []
    excluded = _load_jsonl(excluded_path) if excluded_path.exists() else []
    serialized = json.dumps(
        {"manifest": manifest, "proofs": proofs, "excluded": excluded},
        ensure_ascii=False,
        sort_keys=True,
    ).lower()
    forbidden = (
        "phone",
        "телефон",
        "customer_first_name",
        "customer_last_name",
        "customer name",
        "имя клиента",
        "delivery_address",
        "delivery address",
        "адрес доставки",
    )
    phone_pattern = re.compile(
        r"(?<![0-9a-z])(?:\+?7|8)(?:[\s()\-]*\d){10}(?![0-9a-z])",
        re.IGNORECASE,
    )
    email_pattern = re.compile(r"[a-z0-9._%+\-]+@[a-z0-9.\-]+\.[a-z]{2,}", re.IGNORECASE)
    if any(token in serialized for token in forbidden) or phone_pattern.search(
        serialized
    ) or email_pattern.search(serialized):
        errors.append("PII_FIELD_LEAK")
    report = {
        "schema": "sales_formula_provenance_crm_validation_v2",
        "manifest_path": str(manifest_path),
        "manifest_sha256": expected_manifest_sha,
        "status": "PASS" if not errors else "FAIL",
        "ok": not errors,
        "errors": sorted(set(errors)),
        "proof_count": len(proofs),
        "independently_reconstructed_proof_count": len(constructed["proofs"]),
        "excluded_count": len(constructed["excluded"]),
    }
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, required=True)
    parser.add_argument("--expected-db-sha256", required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--expected-workbook-sha256", required=True)
    parser.add_argument("--base-excluded", type=Path, required=True)
    parser.add_argument("--base-excluded-sha256", required=True)
    parser.add_argument("--since", type=date.fromisoformat, required=True)
    parser.add_argument("--until", type=date.fromisoformat, required=True)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        manifest = build_sidecar(
            db_path=args.db,
            workbook_path=args.workbook,
            expected_db_sha256=args.expected_db_sha256,
            expected_workbook_sha256=args.expected_workbook_sha256,
            base_excluded_path=args.base_excluded,
            base_excluded_sha256=args.base_excluded_sha256,
            since=args.since,
            until=args.until,
            sheet=args.sheet,
            output_dir=args.output_dir,
        )
    except (OSError, ValueError, sqlite3.Error, ProvenanceError) as exc:
        print(f"ERROR: {exc}", file=__import__("sys").stderr)
        return 1
    print(json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
