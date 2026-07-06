#!/usr/bin/env python3
"""Build physical count-day stock sheets from stock_ledger.

The script is file-only: it opens the SQLite database in read-only mode and
generates CSV/XLSX artifacts for a warehouse physical count.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sqlite3
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DB_PATH = PROJECT_ROOT / "db" / "app.db"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "docs" / "inventory" / "count_day_2026-07-10"
DEFAULT_EVIDENCE_DIR = PROJECT_ROOT / "exports" / "validation" / "g_count_prep_20260706"
DEFAULT_COUNT_DATE = "2026-07-10"
ALMATY = ZoneInfo("Asia/Almaty")

LIQUIDATION_CONFIG = PROJECT_ROOT / "config" / "validation" / "liquidation_register.json"
DARK_RELIST_CONFIG = PROJECT_ROOT / "config" / "validation" / "dark_family_relist_state.json"
TEMP_OCR_DECISION = (
    PROJECT_ROOT
    / "exports"
    / "current"
    / "temporary_ocr_stock_override"
    / "temporary_stock_decision_latest.csv"
)

SIZE_ORDER = {
    "XXS": 5,
    "XS": 10,
    "S": 20,
    "M": 30,
    "L": 40,
    "XL": 50,
    "2XL": 60,
    "3XL": 70,
    "4XL": 80,
    "5XL": 90,
    "6XL": 100,
}

DETAIL_COLUMNS = [
    "count_sequence",
    "priority_group",
    "family",
    "model",
    "color",
    "product_type",
    "sku_key",
    "sku_id",
    "size",
    "store_code",
    "expected_balance",
    "physical_count",
    "variance",
    "recount_count",
    "count_photo_ref",
    "counted_by",
    "issue_flags",
    "od017_flag",
    "held_flag",
    "cl_identity_flag",
    "temporary_ocr_recommendation",
    "temporary_ocr_stock",
    "latest_event_date",
    "latest_event_time",
    "event_rows",
    "notes",
]

SUMMARY_COLUMNS = [
    "count_sequence_first",
    "priority_group",
    "family",
    "family_sheet_rows",
    "total_expected_balance",
    "positive_rows",
    "zero_rows",
    "negative_rows",
    "held_rows",
    "cl_identity_rows",
    "latest_event_date",
]


@dataclass(frozen=True)
class PriorityConfig:
    dark_relist_sku_keys: tuple[str, ...]
    canonical_tranche_sku_keys: tuple[str, ...]


def _now_almaty() -> str:
    return datetime.now(ALMATY).replace(microsecond=0).isoformat()


def _resolve_project_path(path: Path) -> Path:
    path = path.expanduser()
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    return path


def _read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def _load_priority_config() -> PriorityConfig:
    dark_config = _read_json(DARK_RELIST_CONFIG)
    dark_keys = tuple(
        str(row.get("sku_key") or "")
        for row in dark_config.get("families", [])
        if row.get("sku_key")
    )
    liq_config = _read_json(LIQUIDATION_CONFIG)
    tranche_keys = tuple(str(value) for value in liq_config.get("canonical_tranche1_sku_keys", []))
    return PriorityConfig(dark_relist_sku_keys=dark_keys, canonical_tranche_sku_keys=tranche_keys)


def _connect_readonly(db_path: Path) -> sqlite3.Connection:
    uri = f"file:{db_path}?mode=ro"
    conn = sqlite3.connect(uri, uri=True)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA query_only = ON")
    return conn


def _sanitize_token(value: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", str(value).upper()).strip("_")


def _size_sort_key(size: str) -> tuple[int, int, str]:
    text = str(size or "").strip().upper()
    if text in SIZE_ORDER:
        return (0, SIZE_ORDER[text], text)
    if "/" in text:
        left, _, right = text.partition("/")
        for token in (right, left):
            if token.isdigit():
                return (1, int(token), text)
    if text.isdigit():
        return (1, int(text), text)
    match = re.search(r"(\d+)", text)
    if match:
        return (2, int(match.group(1)), text)
    return (3, 9999, text)


def _priority_group(sku_key: str, config: PriorityConfig) -> tuple[int, str]:
    if sku_key in config.dark_relist_sku_keys:
        idx = config.dark_relist_sku_keys.index(sku_key) + 1
        return (idx, "01_dark_relist_rush_tshirt_od033")
    if sku_key in config.canonical_tranche_sku_keys:
        idx = config.canonical_tranche_sku_keys.index(sku_key) + 1
        return (100 + idx, "02_canonical_tranche1_od011")
    upper = sku_key.upper()
    if "RUSH" in upper or "T-SHIRT" in upper or "TSHIRT" in upper:
        return (300, "03_other_rush_tshirt")
    return (900, "99_all_other_families")


def _row_sort_key(row: dict[str, Any]) -> tuple[Any, ...]:
    return (
        int(row["priority_rank"]),
        str(row["family"]),
        _size_sort_key(str(row["size"])),
        str(row["sku_id"]),
        str(row["store_code"]),
    )


def _write_csv_atomic(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = path.with_name(f".{path.name}.tmp")
    with tmp_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)
    os.replace(tmp_path, path)


def _write_json_atomic(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = path.with_name(f".{path.name}.tmp")
    tmp_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    os.replace(tmp_path, path)


def _load_exception_flags(conn: sqlite3.Connection) -> list[dict[str, str]]:
    rows = conn.execute(
        """
        SELECT exception_id, reason, recommended_action
        FROM exception_queue
        WHERE status = 'OPEN'
          AND domain = 'STOCK'
        ORDER BY exception_id
        """
    ).fetchall()
    return [
        {
            "exception_id": str(row["exception_id"] or ""),
            "reason": str(row["reason"] or ""),
            "recommended_action": str(row["recommended_action"] or ""),
            "haystack": _sanitize_token(str(row["exception_id"] or "")),
        }
        for row in rows
    ]


def _exception_matches(record: dict[str, str], row: dict[str, Any]) -> bool:
    candidates = [
        _sanitize_token(str(row.get("sku_id") or "")),
        _sanitize_token(f"{row.get('sku_key')}_{row.get('size')}"),
        _sanitize_token(str(row.get("sku_key") or "")),
    ]
    haystack = record["haystack"]
    for candidate in candidates:
        if len(candidate) >= 5 and candidate in haystack:
            return True
    return False


def _exception_codes_for_row(records: list[dict[str, str]], row: dict[str, Any]) -> tuple[list[str], list[str]]:
    codes: list[str] = []
    notes: list[str] = []
    for record in records:
        if not _exception_matches(record, row):
            continue
        reason = record["reason"]
        code = reason.split(":", 1)[0].strip() if reason else record["exception_id"]
        if code:
            codes.append(code)
        if record["recommended_action"]:
            notes.append(record["recommended_action"])
    return sorted(dict.fromkeys(codes)), notes


def _load_temporary_ocr_flags(path: Path) -> dict[tuple[str, str, str], dict[str, str]]:
    if not path.exists():
        return {}
    flags: dict[tuple[str, str, str], dict[str, str]] = {}
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            key = (str(row.get("sku_key") or ""), str(row.get("sku_id") or ""), str(row.get("my_size") or ""))
            flags[key] = {
                "activation_recommendation": str(row.get("activation_recommendation") or ""),
                "temporary_current_stock": str(row.get("temporary_current_stock") or ""),
                "notes": str(row.get("notes") or ""),
            }
    return flags


def _is_cl_identity_placeholder(row: dict[str, Any]) -> bool:
    return str(row.get("sku_key") or "").strip().upper() == "CL"


def _is_temp_held(temp: dict[str, str] | None) -> bool:
    if not temp:
        return False
    recommendation = temp.get("activation_recommendation", "").upper()
    return "OWNER_CONFLICT_HOLD" in recommendation or "PARKED_MAPPING_PENDING" in recommendation


def _load_ledger_rows(
    conn: sqlite3.Connection,
    *,
    priority_config: PriorityConfig,
    exception_records: list[dict[str, str]],
    temp_flags: dict[tuple[str, str, str], dict[str, str]],
) -> list[dict[str, Any]]:
    sql = """
        WITH balances AS (
            SELECT
                sku_key,
                sku_id,
                my_size,
                COALESCE(store_code, '') AS store_code,
                SUM(qty_change) AS expected_balance,
                COUNT(*) AS event_rows,
                MIN(event_date) AS first_event_date,
                MAX(event_date) AS latest_event_date,
                MAX(COALESCE(event_time, '')) AS latest_event_time
            FROM stock_ledger
            GROUP BY sku_key, sku_id, my_size, COALESCE(store_code, '')
        )
        SELECT
            b.*,
            COALESCE(d.model, '') AS model,
            COALESCE(d.color, '') AS color,
            COALESCE(d.product_type, '') AS product_type,
            COALESCE(d.category, '') AS category,
            COALESCE(d.gender, '') AS gender,
            COALESCE(d.active_flag, '') AS dim_active_flag
        FROM balances b
        LEFT JOIN dim_sku d
          ON d.sku_key = b.sku_key
    """
    rows: list[dict[str, Any]] = []
    for db_row in conn.execute(sql):
        sku_key = str(db_row["sku_key"] or "")
        sku_id = str(db_row["sku_id"] or "")
        size = str(db_row["my_size"] or "")
        priority_rank, priority_label = _priority_group(sku_key, priority_config)
        base = {
            "priority_rank": priority_rank,
            "priority_group": priority_label,
            "family": sku_key,
            "model": str(db_row["model"] or ""),
            "color": str(db_row["color"] or ""),
            "product_type": str(db_row["product_type"] or ""),
            "sku_key": sku_key,
            "sku_id": sku_id,
            "size": size,
            "store_code": str(db_row["store_code"] or ""),
            "expected_balance": int(db_row["expected_balance"] or 0),
            "physical_count": "",
            "variance": "",
            "recount_count": "",
            "count_photo_ref": "",
            "counted_by": "",
            "latest_event_date": str(db_row["latest_event_date"] or ""),
            "latest_event_time": str(db_row["latest_event_time"] or ""),
            "event_rows": int(db_row["event_rows"] or 0),
        }
        exception_codes, exception_notes = _exception_codes_for_row(exception_records, base)
        temp = temp_flags.get((sku_key, sku_id, size))
        issue_flags: list[str] = []
        notes: list[str] = []
        held_flag = ""
        od017_flag = ""
        cl_identity_flag = ""

        if int(base["expected_balance"]) < 0:
            issue_flags.append("NEGATIVE_BALANCE")
            od017_flag = "OD-017_NEGATIVE_ROUTE_QUARANTINE_NOT_CLAMP"

        if exception_codes:
            issue_flags.extend(f"EXCEPTION_{code}" for code in exception_codes)
            if any(
                token in code
                for code in exception_codes
                for token in (
                    "OWNER_OOS_ACTIVE_ZERO",
                    "OWNER_OVERRIDE_NO_DOUBLE_REDUCE",
                    "LINE61_4XL_EXCLUDED",
                    "NEGATIVE_RAW_LEDGER_BALANCE",
                )
            ):
                held_flag = "HELD_BY_OPEN_STOCK_EXCEPTION"
                od017_flag = od017_flag or "OD-017_HELD_RECOUNT_OR_QUARANTINE_REVIEW"
            notes.extend(exception_notes)

        if _is_temp_held(temp):
            issue_flags.append("TEMP_OCR_HOLD_OR_MAPPING_PARK")
            held_flag = held_flag or "HELD_BY_TEMP_OCR_LAYER"
            od017_flag = od017_flag or "OD-017_HELD_RECOUNT_OR_QUARANTINE_REVIEW"
            if temp and temp.get("notes"):
                notes.append(temp["notes"])

        if _is_cl_identity_placeholder(base):
            issue_flags.append("CL_IDENTITY_PLACEHOLDER")
            cl_identity_flag = "CL_IDENTITY_PLACEHOLDER_REQUIRES_MAPPING"

        if not issue_flags:
            issue_flags.append("OK")

        base.update(
            {
                "issue_flags": ";".join(dict.fromkeys(issue_flags)),
                "od017_flag": od017_flag,
                "held_flag": held_flag,
                "cl_identity_flag": cl_identity_flag,
                "temporary_ocr_recommendation": temp.get("activation_recommendation", "") if temp else "",
                "temporary_ocr_stock": temp.get("temporary_current_stock", "") if temp else "",
                "notes": " | ".join(dict.fromkeys(note for note in notes if note)),
            }
        )
        rows.append(base)

    rows.sort(key=_row_sort_key)
    for idx, row in enumerate(rows, start=1):
        row["count_sequence"] = idx
    return rows


def _summarize_by_family(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[str(row["family"])].append(row)

    summaries: list[dict[str, Any]] = []
    for family, family_rows in grouped.items():
        balances = [int(row["expected_balance"]) for row in family_rows]
        latest_dates = [str(row["latest_event_date"]) for row in family_rows if row.get("latest_event_date")]
        summaries.append(
            {
                "count_sequence_first": min(int(row["count_sequence"]) for row in family_rows),
                "priority_group": str(family_rows[0]["priority_group"]),
                "family": family,
                "family_sheet_rows": len(family_rows),
                "total_expected_balance": sum(balances),
                "positive_rows": sum(1 for value in balances if value > 0),
                "zero_rows": sum(1 for value in balances if value == 0),
                "negative_rows": sum(1 for value in balances if value < 0),
                "held_rows": sum(1 for row in family_rows if row.get("held_flag")),
                "cl_identity_rows": sum(1 for row in family_rows if row.get("cl_identity_flag")),
                "latest_event_date": max(latest_dates) if latest_dates else "",
            }
        )
    summaries.sort(key=lambda row: int(row["count_sequence_first"]))
    return summaries


def _matrix_rows(rows: list[dict[str, Any]], summaries: list[dict[str, Any]]) -> tuple[list[str], list[dict[str, Any]]]:
    sizes = sorted({str(row["size"]) for row in rows}, key=_size_sort_key)
    balances: dict[tuple[str, str], int] = defaultdict(int)
    for row in rows:
        balances[(str(row["family"]), str(row["size"]))] += int(row["expected_balance"])
    summary_by_family = {str(row["family"]): row for row in summaries}
    fieldnames = [
        "count_sequence_first",
        "priority_group",
        "family",
        "total_expected_balance",
        "negative_rows",
        "held_rows",
        "cl_identity_rows",
        *sizes,
    ]
    matrix: list[dict[str, Any]] = []
    for family in [str(row["family"]) for row in summaries]:
        summary = summary_by_family[family]
        out = {key: summary.get(key, "") for key in fieldnames}
        for size in sizes:
            out[size] = balances.get((family, size), "")
        matrix.append(out)
    return fieldnames, matrix


def _autosize_sheet(ws: Any, max_width: int = 46) -> None:
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        width = min(max_width, max(len(str(cell.value or "")) for cell in column_cells) + 2)
        ws.column_dimensions[letter].width = max(10, width)


def _write_xlsx(path: Path, detail_rows: list[dict[str, Any]], summary_rows: list[dict[str, Any]], matrix_rows: list[dict[str, Any]], matrix_columns: list[str], metadata: dict[str, Any]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "openpyxl is required for XLSX generation; run with the repo .venv Python"
        ) from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "Count Sheets"
    ws.append(DETAIL_COLUMNS)
    for row in detail_rows:
        ws.append([row.get(column, "") for column in DETAIL_COLUMNS])

    summary_ws = wb.create_sheet("Family Summary")
    summary_ws.append(SUMMARY_COLUMNS)
    for row in summary_rows:
        summary_ws.append([row.get(column, "") for column in SUMMARY_COLUMNS])

    matrix_ws = wb.create_sheet("Family Size Matrix")
    matrix_ws.append(matrix_columns)
    for row in matrix_rows:
        matrix_ws.append([row.get(column, "") for column in matrix_columns])

    meta_ws = wb.create_sheet("Source Notes")
    for key, value in metadata.items():
        meta_ws.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    negative_fill = PatternFill("solid", fgColor="F4CCCC")
    held_fill = PatternFill("solid", fgColor="FFF2CC")
    cl_fill = PatternFill("solid", fgColor="FCE4D6")
    priority_fill = PatternFill("solid", fgColor="D9EAF7")

    for sheet in (ws, summary_ws, matrix_ws, meta_ws):
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        _autosize_sheet(sheet)

    issue_col = DETAIL_COLUMNS.index("issue_flags") + 1
    priority_col = DETAIL_COLUMNS.index("priority_group") + 1
    for row_idx in range(2, ws.max_row + 1):
        issue_text = str(ws.cell(row_idx, issue_col).value or "")
        priority_text = str(ws.cell(row_idx, priority_col).value or "")
        fill = None
        if "NEGATIVE_BALANCE" in issue_text:
            fill = negative_fill
        elif "HOLD" in issue_text or "PARK" in issue_text or "EXCEPTION_" in issue_text:
            fill = held_fill
        elif "CL_IDENTITY_PLACEHOLDER" in issue_text:
            fill = cl_fill
        elif priority_text.startswith(("01_", "02_", "03_")):
            fill = priority_fill
        if fill:
            for cell in ws[row_idx]:
                cell.fill = fill

    for sheet in (ws, summary_ws, matrix_ws):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, int):
                    cell.number_format = "#,##0"

    path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = path.with_name(f".{path.name}.tmp")
    wb.save(tmp_path)
    os.replace(tmp_path, path)


def build_count_day_stock_sheets(
    *,
    db_path: Path,
    output_dir: Path,
    evidence_dir: Path,
    count_date: str,
) -> dict[str, Any]:
    db_path = _resolve_project_path(db_path)
    output_dir = _resolve_project_path(output_dir)
    evidence_dir = _resolve_project_path(evidence_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    evidence_dir.mkdir(parents=True, exist_ok=True)

    priority_config = _load_priority_config()
    temp_flags = _load_temporary_ocr_flags(TEMP_OCR_DECISION)
    with _connect_readonly(db_path) as conn:
        exception_records = _load_exception_flags(conn)
        detail_rows = _load_ledger_rows(
            conn,
            priority_config=priority_config,
            exception_records=exception_records,
            temp_flags=temp_flags,
        )

    summary_rows = _summarize_by_family(detail_rows)
    matrix_columns, matrix = _matrix_rows(detail_rows, summary_rows)

    count_tag = count_date.replace("-", "")
    detail_csv = output_dir / f"count_day_expected_balances_{count_tag}.csv"
    summary_csv = output_dir / f"count_day_family_summary_{count_tag}.csv"
    matrix_csv = output_dir / f"count_day_family_size_matrix_{count_tag}.csv"
    xlsx_path = output_dir / f"count_day_expected_balances_{count_tag}.xlsx"

    _write_csv_atomic(detail_csv, detail_rows, DETAIL_COLUMNS)
    _write_csv_atomic(summary_csv, summary_rows, SUMMARY_COLUMNS)
    _write_csv_atomic(matrix_csv, matrix, matrix_columns)

    metadata = {
        "generated_at_almaty": _now_almaty(),
        "count_date": count_date,
        "db_path": str(db_path),
        "source_table": "stock_ledger",
        "db_open_mode": "sqlite readonly URI mode=ro plus PRAGMA query_only=ON",
        "detail_rows": len(detail_rows),
        "family_count": len(summary_rows),
        "negative_rows": sum(1 for row in detail_rows if int(row["expected_balance"]) < 0),
        "held_rows": sum(1 for row in detail_rows if row.get("held_flag")),
        "cl_identity_rows": sum(1 for row in detail_rows if row.get("cl_identity_flag")),
        "priority_sources": {
            "dark_relist_config": str(DARK_RELIST_CONFIG),
            "liquidation_config": str(LIQUIDATION_CONFIG),
        },
        "temporary_ocr_decision": str(TEMP_OCR_DECISION) if TEMP_OCR_DECISION.exists() else "",
    }
    _write_xlsx(xlsx_path, detail_rows, summary_rows, matrix, matrix_columns, metadata)

    summary = {
        **metadata,
        "outputs": {
            "detail_csv": str(detail_csv),
            "summary_csv": str(summary_csv),
            "matrix_csv": str(matrix_csv),
            "xlsx": str(xlsx_path),
        },
        "family_counts": summary_rows,
    }
    _write_json_atomic(evidence_dir / "count_sheet_build_summary.json", summary)
    _write_csv_atomic(evidence_dir / "count_sheet_family_summary.csv", summary_rows, SUMMARY_COLUMNS)
    return summary


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB_PATH)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--evidence-dir", type=Path, default=DEFAULT_EVIDENCE_DIR)
    parser.add_argument("--count-date", default=DEFAULT_COUNT_DATE)
    parser.add_argument("--json", action="store_true", help="Print the build summary JSON")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    summary = build_count_day_stock_sheets(
        db_path=args.db,
        output_dir=args.output_dir,
        evidence_dir=args.evidence_dir,
        count_date=args.count_date,
    )
    if args.json:
        print(json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True))
    else:
        print(f"family_count={summary['family_count']}")
        print(f"detail_rows={summary['detail_rows']}")
        print(f"negative_rows={summary['negative_rows']}")
        print(f"held_rows={summary['held_rows']}")
        print(f"cl_identity_rows={summary['cl_identity_rows']}")
        print(f"xlsx={summary['outputs']['xlsx']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
