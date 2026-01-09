#!/usr/bin/env python3
"""
Import PO-4 approved inbound quantities from PO-4_inbound.xlsx.

Source column: Approved_by_supplier_qty (shipped by supplier, not received yet).
Writes to po_header + po_line and activates SKUs/sizes in dim_sku + dim_sku_size.
"""

from __future__ import annotations

import argparse
from datetime import date, datetime
from pathlib import Path
from typing import Optional

import sys

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from core.db import DEFAULT_DB_PATH, get_db


def _parse_date(val) -> Optional[str]:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date().isoformat()
    if isinstance(val, date):
        return val.isoformat()
    if isinstance(val, str):
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d.%m.%y", "%d/%m/%y"):
            try:
                return datetime.strptime(val.strip(), fmt).date().isoformat()
            except ValueError:
                continue
    return None


def _norm_header(val: object) -> str:
    if val is None:
        return ""
    return str(val).strip().lower()


def _infer_sku_key_from_id(sku_id: str) -> str:
    if "_" not in sku_id:
        return sku_id
    return sku_id.rsplit("_", 1)[0]


def _infer_size_from_id(sku_id: str) -> Optional[str]:
    if "_" not in sku_id:
        return None
    return sku_id.rsplit("_", 1)[-1]


def _infer_product_type(sku_key: str, raw_type: Optional[str]) -> str:
    if raw_type:
        return str(raw_type).strip().upper()
    if sku_key.startswith("ELS_"):
        return "ELS"
    return "CL"


def import_po4_inbound(xlsx_path: Path, po_id: str, db_path: Path) -> dict:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active

    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=20, values_only=True):
        if row and any(row):
            header_row = row
            break

    if not header_row:
        raise RuntimeError("No header row found in PO-4 inbound file")

    header_map = {_norm_header(v): idx for idx, v in enumerate(header_row)}

    def col(name: str) -> Optional[int]:
        return header_map.get(_norm_header(name))

    idx_internal = col("Internal_PO")
    idx_sku_key = col("SKU_key")
    idx_sku_id = col("SKU_ID")
    idx_size = col("MY_SIZE")
    idx_qty = col("Approved_by_supplier_qty")
    idx_msg = col("Message_date")
    idx_ship = col("Ship_date")
    idx_ptype = col("Product_Type")
    idx_base_cny = col("BaseCost_CNY")
    idx_weight = col("Weight_kg")
    idx_unit_cogs = col("COGS_unit_KZT")

    rows = []
    msg_dates = set()
    ship_dates = set()
    po_ids = set()

    for row in ws.iter_rows(min_row=ws.min_row + 1, values_only=True):
        if not row or not any(row):
            continue
        approved = row[idx_qty] if idx_qty is not None else None
        try:
            approved_qty = int(approved) if approved not in (None, "") else 0
        except (ValueError, TypeError):
            approved_qty = 0
        if approved_qty <= 0:
            continue

        raw_sku_key = row[idx_sku_key] if idx_sku_key is not None else None
        raw_sku_id = row[idx_sku_id] if idx_sku_id is not None else None
        raw_size = row[idx_size] if idx_size is not None else None
        raw_type = row[idx_ptype] if idx_ptype is not None else None

        sku_id = str(raw_sku_id).strip() if raw_sku_id not in (None, "") else ""
        sku_key = str(raw_sku_key).strip() if raw_sku_key not in (None, "") else ""
        if not sku_key and sku_id:
            sku_key = _infer_sku_key_from_id(sku_id)
        if not sku_key:
            continue

        my_size = str(raw_size).strip() if raw_size not in (None, "") else ""
        if not my_size and sku_id:
            inferred = _infer_size_from_id(sku_id)
            if inferred:
                my_size = inferred

        product_type = _infer_product_type(sku_key, raw_type)
        if not my_size and product_type == "ELS":
            my_size = "ONE_SIZE"

        if not sku_id and my_size:
            sku_id = f"{sku_key}_{my_size}"

        base_cny = row[idx_base_cny] if idx_base_cny is not None else None
        weight_kg = row[idx_weight] if idx_weight is not None else None
        unit_cogs = row[idx_unit_cogs] if idx_unit_cogs is not None else None

        msg_date = _parse_date(row[idx_msg]) if idx_msg is not None else None
        ship_date = _parse_date(row[idx_ship]) if idx_ship is not None else None
        if msg_date:
            msg_dates.add(msg_date)
        if ship_date:
            ship_dates.add(ship_date)

        internal_po = str(row[idx_internal]).strip() if idx_internal is not None and row[idx_internal] not in (None, "") else ""
        if internal_po:
            po_ids.add(internal_po)

        rows.append({
            "sku_key": sku_key,
            "sku_id": sku_id,
            "my_size": my_size,
            "product_type": product_type,
            "approved_qty": approved_qty,
            "base_cost_cny": float(base_cny) if base_cny not in (None, "") else None,
            "weight_kg": float(weight_kg) if weight_kg not in (None, "") else None,
            "unit_cogs_kzt": float(unit_cogs) if unit_cogs not in (None, "") else None,
        })

    if not rows:
        raise RuntimeError("No approved rows found (Approved_by_supplier_qty > 0)")

    if len(po_ids) == 1:
        po_id = po_ids.pop()

    msg_date = sorted(msg_dates)[0] if msg_dates else None
    ship_date = sorted(ship_dates)[0] if ship_dates else None

    with get_db(db_path) as conn:
        # Activate / create SKU master data
        for row in rows:
            sku_key = row["sku_key"]
            product_type = row["product_type"]
            base_cost_cny = row["base_cost_cny"]
            weight_kg = row["weight_kg"]

            existing = conn.execute(
                "SELECT sku_key, product_type, base_cost_cny, weight_kg FROM dim_sku WHERE sku_key = ?",
                (sku_key,),
            ).fetchone()

            if existing:
                new_product_type = existing[1] or product_type
                new_base = existing[2] if existing[2] not in (None, 0) else (base_cost_cny or 0)
                new_weight = existing[3] if existing[3] not in (None, 0) else (weight_kg or 0)
                conn.execute(
                    """
                    UPDATE dim_sku
                       SET active_flag = 1,
                           product_type = ?,
                           base_cost_cny = ?,
                           weight_kg = ?,
                           updated_at = datetime('now')
                     WHERE sku_key = ?
                    """,
                    (new_product_type, new_base, new_weight, sku_key),
                )
            else:
                conn.execute(
                    """
                    INSERT INTO dim_sku (
                        sku_key, model, color, product_type,
                        base_cost_cny, weight_kg, category, gender,
                        active_flag, created_at, updated_at,
                        cogs_kzt, avg_sell_price_kzt_used, avg_sell_price_source, price_missing_flag
                    ) VALUES (
                        ?, NULL, NULL, ?,
                        ?, ?, NULL, NULL,
                        1, datetime('now'), datetime('now'),
                        NULL, NULL, NULL, 1
                    )
                    """,
                    (sku_key, product_type, base_cost_cny or 0, weight_kg or 0),
                )

            # Ensure size entry exists + active
            my_size = row["my_size"]
            if my_size:
                size_exists = conn.execute(
                    "SELECT 1 FROM dim_sku_size WHERE sku_key = ? AND my_size = ?",
                    (sku_key, my_size),
                ).fetchone()
                if size_exists:
                    conn.execute(
                        """
                        UPDATE dim_sku_size
                           SET active_flag = 1,
                               sku_id = COALESCE(sku_id, ?)
                         WHERE sku_key = ? AND my_size = ?
                        """,
                        (row["sku_id"], sku_key, my_size),
                    )
                else:
                    conn.execute(
                        """
                        INSERT INTO dim_sku_size (
                            sku_key, sku_id, my_size, active_flag, created_at
                        ) VALUES (
                            ?, ?, ?, 1, datetime('now')
                        )
                        """,
                        (sku_key, row["sku_id"], my_size),
                    )

        # Upsert PO header
        units_total = sum(r["approved_qty"] for r in rows)
        conn.execute(
            """
            INSERT INTO po_header (
                po_id, message_date, ship_date_seller, status, units_total, created_at, updated_at
            ) VALUES (
                ?, ?, ?, 'IN_TRANSIT', ?, datetime('now'), datetime('now')
            )
            ON CONFLICT(po_id) DO UPDATE SET
                message_date = excluded.message_date,
                ship_date_seller = excluded.ship_date_seller,
                status = 'IN_TRANSIT',
                units_total = excluded.units_total,
                updated_at = datetime('now')
            """,
            (po_id, msg_date, ship_date, units_total),
        )

        # Replace PO lines for this PO
        conn.execute("DELETE FROM po_line WHERE po_id = ?", (po_id,))
        for row in rows:
            conn.execute(
                """
                INSERT INTO po_line (
                    po_id, sku_key, sku_id, my_size,
                    order_qty, received_qty,
                    unit_cost_cny, unit_cost_kzt, unit_weight_kg,
                    status, created_at, updated_at
                ) VALUES (
                    ?, ?, ?, ?,
                    ?, 0,
                    ?, ?, ?,
                    'IN_TRANSIT', datetime('now'), datetime('now')
                )
                """,
                (
                    po_id,
                    row["sku_key"],
                    row["sku_id"],
                    row["my_size"],
                    row["approved_qty"],
                    row["base_cost_cny"] or 0.0,
                    row["unit_cogs_kzt"],
                    row["weight_kg"],
                ),
            )

    return {
        "po_id": po_id,
        "rows": len(rows),
        "message_date": msg_date,
        "ship_date": ship_date,
        "units_total": units_total,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Import PO-4 inbound approvals")
    parser.add_argument("xlsx_path", type=Path, help="Path to PO-4_inbound.xlsx")
    parser.add_argument("--po-id", default="PO-4", help="PO id to store in DB (default: PO-4)")
    parser.add_argument("--db", dest="db_path", default=DEFAULT_DB_PATH, type=Path, help="Database path")
    args = parser.parse_args()

    result = import_po4_inbound(args.xlsx_path, args.po_id, args.db_path)
    print("Imported PO-4 inbound approvals:")
    print(f"  PO: {result['po_id']}")
    print(f"  Rows: {result['rows']}")
    print(f"  Units total: {result['units_total']}")
    if result["message_date"]:
        print(f"  Message date: {result['message_date']}")
    if result["ship_date"]:
        print(f"  Ship date: {result['ship_date']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
