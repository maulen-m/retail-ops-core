#!/usr/bin/env python3
"""Build a deterministic read-only packet for a derived/raw order-ID conflict.

The packet reads only allowlisted, non-customer workbook fields and safe SQLite
identity/lifecycle fields. It never edits a workbook, SQLite database, or an
external system.
"""

from __future__ import annotations

import argparse
import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import sqlite3
from typing import Any, Iterable

from openpyxl import load_workbook


class PacketError(RuntimeError):
    """Raised when pinned evidence is missing, ambiguous, or inconsistent."""


CANONICAL_HASH_VERSION = "canonical-json-v1-sort-keys-utf8-no-whitespace"
DEFAULT_SHEET = "SALES_KSP_CRM_1"
SAFE_WORKBOOK_HEADERS = (
    "Date",
    "STORE_NAME",
    "Quantity",
    "Kaspi_name_core",
    "OrderID",
    "MY_SIZE",
    "SKU_key",
    "SKU_ID",
    "PLANNED_SHIPPING_DATE",
    "SKU_ID_KSP",
    "Kaspi_name_source",
    "№ заказа",
    "Дата поступления заказа",
    "Название товара в Kaspi Магазине",
    "Название в системе продавца",
    "Артикул",
    "Сумма",
    "Дата изменения статуса",
    "Статус",
    "Причина отмены",
    "Количество",
    "Стоимость доставки для продавца",
    "Плановая дата передачи курьеру",
    "Склад передачи КД",
)
TERMINAL_INTERNAL_STATUSES = {
    "CANCELLED",
    "CANCELED",
    "COMPLETED",
    "DELIVERED",
    "RETURNED",
}


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _canonical_sha(value: Any) -> str:
    return hashlib.sha256(
        json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            default=str,
        ).encode("utf-8")
    ).hexdigest()


def _normalize_scalar(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(Decimal(str(value)).normalize(), "f")
    return _text(value)


def _normalize_id(value: Any, *, field: str) -> str:
    normalized = _normalize_scalar(value)
    if not normalized or not normalized.isdigit():
        raise PacketError(f"{field} must be a non-empty decimal ID, got {normalized!r}")
    return normalized


def _normalize_store(value: Any) -> str:
    raw = _text(value).upper().replace("-", "").replace("_", "")
    aliases = {
        "ACMEWEAR": "ACMEWEAR",
        "UNIVERSAL": "UNIVERSAL",
        "STOREB": "STOREB",
    }
    return aliases.get(raw, raw)


def _decimal_equal(left: Any, right: Any) -> bool:
    try:
        return Decimal(_text(left)) == Decimal(_text(right))
    except (InvalidOperation, ValueError):
        return False


def _connect_read_only(path: Path) -> sqlite3.Connection:
    resolved = path.expanduser().resolve()
    conn = sqlite3.connect(f"file:{resolved}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA query_only=ON")
    return conn


def _object_exists(conn: sqlite3.Connection, kind: str, name: str) -> bool:
    return (
        conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type=? AND name=?",
            (kind, name),
        ).fetchone()
        is not None
    )


def _safe_db_rows(
    conn: sqlite3.Connection,
    *,
    derived_order_id: str,
    raw_order_id: str,
) -> dict[str, list[dict[str, Any]]]:
    ids = (derived_order_id, raw_order_id)
    if not _object_exists(conn, "table", "fact_orders_kaspi"):
        raise PacketError("fact_orders_kaspi is missing")
    order_rows = [
        dict(row)
        for row in conn.execute(
            """
            SELECT id, CAST(order_id AS TEXT) AS order_id, store_code,
                   channel_code, kaspi_offer_name, sku_key, sku_id, my_size,
                   assigned_size, quantity, unit_price_kzt,
                   planned_shipment_date, actual_shipment_date, kaspi_status,
                   internal_status, status_updated_at, source, source_file,
                   line_identity_key, kaspi_article
            FROM fact_orders_kaspi
            WHERE CAST(order_id AS TEXT) IN (?, ?)
            ORDER BY CAST(order_id AS TEXT), id
            """,
            ids,
        )
    ]

    observations: list[dict[str, Any]] = []
    if _object_exists(conn, "table", "fact_order_status_observations"):
        observations = [
            dict(row)
            for row in conn.execute(
                """
                SELECT id, CAST(order_id AS TEXT) AS order_id, store_code,
                       status_internal, observed_at, source, ledger_run_id,
                       source_detail
                FROM fact_order_status_observations
                WHERE CAST(order_id AS TEXT) IN (?, ?)
                ORDER BY CAST(order_id AS TEXT), observed_at, id
                """,
                ids,
            )
        ]

    events: list[dict[str, Any]] = []
    if _object_exists(conn, "table", "order_status_event"):
        events = [
            dict(row)
            for row in conn.execute(
                """
                SELECT event_id, store_code, CAST(order_id AS TEXT) AS order_id,
                       stage_code, event_ts, source, raw_state, raw_status,
                       source_status_change_at, source_run_id, source_row_hash,
                       idempotency_key
                FROM order_status_event
                WHERE CAST(order_id AS TEXT) IN (?, ?)
                ORDER BY CAST(order_id AS TEXT), event_ts, event_id
                """,
                ids,
            )
        ]

    anchors: list[dict[str, Any]] = []
    if _object_exists(conn, "table", "fact_sales_workbook_anchor"):
        anchors = [
            dict(row)
            for row in conn.execute(
                """
                SELECT CAST(order_id AS TEXT) AS order_id, store_code, sale_date,
                       quantity, net_rev_kzt, total_price_kzt, source_file,
                       updated_at
                FROM fact_sales_workbook_anchor
                WHERE CAST(order_id AS TEXT) IN (?, ?)
                ORDER BY CAST(order_id AS TEXT), store_code
                """,
                ids,
            )
        ]

    return {
        "fact_orders_kaspi": order_rows,
        "fact_order_status_observations": observations,
        "order_status_event": events,
        "fact_sales_workbook_anchor": anchors,
    }


def _read_workbook_match(
    path: Path,
    *,
    expected_sha256: str,
    sheet_name: str,
    expected_row_number: int,
    derived_order_id: str,
    raw_order_id: str,
    expected_store: str,
) -> dict[str, Any]:
    before_sha = _sha256(path)
    if before_sha != expected_sha256:
        raise PacketError(
            f"workbook hash mismatch for {path}: expected {expected_sha256}, got {before_sha}"
        )
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise PacketError(f"sheet {sheet_name!r} missing from {path}")
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = list(next(rows))
        except StopIteration as exc:
            raise PacketError(f"empty workbook sheet: {path}:{sheet_name}") from exc
        header_positions: dict[str, int] = {}
        for index, value in enumerate(raw_headers):
            header = _text(value)
            if not header:
                continue
            if header in header_positions:
                raise PacketError(f"duplicate header {header!r} in {path}:{sheet_name}")
            header_positions[header] = index
        missing = [name for name in SAFE_WORKBOOK_HEADERS if name not in header_positions]
        if missing:
            raise PacketError(f"required safe headers missing from {path}: {missing}")

        matches: list[tuple[int, dict[str, str]]] = []
        for row_number, row in enumerate(rows, start=2):
            try:
                derived_value = _normalize_id(
                    row[header_positions["OrderID"]], field="OrderID"
                )
            except PacketError:
                continue
            if derived_value != derived_order_id:
                continue
            raw_value = _normalize_id(row[header_positions["№ заказа"]], field="№ заказа")
            safe_row = {
                header: _normalize_scalar(row[header_positions[header]])
                for header in SAFE_WORKBOOK_HEADERS
            }
            if raw_value != raw_order_id:
                raise PacketError(
                    f"derived ID {derived_order_id} maps to unexpected raw ID {raw_value} in {path}"
                )
            if _normalize_store(safe_row["STORE_NAME"]) != expected_store:
                raise PacketError(
                    f"unexpected store {safe_row['STORE_NAME']!r} for {derived_order_id} in {path}"
                )
            matches.append((row_number, safe_row))
        if len(matches) != 1:
            raise PacketError(
                f"expected one derived/raw match in {path}:{sheet_name}, found {len(matches)}"
            )
        row_number, safe_row = matches[0]
        if row_number != expected_row_number:
            raise PacketError(
                f"expected physical row {expected_row_number} in {path}, got {row_number}"
            )
    finally:
        workbook.close()
    after_sha = _sha256(path)
    if after_sha != before_sha:
        raise PacketError(f"workbook changed during read-only inspection: {path}")
    return {
        "path": str(path.resolve()),
        "sha256_before": before_sha,
        "sha256_after": after_sha,
        "sheet": sheet_name,
        "physical_row": row_number,
        "safe_row": safe_row,
        "safe_row_sha256": _canonical_sha(safe_row),
    }


def _terminal_evidence(
    db_rows: dict[str, list[dict[str, Any]]], order_id: str
) -> list[dict[str, Any]]:
    evidence: list[dict[str, Any]] = []
    for row in db_rows["fact_order_status_observations"]:
        if _text(row["order_id"]) != order_id:
            continue
        if _text(row["status_internal"]).upper() in TERMINAL_INTERNAL_STATUSES:
            evidence.append(
                {
                    "source_table": "fact_order_status_observations",
                    "row_id": row["id"],
                    "status": _text(row["status_internal"]).upper(),
                    "event_date": _text(row["observed_at"])[:10],
                    "row_sha256": _canonical_sha(row),
                }
            )
    for row in db_rows["order_status_event"]:
        if _text(row["order_id"]) != order_id:
            continue
        if _text(row["stage_code"]).upper() in TERMINAL_INTERNAL_STATUSES:
            evidence.append(
                {
                    "source_table": "order_status_event",
                    "row_id": row["event_id"],
                    "status": _text(row["stage_code"]).upper(),
                    "event_date": _text(row["event_ts"])[:10],
                    "row_sha256": _canonical_sha(row),
                }
            )
    return sorted(
        evidence,
        key=lambda item: (item["event_date"], item["source_table"], str(item["row_id"])),
    )


def _rows_for_order(rows: Iterable[dict[str, Any]], order_id: str) -> list[dict[str, Any]]:
    return [row for row in rows if _text(row.get("order_id")) == order_id]


def build_packet(
    *,
    db_path: Path,
    expected_db_sha256: str,
    workbooks: list[tuple[Path, str]],
    derived_order_id: str,
    raw_order_id: str,
    expected_store: str,
    sheet_name: str,
    expected_row_number: int,
    output_dir: Path,
) -> dict[str, Any]:
    db_path = db_path.expanduser().resolve()
    output_dir = output_dir.expanduser().resolve()
    derived_order_id = _normalize_id(derived_order_id, field="derived_order_id")
    raw_order_id = _normalize_id(raw_order_id, field="raw_order_id")
    expected_store = _normalize_store(expected_store)
    if derived_order_id == raw_order_id:
        raise PacketError("derived_order_id and raw_order_id must differ")
    if len(workbooks) < 2:
        raise PacketError("at least two pinned workbook snapshots are required")
    if output_dir.exists():
        raise PacketError(f"output directory already exists: {output_dir}")

    db_sha_before = _sha256(db_path)
    if db_sha_before != expected_db_sha256:
        raise PacketError(
            f"DB hash mismatch: expected {expected_db_sha256}, got {db_sha_before}"
        )
    workbook_rows = [
        _read_workbook_match(
            path.expanduser().resolve(),
            expected_sha256=expected_sha,
            sheet_name=sheet_name,
            expected_row_number=expected_row_number,
            derived_order_id=derived_order_id,
            raw_order_id=raw_order_id,
            expected_store=expected_store,
        )
        for path, expected_sha in workbooks
    ]
    safe_hashes = {row["safe_row_sha256"] for row in workbook_rows}
    if len(safe_hashes) != 1:
        raise PacketError("pinned workbook snapshots disagree on the safe identity row")
    canonical_safe_row = workbook_rows[0]["safe_row"]
    if canonical_safe_row["OrderID"] != derived_order_id:
        raise PacketError("canonical workbook derived ID does not match requested derived ID")
    if canonical_safe_row["№ заказа"] != raw_order_id:
        raise PacketError("canonical workbook raw ID does not match requested raw ID")

    with _connect_read_only(db_path) as conn:
        integrity = _text(conn.execute("PRAGMA integrity_check").fetchone()[0])
        if integrity != "ok":
            raise PacketError(f"source DB integrity_check failed: {integrity}")
        db_rows = _safe_db_rows(
            conn,
            derived_order_id=derived_order_id,
            raw_order_id=raw_order_id,
        )

    derived_orders = _rows_for_order(db_rows["fact_orders_kaspi"], derived_order_id)
    raw_orders = _rows_for_order(db_rows["fact_orders_kaspi"], raw_order_id)
    derived_anchors = _rows_for_order(
        db_rows["fact_sales_workbook_anchor"], derived_order_id
    )
    raw_anchors = _rows_for_order(db_rows["fact_sales_workbook_anchor"], raw_order_id)
    derived_terminal = _terminal_evidence(db_rows, derived_order_id)
    raw_terminal = _terminal_evidence(db_rows, raw_order_id)

    if not derived_orders or not raw_orders:
        raise PacketError("both derived and raw order IDs must exist in fact_orders_kaspi")
    if len(derived_anchors) != 1 or _normalize_store(derived_anchors[0]["store_code"]) != expected_store:
        raise PacketError("expected exactly one derived workbook anchor in the expected store")
    if raw_anchors:
        raise PacketError("raw order unexpectedly already has a workbook anchor")
    if derived_terminal:
        raise PacketError("derived order unexpectedly has terminal evidence")
    if not raw_terminal:
        raise PacketError("raw order lacks terminal evidence")
    matching_raw_products = [
        row
        for row in raw_orders
        if _normalize_store(row["store_code"]) == expected_store
        and _text(row["sku_id"]) == canonical_safe_row["SKU_ID"]
        and _decimal_equal(row["quantity"], canonical_safe_row["Quantity"])
        and _decimal_equal(row["unit_price_kzt"], canonical_safe_row["Сумма"])
    ]
    if len(matching_raw_products) != 1:
        raise PacketError(
            "raw order does not have exactly one DB product row matching workbook store/SKU/quantity/price"
        )

    evidence = {
        "schema": "derived_order_id_conflict_packet_v2",
        "canonical_hash_version": CANONICAL_HASH_VERSION,
        "verdict": "DERIVED_ORDER_ID_CONFLICT_QUARANTINE",
        "derived_order_id": derived_order_id,
        "raw_order_id": raw_order_id,
        "expected_store": expected_store,
        "workbook_snapshot_count": len(workbook_rows),
        "workbook_safe_row_sha256": next(iter(safe_hashes)),
        "workbook_rows": workbook_rows,
        "db": {
            "path": str(db_path),
            "sha256_before": db_sha_before,
            "integrity_check": integrity,
            "rows": {
                table: [
                    {"row": row, "row_sha256": _canonical_sha(row)} for row in rows
                ]
                for table, rows in db_rows.items()
            },
            "derived_anchor_count": len(derived_anchors),
            "raw_anchor_count": len(raw_anchors),
            "derived_terminal_evidence": derived_terminal,
            "raw_terminal_evidence": raw_terminal,
            "matching_raw_product_row_count": len(matching_raw_products),
        },
        "proofs": {
            "derived_and_raw_ids_differ": True,
            "all_pinned_workbooks_match_one_safe_row": True,
            "derived_anchor_exists": True,
            "raw_anchor_absent": True,
            "derived_terminal_evidence_absent": True,
            "raw_terminal_evidence_present": True,
            "raw_product_identity_matches_workbook": True,
            "production_write_authorized": False,
        },
        "next_action": (
            "Quarantine the derived anchor from publication and cash reasoning; bind no values "
            "and make no production change without a separate exact owner-approved repair manifest."
        ),
    }
    db_sha_after = _sha256(db_path)
    if db_sha_after != db_sha_before:
        raise PacketError("source DB changed during read-only packet build")
    evidence["db"]["sha256_after"] = db_sha_after
    evidence["evidence_sha256"] = _canonical_sha(evidence)

    output_dir.mkdir(parents=True, exist_ok=False)
    evidence_path = output_dir / "evidence.json"
    evidence_path.write_text(
        json.dumps(evidence, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
        encoding="utf-8",
    )
    rows_path = output_dir / "workbook_rows.csv"
    with rows_path.open("x", encoding="utf-8", newline="") as handle:
        fieldnames = [
            "path",
            "sha256_before",
            "sha256_after",
            "sheet",
            "physical_row",
            "safe_row_sha256",
            *SAFE_WORKBOOK_HEADERS,
        ]
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in workbook_rows:
            writer.writerow(
                {
                    "path": row["path"],
                    "sha256_before": row["sha256_before"],
                    "sha256_after": row["sha256_after"],
                    "sheet": row["sheet"],
                    "physical_row": row["physical_row"],
                    "safe_row_sha256": row["safe_row_sha256"],
                    **row["safe_row"],
                }
            )
    closeout_path = output_dir / "CLOSEOUT.md"
    closeout_path.write_text(
        "\n".join(
            [
                "# Derived order-ID conflict forensic closeout",
                "",
                "Gate: GREEN",
                "",
                "Verdict: DERIVED_ORDER_ID_CONFLICT_QUARANTINE",
                "",
                f"- Derived workbook ID: `{derived_order_id}`",
                f"- Raw first-party order ID: `{raw_order_id}`",
                f"- Pinned workbook snapshots: `{len(workbook_rows)}`",
                f"- Identical safe-row SHA-256: `{next(iter(safe_hashes))}`",
                f"- Source DB SHA-256 before/after: `{db_sha_before}`",
                f"- Raw terminal evidence rows: `{len(raw_terminal)}`",
                "- Derived terminal evidence rows: `0`",
                "- Production writes: `0`",
                "",
                "The workbook formula-derived ID conflicts with the immutable raw order ID. "
                "The derived anchor must not be treated as a sale awaiting terminal or cash proof. "
                "No production quarantine/delete is authorized by this packet.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    manifest = {
        "schema": "derived_order_id_conflict_packet_manifest_v2",
        "evidence_sha256": evidence["evidence_sha256"],
        "files": {
            "evidence.json": _sha256(evidence_path),
            "workbook_rows.csv": _sha256(rows_path),
            "CLOSEOUT.md": _sha256(closeout_path),
        },
        "source_db_sha256": db_sha_before,
        "source_workbook_sha256": [row["sha256_before"] for row in workbook_rows],
        "verdict": evidence["verdict"],
    }
    manifest_path = output_dir / "manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
        encoding="utf-8",
    )
    return {
        **evidence,
        "output_dir": str(output_dir),
        "evidence_path": str(evidence_path),
        "workbook_rows_path": str(rows_path),
        "closeout_path": str(closeout_path),
        "manifest_path": str(manifest_path),
        "manifest_file_sha256": _sha256(manifest_path),
    }


def _parse_workbook(value: str) -> tuple[Path, str]:
    try:
        path, expected_sha = value.rsplit("=", 1)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("workbook must be PATH=SHA256") from exc
    if len(expected_sha) != 64 or any(ch not in "0123456789abcdef" for ch in expected_sha):
        raise argparse.ArgumentTypeError("workbook SHA-256 must be 64 lowercase hex characters")
    return Path(path), expected_sha


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, required=True)
    parser.add_argument("--expected-db-sha256", required=True)
    parser.add_argument("--workbook", action="append", type=_parse_workbook, required=True)
    parser.add_argument("--derived-order-id", required=True)
    parser.add_argument("--raw-order-id", required=True)
    parser.add_argument("--expected-store", required=True)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--expected-row", type=int, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    packet = build_packet(
        db_path=args.db,
        expected_db_sha256=args.expected_db_sha256,
        workbooks=args.workbook,
        derived_order_id=args.derived_order_id,
        raw_order_id=args.raw_order_id,
        expected_store=args.expected_store,
        sheet_name=args.sheet,
        expected_row_number=args.expected_row,
        output_dir=args.output_dir,
    )
    print(json.dumps(packet, ensure_ascii=False, sort_keys=True, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
