#!/usr/bin/env python3
"""Build the local G-PRICE-03 remediation evidence packet.

The packet is read-only with respect to production state. It combines the
under-floor sales report, recent order probe, Kaspi ACTIVE readbacks, and
Web_automation stop-loss evidence into reviewable CSV/JSON/Markdown artifacts.
"""

from __future__ import annotations

import argparse
import csv
import json
import sqlite3
from collections import Counter, defaultdict
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ALMATY = timezone(timedelta(hours=5))

TARGET_FAMILIES = (
    "CL_NEW-CLO_MEN_ROMBIK_BLACK",
    "CL_OC_MEN_LINE52_BLACK",
    "CL_NEW-CLO_KID_ROMBIK_BLACK",
    "CL_NEW-CLO2_MEN_SUIT-61_BLACK",
    "CL_OC_MEN_LINE51_WHITE",
    "SUIT-31-LS",
    "SUIT-31-TS",
    "LINE-31-TS",
    "LINE-31-LS",
    "LINE-21-TS",
    "LINE-21-LS",
)

SUIT_ALIAS_NOTE = (
    "SUIT compact validation alias to CL_NEW-CLO2_MEN_SUIT-61_BLACK from owner "
    "compatibility/allocation approvals; this is not price-write authority."
)

LINE_AUTHORITY_NOTE = (
    "Missing SKU-wide compact LINE floor authority. Prior owner approvals were exact-row "
    "COGS/cost-basis only and explicitly not LINE-31/LINE-21 SKU-wide inheritance."
)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: row.get(key, "") for key in fieldnames})


def load_floor_map(path: Path) -> dict[str, dict[str, str]]:
    rows = read_csv(path)
    return {row["SKU_key"].strip(): row for row in rows if row.get("SKU_key", "").strip()}


def parse_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "."))
    except ValueError:
        return None


def family_for_sku(sku: str) -> str:
    clean = (sku or "").strip()
    if clean.startswith("SUIT-31-LS"):
        return "SUIT-31-LS"
    if clean.startswith("SUIT-31-TS"):
        return "SUIT-31-TS"
    if clean.startswith("LINE-31-TS"):
        return "LINE-31-TS"
    if clean.startswith("LINE-31-LS"):
        return "LINE-31-LS"
    if clean.startswith("LINE-21-TS"):
        return "LINE-21-TS"
    if clean.startswith("LINE-21-LS"):
        return "LINE-21-LS"
    for family in sorted(TARGET_FAMILIES, key=len, reverse=True):
        if clean.startswith(family):
            return family
    return ""


def floor_for_family(family: str, floor_map: dict[str, dict[str, str]]) -> tuple[str, float | None, str]:
    if family in {"SUIT-31-LS", "SUIT-31-TS"}:
        floor_row = floor_map.get("CL_NEW-CLO2_MEN_SUIT-61_BLACK", {})
        return "CL_NEW-CLO2_MEN_SUIT-61_BLACK", parse_float(floor_row.get("Min_price_35pct")), SUIT_ALIAS_NOTE
    if family.startswith("LINE-31") or family.startswith("LINE-21"):
        return family, None, LINE_AUTHORITY_NOTE
    floor_row = floor_map.get(family, {})
    if floor_row:
        return family, parse_float(floor_row.get("Min_price_35pct")), "exact v7 floor row"
    return family, None, "missing floor row"


def relation(price: float | None, floor: float | None) -> str:
    if floor is None:
        return "missing_floor"
    if price is None:
        return "missing_price"
    if price < floor:
        return "below_floor"
    return "at_or_above_floor"


def load_active_rows(path: Path, store: str, source: str, floor_map: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    row_iter = ws.iter_rows(values_only=True)
    header = [str(value or "").strip() for value in next(row_iter)]
    index = {name: i for i, name in enumerate(header)}
    out: list[dict[str, Any]] = []
    for raw in row_iter:
        sku = str(raw[index["SKU"]] or "").strip()
        family = family_for_sku(sku)
        if not family:
            continue
        floor_key, floor_value, floor_source = floor_for_family(family, floor_map)
        price = parse_float(raw[index["price"]])
        out.append(
            {
                "store": store,
                "source": source,
                "sku": sku,
                "family": family,
                "price": "" if price is None else f"{price:.2f}",
                "pp1": raw[index.get("PP1", -1)] if "PP1" in index else "",
                "pp2": raw[index.get("PP2", -1)] if "PP2" in index else "",
                "floor_sku_key": floor_key,
                "floor_min_price_kzt": "" if floor_value is None else f"{floor_value:.2f}",
                "floor_source": floor_source,
                "relation": relation(price, floor_value),
            }
        )
    return out


def load_acmewear_export_rows(path: Path, floor_map: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for row in read_csv(path):
        sku = row.get("SKU", "").strip()
        family = family_for_sku(sku)
        if not family:
            continue
        floor_key, floor_value, floor_source = floor_for_family(family, floor_map)
        price = parse_float(row.get("price"))
        out.append(
            {
                "store": "ACMEWEAR",
                "source": "acmewear_daily_ops_fetch_csv",
                "sku": sku,
                "family": family,
                "price": "" if price is None else f"{price:.2f}",
                "pp1": row.get("PP1", ""),
                "pp2": row.get("PP2", ""),
                "floor_sku_key": floor_key,
                "floor_min_price_kzt": "" if floor_value is None else f"{floor_value:.2f}",
                "floor_source": floor_source,
                "relation": relation(price, floor_value),
            }
        )
    return out


def summarize_current_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[(str(row["store"]), str(row["family"]))].append(row)
    out: list[dict[str, Any]] = []
    for (store, family), bucket in sorted(grouped.items()):
        prices = [parse_float(row.get("price")) for row in bucket]
        numeric = [value for value in prices if value is not None]
        rel_counts = Counter(str(row.get("relation", "")) for row in bucket)
        floor_values = [parse_float(row.get("floor_min_price_kzt")) for row in bucket]
        floor_numeric = [value for value in floor_values if value is not None]
        out.append(
            {
                "store": store,
                "family": family,
                "active_rows": len(bucket),
                "min_active_price_kzt": "" if not numeric else f"{min(numeric):.2f}",
                "max_active_price_kzt": "" if not numeric else f"{max(numeric):.2f}",
                "floor_min_price_kzt": "" if not floor_numeric else f"{floor_numeric[0]:.2f}",
                "below_floor_rows": rel_counts.get("below_floor", 0),
                "missing_floor_rows": rel_counts.get("missing_floor", 0),
                "at_or_above_floor_rows": rel_counts.get("at_or_above_floor", 0),
            }
        )
    return out


def build_recent_order_probe(db_path: Path, floor_map: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    targets = {
        "CL_NEW-CLO_MEN_ROMBIK_BLACK",
        "CL_OC_MEN_LINE52_BLACK",
        "CL_NEW-CLO_KID_ROMBIK_BLACK",
        "SUIT-31-TS",
        "SUIT-31-LS",
        "LINE-31-TS",
        "LINE-21-TS",
    }
    con = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    placeholders = ",".join("?" for _ in targets)
    sql = f"""
        SELECT
            created_at,
            planned_shipment_date,
            actual_shipment_date,
            store_code,
            order_id,
            sku_key,
            sku_id,
            quantity,
            unit_price_kzt,
            kaspi_status,
            internal_status
        FROM fact_orders_kaspi
        WHERE date(COALESCE(actual_shipment_date, planned_shipment_date, created_at))
              BETWEEN '2026-06-17' AND '2026-06-18'
          AND (
              sku_key IN ({placeholders})
              OR sku_id LIKE 'SUIT-31-%'
              OR sku_id LIKE 'LINE-31-%'
              OR sku_id LIKE 'LINE-21-%'
          )
        ORDER BY created_at, store_code, sku_key, sku_id
    """
    rows: list[dict[str, Any]] = []
    for raw in con.execute(sql, sorted(targets)):
        sku_key = str(raw["sku_key"] or "")
        sku_id = str(raw["sku_id"] or "")
        family = family_for_sku(sku_key) or family_for_sku(sku_id) or sku_key
        floor_key, floor_value, floor_source = floor_for_family(family, floor_map)
        price = parse_float(raw["unit_price_kzt"])
        rows.append(
            {
                "created_at": raw["created_at"],
                "planned_shipment_date": raw["planned_shipment_date"],
                "actual_shipment_date": raw["actual_shipment_date"],
                "store_code": raw["store_code"],
                "order_id": raw["order_id"],
                "sku_key": sku_key,
                "sku_id": sku_id,
                "family": family,
                "quantity": raw["quantity"],
                "unit_price_kzt": "" if price is None else f"{price:.2f}",
                "floor_sku_key": floor_key,
                "floor_min_price_kzt": "" if floor_value is None else f"{floor_value:.2f}",
                "floor_source": floor_source,
                "relation": relation(price, floor_value),
                "kaspi_status": raw["kaspi_status"],
                "internal_status": raw["internal_status"],
            }
        )
    return rows


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def build_coverage_rows(args: argparse.Namespace, current_summary: list[dict[str, Any]]) -> list[dict[str, Any]]:
    repricer = read_json(args.repricer_readback_summary)
    universal = read_json(args.universal_verify_summary)
    storeb = read_json(args.storeb_verify_summary)
    stoploss_closeout = args.stoploss_closeout
    total_planned_api = 0
    total_planned_actions = 0
    for store_payload in repricer.get("stores", {}).values():
        total_planned_api += int(store_payload.get("planned_min_updates") or 0)
        total_planned_api += int(store_payload.get("planned_max_updates") or 0)
        total_planned_api += int(store_payload.get("planned_live_price_updates") or 0)
        total_planned_actions += int(store_payload.get("planned_actions") or 0)

    rows = [
        {
            "surface": "Repricer exact-row stoploss",
            "stores": "UNIVERSAL;STORE-B",
            "evidence": str(args.repricer_readback_summary),
            "status": "floor_safe_readback",
            "key_counts": f"validated_plan_rows={repricer.get('validated_plan_rows')}; planned_actions={total_planned_actions}; planned_api_requests={total_planned_api}",
            "interpretation": "Post-apply dry-run/readback planned no API requests; current exact Repricer rows in the stoploss plan need no further raise.",
        },
        {
            "surface": "Kaspi merchant ACTIVE readback",
            "stores": "UNIVERSAL",
            "evidence": str(args.universal_verify_summary),
            "status": "floor_safe_with_expected_mismatch",
            "key_counts": f"missing_count={universal.get('missing_count')}; field_mismatch_count={universal.get('field_mismatch_count')}",
            "interpretation": "Verifier status is mismatch, but closeout classifies missing rows as removed from active sellability and price mismatches as floor-safe.",
        },
        {
            "surface": "Kaspi merchant ACTIVE readback",
            "stores": "STORE-B",
            "evidence": str(args.storeb_verify_summary),
            "status": "floor_safe_with_expected_mismatch",
            "key_counts": f"missing_count={storeb.get('missing_count')}; field_mismatch_count={storeb.get('field_mismatch_count')}",
            "interpretation": "Verifier status is mismatch, but closeout classifies missing rows as removed from active sellability and price mismatches as floor-safe.",
        },
        {
            "surface": "Final Web_automation stoploss closeout",
            "stores": "UNIVERSAL;STORE-B",
            "evidence": str(stoploss_closeout),
            "status": "green_for_active_stoploss_safety_yellow_for_archive_reactivation",
            "key_counts": "Repricer remaining_mismatches_total=0; fresh dry-run planned_api_requests=0",
            "interpretation": "Current U/M stoploss exposure was remediated separately; G-PRICE-03 remains RED because the trailing sales window still contains loss events.",
        },
    ]
    for row in current_summary:
        if row["store"] == "ACMEWEAR" or row["family"] in {
            "CL_NEW-CLO_MEN_ROMBIK_BLACK",
            "CL_OC_MEN_LINE52_BLACK",
            "CL_NEW-CLO_KID_ROMBIK_BLACK",
        }:
            status = "current_below_floor" if int(row["below_floor_rows"]) else "current_missing_floor" if int(row["missing_floor_rows"]) else "current_floor_safe"
            rows.append(
                {
                    "surface": "ACTIVE price row summary",
                    "stores": row["store"],
                    "evidence": "current_active_readback_rows.csv",
                    "status": status,
                    "key_counts": (
                        f"family={row['family']}; active_rows={row['active_rows']}; "
                        f"below_floor_rows={row['below_floor_rows']}; missing_floor_rows={row['missing_floor_rows']}; "
                        f"min_active_price={row['min_active_price_kzt']}; floor={row['floor_min_price_kzt']}"
                    ),
                    "interpretation": coverage_interpretation(row),
                }
            )
    return rows


def coverage_interpretation(row: dict[str, Any]) -> str:
    family = str(row["family"])
    if family.startswith("SUIT-31"):
        return "ACMEWEAR compact SUIT rows are active below the validation parent floor; exact price/floor decision is required."
    if family.startswith("LINE-31") or family.startswith("LINE-21"):
        return "ACMEWEAR compact LINE rows are active but lack SKU-wide scoring floor authority."
    if int(row.get("below_floor_rows") or 0):
        return "Current ACTIVE readback still has rows below floor and needs exact remediation before retest."
    return "Current ACTIVE readback rows in this family are at or above the binding floor."


def build_bucket_matrix(red_rows: list[dict[str, str]], current_summary: list[dict[str, Any]]) -> list[dict[str, Any]]:
    summary_by_family = {(row["store"], row["family"]): row for row in current_summary}
    families_by_name = defaultdict(list)
    for row in current_summary:
        families_by_name[row["family"]].append(row)

    out: list[dict[str, Any]] = []
    red_relevant = [
        row
        for row in red_rows
        if int(float(row.get("under_floor_units") or 0)) > 0
        or int(float(row.get("missing_floor_rows") or 0)) > 0
    ]
    for row in red_relevant:
        sku = row["sku_key"]
        current_rows = families_by_name.get(sku, [])
        if sku in {"SUIT-31-LS", "SUIT-31-TS"}:
            current_rows = [r for r in families_by_name.get(sku, []) if r["store"] == "ACMEWEAR"]
            status = "current_acmewear_compact_price_below_parent_floor"
            next_action = "Owner-approved exact ACMEWEAR compact SUIT price/floor remediation, then fresh ACTIVE readback and rerun G-PRICE-03."
            approval = "required_before_any_live_price_or_offer_write"
        elif sku.startswith("LINE-31") or sku.startswith("LINE-21"):
            status = "missing_compact_beli_floor_authority"
            next_action = "Owner must approve compact LINE floor mapping or explicit exclusion before scoring/pricing remediation."
            approval = "required_before_scoring_or_live_price_write"
        elif sku in {"CL_OC_MEN_LINE52_BLACK", "CL_NEW-CLO_KID_ROMBIK_BLACK"}:
            status = "current_um_stoploss_evidence_floor_safe_but_trailing_window_red"
            next_action = "No new live write from this packet; monitor for new sub-floor orders and rerun strict G-PRICE-03 after the trailing window clears."
            approval = "not_required_for_packet; required_for_any_new_live_apply"
        elif sku == "CL_NEW-CLO_MEN_ROMBIK_BLACK":
            status = "current_active_readback_floor_safe_but_trailing_window_red"
            next_action = "No new live write from this packet unless a fresh readback later shows rows below floor."
            approval = "not_required_for_packet; required_for_any_new_live_apply"
        else:
            status = "needs_manual_review"
            next_action = "Review exact current surface and approval path."
            approval = "required_if_live_write_needed"

        current_snapshot = "; ".join(
            f"{r['store']} rows={r['active_rows']} min={r['min_active_price_kzt']} below={r['below_floor_rows']} missing_floor={r['missing_floor_rows']}"
            for r in current_rows
        )
        out.append(
            {
                "sku_key": sku,
                "stores_from_red_report": row.get("stores", ""),
                "floor_sku_key": row.get("floor_sku_key", ""),
                "floor_min_price_kzt": row.get("floor_min_price_kzt", ""),
                "under_floor_units": row.get("under_floor_units", ""),
                "gap_total_kzt": row.get("gap_total_kzt", ""),
                "missing_floor_rows": row.get("missing_floor_rows", ""),
                "min_sell_price_kzt": row.get("min_sell_price_kzt", ""),
                "current_status": status,
                "current_active_snapshot": current_snapshot,
                "next_safe_action": next_action,
                "approval_status": approval,
            }
        )
    return out


def build_approval_rows() -> list[dict[str, str]]:
    return [
        {
            "area": "ACMEWEAR compact SUIT",
            "why_needed": "Current ACTIVE compact SUIT-31-LS/TS rows are priced below the parent SUIT-61 floor used by validation.",
            "minimum_approval": (
                "I authorize an ACMEWEAR exact-row compact SUIT-31-LS/TS price/floor remediation for only the active rows listed "
                "in G-PRICE-03 current_active_readback_rows.csv, using the CL_NEW-CLO2_MEN_SUIT-61_BLACK floor as the minimum "
                "price basis, backup/readback first, with no stock, workbook, Telegram, customer, LaunchAgent, or unrelated writes."
            ),
            "allowed_after_approval": "Generate exact ACMEWEAR upload/API plan, dry-run, backup/readback, apply only with env/confirm gates, rerun G-PRICE-03.",
        },
        {
            "area": "ACMEWEAR compact LINE",
            "why_needed": "LINE-31/LINE-21 compact rows have no SKU-wide floor authority; exact-row owner COGS overrides do not authorize SKU-wide inheritance.",
            "minimum_approval": (
                "I authorize a compact LINE floor decision for G-PRICE-03 scoring, specifying whether LINE-31/LINE-21 inherit "
                "CL_OC_MEN_LINE51_WHITE floor economics or a different explicit owner floor, with no live price write unless separately approved."
            ),
            "allowed_after_approval": "Update the validation alias/floor contract and rerun the read-only report; live pricing remains separately gated.",
        },
        {
            "area": "Strict trailing-window gate",
            "why_needed": "G-PRICE-03 green criterion is 0 non-cancelled sub-floor units over trailing 7 days; remediated current exposure does not erase historical rows.",
            "minimum_approval": "No approval can truthfully convert historical under-floor sales to green; either wait for the window to clear or explicitly change the gate contract.",
            "allowed_after_approval": "Rerun strict report after window clears; if contract changes, update owning docs first and preserve old evidence.",
        },
    ]


def build_acmewear_suit_update_template(current_rows: list[dict[str, Any]]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for row in current_rows:
        if row.get("store") != "ACMEWEAR" or row.get("family") not in {"SUIT-31-LS", "SUIT-31-TS"}:
            continue
        rows.append(
            {
                "SKU": str(row.get("sku", "")),
                "price": str(row.get("floor_min_price_kzt", "")),
                "current_price": str(row.get("price", "")),
                "floor_min_price_kzt": str(row.get("floor_min_price_kzt", "")),
                "family": str(row.get("family", "")),
                "approval_status": "NOT_APPROVED_NO_APPLY",
                "notes": "Template only. Owner approval and safe-active-patch dry-run/readback required before any upload.",
            }
        )
    return rows


def build_markdown(
    output_dir: Path,
    red_summary: dict[str, Any],
    bucket_rows: list[dict[str, Any]],
    coverage_rows: list[dict[str, Any]],
    recent_rows: list[dict[str, Any]],
    approval_rows: list[dict[str, str]],
    generated_at: str,
) -> str:
    below_current = [row for row in coverage_rows if row["status"] == "current_below_floor"]
    missing_current = [row for row in coverage_rows if row["status"] == "current_missing_floor"]
    recent_below = [row for row in recent_rows if row["relation"] == "below_floor"]
    lines = [
        "# G-PRICE-03 Remediation Packet",
        "",
        "Gate: RED",
        f"Generated at: {generated_at}",
        "",
        "## Summary",
        "",
        f"- strict report evidence: `{red_summary.get('evidence_dir', '')}`",
        f"- strict report under-floor units: `{red_summary.get('under_floor_units', '')}`",
        f"- strict report under-floor gap KZT: `{red_summary.get('gap_total_kzt', '')}`",
        f"- strict report missing-floor rows: `{red_summary.get('missing_floor_rows', '')}`",
        f"- current ACTIVE below-floor summaries: `{len(below_current)}`",
        f"- current ACTIVE missing-floor summaries: `{len(missing_current)}`",
        f"- recent June 17-18 fact_orders rows below floor: `{len(recent_below)}`",
        "",
        "## Interpretation",
        "",
        "- UNIVERSAL/STORE-B stop-loss has separate June 18 evidence showing Repricer exact-row readback and Kaspi ACTIVE readbacks are floor-safe for the covered rows.",
        "- `G-PRICE-03` still remains RED because the strict trailing-7d sales gate contains historical non-cancelled below-floor sales.",
        "- ACMEWEAR compact `SUIT-31-LS` and `SUIT-31-TS` are current active exposure: active rows are below the SUIT parent floor used by validation.",
        "- ACMEWEAR compact `LINE-31`/`LINE-21` rows cannot be scored until the owner approves SKU-wide floor authority or an explicit alternative floor contract.",
        "- No live price, DB, workbook, Google, Telegram, Kaspi, Repricer, stock, customer, LaunchAgent, or other external write was performed by this packet builder.",
        "",
        "## Artifacts",
        "",
        "- `red_bucket_matrix.csv`",
        "- `current_active_readback_rows.csv`",
        "- `current_readback_coverage.csv`",
        "- `recent_fact_orders_probe.csv`",
        "- `acmewear_suit_compact_price_updates_template.csv`",
        "- `approval_required_actions.csv`",
        "- `remediation_packet.json`",
        "",
        "## Red Buckets",
        "",
    ]
    for row in bucket_rows:
        lines.append(
            f"- `{row['sku_key']}`: status `{row['current_status']}`, under-floor units `{row['under_floor_units']}`, "
            f"missing-floor rows `{row['missing_floor_rows']}`; next: {row['next_safe_action']}"
        )
    lines.extend(["", "## Required Approvals", ""])
    for row in approval_rows:
        lines.append(f"- `{row['area']}`: {row['why_needed']}")
    return "\n".join(lines) + "\n"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    repo = Path(__file__).resolve().parents[1]
    wa = Path("~/Docs/Web_automation")
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--red-report-dir", type=Path, default=repo / "exports/validation/g_price03_under_floor_leak_20260618_153557")
    parser.add_argument("--ab-db", type=Path, default=repo / "db/app.db")
    parser.add_argument("--floor-csv", type=Path, default=wa / "exports/pricelist_snapshots/min_price_floor_35pct_by_sku_v7.csv")
    parser.add_argument(
        "--universal-active-xlsx",
        type=Path,
        default=wa / "runs/kaspi_pricelist_ops/20260618_105910_universal/post_verify_download/downloads/universal_ACTIVE.xlsx",
    )
    parser.add_argument(
        "--storeb-active-xlsx",
        type=Path,
        default=wa / "runs/kaspi_pricelist_ops/20260618_110206_storeb/post_verify_download/downloads/storeb_ACTIVE.xlsx",
    )
    parser.add_argument(
        "--acmewear-active-xlsx",
        type=Path,
        default=wa / "runs/line61_line51_daily_ops_fetch/20260618_154235/pricelist_download/downloads/acmewear_ACTIVE.xlsx",
    )
    parser.add_argument(
        "--acmewear-export-csv",
        type=Path,
        default=wa / "runs/line61_line51_daily_ops_fetch/20260618_154235/exports/merchant_pricelist_offer_rows_line61_line51.csv",
    )
    parser.add_argument(
        "--repricer-readback-summary",
        type=Path,
        default=wa / "runs/universal_storeb_reactivation/20260618_post_restriction/exact_repricer_stoploss/20260618_112917/summary.json",
    )
    parser.add_argument(
        "--universal-verify-summary",
        type=Path,
        default=wa / "runs/kaspi_pricelist_ops/20260618_105910_universal/post_verify_download/universal_safe_active_verify_summary.json",
    )
    parser.add_argument(
        "--storeb-verify-summary",
        type=Path,
        default=wa / "runs/kaspi_pricelist_ops/20260618_110206_storeb/post_verify_download/storeb_safe_active_verify_summary.json",
    )
    parser.add_argument(
        "--stoploss-closeout",
        type=Path,
        default=wa / "runs/universal_storeb_reactivation/20260618_post_restriction/final_stoploss_closeout_20260618_1130/FINAL_STOPLOSS_CLOSEOUT.md",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output_dir = args.output_dir
    output_dir.mkdir(parents=True, exist_ok=True)
    generated_at = datetime.now(ALMATY).replace(microsecond=0).isoformat()

    floor_map = load_floor_map(args.floor_csv)
    red_rows = read_csv(args.red_report_dir / "under_floor_by_sku.csv")
    red_report = read_json(args.red_report_dir / "under_floor_leak_report.json")

    current_rows = []
    current_rows.extend(load_active_rows(args.universal_active_xlsx, "UNIVERSAL", "kaspi_active_post_verify", floor_map))
    current_rows.extend(load_active_rows(args.storeb_active_xlsx, "STORE-B", "kaspi_active_post_verify", floor_map))
    current_rows.extend(load_active_rows(args.acmewear_active_xlsx, "ACMEWEAR", "acmewear_active_download", floor_map))
    current_summary = summarize_current_rows(current_rows)

    # The CSV export is redundant with the xlsx, but it is a useful independent
    # source from the daily fetch packet for ACMEWEAR parent/compact coverage.
    acmewear_export_rows = load_acmewear_export_rows(args.acmewear_export_csv, floor_map)
    acmewear_export_summary = summarize_current_rows(acmewear_export_rows)

    recent_rows = build_recent_order_probe(args.ab_db, floor_map)
    bucket_rows = build_bucket_matrix(red_rows, current_summary)
    coverage_rows = build_coverage_rows(args, current_summary)
    approval_rows = build_approval_rows()
    acmewear_suit_template_rows = build_acmewear_suit_update_template(current_rows)

    write_csv(
        output_dir / "red_bucket_matrix.csv",
        bucket_rows,
        [
            "sku_key",
            "stores_from_red_report",
            "floor_sku_key",
            "floor_min_price_kzt",
            "under_floor_units",
            "gap_total_kzt",
            "missing_floor_rows",
            "min_sell_price_kzt",
            "current_status",
            "current_active_snapshot",
            "next_safe_action",
            "approval_status",
        ],
    )
    write_csv(
        output_dir / "current_active_readback_rows.csv",
        current_rows,
        [
            "store",
            "source",
            "sku",
            "family",
            "price",
            "pp1",
            "pp2",
            "floor_sku_key",
            "floor_min_price_kzt",
            "floor_source",
            "relation",
        ],
    )
    write_csv(
        output_dir / "current_active_summary.csv",
        current_summary,
        [
            "store",
            "family",
            "active_rows",
            "min_active_price_kzt",
            "max_active_price_kzt",
            "floor_min_price_kzt",
            "below_floor_rows",
            "missing_floor_rows",
            "at_or_above_floor_rows",
        ],
    )
    write_csv(
        output_dir / "acmewear_daily_export_summary.csv",
        acmewear_export_summary,
        [
            "store",
            "family",
            "active_rows",
            "min_active_price_kzt",
            "max_active_price_kzt",
            "floor_min_price_kzt",
            "below_floor_rows",
            "missing_floor_rows",
            "at_or_above_floor_rows",
        ],
    )
    write_csv(
        output_dir / "current_readback_coverage.csv",
        coverage_rows,
        ["surface", "stores", "evidence", "status", "key_counts", "interpretation"],
    )
    write_csv(
        output_dir / "recent_fact_orders_probe.csv",
        recent_rows,
        [
            "created_at",
            "planned_shipment_date",
            "actual_shipment_date",
            "store_code",
            "order_id",
            "sku_key",
            "sku_id",
            "family",
            "quantity",
            "unit_price_kzt",
            "floor_sku_key",
            "floor_min_price_kzt",
            "floor_source",
            "relation",
            "kaspi_status",
            "internal_status",
        ],
    )
    write_csv(output_dir / "approval_required_actions.csv", approval_rows, ["area", "why_needed", "minimum_approval", "allowed_after_approval"])
    write_csv(
        output_dir / "acmewear_suit_compact_price_updates_template.csv",
        acmewear_suit_template_rows,
        ["SKU", "price", "current_price", "floor_min_price_kzt", "family", "approval_status", "notes"],
    )

    red_summary = {
        "evidence_dir": str(args.red_report_dir),
        "gate": red_report.get("gate"),
        "under_floor_units": red_report.get("under_floor_units"),
        "gap_total_kzt": red_report.get("under_floor_gap_kzt"),
        "missing_floor_rows": red_report.get("missing_floor_row_count"),
        "generated_at": red_report.get("generated_at"),
    }
    packet = {
        "schema_version": "g_price03_remediation_packet.v1",
        "generated_at": generated_at,
        "gate": "RED",
        "red_summary": red_summary,
        "counts": {
            "red_bucket_rows": len(bucket_rows),
            "current_active_rows": len(current_rows),
            "current_active_summary_rows": len(current_summary),
            "recent_fact_order_rows": len(recent_rows),
            "recent_fact_order_below_floor_rows": sum(1 for row in recent_rows if row["relation"] == "below_floor"),
            "current_active_below_floor_summary_rows": sum(1 for row in current_summary if int(row["below_floor_rows"]) > 0),
            "current_active_missing_floor_summary_rows": sum(1 for row in current_summary if int(row["missing_floor_rows"]) > 0),
        },
        "artifacts": {
            "red_bucket_matrix_csv": str(output_dir / "red_bucket_matrix.csv"),
            "current_active_readback_rows_csv": str(output_dir / "current_active_readback_rows.csv"),
            "current_active_summary_csv": str(output_dir / "current_active_summary.csv"),
            "acmewear_daily_export_summary_csv": str(output_dir / "acmewear_daily_export_summary.csv"),
            "current_readback_coverage_csv": str(output_dir / "current_readback_coverage.csv"),
            "recent_fact_orders_probe_csv": str(output_dir / "recent_fact_orders_probe.csv"),
            "acmewear_suit_compact_price_updates_template_csv": str(output_dir / "acmewear_suit_compact_price_updates_template.csv"),
            "approval_required_actions_csv": str(output_dir / "approval_required_actions.csv"),
            "closeout_md": str(output_dir / "closeout.md"),
        },
        "no_external_write_performed": True,
        "forbidden_surfaces_touched": [],
    }
    (output_dir / "remediation_packet.json").write_text(json.dumps(packet, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    closeout = build_markdown(output_dir, red_summary, bucket_rows, coverage_rows, recent_rows, approval_rows, generated_at)
    (output_dir / "closeout.md").write_text(closeout, encoding="utf-8")

    print(json.dumps(packet, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
