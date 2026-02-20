#!/usr/bin/env python3
"""
Sync ABC_View sheet from Inventory_Core workbook into abc_view_cache table.

Read-only Excel access. Writes only to SQLite DB.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from core.db import get_connection
from core.paths import data_path

DEFAULT_XL = data_path("excel", "Inventory_Core_V18.1_V2.xlsx")
DEFAULT_DB = data_path("db", "app.db")
DEFAULT_SHEET = "ABC_View"

TEXT_COLUMNS = {
    "SKU_key",
    "Product_Type",
    "Status",
    "Lifecycle_flag",
    "Notes",
}


def _normalize_header(value: object, idx: int) -> str:
    if value is None:
        return f"col_{idx}"
    name = str(value).strip().replace(" ", "_")
    name = "".join(ch for ch in name if ch.isalnum() or ch == "_")
    return name or f"col_{idx}"


def _resolve_sheet(wb, target: str) -> str:
    for name in wb.sheetnames:
        if name.lower() == target.lower():
            return name
    raise SystemExit(f"Sheet '{target}' not found. Available: {wb.sheetnames}")


def read_abc_view(path: Path, sheet: str) -> tuple[list[str], list[dict]]:
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    sheet_name = _resolve_sheet(wb, sheet)
    ws = wb[sheet_name]

    rows = ws.iter_rows(values_only=True)
    header_raw = next(rows, None)
    if not header_raw:
        wb.close()
        raise SystemExit("ABC_View sheet has no header row.")

    headers = [_normalize_header(h, i) for i, h in enumerate(header_raw, start=1)]
    data = []
    for row in rows:
        if not row or all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            continue
        record = {col: val for col, val in zip(headers, row)}
        data.append(record)

    wb.close()
    return headers, data


def write_to_db(db_path: Path, headers: list[str], rows: list[dict]) -> dict:
    conn = get_connection(db_path)
    try:
        conn.execute("DROP TABLE IF EXISTS abc_view_cache")
        columns_sql = []
        for col in headers:
            col_type = "TEXT" if col in TEXT_COLUMNS else "REAL"
            columns_sql.append(f'"{col}" {col_type}')
        conn.execute(f"CREATE TABLE abc_view_cache ({', '.join(columns_sql)})")
        if "SKU_key" in headers:
            conn.execute("CREATE INDEX IF NOT EXISTS idx_abc_view_sku ON abc_view_cache (SKU_key)")

        if rows:
            placeholders = ",".join(["?"] * len(headers))
            sql = f"INSERT INTO abc_view_cache ({', '.join([f'"{c}"' for c in headers])}) VALUES ({placeholders})"
            values = [[row.get(h) for h in headers] for row in rows]
            conn.executemany(sql, values)
        conn.commit()
        return {"rows": len(rows)}
    finally:
        conn.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Sync ABC_View sheet into SQLite")
    parser.add_argument("--file", type=Path, default=DEFAULT_XL, help="Inventory_Core workbook path")
    parser.add_argument("--sheet", type=str, default=DEFAULT_SHEET, help="Sheet name (default: ABC_View)")
    parser.add_argument("--db", type=Path, default=DEFAULT_DB, help="SQLite db path")
    parser.add_argument("--dry-run", action="store_true", help="Read only; do not write DB")
    args = parser.parse_args()

    if not args.file.exists():
        raise SystemExit(f"Excel file not found: {args.file}")

    headers, rows = read_abc_view(args.file, args.sheet)
    if args.dry_run:
        print(f"Read {len(rows)} rows with {len(headers)} columns from {args.sheet}")
        print("Columns:", headers)
        return

    result = write_to_db(args.db, headers, rows)
    print(f"✅ abc_view_cache updated: {result['rows']} rows")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit(1)
