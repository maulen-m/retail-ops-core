#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import shutil
import sqlite3
import sys
import tempfile
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Mapping, Optional

from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.integrations.kaspi_order_stage import (  # noqa: E402
    classify_kaspi_stage_from_db_row,
    kaspi_order_to_russian_status,
    stage_to_crm_indicators,
)


DATE_FMT_DDMMYYYY = "%d.%m.%Y"
DATE_FMT_ISO = "%Y-%m-%d"


@dataclass
class DbPatch:
    status: str
    issued_flag: str
    status_change_date: str
    planned_handover_date: str
    seller_fee: str


def _now_ts() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def normalize_order_id(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    if text.endswith(".0"):
        candidate = text[:-2]
        if candidate.isdigit():
            return candidate
    if text.isdigit():
        return text
    return "".join(ch for ch in text if ch.isdigit())


def parse_date_arg(value: str) -> date:
    return datetime.strptime(value, DATE_FMT_ISO).date()


def parse_date_any(value: Any) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        try:
            num = float(value)
        except (TypeError, ValueError):
            num = 0.0
        # Excel serial date window used by CRM workbook.
        if 40000 <= num <= 60000:
            base = datetime(1899, 12, 30)
            return (base + timedelta(days=int(num))).date()
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    text = str(value).strip()
    if not text:
        return None
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M:%S%z",
        "%d.%m.%Y",
        "%d.%m.%Y %H:%M:%S",
    ):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        return None


def parse_datetime_any(value: Any) -> Optional[datetime]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = str(value).strip()
    if not text:
        return None
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M:%S%z",
        "%Y-%m-%d",
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _float_or_none(value: Any) -> Optional[float]:
    if value is None:
        return None
    text = str(value).strip().replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def is_blank_or_zero(value: Any) -> bool:
    if is_blank(value):
        return True
    num = _float_or_none(value)
    if num is None:
        return False
    return abs(num) < 1e-9


def fmt_ddmmyyyy(value: Optional[date]) -> str:
    if value is None:
        return ""
    return value.strftime(DATE_FMT_DDMMYYYY)


def fmt_numeric(value: Any) -> str:
    number = _float_or_none(value)
    if number is None:
        return ""
    if abs(number) < 1e-9:
        return ""
    rounded = round(number)
    if abs(number - rounded) < 1e-9:
        return str(int(rounded))
    text = f"{number:.6f}".rstrip("0").rstrip(".")
    return text


def latest_non_empty(values: set[date]) -> str:
    if not values:
        return ""
    return fmt_ddmmyyyy(max(values))


def in_range(value: Optional[date], date_from: date, date_to: date) -> bool:
    if value is None:
        return False
    return date_from <= value <= date_to


def _backup_file(src: Path, backup_dir: Path, ts: str) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    dst = backup_dir / f"{src.name}.{ts}.bak"
    shutil.copy2(src, dst)
    return dst


def _atomic_write_text(path: Path, content: str, encoding: str = "utf-8") -> None:
    fd, tmp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=str(path.parent))
    tmp_path = Path(tmp_name)
    try:
        with open(fd, "w", encoding=encoding, newline="") as fh:
            fh.write(content)
        tmp_path.replace(path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink(missing_ok=True)


def _atomic_save_workbook(path: Path, save_fn) -> None:
    fd, tmp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=str(path.parent))
    tmp_path = Path(tmp_name)
    Path(tmp_name).unlink(missing_ok=True)
    try:
        save_fn(tmp_path)
        tmp_path.replace(path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink(missing_ok=True)


def _archive_row_missing_period_values(row: Mapping[str, str]) -> bool:
    if is_blank_or_zero(row.get("Стоимость доставки для продавца", "")):
        return True
    if is_blank(row.get("Статус", "")):
        return True
    if is_blank(row.get("Выдал", "")):
        return True
    if is_blank(row.get("Дата изменения статуса", "")):
        return True
    if is_blank(row.get("Плановая дата передачи курьеру", "")):
        return True
    return False


def load_archive_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        header = list(reader.fieldnames or [])
        rows = [dict(row) for row in reader]
    return header, rows


def build_target_order_ids(rows: list[dict[str, str]], date_from: date, date_to: date) -> set[str]:
    out: set[str] = set()
    for row in rows:
        sale_date = parse_date_any(row.get("Дата поступления заказа"))
        if not in_range(sale_date, date_from, date_to):
            continue
        if not _archive_row_missing_period_values(row):
            continue
        oid = normalize_order_id(row.get("№ заказа"))
        if oid:
            out.add(oid)
    return out


def _build_db_patch_from_row(row: Mapping[str, Any]) -> DbPatch:
    attrs = {
        "state": row.get("kaspi_status", ""),
        "status": row.get("kaspi_status_detail", ""),
    }
    stage = classify_kaspi_stage_from_db_row(row)
    status = kaspi_order_to_russian_status({"attributes": attrs})
    indicators = stage_to_crm_indicators(stage)

    status_change = fmt_ddmmyyyy(parse_datetime_any(row.get("status_updated_at")))
    planned_handover = (
        parse_date_any(row.get("courier_transmission_planning_date"))
        or parse_date_any(row.get("planned_shipment_date"))
        or parse_date_any(row.get("courier_transmission_date"))
    )
    seller_fee = fmt_numeric(row.get("delivery_cost_for_seller"))
    return DbPatch(
        status=status,
        issued_flag=indicators.get("Выдал", ""),
        status_change_date=status_change,
        planned_handover_date=fmt_ddmmyyyy(planned_handover),
        seller_fee=seller_fee,
    )


def load_db_patches(db_path: Path, order_ids: set[str]) -> dict[str, DbPatch]:
    if not order_ids:
        return {}
    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    try:
        rows: list[dict[str, Any]] = []
        ids = sorted(order_ids)
        chunk_size = 800
        for i in range(0, len(ids), chunk_size):
            chunk = ids[i : i + chunk_size]
            placeholders = ",".join("?" for _ in chunk)
            query = f"""
                SELECT
                    order_id,
                    kaspi_status,
                    kaspi_status_detail,
                    status_updated_at,
                    planned_shipment_date,
                    courier_transmission_planning_date,
                    courier_transmission_date,
                    delivery_cost_for_seller
                FROM fact_orders_kaspi
                WHERE order_id IN ({placeholders})
            """
            cur = conn.execute(query, chunk)
            rows.extend(dict(r) for r in cur.fetchall())

        latest_by_order: dict[str, dict[str, Any]] = {}
        for row in rows:
            oid = normalize_order_id(row.get("order_id"))
            if not oid:
                continue
            current = latest_by_order.get(oid)
            if current is None:
                latest_by_order[oid] = row
                continue
            current_dt = parse_datetime_any(current.get("status_updated_at")) or datetime.min
            new_dt = parse_datetime_any(row.get("status_updated_at")) or datetime.min
            if new_dt >= current_dt:
                latest_by_order[oid] = row

        return {oid: _build_db_patch_from_row(row) for oid, row in latest_by_order.items()}
    finally:
        conn.close()


def apply_crm_updates(
    ws,
    db_patches: Mapping[str, DbPatch],
    target_order_ids: set[str],
    date_from: date,
    date_to: date,
) -> tuple[dict[str, int], dict[str, set[date]], dict[str, dict[str, str]]]:
    header = {str(ws.cell(row=1, column=c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    order_col_primary = header.get("OrderID")
    order_col_alt = header.get("№ заказа")
    date_col = header.get("Date")
    status_col = header.get("Статус")
    issued_col = header.get("Выдал")
    status_change_col = header.get("Дата изменения статуса")
    planned_col = header.get("Плановая дата передачи курьеру")
    seller_fee_col = header.get("Стоимость доставки для продавца")
    delivery_fee_col = header.get("Delivery_fee_kzt")

    required = {
        "OrderID_or_№заказа": order_col_primary or order_col_alt,
        "Date": date_col,
        "Статус": status_col,
        "Выдал": issued_col,
        "Дата изменения статуса": status_change_col,
        "Плановая дата передачи курьеру": planned_col,
        "Стоимость доставки для продавца": seller_fee_col,
        "Delivery_fee_kzt": delivery_fee_col,
    }
    missing_headers = [name for name, col in required.items() if col is None]
    if missing_headers:
        raise RuntimeError(f"CRM workbook missing required columns: {missing_headers}")

    updates = defaultdict(int)
    crm_send_dates: dict[str, set[date]] = defaultdict(set)
    crm_patch_map: dict[str, dict[str, str]] = {}
    crm_patch_rank: dict[str, date] = {}

    for row_num in range(2, ws.max_row + 1):
        oid = ""
        if order_col_primary:
            oid = normalize_order_id(ws.cell(row=row_num, column=order_col_primary).value)
        if not oid and order_col_alt:
            oid = normalize_order_id(ws.cell(row=row_num, column=order_col_alt).value)
        if not oid:
            continue

        sale_date = parse_date_any(ws.cell(row=row_num, column=date_col).value)
        if sale_date is not None:
            crm_send_dates[oid].add(sale_date)

        if oid in target_order_ids and in_range(sale_date, date_from, date_to):
            patch = db_patches.get(oid)
            if patch:
                if is_blank(ws.cell(row=row_num, column=status_col).value) and patch.status:
                    ws.cell(row=row_num, column=status_col).value = patch.status
                    updates["crm_status"] += 1
                if is_blank(ws.cell(row=row_num, column=issued_col).value) and patch.issued_flag:
                    ws.cell(row=row_num, column=issued_col).value = patch.issued_flag
                    updates["crm_issued"] += 1
                if is_blank(ws.cell(row=row_num, column=status_change_col).value) and patch.status_change_date:
                    ws.cell(row=row_num, column=status_change_col).value = patch.status_change_date
                    updates["crm_status_change_date"] += 1
                if is_blank(ws.cell(row=row_num, column=planned_col).value) and patch.planned_handover_date:
                    ws.cell(row=row_num, column=planned_col).value = patch.planned_handover_date
                    updates["crm_planned_handover"] += 1

            if is_blank_or_zero(ws.cell(row=row_num, column=seller_fee_col).value):
                seller_fee = ""
                if patch and patch.seller_fee:
                    seller_fee = patch.seller_fee
                if not seller_fee:
                    seller_fee = fmt_numeric(ws.cell(row=row_num, column=delivery_fee_col).value)
                if seller_fee:
                    ws.cell(row=row_num, column=seller_fee_col).value = seller_fee
                    updates["crm_seller_fee"] += 1

        rank = sale_date or date.min
        prev_rank = crm_patch_rank.get(oid, date.min)
        if rank < prev_rank:
            continue
        crm_patch_rank[oid] = rank
        crm_patch_map[oid] = {
            "status": str(ws.cell(row=row_num, column=status_col).value or "").strip(),
            "issued_flag": str(ws.cell(row=row_num, column=issued_col).value or "").strip(),
            "status_change_date": str(ws.cell(row=row_num, column=status_change_col).value or "").strip(),
            "planned_handover_date": str(ws.cell(row=row_num, column=planned_col).value or "").strip(),
            "seller_fee": fmt_numeric(ws.cell(row=row_num, column=seller_fee_col).value),
        }

    return dict(updates), crm_send_dates, crm_patch_map


def load_truth_send_date_map(path: Path) -> dict[str, str]:
    out: dict[str, set[date]] = defaultdict(set)
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        for row in reader:
            oid = (
                normalize_order_id(row.get("OrderID"))
                or normalize_order_id(row.get("№ заказа"))
                or normalize_order_id(row.get("norm_order_id"))
            )
            if not oid:
                continue
            d = (
                parse_date_any(row.get("Date"))
                or parse_date_any(row.get("Дата поступления заказа"))
                or parse_date_any(row.get("norm_sale_date"))
            )
            if d is None:
                continue
            out[oid].add(d)
    return {oid: latest_non_empty(dates) for oid, dates in out.items()}


def _coalesce(*values: str) -> str:
    for value in values:
        if value and value.strip():
            return value.strip()
    return ""


def update_archive_rows(
    rows: list[dict[str, str]],
    date_from: date,
    date_to: date,
    target_order_ids: set[str],
    crm_patch_map: Mapping[str, Mapping[str, str]],
    db_patches: Mapping[str, DbPatch],
    crm_send_dates: Mapping[str, set[date]],
    truth_send_dates: Mapping[str, str],
) -> tuple[list[dict[str, str]], dict[str, int]]:
    stats = defaultdict(int)
    updated_rows: list[dict[str, str]] = []
    for row in rows:
        out = dict(row)
        oid = normalize_order_id(out.get("№ заказа"))
        row_date = parse_date_any(out.get("Дата поступления заказа"))
        in_target = oid in target_order_ids and in_range(row_date, date_from, date_to)
        crm_patch = crm_patch_map.get(oid, {})
        db_patch = db_patches.get(oid)

        if in_target:
            if is_blank(out.get("Статус")):
                new_val = _coalesce(crm_patch.get("status", ""), db_patch.status if db_patch else "")
                if new_val:
                    out["Статус"] = new_val
                    stats["archive_status"] += 1

            if is_blank(out.get("Выдал")):
                new_val = _coalesce(crm_patch.get("issued_flag", ""), db_patch.issued_flag if db_patch else "")
                if new_val:
                    out["Выдал"] = new_val
                    stats["archive_issued"] += 1

            if is_blank(out.get("Дата изменения статуса")):
                new_val = _coalesce(
                    crm_patch.get("status_change_date", ""),
                    db_patch.status_change_date if db_patch else "",
                )
                if new_val:
                    out["Дата изменения статуса"] = new_val
                    stats["archive_status_change_date"] += 1

            if is_blank(out.get("Плановая дата передачи курьеру")):
                new_val = _coalesce(
                    crm_patch.get("planned_handover_date", ""),
                    db_patch.planned_handover_date if db_patch else "",
                )
                if new_val:
                    out["Плановая дата передачи курьеру"] = new_val
                    stats["archive_planned_handover"] += 1

            if is_blank_or_zero(out.get("Стоимость доставки для продавца")):
                new_val = _coalesce(
                    fmt_numeric(crm_patch.get("seller_fee", "")),
                    db_patch.seller_fee if db_patch else "",
                )
                if new_val:
                    out["Стоимость доставки для продавца"] = new_val
                    stats["archive_seller_fee"] += 1

        send_date = str(out.get("send_date", "")).strip()
        if not send_date:
            from_crm = latest_non_empty(crm_send_dates.get(oid, set()))
            if from_crm:
                out["send_date"] = from_crm
                stats["send_date_from_crm"] += 1
            else:
                from_truth = truth_send_dates.get(oid, "")
                if from_truth:
                    out["send_date"] = from_truth
                    stats["send_date_from_truth"] += 1

        updated_rows.append(out)
    stats["send_date_filled_total"] = stats["send_date_from_crm"] + stats["send_date_from_truth"]
    return updated_rows, dict(stats)


def write_archive_csv(path: Path, header: list[str], rows: list[dict[str, str]]) -> None:
    fieldnames = list(header)
    if "send_date" not in fieldnames:
        fieldnames.append("send_date")

    fd, tmp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=str(path.parent))
    tmp_path = Path(tmp_name)
    try:
        with open(fd, "w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=fieldnames)
            writer.writeheader()
            for row in rows:
                writer.writerow({k: row.get(k, "") for k in fieldnames})
        tmp_path.replace(path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink(missing_ok=True)


def write_archive_xlsx(path: Path, header: list[str], rows: list[dict[str, str]]) -> None:
    columns = list(header)
    if "send_date" not in columns:
        columns.append("send_date")

    def _save(tmp_path: Path) -> None:
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("ArchiveOrders_Mapped")
        ws.append(columns)
        for row in rows:
            ws.append([row.get(col, "") for col in columns])
        wb.save(tmp_path)

    _atomic_save_workbook(path, _save)


def validate_archive_outputs(csv_path: Path, xlsx_path: Path, expected_rows: int) -> None:
    with csv_path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        csv_header = list(reader.fieldnames or [])
        csv_rows = [dict(row) for row in reader]
    if len(csv_rows) != expected_rows:
        raise RuntimeError(f"CSV row count mismatch: expected={expected_rows} actual={len(csv_rows)}")

    wb = load_workbook(xlsx_path, read_only=True)
    if "ArchiveOrders_Mapped" not in wb.sheetnames:
        raise RuntimeError("XLSX output missing sheet ArchiveOrders_Mapped")
    ws = wb["ArchiveOrders_Mapped"]
    rows_iter = ws.iter_rows(values_only=True)
    header_tuple = next(rows_iter, None)
    if header_tuple is None:
        raise RuntimeError("XLSX output is empty")
    xlsx_header = [str(v) if v is not None else "" for v in header_tuple]
    xlsx_rows = sum(1 for _ in rows_iter)
    if xlsx_rows != expected_rows:
        raise RuntimeError(f"XLSX row count mismatch: expected={expected_rows} actual={xlsx_rows}")
    if csv_header != xlsx_header:
        raise RuntimeError("CSV/XLSX header mismatch after write")


def run(
    *,
    crm_workbook: Path,
    crm_sheet: str,
    db_path: Path,
    archive_final_csv: Path,
    archive_final_xlsx: Path,
    archive_sales_truth_csv: Path,
    date_from: date,
    date_to: date,
    backup_dir: Path,
    dry_run: bool = False,
) -> dict[str, Any]:
    ts = _now_ts()
    header, archive_rows = load_archive_csv(archive_final_csv)
    target_order_ids = build_target_order_ids(archive_rows, date_from, date_to)
    db_patches = load_db_patches(db_path, target_order_ids)

    wb = load_workbook(crm_workbook)
    if crm_sheet not in wb.sheetnames:
        raise RuntimeError(f"Sheet not found in CRM workbook: {crm_sheet}")
    ws = wb[crm_sheet]

    backups = {}
    if not dry_run:
        backups["crm_workbook"] = str(_backup_file(crm_workbook, backup_dir, ts))
        backups["archive_final_csv"] = str(_backup_file(archive_final_csv, backup_dir, ts))
        backups["archive_final_xlsx"] = str(_backup_file(archive_final_xlsx, backup_dir, ts))

    crm_updates, crm_send_dates, crm_patch_map = apply_crm_updates(
        ws=ws,
        db_patches=db_patches,
        target_order_ids=target_order_ids,
        date_from=date_from,
        date_to=date_to,
    )
    truth_send_dates = load_truth_send_date_map(archive_sales_truth_csv)
    updated_rows, archive_updates = update_archive_rows(
        rows=archive_rows,
        date_from=date_from,
        date_to=date_to,
        target_order_ids=target_order_ids,
        crm_patch_map=crm_patch_map,
        db_patches=db_patches,
        crm_send_dates=crm_send_dates,
        truth_send_dates=truth_send_dates,
    )

    if not dry_run:
        _atomic_save_workbook(crm_workbook, wb.save)

        write_archive_csv(archive_final_csv, header, updated_rows)
        write_archive_xlsx(archive_final_xlsx, header, updated_rows)
        validate_archive_outputs(archive_final_csv, archive_final_xlsx, expected_rows=len(archive_rows))

    unresolved_orders: set[str] = set()
    for row in updated_rows:
        d = parse_date_any(row.get("Дата поступления заказа"))
        if not in_range(d, date_from, date_to):
            continue
        if _archive_row_missing_period_values(row):
            oid = normalize_order_id(row.get("№ заказа"))
            if oid:
                unresolved_orders.add(oid)

    return {
        "timestamp": ts,
        "date_from": date_from.strftime(DATE_FMT_ISO),
        "date_to": date_to.strftime(DATE_FMT_ISO),
        "rows_target_period": len(target_order_ids),
        "order_ids_target": len(target_order_ids),
        "db_patches_found": len(db_patches),
        "crm_rows_updated_by_field": crm_updates,
        "archive_rows_updated_by_field": {
            k: v
            for k, v in archive_updates.items()
            if k
            in {
                "archive_status",
                "archive_issued",
                "archive_status_change_date",
                "archive_planned_handover",
                "archive_seller_fee",
            }
        },
        "send_date_filled": archive_updates.get("send_date_filled_total", 0),
        "send_date_from_crm": archive_updates.get("send_date_from_crm", 0),
        "send_date_from_truth": archive_updates.get("send_date_from_truth", 0),
        "unresolved_order_ids_count": len(unresolved_orders),
        "backup_paths": backups,
        "dry_run": dry_run,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Backfill CRM + Archive mapped period values and append send_date.")
    parser.add_argument("--crm-workbook", required=True)
    parser.add_argument("--crm-sheet", default="SALES_KSP_CRM_1")
    parser.add_argument("--db-path", required=True)
    parser.add_argument("--archive-final-csv", required=True)
    parser.add_argument("--archive-final-xlsx", required=True)
    parser.add_argument("--archive-sales-truth-csv", required=True)
    parser.add_argument("--date-from", required=True)
    parser.add_argument("--date-to", required=True)
    parser.add_argument("--backup-dir", required=True)
    parser.add_argument("--summary-out", default="")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    summary = run(
        crm_workbook=Path(args.crm_workbook),
        crm_sheet=args.crm_sheet,
        db_path=Path(args.db_path),
        archive_final_csv=Path(args.archive_final_csv),
        archive_final_xlsx=Path(args.archive_final_xlsx),
        archive_sales_truth_csv=Path(args.archive_sales_truth_csv),
        date_from=parse_date_arg(args.date_from),
        date_to=parse_date_arg(args.date_to),
        backup_dir=Path(args.backup_dir),
        dry_run=bool(args.dry_run),
    )

    if args.summary_out:
        summary_path = Path(args.summary_out)
        summary_path.parent.mkdir(parents=True, exist_ok=True)
        summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
