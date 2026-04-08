#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import os
import shutil
import sqlite3
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

import pandas as pd
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.paths import data_path
from scripts.backfill_crm_archive_period import (
    fmt_numeric,
    is_blank_or_zero,
    load_archive_csv,
    parse_date_any,
    validate_archive_outputs,
    write_archive_csv,
    write_archive_xlsx,
)
from scripts.import_orders_to_crm import (
    KNOWN_BASELINE_INTEGRITY_ERRORS,
    KNOWN_BASELINE_INTEGRITY_ERROR_PREFIXES,
    _refresh_formula_caches_xlwings,
    _resolve_table,
    _restore_preserved_package_parts,
    _snapshot_preserved_package_parts,
    _table_bounds,
)
from scripts.report_import_status import normalize_store_name
from scripts.validate_crm_workbook_integrity import (
    filter_integrity_errors,
    validate_workbook_integrity,
)


DEFAULT_CRM = data_path("excel_ui", "SALES_KSP_CRM_V3.xlsx")
DEFAULT_DB = data_path("db", "app.db")
DEFAULT_ARCHIVE_ROOT = data_path("exports")
DEFAULT_BACKUP_DIR = data_path("excel_ui", "backups")


def _backup_file(path: Path, backup_dir: Path) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = backup_dir / f"{path.name}.{ts}.bak"
    shutil.copy2(path, out)
    return out


def _float_or_zero(value: Any) -> float:
    if value is None:
        return 0.0
    text = str(value).strip().replace(",", ".")
    if not text:
        return 0.0
    try:
        return float(text)
    except (TypeError, ValueError):
        return 0.0


def _load_db_seller_fee_truth(db_path: Path, order_ids: set[str]) -> dict[str, float]:
    if not db_path.exists() or not order_ids:
        return {}
    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    try:
        table = conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name='fact_orders_kaspi'"
        ).fetchone()
        if table is None:
            return {}
        cols = {row[1] for row in conn.execute("PRAGMA table_info(fact_orders_kaspi)").fetchall()}
        if "order_id" not in cols or "delivery_cost_for_seller" not in cols:
            return {}

        out: dict[str, float] = {}
        ordered_ids = sorted(order_ids)
        for i in range(0, len(ordered_ids), 900):
            chunk = ordered_ids[i : i + 900]
            placeholders = ",".join("?" for _ in chunk)
            rows = conn.execute(
                f"""
                SELECT order_id, delivery_cost_for_seller
                FROM fact_orders_kaspi
                WHERE order_id IN ({placeholders})
                """,
                chunk,
            ).fetchall()
            for row in rows:
                order_id = str(row["order_id"]).strip()
                if order_id.endswith(".0"):
                    order_id = order_id[:-2]
                fee = _float_or_zero(row["delivery_cost_for_seller"])
                if order_id and fee != 0.0:
                    out[order_id] = fee
        return out
    finally:
        conn.close()


def _archive_paths(archive_root: Path) -> tuple[Path, Path]:
    if archive_root.is_file():
        csv_path = archive_root
        return csv_path, csv_path.with_suffix(".xlsx")

    candidates: list[tuple[float, int, Path]] = []
    for csv_path in archive_root.rglob("ArchiveOrders_ALL_STORES.csv"):
        try:
            with csv_path.open("r", encoding="utf-8-sig", newline="") as fh:
                reader = csv.DictReader(fh)
                fieldnames = list(reader.fieldnames or [])
                if "Стоимость доставки для продавца" not in fieldnames:
                    continue
                row_count = sum(1 for _ in reader)
        except Exception:
            continue
        candidates.append((csv_path.stat().st_mtime, row_count, csv_path))

    non_empty = [item for item in candidates if item[1] > 0]
    chosen = max(non_empty or candidates, default=None)
    if chosen is None:
        raise FileNotFoundError(
            f"No ArchiveOrders_ALL_STORES.csv with seller fee column found under {archive_root}"
        )
    csv_path = chosen[2]
    return csv_path, csv_path.with_suffix(".xlsx")


def _resolve_date_range(
    crm_workbook: Path,
    crm_sheet: str,
    raw_date_from: Optional[str],
    raw_date_to: Optional[str],
) -> tuple[date, date]:
    if raw_date_from and raw_date_to:
        return datetime.strptime(raw_date_from, "%Y-%m-%d").date(), datetime.strptime(raw_date_to, "%Y-%m-%d").date()

    df = pd.read_excel(crm_workbook, sheet_name=crm_sheet, usecols=["Date"])
    dates = [parse_date_any(value) for value in df["Date"].tolist()]
    valid_dates = [value for value in dates if value is not None]
    if not valid_dates:
        raise RuntimeError("Could not resolve CRM history date range from Date column.")
    full_from = min(valid_dates)
    full_to = max(valid_dates)
    if raw_date_from:
        full_from = datetime.strptime(raw_date_from, "%Y-%m-%d").date()
    if raw_date_to:
        full_to = datetime.strptime(raw_date_to, "%Y-%m-%d").date()
    return full_from, full_to


def _scan_crm_candidates(
    crm_workbook: Path,
    crm_sheet: str,
    db_path: Path,
    *,
    date_from: date,
    date_to: date,
) -> dict[str, Any]:
    df = pd.read_excel(crm_workbook, sheet_name=crm_sheet)
    order_col = "OrderID" if "OrderID" in df.columns else ("№ заказа" if "№ заказа" in df.columns else None)
    if order_col is None:
        raise RuntimeError("CRM workbook missing OrderID/№ заказа column.")
    store_col = "STORE_NAME" if "STORE_NAME" in df.columns else ("Склад передачи КД" if "Склад передачи КД" in df.columns else None)
    seller_col = "Стоимость доставки для продавца"
    if seller_col not in df.columns:
        raise RuntimeError("CRM workbook missing 'Стоимость доставки для продавца' column.")
    raw_fee_col = "Delivery_fee_kzt" if "Delivery_fee_kzt" in df.columns else ("Delivery_fee" if "Delivery_fee" in df.columns else None)

    row_dates = [parse_date_any(value) for value in df["Date"].tolist()]
    target_indices = [idx for idx, row_date in enumerate(row_dates) if row_date is not None and date_from <= row_date <= date_to]
    order_ids: set[str] = set()
    normalized_orders: dict[int, str] = {}
    for idx in target_indices:
        order_id = str(df.iloc[idx][order_col]).strip()
        if order_id.endswith(".0"):
            order_id = order_id[:-2]
        normalized_orders[idx] = order_id
        if order_id:
            order_ids.add(order_id)

    db_truth = _load_db_seller_fee_truth(db_path, order_ids)
    store_updates: dict[str, int] = defaultdict(int)
    seller_fee_updates = 0
    for idx in target_indices:
        row = df.iloc[idx]
        seller_fee = _float_or_zero(row.get(seller_col))
        raw_fee = _float_or_zero(row.get(raw_fee_col)) if raw_fee_col else 0.0
        db_fee = _float_or_zero(db_truth.get(normalized_orders.get(idx, "")))
        if seller_fee == 0.0 and (db_fee != 0.0 or raw_fee != 0.0):
            seller_fee_updates += 1
            store_name = normalize_store_name(row.get(store_col) if store_col else "")
            store_updates[store_name] += 1

    return {
        "rows_considered": len(target_indices),
        "seller_fee_updates": seller_fee_updates,
        "stores": dict(sorted(store_updates.items())),
        "workbook": str(crm_workbook),
    }


def _coerce_iso_date(value: str | date) -> date:
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value), "%Y-%m-%d").date()


def _refresh_formula_caches_if_trusted(workbook_path: Path) -> None:
    resolved = Path(workbook_path).expanduser().resolve()
    try:
        if not resolved.is_relative_to(PROJECT_ROOT):
            return
    except AttributeError:
        if PROJECT_ROOT not in resolved.parents and resolved != PROJECT_ROOT:
            return
    _refresh_formula_caches_xlwings(resolved, verbose=False)


def _apply_crm_seller_fee_backfill_openpyxl(
    *,
    crm_workbook: Path,
    crm_sheet: str,
    db_path: Path,
    date_from: str | date,
    date_to: str | date,
    table_name: str = "tb_SalesRaw",
) -> int:
    resolved_from = _coerce_iso_date(date_from)
    resolved_to = _coerce_iso_date(date_to)
    preserved_parts = _snapshot_preserved_package_parts(crm_workbook)
    tmp_path = crm_workbook.with_name(f"{crm_workbook.stem}.seller_fee_tmp{crm_workbook.suffix}")
    data_wb = load_workbook(filename=str(crm_workbook), read_only=False, data_only=True)
    wb = load_workbook(filename=str(crm_workbook), read_only=False, data_only=False)
    try:
        data_ws = data_wb[crm_sheet]
        ws = wb[crm_sheet]
        table = _resolve_table(ws, table_name)
        start_col, start_row, end_col, end_row = _table_bounds(table)

        header_to_col: dict[str, int] = {}
        for col_num in range(start_col, end_col + 1):
            header = str(ws.cell(row=start_row, column=col_num).value or "").strip()
            if header:
                header_to_col[header] = col_num

        date_col = header_to_col.get("Date") or header_to_col.get("Дата поступления заказа")
        fee_col = header_to_col.get("Delivery_fee_kzt") or header_to_col.get("Delivery_fee")
        seller_col = header_to_col.get("Стоимость доставки для продавца")
        order_col = header_to_col.get("OrderID") or header_to_col.get("№ заказа")
        if not date_col or not fee_col or not seller_col:
            raise RuntimeError("CRM workbook missing required seller-fee backfill columns.")

        pending_rows: list[tuple[int, str, float]] = []
        order_ids: set[str] = set()
        for row_num in range(start_row + 1, end_row + 1):
            parsed_date = parse_date_any(data_ws.cell(row=row_num, column=date_col).value)
            if parsed_date is None or parsed_date < resolved_from or parsed_date > resolved_to:
                continue
            seller_fee = _float_or_zero(data_ws.cell(row=row_num, column=seller_col).value)
            if seller_fee != 0.0:
                continue
            order_id = str(data_ws.cell(row=row_num, column=order_col).value).strip() if order_col else ""
            if order_id.endswith(".0"):
                order_id = order_id[:-2]
            raw_fee = _float_or_zero(data_ws.cell(row=row_num, column=fee_col).value)
            pending_rows.append((row_num, order_id, raw_fee))
            if order_id:
                order_ids.add(order_id)

        seller_fee_by_order = _load_db_seller_fee_truth(db_path, order_ids)
        updates: list[tuple[int, float]] = []
        for row_num, order_id, raw_fee in pending_rows:
            db_fee = _float_or_zero(seller_fee_by_order.get(order_id))
            new_value = db_fee if db_fee != 0.0 else raw_fee
            if new_value == 0.0:
                continue
            ws.cell(row=row_num, column=seller_col).value = new_value
            updates.append((row_num, new_value))

        if not updates:
            return 0

        wb.save(str(tmp_path))
    finally:
        wb.close()
        data_wb.close()

    try:
        _restore_preserved_package_parts(tmp_path, preserved_parts)
        integrity = validate_workbook_integrity(tmp_path)
        blocking_errors, _allowed_errors = filter_integrity_errors(
            integrity.errors,
            allow_exact=KNOWN_BASELINE_INTEGRITY_ERRORS,
            allow_prefix=KNOWN_BASELINE_INTEGRITY_ERROR_PREFIXES,
        )
        if blocking_errors:
            raise RuntimeError(
                "CRM workbook integrity validation failed after seller-fee history backfill: "
                + "; ".join(blocking_errors)
            )
        os.replace(str(tmp_path), str(crm_workbook))
        _refresh_formula_caches_if_trusted(crm_workbook)
    finally:
        if tmp_path.exists():
            tmp_path.unlink()

    return len(updates)


def _scan_archive_candidates(
    csv_path: Path,
    db_path: Path,
    *,
    date_from: date,
    date_to: date,
) -> tuple[dict[str, Any], list[str], list[dict[str, str]], dict[str, float]]:
    header, rows = load_archive_csv(csv_path)
    order_ids: set[str] = set()
    target_rows: list[int] = []
    for idx, row in enumerate(rows):
        row_date = parse_date_any(row.get("Дата поступления заказа"))
        if row_date is None or row_date < date_from or row_date > date_to:
            continue
        target_rows.append(idx)
        order_id = str(row.get("№ заказа", "")).strip()
        if order_id.endswith(".0"):
            order_id = order_id[:-2]
        if order_id:
            order_ids.add(order_id)
    db_truth = _load_db_seller_fee_truth(db_path, order_ids)

    store_updates: dict[str, int] = defaultdict(int)
    seller_fee_updates = 0
    for idx in target_rows:
        row = rows[idx]
        order_id = str(row.get("№ заказа", "")).strip()
        if order_id.endswith(".0"):
            order_id = order_id[:-2]
        seller_fee = _float_or_zero(row.get("Стоимость доставки для продавца"))
        raw_fee = _float_or_zero(row.get("Delivery_fee_kzt"))
        db_fee = _float_or_zero(db_truth.get(order_id))
        if seller_fee == 0.0 and (db_fee != 0.0 or raw_fee != 0.0):
            seller_fee_updates += 1
            store_name = normalize_store_name(
                row.get("STORE_NAME")
                or row.get("store_name")
                or row.get("store_code")
                or row.get("Склад передачи КД")
                or ""
            )
            store_updates[store_name] += 1

    return (
        {
            "rows_considered": len(target_rows),
            "seller_fee_updates": seller_fee_updates,
            "stores": dict(sorted(store_updates.items())),
            "csv_path": str(csv_path),
        },
        header,
        rows,
        db_truth,
    )


def _apply_archive_seller_fee_backfill(
    *,
    header: list[str],
    rows: list[dict[str, str]],
    db_truth: dict[str, float],
    csv_path: Path,
    xlsx_path: Path,
    date_from: date,
    date_to: date,
) -> int:
    updates = 0
    for row in rows:
        row_date = parse_date_any(row.get("Дата поступления заказа"))
        if row_date is None or row_date < date_from or row_date > date_to:
            continue
        if not is_blank_or_zero(row.get("Стоимость доставки для продавца")):
            continue
        order_id = str(row.get("№ заказа", "")).strip()
        if order_id.endswith(".0"):
            order_id = order_id[:-2]
        new_value = fmt_numeric(db_truth.get(order_id))
        if not new_value:
            new_value = fmt_numeric(row.get("Delivery_fee_kzt"))
        if not new_value:
            continue
        row["Стоимость доставки для продавца"] = new_value
        updates += 1

    write_archive_csv(csv_path, header, rows)
    write_archive_xlsx(xlsx_path, header, rows)
    validate_archive_outputs(csv_path, xlsx_path, expected_rows=len(rows))
    return updates


def run_history_backfill(
    *,
    crm_workbook: Path,
    crm_sheet: str,
    db_path: Path,
    archive_root: Path,
    backup_dir: Path,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    apply: bool = False,
) -> dict[str, Any]:
    resolved_from, resolved_to = _resolve_date_range(crm_workbook, crm_sheet, date_from, date_to)
    archive_csv, archive_xlsx = _archive_paths(archive_root)

    crm_summary = _scan_crm_candidates(
        crm_workbook,
        crm_sheet,
        db_path,
        date_from=resolved_from,
        date_to=resolved_to,
    )
    archive_summary, header, rows, db_truth = _scan_archive_candidates(
        archive_csv,
        db_path,
        date_from=resolved_from,
        date_to=resolved_to,
    )

    result: dict[str, Any] = {
        "apply": bool(apply),
        "date_from": resolved_from.isoformat(),
        "date_to": resolved_to.isoformat(),
        "backup_dir": str(backup_dir),
        "targets": {
            "crm": crm_summary,
            "archive": archive_summary,
        },
        "backups": {},
    }

    if apply:
        result["backups"]["crm"] = str(_backup_file(crm_workbook, backup_dir))
        result["backups"]["archive_csv"] = str(_backup_file(archive_csv, backup_dir))
        if archive_xlsx.exists():
            result["backups"]["archive_xlsx"] = str(_backup_file(archive_xlsx, backup_dir))

        crm_updates = _apply_crm_seller_fee_backfill_openpyxl(
            crm_workbook=crm_workbook,
            crm_sheet=crm_sheet,
            db_path=db_path,
            date_from=resolved_from,
            date_to=resolved_to,
        )
        archive_updates = _apply_archive_seller_fee_backfill(
            header=header,
            rows=rows,
            db_truth=db_truth,
            csv_path=archive_csv,
            xlsx_path=archive_xlsx,
            date_from=resolved_from,
            date_to=resolved_to,
        )
        result["targets"]["crm"]["seller_fee_updates"] = crm_updates
        result["targets"]["archive"]["seller_fee_updates"] = archive_updates

    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Backfill historical seller delivery fees in CRM and canonical archive")
    parser.add_argument("--crm-workbook", type=Path, default=DEFAULT_CRM)
    parser.add_argument("--crm-sheet", default="SALES_KSP_CRM_1")
    parser.add_argument("--db-path", type=Path, default=DEFAULT_DB)
    parser.add_argument("--archive-root", type=Path, default=DEFAULT_ARCHIVE_ROOT)
    parser.add_argument("--backup-dir", type=Path, default=DEFAULT_BACKUP_DIR)
    parser.add_argument("--date-from", default=None)
    parser.add_argument("--date-to", default=None)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--summary-out", type=Path, default=None)
    args = parser.parse_args(argv)

    summary = run_history_backfill(
        crm_workbook=args.crm_workbook.expanduser().resolve(),
        crm_sheet=args.crm_sheet,
        db_path=args.db_path.expanduser().resolve(),
        archive_root=args.archive_root.expanduser().resolve(),
        backup_dir=args.backup_dir.expanduser().resolve(),
        date_from=args.date_from,
        date_to=args.date_to,
        apply=bool(args.apply),
    )

    if args.summary_out:
        args.summary_out.parent.mkdir(parents=True, exist_ok=True)
        args.summary_out.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
