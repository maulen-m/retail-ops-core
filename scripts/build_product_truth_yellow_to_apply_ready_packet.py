#!/usr/bin/env python3
"""Build the read-only Product Truth yellow-to-apply-ready evidence packet."""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from datetime import datetime
import hashlib
import json
from pathlib import Path
import shutil
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_PRIOR_PACKET = (
    PROJECT_ROOT
    / "exports"
    / "validation"
    / "product_truth_blocker_closure_live_preflight_20260528_234625"
)
DEFAULT_CONFIG = PROJECT_ROOT / "config" / "product_truth_canonicalization_2026_05_28.json"
SIZES = ("S", "M", "L", "XL", "2XL", "3XL", "4XL")
LINE31_WHALE_BLUE_URL = (
    "https://kaspi.kz/shop/p/sportivnyi-kostjum-acmewear-of-line31-st-wb-l-sinii-l-162287320/"
)


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fieldnames is None:
        fieldnames = sorted({key for row in rows for key in row})
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _to_float(value: Any) -> float:
    text = str(value or "").strip()
    if not text:
        return 0.0
    return float(text.replace(",", ""))


def _size_index(size: str) -> int:
    return SIZES.index(size) if size in SIZES else 99


def _map_line31_article(row: dict[str, str]) -> tuple[str, str, str, str]:
    article = str(row.get("article") or "").strip()
    offer_name = str(row.get("kaspi_offer_name") or "")
    seller_name = str(row.get("seller_system_name") or "")
    haystack = " ".join([article, offer_name, seller_name]).upper()
    if "LINE31" not in haystack and "CL_OF_ARC_WM_LINE31" not in haystack:
        return "", "", "", ""

    if article.startswith("CL_OF_ARC_WM_LINE31") and "_ST_" in article:
        sku_key, raw_size = article.split("_ST_", 1)
        size = raw_size.split("_", 1)[0].upper()
        return sku_key, size, "canonical_article_st_size", ""

    aliases = {
        "OF_LINE31_ST_SB": "CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK",
        "OF_LINE31_ST_WB": "CL_OF_ARC_WM_LINE31_C-025_WHALE-BLUE",
        "OF_LINE31_ST_MB": "CL_OF_ARC_WM_LINE31_C-014_MISTY-BLUE",
        "OF_LINE31_ST_IV": "CL_OF_ARC_WM_LINE31_C-011_IVORY",
        "OF_LINE31_ST_IWSB": "CL_OF_ARC_WM_LINE31_B-C-011_IVORY_J-C-026_WHITE_L-C-023_STARRY-BLACK",
        "OF_LINE31_ST_CGOG": "CL_OF_ARC_WM_LINE31_B-C-005_CARDAMOM-GREEN_J-C-005_CARDAMOM-GREEN_L-C-015_OLIVE-GREEN",
    }
    upper_article = article.upper()
    for prefix, sku_key in aliases.items():
        token = f"{prefix}_"
        if upper_article.startswith(token):
            size = upper_article.removeprefix(token).split("_", 1)[0]
            return sku_key, size, f"alias:{prefix}", ""

    return "", "", "", "LINE31_TEXT_PRESENT_BUT_ARTICLE_NOT_MAPPABLE"


def _current_line31_status_rows(current_webui_csv: Path, prior_status_rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    prior_order_ids = {row.get("order_id") for row in prior_status_rows}
    out: list[dict[str, Any]] = []
    for row in _read_csv(current_webui_csv):
        sku_key, size, mapping_rule, quarantine_reason = _map_line31_article(row)
        if not sku_key:
            continue
        quantity = _to_float(row.get("quantity"))
        delivered = row.get("status_internal") == "DELIVERED"
        duplicate_prior = row.get("order_id") in prior_order_ids
        economic_qty = quantity if delivered and not duplicate_prior else 0.0
        out.append(
            {
                "order_id": row.get("order_id", ""),
                "store_code": row.get("store_code", ""),
                "source_type": "WEBUI_ARCHIVE_STATUS_CHANGE_CURRENT_SUPPLEMENT",
                "order_intake_date": row.get("created_at", ""),
                "status_change_at": row.get("status_change_at", ""),
                "status_raw": row.get("status_raw", ""),
                "status_internal": row.get("status_internal", ""),
                "article": row.get("article", ""),
                "kaspi_offer_name": row.get("kaspi_offer_name", ""),
                "seller_system_name": row.get("seller_system_name", ""),
                "mapped_sku_key": sku_key,
                "mapped_size": size,
                "quantity": quantity,
                "date_basis": "webui_status_change_at_current_refresh",
                "included_in_economic_depletion": "yes" if economic_qty else "no",
                "economic_depletion_qty": economic_qty,
                "included_in_physical_depletion": "no",
                "physical_depletion_qty": 0,
                "mapping_confidence": "HIGH" if mapping_rule else "LOW",
                "mapping_rule": mapping_rule,
                "quarantine_reason": "duplicate_prior_status_ledger" if duplicate_prior else quarantine_reason,
                "source_path": str(current_webui_csv),
                "source_file_sha256": row.get("source_file_sha256", ""),
            }
        )
    return out


def _build_line31_stock_split(
    prior_stock_rows: list[dict[str, str]],
    current_status_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    extra_economic: dict[tuple[str, str], float] = defaultdict(float)
    for row in current_status_rows:
        if row["included_in_economic_depletion"] == "yes":
            extra_economic[(row["mapped_sku_key"], row["mapped_size"])] += float(row["economic_depletion_qty"])

    out: list[dict[str, Any]] = []
    for row in prior_stock_rows:
        sku_key = row["sku_key"]
        color = row["color"]
        size = row["size"]
        anchor_qty = _to_float(row["anchor_qty"])
        po1a_sellable_addback_qty = _to_float(row["po1a_sellable_addback_qty"])
        po1a_no_route_physical_qty = _to_float(row["po1a_no_route_physical_qty"])
        physical_depletion_qty = _to_float(row["physical_depletion_qty"])
        economic_depletion_qty = _to_float(row["economic_depletion_qty"]) + extra_economic[(sku_key, size)]
        open_reserved_exposure_qty = _to_float(row["open_reserved_exposure_qty"])

        is_whale_blue = color == "Whale Blue"
        is_not_for_sale_reserve = po1a_no_route_physical_qty > 0 or is_whale_blue
        physical_total_qty = (
            anchor_qty + po1a_sellable_addback_qty + po1a_no_route_physical_qty - physical_depletion_qty
        )
        economic_final_sales_estimate_qty = (
            anchor_qty + po1a_sellable_addback_qty + po1a_no_route_physical_qty - economic_depletion_qty
        )
        physical_not_for_sale_reserve_qty = physical_total_qty if is_not_for_sale_reserve else 0.0
        physical_sellable_stock_qty = 0.0 if is_not_for_sale_reserve else physical_total_qty
        physical_sellable_available_after_open_reserved_qty = max(
            0.0,
            physical_sellable_stock_qty - open_reserved_exposure_qty,
        )

        if is_whale_blue:
            route_state = "NO_VALID_LIVE_ROUTE"
            retained = "WHALE_BLUE_NON_LIVE_INCORRECT_UNATTACHED_OFFER_IGNORED"
            ignored_url = LINE31_WHALE_BLUE_URL
            classification = "PHYSICAL_NOT_FOR_SALE_RESERVE_NO_VALID_LIVE_ROUTE"
        elif po1a_no_route_physical_qty > 0:
            route_state = "OWNER_CLARIFIED_NOT_FOR_SALE_YET"
            retained = "OWNER_CLARIFIED_RESERVE_NOT_BLOCKER"
            ignored_url = ""
            classification = "PHYSICAL_NOT_FOR_SALE_RESERVE"
        else:
            route_state = "SELLABLE_ROUTE_SOURCE_BACKED_OR_ANCHOR_ROW"
            retained = row.get("retained_blockers", "")
            ignored_url = ""
            classification = "PHYSICAL_SELLABLE_STOCK"

        out.append(
            {
                "sku_key": sku_key,
                "set_name": row.get("set_name", ""),
                "color": color,
                "size": size,
                "stock_classification": classification,
                "anchor_qty": anchor_qty,
                "po1a_sellable_addback_qty": 0.0 if is_whale_blue else po1a_sellable_addback_qty,
                "po1a_physical_not_for_sale_reserve_qty": (
                    po1a_sellable_addback_qty + po1a_no_route_physical_qty if is_whale_blue else po1a_no_route_physical_qty
                ),
                "physical_depletion_qty": physical_depletion_qty,
                "economic_depletion_qty": economic_depletion_qty,
                "economic_depletion_current_webui_supplement_qty": extra_economic[(sku_key, size)],
                "open_reserved_on_delivery_exposure_qty": open_reserved_exposure_qty,
                "physical_total_stock_qty": physical_total_qty,
                "physical_sellable_stock_qty": physical_sellable_stock_qty,
                "physical_sellable_available_after_open_reserved_qty": physical_sellable_available_after_open_reserved_qty,
                "physical_not_for_sale_reserve_qty": physical_not_for_sale_reserve_qty,
                "economic_final_sales_estimate_qty": economic_final_sales_estimate_qty,
                "route_state": route_state,
                "ignored_sellable_route_url": ignored_url,
                "confidence": "HIGH" if not retained else "MEDIUM",
                "retained_blockers_or_disclosures": retained,
                "date_basis": (
                    "physical=ship_date/courier handoff from prior DB shipping ledger; "
                    "economic=WebUI status_change_at through 2026-05-29 current ACMEWEAR refresh"
                ),
            }
        )

    return sorted(out, key=lambda r: (r["stock_classification"], r["sku_key"], r["color"], _size_index(r["size"])))


def _build_po1a_reclassification(prior_rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for row in prior_rows:
        set_color = row["set_color"]
        if set_color in {
            "Bean Paste Pink 3-piece Set",
            "Pomelo Pink 3-piece Set",
            "Eggplant Purple 3-piece Set",
        }:
            classification = "PHYSICAL_NOT_FOR_SALE_RESERVE"
            reason = "owner clarified real physical stock, not for sale yet; no immediate offer creation blocker"
        elif set_color == "Whale Blue 3-piece Set":
            classification = "PHYSICAL_NOT_FOR_SALE_RESERVE_NO_VALID_LIVE_ROUTE"
            reason = "owner clarified Whale Blue URL is incorrect/unattached and ignored for sellable-route proof"
        else:
            classification = row["mapping_classification"]
            reason = row["route_reason"]
        item = dict(row)
        item["apply_ready_classification"] = classification
        item["apply_ready_reason"] = reason
        item["ignored_sellable_route_url"] = LINE31_WHALE_BLUE_URL if set_color == "Whale Blue 3-piece Set" else ""
        out.append(item)
    return out


def _build_rush31_readiness(prior_rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for row in prior_rows:
        item = dict(row)
        item["pp_stock_inventory_authority"] = "false"
        item["current_pp1_stock_evidence_class"] = "PRICELIST_ONLY_NOT_INVENTORY_TRUTH"
        item["canonical_planning_stock_total"] = item.get("owner_stock_total", "290")
        item["canonical_planning_stock_source"] = "OWNER_PHYSICAL_OVERRIDE_2026_05_28_21_38_20"
        item["preflight_readiness"] = "APPLY_READY_REVIEW_ONLY_WITH_EXACT_OWNER_APPROVAL"
        item["readiness_reason"] = (
            "offer/campaign/product found in prior read-only packet; owner physical stock is canonical; "
            "PP/pricelist stock mismatch is ignored as inventory truth; live price/bid/stock changes still need exact approval"
        )
        out.append(item)
    return out


def _write_workbook(path: Path, sheets: dict[str, list[dict[str, Any]]], summary: list[list[Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    for row in summary:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 70
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    for title, rows in sheets.items():
        sheet = wb.create_sheet(title[:31])
        if not rows:
            sheet.append(["status"])
            sheet.append(["no rows"])
            continue
        headers = list(rows[0].keys())
        sheet.append(headers)
        for item in rows:
            sheet.append([item.get(header, "") for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for col_idx, header in enumerate(headers, start=1):
            width = min(max(len(str(header)) + 2, 12), 48)
            sheet.column_dimensions[get_column_letter(col_idx)].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    shutil.move(str(tmp), str(path))
    probe = load_workbook(path, read_only=True, data_only=True)
    probe.close()


def build_packet(args: argparse.Namespace) -> dict[str, Any]:
    output_dir = args.output_dir or (
        PROJECT_ROOT
        / "exports"
        / "validation"
        / f"product_truth_yellow_to_apply_ready_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    )
    output_dir.mkdir(parents=True, exist_ok=True)

    prior_stock = _read_csv(args.prior_packet / "updated_line31_stock_table.csv")
    prior_po1a = _read_csv(args.prior_packet / "po1a_mapping_table.csv")
    prior_status = _read_csv(args.prior_packet / "line31_sale_status_ledger.csv")
    prior_rush = _read_csv(args.prior_packet / "rush31_preflight_table.csv")

    current_status = _current_line31_status_rows(args.current_webui_csv, prior_status)
    all_status = prior_status + current_status
    stock_split = _build_line31_stock_split(prior_stock, current_status)
    po1a_reclass = _build_po1a_reclassification(prior_po1a)
    rush_readiness = _build_rush31_readiness(prior_rush)

    ignored_evidence = [
        {
            "evidence_type": "PRICELIST_PP_STOCK",
            "product_family": "RUSH31",
            "observed_value": prior_rush[0].get("current_pp1_stock", "") if prior_rush else "",
            "authority": "NON_AUTHORITATIVE_FOR_INVENTORY_TRUTH",
            "allowed_use": "prior pricelist/offer-state evidence only",
            "ignored_for": "physical stock, reorder qty, owner stock truth",
            "source_paths": prior_rush[0].get("source_paths", "") if prior_rush else "",
        },
        {
            "evidence_type": "INCORRECT_LINE31_WHALE_BLUE_URL",
            "product_family": "LINE31",
            "observed_value": LINE31_WHALE_BLUE_URL,
            "authority": "IGNORED_FOR_SELLABLE_ROUTE_PROOF",
            "allowed_use": "negative evidence of incorrect/unattached route",
            "ignored_for": "sellable route, live stock route proof",
            "source_paths": "owner clarification 2026-05-29",
        },
    ]
    physical_sellable_rows = [row for row in stock_split if row["physical_sellable_stock_qty"] > 0]
    physical_not_for_sale_rows = [row for row in stock_split if row["physical_not_for_sale_reserve_qty"] > 0]
    economic_final_sales_rows = [row for row in stock_split if row["economic_final_sales_estimate_qty"] != 0]
    open_reserved_rows = [row for row in stock_split if row["open_reserved_on_delivery_exposure_qty"] > 0]
    low_confidence_rows = [
        row
        for row in stock_split
        if row["confidence"] != "HIGH"
        or row["retained_blockers_or_disclosures"]
        or row["route_state"] != "SELLABLE_ROUTE_SOURCE_BACKED_OR_ANCHOR_ROW"
    ]

    disclosures = [
        {
            "domain": "WebUI current refresh",
            "disclosure": "Current WebUI refresh targeted ACMEWEAR for LINE31 2026-05-27..2026-05-29.",
            "blocker": "no",
            "reason": "LINE31 evidence route in this wave is ACMEWEAR-scoped; omitted enabled stores are disclosed, not used as LINE31 blockers.",
        }
    ]
    retained_blockers: list[dict[str, Any]] = []

    csv_outputs = {
        "line31_apply_ready_stock_split.csv": stock_split,
        "line31_status_change_current_supplement.csv": current_status,
        "line31_sale_status_ledger_refreshed.csv": all_status,
        "physical_sellable_stock.csv": physical_sellable_rows,
        "physical_not_for_sale_reserve.csv": physical_not_for_sale_rows,
        "economic_final_sales_estimate.csv": economic_final_sales_rows,
        "open_reserved_on_delivery_exposure.csv": open_reserved_rows,
        "low_confidence_rows.csv": low_confidence_rows,
        "po1a_reserve_reclassification.csv": po1a_reclass,
        "rush31_apply_readiness.csv": rush_readiness,
        "ignored_pricelist_only_evidence.csv": ignored_evidence,
        "scope_disclosures.csv": disclosures,
        "retained_blockers.csv": retained_blockers,
    }
    for filename, rows in csv_outputs.items():
        if filename == "retained_blockers.csv":
            _write_csv(
                output_dir / filename,
                rows,
                fieldnames=["domain", "blocker", "impact", "next_safe_step", "severity"],
            )
        else:
            _write_csv(output_dir / filename, rows)

    totals = {
        "line31_physical_total_stock_qty": sum(row["physical_total_stock_qty"] for row in stock_split),
        "line31_physical_sellable_stock_qty": sum(row["physical_sellable_stock_qty"] for row in stock_split),
        "line31_physical_sellable_available_after_open_reserved_qty": sum(
            row["physical_sellable_available_after_open_reserved_qty"] for row in stock_split
        ),
        "line31_physical_not_for_sale_reserve_qty": sum(row["physical_not_for_sale_reserve_qty"] for row in stock_split),
        "line31_economic_final_sales_estimate_qty": sum(row["economic_final_sales_estimate_qty"] for row in stock_split),
        "line31_open_reserved_on_delivery_exposure_qty": sum(row["open_reserved_on_delivery_exposure_qty"] for row in stock_split),
        "line31_current_webui_economic_depletion_supplement_qty": sum(
            row["economic_depletion_current_webui_supplement_qty"] for row in stock_split
        ),
        "line31_current_webui_supplement_rows": len(current_status),
        "rush31_owner_canonical_stock_total": 290,
        "retained_blocker_count": len(retained_blockers),
    }

    summary_rows = [
        ["metric", "value", "note"],
        ["gate", "GREEN_PENDING_VALIDATION", "Final gate is set in closeout after validation commands pass."],
        ["line31 physical total stock", totals["line31_physical_total_stock_qty"], "Sellable plus not-for-sale reserve."],
        ["line31 physical sellable stock", totals["line31_physical_sellable_stock_qty"], "Valid sellable route only."],
        [
            "line31 available after open reserved",
            totals["line31_physical_sellable_available_after_open_reserved_qty"],
            "Sellable stock less open reserved/on-delivery exposure.",
        ],
        ["line31 not-for-sale reserve", totals["line31_physical_not_for_sale_reserve_qty"], "Whale Blue + pink/purple reserve."],
        ["line31 economic final-sales estimate", totals["line31_economic_final_sales_estimate_qty"], "Status-change depletion view."],
        ["rush31 canonical owner stock", 290, "PP/pricelist stock is non-authoritative."],
        ["retained blockers", len(retained_blockers), "No blockers retained by packet generation."],
    ]
    workbook_path = output_dir / "product_truth_apply_ready_workbook.xlsx"
    _write_workbook(
        workbook_path,
        {
            "LINE31 Stock Split": stock_split,
            "Physical Sellable": physical_sellable_rows,
            "Not For Sale Reserve": physical_not_for_sale_rows,
            "Economic Final Sales": economic_final_sales_rows,
            "Open Reserved": open_reserved_rows,
            "Low Confidence": low_confidence_rows,
            "LINE31 Status Supplement": current_status,
            "PO1A Reclass": po1a_reclass,
            "RUSH31 Readiness": rush_readiness,
            "Ignored Evidence": ignored_evidence,
            "Disclosures": disclosures,
            "Retained Blockers": retained_blockers,
        },
        summary_rows,
    )

    source_paths = [
        args.config,
        args.prior_packet / "updated_line31_stock_table.csv",
        args.prior_packet / "line31_sale_status_ledger.csv",
        args.prior_packet / "line31_ship_ledger.csv",
        args.prior_packet / "po1a_mapping_table.csv",
        args.prior_packet / "rush31_preflight_table.csv",
        args.current_webui_csv,
    ]
    if args.current_webui_manifest:
        source_paths.append(args.current_webui_manifest)

    manifest = {
        "schema_version": "product_truth_yellow_to_apply_ready.v1",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "gate": "GREEN_PENDING_VALIDATION",
        "run_dir": str(output_dir),
        "authority": "read-only/copied-temp/repo-docs-tests-only; no production or external writes",
        "sources": [
            {"path": str(path), "sha256": _sha256(path), "exists": path.exists()} for path in source_paths if path.exists()
        ],
        "outputs": sorted(csv_outputs.keys())
        + [
            "product_truth_apply_ready_workbook.xlsx",
            "product_truth_yellow_to_apply_ready_manifest.json",
            "future_owner_approval_phrases.md",
            "closeout.md",
        ],
        "totals": totals,
        "owner_clarifications_applied": {
            "pricelist_pp_stock_inventory_truth": False,
            "line31_pink_purple_physical_not_for_sale_reserve": True,
            "line31_whale_blue_single_size_url_ignored": LINE31_WHALE_BLUE_URL,
            "rush31_owner_stock_canonical_over_pricelist": True,
        },
    }
    (output_dir / "product_truth_yellow_to_apply_ready_manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    (output_dir / "future_owner_approval_phrases.md").write_text(
        """# Future Owner Approval Phrases

Do not use these unless a fresh read-only live preflight immediately before action still matches this packet.

## RUSH31 Price

I approve setting RUSH31 offer `165486887` / SKU `21282790b` price to `9,990 KZT` as a live Kaspi price action after fresh preflight confirms the offer, stock, and campaign mapping. This approval does not authorize any other product, stock, bid, budget, or campaign change.

## RUSH31 Bid Test

I approve a live RUSH31 campaign `2794142` BID test at `>=100 KZT` only after fresh read-only preflight confirms current price, stock, campaign state, budget, and no conflicting blocker. This approval does not authorize budget changes, stock changes, or other campaigns.

## LINE31 Sellable Route Repair

I approve a future LINE31 sellable-route repair preflight for physical reserve stock, including Whale Blue and pink/purple reserve rows, with no offer creation or activation unless a separate exact live-action approval is issued after evidence review.

## Stock/Price/Campaign Boundary

I approve only the exact live action named above after fresh preflight. No production DB write, workbook write, scheduler change, source-pointer change, Kaspi/API/WebUI/Meta/CRM mutation, stock change, price change, bid/budget/state change, cash movement, supplier payment, PO commitment, owner publication, production preflight, or production apply is authorized beyond that exact phrase.
""",
        encoding="utf-8",
    )
    (output_dir / "closeout.md").write_text(
        f"""# Product Truth Yellow-To-Apply-Ready Wave Closeout

Gate: `GREEN_PENDING_VALIDATION`

Generated at: `{datetime.now().isoformat(timespec='seconds')}`

## Applied Clarifications

- Kaspi pricelist `PP*` stock is non-authoritative for inventory truth and retained only as pricelist/offer-state evidence.
- LINE31 Bean Paste Pink, Pomelo Pink, and Eggplant Purple PO-1A units are real physical stock, not for sale yet, and no longer blockers requiring immediate offer creation.
- LINE31 Whale Blue has no valid live route; `{LINE31_WHALE_BLUE_URL}` is ignored as incorrect/unattached single-size evidence.
- RUSH31 uses owner physical stock total `290` as canonical over pricelist stock.

## Key Totals

- LINE31 physical total stock: `{totals['line31_physical_total_stock_qty']}`
- LINE31 physical sellable stock: `{totals['line31_physical_sellable_stock_qty']}`
- LINE31 physical sellable available after open reserved: `{totals['line31_physical_sellable_available_after_open_reserved_qty']}`
- LINE31 physical not-for-sale reserve: `{totals['line31_physical_not_for_sale_reserve_qty']}`
- LINE31 economic final-sales estimate: `{totals['line31_economic_final_sales_estimate_qty']}`
- LINE31 current WebUI economic depletion supplement: `{totals['line31_current_webui_economic_depletion_supplement_qty']}`
- RUSH31 canonical owner stock: `290`

## Boundary

No production DB writes, production workbook writes, source-pointer writes,
scheduler changes, Web_automation writes, Kaspi/API/WebUI/Meta/CRM mutations,
campaign bid/budget/state changes, price changes, stock changes, supplier
messages, payment, PO commitment, cash movement, owner publication, production
preflight, production apply, or external writes were performed.
""",
        encoding="utf-8",
    )
    return manifest


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--prior-packet", type=Path, default=DEFAULT_PRIOR_PACKET)
    parser.add_argument("--current-webui-csv", type=Path, required=True)
    parser.add_argument("--current-webui-manifest", type=Path, default=None)
    parser.add_argument("--output-dir", type=Path, default=None)
    args = parser.parse_args(argv)
    manifest = build_packet(args)
    print(json.dumps({"ok": True, "run_dir": manifest["run_dir"], "totals": manifest["totals"]}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
