from __future__ import annotations

import json
from pathlib import Path
import zipfile

from openpyxl import Workbook
import yaml

from core.ops.kaspi_offer_manifest import (
    build_offer_packages,
    build_workbook_ingest_payload,
    load_offer_manifest,
    validate_offer_package_layout,
)


def _write_template(path: Path, *, category: str) -> None:
    wb = Workbook()
    intro = wb.active
    intro.title = "intro"
    intro["A1"] = category

    attr = wb.create_sheet("attributes")
    columns = [
        ("merchant_sku", "Артикул"),
        ("name", "Название товара"),
        ("brand", "Бренд"),
        ("image_code", "Код изображений"),
        ("description", "Описание (мин. 100 символов, макс. 7 000 символов)"),
        ("family_id", "Объединить в одну карточку"),
        ("manufacturer_size", "Размер производителя"),
        ("manufacturer_code", "Артикул производителя"),
        ("colour", "Цвет"),
        ("collection", "Коллекция"),
    ]
    if category == "women-sport-suits":
        columns.extend(
            [
                ("style", "Стиль"),
                ("sport", "Вид спорта"),
            ]
        )
    else:
        columns.extend(
            [
                ("model", "Модель"),
                ("gender", "Пол"),
            ]
        )

    for idx, (machine_key, display_name) in enumerate(columns, start=1):
        attr.cell(1, idx, "- текстовое значение\n- обязательное поле")
        attr.cell(2, idx, machine_key)
        attr.cell(3, idx, display_name)

    values = wb.create_sheet("values")
    values.cell(1, 1, "Цвет")
    values.cell(2, 1, "черный")
    values.cell(1, 2, "Коллекция")
    values.cell(2, 2, "Весна-Лето 2026")
    values.cell(1, 3, "Вид спорта")
    values.cell(2, 3, "универсальный")
    values.cell(1, 4, "Стиль")
    values.cell(2, 4, "спортивный")
    values.cell(1, 5, "Модель")
    values.cell(2, 5, "комплект")
    values.cell(1, 6, "Пол")
    values.cell(2, 6, "женский")

    wb.save(path)


def _write_source_images(source_dir: Path, *, count: int) -> None:
    png_bytes = (
        b"\x89PNG\r\n\x1a\n"
        b"\x00\x00\x00\rIHDR"
        b"\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00"
        b"\x1f\x15\xc4\x89"
        b"\x00\x00\x00\x0cIDATx\x9cc```\x00\x00\x00\x04\x00\x01"
        b"\x0b\xe7\x02\x9d"
        b"\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    for idx in range(1, count + 1):
        folder = source_dir / f"{idx:02d}"
        folder.mkdir(parents=True, exist_ok=True)
        (folder / f"source_{idx:02d}.png").write_bytes(png_bytes)


def _write_renamed_images(source_dir: Path, *, count: int) -> None:
    png_bytes = (
        b"\x89PNG\r\n\x1a\n"
        b"\x00\x00\x00\rIHDR"
        b"\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00"
        b"\x1f\x15\xc4\x89"
        b"\x00\x00\x00\x0cIDATx\x9cc```\x00\x00\x00\x04\x00\x01"
        b"\x0b\xe7\x02\x9d"
        b"\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    source_dir.mkdir(parents=True, exist_ok=True)
    for idx in range(1, count + 1):
        (source_dir / f"{idx}.png").write_bytes(png_bytes)


def _write_manifest(base_dir: Path) -> Path:
    product_root = base_dir / "Product_offers" / "LINE31" / "Starry_Black"
    source_dir = base_dir / "source_images"
    source_dir.mkdir(parents=True, exist_ok=True)
    _write_source_images(source_dir, count=6)

    description_path = base_dir / "desc.md"
    description_path.write_text("ACMEWEAR LINE31 description", encoding="utf-8")

    sport_template = base_dir / "Women-sport-suits-import-template.xlsm"
    thermal_template = base_dir / "Women-thermal-underwear-import-template.xlsm"
    _write_template(sport_template, category="women-sport-suits")
    _write_template(thermal_template, category="women-thermal-underwear")

    manifest = {
        "schema_version": 1,
        "product_root": str(product_root),
        "product": {
            "store": "ACMEWEAR",
            "brand": "ACMEWEAR",
            "base_model": "LINE31",
            "color_slug": "Starry_Black",
            "color_label": "Starry Black",
            "color_kaspi": "черный",
            "internal_base_sku_key": "CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK",
            "kaspi_name_core": "LINE31_Starry_Black",
            "display_title": "ACMEWEAR Женский спортивный костюм 3 в 1 — компрессионный комплект для тренировок",
            "family_id": "1n",
            "description_path": str(description_path),
            "vendor_id": "ARC",
            "product_type": "CL",
            "gender": "Women",
            "gender_rus": "женский",
            "season": "All",
            "base_cost_cny": 51,
            "base_cost_kzt": 3978,
            "weight_kg": 0.65,
            "sell_price_kzt": 14990,
            "stock_entered_kzt": 10000,
        },
        "images": {
            "source_dir": str(source_dir),
            "selection_mode": "folder_sequence_first_file",
            "folder_names": [f"{idx:02d}" for idx in range(1, 7)],
            "upload_count": 5,
        },
        "sizes": [
            {"size_label": "S", "numeric_size": 44},
            {"size_label": "M", "numeric_size": 46},
        ],
        "categories": [
            {
                "slug": "women-sport-suits",
                "short_code": "ST",
                "template_path": str(sport_template),
                "package_slug": "LINE31_STARRY_BLACK_ST_V1",
                "zip_slug": "ACMEWEAR_LINE31_STARRY_BLACK_ST_V1_UPLOAD",
                "group_token": "LINE31_ST_SB",
                "internal_article_pattern": "CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK_ST_{size_label}",
                "external_sku_pattern": "OF_LINE31_ST_SB_{size_label}",
                "offer_name_pattern": "{display_title} черный {size_label}",
                "shared_fields": {
                    "Название товара": "{display_title}",
                    "Бренд": "{brand}",
                    "Код изображений": "{shared_image_code}",
                    "Описание (мин. 100 символов, макс. 7 000 символов)": "{description_text}",
                    "Объединить в одну карточку": "{family_id}",
                    "Цвет": "{color_kaspi}",
                    "Коллекция": "Весна-Лето 2026",
                    "Стиль": "спортивный",
                    "Вид спорта": "универсальный",
                },
                "row_fields": {
                    "Артикул": "{internal_article}",
                    "Артикул производителя": "{external_sku}",
                    "Размер производителя": "{size_label}",
                },
            },
            {
                "slug": "women-thermal-underwear",
                "short_code": "TRM",
                "template_path": str(thermal_template),
                "package_slug": "LINE31_STARRY_BLACK_TRM_V1",
                "zip_slug": "ACMEWEAR_LINE31_STARRY_BLACK_TRM_V1_UPLOAD",
                "group_token": "LINE31_TRM_SB",
                "internal_article_pattern": "CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK_TRM_{size_label}",
                "external_sku_pattern": "OF_LINE31_TRM_SB_{size_label}",
                "offer_name_pattern": "{display_title} черный {size_label}",
                "shared_fields": {
                    "Название товара": "{display_title}",
                    "Бренд": "{brand}",
                    "Код изображений": "{shared_image_code}",
                    "Описание (мин. 100 символов, макс. 7 000 символов)": "{description_text}",
                    "Объединить в одну карточку": "{family_id}",
                    "Цвет": "{color_kaspi}",
                    "Коллекция": "Весна-Лето 2026",
                    "Модель": "комплект",
                    "Пол": "{gender_rus}",
                },
                "row_fields": {
                    "Артикул": "{internal_article}",
                    "Артикул производителя": "{external_sku}",
                    "Размер производителя": "{size_label}",
                },
            },
        ],
        "workbook": {
            "workbook_path": str(base_dir / "SALES_KSP_CRM_V3.xlsx"),
            "agent_session_title": "2026-03-19 — LINE31 manifest workflow",
        },
    }
    manifest_path = product_root / "offer_manifest.yaml"
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(yaml.safe_dump(manifest, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return manifest_path


def _write_batch_manifest(base_dir: Path) -> Path:
    product_root = base_dir / "Product_offers" / "LINE31" / "_final_color_bundles"
    description_path = base_dir / "desc.md"
    description_path.write_text("ACMEWEAR LINE31 multi-color description", encoding="utf-8")

    sport_template = base_dir / "Women-sport-suits-import-template.xlsm"
    thermal_template = base_dir / "Women-thermal-underwear-import-template.xlsm"
    _write_template(sport_template, category="women-sport-suits")
    _write_template(thermal_template, category="women-thermal-underwear")

    misty_dir = base_dir / "Misty_Blue"
    whale_dir = base_dir / "Whale_Blue"
    _write_renamed_images(misty_dir, count=3)
    _write_renamed_images(whale_dir, count=3)

    manifest = {
        "schema_version": 2,
        "product_root": str(product_root),
        "build_root": str(product_root / "final_uploads"),
        "product": {
            "store": "ACMEWEAR",
            "brand": "ACMEWEAR",
            "base_model": "LINE31",
            "display_title": "ACMEWEAR Женский спортивный костюм 3 в 1 — компрессионный комплект для тренировок",
            "family_id": "1n",
            "description_path": str(description_path),
            "vendor_id": "ARC",
            "product_type": "CL",
            "gender": "Women",
            "gender_rus": "женский",
            "season": "All",
            "base_cost_cny": 51,
            "base_cost_kzt": 3978,
            "weight_kg": 0.65,
            "sell_price_kzt": 14990,
            "stock_entered_kzt": 10000,
        },
        "variants": [
            {
                "color_slug": "Misty_Blue",
                "color_label": "Misty Blue",
                "color_title": "голубой",
                "color_kaspi": "голубой",
                "internal_base_sku_key": "CL_OF_ARC_WM_LINE31_C-014_MISTY-BLUE",
                "kaspi_name_core": "LINE31_Misty_Blue",
                "external_token": "MB",
                "description_path": str(description_path),
                "images": {
                    "source_dir": str(misty_dir),
                    "selection_mode": "explicit_filenames",
                    "filenames": ["1.png", "2.png", "3.png"],
                    "upload_count": 3,
                },
                "sizes": [
                    {"size_label": "M", "numeric_size": 46},
                    {"size_label": "L", "numeric_size": 48},
                ],
            },
            {
                "color_slug": "Whale_Blue",
                "color_label": "Whale Blue",
                "color_title": "синий",
                "color_kaspi": "синий",
                "internal_base_sku_key": "CL_OF_ARC_WM_LINE31_C-025_WHALE-BLUE",
                "kaspi_name_core": "LINE31_Whale_Blue",
                "external_token": "WB",
                "description_path": str(description_path),
                "images": {
                    "source_dir": str(whale_dir),
                    "selection_mode": "explicit_filenames",
                    "filenames": ["1.png", "2.png", "3.png"],
                    "upload_count": 3,
                },
                "sizes": [
                    {"size_label": "M", "numeric_size": 46},
                    {"size_label": "L", "numeric_size": 48},
                ],
            },
        ],
        "categories": [
            {
                "slug": "women-sport-suits",
                "short_code": "ST",
                "template_path": str(sport_template),
                "package_slug": "LINE31_FINAL_COLORS_ST_V1",
                "zip_slug": "ACMEWEAR_LINE31_FINAL_COLORS_ST_V1_UPLOAD",
                "group_token": "LINE31_ST_FINAL",
                "internal_article_pattern": "{internal_base_sku_key}_{short_code}_{size_label}",
                "external_sku_pattern": "OF_LINE31_ST_{external_token}_{size_label}",
                "offer_name_pattern": "{display_title} {color_title} {size_label}",
                "shared_fields": {
                    "Название товара": "{display_title}",
                    "Бренд": "{brand}",
                    "Код изображений": "{shared_image_code}",
                    "Описание (мин. 100 символов, макс. 7 000 символов)": "{description_text}",
                    "Объединить в одну карточку": "{family_id}",
                    "Цвет": "{color_kaspi}",
                    "Коллекция": "Весна-Лето 2026",
                    "Стиль": "спортивный",
                    "Вид спорта": "универсальный",
                },
                "row_fields": {
                    "Артикул": "{internal_article}",
                    "Артикул производителя": "{external_sku}",
                    "Размер производителя": "{size_label}",
                },
            },
            {
                "slug": "women-thermal-underwear",
                "short_code": "TRM",
                "template_path": str(thermal_template),
                "package_slug": "LINE31_FINAL_COLORS_TRM_V1",
                "zip_slug": "ACMEWEAR_LINE31_FINAL_COLORS_TRM_V1_UPLOAD",
                "group_token": "LINE31_TRM_FINAL",
                "internal_article_pattern": "{internal_base_sku_key}_{short_code}_{size_label}",
                "external_sku_pattern": "OF_LINE31_TRM_{external_token}_{size_label}",
                "offer_name_pattern": "{display_title} {color_title} {size_label}",
                "shared_fields": {
                    "Название товара": "{display_title}",
                    "Бренд": "{brand}",
                    "Код изображений": "{shared_image_code}",
                    "Описание (мин. 100 символов, макс. 7 000 символов)": "{description_text}",
                    "Объединить в одну карточку": "{family_id}",
                    "Цвет": "{color_kaspi}",
                    "Коллекция": "Весна-Лето 2026",
                    "Модель": "комплект",
                    "Пол": "{gender_rus}",
                },
                "row_fields": {
                    "Артикул": "{internal_article}",
                    "Артикул производителя": "{external_sku}",
                    "Размер производителя": "{size_label}",
                },
            },
        ],
        "workbook": {
            "workbook_path": str(base_dir / "SALES_KSP_CRM_V3.xlsx"),
            "agent_session_title": "2026-03-24 — LINE31 final colors manifest workflow",
            "candidate_color": "multi",
        },
    }
    manifest_path = product_root / "offer_batch_manifest.yaml"
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(yaml.safe_dump(manifest, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return manifest_path


def test_load_offer_manifest_resolves_core_paths(tmp_path: Path) -> None:
    manifest_path = _write_manifest(tmp_path)
    manifest = load_offer_manifest(manifest_path)

    assert manifest["product_root"] == manifest_path.parent
    assert manifest["product"]["description_path"].is_file()
    assert manifest["categories"][0]["template_path"].is_file()
    assert manifest["images"]["source_dir"].is_dir()


def test_build_offer_packages_from_manifest_creates_expected_artifacts(tmp_path: Path) -> None:
    manifest_path = _write_manifest(tmp_path)
    manifest = load_offer_manifest(manifest_path)

    report = build_offer_packages(manifest, timestamp="20260319_111500")

    source_root = manifest["product_root"] / "sources"
    assert (source_root / "original_6").is_dir()
    assert (source_root / "renamed_1_6").is_dir()
    assert (source_root / "images_part_2" / "6.png").is_file()
    assert (source_root / "IMAGE_MIGRATION_MAP.md").is_file()

    assert len(report["packages"]) == 2
    for package in report["packages"]:
        package_dir = Path(package["package_dir"])
        zip_path = Path(package["zip_path"])
        shared_image_code = package["rows"][0]["internal_article"]
        assert package_dir.is_dir()
        assert zip_path.is_file()
        assert (package_dir / "ROW_MAPPING.csv").is_file()
        assert (package_dir / "BUILD_LOG.md").is_file()
        assert (package_dir / "images" / shared_image_code / "1.png").is_file()
        assert (package_dir / "images" / shared_image_code / "5.png").is_file()
        with zipfile.ZipFile(zip_path) as zf:
            names = sorted(zf.namelist())
            assert any(name.endswith(".xlsm") for name in names)
            assert f"images/{shared_image_code}/1.png" in names
            assert f"images/{shared_image_code}/5.png" in names

        validation = validate_offer_package_layout(package_dir=package_dir, zip_path=zip_path)
        assert validation["ok"] is True
        assert validation["errors"] == []


def test_build_workbook_ingest_payload_matches_manifest_contract(tmp_path: Path) -> None:
    manifest_path = _write_manifest(tmp_path)
    manifest = load_offer_manifest(manifest_path)
    report = build_offer_packages(manifest, timestamp="20260319_111500")

    payload = build_workbook_ingest_payload(manifest, report)

    assert payload["m02_rows"][0][9] == "OF_LINE31_ST_SB_S"
    assert payload["m02_rows"][2][9] == "OF_LINE31_TRM_SB_S"
    assert payload["sku_map_rows"][0][0] == "CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK_S"


def test_build_offer_packages_batch_variants_create_multi_color_category_packages(tmp_path: Path) -> None:
    manifest_path = _write_batch_manifest(tmp_path)
    manifest = load_offer_manifest(manifest_path)

    report = build_offer_packages(manifest, timestamp="20260324_101500")

    assert len(report["packages"]) == 2
    assert Path(report["packages"][0]["package_dir"]).parts[-4:-1] == (
        "final_uploads",
        "20260324_101500",
        report["packages"][0]["category_slug"],
    )

    for package in report["packages"]:
        package_dir = Path(package["package_dir"])
        zip_path = Path(package["zip_path"])
        assert len(package["rows"]) == 4
        image_dirs = sorted(path.name for path in (package_dir / "images").iterdir() if path.is_dir())
        assert image_dirs == [
            "CL_OF_ARC_WM_LINE31_C-014_MISTY-BLUE_" + package["short_code"] + "_M",
            "CL_OF_ARC_WM_LINE31_C-025_WHALE-BLUE_" + package["short_code"] + "_M",
        ]
        with zipfile.ZipFile(zip_path) as zf:
            names = sorted(zf.namelist())
            assert f"images/{image_dirs[0]}/1.png" in names
            assert f"images/{image_dirs[0]}/3.png" in names
            assert f"images/{image_dirs[1]}/1.png" in names
            assert f"images/{image_dirs[1]}/3.png" in names

    payload = build_workbook_ingest_payload(manifest, report)
    assert [row[9] for row in payload["m02_rows"][:4]] == [
        "OF_LINE31_ST_MB_M",
        "OF_LINE31_ST_MB_L",
        "OF_LINE31_ST_WB_M",
        "OF_LINE31_ST_WB_L",
    ]
    assert payload["sku_map_rows"][0][0] == "CL_OF_ARC_WM_LINE31_C-014_MISTY-BLUE_M"
    assert payload["sku_map_rows"][2][0] == "CL_OF_ARC_WM_LINE31_C-025_WHALE-BLUE_M"
    assert payload["ingest_candidates"][0][0] == "LINE31_ST_FINAL"
    assert payload["size_engine_rows"][0][13] == "M"
    assert payload["offer_map_blocks"][0][0][0] == "NEW OFFER GROUP — LINE31_ST_FINAL — Misty_Blue"
    assert payload["offer_map_blocks"][2][0][0] == "NEW OFFER GROUP — LINE31_TRM_FINAL — Misty_Blue"

    report_path = manifest["product_root"] / "sources" / "reports" / "workbook_payload.json"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    assert report_path.is_file()


def test_validate_offer_package_layout_rejects_flat_images_directory(tmp_path: Path) -> None:
    package_dir = tmp_path / "flat_package"
    package_dir.mkdir(parents=True, exist_ok=True)
    template_path = package_dir / "Women-sport-suits-import-template_FIXED.xlsm"
    _write_template(template_path, category="women-sport-suits")

    images_dir = package_dir / "images"
    images_dir.mkdir(parents=True, exist_ok=True)
    _write_source_images(tmp_path / "src", count=1)
    src_image = next((tmp_path / "src" / "01").glob("*.png"))
    (images_dir / "1.png").write_bytes(src_image.read_bytes())

    zip_path = package_dir / "flat_package.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.write(template_path, arcname=template_path.name)
        zf.write(images_dir / "1.png", arcname="images/1.png")

    validation = validate_offer_package_layout(package_dir=package_dir, zip_path=zip_path)
    assert validation["ok"] is False
    assert any("images/<image_code>/" in error for error in validation["errors"])
