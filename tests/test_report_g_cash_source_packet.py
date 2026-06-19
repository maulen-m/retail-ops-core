from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

from scripts.report_g_cash_source_packet import build_cash_source_packet


def _write_cash_workbook(path: Path, *, timestamp: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cash_Balances"
    ws["A4"] = "Store"
    ws["B4"] = "Account"
    ws["C4"] = "Currency"
    ws["D4"] = timestamp
    ws["A5"] = "UNIVERSAL"
    ws["B5"] = "kaspi_gold"
    ws["C5"] = "KZT"
    ws["D5"] = 1703000
    ws["A6"] = "ACMEWEAR"
    ws["B6"] = "cash_usd"
    ws["C6"] = "USD"
    ws["D6"] = 748
    wb.save(path)


def _write_history(path: Path) -> None:
    path.write_text(
        """entries:
- as_of: 2026-06-13 01:01:26 GMT+5
  source: manual snapshot (Cash_Balances)
  balances: []
""",
        encoding="utf-8",
    )


def test_cash_source_packet_reads_today_snapshot_and_receipts(tmp_path: Path) -> None:
    workbook = tmp_path / "Inbound_calendar.xlsx"
    history = tmp_path / "bank_accounts_history.yaml"
    po51 = tmp_path / "PO_5.1_24.03.2026"
    arc = tmp_path / "ARC"
    nested_arc = arc / "PO-1B_20.05.2026_13_18_46"
    po51.mkdir()
    nested_arc.mkdir(parents=True)
    (po51 / "1_5000_PO_5.png").write_bytes(b"po51-a")
    (po51 / "2_7000_01.06.2026_21_08_22_PO_5.png").write_bytes(b"po51-b")
    (nested_arc / "1_7000.png").write_bytes(b"arc-a")
    (arc / "PO-1O-olivr-pt2_10.05.2026_16_09_33.png").write_bytes(b"arc-unparsed")
    _write_cash_workbook(workbook, timestamp="19.06.2026_11_44_56")
    _write_history(history)

    report = build_cash_source_packet(
        inbound_workbook=workbook,
        po51_receipts=po51,
        arc_receipts=arc,
        history_path=history,
        output_root=tmp_path / "out",
        generated_at="2026-06-19T12:30:00+05:00",
    )

    assert report["gate"] == "ARMED"
    assert report["cash_snapshot"]["as_of"] == "2026-06-19 11:44:56 GMT+5"
    assert report["cash_snapshot"]["latest_config_history_as_of"] == "2026-06-13 01:01:26 GMT+5"
    assert report["workbook_written"] is False
    assert report["external_writes_performed"] is False
    groups = {row["group"]: row for row in report["receipt_groups"]}
    assert groups["PO_5.1"]["image_file_count"] == 2
    assert groups["PO_5.1"]["parsed_amount_total_cny"] == "12000"
    assert groups["ARC"]["image_file_count"] == 2
    assert groups["ARC"]["parsed_amount_total_cny"] == "7000"
    assert groups["ARC"]["unparsed_amount_file_count"] == 1

    manifest_rows = list(csv.DictReader(Path(report["receipt_manifest_csv"]).open(encoding="utf-8")))
    assert len(manifest_rows) == 4
    assert {row["amount_parse_status"] for row in manifest_rows} == {
        "PARSED_LEADING_SEQUENCE_AMOUNT",
        "UNPARSED_NO_LEADING_SEQUENCE_AMOUNT",
    }


def test_cash_source_packet_blocks_stale_cash_snapshot(tmp_path: Path) -> None:
    workbook = tmp_path / "Inbound_calendar.xlsx"
    history = tmp_path / "bank_accounts_history.yaml"
    po51 = tmp_path / "PO_5.1_24.03.2026"
    arc = tmp_path / "ARC"
    po51.mkdir()
    arc.mkdir()
    (po51 / "1_5000_PO_5.png").write_bytes(b"po51-a")
    (arc / "1_7000.png").write_bytes(b"arc-a")
    _write_cash_workbook(workbook, timestamp="18.06.2026_11_44_56")
    _write_history(history)

    report = build_cash_source_packet(
        inbound_workbook=workbook,
        po51_receipts=po51,
        arc_receipts=arc,
        history_path=history,
        output_root=tmp_path / "out",
        generated_at="2026-06-19T12:30:00+05:00",
    )

    assert report["gate"] == "RED"
    assert "cash_snapshot_not_today:2026-06-18 11:44:56 GMT+5" in report["blockers"]
    assert report["cash_config_written"] is False
