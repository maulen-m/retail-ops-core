from __future__ import annotations

from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.table import TableFormula

from scripts.backfill_line61_kaspi_core import (
    CANONICAL_LINE61_CORE,
    backfill_line61_kaspi_core,
)

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _build_fixture_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SALES_KSP_CRM_1"
    headers = ["Date", "Kaspi_name_core", "SKU_key", "SKU_ID_KSP"]
    ws.append(headers)
    ws.append(["08.02.2026", "Спортивный_костюм_ACMEWEAR", "CL_NEW-CLO2_MEN_SUIT-61_BLACK", "OF_SUIT-61_BLK_XL"])
    ws.append(["08.02.2026", "Line51", "", "OF_SUIT-61_BLK_3XL"])
    ws.append(["08.02.2026", "Line51", "CL_OC_MEN_LINE52_BLACK", "OF_PRINT-51_BLK_L"])
    ws.cell(row=2, column=3, value=ArrayFormula(ref="C2", text='=D2&"_SKU"'))
    table = Table(displayName="tb_SalesRaw", ref="A1:D4")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    for table_col in table.tableColumns:
        if table_col.name == "SKU_key":
            table_col.calculatedColumnFormula = TableFormula(array=True, attr_text='D2&"_SKU"')
            break
    fill = PatternFill(fill_type="solid", fgColor="FFFF7E79")
    ws.conditional_formatting.add(
        "O2:Q2",
        FormulaRule(formula=['$M2="CL_OC_MEN_LINE52_BLACK"'], fill=fill),
    )
    wb.save(path)
    wb.close()

    with zipfile.ZipFile(path, "r") as zin:
        table_xml = ET.fromstring(zin.read("xl/tables/table1.xml"))
        columns = table_xml.find(f"{{{MAIN_NS}}}tableColumns")
        assert columns is not None
        for table_col in columns.findall(f"{{{MAIN_NS}}}tableColumn"):
            if table_col.attrib.get("name") != "SKU_key":
                continue
            calc = table_col.find(f"{{{MAIN_NS}}}calculatedColumnFormula")
            if calc is None:
                calc = ET.SubElement(table_col, f"{{{MAIN_NS}}}calculatedColumnFormula")
            calc.set("array", "1")
            calc.text = 'D2&"_SKU"'
            break
        modified = ET.tostring(table_xml, encoding="utf-8", xml_declaration=True)

        tmp = path.with_name(f"{path.stem}_table_formula_tmp.xlsx")
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/tables/table1.xml":
                    zout.writestr(item, modified)
                else:
                    zout.writestr(item, zin.read(item.filename))
    tmp.replace(path)


def test_backfill_line61_kaspi_core_dry_run_reports_without_writing(tmp_path: Path) -> None:
    workbook = tmp_path / "crm.xlsx"
    _build_fixture_workbook(workbook)

    stats = backfill_line61_kaspi_core(
        workbook=workbook,
        sheet_name="SALES_KSP_CRM_1",
        table_name="tb_SalesRaw",
        apply=False,
        backup_dir=tmp_path / "backups",
    )
    assert stats.line61_rows == 2
    assert stats.updated_rows == 2

    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb["SALES_KSP_CRM_1"]
    assert ws.cell(2, 2).value == "Спортивный_костюм_ACMEWEAR"
    assert ws.cell(3, 2).value == "Line51"
    wb.close()


def test_backfill_line61_kaspi_core_apply_updates_and_creates_backup(tmp_path: Path) -> None:
    workbook = tmp_path / "crm.xlsx"
    backup_dir = tmp_path / "backups"
    _build_fixture_workbook(workbook)

    stats = backfill_line61_kaspi_core(
        workbook=workbook,
        sheet_name="SALES_KSP_CRM_1",
        table_name="tb_SalesRaw",
        apply=True,
        backup_dir=backup_dir,
    )
    assert stats.line61_rows == 2
    assert stats.updated_rows == 2
    assert stats.backup_path is not None
    assert stats.backup_path.exists()
    assert stats.backup_path.parent == backup_dir

    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb["SALES_KSP_CRM_1"]
    assert ws.cell(2, 2).value == CANONICAL_LINE61_CORE
    assert ws.cell(3, 2).value == CANONICAL_LINE61_CORE
    assert ws.cell(4, 2).value == "Line51"
    wb.close()
