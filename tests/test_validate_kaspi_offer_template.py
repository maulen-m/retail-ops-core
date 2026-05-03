from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import Workbook

from scripts.validate_kaspi_offer_template import validate_kaspi_offer_template


def _setup_db(path: Path) -> None:
    conn = sqlite3.connect(path)
    conn.executescript(
        """
        CREATE TABLE dim_sku (
            sku_key TEXT PRIMARY KEY
        );

        CREATE TABLE dim_kaspi_article_map (
            store_code TEXT,
            kaspi_article TEXT,
            sku_key TEXT
        );

        CREATE TABLE fact_offer_stock_mapper_current (
            store_code TEXT,
            kaspi_article TEXT,
            offer_id TEXT,
            sku_key TEXT,
            mapping_method TEXT
        );
        """
    )
    conn.commit()
    conn.close()


def _insert_dim_sku(path: Path, sku_keys: list[str]) -> None:
    conn = sqlite3.connect(path)
    conn.executemany("INSERT INTO dim_sku (sku_key) VALUES (?)", [(x,) for x in sku_keys])
    conn.commit()
    conn.close()


def _insert_article_map(path: Path, store: str, article: str, sku_key: str) -> None:
    conn = sqlite3.connect(path)
    conn.execute(
        """
        INSERT INTO dim_kaspi_article_map (store_code, kaspi_article, sku_key)
        VALUES (?, ?, ?)
        """,
        (store, article, sku_key),
    )
    conn.commit()
    conn.close()


def _build_offer_xlsm(path: Path, *, merchant_sku: str, manufacturer_code: str, color: str, collection: str) -> None:
    wb = Workbook()
    ws_intro = wb.active
    ws_intro.title = "intro"

    ws_attr = wb.create_sheet("attributes")
    ws_attr.cell(1, 1, "- текстовое значение\n- обязательное поле")
    ws_attr.cell(1, 2, "- текстовое значение\n- обязательное поле")
    ws_attr.cell(1, 3, "- текстовое значение\n- обязательное поле")
    ws_attr.cell(1, 4, "- множество значений из списка\n- обязательное поле")
    ws_attr.cell(1, 5, "- множество значений из списка\n- обязательное поле")

    ws_attr.cell(2, 1, "merchant_sku")
    ws_attr.cell(2, 2, "name")
    ws_attr.cell(2, 3, "Clothes*General.clothes*manufacturer code")
    ws_attr.cell(2, 4, "Clothes*General.clothes*colour")
    ws_attr.cell(2, 5, "Clothes*General.clothes*collection")

    ws_attr.cell(3, 1, "Артикул")
    ws_attr.cell(3, 2, "Название товара")
    ws_attr.cell(3, 3, "Артикул производителя")
    ws_attr.cell(3, 4, "Цвет")
    ws_attr.cell(3, 5, "Коллекция")

    ws_attr.cell(4, 1, merchant_sku)
    ws_attr.cell(4, 2, "Спортивный костюм ACMEWEAR LINE51")
    ws_attr.cell(4, 3, manufacturer_code)
    ws_attr.cell(4, 4, color)
    ws_attr.cell(4, 5, collection)

    ws_values = wb.create_sheet("values")
    ws_values.cell(1, 1, "Цвет")
    ws_values.cell(1, 2, "Коллекция")
    ws_values.cell(2, 1, "черный")
    ws_values.cell(3, 1, "белый")
    ws_values.cell(2, 2, "Весна-Лето 2026")
    ws_values.cell(3, 2, "Осень-Зима 2026")

    wb.save(path)


def test_validate_offer_template_passes_for_line51_mapping(tmp_path: Path) -> None:
    db_path = tmp_path / "app.db"
    xlsm_path = tmp_path / "offer.xlsm"
    _setup_db(db_path)

    _insert_dim_sku(
        db_path,
        ["CL_OC_MEN_LINE51_WHITE", "CL_NEW-CLO2_MEN_SUIT-61_BLACK"],
    )
    _insert_article_map(db_path, "ACMEWEAR", "OF_LINE51_K-O_XL_48", "CL_OC_MEN_LINE51_WHITE")
    _build_offer_xlsm(
        xlsm_path,
        merchant_sku="OF_LINE51_K-O_XL_48",
        manufacturer_code="OF_LINE51_K-O_XL_48",
        color="черный, белый",
        collection="Весна-Лето 2026",
    )

    report = validate_kaspi_offer_template(
        xlsm_path=xlsm_path,
        db_path=db_path,
        category="men-sport-suits",
        store_code="ACMEWEAR",
        expect_sku_key="CL_OC_MEN_LINE51_WHITE",
        mode="fast",
    )
    assert report["ok"] is True
    assert report["error_count"] == 0
    assert report["mode"] == "fast"
    assert report["value_dict_scope"] == "skipped"


def test_validate_offer_template_passes_for_compact_bundle_token(tmp_path: Path) -> None:
    db_path = tmp_path / "app.db"
    xlsm_path = tmp_path / "offer.xlsm"
    _setup_db(db_path)

    _insert_dim_sku(db_path, ["SUIT-21-LS"])
    _insert_article_map(db_path, "ACMEWEAR", "SUIT-21-LS-ST-XL-48", "SUIT-21-LS")
    _build_offer_xlsm(
        xlsm_path,
        merchant_sku="SUIT-21-LS-ST-XL-48",
        manufacturer_code="SUIT-21-LS-ST-XL-48",
        color="черный",
        collection="Весна-Лето 2026",
    )

    report = validate_kaspi_offer_template(
        xlsm_path=xlsm_path,
        db_path=db_path,
        category="men-sport-suits",
        store_code="ACMEWEAR",
        mode="fast",
    )

    assert report["ok"] is True
    assert report["error_count"] == 0


def test_validate_offer_template_fails_on_model_mismatch(tmp_path: Path) -> None:
    db_path = tmp_path / "app.db"
    xlsm_path = tmp_path / "offer.xlsm"
    _setup_db(db_path)

    _insert_dim_sku(
        db_path,
        ["CL_OC_MEN_LINE51_WHITE", "CL_NEW-CLO2_MEN_SUIT-61_BLACK"],
    )
    _insert_article_map(db_path, "ACMEWEAR", "OF_LINE51_K-O_XL_48", "CL_NEW-CLO2_MEN_SUIT-61_BLACK")
    _build_offer_xlsm(
        xlsm_path,
        merchant_sku="OF_LINE51_K-O_XL_48",
        manufacturer_code="OF_LINE51_K-O_XL_48",
        color="черный, белый",
        collection="Весна-Лето 2026",
    )

    report = validate_kaspi_offer_template(
        xlsm_path=xlsm_path,
        db_path=db_path,
        category="men-sport-suits",
        store_code="ACMEWEAR",
        mode="fast",
    )
    assert report["ok"] is False
    assert any("model token mismatch" in err.lower() for err in report["errors"])


def test_validate_offer_template_fails_on_invalid_dictionary_value(tmp_path: Path) -> None:
    db_path = tmp_path / "app.db"
    xlsm_path = tmp_path / "offer.xlsm"
    _setup_db(db_path)

    _insert_dim_sku(db_path, ["CL_OC_MEN_LINE51_WHITE"])
    _insert_article_map(db_path, "ACMEWEAR", "OF_LINE51_K-O_XL_48", "CL_OC_MEN_LINE51_WHITE")
    _build_offer_xlsm(
        xlsm_path,
        merchant_sku="OF_LINE51_K-O_XL_48",
        manufacturer_code="OF_LINE51_K-O_XL_48",
        color="черный, зеленый",
        collection="Весна-Лето 2026",
    )

    report_fast = validate_kaspi_offer_template(
        xlsm_path=xlsm_path,
        db_path=db_path,
        category="men-sport-suits",
        store_code="ACMEWEAR",
        mode="fast",
    )
    assert report_fast["ok"] is True
    assert report_fast["value_dict_scope"] == "skipped"

    report_balanced = validate_kaspi_offer_template(
        xlsm_path=xlsm_path,
        db_path=db_path,
        category="men-sport-suits",
        store_code="ACMEWEAR",
        mode="balanced",
    )
    assert report_balanced["ok"] is False
    assert report_balanced["value_dict_scope"] == "populated_list_bound_columns"
    assert any("not present in values dictionary" in err.lower() for err in report_balanced["errors"])

    report_strict = validate_kaspi_offer_template(
        xlsm_path=xlsm_path,
        db_path=db_path,
        category="men-sport-suits",
        store_code="ACMEWEAR",
        mode="strict",
    )
    assert report_strict["ok"] is False
    assert report_strict["value_dict_scope"] == "all_list_bound_columns"
    assert any("not present in values dictionary" in err.lower() for err in report_strict["errors"])


def test_validate_offer_template_fails_on_missing_required_field(tmp_path: Path) -> None:
    db_path = tmp_path / "app.db"
    xlsm_path = tmp_path / "offer.xlsm"
    _setup_db(db_path)

    _insert_dim_sku(db_path, ["CL_OC_MEN_LINE51_WHITE"])
    _insert_article_map(db_path, "ACMEWEAR", "OF_LINE51_K-O_XL_48", "CL_OC_MEN_LINE51_WHITE")
    _build_offer_xlsm(
        xlsm_path,
        merchant_sku="OF_LINE51_K-O_XL_48",
        manufacturer_code="",
        color="черный, белый",
        collection="Весна-Лето 2026",
    )

    report = validate_kaspi_offer_template(
        xlsm_path=xlsm_path,
        db_path=db_path,
        category="men-sport-suits",
        store_code="ACMEWEAR",
        mode="fast",
    )
    assert report["ok"] is False
    assert any("missing required value" in err.lower() for err in report["errors"])


def test_validate_offer_template_fails_on_ambiguous_mapping(tmp_path: Path) -> None:
    db_path = tmp_path / "app.db"
    xlsm_path = tmp_path / "offer.xlsm"
    _setup_db(db_path)

    _insert_dim_sku(
        db_path,
        ["CL_OC_MEN_LINE51_WHITE", "CL_NEW-CLO2_MEN_SUIT-61_BLACK"],
    )
    _insert_article_map(db_path, "ACMEWEAR", "OF_LINE51_K-O_XL_48", "CL_OC_MEN_LINE51_WHITE")
    _insert_article_map(db_path, "ACMEWEAR", "OF_LINE51_K-O_XL_48", "CL_NEW-CLO2_MEN_SUIT-61_BLACK")
    _build_offer_xlsm(
        xlsm_path,
        merchant_sku="OF_LINE51_K-O_XL_48",
        manufacturer_code="OF_LINE51_K-O_XL_48",
        color="черный, белый",
        collection="Весна-Лето 2026",
    )

    report = validate_kaspi_offer_template(
        xlsm_path=xlsm_path,
        db_path=db_path,
        category="men-sport-suits",
        store_code="ACMEWEAR",
        mode="fast",
    )
    assert report["ok"] is False
    assert any("ambiguous sku mapping" in err.lower() for err in report["errors"])


def test_validate_offer_template_passes_for_women_sport_suits_category(tmp_path: Path) -> None:
    db_path = tmp_path / "app.db"
    xlsm_path = tmp_path / "offer_line31_sport.xlsm"
    _setup_db(db_path)

    _insert_dim_sku(db_path, ["CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK"])
    _insert_article_map(
        db_path,
        "ACMEWEAR",
        "OF_LINE31_ST_SB_S",
        "CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK",
    )
    _build_offer_xlsm(
        xlsm_path,
        merchant_sku="CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK_ST_S",
        manufacturer_code="OF_LINE31_ST_SB_S",
        color="черный",
        collection="Весна-Лето 2026",
    )

    report = validate_kaspi_offer_template(
        xlsm_path=xlsm_path,
        db_path=db_path,
        category="women-sport-suits",
        store_code="ACMEWEAR",
        expect_sku_key="CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK",
        mode="fast",
    )
    assert report["ok"] is True
    assert report["error_count"] == 0


def test_validate_offer_template_passes_for_women_thermal_underwear_category(tmp_path: Path) -> None:
    db_path = tmp_path / "app.db"
    xlsm_path = tmp_path / "offer_line31_thermal.xlsm"
    _setup_db(db_path)

    _insert_dim_sku(db_path, ["CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK"])
    _insert_article_map(
        db_path,
        "ACMEWEAR",
        "OF_LINE31_TRM_SB_S",
        "CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK",
    )
    _build_offer_xlsm(
        xlsm_path,
        merchant_sku="CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK_TRM_S",
        manufacturer_code="OF_LINE31_TRM_SB_S",
        color="черный",
        collection="Весна-Лето 2026",
    )

    report = validate_kaspi_offer_template(
        xlsm_path=xlsm_path,
        db_path=db_path,
        category="women-thermal-underwear",
        store_code="ACMEWEAR",
        expect_sku_key="CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK",
        mode="fast",
    )
    assert report["ok"] is True
    assert report["error_count"] == 0


def test_validate_offer_template_accepts_new_external_offer_token_when_expect_sku_key_is_locked(
    tmp_path: Path,
) -> None:
    db_path = tmp_path / "app.db"
    xlsm_path = tmp_path / "offer_line31_new_token.xlsm"
    _setup_db(db_path)

    _insert_dim_sku(db_path, ["CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK"])
    _build_offer_xlsm(
        xlsm_path,
        merchant_sku="CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK_ST_S",
        manufacturer_code="OF_LINE31_ST_SB_S",
        color="черный",
        collection="Весна-Лето 2026",
    )

    report = validate_kaspi_offer_template(
        xlsm_path=xlsm_path,
        db_path=db_path,
        category="women-sport-suits",
        store_code="ACMEWEAR",
        expect_sku_key="CL_OF_ARC_WM_LINE31_C-023_STARRY-BLACK",
        mode="fast",
    )
    assert report["ok"] is True


def test_validate_offer_template_balanced_checks_only_populated_list_bound_columns(tmp_path: Path) -> None:
    db_path = tmp_path / "app.db"
    xlsm_path = tmp_path / "offer_balanced.xlsm"
    _setup_db(db_path)
    _insert_dim_sku(db_path, ["CL_OC_MEN_LINE51_WHITE"])
    _insert_article_map(db_path, "ACMEWEAR", "OF_LINE51_K-O_XL_48", "CL_OC_MEN_LINE51_WHITE")
    _build_offer_xlsm(
        xlsm_path,
        merchant_sku="OF_LINE51_K-O_XL_48",
        manufacturer_code="OF_LINE51_K-O_XL_48",
        color="черный",
        collection="Весна-Лето 2026",
    )

    report = validate_kaspi_offer_template(
        xlsm_path=xlsm_path,
        db_path=db_path,
        category="men-sport-suits",
        store_code="ACMEWEAR",
        mode="balanced",
    )
    assert report["ok"] is True
    assert report["value_dict_scope"] == "populated_list_bound_columns"
    assert report["error_count"] == 0
