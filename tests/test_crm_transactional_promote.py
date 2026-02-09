from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from scripts.import_orders_to_crm import _promote_candidate_workbook


def _write_xlsx(path: Path, marker: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = marker
    wb.save(path)


def _read_xlsx_marker(path: Path) -> str:
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb.active
        return str(ws["A1"].value or "")
    finally:
        wb.close()


def test_promote_candidate_workbook_replaces_source_on_success(monkeypatch, tmp_path: Path):
    source = tmp_path / "source.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    failed_dir = tmp_path / "failed"
    _write_xlsx(source, "old")
    _write_xlsx(candidate, "new")

    monkeypatch.setattr(
        "scripts.import_orders_to_crm._verify_candidate_workbook",
        lambda *_args, **_kwargs: None,
        raising=False,
    )

    _promote_candidate_workbook(
        source_path=source,
        candidate_path=candidate,
        failed_dir=failed_dir,
        strict_excel=True,
        verbose=False,
    )

    assert _read_xlsx_marker(source) == "new"
    assert not candidate.exists()


def test_promote_candidate_workbook_keeps_source_and_moves_failed_candidate(monkeypatch, tmp_path: Path):
    source = tmp_path / "source.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    failed_dir = tmp_path / "failed"
    _write_xlsx(source, "stable")
    _write_xlsx(candidate, "bad")

    def _raise_verify(*_args, **_kwargs):
        raise RuntimeError("integrity failed")

    monkeypatch.setattr(
        "scripts.import_orders_to_crm._verify_candidate_workbook",
        _raise_verify,
        raising=False,
    )

    with pytest.raises(RuntimeError, match="integrity failed"):
        _promote_candidate_workbook(
            source_path=source,
            candidate_path=candidate,
            failed_dir=failed_dir,
            strict_excel=True,
            verbose=False,
        )

    assert _read_xlsx_marker(source) == "stable"
    failed_candidates = sorted(failed_dir.glob("candidate.failed_*.xlsx"))
    assert failed_candidates
    assert not candidate.exists()
