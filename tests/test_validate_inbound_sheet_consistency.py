from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import openpyxl


SCRIPT = Path(__file__).resolve().parent.parent / "scripts" / "validate_inbound_sheet_consistency.py"


def _make_workbook(path: Path, *, cargo_qty: int, actual_qty: int, ordered_qty: int = 3980) -> None:
    wb = openpyxl.Workbook()

    ws_inbounds = wb.active
    ws_inbounds.title = "Inbounds_sheet"
    ws_inbounds.append(["PO_part_id", "SKU_key", "Qty", "Actual_qty"])
    ws_inbounds.append(["PO-5.2", "CL_OC_MEN_LINE52_BLACK", ordered_qty, actual_qty])

    ws_totals = wb.create_sheet("PO_part_id_Totals")
    ws_totals.append(["PO_part_id", "Total Units"])
    ws_totals.append(["PO-5.2", actual_qty])

    ws_cargo = wb.create_sheet("Cargo_send_31.1.2026_PO-5.2_4.2")
    ws_cargo.append(["PO_part_id", "SKU_key", "Qty"])
    ws_cargo.append(["PO-5.2", "CL_OC_MEN_LINE52_BLACK", cargo_qty])

    wb.save(path)


def _run(xlsx: Path, *extra_args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--xlsx",
            str(xlsx),
            "--json",
            *extra_args,
        ],
        check=False,
        capture_output=True,
        text=True,
    )


def _make_styled_workbook(path: Path, *, qty: int = 3980) -> None:
    wb = openpyxl.Workbook()

    ws_inbounds = wb.active
    ws_inbounds.title = "Inbounds_sheet"
    ws_inbounds.append(["PO_part_id", "SKU_key", "Qty", "Actual_qty"])
    ws_inbounds.append(["PO-5.2", "CL_OC_MEN_LINE52_BLACK", qty, qty])

    ws_totals = wb.create_sheet("PO_part_id_Totals")
    ws_totals.append(["PO_part_id", "Total Units"])
    ws_totals.append(["PO-5.2", qty])

    ws_cargo = wb.create_sheet("Cargo_send_31.1.2026_PO-5.2_4.2")
    ws_cargo["A1"] = "SHIPMENT INFO"
    ws_cargo["A3"] = "PO_part_id"
    ws_cargo["B3"] = "PO-5.2"
    ws_cargo["A17"] = "SKU_Key"
    ws_cargo["B17"] = "Size"
    ws_cargo["C17"] = "Qty"
    ws_cargo["F17"] = "PO Name"
    ws_cargo["A18"] = "CL_OC_MEN_LINE52_BLACK"
    ws_cargo["B18"] = "M"
    ws_cargo["C18"] = qty
    ws_cargo["F18"] = "PO-5.2"

    wb.save(path)


def _make_line61_accepted_shortage_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()

    ws_inbounds = wb.active
    ws_inbounds.title = "Inbounds_sheet"
    ws_inbounds.append(["PO_part_id", "SKU_key", "Qty", "Actual_qty"])
    ws_inbounds.append(["PO-4.0", "CL_NEW-CLO2_MEN_SUIT-61_BLACK", 115, 92])
    ws_inbounds.append(["PO-4.0", "CL_OTHER_MEN_REFERENCE_BLACK", 1810, 1810])

    ws_totals = wb.create_sheet("PO_part_id_Totals")
    ws_totals.append(["PO_part_id", "Total Units"])
    ws_totals.append(["PO-4.0", 1925])

    ws_cargo = wb.create_sheet("Cargo_send_1.1.2026_PO-4.0")
    ws_cargo.append(["PO_part_id", "SKU_key", "Qty"])
    ws_cargo.append(["PO-4.0", "CL_NEW-CLO2_MEN_SUIT-61_BLACK", 115])
    ws_cargo.append(["PO-4.0", "CL_OTHER_MEN_REFERENCE_BLACK", 1810])

    wb.save(path)


def test_flags_qty_mismatch_between_cargo_and_inbounds(tmp_path: Path) -> None:
    xlsx = tmp_path / "inbound.xlsx"
    _make_workbook(xlsx, cargo_qty=3940, actual_qty=3980)

    proc = _run(xlsx)
    assert proc.returncode == 1
    payload = json.loads(proc.stdout)
    mismatches = payload["mismatches"]
    assert mismatches
    first = mismatches[0]
    assert first["po_part_id"] == "PO-5.2"
    assert first["sku_key"] == "CL_OC_MEN_LINE52_BLACK"
    assert first["expected_qty"] == 3980
    assert first["observed_qty"] == 3940


def test_reports_authoritative_sheet_precedence(tmp_path: Path) -> None:
    xlsx = tmp_path / "inbound.xlsx"
    # Qty != Actual_qty; validator must prefer Actual_qty as authoritative.
    _make_workbook(xlsx, cargo_qty=3980, actual_qty=3980, ordered_qty=4200)

    proc = _run(xlsx)
    assert proc.returncode == 0
    payload = json.loads(proc.stdout)
    assert payload["authoritative_sheet"] == "Inbounds_sheet"
    assert payload["authoritative_quantity_column"] == "Actual_qty"


def test_parses_styled_cargo_sheet_layout(tmp_path: Path) -> None:
    xlsx = tmp_path / "inbound_styled.xlsx"
    _make_styled_workbook(xlsx, qty=3980)

    proc = _run(xlsx)
    assert proc.returncode == 0
    payload = json.loads(proc.stdout)
    assert payload["ok"] is True
    assert payload["checked_keys"] == 1
    assert payload["mismatch_count"] == 0


def test_fails_when_no_comparable_cargo_keys_exist(tmp_path: Path) -> None:
    xlsx = tmp_path / "inbound_contract_breach.xlsx"
    _make_workbook(xlsx, cargo_qty=0, actual_qty=3980)

    proc = _run(xlsx)
    assert proc.returncode != 0
    payload = json.loads(proc.stdout)
    types = {m["type"] for m in payload["mismatches"]}
    assert "cargo_parse_contract" in types


def test_exit_nonzero_on_contract_breach(tmp_path: Path) -> None:
    xlsx = tmp_path / "inbound.xlsx"
    _make_workbook(xlsx, cargo_qty=3950, actual_qty=3980)

    proc = _run(xlsx)
    assert proc.returncode != 0
    payload = json.loads(proc.stdout)
    assert payload["ok"] is False
    assert payload["mismatch_count"] >= 1


def test_classifies_exact_line61_shortage_as_visible_production_safe_truth(tmp_path: Path) -> None:
    xlsx = tmp_path / "line61_shortage.xlsx"
    _make_line61_accepted_shortage_workbook(xlsx)

    proc = _run(xlsx)
    assert proc.returncode == 0
    payload = json.loads(proc.stdout)

    assert payload["ok"] is True
    assert payload["mismatch_count"] == 2
    assert payload["unknown_mismatch_count"] == 0
    assert payload["accepted_shortage_count"] == 2
    assert payload["unknown_mismatches"] == []
    assert payload["production_safe_accepted_shortage_applied"] is True

    accepted = payload["accepted_shortages"]
    assert {row["type"] for row in accepted} == {"cargo_vs_inbounds", "totals_vs_inbounds"}
    assert {
        row["classification_id"] for row in accepted
    } == {"PO_ACCEPTED_REAL_SHORTAGE_LINE61_2026_05_OWNER_CONFIRMED"}
    assert all(row["production_authority"] is True for row in accepted)
    assert all(row["clears_inbound_sheet_consistency"] is True for row in accepted)
    assert all(row["clears_po_money_gate"] is False for row in accepted)


def test_copied_temp_allowance_passes_only_exact_line61_shortage(tmp_path: Path) -> None:
    xlsx = tmp_path / "line61_shortage.xlsx"
    _make_line61_accepted_shortage_workbook(xlsx)

    proc = _run(xlsx, "--allow-accepted-shortages-for-copied-temp")
    assert proc.returncode == 0
    payload = json.loads(proc.stdout)

    assert payload["ok"] is True
    assert payload["mismatch_count"] == 2
    assert payload["unknown_mismatch_count"] == 0
    assert payload["accepted_shortage_count"] == 2
    assert payload["copied_temp_accepted_shortage_allowance_applied"] is True
    assert payload["production_safe_accepted_shortage_applied"] is True
    assert all(row["clears_po_money_gate"] is False for row in payload["accepted_shortages"])
