from __future__ import annotations

from datetime import date
import json
import sqlite3
from pathlib import Path

from openpyxl import Workbook

from scripts.validate_weekly_cash_reanchor_cadence import (
    validate_weekly_cash_reanchor_cadence,
)


def _write_cash_workbook(path: Path, *, snapshot_label: str = "13.06.2026_01_01_26") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cash_Balances"
    ws.cell(4, 1, "Store")
    ws.cell(4, 2, "Account")
    ws.cell(4, 3, "Currency")
    ws.cell(4, 13, snapshot_label)
    rows = [
        ("UNIVERSAL", "kaspi_gold", "KZT", 1),
        ("UNIVERSAL", "kaspi_pay", "KZT", 1),
        ("UNIVERSAL", "bcc", "KZT", 1),
        ("UNIVERSAL", "freedom", "KZT", 1),
        ("UNIVERSAL", "cash_kzt", "KZT", 1),
        ("UNIVERSAL", "cash_usd", "USD", 1),
        ("UNIVERSAL", "cash_rub", "RUB", 1),
        ("UNIVERSAL", "binance_usdt", "USDT", 1),
        ("11KZ", "kaspi_gold", "KZT", 1),
        ("11KZ", "kaspi_pay", "KZT", 1),
        ("11KZ", "binance_usdt", "USDT", 1),
        ("STOREB", "kaspi_gold", "KZT", 1),
        ("STOREB", "kaspi_pay", "KZT", 1),
        ("ACMEWEAR", "kaspi_gold", "KZT", 1),
        ("ACMEWEAR", "kaspi_pay", "KZT", 1),
        ("ACMEWEAR", "cash_kzt", "KZT", 1),
        ("MELVIS", "kaspi_gold", "KZT", 1),
        ("MELVIS", "kaspi_pay", "KZT", 1),
    ]
    for row_idx, row in enumerate(rows, start=5):
        for col_idx, value in enumerate(row[:3], start=1):
            ws.cell(row_idx, col_idx, value)
        ws.cell(row_idx, 13, row[3])
    wb.save(path)


def _write_db(db_path: Path, *, anchor_date: str = "2026-06-13", anchor_total: float = 3937364.0) -> None:
    conn = sqlite3.connect(db_path)
    try:
        conn.execute(
            """
            CREATE TABLE cashflow_cash_anchor (
                anchor_date TEXT,
                source_store_name TEXT,
                anchor_closing_balance_kzt REAL,
                created_by_run_id TEXT
            )
            """
        )
        for idx in range(18):
            conn.execute(
                """
                INSERT INTO cashflow_cash_anchor (
                    anchor_date, source_store_name, anchor_closing_balance_kzt, created_by_run_id
                ) VALUES (?, ?, ?, ?)
                """,
                (anchor_date, f"STORE:acct{idx}:KZT", anchor_total / 18, "test_cash_anchor"),
            )
        conn.execute(
            """
            CREATE TABLE fact_cashflow_daily (
                date TEXT PRIMARY KEY,
                cash_close REAL,
                run_id TEXT
            )
            """
        )
        conn.execute(
            "INSERT INTO fact_cashflow_daily (date, cash_close, run_id) VALUES (?, ?, ?)",
            (anchor_date, anchor_total, "test_cash_anchor"),
        )
        conn.commit()
    finally:
        conn.close()


def test_weekly_cash_reanchor_cadence_passes_for_current_anchor(tmp_path: Path) -> None:
    db = tmp_path / "app.db"
    workbook = tmp_path / "cash.xlsx"
    _write_db(db)
    _write_cash_workbook(workbook)

    report = validate_weekly_cash_reanchor_cadence(
        db_path=db,
        workbook_path=workbook,
        as_of=date(2026, 6, 14),
        max_age_days=7,
        expected_anchor_rows=18,
        output_root=tmp_path / "out",
    )

    assert report["status"] == "PASS"
    assert report["anchor"]["anchor_date"] == "2026-06-13"
    assert report["workbook"]["snapshot_date"] == "2026-06-13"
    saved = json.loads(Path(report["json_path"]).read_text(encoding="utf-8"))
    assert saved["status"] == "PASS"


def test_weekly_cash_reanchor_cadence_fails_when_anchor_is_stale(tmp_path: Path) -> None:
    db = tmp_path / "app.db"
    workbook = tmp_path / "cash.xlsx"
    _write_db(db)
    _write_cash_workbook(workbook)

    report = validate_weekly_cash_reanchor_cadence(
        db_path=db,
        workbook_path=workbook,
        as_of=date(2026, 6, 22),
        max_age_days=7,
        expected_anchor_rows=18,
        output_root=tmp_path / "out",
    )

    assert report["status"] == "FAIL"
    assert any("cash anchor is stale" in error for error in report["errors"])


def test_weekly_cash_reanchor_cadence_fails_when_workbook_is_newer_than_anchor(
    tmp_path: Path,
) -> None:
    db = tmp_path / "app.db"
    workbook = tmp_path / "cash.xlsx"
    _write_db(db)
    _write_cash_workbook(workbook, snapshot_label="14.06.2026_01_01_26")

    report = validate_weekly_cash_reanchor_cadence(
        db_path=db,
        workbook_path=workbook,
        as_of=date(2026, 6, 14),
        max_age_days=7,
        expected_anchor_rows=18,
        output_root=tmp_path / "out",
    )

    assert report["status"] == "FAIL"
    assert any("newer than latest DB anchor" in error for error in report["errors"])


def test_weekly_cash_reanchor_cadence_fails_on_cash_close_mismatch(tmp_path: Path) -> None:
    db = tmp_path / "app.db"
    workbook = tmp_path / "cash.xlsx"
    _write_db(db, anchor_total=1000.0)
    _write_cash_workbook(workbook)
    with sqlite3.connect(db) as conn:
        conn.execute("UPDATE fact_cashflow_daily SET cash_close=999 WHERE date='2026-06-13'")
        conn.commit()

    report = validate_weekly_cash_reanchor_cadence(
        db_path=db,
        workbook_path=workbook,
        as_of=date(2026, 6, 14),
        max_age_days=7,
        expected_anchor_rows=18,
        output_root=tmp_path / "out",
    )

    assert report["status"] == "FAIL"
    assert any("cash_close does not match" in error for error in report["errors"])
