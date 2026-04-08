from __future__ import annotations

import re
import zipfile
import zlib
from pathlib import Path
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.formula import ArrayFormula

from scripts.validate_crm_workbook_integrity import (
    filter_integrity_errors,
    repair_missing_shared_strings_part,
    validate_workbook_integrity,
)


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def _make_workbook_with_table(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SALES_KSP_CRM_1"
    ws["A1"] = "id"
    ws["B1"] = "value"
    ws["A2"] = "1"
    ws["B2"] = "x"
    table = Table(displayName="tb_SalesRaw", ref="A1:B2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    wb.save(path)


def _make_sales_workbook_with_sku_key_formula(path: Path, *, array_formula: bool) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SALES_KSP_CRM_1"
    headers = ["Date", "Phone", "Spacer", "Статус", "SKU_key", "SKU_ID_KSP"]
    for idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=idx, value=header)
    ws.cell(row=2, column=1, value="2026-04-04")
    ws.cell(row=2, column=2, value="77770000000")
    ws.cell(row=2, column=3, value="seed")
    ws.cell(row=2, column=4, value="Новый")
    if array_formula:
        ws.cell(row=2, column=5, value=ArrayFormula(ref="E2", text='=F2&"_SKU"'))
    else:
        ws.cell(row=2, column=5, value='=F2&"_SKU"')
    ws.cell(row=2, column=6, value="OF_LINE31_XL")
    table = Table(displayName="tb_SalesRaw", ref="A1:F2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    for table_col in table.tableColumns:
        if table_col.name == "SKU_key":
            if array_formula:
                from openpyxl.worksheet.table import TableFormula

                table_col.calculatedColumnFormula = TableFormula(array=True, attr_text='F2&"_SKU"')
    wb.save(path)
    wb.close()

    if array_formula:
        with zipfile.ZipFile(path, "r") as zin:
            table_xml = ET.fromstring(zin.read("xl/tables/table1.xml"))
            columns = table_xml.find(f"{{{MAIN_NS}}}tableColumns")
            assert columns is not None
            for table_col in columns.findall(f"{{{MAIN_NS}}}tableColumn"):
                if table_col.attrib.get("name") != "SKU_key":
                    continue
                calc = ET.SubElement(table_col, f"{{{MAIN_NS}}}calculatedColumnFormula")
                calc.set("array", "1")
                calc.text = 'F2&"_SKU"'
                break
            modified = ET.tostring(table_xml, encoding="utf-8", xml_declaration=True)

            tmp = path.with_name(f"{path.stem}_sku_formula_tmp.xlsx")
            with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    if item.filename == "xl/tables/table1.xml":
                        zout.writestr(item, modified)
                    else:
                        zout.writestr(item, zin.read(item.filename))
        tmp.replace(path)


def _make_sales_workbook_with_kaspi_name_core_cf(path: Path, *, include_cf_sentinel: bool) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SALES_KSP_CRM_1"
    headers = ["Date", "Phone", "Spacer", "Статус", "Kaspi_name_core", "SKU_key", "SKU_ID_KSP"]
    for idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=idx, value=header)
    ws.cell(row=2, column=1, value="2026-04-04")
    ws.cell(row=2, column=2, value="77770000000")
    ws.cell(row=2, column=3, value="seed")
    ws.cell(row=2, column=4, value="Новый")
    ws.cell(row=2, column=5, value="Line51")
    ws.cell(row=2, column=6, value=ArrayFormula(ref="F2", text='=G2&"_SKU"'))
    ws.cell(row=2, column=7, value="OF_LINE31_XL")
    table = Table(displayName="tb_SalesRaw", ref="A1:G2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    if include_cf_sentinel:
        fill = PatternFill(fill_type="solid", fgColor="FFFF7E79")
        ws.conditional_formatting.add(
            "O2:Q2",
            FormulaRule(formula=['$M2="CL_OC_MEN_LINE52_BLACK"'], fill=fill),
        )
    wb.save(path)
    wb.close()


def _inject_array_formula_ref_mismatch(
    path: Path,
    *,
    cell_ref: str,
    bad_ref: str,
    sheet_xml_path: str = "xl/worksheets/sheet1.xml",
) -> None:
    with zipfile.ZipFile(path, "r") as zin:
        sheet_xml = zin.read(sheet_xml_path).decode("utf-8", "ignore")
        pattern = re.compile(
            rf'(<c\b[^>]*\br="{re.escape(cell_ref)}"[^>]*>.*?<f\b[^>]*\bt="array"[^>]*\bref=")([^"]*)(")',
            re.S,
        )
        updated, count = pattern.subn(lambda m: f"{m.group(1)}{bad_ref}{m.group(3)}", sheet_xml, count=1)
        assert count == 1

        tmp = path.with_name(f"{path.stem}_bad_array_ref_tmp.xlsx")
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == sheet_xml_path:
                    zout.writestr(item, updated.encode("utf-8"))
                else:
                    zout.writestr(item, zin.read(item.filename))
    tmp.replace(path)


def _inject_self_external_sales_formula(
    path: Path,
    *,
    cell_ref: str,
    external_index: int = 4,
    sheet_xml_path: str = "xl/worksheets/sheet1.xml",
) -> None:
    with zipfile.ZipFile(path, "r") as zin:
        sheet_xml = zin.read(sheet_xml_path).decode("utf-8", "ignore")
        pattern = re.compile(
            rf'(<c\b[^>]*\br="{re.escape(cell_ref)}"[^>]*>.*?<f\b[^>]*>)(.*?)(</f>)',
            re.S,
        )
        replacement = (
            rf"\1_xlfn.XLOOKUP([{external_index}]!tb_SalesRaw[[#This Row],[SKU_ID_KSP]],"
            rf"[{external_index}]!CRM_3[SKU_ID_KSP_v2],"
            rf"[{external_index}]!CRM_3[SKU_key],\"\",0,1)\3"
        )
        updated, count = pattern.subn(replacement, sheet_xml, count=1)
        assert count == 1

        tmp = path.with_name(f"{path.stem}_self_external_formula_tmp.xlsx")
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == sheet_xml_path:
                    zout.writestr(item, updated.encode("utf-8"))
                else:
                    zout.writestr(item, zin.read(item.filename))
    tmp.replace(path)


def _inject_external_link_with_missing_primary_relationship(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as zin:
        workbook_xml = zin.read("xl/workbook.xml").decode("utf-8", "ignore")
        workbook_rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8", "ignore")

        workbook_xml = workbook_xml.replace(
            "</workbook>",
            (
                '<externalReferences xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '<externalReference r:id="rIdExt1"/>'
                "</externalReferences></workbook>"
            ),
            1,
        )
        workbook_rels = workbook_rels.replace(
            "</Relationships>",
            (
                '<Relationship Id="rIdExt1" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink" '
                'Target="externalLinks/externalLink1.xml"/>'
                "</Relationships>"
            ),
            1,
        )

        external_link_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'mc:Ignorable="x14 xxl21" '
            'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
            'xmlns:xxl21="http://schemas.microsoft.com/office/spreadsheetml/2021/extlinks2021">'
            '<externalBook xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1">'
            '<xxl21:alternateUrls><xxl21:absoluteUrl r:id="rId2"/></xxl21:alternateUrls>'
            '<sheetNames><sheetName val="Sheet1"/></sheetNames>'
            '<sheetDataSet><sheetData sheetId="0"/></sheetDataSet>'
            "</externalBook></externalLink>"
        ).encode("utf-8")
        external_link_rels = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId2" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath" '
            'Target="/tmp/reference.xlsx" TargetMode="External"/>'
            "</Relationships>"
        ).encode("utf-8")

        tmp = path.with_name(f"{path.stem}_broken_external_link_tmp.xlsx")
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/workbook.xml":
                    zout.writestr(item, workbook_xml.encode("utf-8"))
                elif item.filename == "xl/_rels/workbook.xml.rels":
                    zout.writestr(item, workbook_rels.encode("utf-8"))
                else:
                    zout.writestr(item, zin.read(item.filename))
            zout.writestr("xl/externalLinks/externalLink1.xml", external_link_xml)
            zout.writestr("xl/externalLinks/_rels/externalLink1.xml.rels", external_link_rels)
    tmp.replace(path)


def _make_workbook_with_sales_and_secondary_sku_table(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SALES_KSP_CRM_1"
    headers = ["Date", "Phone", "Spacer", "Статус", "Kaspi_name_core", "SKU_key", "SKU_ID_KSP"]
    for idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=idx, value=header)
    ws.cell(row=2, column=1, value="2026-04-04")
    ws.cell(row=2, column=2, value="77770000000")
    ws.cell(row=2, column=3, value="seed")
    ws.cell(row=2, column=4, value="Новый")
    ws.cell(row=2, column=5, value="Line51")
    ws.cell(row=2, column=6, value=ArrayFormula(ref="F2", text='=G2&"_SKU"'))
    ws.cell(row=2, column=7, value="OF_LINE31_XL")
    sales_table = Table(displayName="tb_SalesRaw", ref="A1:G2")
    sales_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(sales_table)
    fill = PatternFill(fill_type="solid", fgColor="FFFF7E79")
    ws.conditional_formatting.add(
        "O2:Q2",
        FormulaRule(formula=['$M2="CL_OC_MEN_LINE52_BLACK"'], fill=fill),
    )

    ws2 = wb.create_sheet("M02_SKU_CATALOG_NC")
    ws2.append(["SKU_key", "Value"])
    ws2.append(['=B2&"_CATALOG"', "seed"])
    secondary_table = Table(displayName="CRM_3", ref="A1:B2")
    secondary_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws2.add_table(secondary_table)

    wb.save(path)
    wb.close()

    with zipfile.ZipFile(path, "r") as zin:
        table_xml = ET.fromstring(zin.read("xl/tables/table1.xml"))
        columns = table_xml.find(f"{{{MAIN_NS}}}tableColumns")
        assert columns is not None
        for table_col in columns.findall(f"{{{MAIN_NS}}}tableColumn"):
            if table_col.attrib.get("name") != "SKU_key":
                continue
            calc = ET.SubElement(table_col, f"{{{MAIN_NS}}}calculatedColumnFormula")
            calc.set("array", "1")
            calc.text = 'G2&"_SKU"'
            break
        modified = ET.tostring(table_xml, encoding="utf-8", xml_declaration=True)

        tmp = path.with_name(f"{path.stem}_mixed_tables_tmp.xlsx")
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/tables/table1.xml":
                    zout.writestr(item, modified)
                else:
                    zout.writestr(item, zin.read(item.filename))
    tmp.replace(path)


def _inject_broken_defined_name(path: Path, name: str = "BROKEN_NAME") -> None:
    with zipfile.ZipFile(path, "r") as zin:
        wb_xml = ET.fromstring(zin.read("xl/workbook.xml"))
        dns = wb_xml.find(f"{{{MAIN_NS}}}definedNames")
        if dns is None:
            dns = ET.SubElement(wb_xml, f"{{{MAIN_NS}}}definedNames")
        dn = ET.SubElement(dns, f"{{{MAIN_NS}}}definedName", {"name": name})
        dn.text = "#REF!"
        wb_modified = ET.tostring(wb_xml, encoding="utf-8", xml_declaration=True)

        tmp = path.with_name(f"{path.stem}_broken_names_tmp.xlsx")
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/workbook.xml":
                    zout.writestr(item, wb_modified)
                else:
                    zout.writestr(item, zin.read(item.filename))
    tmp.replace(path)


def _inject_missing_workbook_relationship(
    path: Path,
    *,
    rel_id: str,
    rel_type: str,
    target: str,
) -> None:
    with zipfile.ZipFile(path, "r") as zin:
        rels_root = ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
        ET.SubElement(
            rels_root,
            f"{{{PKG_REL_NS}}}Relationship",
            {
                "Id": rel_id,
                "Type": rel_type,
                "Target": target,
            },
        )
        rels_modified = ET.tostring(rels_root, encoding="utf-8", xml_declaration=True)

        tmp = path.with_name(f"{path.stem}_missing_workbook_rel_tmp.xlsx")
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/_rels/workbook.xml.rels":
                    zout.writestr(item, rels_modified)
                else:
                    zout.writestr(item, zin.read(item.filename))
    tmp.replace(path)


def _inject_undeclared_ignorable_prefixes(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as zin:
        workbook_xml = zin.read("xl/workbook.xml").decode("utf-8")
        workbook_xml = workbook_xml.replace(
            "<workbook ",
            '<workbook xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'mc:Ignorable="x15 xr xr6 xr10 xr2" ',
            1,
        )
        workbook_xml = workbook_xml.replace('xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main" ', "")
        workbook_xml = workbook_xml.replace('xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision" ', "")
        workbook_xml = workbook_xml.replace('xmlns:xr6="http://schemas.microsoft.com/office/spreadsheetml/2016/revision6" ', "")
        workbook_xml = workbook_xml.replace('xmlns:xr10="http://schemas.microsoft.com/office/spreadsheetml/2016/revision10" ', "")
        workbook_xml = workbook_xml.replace('xmlns:xr2="http://schemas.microsoft.com/office/spreadsheetml/2015/revision2" ', "")

        tmp = path.with_name(f"{path.stem}_broken_ignorable_tmp.xlsx")
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/workbook.xml":
                    zout.writestr(item, workbook_xml.encode("utf-8"))
                else:
                    zout.writestr(item, zin.read(item.filename))
    tmp.replace(path)


def _inject_missing_shared_strings_reference(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as zin:
        wb_rels = ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
        rid = "rIdSharedStrings"
        ET.SubElement(
            wb_rels,
            f"{{{PKG_REL_NS}}}Relationship",
            {
                "Id": rid,
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
                "Target": "sharedStrings.xml",
            },
        )
        wb_rels_modified = ET.tostring(wb_rels, encoding="utf-8", xml_declaration=True)

        ct_root = ET.fromstring(zin.read("[Content_Types].xml"))
        ET.SubElement(
            ct_root,
            f"{{{CT_NS}}}Override",
            {
                "PartName": "/xl/sharedStrings.xml",
                "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
            },
        )
        ct_modified = ET.tostring(ct_root, encoding="utf-8", xml_declaration=True)

        tmp = path.with_name(f"{path.stem}_missing_ss_tmp.xlsx")
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/_rels/workbook.xml.rels":
                    zout.writestr(item, wb_rels_modified)
                elif item.filename == "[Content_Types].xml":
                    zout.writestr(item, ct_modified)
                else:
                    zout.writestr(item, zin.read(item.filename))
    tmp.replace(path)


def _inject_stale_calcchain_reference(
    path: Path,
    *,
    sheet_id: str = "1",
    cell_ref: str = "Z999",
) -> None:
    with zipfile.ZipFile(path, "r") as zin:
        wb_rels = ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
        existing = {
            rel.attrib.get("Type", "")
            for rel in wb_rels.findall(f"{{{PKG_REL_NS}}}Relationship")
        }
        calc_rel_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain"
        if calc_rel_type not in existing:
            ET.SubElement(
                wb_rels,
                f"{{{PKG_REL_NS}}}Relationship",
                {
                    "Id": "rIdCalcChain",
                    "Type": calc_rel_type,
                    "Target": "calcChain.xml",
                },
            )
        wb_rels_modified = ET.tostring(wb_rels, encoding="utf-8", xml_declaration=True)

        tmp = path.with_name(f"{path.stem}_stale_calcchain_tmp.xlsx")
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/_rels/workbook.xml.rels":
                    zout.writestr(item, wb_rels_modified)
                else:
                    zout.writestr(item, zin.read(item.filename))
            zout.writestr(
                "xl/calcChain.xml",
                (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                    f'<c r="{cell_ref}" i="{sheet_id}"/>'
                    "</calcChain>"
                ),
            )
    tmp.replace(path)


def test_validate_workbook_integrity_ok(tmp_path: Path):
    wb_path = tmp_path / "ok.xlsx"
    _make_workbook_with_table(wb_path)

    result = validate_workbook_integrity(wb_path)
    assert result.errors == []


def test_validate_workbook_integrity_detects_table_column_mismatch(tmp_path: Path):
    wb_path = tmp_path / "bad_table.xlsx"
    _make_workbook_with_table(wb_path)

    with zipfile.ZipFile(wb_path, "r") as zin:
        table_xml = ET.fromstring(zin.read("xl/tables/table1.xml"))
        table_cols = table_xml.find(f"{{{MAIN_NS}}}tableColumns")
        assert table_cols is not None
        table_cols.set("count", "3")
        modified = ET.tostring(table_xml, encoding="utf-8", xml_declaration=True)

        tmp = tmp_path / "bad_table_tmp.xlsx"
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/tables/table1.xml":
                    zout.writestr(item, modified)
                else:
                    zout.writestr(item, zin.read(item.filename))
    tmp.replace(wb_path)

    result = validate_workbook_integrity(wb_path)
    assert any("table column mismatch" in err for err in result.errors)


def test_validate_workbook_integrity_detects_missing_table_target(tmp_path: Path):
    wb_path = tmp_path / "missing_target.xlsx"
    _make_workbook_with_table(wb_path)

    with zipfile.ZipFile(wb_path, "r") as zin:
        tmp = tmp_path / "missing_target_tmp.xlsx"
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/tables/table1.xml":
                    continue
                zout.writestr(item, zin.read(item.filename))
    tmp.replace(wb_path)

    result = validate_workbook_integrity(wb_path)
    assert any("tablePart target missing" in err for err in result.errors)


def test_validate_workbook_integrity_accepts_working_array_formula_contract_on_sku_key(tmp_path: Path):
    wb_path = tmp_path / "sku_key_array.xlsx"
    _make_sales_workbook_with_sku_key_formula(wb_path, array_formula=True)

    result = validate_workbook_integrity(wb_path)
    assert not any("SKU_key calculatedColumnFormula" in err for err in result.errors)
    assert not any("SKU_key contains" in err for err in result.errors)


def test_validate_workbook_integrity_detects_missing_working_array_formula_contract_on_sku_key(tmp_path: Path):
    wb_path = tmp_path / "sku_key_plain.xlsx"
    _make_sales_workbook_with_sku_key_formula(wb_path, array_formula=False)

    result = validate_workbook_integrity(wb_path)
    assert any("SKU_key calculatedColumnFormula must use array=1" in err for err in result.errors)
    assert any("SKU_key must preserve per-row array formulas" in err for err in result.errors)


def test_validate_workbook_integrity_detects_mismatched_sku_key_array_formula_refs(tmp_path: Path):
    wb_path = tmp_path / "sku_key_bad_ref.xlsx"
    _make_sales_workbook_with_sku_key_formula(wb_path, array_formula=True)
    _inject_array_formula_ref_mismatch(wb_path, cell_ref="E2", bad_ref="E3")

    result = validate_workbook_integrity(wb_path)
    assert any("SKU_key array formula refs must match owning cell" in err for err in result.errors)


def test_validate_workbook_integrity_detects_missing_kaspi_name_core_cf_sentinel(tmp_path: Path):
    wb_path = tmp_path / "kaspi_name_core_cf_missing.xlsx"
    _make_sales_workbook_with_kaspi_name_core_cf(wb_path, include_cf_sentinel=False)

    result = validate_workbook_integrity(wb_path)
    assert any("Kaspi_name_core conditional-formatting sentinel missing" in err for err in result.errors)


def test_validate_workbook_integrity_accepts_kaspi_name_core_cf_sentinel(tmp_path: Path):
    wb_path = tmp_path / "kaspi_name_core_cf_ok.xlsx"
    _make_sales_workbook_with_kaspi_name_core_cf(wb_path, include_cf_sentinel=True)

    result = validate_workbook_integrity(wb_path)
    assert not any("Kaspi_name_core conditional-formatting sentinel missing" in err for err in result.errors)


def test_validate_workbook_integrity_only_enforces_array_formula_contract_on_sales_table(tmp_path: Path):
    wb_path = tmp_path / "mixed_tables.xlsx"
    _make_workbook_with_sales_and_secondary_sku_table(wb_path)

    result = validate_workbook_integrity(wb_path)
    assert not any("SKU_key calculatedColumnFormula must use array=1" in err for err in result.errors)
    assert not any("SKU_key must preserve per-row array formulas" in err for err in result.errors)
    assert not any("Kaspi_name_core conditional-formatting sentinel missing" in err for err in result.errors)


def test_validate_workbook_integrity_treats_ref_named_ranges_as_errors(tmp_path: Path):
    wb_path = tmp_path / "broken_name.xlsx"
    _make_workbook_with_table(wb_path)
    _inject_broken_defined_name(wb_path)

    result = validate_workbook_integrity(wb_path)
    assert any("named range contains #REF!" in err for err in result.errors)


def test_validate_workbook_integrity_detects_missing_shared_strings_part(tmp_path: Path):
    wb_path = tmp_path / "missing_shared_strings.xlsx"
    _make_workbook_with_table(wb_path)
    _inject_missing_shared_strings_reference(wb_path)

    result = validate_workbook_integrity(wb_path)
    assert any("sharedStrings target missing" in err for err in result.errors)


def test_validate_workbook_integrity_detects_missing_calcchain_target(tmp_path: Path):
    wb_path = tmp_path / "missing_calcchain.xlsx"
    _make_workbook_with_table(wb_path)
    _inject_missing_workbook_relationship(
        wb_path,
        rel_id="rId999",
        rel_type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain",
        target="calcChain.xml",
    )

    result = validate_workbook_integrity(wb_path)
    assert any("workbook relationship target missing" in err and "calcChain.xml" in err for err in result.errors)


def test_validate_workbook_integrity_detects_stale_calcchain_refs(tmp_path: Path):
    wb_path = tmp_path / "stale_calcchain.xlsx"
    _make_workbook_with_table(wb_path)
    _inject_stale_calcchain_reference(wb_path)

    result = validate_workbook_integrity(wb_path)
    assert any("calcChain contains refs to non-formula cells" in err for err in result.errors)


def test_validate_workbook_integrity_detects_self_externalized_sales_formulas(tmp_path: Path):
    wb_path = tmp_path / "self_externalized.xlsx"
    _make_sales_workbook_with_sku_key_formula(wb_path, array_formula=True)
    _inject_self_external_sales_formula(wb_path, cell_ref="E2", external_index=4)

    result = validate_workbook_integrity(wb_path)
    assert any("CRM sales sheet contains self-externalized formulas" in err for err in result.errors)


def test_validate_workbook_integrity_detects_external_link_relationship_id_mismatch(tmp_path: Path):
    wb_path = tmp_path / "broken_external_link.xlsx"
    _make_workbook_with_table(wb_path)
    _inject_external_link_with_missing_primary_relationship(wb_path)

    result = validate_workbook_integrity(wb_path)
    assert any(
        "externalLink relationship ids missing" in err and "externalLink1.xml" in err and "rId1" in err
        for err in result.errors
    )


def test_validate_workbook_integrity_detects_missing_sheetmetadata_target(tmp_path: Path):
    wb_path = tmp_path / "missing_metadata.xlsx"
    _make_workbook_with_table(wb_path)
    _inject_missing_workbook_relationship(
        wb_path,
        rel_id="rId998",
        rel_type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sheetMetadata",
        target="metadata.xml",
    )

    result = validate_workbook_integrity(wb_path)
    assert any("workbook relationship target missing" in err and "metadata.xml" in err for err in result.errors)


def test_validate_workbook_integrity_detects_undeclared_ignorable_prefixes(tmp_path: Path):
    wb_path = tmp_path / "broken_ignorable.xlsx"
    _make_workbook_with_table(wb_path)
    _inject_undeclared_ignorable_prefixes(wb_path)

    result = validate_workbook_integrity(wb_path)
    assert any("workbook mc:Ignorable references undeclared prefixes" in err for err in result.errors)


def test_repair_missing_shared_strings_part_adds_empty_part(tmp_path: Path):
    wb_path = tmp_path / "repair_shared_strings.xlsx"
    _make_workbook_with_table(wb_path)
    _inject_missing_shared_strings_reference(wb_path)

    repaired, backup = repair_missing_shared_strings_part(
        wb_path,
        backup_dir=tmp_path / "backups",
    )

    assert repaired is True
    assert backup is not None and backup.exists()

    with zipfile.ZipFile(wb_path, "r") as zf:
        assert "xl/sharedStrings.xml" in set(zf.namelist())

    wb = openpyxl.load_workbook(wb_path, read_only=True)
    wb.close()

    result = validate_workbook_integrity(wb_path)
    assert not any("sharedStrings target missing" in err for err in result.errors)


def test_filter_integrity_errors_allows_prefix():
    errors = [
        "named range contains #REF!: B",
        "table column mismatch: xl/tables/table1.xml width=2 count_attr=3 count_nodes=2",
    ]
    blocking, allowed = filter_integrity_errors(
        errors,
        allow_exact=[],
        allow_prefix=["named range contains #REF!:"],
    )
    assert allowed == ["named range contains #REF!: B"]
    assert blocking == [errors[1]]


def test_filter_integrity_errors_allows_exact_only():
    errors = [
        "named range contains #REF!: B",
        "named range contains #REF!: SS_TOTAL",
    ]
    blocking, allowed = filter_integrity_errors(
        errors,
        allow_exact=["named range contains #REF!: SS_TOTAL"],
        allow_prefix=[],
    )
    assert allowed == ["named range contains #REF!: SS_TOTAL"]
    assert blocking == ["named range contains #REF!: B"]


def test_validate_workbook_integrity_reports_unreadable_zip_member(tmp_path: Path, monkeypatch):
    wb_path = tmp_path / "zip_corrupt_like.xlsx"
    _make_workbook_with_table(wb_path)

    original_read = zipfile.ZipFile.read

    def _patched_read(self, name, *args, **kwargs):  # type: ignore[no-untyped-def]
        if name == "xl/workbook.xml":
            raise zlib.error("Error -3 while decompressing data: invalid code lengths set")
        return original_read(self, name, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "read", _patched_read)

    result = validate_workbook_integrity(wb_path)
    assert any("zip entry unreadable: xl/workbook.xml" in err for err in result.errors)


def test_validate_workbook_integrity_reports_unreadable_non_core_xl_part(
    tmp_path: Path, monkeypatch
):
    wb_path = tmp_path / "zip_corrupt_pivot_cache_like.xlsx"
    _make_workbook_with_table(wb_path)

    # Add an extra xl/ part to simulate a workbook member that openpyxl/Excel may touch later.
    with zipfile.ZipFile(wb_path, "r") as zin:
        tmp = tmp_path / "zip_corrupt_pivot_cache_like_tmp.xlsx"
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, zin.read(item.filename))
            zout.writestr("xl/pivotCache/pivotCacheRecords1.xml", b"<pivotCacheRecords/>")
    tmp.replace(wb_path)

    original_read = zipfile.ZipFile.read

    def _patched_read(self, name, *args, **kwargs):  # type: ignore[no-untyped-def]
        if name == "xl/pivotCache/pivotCacheRecords1.xml":
            raise zlib.error("Error -3 while decompressing data: invalid code lengths set")
        return original_read(self, name, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "read", _patched_read)

    result = validate_workbook_integrity(wb_path)
    assert any(
        "zip entry unreadable: xl/pivotCache/pivotCacheRecords1.xml" in err
        for err in result.errors
    )
