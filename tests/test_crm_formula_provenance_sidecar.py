from __future__ import annotations

import hashlib
import json
import shutil
import sqlite3
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Callable

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.formula import ArrayFormula

from scripts.build_crm_formula_provenance_sidecar import (
    _canonical_cell,
    _is_formula_cell,
    _read_workbook_rows,
    build_sidecar,
    validate_manifest,
)
from scripts.apply_crm_formula_provenance_sidecar import promote_crm_sidecar
from scripts.apply_sales_formula_provenance_sidecar import SidecarApplyError
from scripts.compare_sales_formula_promotion_db_state import compare as compare_promotion


HEADERS = [
    "Date",
    "STORE_NAME",
    "Quantity",
    "OrderID",
    "MY_SIZE",
    "SKU_key",
    "SKU_ID",
    "Sell_price_kzt",
    "Delivery_fee_kzt",
    "№ заказа",
    "Дата поступления заказа",
    "Артикул",
    "Сумма",
    "Статус",
    "Причина отмены",
    "Количество",
    "Стоимость доставки для продавца",
    "Phone",
    "Customer First Name",
    "Delivery Address",
]


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _init_db(path: Path) -> None:
    with sqlite3.connect(path) as conn:
        conn.executescript(
            """
            CREATE TABLE sales_fact_v2 (
                sale_id INTEGER PRIMARY KEY,
                order_id TEXT,
                order_date TEXT,
                sku_key TEXT,
                sku_id TEXT,
                my_size TEXT,
                store_code TEXT,
                quantity REAL,
                sell_price_kzt REAL,
                delivery_fee REAL,
                cogs REAL,
                net_rev REAL,
                profit REAL,
                status TEXT,
                return_flag INTEGER,
                source_file TEXT,
                source_entry_id TEXT,
                kaspi_article TEXT,
                line_identity_key TEXT
            );
            CREATE VIEW view_sales_line_truth AS
            SELECT order_id, date(order_date) AS sale_date, store_code,
                   sku_key, sku_id, my_size, quantity AS units,
                   net_rev AS net_rev_kzt, cogs AS cogs_kzt,
                   profit AS profit_kzt, 'legacy' AS cogs_source,
                   'sales_fact_v2' AS source_table,
                   sku_key AS source_sku_key, sku_id AS source_sku_id,
                   quantity AS source_units, net_rev AS source_net_rev_kzt,
                   cogs AS source_cogs_kzt, profit AS source_profit_kzt
            FROM sales_fact_v2
            WHERE status='DELIVERED' AND COALESCE(return_flag, 0)=0;
            CREATE TABLE fact_order_status_observations (
                id INTEGER PRIMARY KEY,
                order_id TEXT,
                store_code TEXT,
                status_internal TEXT,
                observed_at TEXT,
                source TEXT,
                ledger_run_id TEXT,
                source_detail TEXT,
                created_at TEXT
            );
            CREATE TABLE fact_orders_kaspi (
                id INTEGER PRIMARY KEY,
                order_id TEXT,
                store_code TEXT,
                internal_status TEXT,
                kaspi_status TEXT,
                source TEXT,
                status_updated_at TEXT,
                updated_at TEXT,
                imported_at TEXT,
                created_at TEXT
            );
            CREATE TABLE order_status_event (
                event_id INTEGER PRIMARY KEY,
                store_code TEXT,
                order_id TEXT,
                stage_code TEXT,
                event_ts TEXT,
                source TEXT,
                raw_state TEXT,
                raw_status TEXT,
                source_status_change_at TEXT,
                source_run_id TEXT,
                flags_json TEXT,
                source_row_hash TEXT,
                idempotency_key TEXT,
                observed_at TEXT,
                created_at TEXT
            );
            """
        )
        conn.execute(
            """
            INSERT INTO sales_fact_v2 VALUES (
                1, 'O1', '2025-12-31', 'SKU_A', 'SKU_A_M', 'M', 'ACMEWEAR',
                1, 10000, 1000, 2000, 10000, 8000, 'DELIVERED', 0,
                'SALES_KSP_CRM_V3.xlsx', NULL, 'ART', NULL
            )
            """
        )
        conn.execute(
            """
            INSERT INTO fact_order_status_observations VALUES (
                1, 'O1', 'ACMEWEAR', 'COMPLETED', '2026-01-03T10:00:00+05:00',
                'API', 'run-1', '{"status":"ARCHIVE"}',
                '2026-01-03T10:00:01+05:00'
            )
            """
        )


def _workbook_row(*, raw_fee: object = 1000, legacy_fee: object = 9999) -> list[object]:
    return [
        datetime(2025, 12, 31),
        "ACMEWEAR",
        1,
        "O1",
        "M",
        "SKU_A",
        "SKU_A_M",
        10000,
        legacy_fee,
        "O1",
        datetime(2025, 12, 31),
        "ART",
        10000,
        "Завершен",
        "",
        1,
        raw_fee,
        "+77000000000",
        "Private Person",
        "Private Address",
    ]


def _write_workbook(path: Path, rows: list[list[object]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "SALES_KSP_CRM_1"
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()


def _write_base_excluded(path: Path) -> None:
    path.write_text(
        json.dumps({"order_id": "O1", "store_code": "ACMEWEAR", "reasons": ["FORMULA_UNPROVEN"]}, sort_keys=True)
        + "\n",
        encoding="utf-8",
    )


def _build(
    tmp_path: Path,
    *,
    rows: list[list[object]] | None = None,
    db_mutator: Callable[[Path], None] | None = None,
) -> tuple[dict, Path, Path, Path]:
    db = tmp_path / "app.db"
    workbook = tmp_path / "SALES_KSP_CRM_V3.xlsx"
    base = tmp_path / "base_excluded.jsonl"
    out = tmp_path / "out"
    _init_db(db)
    if db_mutator:
        db_mutator(db)
    _write_workbook(workbook, rows or [_workbook_row()])
    _write_base_excluded(base)
    manifest = build_sidecar(
        db_path=db,
        workbook_path=workbook,
        expected_db_sha256=_sha256(db),
        expected_workbook_sha256=_sha256(workbook),
        base_excluded_path=base,
        base_excluded_sha256=_sha256(base),
        since=date(2025, 6, 6),
        until=date(2026, 7, 14),
        sheet="SALES_KSP_CRM_1",
        output_dir=out,
    )
    return manifest, db, workbook, out


def test_build_is_deterministic_valid_and_pii_excluding(tmp_path: Path) -> None:
    manifest, _, _, out = _build(tmp_path)
    proof = json.loads((out / "formula_provenance.jsonl").read_text(encoding="utf-8").strip())
    serialized = "\n".join(
        (out / name).read_text(encoding="utf-8")
        for name in ("manifest.json", "formula_provenance.jsonl", "excluded.jsonl")
    )
    assert manifest["proof_count"] == 1
    assert proof["seller_delivery_fee_total_kzt"] == "1000.00"
    assert "9999" not in proof["seller_delivery_fee_total_kzt"]
    assert "+77000000000" not in serialized
    assert "Private Person" not in serialized
    assert "Private Address" not in serialized
    assert validate_manifest(out / "manifest.json")["ok"] is True


def test_duplicate_exact_workbook_rows_fail_closed(tmp_path: Path) -> None:
    manifest, _, _, out = _build(tmp_path, rows=[_workbook_row(), _workbook_row()])
    excluded = json.loads((out / "excluded.jsonl").read_text(encoding="utf-8").strip())
    assert manifest["proof_count"] == 0
    assert "DUPLICATE_EXACT_WORKBOOK_ROWS" in excluded["reasons"]


def test_missing_raw_fee_never_falls_back_to_legacy_fee(tmp_path: Path) -> None:
    manifest, _, _, out = _build(tmp_path, rows=[_workbook_row(raw_fee=None, legacy_fee=1000)])
    excluded = json.loads((out / "excluded.jsonl").read_text(encoding="utf-8").strip())
    assert manifest["proof_count"] == 0
    assert "NO_EXACT_RAW_FEE_WORKBOOK_ROW" in excluded["reasons"]


def test_raw_fee_must_corroborate_stored_db_fee(tmp_path: Path) -> None:
    manifest, _, _, out = _build(tmp_path, rows=[_workbook_row(raw_fee=173, legacy_fee=173)])
    excluded = json.loads((out / "excluded.jsonl").read_text(encoding="utf-8").strip())
    assert manifest["proof_count"] == 0
    assert "NO_EXACT_RAW_FEE_WORKBOOK_ROW" in excluded["reasons"]


def test_derived_sell_price_must_match_db_unit_price(tmp_path: Path) -> None:
    row = _workbook_row()
    row[7] = 9999
    manifest, _, _, out = _build(tmp_path, rows=[row])
    excluded = json.loads((out / "excluded.jsonl").read_text(encoding="utf-8").strip())
    assert manifest["proof_count"] == 0
    assert "NO_EXACT_RAW_FEE_WORKBOOK_ROW" in excluded["reasons"]


def test_formula_in_raw_fee_cell_is_rejected_even_as_raw_field(tmp_path: Path) -> None:
    row = _workbook_row()
    row[16] = "=1000"
    manifest, _, workbook, out = _build(tmp_path, rows=[row])
    workbook_rows, _ = _read_workbook_rows(workbook, "SALES_KSP_CRM_1")
    excluded = json.loads((out / "excluded.jsonl").read_text(encoding="utf-8").strip())
    assert workbook_rows[0]["raw_formula_headers"] == ["Стоимость доставки для продавца"]
    assert manifest["proof_count"] == 0
    assert "NO_EXACT_RAW_FEE_WORKBOOK_ROW" in excluded["reasons"]


def test_missing_first_party_lifecycle_fails_closed(tmp_path: Path) -> None:
    def remove_lifecycle(db: Path) -> None:
        with sqlite3.connect(db) as conn:
            conn.execute("DELETE FROM fact_order_status_observations")

    manifest, _, _, out = _build(tmp_path, db_mutator=remove_lifecycle)
    excluded = json.loads((out / "excluded.jsonl").read_text(encoding="utf-8").strip())
    assert manifest["proof_count"] == 0
    assert "NO_FIRST_PARTY_TERMINAL_POSITIVE" in excluded["reasons"]


def test_latest_negative_observation_cannot_be_masked_by_positive_fact_order(
    tmp_path: Path,
) -> None:
    def add_conflict(db: Path) -> None:
        with sqlite3.connect(db) as conn:
            conn.execute(
                """
                UPDATE fact_order_status_observations
                SET status_internal='CANCELLED', observed_at='2026-01-04T10:00:00+05:00'
                """
            )
            conn.execute(
                """
                INSERT INTO fact_orders_kaspi (
                    id, order_id, store_code, internal_status, kaspi_status,
                    source, status_updated_at, updated_at, imported_at, created_at
                ) VALUES (
                    1, 'O1', 'ACMEWEAR', 'COMPLETED', 'ARCHIVE', 'API',
                    '2026-01-03T10:00:00+05:00', '2026-01-03T10:00:00+05:00',
                    '2026-01-03T10:00:00+05:00', '2026-01-03T10:00:00+05:00'
                )
                """
            )

    manifest, _, _, out = _build(tmp_path, db_mutator=add_conflict)
    excluded = json.loads((out / "excluded.jsonl").read_text(encoding="utf-8").strip())
    assert manifest["proof_count"] == 0
    assert "TERMINAL_NEGATIVE_CONTRADICTION" in excluded["reasons"]


def test_negative_fact_order_cannot_be_masked_by_positive_observation(tmp_path: Path) -> None:
    def add_conflict(db: Path) -> None:
        with sqlite3.connect(db) as conn:
            conn.execute(
                """
                INSERT INTO fact_orders_kaspi (
                    id, order_id, store_code, internal_status, kaspi_status,
                    source, status_updated_at, updated_at, imported_at, created_at
                ) VALUES (
                    1, 'O1', 'ACMEWEAR', 'CANCELLED', 'CANCELLED', 'API',
                    '2026-01-04T10:00:00+05:00', '2026-01-04T10:00:00+05:00',
                    '2026-01-04T10:00:00+05:00', '2026-01-04T10:00:00+05:00'
                )
                """
            )

    manifest, _, _, out = _build(tmp_path, db_mutator=add_conflict)
    excluded = json.loads((out / "excluded.jsonl").read_text(encoding="utf-8").strip())
    assert manifest["proof_count"] == 0
    assert "TERMINAL_NEGATIVE_CONTRADICTION" in excluded["reasons"]


def test_later_cancelled_event_excludes_candidate(tmp_path: Path) -> None:
    db = tmp_path / "app.db"
    workbook = tmp_path / "SALES_KSP_CRM_V3.xlsx"
    base = tmp_path / "base_excluded.jsonl"
    out = tmp_path / "out"
    _init_db(db)
    with sqlite3.connect(db) as conn:
        conn.execute(
            """
            INSERT INTO order_status_event VALUES (
                1, 'ACMEWEAR', 'O1', 'CANCELLED', '2026-01-04T10:00:00+05:00',
                'api', '', '', '2026-01-04T10:00:00+05:00', 'run-2', '{}',
                ?, 'idem-1', '2026-01-04T10:00:01+05:00',
                '2026-01-04T10:00:01+05:00'
            )
            """,
            ("a" * 64,),
        )
    _write_workbook(workbook, [_workbook_row()])
    _write_base_excluded(base)
    manifest = build_sidecar(
        db_path=db,
        workbook_path=workbook,
        expected_db_sha256=_sha256(db),
        expected_workbook_sha256=_sha256(workbook),
        base_excluded_path=base,
        base_excluded_sha256=_sha256(base),
        since=date(2025, 6, 6),
        until=date(2026, 7, 14),
        sheet="SALES_KSP_CRM_1",
        output_dir=out,
    )
    excluded = json.loads((out / "excluded.jsonl").read_text(encoding="utf-8").strip())
    assert manifest["proof_count"] == 0
    assert "TERMINAL_NEGATIVE_CONTRADICTION" in excluded["reasons"]


def test_workbook_drift_is_rejected(tmp_path: Path) -> None:
    _, _, workbook, out = _build(tmp_path)
    wb = load_workbook(workbook)
    wb["SALES_KSP_CRM_1"]["Q2"] = 1200
    wb.save(workbook)
    wb.close()
    report = validate_manifest(out / "manifest.json")
    assert report["ok"] is False
    assert "SOURCE_DB_RECONSTRUCTION_FAILED" in report["errors"]


def test_numeric_excel_serial_dates_are_epoch_decoded(tmp_path: Path) -> None:
    row = _workbook_row()
    row[0] = to_excel(datetime(2025, 12, 31))
    row[10] = to_excel(datetime(2025, 12, 31))
    manifest, _, _, out = _build(tmp_path, rows=[row])
    proof = json.loads((out / "formula_provenance.jsonl").read_text(encoding="utf-8").strip())
    assert manifest["proof_count"] == 1
    assert proof["source_allowlisted_fields"]["order_date"] == "2025-12-31"


def test_standalone_validator_cli_runs_from_repo_root(tmp_path: Path) -> None:
    _, _, _, out = _build(tmp_path)
    repo_root = Path(__file__).resolve().parents[1]
    completed = subprocess.run(
        [
            sys.executable,
            "scripts/validate_crm_formula_provenance_sidecar.py",
            "--manifest",
            str(out / "manifest.json"),
        ],
        cwd=repo_root,
        check=False,
        capture_output=True,
        text=True,
    )
    assert completed.returncode == 0, completed.stderr
    assert json.loads(completed.stdout)["ok"] is True


def test_array_formula_canonicalization_has_no_object_identity() -> None:
    first = ArrayFormula(ref="M2", text="=SUM(A2:B2)")
    second = ArrayFormula(ref="M2", text="=SUM(A2:B2)")
    assert first is not second
    assert _canonical_cell(first) == _canonical_cell(second)
    assert "object at" not in json.dumps(_canonical_cell(first), sort_keys=True)
    assert _is_formula_cell(first) is True
    assert _is_formula_cell("  =SUM(A2:B2)") is True
    assert _is_formula_cell("literal") is False


def test_crm_copied_db_promotion_is_gated_backup_first_and_target_only(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    manifest, source_db, _, out = _build(tmp_path)
    copied_db = tmp_path / "promotion.db"
    shutil.copy2(source_db, copied_db)
    pre_sha = _sha256(copied_db)

    dry = promote_crm_sidecar(
        manifest_path=out / "manifest.json",
        db_path=copied_db,
        expected_manifest_sha256=manifest["manifest_sha256"],
        expected_db_sha256=pre_sha,
        apply=False,
    )
    assert dry["status"] == "PASS"
    assert _sha256(copied_db) == pre_sha

    monkeypatch.delenv("AB_ALLOW_COPIED_DB_SALES_FORMULA_REPAIR", raising=False)
    with pytest.raises(SidecarApplyError, match="requires AB_ALLOW"):
        promote_crm_sidecar(
            manifest_path=out / "manifest.json",
            db_path=copied_db,
            expected_manifest_sha256=manifest["manifest_sha256"],
            expected_db_sha256=pre_sha,
            apply=True,
            backup_dir=tmp_path / "backups",
        )

    monkeypatch.setenv("AB_ALLOW_COPIED_DB_SALES_FORMULA_REPAIR", "1")
    applied = promote_crm_sidecar(
        manifest_path=out / "manifest.json",
        db_path=copied_db,
        expected_manifest_sha256=manifest["manifest_sha256"],
        expected_db_sha256=pre_sha,
        apply=True,
        backup_dir=tmp_path / "backups",
    )
    assert applied["status"] == "PASS"
    assert applied["target_count"] == 1
    assert applied["backup_sha256"] == pre_sha
    rollback = Path(applied["rollback_artifact_path"])
    assert rollback.exists()
    rollback_payload = json.loads(rollback.read_text(encoding="utf-8"))
    assert rollback_payload["backup_sha256"] == pre_sha
    assert rollback_payload["rollback_command"] == applied["rollback_command"]
    assert applied["non_target_sales_fact_v2_sha256_before"] == applied[
        "non_target_sales_fact_v2_sha256_after"
    ]
    with sqlite3.connect(copied_db) as conn:
        fee, net, cogs, profit = conn.execute(
            "SELECT delivery_fee, net_rev, cogs, profit FROM sales_fact_v2 WHERE sale_id=1"
        ).fetchone()
    assert fee == 1000.0
    assert net != 10000.0
    assert profit == pytest.approx(net - cogs)

    comparison = compare_promotion(
        before_path=source_db,
        after_path=copied_db,
        manifest_path=out / "manifest.json",
        expected_manifest_sha256=manifest["manifest_sha256"],
        mode="applied",
    )
    assert comparison["gate"] == "GREEN"
    assert comparison["changed_target_count"] == 1
    assert comparison["other_table_logical_mismatches"] == []

    with sqlite3.connect(copied_db) as conn:
        conn.execute(
            "UPDATE fact_order_status_observations SET source='WEBUI' WHERE id=1"
        )
    contaminated = compare_promotion(
        before_path=source_db,
        after_path=copied_db,
        manifest_path=out / "manifest.json",
        expected_manifest_sha256=manifest["manifest_sha256"],
        mode="applied",
    )
    assert contaminated["gate"] == "RED"
    assert contaminated["other_table_logical_mismatches"] == [
        "fact_order_status_observations"
    ]
