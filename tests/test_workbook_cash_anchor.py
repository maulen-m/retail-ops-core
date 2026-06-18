import hashlib
import json
import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook

from core.cashflow.workbook_cash_anchor import (
    SHEET_NAME,
    WorkbookCashAnchorError,
    apply_workbook_cash_anchor,
    parse_workbook_cash_snapshot,
)


SNAPSHOT = "13.06.2026_01_01_26"


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_od002_workbook(path: Path, *, reserve: int = 1_500_000) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.cell(4, 1, "Store")
    ws.cell(4, 2, "Account")
    ws.cell(4, 3, "Currency")
    ws.cell(4, 13, SNAPSHOT)
    rows = [
        ("UNIVERSAL", "kaspi_gold", "KZT", 79600),
        ("UNIVERSAL", "kaspi_pay", "KZT", 81512),
        ("UNIVERSAL", "bcc", "KZT", 4000),
        ("UNIVERSAL", "freedom", "KZT", 27000),
        ("UNIVERSAL", "cash_kzt", "KZT", 570000),
        ("UNIVERSAL", "cash_usd", "USD", 387),
        ("UNIVERSAL", "cash_rub", "RUB", 10000),
        ("UNIVERSAL", "binance_usdt", "USDT", 3134),
        ("11KZ", "kaspi_gold", "KZT", None),
        ("11KZ", "kaspi_pay", "KZT", None),
        ("11KZ", "binance_usdt", "USDT", None),
        ("STOREB", "kaspi_gold", "KZT", 70000),
        ("STOREB", "kaspi_pay", "KZT", 119644),
        ("ACMEWEAR", "kaspi_gold", "KZT", 950000),
        ("ACMEWEAR", "kaspi_pay", "KZT", 267923),
        ("ACMEWEAR", "cash_kzt", "KZT", None),
        ("MELVIS", "kaspi_gold", "KZT", None),
        ("MELVIS", "kaspi_pay", "KZT", None),
        ("Reserve", "cash_kzt", "KZT", reserve),
    ]
    for row_idx, row in enumerate(rows, start=5):
        for col_idx, value in enumerate(row[:3], start=1):
            ws.cell(row_idx, col_idx, value)
        ws.cell(row_idx, 13, row[3])

    ws.cell(25, 1, "STORE TOTALS (KZT only)")
    for row_idx, row in enumerate(
        [
            ("UNIVERSAL", 762112),
            ("11KZ", 0),
            ("STOREB", 189644),
            ("ACMEWEAR", 1217923),
            ("MELVIS", 0),
        ],
        start=26,
    ):
        ws.cell(row_idx, 1, row[0])
        ws.cell(row_idx, 13, row[1])

    ws.cell(32, 1, "CURRENCY TOTALS")
    for row_idx, row in enumerate(
        [
            ("KZT", 2169679),
            ("RUB", 60000),
            ("USD", 187695),
            ("USDT", 1519990),
        ],
        start=33,
    ):
        ws.cell(row_idx, 1, row[0])
        ws.cell(row_idx, 13, row[1])
    ws.cell(38, 1, "GRAND TOTAL KZT")
    ws.cell(38, 13, 3937364)
    ws.cell(39, 1, "GRAND TOTAL +reserve KZT")
    ws.cell(39, 13, 5437364 if reserve == 1_500_000 else 3937364 + reserve)
    wb.save(path)


def _init_db(db_path: Path) -> None:
    conn = sqlite3.connect(str(db_path))
    try:
        conn.execute(
            """
            CREATE TABLE fact_cashflow_daily (
                date TEXT PRIMARY KEY,
                cash_open REAL NOT NULL DEFAULT 0,
                cash_close REAL NOT NULL DEFAULT 0,
                receivables_open REAL NOT NULL DEFAULT 0,
                receivables_close REAL NOT NULL DEFAULT 0,
                inventory_cost_open REAL NOT NULL DEFAULT 0,
                inventory_cost_close REAL NOT NULL DEFAULT 0,
                capital_close REAL NOT NULL DEFAULT 0,
                sales_accrued_kzt REAL NOT NULL DEFAULT 0,
                payouts_received_kzt REAL NOT NULL DEFAULT 0,
                refunds_kzt REAL NOT NULL DEFAULT 0,
                po_payments_kzt REAL NOT NULL DEFAULT 0,
                expenses_kzt REAL NOT NULL DEFAULT 0,
                cogs_kzt REAL NOT NULL DEFAULT 0,
                cash_flow_kzt REAL NOT NULL DEFAULT 0,
                receivables_flow_kzt REAL NOT NULL DEFAULT 0,
                inventory_cost_flow_kzt REAL NOT NULL DEFAULT 0,
                profit_accrual_kzt REAL NOT NULL DEFAULT 0,
                run_id TEXT,
                created_at TEXT DEFAULT (datetime('now'))
            )
            """
        )
        conn.execute(
            """
            INSERT INTO fact_cashflow_daily (date, cash_open, cash_close, capital_close, run_id)
            VALUES ('2026-06-13', 10, 20, 20, 'before')
            """
        )
        conn.commit()
    finally:
        conn.close()


def test_parse_od002_snapshot_preserves_operating_reserve_and_fx(tmp_path: Path) -> None:
    workbook = tmp_path / "cash.xlsx"
    _write_od002_workbook(workbook)

    snapshot = parse_workbook_cash_snapshot(
        workbook_path=workbook,
        sheet_name=SHEET_NAME,
        snapshot_label=SNAPSHOT,
    )

    assert snapshot["operating_total_kzt"] == 3_937_364
    assert snapshot["reserve_context"]["reserve_kzt"] == 1_500_000
    assert snapshot["grand_total_with_reserve_kzt"] == 5_437_364
    assert snapshot["store_totals_kzt_only"]["UNIVERSAL"] == 762_112
    assert snapshot["store_totals_kzt_equiv"]["UNIVERSAL"] == 2_529_797
    assert snapshot["currency_totals"]["USDT"] == 1_519_990
    assert snapshot["operating_account_row_count"] == 18
    assert snapshot["reserve_account_row_count"] == 1
    assert snapshot["zero_balance_operating_account_rows"] == 6


def test_reserve_mismatch_fails_closed(tmp_path: Path) -> None:
    workbook = tmp_path / "cash.xlsx"
    _write_od002_workbook(workbook, reserve=1_400_000)

    with pytest.raises(WorkbookCashAnchorError, match="expected reserve"):
        parse_workbook_cash_snapshot(
            workbook_path=workbook,
            sheet_name=SHEET_NAME,
            snapshot_label=SNAPSHOT,
        )


def test_dry_run_writes_evidence_without_db_write(tmp_path: Path) -> None:
    workbook = tmp_path / "cash.xlsx"
    db_path = tmp_path / "app.db"
    output = tmp_path / "out"
    _write_od002_workbook(workbook)
    _init_db(db_path)
    before = _sha256(db_path)

    summary = apply_workbook_cash_anchor(
        db_path=db_path,
        workbook_path=workbook,
        sheet_name=SHEET_NAME,
        snapshot_label=SNAPSHOT,
        output_root=output,
        run_id="test-run",
        apply=False,
    )

    assert summary["apply"]["applied"] is False
    assert summary["apply"]["would_insert_anchor_records"] == 18
    assert _sha256(db_path) == before
    assert (output / "summary.json").exists()
    balance_csv = (output / "cash_balance_check_od002_operating.csv").read_text(encoding="utf-8")
    assert "2026-06-13,3937364" in balance_csv


def test_apply_requires_env_gate_before_schema_write(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    workbook = tmp_path / "cash.xlsx"
    db_path = tmp_path / "app.db"
    _write_od002_workbook(workbook)
    _init_db(db_path)
    before = _sha256(db_path)
    monkeypatch.delenv("ENABLE_CASHFLOW_ANCHOR_WRITE", raising=False)

    with pytest.raises(RuntimeError, match="ENABLE_CASHFLOW_ANCHOR_WRITE=1"):
        apply_workbook_cash_anchor(
            db_path=db_path,
            workbook_path=workbook,
            sheet_name=SHEET_NAME,
            snapshot_label=SNAPSHOT,
            output_root=tmp_path / "out",
            run_id="test-run",
            apply=True,
        )

    assert _sha256(db_path) == before


def test_apply_inserts_exact_operating_account_anchors_and_keeps_reserve_out(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = tmp_path / "cash.xlsx"
    db_path = tmp_path / "app.db"
    backup = tmp_path / "backup.db"
    _write_od002_workbook(workbook)
    _init_db(db_path)
    backup.write_bytes(db_path.read_bytes())
    monkeypatch.setenv("ENABLE_CASHFLOW_ANCHOR_WRITE", "1")

    summary = apply_workbook_cash_anchor(
        db_path=db_path,
        workbook_path=workbook,
        sheet_name=SHEET_NAME,
        snapshot_label=SNAPSHOT,
        output_root=tmp_path / "out",
        run_id="test-run",
        backup_path=backup,
        apply=True,
    )

    assert summary["apply"]["inserted_anchor_records"] == 18
    with sqlite3.connect(db_path) as conn:
        count, total = conn.execute(
            """
            SELECT COUNT(*), ROUND(SUM(anchor_closing_balance_kzt), 2)
            FROM cashflow_cash_anchor
            WHERE anchor_date='2026-06-13'
            """
        ).fetchone()
        reserve_rows = conn.execute(
            "SELECT COUNT(*) FROM cashflow_cash_anchor WHERE source_store_name LIKE 'RESERVE:%'"
        ).fetchone()[0]
        daily = conn.execute(
            "SELECT cash_open, cash_close, run_id FROM fact_cashflow_daily WHERE date='2026-06-13'"
        ).fetchone()

    assert count == 18
    assert total == 3_937_364
    assert reserve_rows == 0
    assert daily == (10, 20, "before")
    saved = json.loads((tmp_path / "out" / "summary.json").read_text(encoding="utf-8"))
    assert saved["rollback"]["backup_path"] == str(backup)
