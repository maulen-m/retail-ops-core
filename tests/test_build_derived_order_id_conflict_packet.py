from __future__ import annotations

import hashlib
import json
from pathlib import Path
import sqlite3

from openpyxl import Workbook, load_workbook
import pytest

from scripts.build_derived_order_id_conflict_packet import PacketError, build_packet


SAFE_HEADERS = [
    "Date",
    "STORE_NAME",
    "Quantity",
    "Kaspi_name_core",
    "OrderID",
    "MY_SIZE",
    "SKU_key",
    "SKU_ID",
    "PLANNED_SHIPPING_DATE",
    "SKU_ID_KSP",
    "Kaspi_name_source",
    "№ заказа",
    "Дата поступления заказа",
    "Название товара в Kaspi Магазине",
    "Название в системе продавца",
    "Артикул",
    "Сумма",
    "Дата изменения статуса",
    "Статус",
    "Причина отмены",
    "Количество",
    "Стоимость доставки для продавца",
    "Плановая дата передачи курьеру",
    "Склад передачи КД",
]


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _make_workbook(path: Path, *, duplicate: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SALES_KSP_CRM_1"
    sheet.append(SAFE_HEADERS)
    values = {
        "Date": "2025-08-12",
        "STORE_NAME": "AcmeWear",
        "Quantity": 1,
        "Kaspi_name_core": "Leggings black",
        "OrderID": 610999345,
        "MY_SIZE": "3XL",
        "SKU_key": "CL_NEW-CLO_MEN_LEG_BLACK",
        "SKU_ID": "CL_NEW-CLO_MEN_LEG_BLACK_3XL",
        "PLANNED_SHIPPING_DATE": "12.08.2025",
        "SKU_ID_KSP": "CL_NEW-CLO_MEN_LEG_BLACK_3XL",
        "Kaspi_name_source": "Leggings 3XL",
        "№ заказа": 610619543,
        "Дата поступления заказа": "11.08.2025",
        "Название товара в Kaspi Магазине": "Leggings black 3XL",
        "Название в системе продавца": "Leggings 3XL",
        "Артикул": "CL_NEW-CLO_MEN_LEG_BLACK_3XL",
        "Сумма": 1398,
        "Дата изменения статуса": "11.08.2025",
        "Статус": "Ожидает передачи курьеру",
        "Причина отмены": None,
        "Количество": 1,
        "Стоимость доставки для продавца": 0,
        "Плановая дата передачи курьеру": "12.08.2025",
        "Склад передачи КД": "30137883_PP1",
    }
    sheet.append([values[name] for name in SAFE_HEADERS])
    if duplicate:
        sheet.append([values[name] for name in SAFE_HEADERS])
    workbook.save(path)
    workbook.close()


def _make_db(path: Path) -> None:
    with sqlite3.connect(path) as conn:
        conn.executescript(
            """
            CREATE TABLE fact_orders_kaspi (
                id INTEGER PRIMARY KEY, order_id TEXT, store_code TEXT,
                channel_code TEXT, kaspi_offer_name TEXT, sku_key TEXT,
                sku_id TEXT, my_size TEXT, assigned_size TEXT, quantity REAL,
                unit_price_kzt REAL, planned_shipment_date TEXT,
                actual_shipment_date TEXT, kaspi_status TEXT,
                internal_status TEXT, status_updated_at TEXT, source TEXT,
                source_file TEXT, line_identity_key TEXT, kaspi_article TEXT
            );
            CREATE TABLE fact_order_status_observations (
                id INTEGER PRIMARY KEY, order_id TEXT, store_code TEXT,
                status_internal TEXT, observed_at TEXT, source TEXT,
                ledger_run_id TEXT, source_detail TEXT
            );
            CREATE TABLE order_status_event (
                event_id INTEGER PRIMARY KEY, store_code TEXT, order_id TEXT,
                stage_code TEXT, event_ts TEXT, source TEXT, raw_state TEXT,
                raw_status TEXT, source_status_change_at TEXT,
                source_run_id TEXT, source_row_hash TEXT, idempotency_key TEXT
            );
            CREATE TABLE fact_sales_workbook_anchor (
                order_id TEXT, store_code TEXT, sale_date TEXT, quantity REAL,
                net_rev_kzt REAL, total_price_kzt REAL, source_file TEXT,
                updated_at TEXT
            );
            INSERT INTO fact_orders_kaspi VALUES
                (1, '610999345', '30137883_PP1', 'KSP', NULL, NULL, NULL,
                 NULL, NULL, 1, NULL, '2025-08-12', NULL, NULL, 'NEW', NULL,
                 'CRM_MANUAL', 'SALES_KSP_CRM_V3.xlsx', 'legacy-row:1', NULL),
                (2, '610619543', 'ACMEWEAR', 'KASPI', 'Leggings black 3XL',
                 'CL_NEW-CLO_MEN_LEG_BLACK', 'CL_NEW-CLO_MEN_LEG_BLACK_3XL',
                 '3XL', NULL, 1, 1398, '2025-08-12', NULL, 'Выдан', 'NEW',
                 '2025-08-13', 'kaspi_archive', 'archive.xlsx', 'raw-line', NULL);
            INSERT INTO fact_order_status_observations VALUES
                (1, '610619543', 'ACMEWEAR', 'DELIVERED', '2025-08-13',
                 'WEBUI', 'run', 'webui_status_ledger');
            INSERT INTO order_status_event VALUES
                (1, 'ACMEWEAR', '610619543', 'COMPLETED', '2025-08-13',
                 'LOCAL_FACT_ORDER_STATUS_OBSERVATIONS', NULL, 'DELIVERED',
                 '2025-08-13', 'run', 'rowhash', 'idempotency');
            INSERT INTO fact_sales_workbook_anchor VALUES
                ('610999345', 'ACMEWEAR', '2025-08-12', 1, 1186.55, 1398,
                 'SALES_KSP_CRM_V3.xlsx', '2026-03-07');
            """
        )


def _build(tmp_path: Path, *, duplicate: bool = False):
    db = tmp_path / "app.db"
    _make_db(db)
    workbooks = []
    for index in range(2):
        workbook = tmp_path / f"snapshot-{index}.xlsx"
        _make_workbook(workbook, duplicate=duplicate)
        workbooks.append((workbook, _sha256(workbook)))
    db_before = db.read_bytes()
    workbook_before = [path.read_bytes() for path, _ in workbooks]
    packet = build_packet(
        db_path=db,
        expected_db_sha256=_sha256(db),
        workbooks=workbooks,
        derived_order_id="610999345",
        raw_order_id="610619543",
        expected_store="ACMEWEAR",
        sheet_name="SALES_KSP_CRM_1",
        expected_row_number=2,
        output_dir=tmp_path / "packet",
    )
    return db, workbooks, db_before, workbook_before, packet


def test_builds_read_only_conflict_quarantine_packet(tmp_path: Path) -> None:
    db, workbooks, db_before, workbook_before, packet = _build(tmp_path)
    assert db.read_bytes() == db_before
    assert [path.read_bytes() for path, _ in workbooks] == workbook_before
    assert packet["verdict"] == "DERIVED_ORDER_ID_CONFLICT_QUARANTINE"
    assert packet["schema"] == "derived_order_id_conflict_packet_v2"
    assert packet["proofs"] == {
        "derived_and_raw_ids_differ": True,
        "all_pinned_workbooks_match_one_safe_row": True,
        "derived_anchor_exists": True,
        "raw_anchor_absent": True,
        "derived_terminal_evidence_absent": True,
        "raw_terminal_evidence_present": True,
        "raw_product_identity_matches_workbook": True,
        "production_write_authorized": False,
    }
    assert packet["db"]["sha256_before"] == packet["db"]["sha256_after"]
    evidence = json.loads(Path(packet["evidence_path"]).read_text())
    claimed = evidence.pop("evidence_sha256")
    recomputed = hashlib.sha256(
        json.dumps(
            evidence,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            default=str,
        ).encode("utf-8")
    ).hexdigest()
    assert claimed == recomputed
    assert Path(packet["manifest_path"]).is_file()
    assert "Gate: GREEN" in Path(packet["closeout_path"]).read_text()


def test_rejects_equal_derived_and_raw_id(tmp_path: Path) -> None:
    db = tmp_path / "app.db"
    _make_db(db)
    workbook = tmp_path / "snapshot.xlsx"
    _make_workbook(workbook)
    with pytest.raises(PacketError, match="must differ"):
        build_packet(
            db_path=db,
            expected_db_sha256=_sha256(db),
            workbooks=[(workbook, _sha256(workbook)), (workbook, _sha256(workbook))],
            derived_order_id="610619543",
            raw_order_id="610619543",
            expected_store="ACMEWEAR",
            sheet_name="SALES_KSP_CRM_1",
            expected_row_number=2,
            output_dir=tmp_path / "packet",
        )


def test_rejects_ambiguous_workbook_match(tmp_path: Path) -> None:
    db = tmp_path / "app.db"
    _make_db(db)
    workbooks = []
    for index in range(2):
        workbook = tmp_path / f"snapshot-{index}.xlsx"
        _make_workbook(workbook, duplicate=True)
        workbooks.append((workbook, _sha256(workbook)))
    with pytest.raises(PacketError, match="found 2"):
        build_packet(
            db_path=db,
            expected_db_sha256=_sha256(db),
            workbooks=workbooks,
            derived_order_id="610999345",
            raw_order_id="610619543",
            expected_store="ACMEWEAR",
            sheet_name="SALES_KSP_CRM_1",
            expected_row_number=2,
            output_dir=tmp_path / "packet",
        )


def test_generated_workbook_fixture_remains_readable(tmp_path: Path) -> None:
    workbook_path = tmp_path / "snapshot.xlsx"
    _make_workbook(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        assert workbook["SALES_KSP_CRM_1"]["E2"].value == 610999345
    finally:
        workbook.close()
