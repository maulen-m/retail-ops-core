from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.integrations.kaspi_api_client import APIResponse, KaspiAPIClient, KaspiAPIError, KaspiAuthError
from scripts import export_api_orders as mod
from scripts import validate_activeorders_columns as validator_mod
from scripts.export_api_orders import filter_rows_by_planned_date


def _delivery_order(
    code: str,
    *,
    assembled: bool = False,
    courier_transmission_date=None,
    creation_date=None,
    courier_transmission_planning_date=1772920800000,
):
    attrs = {
        "code": code,
        "state": "KASPI_DELIVERY",
        "status": "ACCEPTED_BY_MERCHANT",
        "assembled": assembled,
        "creationDate": creation_date,
        "kaspiDelivery": {
            "courierTransmissionPlanningDate": courier_transmission_planning_date,
        },
        "customer": {},
    }
    if courier_transmission_date is not None:
        attrs["courierTransmissionDate"] = courier_transmission_date
    return {"id": f"base64-{code}", "attributes": attrs}


def test_filter_rows_by_planned_date_exact():
    rows = [
        {"Плановая дата передачи курьеру": "01.01.2026"},
        {"Плановая дата передачи курьеру": "03.01.2026"},
    ]

    filtered = filter_rows_by_planned_date(
        rows, target_date="03.01.2026", include_overdue=False
    )

    assert len(filtered) == 1
    assert filtered[0]["Плановая дата передачи курьеру"] == "03.01.2026"


def test_filter_rows_by_planned_date_include_overdue():
    rows = [
        {"Плановая дата передачи курьеру": "01.01.2026"},
        {"Плановая дата передачи курьеру": "03.01.2026"},
        {"Плановая дата передачи курьеру": "05.01.2026"},
    ]

    filtered = filter_rows_by_planned_date(
        rows, target_date="03.01.2026", include_overdue=True
    )

    planned = [row["Плановая дата передачи курьеру"] for row in filtered]
    assert planned == ["01.01.2026", "03.01.2026"]


def test_write_excel_materializes_canonical_header_only_workbook(
    tmp_path: Path, monkeypatch
) -> None:
    output_path = tmp_path / "ActiveOrders.xlsx"

    assert mod.write_excel([], output_path) == 0

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    assert headers == mod.EXCEL_COLUMNS
    assert worksheet.max_row == 1
    monkeypatch.setattr(
        validator_mod.sys,
        "argv",
        ["validate_activeorders_columns.py", str(output_path)],
    )
    assert validator_mod.main() == 0


def test_main_writes_header_only_workbook_after_complete_zero_order_refresh(
    tmp_path: Path, monkeypatch
) -> None:
    output_path = tmp_path / "ActiveOrders.xlsx"
    captured: dict[str, object] = {}

    def fake_export_all_stores(**kwargs):
        captured.update(kwargs)
        return []

    monkeypatch.setattr(mod, "export_all_stores", fake_export_all_stores)
    monkeypatch.setattr(
        mod.sys,
        "argv",
        [
            "export_api_orders.py",
            "--all-stores",
            "--require-complete",
            "--no-archive",
            "--output",
            str(output_path),
        ],
    )

    mod.main()

    assert captured["require_complete"] is True
    workbook = load_workbook(output_path, read_only=True, data_only=True)
    assert [cell.value for cell in workbook.active[1]] == mod.EXCEL_COLUMNS
    assert workbook.active.max_row == 1


def test_main_preserves_existing_workbook_when_zero_export_is_not_complete(
    tmp_path: Path, monkeypatch
) -> None:
    output_path = tmp_path / "ActiveOrders.xlsx"
    output_path.write_bytes(b"existing-source")
    monkeypatch.setattr(mod, "export_all_stores", lambda **_kwargs: [])
    monkeypatch.setattr(
        mod.sys,
        "argv",
        [
            "export_api_orders.py",
            "--all-stores",
            "--no-archive",
            "--output",
            str(output_path),
        ],
    )

    mod.main()

    assert output_path.read_bytes() == b"existing-source"


def test_require_complete_rejects_single_store_before_export_or_write(
    tmp_path: Path, monkeypatch
) -> None:
    output_path = tmp_path / "ActiveOrders.xlsx"
    output_path.write_bytes(b"existing-source")

    def fail_if_called(**_kwargs):
        raise AssertionError("single-store export must not run in authoritative mode")

    monkeypatch.setattr(mod, "export_store_orders", fail_if_called)
    monkeypatch.setattr(
        mod.sys,
        "argv",
        [
            "export_api_orders.py",
            "--store",
            "UNIVERSAL",
            "--require-complete",
            "--output",
            str(output_path),
        ],
    )

    with pytest.raises(SystemExit) as exc_info:
        mod.main()

    assert exc_info.value.code == 2
    assert output_path.read_bytes() == b"existing-source"


def test_complete_export_propagates_missing_store_auth(monkeypatch) -> None:
    class MissingAuthClient:
        def __init__(self, store_code: str):
            raise KaspiAuthError(f"missing token for {store_code}")

    monkeypatch.setattr(mod, "KaspiAPIClient", MissingAuthClient)

    with pytest.raises(KaspiAuthError, match="missing token"):
        mod.export_store_orders(
            store_code="UNIVERSAL",
            state="KASPI_DELIVERY",
            include_archive=False,
            require_complete=True,
        )


def test_complete_export_rejects_enabled_store_without_token_mapping(monkeypatch) -> None:
    monkeypatch.setattr(mod, "load_sync_enabled_kaspi_store_codes", lambda: ["UNKNOWN_STORE"])

    with pytest.raises(RuntimeError, match="missing API token mappings: UNKNOWN_STORE"):
        mod.export_all_stores(require_complete=True)


def test_list_all_orders_can_fail_closed_on_incomplete_pagination(monkeypatch) -> None:
    client = object.__new__(KaspiAPIClient)
    monkeypatch.setattr(
        client,
        "list_orders",
        lambda **_kwargs: APIResponse(success=False, error="temporary upstream failure"),
    )

    with pytest.raises(KaspiAPIError, match="Failed to fetch page 0"):
        client.list_all_orders(state="KASPI_DELIVERY", raise_on_error=True)


def test_list_all_orders_can_fail_closed_on_pagination_safety_limit(monkeypatch) -> None:
    client = object.__new__(KaspiAPIClient)
    monkeypatch.setattr(
        client,
        "list_orders",
        lambda **_kwargs: APIResponse(
            success=True,
            data={"data": [{}], "meta": {"pageCount": 2, "totalCount": 2}},
        ),
    )

    with pytest.raises(KaspiAPIError, match="Advertised pageCount 2 exceeds safety limit 1"):
        client.list_all_orders(state="KASPI_DELIVERY", max_pages=1, raise_on_error=True)


def test_list_all_orders_complete_mode_fetches_all_advertised_short_pages(
    monkeypatch,
) -> None:
    client = object.__new__(KaspiAPIClient)
    requested_pages: list[int] = []

    def fake_list_orders(**kwargs):
        page_number = int(kwargs["page_number"])
        requested_pages.append(page_number)
        return APIResponse(
            success=True,
            data={
                "data": [{"id": f"order-{page_number}"}],
                "meta": {"pageCount": 2, "totalCount": 2},
            },
        )

    monkeypatch.setattr(client, "list_orders", fake_list_orders)

    rows = client.list_all_orders(state="KASPI_DELIVERY", raise_on_error=True)

    assert requested_pages == [0, 1]
    assert [row["id"] for row in rows] == ["order-0", "order-1"]


def test_list_all_orders_complete_mode_rejects_total_count_mismatch(monkeypatch) -> None:
    client = object.__new__(KaspiAPIClient)
    requested_pages: list[int] = []

    def fake_list_orders(**kwargs):
        page_number = int(kwargs["page_number"])
        requested_pages.append(page_number)
        return APIResponse(
            success=True,
            data={
                "data": [{"id": f"order-{page_number}"}],
                "meta": {"pageCount": 2, "totalCount": 3},
            },
        )

    monkeypatch.setattr(client, "list_orders", fake_list_orders)

    with pytest.raises(KaspiAPIError, match="total-count mismatch"):
        client.list_all_orders(state="KASPI_DELIVERY", raise_on_error=True)

    assert requested_pages == [0, 1]


def test_list_all_orders_complete_mode_rejects_changed_page_meta(monkeypatch) -> None:
    client = object.__new__(KaspiAPIClient)

    def fake_list_orders(**kwargs):
        page_number = int(kwargs["page_number"])
        total_count = 2 if page_number == 0 else 3
        return APIResponse(
            success=True,
            data={
                "data": [{"id": f"order-{page_number}"}],
                "meta": {"pageCount": 2, "totalCount": total_count},
            },
        )

    monkeypatch.setattr(client, "list_orders", fake_list_orders)

    with pytest.raises(KaspiAPIError, match="Pagination meta changed on page 1"):
        client.list_all_orders(state="KASPI_DELIVERY", raise_on_error=True)


def test_list_all_orders_complete_mode_rejects_empty_advertised_page(monkeypatch) -> None:
    client = object.__new__(KaspiAPIClient)

    def fake_list_orders(**kwargs):
        page_number = int(kwargs["page_number"])
        page_rows = [{"id": "order-0"}] if page_number == 0 else []
        return APIResponse(
            success=True,
            data={
                "data": page_rows,
                "meta": {"pageCount": 2, "totalCount": 2},
            },
        )

    monkeypatch.setattr(client, "list_orders", fake_list_orders)

    with pytest.raises(KaspiAPIError, match="advertised page 1 is empty"):
        client.list_all_orders(state="KASPI_DELIVERY", raise_on_error=True)


def test_list_all_orders_complete_mode_accepts_advertised_zero(monkeypatch) -> None:
    client = object.__new__(KaspiAPIClient)
    requested_pages: list[int] = []

    def fake_list_orders(**kwargs):
        requested_pages.append(int(kwargs["page_number"]))
        return APIResponse(
            success=True,
            data={"data": [], "meta": {"pageCount": 0, "totalCount": 0}},
        )

    monkeypatch.setattr(client, "list_orders", fake_list_orders)

    assert client.list_all_orders(state="KASPI_DELIVERY", raise_on_error=True) == []
    assert requested_pages == [0]


def test_list_all_orders_complete_mode_accepts_single_empty_page_convention(
    monkeypatch,
) -> None:
    client = object.__new__(KaspiAPIClient)
    requested_pages: list[int] = []

    def fake_list_orders(**kwargs):
        requested_pages.append(int(kwargs["page_number"]))
        return APIResponse(
            success=True,
            data={"data": [], "meta": {"pageCount": 1, "totalCount": 0}},
        )

    monkeypatch.setattr(client, "list_orders", fake_list_orders)

    assert client.list_all_orders(state="KASPI_DELIVERY", raise_on_error=True) == []
    assert requested_pages == [0]


def test_list_all_orders_complete_mode_rejects_nonempty_zero_total_convention(
    monkeypatch,
) -> None:
    client = object.__new__(KaspiAPIClient)
    monkeypatch.setattr(
        client,
        "list_orders",
        lambda **_kwargs: APIResponse(
            success=True,
            data={
                "data": [{"id": "unexpected-order"}],
                "meta": {"pageCount": 1, "totalCount": 0},
            },
        ),
    )

    with pytest.raises(KaspiAPIError, match="total-count mismatch"):
        client.list_all_orders(state="KASPI_DELIVERY", raise_on_error=True)


def test_export_store_orders_excludes_in_delivery_orders_from_pending_export(monkeypatch):
    class FakeClient:
        def __init__(self, store_code: str):
            self.store_code = store_code

        def list_all_orders(self, **_kwargs):
            return [
                _delivery_order("PENDING", assembled=False),
                _delivery_order("SHIPPED", assembled=True, courier_transmission_date=1772973600000),
            ]

        def get_order_entries_by_id(self, _order_id: str):
            return type("Resp", (), {"success": True, "data": {"data": []}})()

    monkeypatch.setattr(mod, "KaspiAPIClient", FakeClient)
    monkeypatch.setattr(
        mod,
        "order_to_rows",
        lambda order, entries, store_code, client=None: [
            {"№ заказа": order.get("attributes", {}).get("code", ""), "store": store_code, "entries": len(entries)}
        ],
    )

    rows = mod.export_store_orders(
        store_code="UNIVERSAL",
        state="KASPI_DELIVERY",
        days=3,
        include_archive=False,
    )

    assert [row["№ заказа"] for row in rows] == ["PENDING"]


def test_order_to_rows_keeps_pending_delivery_rows_unissued():
    rows = mod.order_to_rows(
        _delivery_order("PENDING", assembled=True),
        entries=[{"attributes": {"offer": {"name": "Item", "merchantProductId": "SKU1"}, "quantity": 1, "basePrice": 1000}}],
        store_code="UNIVERSAL",
    )

    assert len(rows) == 1
    assert rows[0]["Статус"] == "Ожидает передачи курьеру"
    assert rows[0]["Выдал"] == ""


def test_order_to_rows_uses_raw_courier_planning_date_for_next_day_orders():
    rows = mod.order_to_rows(
        _delivery_order(
            "NEXTDAY",
            creation_date=1772877981000,  # 2026-03-07 15:06:21 +05:00
            courier_transmission_planning_date=1772982000000,  # 2026-03-08 20:00:00 +05:00
        ),
        entries=[
            {
                "attributes": {
                    "offer": {"name": "Item", "merchantProductId": "SKU1"},
                    "quantity": 1,
                    "basePrice": 1000,
                }
            }
        ],
        store_code="UNIVERSAL",
    )

    assert len(rows) == 1
    assert rows[0]["Плановая дата передачи курьеру"] == "08.03.2026"


def test_order_to_rows_passes_store_code_into_planned_date_resolution(monkeypatch):
    captured = {}

    def _fake_planned_date_from_order(order, *, store_code=None):
        captured["store_code"] = store_code
        return date(2026, 3, 14)

    monkeypatch.setattr(mod, "planned_date_from_order", _fake_planned_date_from_order)

    rows = mod.order_to_rows(
        _delivery_order("ACMEWEAR-CUTOFF", courier_transmission_planning_date=None),
        entries=[],
        store_code="ACMEWEAR",
    )

    assert rows[0]["Плановая дата передачи курьеру"] == "14.03.2026"
    assert captured["store_code"] == "ACMEWEAR"
