import csv
import hashlib
import json
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from scripts.build_track_e_owner_count_workbook import (
    AUTHORIZED_LANES,
    build_rows,
    generate_bundle,
    resolve_estimate,
)


AS_OF = "2026-08-10T23:59:59+05:00"


def test_zero_and_blank_owner_corrections_preserve_estimate_semantics():
    assert resolve_estimate(17, 17, 0) == (0, 0)
    assert resolve_estimate(81, 81, None) == (81, 81)
    assert resolve_estimate(522, 716, "") == (522, 716)


def test_only_authorized_lanes_and_beli_4xl_zero_are_present():
    rows = build_rows(AS_OF)
    assert {row["candidate_lane"] for row in rows} == set(AUTHORIZED_LANES)
    beli_4xl = next(
        row for row in rows if row["family"] == "LINE51" and row["size"] == "4XL"
    )
    assert beli_4xl["estimate_low"] == 0
    assert beli_4xl["estimate_high"] == 0
    assert beli_4xl["blank_semantic"] == "EXPLICIT_ZERO_ESTIMATE"


def test_unresolved_evidence_never_approves_quantity_or_economics():
    for row in build_rows(AS_OF):
        assert row["approved_lot_qty"] == ""
        assert row["frozen_landed_cogs"] == ""
        assert row["opportunity_cost"] == ""
        assert row["outreach_status"] == "UNKNOWN/STOP"


def test_workbook_formulas_validations_and_format_never_authorize_outreach(tmp_path):
    bundle = generate_bundle(tmp_path, AS_OF, strict=True)
    workbook = openpyxl.load_workbook(bundle["workbook"], data_only=False)
    assert workbook.sheetnames == [
        "README",
        "COUNT_ENTRY",
        "REPLAY_AND_GAPS",
        "COGS_AND_FLOORS",
        "LOT_PREVIEW",
    ]
    assert workbook.calculation.fullCalcOnLoad is True
    assert not workbook._external_links
    assert workbook.vba_archive is None
    for sheet in workbook.worksheets:
        assert sheet.freeze_panes is not None
        assert any(dimension.width for dimension in sheet.column_dimensions.values())

    count_entry = workbook["COUNT_ENTRY"]
    assert count_entry.data_validations.count >= 4
    headers = {cell.value: cell.column for cell in count_entry[1]}
    outreach_col = headers["outreach_status"]
    approved_col = headers["approved_lot_qty"]
    for row_number in range(2, count_entry.max_row + 1):
        assert count_entry.cell(row_number, approved_col).data_type == "f"
        formula = count_entry.cell(row_number, outreach_col).value
        assert formula.startswith("=")
        assert '"UNKNOWN/STOP"' in formula
        assert '"AUTHORIZED"' not in formula


def test_approved_quantity_requires_exact_count_cogs_proof_and_price_order(tmp_path):
    bundle = generate_bundle(tmp_path, AS_OF, strict=True)
    workbook = openpyxl.load_workbook(bundle["workbook"], data_only=False)
    sheet = workbook["COUNT_ENTRY"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    approved_formula = sheet.cell(2, headers["approved_lot_qty"]).value

    required_inputs = (
        "physical_count",
        "count_timestamp_+05",
        "counter",
        "photo_hash",
        "fx",
        "cogs_source",
    )
    for name in required_inputs:
        coordinate = f'{get_column_letter(headers[name])}2'
        assert coordinate in approved_formula

    count_low = f'{get_column_letter(headers["count_low"])}2'
    count_high = f'{get_column_letter(headers["count_high"])}2'
    physical_count = f'{get_column_letter(headers["physical_count"])}2'
    opening_price = f'{get_column_letter(headers["opening_price"])}2'
    target_price = f'{get_column_letter(headers["target_price"])}2'
    hard_floor = f'{get_column_letter(headers["hard_floor"])}2'
    assert f"{count_low}<>{count_high}" in approved_formula
    assert f"{physical_count}<>{count_low}" in approved_formula
    assert f"{opening_price}<{target_price}" in approved_formula
    assert f"{target_price}<{hard_floor}" in approved_formula

    validated_price_columns = set()
    for validation in sheet.data_validations.dataValidation:
        if validation.type != "custom":
            continue
        for cell_range in validation.ranges.ranges:
            validated_price_columns.add(cell_range.min_col)
    assert {
        headers["opening_price"],
        headers["target_price"],
        headers["owner_hard_floor"],
    }.issubset(validated_price_columns)


def test_bundle_is_deterministic_and_manifest_hashes_reopen(tmp_path):
    first = generate_bundle(tmp_path / "first", AS_OF, strict=True)
    second = generate_bundle(tmp_path / "second", AS_OF, strict=True)
    for key in ("workbook", "csv", "json"):
        assert Path(first[key]).read_bytes() == Path(second[key]).read_bytes()

    manifest = json.loads(Path(first["manifest"]).read_text(encoding="utf-8"))
    assert manifest["as_of"] == AS_OF
    assert manifest["external_writes"] == 0
    for artifact in manifest["artifacts"]:
        path = Path(first["manifest"]).parent / artifact["name"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == artifact["sha256"]

    with Path(first["csv"]).open(newline="", encoding="utf-8") as handle:
        assert len(list(csv.DictReader(handle))) == manifest["row_count"]
    assert len(json.loads(Path(first["json"]).read_text(encoding="utf-8"))["rows"]) == manifest["row_count"]
