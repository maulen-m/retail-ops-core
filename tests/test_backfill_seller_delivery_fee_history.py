from __future__ import annotations

import json
import sqlite3
from pathlib import Path
from types import SimpleNamespace

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
import pytest

from scripts import backfill_seller_delivery_fee_history as hist_mod


def test_history_backfill_dry_run_emits_summary(monkeypatch, tmp_path):
    summary = tmp_path / "summary.json"

    monkeypatch.setattr(
        hist_mod,
        "run_history_backfill",
        lambda *args, **kwargs: {
            "apply": False,
            "targets": {
                "crm": {"rows_considered": 3, "seller_fee_updates": 2},
                "archive": {"rows_considered": 3, "seller_fee_updates": 2},
            },
        },
    )

    rc = hist_mod.main(["--summary-out", str(summary)])
    assert rc == 0
    payload = json.loads(summary.read_text(encoding="utf-8"))
    assert payload["apply"] is False
    assert payload["targets"]["crm"]["seller_fee_updates"] == 2
    assert payload["targets"]["archive"]["seller_fee_updates"] == 2


def test_history_backfill_apply_requires_backup_and_reports_targets(monkeypatch, tmp_path):
    summary = tmp_path / "summary.json"
    backup_dir = tmp_path / "backups"

    monkeypatch.setattr(
        hist_mod,
        "run_history_backfill",
        lambda *args, **kwargs: {
            "apply": True,
            "backup_dir": str(backup_dir),
            "targets": {
                "crm": {"rows_considered": 4, "seller_fee_updates": 3},
                "archive": {"rows_considered": 4, "seller_fee_updates": 3},
            },
        },
    )

    rc = hist_mod.main(["--apply", "--summary-out", str(summary), "--backup-dir", str(backup_dir)])
    assert rc == 0
    payload = json.loads(summary.read_text(encoding="utf-8"))
    assert payload["apply"] is True
    assert payload["backup_dir"] == str(backup_dir)
    assert payload["targets"]["crm"]["seller_fee_updates"] == 3
    assert payload["targets"]["archive"]["seller_fee_updates"] == 3


def _build_history_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SALES_KSP_CRM_1"
    ws.append(
        [
            "Date",
            "OrderID",
            "STORE_NAME",
            "Delivery_fee_kzt",
            "Стоимость доставки для продавца",
        ]
    )
    ws.append(["01.04.2026", "1001", "Universal", 700, 0])
    ws.append(["01.04.2026", "1002", "AcmeWear", 450, 0])
    ws.append(["30.03.2026", "1003", "STORE-B", 300, 0])

    table = Table(displayName="tb_SalesRaw", ref="A1:E4")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    wb.save(path)
    wb.close()


def _build_db_with_seller_fees(path: Path) -> None:
    conn = sqlite3.connect(path)
    try:
        conn.execute(
            """
            CREATE TABLE fact_orders_kaspi (
                order_id TEXT,
                delivery_cost_for_seller REAL
            )
            """
        )
        conn.executemany(
            "INSERT INTO fact_orders_kaspi(order_id, delivery_cost_for_seller) VALUES (?, ?)",
            [("1001", 1111.0), ("1002", 0.0), ("1003", 333.0)],
        )
        conn.commit()
    finally:
        conn.close()


def test_apply_crm_seller_fee_backfill_openpyxl_updates_db_and_raw_fee(monkeypatch, tmp_path):
    workbook = tmp_path / "crm.xlsx"
    db_path = tmp_path / "app.db"
    _build_history_workbook(workbook)
    _build_db_with_seller_fees(db_path)
    monkeypatch.setattr(
        hist_mod,
        "_refresh_formula_caches_xlwings",
        lambda workbook_path, verbose=False: None,
    )

    updated = hist_mod._apply_crm_seller_fee_backfill_openpyxl(
        crm_workbook=workbook,
        crm_sheet="SALES_KSP_CRM_1",
        db_path=db_path,
        date_from="2026-04-01",
        date_to="2026-04-01",
    )

    assert updated == 2

    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb["SALES_KSP_CRM_1"]
    assert ws["E2"].value == 1111
    assert ws["E3"].value == 450
    assert ws["E4"].value == 0
    wb.close()


def test_run_history_backfill_apply_records_actual_crm_updates(monkeypatch, tmp_path):
    crm = tmp_path / "crm.xlsx"
    archive_csv = tmp_path / "ArchiveOrders_ALL_STORES.csv"
    archive_xlsx = tmp_path / "ArchiveOrders_ALL_STORES.xlsx"
    crm.write_text("crm", encoding="utf-8")
    archive_csv.write_text("csv", encoding="utf-8")
    archive_xlsx.write_text("xlsx", encoding="utf-8")

    monkeypatch.setattr(hist_mod, "_resolve_date_range", lambda *args, **kwargs: (
        hist_mod.datetime.strptime("2026-04-01", "%Y-%m-%d").date(),
        hist_mod.datetime.strptime("2026-04-02", "%Y-%m-%d").date(),
    ))
    monkeypatch.setattr(hist_mod, "_archive_paths", lambda root: (archive_csv, archive_xlsx))
    monkeypatch.setattr(
        hist_mod,
        "_scan_crm_candidates",
        lambda *args, **kwargs: {
            "rows_considered": 5,
            "seller_fee_updates": 9,
            "stores": {"Universal": 9},
            "workbook": str(crm),
        },
    )
    monkeypatch.setattr(
        hist_mod,
        "_scan_archive_candidates",
        lambda *args, **kwargs: (
            {"rows_considered": 4, "seller_fee_updates": 7, "stores": {"Universal": 7}, "csv_path": str(archive_csv)},
            ["h1"],
            [{"h1": "v1"}],
            {},
        ),
    )
    monkeypatch.setattr(hist_mod, "_backup_file", lambda path, backup_dir: backup_dir / f"{path.name}.bak")
    monkeypatch.setattr(hist_mod, "_apply_crm_seller_fee_backfill_openpyxl", lambda **kwargs: 3)
    monkeypatch.setattr(hist_mod, "_apply_archive_seller_fee_backfill", lambda **kwargs: 6)

    summary = hist_mod.run_history_backfill(
        crm_workbook=crm,
        crm_sheet="SALES_KSP_CRM_1",
        db_path=tmp_path / "app.db",
        archive_root=tmp_path,
        backup_dir=tmp_path / "backups",
        apply=True,
    )

    assert summary["targets"]["crm"]["seller_fee_updates"] == 3
    assert summary["targets"]["archive"]["seller_fee_updates"] == 6


def test_apply_crm_seller_fee_backfill_openpyxl_allows_baseline_integrity_errors(monkeypatch, tmp_path):
    workbook = tmp_path / "crm.xlsx"
    db_path = tmp_path / "app.db"
    _build_history_workbook(workbook)
    _build_db_with_seller_fees(db_path)

    monkeypatch.setattr(
        hist_mod,
        "validate_workbook_integrity",
        lambda path: SimpleNamespace(
            errors=["named range contains #REF!: B", "named range contains #REF!: _02.12.2025"],
            warnings=[],
        ),
    )
    monkeypatch.setattr(
        hist_mod,
        "_refresh_formula_caches_xlwings",
        lambda workbook_path, verbose=False: None,
    )

    updated = hist_mod._apply_crm_seller_fee_backfill_openpyxl(
        crm_workbook=workbook,
        crm_sheet="SALES_KSP_CRM_1",
        db_path=db_path,
        date_from="2026-04-01",
        date_to="2026-04-01",
    )

    assert updated == 2


def test_apply_crm_seller_fee_backfill_openpyxl_rejects_new_integrity_errors(monkeypatch, tmp_path):
    workbook = tmp_path / "crm.xlsx"
    db_path = tmp_path / "app.db"
    _build_history_workbook(workbook)
    _build_db_with_seller_fees(db_path)

    monkeypatch.setattr(
        hist_mod,
        "validate_workbook_integrity",
        lambda path: SimpleNamespace(errors=["tablePart target missing: xl/tables/table9.xml"], warnings=[]),
    )
    monkeypatch.setattr(
        hist_mod,
        "_refresh_formula_caches_xlwings",
        lambda workbook_path, verbose=False: None,
    )

    with pytest.raises(RuntimeError, match="integrity validation failed"):
        hist_mod._apply_crm_seller_fee_backfill_openpyxl(
            crm_workbook=workbook,
            crm_sheet="SALES_KSP_CRM_1",
            db_path=db_path,
            date_from="2026-04-01",
            date_to="2026-04-01",
        )


def test_apply_crm_seller_fee_backfill_openpyxl_reads_cached_values_and_skips_external_excel_refresh(
    monkeypatch, tmp_path
):
    workbook = tmp_path / "crm.xlsx"
    db_path = tmp_path / "app.db"
    _build_history_workbook(workbook)
    _build_db_with_seller_fees(db_path)

    real_load_workbook = openpyxl.load_workbook
    data_only_calls: list[bool] = []
    refresh_calls: list[Path] = []

    def tracking_load_workbook(*args, **kwargs):
        data_only_calls.append(bool(kwargs.get("data_only", False)))
        return real_load_workbook(*args, **kwargs)

    monkeypatch.setattr(hist_mod, "load_workbook", tracking_load_workbook)
    monkeypatch.setattr(
        hist_mod,
        "validate_workbook_integrity",
        lambda path: SimpleNamespace(errors=[], warnings=[]),
    )
    monkeypatch.setattr(
        hist_mod,
        "_refresh_formula_caches_xlwings",
        lambda workbook_path, verbose=False: refresh_calls.append(Path(workbook_path)),
    )

    updated = hist_mod._apply_crm_seller_fee_backfill_openpyxl(
        crm_workbook=workbook,
        crm_sheet="SALES_KSP_CRM_1",
        db_path=db_path,
        date_from="2026-04-01",
        date_to="2026-04-01",
    )

    assert updated == 2
    assert sorted(set(data_only_calls)) == [False, True]
    assert refresh_calls == []


def test_apply_crm_seller_fee_backfill_openpyxl_refreshes_only_canonical_workbook_path(
    monkeypatch, tmp_path
):
    workbook = tmp_path / "crm.xlsx"
    db_path = tmp_path / "app.db"
    _build_history_workbook(workbook)
    _build_db_with_seller_fees(db_path)
    refresh_calls: list[Path] = []

    monkeypatch.setattr(hist_mod, "PROJECT_ROOT", tmp_path)
    monkeypatch.setattr(
        hist_mod,
        "validate_workbook_integrity",
        lambda path: SimpleNamespace(errors=[], warnings=[]),
    )
    monkeypatch.setattr(
        hist_mod,
        "_refresh_formula_caches_xlwings",
        lambda workbook_path, verbose=False: refresh_calls.append(Path(workbook_path)),
    )

    updated = hist_mod._apply_crm_seller_fee_backfill_openpyxl(
        crm_workbook=workbook,
        crm_sheet="SALES_KSP_CRM_1",
        db_path=db_path,
        date_from="2026-04-01",
        date_to="2026-04-01",
    )

    assert updated == 2
    assert refresh_calls == [workbook]
    assert all(".seller_fee_tmp" not in str(path) for path in refresh_calls)


def test_refresh_formula_caches_if_trusted_calls_excel_for_project_paths(monkeypatch, tmp_path):
    workbook = tmp_path / "crm.xlsx"
    workbook.write_text("x", encoding="utf-8")
    refresh_calls: list[Path] = []

    monkeypatch.setattr(hist_mod, "PROJECT_ROOT", tmp_path)
    monkeypatch.setattr(
        hist_mod,
        "_refresh_formula_caches_xlwings",
        lambda workbook_path, verbose=False: refresh_calls.append(Path(workbook_path)),
    )

    hist_mod._refresh_formula_caches_if_trusted(workbook)

    assert refresh_calls == [workbook]
