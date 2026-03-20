from __future__ import annotations

import subprocess
import sys
import sqlite3
from pathlib import Path

from openpyxl import load_workbook


SCRIPT = Path("scripts/prepare_ci_headless_fixture.py")


def test_prepare_ci_headless_fixture_script_exists() -> None:
    assert SCRIPT.exists(), "missing scripts/prepare_ci_headless_fixture.py"


def test_prepare_ci_headless_fixture_creates_anchors_and_workbooks(tmp_path: Path) -> None:
    project_root = tmp_path / "repo"
    (project_root / "config" / "anchors").mkdir(parents=True, exist_ok=True)

    completed = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--project-root",
            str(project_root),
            "--as-of",
            "2026-02-20",
        ],
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stdout + completed.stderr

    crm_anchor = project_root / "config" / "anchors" / "SALES_KSP_CRM_LATEST.xlsx"
    inbound_anchor = project_root / "config" / "anchors" / "INBOUND_CALENDAR_LATEST.xlsx"
    stock_anchor = project_root / "config" / "anchors" / "STOCK_SNAPSHOT_LATEST.xlsx"
    assert crm_anchor.is_symlink()
    assert inbound_anchor.is_symlink()
    assert stock_anchor.is_symlink()
    assert crm_anchor.resolve(strict=True).exists()
    assert inbound_anchor.resolve(strict=True).exists()
    assert stock_anchor.resolve(strict=True).exists()

    wb = load_workbook(crm_anchor.resolve(strict=True), read_only=True, data_only=True)
    try:
        assert "SALES_KSP_CRM_1" in wb.sheetnames
    finally:
        wb.close()


def test_prepare_ci_headless_fixture_creates_strict_gate_artifacts(tmp_path: Path) -> None:
    project_root = tmp_path / "repo"
    (project_root / "config" / "anchors").mkdir(parents=True, exist_ok=True)

    completed = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--project-root",
            str(project_root),
            "--as-of",
            "2026-02-20",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr

    db_path = project_root / "db" / "app.db"
    assert db_path.exists(), "fixture must create db/app.db for strict validators"

    with sqlite3.connect(str(db_path)) as conn:
        rows = conn.execute("SELECT COUNT(*) FROM sales_fact_v2").fetchone()[0]
    assert rows >= 14, "fixture should seed overlap rows for workbook anchor validation"

    dashboard_path = project_root / "exports" / "po_dashboard_data.json"
    assert dashboard_path.exists(), "fixture must create dashboard payload for strict validators"

    snapshot = project_root / "config" / "business_insides" / "BUSINESS_INSIDES_2026-02-20.md"
    assert snapshot.exists(), "fixture must create business-insides snapshot for strict validators"

    dim_sku_light = project_root / "config" / "anchors" / "fixtures" / "DIM_SKU_LIGHT_V5.fixture.xlsx"
    assert dim_sku_light.exists(), "fixture must create dim-sku-light workbook for alignment validator"


def test_prepare_ci_headless_fixture_creates_inbounds_sheet_for_strict_validators(tmp_path: Path) -> None:
    project_root = tmp_path / "repo"
    (project_root / "config" / "anchors").mkdir(parents=True, exist_ok=True)

    completed = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--project-root",
            str(project_root),
            "--as-of",
            "2026-02-20",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr

    inbound_anchor = project_root / "config" / "anchors" / "INBOUND_CALENDAR_LATEST.xlsx"
    wb = load_workbook(inbound_anchor.resolve(strict=True), read_only=True, data_only=True)
    try:
        assert "Inbounds_sheet" in wb.sheetnames
        ws = wb["Inbounds_sheet"]
        header = [cell.value for cell in ws[1]]
    finally:
        wb.close()

    assert {"PO_part_id", "SKU_key"}.issubset(set(header))
    assert ("Qty" in header) or ("Actual_qty" in header)


def test_prepare_ci_headless_fixture_includes_strict_single_truth_columns_and_sheets(tmp_path: Path) -> None:
    project_root = tmp_path / "repo"
    (project_root / "config" / "anchors").mkdir(parents=True, exist_ok=True)

    completed = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--project-root",
            str(project_root),
            "--as-of",
            "2026-02-20",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr

    inbound_anchor = project_root / "config" / "anchors" / "INBOUND_CALENDAR_LATEST.xlsx"
    wb = load_workbook(inbound_anchor.resolve(strict=True), read_only=True, data_only=True)
    try:
        assert "2.3.26_astana_totals" in wb.sheetnames
        assert "dlv_payment_2.3.26" in wb.sheetnames
        assert "cargo_send_ci" in wb.sheetnames
        totals = wb["PO_part_id_Totals"]
        header = [cell.value for cell in totals[1]]
    finally:
        wb.close()

    assert {
        "Cargo_freight_id",
        "Actual_DLV_PAY_date",
        "Actual_Weight_kg",
        "Paid_DLV_USD",
        "Paid_DLV_KZT",
        "Final_USD_per_kg",
        "USD_KZT_rate",
        "Actual_DLV_days",
    }.issubset(set(header))


def test_prepare_ci_headless_fixture_po_part_schema_matches_strict_validator(tmp_path: Path) -> None:
    project_root = tmp_path / "repo"
    (project_root / "config" / "anchors").mkdir(parents=True, exist_ok=True)

    completed = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--project-root",
            str(project_root),
            "--as-of",
            "2026-02-20",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr

    db_path = project_root / "db" / "app.db"
    with sqlite3.connect(str(db_path)) as conn:
        cols = {row[1] for row in conn.execute("PRAGMA table_info(po_part)").fetchall()}

    assert {
        "cargo_freight_id",
        "actual_dlv_pay_date",
        "actual_weight_kg",
        "paid_dlv_usd",
        "paid_dlv_kzt",
        "final_usd_per_kg",
        "usd_kzt_rate",
        "actual_dlv_days",
    }.issubset(cols)
