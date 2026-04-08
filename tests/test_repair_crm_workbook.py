from __future__ import annotations

import re
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo

from scripts.repair_crm_workbook import repair_workbook_package


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def _inject_bad_metadata(xlsx_path: Path) -> None:
    with TemporaryDirectory() as td:
        td_path = Path(td)
        with zipfile.ZipFile(xlsx_path, "r") as zf:
            zf.extractall(td_path)

        wb_xml = td_path / "xl" / "workbook.xml"
        wb_root = ET.fromstring(wb_xml.read_bytes())

        ext_refs = ET.SubElement(wb_root, f"{{{MAIN_NS}}}externalReferences")
        ext_ref = ET.SubElement(ext_refs, f"{{{MAIN_NS}}}externalReference")
        ext_ref.set(f"{{{REL_NS}}}id", "rIdExt1")

        dns = wb_root.find(f"{{{MAIN_NS}}}definedNames")
        if dns is None:
            dns = ET.SubElement(wb_root, f"{{{MAIN_NS}}}definedNames")
        dn = ET.SubElement(dns, f"{{{MAIN_NS}}}definedName", {"name": "BROKEN_NAME"})
        dn.text = "[1]ExternalBook.xlsx!#REF!"
        wb_xml.write_bytes(ET.tostring(wb_root, encoding="utf-8", xml_declaration=True))

        rels_xml = td_path / "xl" / "_rels" / "workbook.xml.rels"
        rels_root = ET.fromstring(rels_xml.read_bytes())
        ET.SubElement(
            rels_root,
            f"{{{PKG_REL_NS}}}Relationship",
            {
                "Id": "rIdExt1",
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink",
                "Target": "/xl/externalLinks/externalLink1.xml",
            },
        )
        rels_xml.write_bytes(ET.tostring(rels_root, encoding="utf-8", xml_declaration=True))

        ext_dir = td_path / "xl" / "externalLinks"
        ext_rels_dir = ext_dir / "_rels"
        ext_dir.mkdir(parents=True, exist_ok=True)
        ext_rels_dir.mkdir(parents=True, exist_ok=True)
        (ext_dir / "externalLink1.xml").write_text(
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>',
            encoding="utf-8",
        )
        (ext_rels_dir / "externalLink1.xml.rels").write_text(
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>',
            encoding="utf-8",
        )

        ct_xml = td_path / "[Content_Types].xml"
        ct_root = ET.fromstring(ct_xml.read_bytes())
        ET.SubElement(
            ct_root,
            f"{{{CT_NS}}}Override",
            {
                "PartName": "/xl/externalLinks/externalLink1.xml",
                "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml",
            },
        )
        ct_xml.write_bytes(ET.tostring(ct_root, encoding="utf-8", xml_declaration=True))

        with zipfile.ZipFile(xlsx_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for file_path in td_path.rglob("*"):
                if file_path.is_file():
                    zf.write(file_path, file_path.relative_to(td_path).as_posix())


def _inject_mc_ignorable_namespace_contract(xlsx_path: Path) -> None:
    workbook_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
 xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
 mc:Ignorable="x15 xr xr6 xr10 xr2"
 xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main"
 xmlns:x15ac="http://schemas.microsoft.com/office/spreadsheetml/2010/11/ac"
 xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision"
 xmlns:xr6="http://schemas.microsoft.com/office/spreadsheetml/2016/revision6"
 xmlns:xr10="http://schemas.microsoft.com/office/spreadsheetml/2016/revision10"
 xmlns:xr2="http://schemas.microsoft.com/office/spreadsheetml/2015/revision2">
<fileVersion appName="xl" lastEdited="7" lowestEdited="7" rupBuild="10125"/>
<workbookPr defaultThemeVersion="166925"/>
<mc:AlternateContent>
  <mc:Choice Requires="x15">
    <x15ac:absPath url="/tmp/"/>
  </mc:Choice>
</mc:AlternateContent>
<xr:revisionPtr revIDLastSave="0" documentId="doc" xr6:coauthVersionLast="47" xr6:coauthVersionMax="47" xr10:uidLastSave="{00000000-0000-0000-0000-000000000000}"/>
<bookViews><workbookView xWindow="0" yWindow="0" windowWidth="12000" windowHeight="8000"/></bookViews>
<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets>
<definedNames><definedName name="BROKEN_NAME">[1]ExternalBook.xlsx!#REF!</definedName></definedNames>
<externalReferences><externalReference r:id="rIdExt1"/></externalReferences>
<calcPr calcId="191029"/>
</workbook>
"""
    workbook_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
  <Relationship Id="rIdExt1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink" Target="externalLinks/externalLink1.xml"/>
</Relationships>
"""

    with TemporaryDirectory() as td:
        td_path = Path(td)
        with zipfile.ZipFile(xlsx_path, "r") as zf:
            zf.extractall(td_path)

        (td_path / "xl" / "workbook.xml").write_text(workbook_xml, encoding="utf-8")
        (td_path / "xl" / "_rels" / "workbook.xml.rels").write_text(workbook_rels, encoding="utf-8")

        ext_dir = td_path / "xl" / "externalLinks"
        ext_rels_dir = ext_dir / "_rels"
        ext_dir.mkdir(parents=True, exist_ok=True)
        ext_rels_dir.mkdir(parents=True, exist_ok=True)
        (ext_dir / "externalLink1.xml").write_text(
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>',
            encoding="utf-8",
        )
        (ext_rels_dir / "externalLink1.xml.rels").write_text(
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>',
            encoding="utf-8",
        )

        ct_xml = td_path / "[Content_Types].xml"
        ct_root = ET.fromstring(ct_xml.read_bytes())
        ET.SubElement(
            ct_root,
            f"{{{CT_NS}}}Override",
            {
                "PartName": "/xl/externalLinks/externalLink1.xml",
                "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml",
            },
        )
        ct_xml.write_bytes(ET.tostring(ct_root, encoding="utf-8", xml_declaration=True))

        with zipfile.ZipFile(xlsx_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for file_path in td_path.rglob("*"):
                if file_path.is_file():
                    zf.write(file_path, file_path.relative_to(td_path).as_posix())


def _inject_array_formula_ref_mismatch(
    xlsx_path: Path,
    *,
    cell_ref: str,
    bad_ref: str,
    sheet_xml_path: str = "xl/worksheets/sheet1.xml",
) -> None:
    with zipfile.ZipFile(xlsx_path, "r") as zin:
        sheet_xml = zin.read(sheet_xml_path).decode("utf-8", "ignore")
        pattern = re.compile(
            rf'(<c\b[^>]*\br="{re.escape(cell_ref)}"[^>]*>.*?<f\b[^>]*\bt="array"[^>]*\bref=")([^"]*)(")',
            re.S,
        )
        updated, count = pattern.subn(lambda m: f"{m.group(1)}{bad_ref}{m.group(3)}", sheet_xml, count=1)
        assert count == 1

        tmp = xlsx_path.with_name(f"{xlsx_path.stem}_bad_array_ref_tmp.xlsx")
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == sheet_xml_path:
                    zout.writestr(item, updated.encode("utf-8"))
                else:
                    zout.writestr(item, zin.read(item.filename))
    tmp.replace(xlsx_path)


def _inject_stale_calcchain_ref(
    xlsx_path: Path,
    *,
    sheet_id: str = "1",
    valid_ref: str = "A2",
    stale_ref: str = "Z999",
) -> None:
    with zipfile.ZipFile(xlsx_path, "r") as zin:
        rels_root = ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
        ET.SubElement(
            rels_root,
            f"{{{PKG_REL_NS}}}Relationship",
            {
                "Id": "rIdCalcChain",
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain",
                "Target": "calcChain.xml",
            },
        )
        rels_modified = ET.tostring(rels_root, encoding="utf-8", xml_declaration=True)

        tmp = xlsx_path.with_name(f"{xlsx_path.stem}_stale_calcchain_tmp.xlsx")
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/_rels/workbook.xml.rels":
                    zout.writestr(item, rels_modified)
                else:
                    zout.writestr(item, zin.read(item.filename))
            zout.writestr(
                "xl/calcChain.xml",
                (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                    f'<c r="{valid_ref}" i="{sheet_id}"/>'
                    f'<c r="{stale_ref}" i="{sheet_id}"/>'
                    "</calcChain>"
                ),
            )
    tmp.replace(xlsx_path)


def test_repair_workbook_package_removes_external_links_and_broken_names():
    with TemporaryDirectory() as td:
        td_path = Path(td)
        src = td_path / "sample.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "ok"
        wb.save(src)
        wb.close()

        _inject_bad_metadata(src)
        out = td_path / "sample_repaired.xlsx"
        stats = repair_workbook_package(src, out)
        assert stats["external_links_removed"] >= 1
        assert stats["broken_defined_names_removed"] >= 1

        with zipfile.ZipFile(out, "r") as zf:
            names = set(zf.namelist())
            assert "xl/externalLinks/externalLink1.xml" not in names
            wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
            assert wb_root.find(f".//{{{MAIN_NS}}}externalReferences") is None
            for dn in wb_root.findall(f".//{{{MAIN_NS}}}definedName"):
                assert "#REF!" not in (dn.text or "")

        wb2 = load_workbook(out, read_only=True)
        assert "Sheet1" in wb2.sheetnames
        wb2.close()


def test_repair_workbook_package_preserves_excel_namespace_contract():
    with TemporaryDirectory() as td:
        td_path = Path(td)
        src = td_path / "sample.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "ok"
        wb.save(src)
        wb.close()

        _inject_mc_ignorable_namespace_contract(src)
        out = td_path / "sample_repaired.xlsx"
        repair_workbook_package(src, out)

        with zipfile.ZipFile(out, "r") as zf:
            workbook_xml = zf.read("xl/workbook.xml").decode("utf-8")
            assert 'mc:Ignorable="x15 xr xr6 xr10 xr2"' in workbook_xml
            assert 'xmlns:x15=' in workbook_xml
            assert 'xmlns:xr=' in workbook_xml
            assert 'xmlns:xr6=' in workbook_xml
            assert 'xmlns:xr10=' in workbook_xml
            assert 'xmlns:xr2=' in workbook_xml
            assert '<xr:revisionPtr' in workbook_xml
            assert '<externalReferences' not in workbook_xml
            assert '#REF!' not in workbook_xml
            assert '<ns0:workbook' not in workbook_xml


def test_repair_workbook_package_normalizes_mismatched_sku_key_array_refs():
    with TemporaryDirectory() as td:
        td_path = Path(td)
        src = td_path / "sales.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "SALES_KSP_CRM_1"
        headers = ["Date", "Phone", "Spacer", "Статус", "SKU_key", "SKU_ID_KSP"]
        for idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=idx, value=header)
        ws.cell(row=2, column=1, value="2026-04-05")
        ws.cell(row=2, column=2, value="77770000000")
        ws.cell(row=2, column=3, value="seed")
        ws.cell(row=2, column=4, value="Новый")
        ws.cell(row=2, column=5, value=ArrayFormula(ref="E2", text='=F2&"_SKU"'))
        ws.cell(row=2, column=6, value="OF_LINE31_XL")
        table = Table(displayName="tb_SalesRaw", ref="A1:F2")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
        wb.save(src)
        wb.close()

        _inject_array_formula_ref_mismatch(src, cell_ref="E2", bad_ref="E3")

        out = td_path / "sales_repaired.xlsx"
        stats = repair_workbook_package(src, out)
        assert stats["sku_key_array_refs_normalized"] == 1

        repaired = load_workbook(out, data_only=False)
        try:
            formula = repaired["SALES_KSP_CRM_1"]["E2"].value
            assert isinstance(formula, ArrayFormula)
            assert formula.ref == "E2"
        finally:
            repaired.close()


def test_repair_workbook_package_prunes_stale_calcchain_refs():
    with TemporaryDirectory() as td:
        td_path = Path(td)
        src = td_path / "calcchain.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = 1
        ws["A2"] = "=A1"
        wb.save(src)
        wb.close()

        _inject_stale_calcchain_ref(src)

        out = td_path / "calcchain_repaired.xlsx"
        stats = repair_workbook_package(src, out)
        assert stats["stale_calcchain_refs_removed"] == 1

        with zipfile.ZipFile(out, "r") as zf:
            calc_chain = zf.read("xl/calcChain.xml").decode("utf-8")
            assert 'r="A2"' in calc_chain
            assert 'r="Z999"' not in calc_chain
