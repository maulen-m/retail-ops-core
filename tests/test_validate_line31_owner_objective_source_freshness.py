from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from scripts.validate_line31_owner_objective_source_freshness import (
    validate_owner_objective_source_freshness,
)


def _write_cash_workbook(path: Path, *, latest_header: str = "01.06.2026_09_06_51") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cash_Balances"
    ws.append(["CASH BALANCES"])
    ws.append([None])
    ws.append([None])
    ws.append(["Store", "Account", "Currency", "27.05.2026_17_03_55", latest_header])
    ws.append(["ACMEWEAR", "kaspi_pay", "KZT", 1000, 2000])

    base = wb.create_sheet("base_payment_SHR_log")
    base.append(["BASE PAYMENT LOG — SHR Supplier"])
    base.append([None])
    base.append(["SECTION 1", None, None, None, None, None, None, "Actual paid (receipts + approval)", 261106])
    base.append([None, None, None, None, None, None, None, "Actual remaining CNY", 61101])
    base.append(["Платёж №", "Дата", "Сумма (CNY)"])
    base.append([18, "30.05.2026", 7000])

    ledger = wb.create_sheet("SHR_Receipt_Ledger")
    ledger.append(["receipt_id", "supplier_id", "receipt_group", "po_pool", "receipt_file", "source_folder", "payment_date", "payment_time", "amount_cny"])
    ledger.append(["PO5_018", "SHR", "PO-5 receipt pool", "PO-5+", "18_7000.png", "PO_5.1_24.03.2026", "30.05.2026", None, 7000])

    wb.create_sheet("SHR_Pay_Lock_105952")
    wb.save(path)


def _write_line31_packet(path: Path) -> None:
    path.mkdir(parents=True)
    totals = {
        "line31_physical_total_stock_qty": 436.0,
        "line31_physical_sellable_stock_qty": 354.0,
        "line31_physical_sellable_available_after_open_reserved_qty": 348.0,
        "line31_physical_not_for_sale_reserve_qty": 82.0,
        "line31_economic_final_sales_estimate_qty": 432.0,
        "line31_open_reserved_on_delivery_exposure_qty": 6.0,
    }
    (path / "product_truth_yellow_to_apply_ready_manifest.json").write_text(
        json.dumps({"gate": "GREEN", "totals": totals}),
        encoding="utf-8",
    )
    (path / "validation_summary.json").write_text(
        json.dumps({"gate": "GREEN"}),
        encoding="utf-8",
    )
    (path / "closeout.md").write_text("Gate: `GREEN`\n", encoding="utf-8")


def _write_owner_facts(path: Path, workbook: Path, packet: Path) -> None:
    path.write_text(
        json.dumps(
            {
                "cash": {
                    "source_workbook": str(workbook),
                    "current_timestamp": "2026-06-01 09:06:51 GMT+5",
                    "protected_reserve_min_kzt": 800000,
                },
                "shr": {
                    "paid_base_total_cny": 261106,
                    "remaining_base_payable_cny": 61101,
                    "po5_payment_18_cny": 7000,
                    "po5_payment_18_final_exchanger_receipt_pending": True,
                },
                "line31_stock": {
                    "source_packet": str(packet),
                    "basis": "owner-approved exact rebuild from April leftovers plus PO1-A arrival",
                    "physical_total_stock_qty": 436,
                    "physical_sellable_stock_qty": 354,
                    "open_reserved_on_delivery_exposure_qty": 6,
                    "physical_sellable_available_after_open_reserved_qty": 348,
                    "physical_not_for_sale_reserve_qty": 82,
                    "economic_final_sales_estimate_qty": 432,
                },
            }
        ),
        encoding="utf-8",
    )


def test_owner_objective_source_freshness_green(tmp_path: Path) -> None:
    workbook = tmp_path / "Inbound_calendar_V10.002.xlsx"
    packet = tmp_path / "line31_packet"
    owner_facts = tmp_path / "CURRENT_OWNER_CLARIFICATION_FACTS.json"
    _write_cash_workbook(workbook)
    _write_line31_packet(packet)
    _write_owner_facts(owner_facts, workbook, packet)

    result = validate_owner_objective_source_freshness(owner_facts)

    assert result["ok"] is True
    assert result["gate"] == "GREEN_SOURCE_FRESHNESS"
    assert result["protected_reserve"] == {
        "expected_min_kzt": 800000,
        "owner_fact_min_kzt": 800000,
        "ok": True,
    }
    assert result["cash_and_shr"]["latest_cash_balance_timestamp"] == "2026-06-01 09:06:51 GMT+5"
    assert result["cash_and_shr"]["shr_payment_18_base_log_present"] is True
    assert result["line31_stock"]["manifest_gate"] == "GREEN"


def test_owner_objective_source_freshness_fails_on_stale_cash_timestamp(tmp_path: Path) -> None:
    workbook = tmp_path / "Inbound_calendar_V10.002.xlsx"
    packet = tmp_path / "line31_packet"
    owner_facts = tmp_path / "CURRENT_OWNER_CLARIFICATION_FACTS.json"
    _write_cash_workbook(workbook, latest_header="31.05.2026_09_06_51")
    _write_line31_packet(packet)
    _write_owner_facts(owner_facts, workbook, packet)

    result = validate_owner_objective_source_freshness(owner_facts)

    assert result["ok"] is False
    assert result["gate"] == "YELLOW_SOURCE_WEAK"
    assert any("Cash_Balances latest timestamp mismatch" in error for error in result["errors"])


def test_owner_objective_source_freshness_fails_on_line31_total_mismatch(tmp_path: Path) -> None:
    workbook = tmp_path / "Inbound_calendar_V10.002.xlsx"
    packet = tmp_path / "line31_packet"
    owner_facts = tmp_path / "CURRENT_OWNER_CLARIFICATION_FACTS.json"
    _write_cash_workbook(workbook)
    _write_line31_packet(packet)
    manifest_path = packet / "product_truth_yellow_to_apply_ready_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["totals"]["line31_physical_total_stock_qty"] = 435
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    _write_owner_facts(owner_facts, workbook, packet)

    result = validate_owner_objective_source_freshness(owner_facts)

    assert result["ok"] is False
    assert any("LINE31 stock total mismatch" in error for error in result["errors"])


def test_owner_objective_source_freshness_fails_on_wrong_protected_reserve(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "Inbound_calendar_V10.002.xlsx"
    packet = tmp_path / "line31_packet"
    owner_facts = tmp_path / "CURRENT_OWNER_CLARIFICATION_FACTS.json"
    _write_cash_workbook(workbook)
    _write_line31_packet(packet)
    _write_owner_facts(owner_facts, workbook, packet)
    facts = json.loads(owner_facts.read_text(encoding="utf-8"))
    facts["cash"]["protected_reserve_min_kzt"] = 1500000
    owner_facts.write_text(json.dumps(facts), encoding="utf-8")

    result = validate_owner_objective_source_freshness(owner_facts)

    assert result["ok"] is False
    assert result["protected_reserve"] == {
        "expected_min_kzt": 800000,
        "owner_fact_min_kzt": 1500000,
        "ok": False,
    }
    assert any("protected reserve is not exactly 800000 KZT" in error for error in result["errors"])
