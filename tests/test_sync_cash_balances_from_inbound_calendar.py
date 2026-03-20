from __future__ import annotations

from pathlib import Path

import openpyxl

from scripts.sync_cash_balances_from_inbound_calendar import sync_cash_balances


def _write_cash_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cash_Balances"
    ws["A4"] = "Store"
    ws["B4"] = "Account"
    ws["C4"] = "Currency"
    ws["D4"] = "02.03.2026\n11:45:41"
    ws["E4"] = "02.03.2026\n14:41:00"

    ws["A5"] = "UNIVERSAL"
    ws["B5"] = "kaspi_gold"
    ws["C5"] = "KZT"
    ws["D5"] = 632000
    ws["E5"] = 46780

    ws["A6"] = "UNIVERSAL"
    ws["B6"] = "cash_usd"
    ws["C6"] = "USD"
    ws["D6"] = 1500
    ws["E6"] = 7

    wb.save(path)


def test_sync_cash_balances_defaults_to_latest_snapshot(tmp_path: Path) -> None:
    xlsx = tmp_path / "inbound.xlsx"
    history = tmp_path / "bank_accounts_history.yaml"
    snapshot = tmp_path / "bank_accounts.yaml"
    history_totals = tmp_path / "bank_accounts_history_totals.md"
    _write_cash_workbook(xlsx)
    history.write_text("entries: []\n", encoding="utf-8")

    report = sync_cash_balances(
        xlsx_path=xlsx,
        history_path=history,
        snapshot_path=snapshot,
        history_totals_path=history_totals,
        db_path=tmp_path / "missing.db",
        snapshot_ts=None,
        apply=False,
    )

    assert report["as_of"] == "2026-03-02 14:41:00 GMT+5"
    assert report["selected_column"] == 5
    assert report["rows"] == 2


def test_sync_cash_balances_can_target_specific_timestamp(tmp_path: Path) -> None:
    xlsx = tmp_path / "inbound.xlsx"
    history = tmp_path / "bank_accounts_history.yaml"
    snapshot = tmp_path / "bank_accounts.yaml"
    history_totals = tmp_path / "bank_accounts_history_totals.md"
    _write_cash_workbook(xlsx)
    history.write_text("entries: []\n", encoding="utf-8")

    report = sync_cash_balances(
        xlsx_path=xlsx,
        history_path=history,
        snapshot_path=snapshot,
        history_totals_path=history_totals,
        db_path=tmp_path / "missing.db",
        snapshot_ts="02.03.2026 11:45:41",
        apply=False,
    )

    assert report["as_of"] == "2026-03-02 11:45:41 GMT+5"
    assert report["selected_column"] == 4
