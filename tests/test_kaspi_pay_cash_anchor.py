import json
import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook

from core.cashflow.kaspi_pay_cash_anchor import (
    DEFAULT_EXPECTED_STORES,
    CashAnchorError,
    apply_cash_anchor,
    build_cash_anchor_preview,
)


SALES_HEADERS = [
    "#",
    "Номер заказа (ID/RRN)",
    "Дата операции",
    "Дата учета операции",
    "Тип операции",
    "Тип оплаты",
    "Сумма операции (т)",
    "Сумма к зачислению/ списанию (т)",
    "Комиссия за операции (т)",
    "Комиссия Kaspi Pay (т)",
    "Стоимость услуги за Kaspi Доставку",
    "Номер карты",
    "Детали покупки",
]


def _write_statement(
    path: Path,
    *,
    account_id: str,
    opening: str = "1000,00",
    closing: str = "1400,00",
    bad_closing: bool = False,
) -> None:
    close = "1401,00" if bad_closing else closing
    path.write_text(
        "\n".join(
            [
                ":20:SANITIZED",
                f":25:{account_id}",
                f":60F:C260101KZT{opening}",
                ":61:2601020102C500,00NTRF//SAFE001",
                ":86:190 sanitized purchase detail must not leak",
                ":61:2605040104D100,00NTRF//SAFE002",
                ":86:010 sanitized current day partial must not leak",
                f":62F:C260504KZT{close}",
                "",
            ]
        ),
        encoding="utf-8",
    )


def _write_sales_report(path: Path, *, after_cutoff: bool = True) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for col, header in enumerate(SALES_HEADERS, start=1):
        ws.cell(row=7, column=col, value=header)
    rows = [
        [
            1,
            "ORDER-RAW-001",
            "02.01.2026",
            "02.01.2026",
            "Покупка",
            "Kaspi Pay",
            500,
            500,
            -50,
            -5,
            -30,
            "4400123412341234",
            "Детали покупки raw private text",
        ]
    ]
    if after_cutoff:
        rows.append(
            [
                2,
                "ORDER-RAW-002",
                "04.05.2026",
                "04.05.2026",
                "Покупка",
                "Kaspi Pay",
                100,
                100,
                -10,
                -1,
                -5,
                "4400999988887777",
                "Детали покупки raw current day",
            ]
        )
    for row_index, row in enumerate(rows, start=8):
        for col, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=col, value=value)
    wb.save(path)


def _write_package(root: Path, *, missing_report_store: str | None = None, bad_statement_store: str | None = None) -> None:
    root.mkdir(parents=True, exist_ok=True)
    for index, store in enumerate(DEFAULT_EXPECTED_STORES, start=1):
        store_dir = root / store
        store_dir.mkdir()
        account_id = f"KZ{index:018d}"
        _write_statement(
            store_dir / f"{store}.txt",
            account_id=account_id,
            bad_closing=store == bad_statement_store,
        )
        if store != missing_report_store:
            _write_sales_report(store_dir / f"{store}.xlsx")


def _read_all_artifacts(root: Path) -> str:
    chunks = []
    for path in sorted(root.rglob("*")):
        if path.is_file():
            chunks.append(path.read_text(encoding="utf-8", errors="ignore"))
    return "\n".join(chunks)


def test_complete_five_store_package_passes(tmp_path: Path) -> None:
    source_root = tmp_path / "source"
    output_root = tmp_path / "preview"
    _write_package(source_root)

    summary = build_cash_anchor_preview(
        source_root=source_root,
        cutoff="2026-05-03",
        output_root=output_root,
        strict=True,
        redact=True,
        run_id="test-preview",
    )

    assert summary["status"] == "PASS"
    assert summary["store_count"] == 5
    assert [row["store_code"] for row in summary["records"]] == [
        "11KZ",
        "MELVIS",
        "STOREB",
        "ACMEWEAR",
        "UNIVERSAL",
    ]
    assert all(row["statement_bridge_error_kzt"] == 0 for row in summary["records"])


def test_missing_store_report_fails(tmp_path: Path) -> None:
    source_root = tmp_path / "source"
    _write_package(source_root, missing_report_store="STOREB")

    with pytest.raises(CashAnchorError, match="one sales report"):
        build_cash_anchor_preview(
            source_root=source_root,
            cutoff="2026-05-03",
            output_root=tmp_path / "preview",
            strict=True,
            redact=True,
            run_id="test-preview",
        )


def test_partial_current_day_is_excluded_from_anchor(tmp_path: Path) -> None:
    source_root = tmp_path / "source"
    _write_package(source_root)

    summary = build_cash_anchor_preview(
        source_root=source_root,
        cutoff="2026-05-03",
        output_root=tmp_path / "preview",
        strict=True,
        redact=True,
        run_id="test-preview",
    )

    for row in summary["records"]:
        assert row["anchor_closing_balance_kzt"] == 1500.0
        assert row["source_statement_closing_balance_kzt"] == 1400.0
        assert row["post_cutoff_txn_count"] == 1
        assert row["post_cutoff_txn_sum_kzt"] == -100.0
        assert row["sales_report_rows_after_cutoff"] == 1


def test_statement_bridge_error_must_be_zero(tmp_path: Path) -> None:
    source_root = tmp_path / "source"
    _write_package(source_root, bad_statement_store="ACMEWEAR")

    with pytest.raises(CashAnchorError, match="statement bridge error"):
        build_cash_anchor_preview(
            source_root=source_root,
            cutoff="2026-05-03",
            output_root=tmp_path / "preview",
            strict=True,
            redact=True,
            run_id="test-preview",
        )


def test_generated_artifacts_are_redacted(tmp_path: Path) -> None:
    source_root = tmp_path / "source"
    output_root = tmp_path / "preview"
    _write_package(source_root)

    build_cash_anchor_preview(
        source_root=source_root,
        cutoff="2026-05-03",
        output_root=output_root,
        strict=True,
        redact=True,
        run_id="test-preview",
    )

    artifact_text = _read_all_artifacts(output_root)
    assert "KZ000000000000000001" not in artifact_text
    assert "4400123412341234" not in artifact_text
    assert "ORDER-RAW-001" not in artifact_text
    assert "sanitized purchase detail" not in artifact_text
    assert "Номер карты" not in artifact_text
    assert "Детали покупки" not in artifact_text


def test_balance_anchor_cannot_create_order_level_cash_in(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    db_path = tmp_path / "cash_anchor.sqlite"
    source_root = tmp_path / "source"
    _write_package(source_root)
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            """
            CREATE TABLE fact_cashflow_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                event_date TEXT NOT NULL,
                event_type TEXT NOT NULL,
                account TEXT NOT NULL,
                amount_kzt REAL NOT NULL,
                ref_type TEXT,
                ref_id TEXT,
                source TEXT,
                event_hash TEXT
            )
            """
        )

    monkeypatch.setenv("ENABLE_CASHFLOW_ANCHOR_WRITE", "1")
    apply_cash_anchor(
        db_path=db_path,
        source_root=source_root,
        cutoff="2026-05-03",
        run_id="test-run",
        output_root=tmp_path / "apply",
        strict=True,
        redact=True,
        apply=True,
    )

    with sqlite3.connect(db_path) as conn:
        cash_in_count = conn.execute(
            "SELECT COUNT(*) FROM fact_cashflow_events WHERE event_type='CASH_IN'"
        ).fetchone()[0]
        assert cash_in_count == 0


def test_temp_db_apply_idempotency_and_post_apply_zero(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    db_path = tmp_path / "cash_anchor.sqlite"
    db_path.touch()
    source_root = tmp_path / "source"
    _write_package(source_root)

    monkeypatch.setenv("ENABLE_CASHFLOW_ANCHOR_WRITE", "1")
    first = apply_cash_anchor(
        db_path=db_path,
        source_root=source_root,
        cutoff="2026-05-03",
        run_id="test-run",
        output_root=tmp_path / "apply",
        strict=True,
        redact=True,
        apply=True,
    )
    post_apply = apply_cash_anchor(
        db_path=db_path,
        source_root=source_root,
        cutoff="2026-05-03",
        run_id="test-run",
        output_root=tmp_path / "post_apply",
        strict=True,
        redact=True,
        apply=False,
    )

    assert first["apply"]["inserted_anchor_records"] == 5
    assert first["apply"]["would_insert_anchor_records"] == 5
    assert post_apply["apply"]["would_insert_anchor_records"] == 0
    assert post_apply["apply"]["existing_anchor_records"] == 5
    assert json.loads((tmp_path / "post_apply" / "summary.json").read_text())["apply"][
        "would_insert_anchor_records"
    ] == 0
