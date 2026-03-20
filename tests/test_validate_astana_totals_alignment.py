from __future__ import annotations

from pathlib import Path

import openpyxl

from scripts.validate_astana_totals_alignment import validate_astana_totals_alignment


def _write_workbook(path: Path, *, astana_units: int = 100) -> None:
    wb = openpyxl.Workbook()

    ws_astana = wb.active
    ws_astana.title = "2.3.26_astana_totals"
    ws_astana["B3"] = "TOTAL"
    ws_astana["C3"] = "PO-1.0"
    ws_astana["B8"] = 5
    ws_astana["C8"] = 5
    ws_astana["B9"] = astana_units
    ws_astana["C9"] = astana_units
    ws_astana["B10"] = 20
    ws_astana["C10"] = 20
    ws_astana["B11"] = 50
    ws_astana["C11"] = 50
    ws_astana["B12"] = 25000
    ws_astana["C12"] = 25000
    ws_astana["A56"] = "GRAND TOTAL"
    ws_astana["E56"] = astana_units
    ws_astana["F56"] = 21  # estimate-based by-bag weight (warning only)
    ws_astana["A59"] = "PO_part_id"
    ws_astana["A60"] = "PO-1.0"
    ws_astana["F60"] = 1000

    ws_totals = wb.create_sheet("PO_part_id_Totals")
    ws_totals.append(
        [
            "PO_part_id",
            "Total Units",
            "Total Bags",
            "Base_cost_CNY",
            "Actual_Weight_kg",
            "Paid_DLV_USD",
            "Paid_DLV_KZT",
        ]
    )
    ws_totals.append(["PO-1.0", 100, 5, 1000, 20, 50, 25000])

    ws_dlv = wb.create_sheet("dlv_payment_2.3.26")
    ws_dlv["A13"] = "Total Actual Weight (kg)"
    ws_dlv["B13"] = 20
    ws_dlv["A54"] = "TOTAL DELIVERY COST"
    ws_dlv["B54"] = 50
    ws_dlv["C54"] = 25000

    wb.save(path)


def test_validate_astana_totals_alignment_passes_with_warning(tmp_path: Path) -> None:
    xlsx = tmp_path / "astana_ok.xlsx"
    _write_workbook(xlsx, astana_units=100)

    report = validate_astana_totals_alignment(workbook_path=xlsx)
    assert report["ok"] is True
    assert report["errors"] == []
    assert report["warnings"]


def test_validate_astana_totals_alignment_fails_on_units_mismatch(tmp_path: Path) -> None:
    xlsx = tmp_path / "astana_fail.xlsx"
    _write_workbook(xlsx, astana_units=99)

    report = validate_astana_totals_alignment(workbook_path=xlsx)
    assert report["ok"] is False
    assert any("units" in err for err in report["errors"])
