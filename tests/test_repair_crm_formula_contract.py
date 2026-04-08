from __future__ import annotations

import json
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo

from scripts import repair_crm_formula_contract as repair_mod


def _make_formula_contract_workbook(path: Path, *, array_contract: bool, include_cf_sentinel: bool = False) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SALES_KSP_CRM_1"
    headers = ["Date", "Phone", "Spacer", "Статус", "SKU_key", "SKU_ID_KSP"]
    for idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=idx, value=header)
    for row_num in (2, 3):
        ws.cell(row=row_num, column=1, value="2026-04-04")
        ws.cell(row=row_num, column=2, value=f"7777000000{row_num}")
        ws.cell(row=row_num, column=3, value="seed")
        ws.cell(row=row_num, column=4, value="Новый")
        if array_contract:
            ws.cell(row=row_num, column=5, value=ArrayFormula(ref=f"E{row_num}", text=f'=F{row_num}&"_SKU"'))
        else:
            ws.cell(row=row_num, column=5, value=f'=F{row_num}&"_SKU"')
        ws.cell(row=row_num, column=6, value=f"OF_LINE31_{row_num}")
    table = Table(displayName="tb_SalesRaw", ref="A1:F3")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    for table_col in table.tableColumns:
        if table_col.name == "SKU_key":
            table_col.calculatedColumnFormula = TableFormula(
                attr_text='F2&"_SKU"',
                array=array_contract,
            )
    wb.save(path)
    wb.close()

    if include_cf_sentinel:
        wb = openpyxl.load_workbook(path)
        try:
            ws = wb["SALES_KSP_CRM_1"]
            fill = PatternFill(fill_type="solid", fgColor="FFFF7E79")
            ws.conditional_formatting.add(
                "O2:Q3",
                FormulaRule(formula=['$M2="CL_OC_MEN_LINE52_BLACK"'], fill=fill),
            )
            wb.save(path)
        finally:
            wb.close()

    if array_contract:
        main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with zipfile.ZipFile(path, "r") as zin:
            table_xml = ET.fromstring(zin.read("xl/tables/table1.xml"))
            columns = table_xml.find(f"{{{main_ns}}}tableColumns")
            assert columns is not None
            for table_col in columns.findall(f"{{{main_ns}}}tableColumn"):
                if table_col.attrib.get("name") != "SKU_key":
                    continue
                calc = ET.SubElement(table_col, f"{{{main_ns}}}calculatedColumnFormula")
                calc.set("array", "1")
                calc.text = 'F2&"_SKU"'
                break
            modified = ET.tostring(table_xml, encoding="utf-8", xml_declaration=True)

            tmp = path.with_name(f"{path.stem}_formula_contract_tmp.xlsx")
            with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    if item.filename == "xl/tables/table1.xml":
                        zout.writestr(item, modified)
                    else:
                        zout.writestr(item, zin.read(item.filename))
        tmp.replace(path)


def test_repair_formula_contract_dry_run_reports_array_contract(tmp_path):
    workbook = tmp_path / "crm.xlsx"
    template = tmp_path / "template.xlsx"
    summary = tmp_path / "summary.json"
    _make_formula_contract_workbook(workbook, array_contract=True)
    _make_formula_contract_workbook(template, array_contract=True)

    rc = repair_mod.main(
        [
            "--workbook",
            str(workbook),
            "--template",
            str(template),
            "--summary-out",
            str(summary),
        ]
    )

    assert rc == 0
    payload = json.loads(summary.read_text(encoding="utf-8"))
    assert payload["apply"] is False
    assert payload["targets"]["workbook"]["sku_key_table_formula_array"] is True
    assert payload["targets"]["workbook"]["sku_key_array_cells"] == 2
    assert payload["targets"]["template"]["sku_key_table_formula_array"] is True
    assert payload["targets"]["template"]["sku_key_array_cells"] == 2


def test_repair_formula_contract_apply_restores_array_contract_from_reference(tmp_path):
    workbook = tmp_path / "crm.xlsx"
    template = tmp_path / "template.xlsx"
    reference = tmp_path / "reference.xlsx"
    summary = tmp_path / "summary.json"
    backup_dir = tmp_path / "backups"
    _make_formula_contract_workbook(workbook, array_contract=False)
    _make_formula_contract_workbook(template, array_contract=False)
    _make_formula_contract_workbook(reference, array_contract=True, include_cf_sentinel=True)

    rc = repair_mod.main(
        [
            "--workbook",
            str(workbook),
            "--template",
            str(template),
            "--reference",
            str(reference),
            "--apply",
            "--backup-dir",
            str(backup_dir),
            "--summary-out",
            str(summary),
        ]
    )

    assert rc == 0
    assert list(backup_dir.glob("*.bak")) != []

    wb = openpyxl.load_workbook(workbook, data_only=False)
    try:
        ws = wb["SALES_KSP_CRM_1"]
        assert isinstance(ws["E2"].value, ArrayFormula)
        assert isinstance(ws["E3"].value, ArrayFormula)
        assert ws["E2"].value.text == '=F2&"_SKU"'
        assert ws["E3"].value.text == '=F3&"_SKU"'
        table_col = next(col for col in ws.tables["tb_SalesRaw"].tableColumns if col.name == "SKU_key")
        assert table_col.calculatedColumnFormula is not None
        assert table_col.calculatedColumnFormula.array is True
    finally:
        wb.close()

    payload = json.loads(summary.read_text(encoding="utf-8"))
    assert payload["targets"]["workbook"]["has_kaspi_name_core_cf_sentinel"] is True
