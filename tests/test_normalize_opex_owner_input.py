from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from core.cashflow.opex_owner_input import (
    OpexOwnerInputError,
    build_all_commitments,
    compute_floor_proposal,
    normalize_workbook,
)


OPEX_HEADERS = [
    "owner_action",
    "expense_name",
    "account",
    "expense_type",
    "schedule",
    "payment_day_of_month",
    "source_monthly_kzt",
    "owner_monthly_kzt",
    "source_status",
    "include_in_cashflow",
    "months_left",
    "principal_left_kzt",
    "review_priority",
    "owner_notes",
    "source_ref",
]

LOAN_HEADERS = [
    "owner_action",
    "loan_name",
    "source_status",
    "bank",
    "account",
    "opex_name",
    "start_date",
    "source_payment_day",
    "next_due_date_after_2026_07_02",
    "source_monthly_payment_kzt",
    "owner_monthly_payment_kzt",
    "principal_kzt",
    "balance_now_source_kzt",
    "months_total",
    "last_schedule_date",
    "effective_pa_source",
    "close_date",
    "include_in_cashflow",
    "review_priority",
    "owner_notes",
    "source_ref",
]


def _write_headers(ws, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        ws.cell(6, col).value = header


def _write_row(ws, row_number: int, headers: list[str], payload: dict) -> None:
    for col, header in enumerate(headers, start=1):
        ws.cell(row_number, col).value = payload.get(header)


def _write_fixture(path: Path, *, add_macbook_row: bool = False) -> None:
    wb = Workbook()
    opex = wb.active
    opex.title = "OPEX_INPUT"
    loans = wb.create_sheet("LOANS_INPUT")
    additions = wb.create_sheet("ADDITIONS")
    _write_headers(opex, OPEX_HEADERS)
    _write_headers(loans, LOAN_HEADERS)

    opex_rows = [
        {
            "owner_action": "KEEP",
            "expense_name": "Internet",
            "account": "Business",
            "expense_type": "work_subscriptions",
            "schedule": "monthly",
            "payment_day_of_month": None,
            "source_monthly_kzt": 25000,
            "owner_monthly_kzt": 25000,
            "include_in_cashflow": "YES",
            "source_ref": "fixture",
        },
        {
            "owner_action": "CLOSED",
            "expense_name": "LLM_6",
            "account": "Business",
            "expense_type": "work_subscriptions",
            "schedule": "monthly",
            "payment_day_of_month": 30,
            "source_monthly_kzt": 10000,
            "owner_monthly_kzt": 10000,
            "include_in_cashflow": "YES",
            "source_ref": "fixture",
        },
        {
            "owner_action": "KEEP",
            "expense_name": "Gym_memberships",
            "account": "Life",
            "expense_type": "health",
            "schedule": "monthly",
            "payment_day_of_month": 1,
            "source_monthly_kzt": 30000,
            "owner_monthly_kzt": 90000,
            "include_in_cashflow": "YES",
            "source_ref": "fixture",
        },
        {
            "owner_action": "KEEP",
            "expense_name": "GOLD_store-d",
            "account": "Life",
            "expense_type": "lawn",
            "schedule": "monthly",
            "payment_day_of_month": 17,
            "source_monthly_kzt": 100000,
            "owner_monthly_kzt": 100000,
            "include_in_cashflow": "YES",
            "months_left": 7,
            "source_ref": "fixture",
        },
        {
            "owner_action": "KEEP",
            "expense_name": "GOLD_universal",
            "account": "Life",
            "expense_type": "lawn",
            "schedule": "monthly",
            "payment_day_of_month": 17,
            "source_monthly_kzt": 170000,
            "owner_monthly_kzt": 170000,
            "include_in_cashflow": "YES",
            "months_left": 7,
            "source_ref": "fixture",
        },
        {
            "owner_action": "KEEP",
            "expense_name": "BCC_universal_5M",
            "account": "Business",
            "expense_type": "lawn",
            "schedule": "monthly",
            "payment_day_of_month": 9,
            "source_monthly_kzt": 166451,
            "owner_monthly_kzt": 166451,
            "include_in_cashflow": "YES",
            "months_left": 60,
            "source_ref": "fixture",
        },
        {
            "owner_action": "KEEP",
            "expense_name": "food_1",
            "account": "Life",
            "expense_type": "food",
            "schedule": "daily",
            "payment_day_of_month": "daily",
            "source_monthly_kzt": 30000,
            "owner_monthly_kzt": 30000,
            "include_in_cashflow": "YES",
            "source_ref": "fixture",
        },
    ]
    if add_macbook_row:
        opex_rows.append(
            {
                "owner_action": "KEEP",
                "expense_name": "MacBook installment",
                "account": "Life",
                "expense_type": "lawn",
                "schedule": "monthly",
                "payment_day_of_month": 20,
                "source_monthly_kzt": 100000,
                "owner_monthly_kzt": 100000,
                "include_in_cashflow": "YES",
                "source_ref": "bad fixture",
            }
        )
    for idx, payload in enumerate(opex_rows, start=7):
        _write_row(opex, idx, OPEX_HEADERS, payload)

    loan_rows = [
        {
            "owner_action": "CHANGE",
            "loan_name": "KaspiGOLD_KZ",
            "source_status": "ACTIVE",
            "bank": "Kaspi_Gold",
            "account": "store-d",
            "opex_name": "GOLD_store-d",
            "start_date": datetime(2025, 8, 17),
            "source_payment_day": 17,
            "next_due_date_after_2026_07_02": datetime(2026, 7, 17),
            "source_monthly_payment_kzt": 233222,
            "owner_monthly_payment_kzt": 80000,
            "principal_kzt": 1108320,
            "months_total": 11,
            "last_schedule_date": datetime(2026, 6, 17),
            "include_in_cashflow": "REVIEW",
            "source_ref": "fixture",
        },
        {
            "owner_action": "CHANGE",
            "loan_name": "KaspiGOLD_Uni",
            "source_status": "ACTIVE",
            "bank": "Kaspi_Gold",
            "account": "Universal",
            "opex_name": "GOLD_universal",
            "source_payment_day": 20,
            "next_due_date_after_2026_07_02": datetime(2026, 7, 20),
            "source_monthly_payment_kzt": 11981,
            "owner_monthly_payment_kzt": 165763,
            "principal_kzt": 1420388,
            "months_total": 21,
            "last_schedule_date": datetime(2027, 4, 20),
            "include_in_cashflow": "YES",
            "source_ref": "fixture",
        },
        {
            "owner_action": "KEEP",
            "loan_name": "BCC_5M_CASH_OPP_2026_F_L_001517",
            "source_status": "ACTIVE",
            "bank": "BCC",
            "account": "Universal",
            "opex_name": "BCC_universal_5M",
            "source_payment_day": 9,
            "next_due_date_after_2026_07_02": datetime(2026, 7, 9),
            "source_monthly_payment_kzt": 166451,
            "owner_monthly_payment_kzt": 166451,
            "principal_kzt": 5000000,
            "balance_now_source_kzt": 4925184.11,
            "months_total": 60,
            "last_schedule_date": datetime(2031, 4, 11),
            "include_in_cashflow": "YES",
            "source_ref": "fixture",
        },
    ]
    for idx, payload in enumerate(loan_rows, start=7):
        _write_row(loans, idx, LOAN_HEADERS, payload)

    additions.cell(1, 1).value = "time-record"
    additions.cell(1, 6).value = "PAYMENT_SCHEDULE"
    additions.cell(1, 7).value = "PAYMENT_AMOUNT_KZT"
    additions.cell(1, 8).value = "status"
    additions.cell(1, 9).value = "NOTES"
    additions.cell(2, 1).value = "KaspiGOLD_Uni"
    additions.cell(2, 2).value = "ACTIVE"
    additions.cell(2, 3).value = "Kaspi_Gold"
    additions.cell(2, 4).value = "Universal"
    additions.cell(2, 5).value = "GOLD_universal"
    additions.cell(2, 6).value = datetime(2025, 8, 20)
    for row in range(2, 13):
        additions.cell(row, 7).value = 100000
        additions.cell(row, 8).value = "paid"
    additions.cell(13, 7).value = 166000
    additions.cell(13, 8).value = "UNPAID"
    additions.cell(13, 9).value = "MacBook installment already included"
    additions.cell(14, 7).value = 150000
    additions.cell(14, 8).value = "UNPAID"

    additions.cell(38, 1).value = "BCC_5M_CASH_OPP_2026_F_L_001517"
    additions.cell(38, 2).value = "ACTIVE"
    additions.cell(38, 3).value = "BCC"
    additions.cell(38, 4).value = "Universal"
    additions.cell(38, 5).value = "BCC_universal_5M"
    for row, month, status in [(39, 5, "paid"), (40, 6, "paid"), (41, 7, "TO_BE_PAID")]:
        additions.cell(row, 6).value = datetime(2026, month, 9)
        additions.cell(row, 7).value = 166451
        additions.cell(row, 8).value = status
    facts = {
        48: ("contract_number", "OPP/2026/F/L/001517"),
        49: ("screenshot_date_received", datetime(2026, 4, 11)),
        50: ("owner_stated_taken_date", datetime(2026, 4, 4)),
        51: ("end_date", datetime(2031, 4, 11)),
        52: ("term_months", 60),
        53: ("contract_amount_kzt", 5000000),
        54: ("outstanding_principal_kzt", 4925184.11),
    }
    for row, (key, value) in facts.items():
        additions.cell(row, 6).value = key
        additions.cell(row, 7).value = value

    wb.save(path)


def _line(normalized: dict, name: str) -> dict:
    for item in normalized["opex_lines"]:
        if item.name == name:
            return item
    raise AssertionError(f"missing line {name}")


def test_normalize_precedence_exclusions_missing_day_and_amount_authority(tmp_path: Path) -> None:
    workbook = tmp_path / "owner.xlsx"
    _write_fixture(workbook)

    normalized = normalize_workbook(workbook, as_of=date(2026, 7, 2), horizon_days=365)

    internet = _line(normalized, "Internet")
    assert internet.day == 27
    assert "ASSUMED_DAY_27" in internet.flags

    llm = _line(normalized, "LLM_6")
    assert llm.included is False
    assert "OWNER_ACTION_OVERRIDES_INCLUDE_FLAG" in llm.flags

    gym = _line(normalized, "Gym_memberships")
    assert gym.monthly_amount_kzt == 90000
    assert gym.amount_authority == "owner_monthly_kzt"
    assert "OWNER_SOURCE_DELTA" in gym.flags


def test_review_variants_dated_overrides_and_floor_math(tmp_path: Path) -> None:
    workbook = tmp_path / "owner.xlsx"
    _write_fixture(workbook)
    as_of = date(2026, 7, 2)
    normalized = normalize_workbook(workbook, as_of=as_of, horizon_days=365)

    gold = _line(normalized, "GOLD_store-d")
    assert gold.owner_action_required is True
    assert "REVIEW_VARIANTS_REQUIRED" in gold.flags

    commitments = build_all_commitments(normalized, as_of=as_of, horizon_days=365)
    v1 = commitments["V1_sheet100k"]
    assert any(row["commit_date"] == "2026-07-21" and row["amount_kzt"] == 166000 for row in v1)
    assert not any(row["amount_kzt"] == 170000 for row in v1)
    assert any(row["commit_date"] == "2026-07-09" and row["amount_kzt"] == 166451 for row in v1)

    proposal = compute_floor_proposal(
        commitments,
        as_of=as_of,
        current_context={
            "current_opex_monthly_kzt": 0,
            "current_conservative_floor_kzt": 0,
            "current_conservative_min_cash_kzt": 3668632.39,
        },
        sidecar=None,
    )
    assert proposal["variants"]["V1_sheet100k"]["opex_monthly_kzt"] == 578451
    assert proposal["variants"]["V2_owner80k"]["opex_monthly_kzt"] == 558451
    assert proposal["variants"]["V3_excluded"]["opex_monthly_kzt"] == 478451
    assert proposal["variants"]["V1_sheet100k"]["conservative_floor_kzt"] == 1367676.5


def test_macbook_separate_row_fails(tmp_path: Path) -> None:
    workbook = tmp_path / "owner.xlsx"
    _write_fixture(workbook, add_macbook_row=True)

    with pytest.raises(OpexOwnerInputError, match="MacBook invariant failed"):
        normalize_workbook(workbook, as_of=date(2026, 7, 2), horizon_days=365)
