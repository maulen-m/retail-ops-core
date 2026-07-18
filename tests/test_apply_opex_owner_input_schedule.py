from __future__ import annotations

import json
import os
import sqlite3
from datetime import date, datetime
from pathlib import Path

import yaml
from openpyxl import Workbook
import pytest

from scripts.apply_opex_owner_input_schedule import (
    DEFAULT_PAYMENT_OVERRIDE,
    DEFAULT_VARIANT,
    OpexOwnerApplyError,
    _apply_payment_override_manifest,
    apply_opex_owner_input_schedule,
    build_owner_approved_commitments,
)


OPEX_HEADERS = [
    "owner_action",
    "expense_name",
    "account",
    "expense_type",
    "schedule",
    "payment_day_of_month",
    "source_monthly_kzt",
    "owner_monthly_kzt",
    "source_status",
    "include_in_cashflow",
    "months_left",
    "principal_left_kzt",
    "review_priority",
    "owner_notes",
    "source_ref",
]

LOAN_HEADERS = [
    "owner_action",
    "loan_name",
    "source_status",
    "bank",
    "account",
    "opex_name",
    "start_date",
    "source_payment_day",
    "next_due_date_after_2026_07_02",
    "source_monthly_payment_kzt",
    "owner_monthly_payment_kzt",
    "principal_kzt",
    "balance_now_source_kzt",
    "months_total",
    "last_schedule_date",
    "effective_pa_source",
    "close_date",
    "include_in_cashflow",
    "review_priority",
    "owner_notes",
    "source_ref",
]


def _write_headers(ws, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        ws.cell(6, col).value = header


def _write_row(ws, row_number: int, headers: list[str], payload: dict) -> None:
    for col, header in enumerate(headers, start=1):
        ws.cell(row_number, col).value = payload.get(header)


def _write_owner_fixture(path: Path) -> None:
    wb = Workbook()
    opex = wb.active
    opex.title = "OPEX_INPUT"
    loans = wb.create_sheet("LOANS_INPUT")
    additions = wb.create_sheet("ADDITIONS")
    _write_headers(opex, OPEX_HEADERS)
    _write_headers(loans, LOAN_HEADERS)

    opex_rows = [
        {
            "owner_action": "KEEP",
            "expense_name": "Internet",
            "account": "Business",
            "expense_type": "work_subscriptions",
            "schedule": "monthly",
            "source_monthly_kzt": 25000,
            "owner_monthly_kzt": 25000,
            "include_in_cashflow": "YES",
            "source_ref": "fixture",
        },
        {
            "owner_action": "CLOSED",
            "expense_name": "LLM_6",
            "account": "Business",
            "expense_type": "work_subscriptions",
            "schedule": "monthly",
            "payment_day_of_month": 30,
            "source_monthly_kzt": 10000,
            "owner_monthly_kzt": 10000,
            "include_in_cashflow": "YES",
            "source_ref": "fixture",
        },
        {
            "owner_action": "KEEP",
            "expense_name": "Gym_memberships",
            "account": "Life",
            "expense_type": "health",
            "schedule": "monthly",
            "payment_day_of_month": 1,
            "source_monthly_kzt": 30000,
            "owner_monthly_kzt": 90000,
            "include_in_cashflow": "YES",
            "source_ref": "fixture",
        },
        {
            "owner_action": "KEEP",
            "expense_name": "GOLD_store-d",
            "account": "Life",
            "expense_type": "lawn",
            "schedule": "monthly",
            "payment_day_of_month": 17,
            "source_monthly_kzt": 100000,
            "owner_monthly_kzt": 100000,
            "include_in_cashflow": "YES",
            "months_left": 7,
            "source_ref": "fixture",
        },
        {
            "owner_action": "KEEP",
            "expense_name": "GOLD_Acmewear",
            "account": "Life",
            "expense_type": "lawn",
            "schedule": "monthly",
            "payment_day_of_month": 3,
            "source_monthly_kzt": 176511,
            "owner_monthly_kzt": 176511,
            "include_in_cashflow": "YES",
            "months_left": 6,
            "source_ref": "fixture",
        },
        {
            "owner_action": "KEEP",
            "expense_name": "GOLD_universal",
            "account": "Life",
            "expense_type": "lawn",
            "schedule": "monthly",
            "payment_day_of_month": 17,
            "source_monthly_kzt": 170000,
            "owner_monthly_kzt": 170000,
            "include_in_cashflow": "YES",
            "months_left": 7,
            "source_ref": "fixture",
        },
        {
            "owner_action": "KEEP",
            "expense_name": "BCC_universal_5M",
            "account": "Business",
            "expense_type": "lawn",
            "schedule": "monthly",
            "payment_day_of_month": 9,
            "source_monthly_kzt": 166451,
            "owner_monthly_kzt": 166451,
            "include_in_cashflow": "YES",
            "months_left": 60,
            "source_ref": "fixture",
        },
    ]
    for idx, row in enumerate(opex_rows, start=7):
        _write_row(opex, idx, OPEX_HEADERS, row)

    loan_rows = [
        {
            "owner_action": "CHANGE",
            "loan_name": "KaspiGOLD_KZ",
            "source_status": "ACTIVE",
            "bank": "Kaspi_Gold",
            "account": "store-d",
            "opex_name": "GOLD_store-d",
            "source_payment_day": 17,
            "next_due_date_after_2026_07_02": datetime(2026, 7, 17),
            "source_monthly_payment_kzt": 233222,
            "owner_monthly_payment_kzt": 80000,
            "months_total": 11,
            "last_schedule_date": datetime(2026, 6, 17),
            "include_in_cashflow": "REVIEW",
            "source_ref": "fixture",
        },
        {
            "owner_action": "KEEP",
            "loan_name": "KaspiGOLD_OF",
            "source_status": "ACTIVE",
            "bank": "Kaspi_Gold",
            "account": "Acmewear",
            "opex_name": "GOLD_Acmewear",
            "source_payment_day": 3,
            "next_due_date_after_2026_07_02": datetime(2026, 7, 3),
            "source_monthly_payment_kzt": 92049,
            "owner_monthly_payment_kzt": 92049,
            "months_total": 18,
            "last_schedule_date": datetime(2027, 1, 3),
            "include_in_cashflow": "YES",
            "source_ref": "fixture",
        },
        {
            "owner_action": "CHANGE",
            "loan_name": "KaspiGOLD_Uni",
            "source_status": "ACTIVE",
            "bank": "Kaspi_Gold",
            "account": "Universal",
            "opex_name": "GOLD_universal",
            "source_payment_day": 20,
            "next_due_date_after_2026_07_02": datetime(2026, 7, 20),
            "source_monthly_payment_kzt": 11981,
            "owner_monthly_payment_kzt": 165763,
            "months_total": 21,
            "last_schedule_date": datetime(2027, 4, 20),
            "include_in_cashflow": "YES",
            "source_ref": "fixture",
        },
        {
            "owner_action": "KEEP",
            "loan_name": "BCC_5M_CASH_OPP_2026_F_L_001517",
            "source_status": "ACTIVE",
            "bank": "BCC",
            "account": "Universal",
            "opex_name": "BCC_universal_5M",
            "source_payment_day": 9,
            "next_due_date_after_2026_07_02": datetime(2026, 7, 9),
            "source_monthly_payment_kzt": 166451,
            "owner_monthly_payment_kzt": 166451,
            "months_total": 60,
            "last_schedule_date": datetime(2031, 4, 11),
            "include_in_cashflow": "YES",
            "source_ref": "fixture",
        },
    ]
    for idx, row in enumerate(loan_rows, start=7):
        _write_row(loans, idx, LOAN_HEADERS, row)

    additions.cell(1, 6).value = "PAYMENT_SCHEDULE"
    additions.cell(1, 7).value = "PAYMENT_AMOUNT_KZT"
    additions.cell(1, 8).value = "status"
    additions.cell(1, 9).value = "NOTES"
    additions.cell(2, 1).value = "KaspiGOLD_Uni"
    additions.cell(2, 6).value = datetime(2025, 8, 20)
    for row in range(2, 13):
        additions.cell(row, 7).value = 100000
        additions.cell(row, 8).value = "PAID"
    additions.cell(13, 7).value = 166000
    additions.cell(13, 8).value = "UNPAID"
    additions.cell(13, 9).value = "MacBook installment already included"

    additions.cell(38, 1).value = "BCC_5M_CASH_OPP_2026_F_L_001517"
    for row, month, status in [(39, 5, "PAID"), (40, 6, "PAID"), (41, 7, "TO_BE_PAID")]:
        additions.cell(row, 6).value = datetime(2026, month, 9)
        additions.cell(row, 7).value = 166451
        additions.cell(row, 8).value = status
    additions.cell(48, 6).value = "term_months"
    additions.cell(48, 7).value = 60
    wb.save(path)


def _init_db(db_path: Path) -> None:
    with sqlite3.connect(str(db_path)) as conn:
        conn.executescript(
            """
            CREATE TABLE fact_cashflow_commitments (
                commit_id INTEGER PRIMARY KEY AUTOINCREMENT,
                commit_date TEXT NOT NULL,
                commit_type TEXT NOT NULL,
                amount_kzt REAL NOT NULL,
                probability REAL,
                scenario_tag TEXT,
                ref_id TEXT,
                notes TEXT,
                created_at TEXT DEFAULT (datetime('now'))
            );
            INSERT INTO fact_cashflow_commitments
              (commit_date, commit_type, amount_kzt, scenario_tag, ref_id, notes)
            VALUES
              ('2026-07-01', 'OPEX', 123.0, 'base', 'HISTORY', 'must remain'),
              ('2026-07-03', 'OPEX', 999.0, 'base', 'OLD_FUTURE', 'must be replaced');
            """
        )


def test_build_owner_approved_commitments_applies_stage_b_overrides(tmp_path: Path) -> None:
    workbook = tmp_path / "owner.xlsx"
    _write_owner_fixture(workbook)

    rows = build_owner_approved_commitments(
        workbook_path=workbook,
        as_of=date(2026, 7, 2),
        horizon_days=45,
        variant=DEFAULT_VARIANT,
    )

    assert any(row["ref_id"] and row["commit_date"] == "2026-07-17" and row["amount_kzt"] == 80000 for row in rows)
    assert any(row["commit_date"] == "2026-07-03" and row["amount_kzt"] == 126693 for row in rows)
    assert any("OWNER_INTERIM_REVISIT_20260703" in row["notes"] for row in rows)
    assert any(row["commit_date"] == "2026-08-01" and row["amount_kzt"] == 30000 for row in rows)
    assert not any(row["amount_kzt"] == 90000 for row in rows)
    assert any(row["commit_date"] == "2026-07-21" and row["amount_kzt"] == 166000 for row in rows)


def test_apply_requires_env_gate_and_preserves_history(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    workbook = tmp_path / "owner.xlsx"
    decision = tmp_path / "decision.json"
    db_path = tmp_path / "app.db"
    out_dir = tmp_path / "opex"
    _write_owner_fixture(workbook)
    decision.write_text("{}", encoding="utf-8")
    _init_db(db_path)
    monkeypatch.delenv("ENABLE_CASHFLOW_WRITE", raising=False)

    with pytest.raises(Exception, match="ENABLE_CASHFLOW_WRITE=1"):
        apply_opex_owner_input_schedule(
            workbook_path=workbook,
            db_path=db_path,
            output_dir=out_dir,
            owner_decision_path=decision,
            as_of=date(2026, 7, 2),
            horizon_days=45,
            apply=True,
        )

    with sqlite3.connect(str(db_path)) as conn:
        assert conn.execute("SELECT COUNT(*) FROM fact_cashflow_commitments").fetchone()[0] == 2

    monkeypatch.setenv("ENABLE_CASHFLOW_WRITE", "1")
    report = apply_opex_owner_input_schedule(
        workbook_path=workbook,
        db_path=db_path,
        output_dir=out_dir,
        owner_decision_path=decision,
        as_of=date(2026, 7, 2),
        horizon_days=45,
        apply=True,
    )

    with sqlite3.connect(str(db_path)) as conn:
        old_row = conn.execute(
            "SELECT amount_kzt, ref_id FROM fact_cashflow_commitments WHERE commit_date='2026-07-01'"
        ).fetchone()
        future_old = conn.execute(
            "SELECT COUNT(*) FROM fact_cashflow_commitments WHERE ref_id='OLD_FUTURE'"
        ).fetchone()[0]
        future_rows = conn.execute(
            "SELECT COUNT(*) FROM fact_cashflow_commitments WHERE commit_date >= '2026-07-02'"
        ).fetchone()[0]

    assert old_row == (123.0, "HISTORY")
    assert future_old == 0
    assert future_rows == report["db_rows_inserted"]
    assert report["db_rows_deleted"] == 1
    assert yaml.safe_load((out_dir / "opex_schedule.yaml").read_text(encoding="utf-8"))["source_xlsx"] == str(workbook)
    assert (out_dir / "opex_commitments.csv").exists()


def test_actual_owner_workbook_monthly_opex_matches_decision_if_present() -> None:
    workbook = Path("exports/opex_owner_input/2026-07-02/OPEX_and_Loans_OWNER_INPUT_MINIMAL_20260702.xlsx")
    if not workbook.exists():
        pytest.skip("owner workbook is not present in this checkout")

    rows = build_owner_approved_commitments(
        workbook_path=workbook,
        as_of=date(2026, 7, 2),
        horizon_days=365,
        variant=DEFAULT_VARIANT,
        payment_override_path=DEFAULT_PAYMENT_OVERRIDE,
    )
    monthly = sum(
        float(row["amount_kzt"])
        for row in rows
        if date(2026, 7, 2) <= date.fromisoformat(row["commit_date"]) <= date(2026, 8, 1)
    )

    assert monthly == 2723967.0


def _write_payment_override_fixture(path: Path, source_path: Path, *, preimage_amount: int = 80000) -> None:
    payload = {
        "schema_version": "cashflow_commitment_owner_override.v1",
        "decision_id": "TEST_OWNER_PAYMENT_OVERRIDE",
        "captured_at_local": "2026-07-17T21:55:00+05:00",
        "run_id": "test_owner_payment_override",
        "scope": "test",
        "source_file": str(source_path),
        "owner_pay_date_convention": {
            "date_authority": "owner_stated_pay_date",
            "cashflow_treatment": "planned_outflow_date",
            "loan_withdrawal_constraint": "at least one calendar day after owner-stated date",
        },
        "window_start": "2026-07-18",
        "window_end": "2026-07-31",
        "expected_obligation_count": 1,
        "expected_total_kzt": 184000,
        "obligations": [
            {
                "obligation_id": "kaspi_store-d_pay_gold",
                "display_name": "Kaspi 11KZ (pay+gold)",
                "category": "loan_payment",
                "commit_date": "2026-07-18",
                "amount_kzt": 184000,
                "commit_type": "OPEX",
                "scenario_tag": "base",
                "ref_id": "OPEX_OWNER_TEST_11KZ_PAY_GOLD",
                "minimum_bank_withdrawal_date": "2026-07-19",
                "supersedes": [
                    {
                        "commit_date": "2026-07-17",
                        "amount_kzt": preimage_amount,
                        "ref_id": "PAY_11KZ",
                    },
                    {
                        "commit_date": "2026-07-17",
                        "amount_kzt": 80000,
                        "ref_id": "GOLD_11KZ",
                    },
                ],
            }
        ],
    }
    path.write_text(json.dumps(payload), encoding="utf-8")


def test_payment_override_replaces_exact_preimages_and_tags_inserted_row(tmp_path: Path) -> None:
    source = tmp_path / "owner.md"
    source.write_text("owner input", encoding="utf-8")
    manifest = tmp_path / "override.json"
    _write_payment_override_fixture(manifest, source)
    rows = [
        {
            "commit_date": "2026-07-17",
            "commit_type": "OPEX",
            "amount_kzt": 80000.0,
            "scenario_tag": "base",
            "ref_id": "PAY_11KZ",
            "notes": "old pay",
        },
        {
            "commit_date": "2026-07-17",
            "commit_type": "OPEX",
            "amount_kzt": 80000.0,
            "scenario_tag": "base",
            "ref_id": "GOLD_11KZ",
            "notes": "old gold",
        },
    ]

    updated, report = _apply_payment_override_manifest(rows, manifest)

    assert len(updated) == 1
    assert updated[0]["commit_date"] == "2026-07-18"
    assert updated[0]["amount_kzt"] == 184000.0
    assert "tags=owner_stated_pay_date" in updated[0]["notes"]
    assert "run_id=test_owner_payment_override" in updated[0]["notes"]
    assert report["removed_count"] == 2
    assert report["obligation_total_kzt"] == 184000.0


def test_payment_override_fails_closed_on_preimage_drift(tmp_path: Path) -> None:
    source = tmp_path / "owner.md"
    source.write_text("owner input", encoding="utf-8")
    manifest = tmp_path / "override.json"
    _write_payment_override_fixture(manifest, source, preimage_amount=81000)
    rows = [
        {
            "commit_date": "2026-07-17",
            "commit_type": "OPEX",
            "amount_kzt": 80000.0,
            "scenario_tag": "base",
            "ref_id": "PAY_11KZ",
            "notes": "old pay",
        },
        {
            "commit_date": "2026-07-17",
            "commit_type": "OPEX",
            "amount_kzt": 80000.0,
            "scenario_tag": "base",
            "ref_id": "GOLD_11KZ",
            "notes": "old gold",
        },
    ]

    with pytest.raises(OpexOwnerApplyError, match="preimage mismatch"):
        _apply_payment_override_manifest(rows, manifest)
