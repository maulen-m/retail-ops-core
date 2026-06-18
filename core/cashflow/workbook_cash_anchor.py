"""Governed workbook snapshot cash anchor helpers."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import sqlite3
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from core.cashflow.kaspi_pay_cash_anchor import ensure_cash_anchor_schema


SHEET_NAME = "Cash_Balances"
DEFAULT_EXPECTED_STORES = ("UNIVERSAL", "11KZ", "STOREB", "ACMEWEAR", "MELVIS")
DEFAULT_EXPECTED_OPERATING_TOTAL_KZT = Decimal("3937364")
DEFAULT_EXPECTED_RESERVE_KZT = Decimal("1500000")
DEFAULT_EXPECTED_GRAND_TOTAL_WITH_RESERVE_KZT = Decimal("5437364")
DEFAULT_FX_BASIS = {
    "KZT": Decimal("1"),
    "RUB": Decimal("6"),
    "USD": Decimal("485"),
    "USDT": Decimal("485"),
    "CNY": Decimal("72"),
}
DEFAULT_IN_FLIGHT_CONTEXT = {
    "shr_true_remaining_cny": Decimal("44101"),
    "po1b_remaining_cny_approx": Decimal("17621"),
    "treatment": "explanation_context_only_no_bank_or_binance_movements",
}


class WorkbookCashAnchorError(RuntimeError):
    """Raised when workbook cash anchor evidence violates the OD-002 contract."""


@dataclass(frozen=True)
class WorkbookAccountBalance:
    row_number: int
    store_code: str
    account_name: str
    currency: str
    amount: Decimal
    kzt_equiv: Decimal
    reserve: bool = False


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _money(value: Any) -> Decimal:
    if value is None:
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace("\xa0", "").replace(" ", "")
    if not text:
        return Decimal("0")
    text = re.sub(r"[^0-9,.\-]", "", text)
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    return Decimal(text or "0")


def _round_kzt(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _json_money(value: Decimal) -> int | float:
    rounded = _round_kzt(value)
    if rounded == rounded.to_integral_value():
        return int(rounded)
    return float(rounded)


def _parse_sheet_ts(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value.replace(microsecond=0)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    raw = str(value or "").replace("\n", " ").strip()
    if not raw:
        raise ValueError("empty timestamp header")
    underscore_match = re.fullmatch(
        r"(\d{1,2}\.\d{1,2}\.\d{4})_(\d{1,2})_(\d{1,2})_(\d{1,2})",
        raw,
    )
    if underscore_match:
        raw = (
            f"{underscore_match.group(1)} "
            f"{int(underscore_match.group(2)):02d}:"
            f"{int(underscore_match.group(3)):02d}:"
            f"{int(underscore_match.group(4)):02d}"
        )
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    raise ValueError(f"Unrecognized sheet timestamp: {value!r}")


def _select_snapshot_column(ws: Any, snapshot_label: str) -> tuple[int, str, datetime]:
    target = snapshot_label.strip()
    try:
        target_dt = _parse_sheet_ts(target)
    except ValueError:
        target_dt = None
    for col in range(4, ws.max_column + 1):
        raw = ws.cell(4, col).value
        if raw is None or str(raw).strip() == "":
            continue
        label = str(raw).replace("\n", " ").strip()
        try:
            parsed_dt = _parse_sheet_ts(raw)
        except ValueError:
            continue
        if label == target or (target_dt is not None and parsed_dt == target_dt):
            return col, label, parsed_dt
    raise WorkbookCashAnchorError(f"Requested snapshot timestamp not found: {snapshot_label}")


def _find_row(ws: Any, label: str) -> int:
    wanted = label.strip().upper()
    for row_number in range(1, ws.max_row + 1):
        value = str(ws.cell(row_number, 1).value or "").strip().upper()
        if value == wanted:
            return row_number
    raise WorkbookCashAnchorError(f"Workbook missing marker row: {label}")


def _sum_money(values: list[Decimal]) -> Decimal:
    return _round_kzt(sum(values, Decimal("0")))


def _decimal_to_str(value: Decimal) -> str:
    rounded = _round_kzt(value)
    if rounded == rounded.to_integral_value():
        return str(int(rounded))
    return format(rounded, "f")


def _canonical_payload_hash(snapshot: dict[str, Any]) -> str:
    payload = {
        "snapshot_label": snapshot["snapshot_label"],
        "snapshot_date": snapshot["snapshot_date"],
        "operating_accounts": snapshot["operating_accounts"],
        "reserve_context": snapshot["reserve_context"],
        "store_totals": snapshot["store_totals"],
        "currency_totals": snapshot["currency_totals"],
        "fx_basis": snapshot["fx_basis"],
    }
    return _sha256_text(json.dumps(payload, ensure_ascii=True, sort_keys=True))


def parse_workbook_cash_snapshot(
    *,
    workbook_path: Path,
    sheet_name: str,
    snapshot_label: str,
    expected_stores: tuple[str, ...] = DEFAULT_EXPECTED_STORES,
    fx_basis: dict[str, Decimal] | None = None,
    expected_operating_total_kzt: Decimal = DEFAULT_EXPECTED_OPERATING_TOTAL_KZT,
    expected_reserve_kzt: Decimal = DEFAULT_EXPECTED_RESERVE_KZT,
    expected_grand_total_with_reserve_kzt: Decimal = DEFAULT_EXPECTED_GRAND_TOTAL_WITH_RESERVE_KZT,
) -> dict[str, Any]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    fx = dict(DEFAULT_FX_BASIS)
    if fx_basis:
        fx.update({key.upper(): Decimal(str(value)) for key, value in fx_basis.items()})

    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise WorkbookCashAnchorError(f"Workbook missing sheet: {sheet_name}")
        ws = wb[sheet_name]
        selected_col, selected_label, selected_dt = _select_snapshot_column(ws, snapshot_label)
        store_totals_row = _find_row(ws, "STORE TOTALS (KZT only)")
        currency_totals_row = _find_row(ws, "CURRENCY TOTALS")
        grand_total_row = _find_row(ws, "GRAND TOTAL KZT")
        grand_total_with_reserve_row = _find_row(ws, "GRAND TOTAL +reserve KZT")

        operating_accounts: list[WorkbookAccountBalance] = []
        reserve_accounts: list[WorkbookAccountBalance] = []
        for row_number in range(5, store_totals_row):
            store = str(ws.cell(row_number, 1).value or "").strip()
            account = str(ws.cell(row_number, 2).value or "").strip()
            currency = str(ws.cell(row_number, 3).value or "").strip().upper()
            if not store or not account or not currency:
                continue
            if store not in expected_stores and store.upper() != "RESERVE":
                continue
            amount = _money(ws.cell(row_number, selected_col).value)
            if currency not in fx:
                raise WorkbookCashAnchorError(f"Unsupported currency {currency} at row {row_number}")
            balance = WorkbookAccountBalance(
                row_number=row_number,
                store_code=store if store in expected_stores else "RESERVE",
                account_name=account,
                currency=currency,
                amount=amount,
                kzt_equiv=_round_kzt(amount * fx[currency]),
                reserve=(store.upper() == "RESERVE"),
            )
            if balance.reserve:
                reserve_accounts.append(balance)
            else:
                operating_accounts.append(balance)

        if not operating_accounts:
            raise WorkbookCashAnchorError("No operating account rows parsed")
        observed_stores = {row.store_code for row in operating_accounts}
        missing_stores = sorted(set(expected_stores) - observed_stores)
        if missing_stores:
            raise WorkbookCashAnchorError(f"Missing operating store account rows: {missing_stores}")

        workbook_store_totals: dict[str, Decimal] = {}
        for row_number in range(store_totals_row + 1, currency_totals_row):
            store = str(ws.cell(row_number, 1).value or "").strip()
            if store in expected_stores:
                workbook_store_totals[store] = _round_kzt(_money(ws.cell(row_number, selected_col).value))
        missing_store_totals = sorted(set(expected_stores) - set(workbook_store_totals))
        if missing_store_totals:
            raise WorkbookCashAnchorError(f"Missing workbook store totals: {missing_store_totals}")

        computed_store_totals_kzt_equiv = {
            store: _sum_money([row.kzt_equiv for row in operating_accounts if row.store_code == store])
            for store in expected_stores
        }
        computed_store_totals_kzt_only = {
            store: _sum_money(
                [
                    row.kzt_equiv
                    for row in operating_accounts
                    if row.store_code == store and row.currency == "KZT"
                ]
            )
            for store in expected_stores
        }
        for store in expected_stores:
            if computed_store_totals_kzt_only[store] != workbook_store_totals[store]:
                raise WorkbookCashAnchorError(
                    f"{store} KZT-only total mismatch: accounts={computed_store_totals_kzt_only[store]} "
                    f"workbook={workbook_store_totals[store]}"
                )

        workbook_currency_totals: dict[str, Decimal] = {}
        for row_number in range(currency_totals_row + 1, grand_total_row):
            currency = str(ws.cell(row_number, 1).value or "").strip().upper()
            if currency:
                workbook_currency_totals[currency] = _round_kzt(_money(ws.cell(row_number, selected_col).value))

        computed_currency_totals = {
            currency: _sum_money([row.kzt_equiv for row in operating_accounts if row.currency == currency])
            for currency in sorted({row.currency for row in operating_accounts})
        }
        for currency, computed in computed_currency_totals.items():
            observed = workbook_currency_totals.get(currency)
            if observed is None:
                raise WorkbookCashAnchorError(f"Missing workbook currency total for {currency}")
            if computed != observed:
                raise WorkbookCashAnchorError(
                    f"{currency} total mismatch: accounts={computed} workbook={observed}"
                )

        operating_total = _sum_money([row.kzt_equiv for row in operating_accounts])
        reserve_total = _sum_money([row.kzt_equiv for row in reserve_accounts])
        workbook_operating_total = _round_kzt(_money(ws.cell(grand_total_row, selected_col).value))
        workbook_grand_total_with_reserve = _round_kzt(
            _money(ws.cell(grand_total_with_reserve_row, selected_col).value)
        )
        if operating_total != workbook_operating_total:
            raise WorkbookCashAnchorError(
                f"Operating total mismatch: accounts={operating_total} workbook={workbook_operating_total}"
            )
        if operating_total != expected_operating_total_kzt:
            raise WorkbookCashAnchorError(
                f"Operating total does not match owner OD-002 expected total: "
                f"observed={operating_total} expected={expected_operating_total_kzt}"
            )
        if reserve_total != expected_reserve_kzt:
            raise WorkbookCashAnchorError(
                f"Reserve total does not match owner OD-002 expected reserve: "
                f"observed={reserve_total} expected={expected_reserve_kzt}"
            )
        if workbook_grand_total_with_reserve != expected_grand_total_with_reserve_kzt:
            raise WorkbookCashAnchorError(
                "Grand total with reserve does not match owner OD-002 expected total: "
                f"observed={workbook_grand_total_with_reserve} "
                f"expected={expected_grand_total_with_reserve_kzt}"
            )
        if operating_total + reserve_total != workbook_grand_total_with_reserve:
            raise WorkbookCashAnchorError(
                "Operating plus reserve does not equal workbook reserve-inclusive total"
            )

        def account_to_dict(row: WorkbookAccountBalance) -> dict[str, Any]:
            return {
                "row_number": row.row_number,
                "store_code": row.store_code,
                "account_name": row.account_name,
                "currency": row.currency,
                "amount": _json_money(row.amount),
                "kzt_equiv": _json_money(row.kzt_equiv),
                "reserve": row.reserve,
            }

        snapshot: dict[str, Any] = {
            "source_workbook_path": str(workbook_path),
            "sheet_name": sheet_name,
            "snapshot_label": selected_label,
            "snapshot_column": selected_col,
            "snapshot_timestamp_almaty": f"{selected_dt.strftime('%Y-%m-%d %H:%M:%S')} GMT+5",
            "snapshot_date": selected_dt.date().isoformat(),
            "operating_accounts": [account_to_dict(row) for row in operating_accounts],
            "reserve_accounts": [account_to_dict(row) for row in reserve_accounts],
            "store_totals": {
                store: _json_money(value) for store, value in computed_store_totals_kzt_equiv.items()
            },
            "store_totals_kzt_equiv": {
                store: _json_money(value) for store, value in computed_store_totals_kzt_equiv.items()
            },
            "store_totals_kzt_only": {
                store: _json_money(value) for store, value in computed_store_totals_kzt_only.items()
            },
            "currency_totals": {
                currency: _json_money(value) for currency, value in computed_currency_totals.items()
            },
            "operating_total_kzt": _json_money(operating_total),
            "reserve_context": {
                "reserve_kzt": _json_money(reserve_total),
                "treatment": "non_operating_context_excluded_from_operating_cash_close",
            },
            "grand_total_with_reserve_kzt": _json_money(workbook_grand_total_with_reserve),
            "fx_basis": {currency: _json_money(rate) for currency, rate in fx.items()},
            "account_row_count": len(operating_accounts) + len(reserve_accounts),
            "operating_account_row_count": len(operating_accounts),
            "reserve_account_row_count": len(reserve_accounts),
            "zero_balance_operating_account_rows": sum(1 for row in operating_accounts if row.kzt_equiv == 0),
        }
        snapshot["snapshot_payload_sha256"] = _canonical_payload_hash(snapshot)
        snapshot["workbook_sha256"] = _sha256_file(workbook_path)
        return snapshot
    finally:
        wb.close()


def _anchor_id(record: dict[str, Any]) -> str:
    parts = [
        record["store_code"],
        record["account_identity_hash"],
        record["anchor_date"],
        record["statement_file_sha256"],
        record["sales_report_file_sha256"],
        record["created_by_run_id"],
    ]
    return _sha256_text("|".join(parts))


def build_workbook_anchor_records(snapshot: dict[str, Any], run_id: str) -> list[dict[str, Any]]:
    anchor_date = str(snapshot["snapshot_date"])
    source_root = str(snapshot["source_workbook_path"])
    source_scope = f"{source_root} :: sheet {snapshot['sheet_name']} :: {snapshot['snapshot_label']}"
    workbook_sha = str(snapshot["workbook_sha256"])
    payload_sha = str(snapshot["snapshot_payload_sha256"])
    notes = (
        "OD-002 workbook Cash_Balances snapshot; operating anchor excludes reserve; "
        f"reserve_kzt={snapshot['reserve_context']['reserve_kzt']}; "
        f"grand_total_with_reserve_kzt={snapshot['grand_total_with_reserve_kzt']}; "
        "FX USD/USDT=485 RUB=6 CNY=72; SHR remaining 44101 CNY and PO-1B "
        "remaining approx 17621 CNY are commitment context only."
    )
    records: list[dict[str, Any]] = []
    for account in snapshot["operating_accounts"]:
        account_key = (
            f"{source_scope}|{account['store_code']}|{account['account_name']}|{account['currency']}"
        )
        record = {
            "anchor_id": "",
            "source_root": source_root,
            "source_store_dir": source_scope,
            "source_store_name": (
                f"{account['store_code']}:{account['account_name']}:{account['currency']}"
            ),
            "store_code": account["store_code"],
            "account_identity_hash": _sha256_text(account_key),
            "account_mask": f"{account['account_name']}:{account['currency']}",
            "statement_file_sha256": workbook_sha,
            "sales_report_file_sha256": payload_sha,
            "source_opening_date": anchor_date,
            "source_opening_balance_kzt": float(account["kzt_equiv"]),
            "anchor_date": anchor_date,
            "anchor_closing_balance_kzt": float(account["kzt_equiv"]),
            "source_statement_closing_date": anchor_date,
            "source_statement_closing_balance_kzt": float(account["kzt_equiv"]),
            "post_cutoff_txn_count": 0,
            "post_cutoff_txn_sum_kzt": 0.0,
            "statement_txn_count": 0,
            "statement_txn_count_through_anchor": 0,
            "sales_report_rows_total": int(snapshot["operating_account_row_count"]),
            "sales_report_rows_through_anchor": int(snapshot["operating_account_row_count"]),
            "sales_report_rows_after_cutoff": 0,
            "sales_report_distinct_order_refs": 0,
            "statement_bridge_error_kzt": 0.0,
            "anchor_to_source_bridge_error_kzt": 0.0,
            "partial_day_excluded": 0,
            "reconciliation_status": "RECONCILED",
            "trust_class": "ACTUAL_ANCHOR",
            "notes_redacted": notes,
            "created_by_run_id": run_id,
        }
        record["anchor_id"] = _anchor_id(record)
        records.append(record)
    return records


def _connect_existing(db_path: Path, *, apply: bool) -> sqlite3.Connection:
    if apply:
        return sqlite3.connect(str(db_path))
    return sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)


def _table_exists(conn: sqlite3.Connection, table: str) -> bool:
    return conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name=?",
        (table,),
    ).fetchone() is not None


def _existing_anchor_ids(conn: sqlite3.Connection, anchor_ids: list[str]) -> set[str]:
    if not anchor_ids or not _table_exists(conn, "cashflow_cash_anchor"):
        return set()
    placeholders = ",".join("?" * len(anchor_ids))
    return {
        row[0]
        for row in conn.execute(
            f"SELECT anchor_id FROM cashflow_cash_anchor WHERE anchor_id IN ({placeholders})",
            anchor_ids,
        ).fetchall()
    }


def _insert_anchor(conn: sqlite3.Connection, record: dict[str, Any]) -> None:
    conn.execute(
        """
        INSERT OR IGNORE INTO cashflow_cash_anchor (
            anchor_id, source_root, source_store_dir, source_store_name, store_code,
            account_identity_hash, account_mask, statement_file_sha256, sales_report_file_sha256,
            source_opening_date, source_opening_balance_kzt, anchor_date, anchor_closing_balance_kzt,
            source_statement_closing_date, source_statement_closing_balance_kzt,
            post_cutoff_txn_count, post_cutoff_txn_sum_kzt, statement_txn_count,
            statement_txn_count_through_anchor, sales_report_rows_total,
            sales_report_rows_through_anchor, sales_report_rows_after_cutoff,
            sales_report_distinct_order_refs, statement_bridge_error_kzt,
            anchor_to_source_bridge_error_kzt, partial_day_excluded, reconciliation_status,
            trust_class, notes_redacted, created_by_run_id
        ) VALUES (
            :anchor_id, :source_root, :source_store_dir, :source_store_name, :store_code,
            :account_identity_hash, :account_mask, :statement_file_sha256, :sales_report_file_sha256,
            :source_opening_date, :source_opening_balance_kzt, :anchor_date, :anchor_closing_balance_kzt,
            :source_statement_closing_date, :source_statement_closing_balance_kzt,
            :post_cutoff_txn_count, :post_cutoff_txn_sum_kzt, :statement_txn_count,
            :statement_txn_count_through_anchor, :sales_report_rows_total,
            :sales_report_rows_through_anchor, :sales_report_rows_after_cutoff,
            :sales_report_distinct_order_refs, :statement_bridge_error_kzt,
            :anchor_to_source_bridge_error_kzt, :partial_day_excluded, :reconciliation_status,
            :trust_class, :notes_redacted, :created_by_run_id
        )
        """,
        record,
    )


def _write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fieldnames})


def _write_artifacts(summary: dict[str, Any], output_root: Path) -> None:
    output_root.mkdir(parents=True, exist_ok=True)
    (output_root / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=True, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    account_fields = [
        "row_number",
        "store_code",
        "account_name",
        "currency",
        "amount",
        "kzt_equiv",
        "reserve",
    ]
    _write_csv(output_root / "account_coverage.csv", summary["snapshot"]["operating_accounts"] + summary["snapshot"]["reserve_accounts"], account_fields)
    anchor_fields = [
        "store_code",
        "account_mask",
        "anchor_date",
        "anchor_closing_balance_kzt",
        "reconciliation_status",
        "trust_class",
        "created_by_run_id",
    ]
    _write_csv(output_root / "anchor_preview.csv", summary["records"], anchor_fields)
    balance_check_path = output_root / "cash_balance_check_od002_operating.csv"
    _write_csv(
        balance_check_path,
        [
            {
                "date": summary["snapshot"]["snapshot_date"],
                "cash_balance_kzt": summary["snapshot"]["operating_total_kzt"],
                "notes": (
                    "OD-002 Cash_Balances operating cash excluding Reserve; reserve "
                    f"{summary['snapshot']['reserve_context']['reserve_kzt']} KZT recorded separately"
                ),
                "source": "OWNER_OD002_CASH_BALANCES",
            }
        ],
        ["date", "cash_balance_kzt", "notes", "source"],
    )

    lines = [
        "# OD-002 Workbook Cash Anchor",
        "",
        f"Status: {summary['status']}",
        f"Applied: {str(summary['apply']['applied']).lower()}",
        f"Run id: {summary['run_id']}",
        f"Source workbook: `{summary['snapshot']['source_workbook_path']}`",
        f"Sheet: `{summary['snapshot']['sheet_name']}`",
        f"Snapshot header: `{summary['snapshot']['snapshot_label']}`",
        f"Anchor date: `{summary['snapshot']['snapshot_date']}`",
        f"Operating cash KZT: `{summary['snapshot']['operating_total_kzt']}`",
        f"Reserve KZT: `{summary['snapshot']['reserve_context']['reserve_kzt']}`",
        f"Grand total with reserve KZT: `{summary['snapshot']['grand_total_with_reserve_kzt']}`",
        f"Candidate anchor rows: `{summary['apply']['candidate_anchor_records']}`",
        f"Would insert anchor rows: `{summary['apply']['would_insert_anchor_records']}`",
        f"Inserted anchor rows: `{summary['apply']['inserted_anchor_records']}`",
        "",
        "Reserve treatment: non-operating context, excluded from operating cash close.",
        "In-flight context: SHR 44101 CNY and PO-1B approx 17621 CNY are explanation only.",
    ]
    if summary.get("rollback"):
        lines.extend(["", "## Rollback", ""])
        for command in summary["rollback"].get("commands", []):
            lines.append(f"- `{command}`")
    (output_root / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def apply_workbook_cash_anchor(
    *,
    db_path: Path,
    workbook_path: Path,
    sheet_name: str,
    snapshot_label: str,
    output_root: Path,
    run_id: str,
    apply: bool = False,
    expected_stores: tuple[str, ...] = DEFAULT_EXPECTED_STORES,
    fx_basis: dict[str, Decimal] | None = None,
    expected_operating_total_kzt: Decimal = DEFAULT_EXPECTED_OPERATING_TOTAL_KZT,
    expected_reserve_kzt: Decimal = DEFAULT_EXPECTED_RESERVE_KZT,
    expected_grand_total_with_reserve_kzt: Decimal = DEFAULT_EXPECTED_GRAND_TOTAL_WITH_RESERVE_KZT,
    backup_path: Path | None = None,
) -> dict[str, Any]:
    if not db_path.exists():
        raise FileNotFoundError(f"DB not found: {db_path}")
    if apply and os.environ.get("ENABLE_CASHFLOW_ANCHOR_WRITE") != "1":
        raise RuntimeError("ENABLE_CASHFLOW_ANCHOR_WRITE=1 is required to apply cash anchor writes.")

    snapshot = parse_workbook_cash_snapshot(
        workbook_path=workbook_path,
        sheet_name=sheet_name,
        snapshot_label=snapshot_label,
        expected_stores=expected_stores,
        fx_basis=fx_basis,
        expected_operating_total_kzt=expected_operating_total_kzt,
        expected_reserve_kzt=expected_reserve_kzt,
        expected_grand_total_with_reserve_kzt=expected_grand_total_with_reserve_kzt,
    )
    records = build_workbook_anchor_records(snapshot, run_id)
    anchor_ids = [record["anchor_id"] for record in records]
    inserted = 0
    with _connect_existing(db_path, apply=apply) as conn:
        existing_ids = _existing_anchor_ids(conn, anchor_ids)
        to_insert = [record for record in records if record["anchor_id"] not in existing_ids]
        if apply:
            ensure_cash_anchor_schema(conn)
            for record in to_insert:
                before = conn.total_changes
                _insert_anchor(conn, record)
                inserted += conn.total_changes - before
            conn.commit()

    rollback = {}
    if backup_path is not None:
        rollback = {
            "backup_path": str(backup_path),
            "commands": [
                f"sqlite3 {db_path} \".restore '{backup_path}'\"",
                f"sqlite3 -readonly {db_path} \"PRAGMA integrity_check;\"",
            ],
        }
    summary = {
        "status": "PASS",
        "run_id": run_id,
        "db_path": str(db_path),
        "snapshot": snapshot,
        "records": records,
        "apply": {
            "applied": bool(apply),
            "candidate_anchor_records": len(records),
            "existing_anchor_records": len(existing_ids),
            "would_insert_anchor_records": len(to_insert),
            "inserted_anchor_records": inserted,
            "cashflow_events_created": 0,
        },
        "rollback": rollback,
        "balance_check_csv": str(output_root / "cash_balance_check_od002_operating.csv"),
    }
    _write_artifacts(summary, output_root)
    return summary
