"""Redaction-safe Kaspi Pay cash anchor preview and persistence helpers."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import sqlite3
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


DEFAULT_EXPECTED_STORES = ("11KZ", "MELVIS", "STOREB", "ACMEWEAR", "Universal")
STORE_CODE_BY_FOLDER = {
    "11KZ": "11KZ",
    "MELVIS": "MELVIS",
    "STOREB": "STOREB",
    "ACMEWEAR": "ACMEWEAR",
    "Universal": "UNIVERSAL",
    "UNIVERSAL": "UNIVERSAL",
}
REQUIRED_SALES_HEADERS = (
    "Номер заказа (ID/RRN)",
    "Дата операции",
    "Дата учета операции",
    "Тип операции",
    "Тип оплаты",
    "Сумма операции (т)",
    "Сумма к зачислению/ списанию (т)",
    "Комиссия за операции (т)",
    "Комиссия Kaspi Pay (т)",
    "Стоимость услуги за Kaspi Доставку",
)
SENSITIVE_SALES_HEADERS = {"Номер карты", "Детали покупки"}


class CashAnchorError(RuntimeError):
    """Raised when a cash anchor package violates the strict source contract."""


@dataclass(frozen=True)
class StatementBalances:
    opening: float
    closing: float
    opening_date: str
    closing_date: str


@dataclass(frozen=True)
class StatementTxn:
    value_date: str
    dc_mark: str
    amount_kzt: float
    code_86: str | None = None
    description: str = ""


@dataclass(frozen=True)
class SalesReportSummary:
    rows_total: int
    rows_through_anchor: int
    rows_after_cutoff: int
    distinct_order_refs_through_anchor: int


def _round_money(value: float) -> float:
    rounded = round(float(value or 0.0), 2)
    return 0.0 if rounded == 0 else rounded


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _mask_account(account_id: str) -> str:
    account_id = str(account_id or "").strip()
    if len(account_id) <= 6:
        return "<redacted>"
    prefix = account_id[:2] if account_id.upper().startswith("KZ") else ""
    suffix = account_id[-4:]
    return f"{prefix}{'*' * 14}{suffix}"


def _parse_yy_mm_dd(token: str) -> str | None:
    if not token or len(token) != 6:
        return None
    try:
        year = int(token[0:2])
        month = int(token[2:4])
        day = int(token[4:6])
        year += 2000 if year < 70 else 1900
        return date(year, month, day).isoformat()
    except ValueError:
        return None


def _parse_money(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return 0.0
    text = text.replace("\xa0", "").replace(" ", "")
    text = re.sub(r"[^0-9,\.\-]", "", text)
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    return float(text or 0.0)


def _parse_balance_line(line: str) -> tuple[float, str] | None:
    match = re.match(r":6[02]F:([DC])(\d{6})([A-Z]{3})([0-9,\.]+)", line)
    if not match:
        return None
    value_date = _parse_yy_mm_dd(match.group(2))
    if not value_date:
        return None
    amount = _parse_money(match.group(4))
    return (amount if match.group(1) == "C" else -amount), value_date


def _parse_61_line(line: str) -> StatementTxn | None:
    match = re.match(r":61:(\d{6})(\d{4})?([DC])([A-Z])?([0-9,\.]+)(.*)", line)
    if not match:
        return None
    value_date = _parse_yy_mm_dd(match.group(1))
    if not value_date:
        return None
    amount = _parse_money(match.group(5))
    if match.group(3) == "D":
        amount = -amount
    return StatementTxn(value_date=value_date, dc_mark=match.group(3), amount_kzt=amount)


def _clean_86_text(text: str) -> str:
    cleaned = re.sub(r"\?\d{2}", " ", text)
    return re.sub(r"\s+", " ", cleaned).strip()


def _parse_86_line(line: str) -> tuple[str | None, str]:
    payload = line[4:] if line.startswith(":86:") else line
    payload = payload.strip()
    match = re.match(r"(\d{3})", payload)
    return (match.group(1) if match else None), _clean_86_text(payload)


def _extract_account_id(line: str) -> str | None:
    raw = line.split(":", 1)[-1].strip()
    if "/" in raw:
        raw = raw.split("/")[-1].strip()
    return raw or None


def parse_mt940_statement(path: Path) -> tuple[str, StatementBalances, list[StatementTxn]]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    if ":25:" not in text or ":60F:" not in text or ":62F:" not in text or ":61:" not in text:
        raise CashAnchorError(f"Statement is not MT940-like: {path.parent.name}")

    account_id: str | None = None
    opening: float | None = None
    opening_date: str | None = None
    closing: float | None = None
    closing_date: str | None = None
    txns: list[StatementTxn] = []
    current_txn: StatementTxn | None = None
    desc_parts: list[str] = []
    in_86 = False

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith(":25:"):
            account_id = _extract_account_id(line)
            in_86 = False
            continue
        if line.startswith(":60F:"):
            parsed = _parse_balance_line(line)
            if parsed:
                opening, opening_date = parsed
            in_86 = False
            continue
        if line.startswith(":62F:"):
            parsed = _parse_balance_line(line)
            if parsed:
                closing, closing_date = parsed
            in_86 = False
            continue
        if line.startswith(":61:"):
            if current_txn:
                txns.append(
                    StatementTxn(
                        value_date=current_txn.value_date,
                        dc_mark=current_txn.dc_mark,
                        amount_kzt=current_txn.amount_kzt,
                        code_86=current_txn.code_86,
                        description=" ".join(desc_parts).strip(),
                    )
                )
            current_txn = _parse_61_line(line)
            desc_parts = []
            in_86 = False
            continue
        if line.startswith(":86:"):
            if current_txn is None:
                continue
            code, desc = _parse_86_line(line)
            current_txn = StatementTxn(
                value_date=current_txn.value_date,
                dc_mark=current_txn.dc_mark,
                amount_kzt=current_txn.amount_kzt,
                code_86=code or current_txn.code_86,
                description=current_txn.description,
            )
            if desc:
                desc_parts.append(desc)
            in_86 = True
            continue
        if line.startswith(":"):
            in_86 = False
            continue
        if in_86 and current_txn:
            desc = _clean_86_text(line)
            if desc:
                desc_parts.append(desc)

    if current_txn:
        txns.append(
            StatementTxn(
                value_date=current_txn.value_date,
                dc_mark=current_txn.dc_mark,
                amount_kzt=current_txn.amount_kzt,
                code_86=current_txn.code_86,
                description=" ".join(desc_parts).strip(),
            )
        )

    if not account_id or opening is None or closing is None or not opening_date or not closing_date:
        raise CashAnchorError(f"Statement missing account or balance fields: {path.parent.name}")
    return account_id, StatementBalances(opening, closing, opening_date, closing_date), txns


def _parse_report_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_sales_report(path: Path, cutoff_date: date) -> SalesReportSummary:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in worksheet[7]]
        header_to_index = {header: index for index, header in enumerate(headers) if header}
        missing = [header for header in REQUIRED_SALES_HEADERS if header not in header_to_index]
        if missing:
            raise CashAnchorError(f"Sales report missing required columns for {path.parent.name}: {missing}")

        order_ref_index = header_to_index["Номер заказа (ID/RRN)"]
        accounting_date_index = header_to_index["Дата учета операции"]
        rows_total = 0
        rows_through_anchor = 0
        rows_after_cutoff = 0
        distinct_order_refs: set[str] = set()

        for row in worksheet.iter_rows(min_row=8, values_only=True):
            if not row or not any(value not in (None, "") for value in row):
                continue
            rows_total += 1
            row_date = _parse_report_date(row[accounting_date_index])
            if row_date is None:
                continue
            if row_date <= cutoff_date:
                rows_through_anchor += 1
                order_ref = row[order_ref_index]
                if order_ref not in (None, ""):
                    distinct_order_refs.add(str(order_ref).strip())
            else:
                rows_after_cutoff += 1
        return SalesReportSummary(
            rows_total=rows_total,
            rows_through_anchor=rows_through_anchor,
            rows_after_cutoff=rows_after_cutoff,
            distinct_order_refs_through_anchor=len(distinct_order_refs),
        )
    finally:
        workbook.close()


def _store_code(store_folder: str) -> str:
    return STORE_CODE_BY_FOLDER.get(store_folder, store_folder.upper())


def _one_file(store_dir: Path, suffix: str, label: str) -> Path:
    files = sorted(path for path in store_dir.iterdir() if path.is_file() and path.suffix.lower() == suffix)
    if len(files) != 1:
        raise CashAnchorError(f"{store_dir.name} must contain exactly one {label}; found {len(files)}")
    return files[0]


def _discover_store_dirs(source_root: Path, expected_stores: Iterable[str]) -> list[Path]:
    if not source_root.exists():
        raise FileNotFoundError(f"Source root not found: {source_root}")
    expected = list(expected_stores)
    present = sorted(path.name for path in source_root.iterdir() if path.is_dir())
    if set(present) != set(expected):
        missing = sorted(set(expected) - set(present))
        extra = sorted(set(present) - set(expected))
        raise CashAnchorError(f"Expected exactly stores {expected}; missing={missing}; extra={extra}")
    return [source_root / store for store in expected]


def _anchor_id(record: dict[str, Any], run_id: str) -> str:
    parts = [
        record["store_code"],
        record["account_identity_hash"],
        record["anchor_date"],
        record["statement_file_sha256"],
        record["sales_report_file_sha256"],
        run_id,
    ]
    return _sha256_text("|".join(parts))


def _build_record(source_root: Path, store_dir: Path, cutoff_date: date, run_id: str) -> dict[str, Any]:
    statement_path = _one_file(store_dir, ".txt", "statement")
    sales_report_path = _one_file(store_dir, ".xlsx", "sales report")
    account_id, balances, txns = parse_mt940_statement(statement_path)
    sales_summary = parse_sales_report(sales_report_path, cutoff_date)

    through_txns = [txn for txn in txns if date.fromisoformat(txn.value_date) <= cutoff_date]
    post_cutoff_txns = [txn for txn in txns if date.fromisoformat(txn.value_date) > cutoff_date]
    all_txn_sum = _round_money(sum(txn.amount_kzt for txn in txns))
    through_sum = _round_money(sum(txn.amount_kzt for txn in through_txns))
    post_cutoff_sum = _round_money(sum(txn.amount_kzt for txn in post_cutoff_txns))
    anchor_closing = _round_money(balances.opening + through_sum)
    statement_bridge_error = _round_money((balances.opening + all_txn_sum) - balances.closing)
    anchor_to_source_bridge_error = _round_money(anchor_closing + post_cutoff_sum - balances.closing)
    status = "RECONCILED" if statement_bridge_error == 0 and anchor_to_source_bridge_error == 0 else "BLOCKED"

    record: dict[str, Any] = {
        "anchor_id": "",
        "source_root": str(source_root),
        "source_store_dir": str(store_dir),
        "source_store_name": store_dir.name,
        "store_code": _store_code(store_dir.name),
        "account_identity_hash": _sha256_text(account_id),
        "account_mask": _mask_account(account_id),
        "statement_file_sha256": _sha256_file(statement_path),
        "sales_report_file_sha256": _sha256_file(sales_report_path),
        "source_opening_date": balances.opening_date,
        "source_opening_balance_kzt": _round_money(balances.opening),
        "anchor_date": cutoff_date.isoformat(),
        "anchor_closing_balance_kzt": anchor_closing,
        "source_statement_closing_date": balances.closing_date,
        "source_statement_closing_balance_kzt": _round_money(balances.closing),
        "post_cutoff_txn_count": len(post_cutoff_txns),
        "post_cutoff_txn_sum_kzt": post_cutoff_sum,
        "statement_txn_count": len(txns),
        "statement_txn_count_through_anchor": len(through_txns),
        "sales_report_rows_total": sales_summary.rows_total,
        "sales_report_rows_through_anchor": sales_summary.rows_through_anchor,
        "sales_report_rows_after_cutoff": sales_summary.rows_after_cutoff,
        "sales_report_distinct_order_refs": sales_summary.distinct_order_refs_through_anchor,
        "statement_bridge_error_kzt": statement_bridge_error,
        "anchor_to_source_bridge_error_kzt": anchor_to_source_bridge_error,
        "partial_day_excluded": bool(post_cutoff_txns or sales_summary.rows_after_cutoff or balances.closing_date > cutoff_date.isoformat()),
        "reconciliation_status": status,
        "trust_class": "ACTUAL_ANCHOR",
        "notes_redacted": "One-time Kaspi Pay cash anchor; no raw account IDs, order refs, cards, or transaction text retained.",
        "created_by_run_id": run_id,
    }
    record["anchor_id"] = _anchor_id(record, run_id)
    return record


def _redaction_scan_text(summary: dict[str, Any]) -> list[str]:
    text = json.dumps(summary, ensure_ascii=False, sort_keys=True)
    patterns = [
        r"KZ[0-9A-Z]{10,}",
        r"\b\d{16,19}\b",
        r"ИИН",
        r"БИН",
        r"Номер карты",
        r"Детали покупки",
    ]
    return [pattern for pattern in patterns if re.search(pattern, text)]


def _write_artifacts(summary: dict[str, Any], output_root: Path) -> None:
    output_root.mkdir(parents=True, exist_ok=True)
    (output_root / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=True, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )

    csv_columns = [
        "store_code",
        "account_mask",
        "anchor_date",
        "anchor_closing_balance_kzt",
        "source_statement_closing_date",
        "source_statement_closing_balance_kzt",
        "post_cutoff_txn_count",
        "post_cutoff_txn_sum_kzt",
        "statement_txn_count_through_anchor",
        "sales_report_rows_through_anchor",
        "sales_report_rows_after_cutoff",
        "statement_bridge_error_kzt",
        "anchor_to_source_bridge_error_kzt",
        "reconciliation_status",
        "trust_class",
    ]
    with (output_root / "anchor_preview.csv").open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=csv_columns)
        writer.writeheader()
        for record in summary["records"]:
            writer.writerow({column: record.get(column, "") for column in csv_columns})

    lines = [
        "# Kaspi Pay Cash Anchor Preview",
        "",
        f"Status: {summary['status']}",
        f"Cutoff: {summary['cutoff']}",
        f"Run id: {summary['run_id']}",
        "",
        "| store | anchor balance | excluded txns | excluded txn sum | report rows through cutoff | report rows after cutoff | recon |",
        "|---|---:|---:|---:|---:|---:|---|",
    ]
    for record in summary["records"]:
        lines.append(
            "| {store_code} | {anchor_closing_balance_kzt:.2f} | {post_cutoff_txn_count} | "
            "{post_cutoff_txn_sum_kzt:.2f} | {sales_report_rows_through_anchor} | "
            "{sales_report_rows_after_cutoff} | {reconciliation_status} |".format(**record)
        )
    (output_root / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_cash_anchor_preview(
    *,
    source_root: Path,
    cutoff: str,
    output_root: Path | None = None,
    strict: bool = False,
    redact: bool = True,
    run_id: str = "cash-anchor-preview",
    expected_stores: Iterable[str] = DEFAULT_EXPECTED_STORES,
) -> dict[str, Any]:
    cutoff_date = date.fromisoformat(cutoff)
    store_dirs = _discover_store_dirs(source_root, expected_stores)
    records = [_build_record(source_root, store_dir, cutoff_date, run_id) for store_dir in store_dirs]
    errors: list[str] = []
    for record in records:
        if record["statement_bridge_error_kzt"] != 0:
            errors.append(f"{record['store_code']} statement bridge error {record['statement_bridge_error_kzt']}")
        if record["anchor_to_source_bridge_error_kzt"] != 0:
            errors.append(f"{record['store_code']} anchor bridge error {record['anchor_to_source_bridge_error_kzt']}")
        if record["reconciliation_status"] != "RECONCILED":
            errors.append(f"{record['store_code']} reconciliation status {record['reconciliation_status']}")

    summary: dict[str, Any] = {
        "status": "FAIL" if errors else "PASS",
        "run_id": run_id,
        "source_root": str(source_root),
        "cutoff": cutoff_date.isoformat(),
        "store_count": len(records),
        "expected_stores": list(expected_stores),
        "redacted": bool(redact),
        "records": records,
        "errors": errors,
    }
    redaction_hits = _redaction_scan_text(summary) if redact else []
    summary["redaction_scan"] = {"status": "PASS" if not redaction_hits else "FAIL", "matches": redaction_hits}
    if redaction_hits:
        errors.extend([f"redaction scan matched {pattern}" for pattern in redaction_hits])
        summary["status"] = "FAIL"

    if output_root is not None:
        _write_artifacts(summary, output_root)
    if strict and errors:
        raise CashAnchorError("; ".join(errors))
    return summary


CASH_ANCHOR_TABLE_SQL = """
CREATE TABLE IF NOT EXISTS cashflow_cash_anchor (
    anchor_id TEXT PRIMARY KEY,
    source_root TEXT NOT NULL,
    source_store_dir TEXT NOT NULL,
    source_store_name TEXT NOT NULL,
    store_code TEXT NOT NULL,
    account_identity_hash TEXT NOT NULL,
    account_mask TEXT NOT NULL,
    statement_file_sha256 TEXT NOT NULL,
    sales_report_file_sha256 TEXT NOT NULL,
    source_opening_date TEXT NOT NULL,
    source_opening_balance_kzt REAL NOT NULL,
    anchor_date TEXT NOT NULL,
    anchor_closing_balance_kzt REAL NOT NULL,
    source_statement_closing_date TEXT NOT NULL,
    source_statement_closing_balance_kzt REAL NOT NULL,
    post_cutoff_txn_count INTEGER NOT NULL,
    post_cutoff_txn_sum_kzt REAL NOT NULL,
    statement_txn_count INTEGER NOT NULL,
    statement_txn_count_through_anchor INTEGER NOT NULL,
    sales_report_rows_total INTEGER NOT NULL,
    sales_report_rows_through_anchor INTEGER NOT NULL,
    sales_report_rows_after_cutoff INTEGER NOT NULL,
    sales_report_distinct_order_refs INTEGER NOT NULL,
    statement_bridge_error_kzt REAL NOT NULL,
    anchor_to_source_bridge_error_kzt REAL NOT NULL,
    partial_day_excluded INTEGER NOT NULL,
    reconciliation_status TEXT NOT NULL,
    trust_class TEXT NOT NULL,
    notes_redacted TEXT,
    created_by_run_id TEXT NOT NULL,
    created_at TEXT DEFAULT (datetime('now')),
    UNIQUE (
        store_code,
        account_identity_hash,
        anchor_date,
        statement_file_sha256,
        sales_report_file_sha256,
        created_by_run_id
    )
)
"""


def ensure_cash_anchor_schema(conn: sqlite3.Connection) -> None:
    conn.execute(CASH_ANCHOR_TABLE_SQL)
    conn.execute("CREATE INDEX IF NOT EXISTS idx_cashflow_cash_anchor_store_date ON cashflow_cash_anchor (store_code, anchor_date)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_cashflow_cash_anchor_run ON cashflow_cash_anchor (created_by_run_id)")


def _table_exists(conn: sqlite3.Connection, table: str) -> bool:
    return conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name=?",
        (table,),
    ).fetchone() is not None


def _existing_anchor_ids(conn: sqlite3.Connection, anchor_ids: list[str]) -> set[str]:
    if not anchor_ids or not _table_exists(conn, "cashflow_cash_anchor"):
        return set()
    placeholders = ",".join("?" * len(anchor_ids))
    return {
        row[0]
        for row in conn.execute(
            f"SELECT anchor_id FROM cashflow_cash_anchor WHERE anchor_id IN ({placeholders})",
            anchor_ids,
        ).fetchall()
    }


def _insert_anchor(conn: sqlite3.Connection, record: dict[str, Any]) -> None:
    conn.execute(
        """
        INSERT OR IGNORE INTO cashflow_cash_anchor (
            anchor_id, source_root, source_store_dir, source_store_name, store_code,
            account_identity_hash, account_mask, statement_file_sha256, sales_report_file_sha256,
            source_opening_date, source_opening_balance_kzt, anchor_date, anchor_closing_balance_kzt,
            source_statement_closing_date, source_statement_closing_balance_kzt,
            post_cutoff_txn_count, post_cutoff_txn_sum_kzt, statement_txn_count,
            statement_txn_count_through_anchor, sales_report_rows_total,
            sales_report_rows_through_anchor, sales_report_rows_after_cutoff,
            sales_report_distinct_order_refs, statement_bridge_error_kzt,
            anchor_to_source_bridge_error_kzt, partial_day_excluded, reconciliation_status,
            trust_class, notes_redacted, created_by_run_id
        ) VALUES (
            :anchor_id, :source_root, :source_store_dir, :source_store_name, :store_code,
            :account_identity_hash, :account_mask, :statement_file_sha256, :sales_report_file_sha256,
            :source_opening_date, :source_opening_balance_kzt, :anchor_date, :anchor_closing_balance_kzt,
            :source_statement_closing_date, :source_statement_closing_balance_kzt,
            :post_cutoff_txn_count, :post_cutoff_txn_sum_kzt, :statement_txn_count,
            :statement_txn_count_through_anchor, :sales_report_rows_total,
            :sales_report_rows_through_anchor, :sales_report_rows_after_cutoff,
            :sales_report_distinct_order_refs, :statement_bridge_error_kzt,
            :anchor_to_source_bridge_error_kzt, :partial_day_excluded, :reconciliation_status,
            :trust_class, :notes_redacted, :created_by_run_id
        )
        """,
        {**record, "partial_day_excluded": int(bool(record["partial_day_excluded"]))},
    )


def apply_cash_anchor(
    *,
    db_path: Path,
    source_root: Path,
    cutoff: str,
    run_id: str,
    output_root: Path,
    strict: bool = False,
    redact: bool = True,
    apply: bool = False,
    expected_stores: Iterable[str] = DEFAULT_EXPECTED_STORES,
) -> dict[str, Any]:
    if not db_path.exists():
        raise FileNotFoundError(f"DB not found: {db_path}")

    summary = build_cash_anchor_preview(
        source_root=source_root,
        cutoff=cutoff,
        output_root=None,
        strict=strict,
        redact=redact,
        run_id=run_id,
        expected_stores=expected_stores,
    )
    anchor_ids = [record["anchor_id"] for record in summary["records"]]
    inserted = 0
    with sqlite3.connect(str(db_path)) as conn:
        existing_ids = _existing_anchor_ids(conn, anchor_ids)
        to_insert = [record for record in summary["records"] if record["anchor_id"] not in existing_ids]
        if apply:
            if os.environ.get("ENABLE_CASHFLOW_ANCHOR_WRITE") != "1":
                raise RuntimeError("ENABLE_CASHFLOW_ANCHOR_WRITE=1 is required to apply cash anchor writes.")
            ensure_cash_anchor_schema(conn)
            for record in to_insert:
                before = conn.total_changes
                _insert_anchor(conn, record)
                inserted += conn.total_changes - before
            conn.commit()

    summary["apply"] = {
        "applied": bool(apply),
        "candidate_anchor_records": len(summary["records"]),
        "existing_anchor_records": len(existing_ids),
        "would_insert_anchor_records": len(to_insert),
        "inserted_anchor_records": inserted,
        "db_path": str(db_path),
        "cashflow_events_created": 0,
    }
    _write_artifacts(summary, output_root)
    if strict and summary["status"] != "PASS":
        raise CashAnchorError("; ".join(summary.get("errors") or ["cash anchor preview failed"]))
    return summary
