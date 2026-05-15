"""Approved stock anchor and one-time 20 percent baseline adjustment workflow."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import sqlite3
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_FLOOR, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


STOCK_ANCHOR_REFERENCE_TYPE = "STOCK_ANCHOR"
BASELINE_REFERENCE_TYPE = "BASELINE_20PCT_DECREASE"
ROLLBACK_REFERENCE_TYPE = "ROLLBACK_BASELINE_20PCT_DECREASE"
METHOD_VERSION = "stock_anchor_20pct_v1"
CANONICAL_STORE_CODE = "UNIVERSAL"

OWNER_OVERRIDE_NO_DOUBLE_REDUCE_SKU_KEYS = frozenset(
    {
        "CL_OC_MEN_LINE51_WHITE",
        "CL_NEW-CLO2_MEN_SUIT-61_BLACK",
    }
)

OWNER_OOS_ACTIVE_ZERO_KEYS = frozenset(
    {
        ("CL_OC_MEN_LINE52_BLACK", "4XL"),
        ("CL_NEW-CLO_MEN_T-SHIRT_BLACK", "S"),
        ("CL_NEW-CLO_MEN_T-SHIRT_BLACK", "M"),
        ("CL_NEW-CLO_MEN_T-SHIRT_BLACK", "L"),
    }
)

LINE61_SKU_KEY = "CL_NEW-CLO2_MEN_SUIT-61_BLACK"


@dataclass(frozen=True)
class AnchorStockRow:
    sku_key: str
    sku_id: str
    my_size: str
    qty: int
    model: str | None = None
    color: str | None = None
    product_type: str | None = None


@dataclass(frozen=True)
class AdjustmentRow:
    sku_key: str
    sku_id: str
    my_size: str
    original_qty: int
    exact_reduction: str
    fractional_remainder: str
    reduction_qty: int
    post_adjustment_qty: int
    allocation_rank: int | None = None
    excluded_reason: str | None = None


@dataclass(frozen=True)
class AdjustmentPlan:
    reduction_rate: str
    target_reduction_units: int
    generated_reduction_units: int
    eligible_units: int
    rows: list[AdjustmentRow]


@dataclass(frozen=True)
class ApplyEventsResult:
    anchor_inserted: bool
    anchor_events_inserted: int
    adjustment_events_inserted: int
    already_applied: bool


@dataclass(frozen=True)
class SnapshotRow:
    sku_key: str
    sku_id: str
    my_size: str
    current_stock: int
    inbound_stock: int = 0
    raw_current_stock: int | None = None


@dataclass(frozen=True)
class SnapshotBuildResult:
    snapshot_date: str
    rows: list[SnapshotRow]
    negative_rows: list[SnapshotRow]
    raw_total_units: int
    current_stock_total: int
    inbound_stock_total: int


@dataclass(frozen=True)
class ExceptionSpec:
    exception_id: str
    domain: str
    severity: str
    reason_code: str
    reason: str
    evidence: dict


@dataclass(frozen=True)
class OfferAvailabilityRow:
    snapshot_id: str
    snapshot_date: str
    store_code: str
    sku_key: str
    sku_id: str
    my_size: str
    offer_available_qty: int | None
    physical_stock_qty: int
    status: str
    source_manifest_id: str


def _normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def _slug(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", value.upper()).strip("_")
    return slug[:120] or "UNKNOWN"


def _iso_date(value: str | date) -> str:
    return value.isoformat() if isinstance(value, date) else str(value)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def deterministic_anchor_id(anchor_date: str, source_sha256: str) -> str:
    return f"PHYSICAL_STOCK_ANCHOR_{anchor_date.replace('-', '')}_{source_sha256[:12].upper()}"


def deterministic_batch_id(anchor_date: str, source_sha256: str) -> str:
    return f"BASELINE_20PCT_DECREASE_{anchor_date.replace('-', '')}_{source_sha256[:12].upper()}"


def anchor_event_date_for(anchor_date: str | date) -> str:
    parsed = date.fromisoformat(_iso_date(anchor_date))
    return (parsed - timedelta(days=1)).isoformat()


def load_anchor_rows_from_workbook(path: Path, sheet_name: str) -> list[AnchorStockRow]:
    """Read positive physical stock rows from the approved anchor workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Anchor workbook missing sheet: {sheet_name}")

    worksheet = workbook[sheet_name]
    iterator = worksheet.iter_rows(values_only=True)
    try:
        headers = [_normalize_text(cell) for cell in next(iterator)]
    except StopIteration as exc:
        raise ValueError("Anchor workbook sheet is empty") from exc

    required = {"sku_key", "model", "color", "product_type", "total_stock"}
    missing = sorted(required.difference(headers))
    if missing:
        raise ValueError(f"Anchor workbook missing required columns: {missing}")

    size_columns: list[tuple[int, str]] = []
    for index, header in enumerate(headers):
        if header.startswith("Size_"):
            size = header.removeprefix("Size_")
            if size:
                size_columns.append((index, size))
    if not size_columns:
        raise ValueError("Anchor workbook has no Size_* columns")

    index_by_header = {header: idx for idx, header in enumerate(headers)}
    rows: list[AnchorStockRow] = []
    seen: set[str] = set()

    for raw_row in iterator:
        if not raw_row:
            continue
        sku_key = _normalize_text(raw_row[index_by_header["sku_key"]])
        if not sku_key:
            continue
        model = _normalize_text(raw_row[index_by_header["model"]]) or None
        color = _normalize_text(raw_row[index_by_header["color"]]) or None
        product_type = _normalize_text(raw_row[index_by_header["product_type"]]) or None
        for index, size in size_columns:
            qty = int(raw_row[index] or 0)
            if qty <= 0:
                continue
            sku_id = f"{sku_key}_{size}"
            if sku_id in seen:
                raise ValueError(f"Duplicate anchor SKU-size row: {sku_id}")
            seen.add(sku_id)
            rows.append(
                AnchorStockRow(
                    sku_key=sku_key,
                    sku_id=sku_id,
                    my_size=size,
                    qty=qty,
                    model=model,
                    color=color,
                    product_type=product_type,
                )
            )

    if not rows:
        raise ValueError("Anchor workbook has no positive stock rows")
    return rows


def compute_adjustment_plan(
    rows: Iterable[AnchorStockRow],
    *,
    reduction_rate: Decimal = Decimal("0.20"),
    excluded_sku_keys: frozenset[str] = OWNER_OVERRIDE_NO_DOUBLE_REDUCE_SKU_KEYS,
) -> AdjustmentPlan:
    """Compute deterministic 20 percent reductions by largest fractional remainder."""
    source_rows = list(rows)
    reductions: dict[str, int] = {}
    exact_by_sku: dict[str, Decimal] = {}
    fraction_by_sku: dict[str, Decimal] = {}
    excluded: dict[str, str] = {}
    eligible: list[AnchorStockRow] = []

    for row in source_rows:
        reductions[row.sku_id] = 0
        if row.qty <= 0:
            excluded[row.sku_id] = "NON_POSITIVE_STOCK"
            continue
        if row.sku_key in excluded_sku_keys:
            excluded[row.sku_id] = "OWNER_OVERRIDE_NO_DOUBLE_REDUCE"
            continue
        exact = Decimal(row.qty) * reduction_rate
        floored = int(exact.to_integral_value(rounding=ROUND_FLOOR))
        reductions[row.sku_id] = min(floored, row.qty)
        exact_by_sku[row.sku_id] = exact
        fraction_by_sku[row.sku_id] = exact - Decimal(floored)
        eligible.append(row)

    eligible_units = sum(row.qty for row in eligible)
    target = int((Decimal(eligible_units) * reduction_rate).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    remaining = target - sum(reductions[row.sku_id] for row in eligible)

    ranked = sorted(
        eligible,
        key=lambda row: (-fraction_by_sku[row.sku_id], row.sku_key, row.my_size, row.sku_id),
    )
    allocation_rank_by_sku: dict[str, int] = {}
    for rank, row in enumerate(ranked, start=1):
        if remaining <= 0:
            break
        if reductions[row.sku_id] < row.qty:
            reductions[row.sku_id] += 1
            allocation_rank_by_sku[row.sku_id] = rank
            remaining -= 1

    plan_rows: list[AdjustmentRow] = []
    for row in sorted(source_rows, key=lambda item: (item.sku_key, item.my_size, item.sku_id)):
        reduction = reductions[row.sku_id]
        exact = exact_by_sku.get(row.sku_id, Decimal("0"))
        fraction = fraction_by_sku.get(row.sku_id, Decimal("0"))
        plan_rows.append(
            AdjustmentRow(
                sku_key=row.sku_key,
                sku_id=row.sku_id,
                my_size=row.my_size,
                original_qty=row.qty,
                exact_reduction=str(exact),
                fractional_remainder=str(fraction),
                reduction_qty=reduction,
                post_adjustment_qty=max(0, row.qty - reduction),
                allocation_rank=allocation_rank_by_sku.get(row.sku_id),
                excluded_reason=excluded.get(row.sku_id),
            )
        )

    generated = sum(row.reduction_qty for row in plan_rows)
    return AdjustmentPlan(
        reduction_rate=str(reduction_rate),
        target_reduction_units=target,
        generated_reduction_units=generated,
        eligible_units=eligible_units,
        rows=plan_rows,
    )


def _table_exists(conn: sqlite3.Connection, table: str) -> bool:
    row = conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name=?",
        (table,),
    ).fetchone()
    return row is not None


def _columns(conn: sqlite3.Connection, table: str) -> set[str]:
    return {row[1] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()}


def apply_anchor_and_adjustment_events(
    conn: sqlite3.Connection,
    *,
    anchor_id: str,
    batch_id: str,
    anchor_rows: list[AnchorStockRow],
    adjustment_plan: AdjustmentPlan,
    source_path: str,
    source_sha256: str,
    anchor_snapshot_date: str,
    anchor_event_date: str,
    adjustment_event_date: str,
    approved_by: str,
    approved_at: str,
    dry_run_report_path: str,
) -> ApplyEventsResult:
    """Append immutable anchor and baseline adjustment events with idempotency keys."""
    existing_anchor = conn.execute(
        "SELECT * FROM stock_anchor WHERE anchor_id=?",
        (anchor_id,),
    ).fetchone()
    anchor_inserted = False
    if existing_anchor is not None:
        if (
            existing_anchor["source_sha256"] != source_sha256
            or existing_anchor["snapshot_date"] != anchor_snapshot_date
            or existing_anchor["status"] != "APPROVED"
        ):
            raise RuntimeError(f"Existing stock_anchor {anchor_id} does not match approved artifact")
    else:
        conn.execute(
            """
            INSERT INTO stock_anchor (
                anchor_id, anchor_type, source_path, source_sha256, snapshot_date,
                as_of_date, row_count, total_units, approved_by, approved_at, status, notes
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                anchor_id,
                "PHYSICAL_WAREHOUSE_STOCK",
                source_path,
                source_sha256,
                anchor_snapshot_date,
                anchor_snapshot_date,
                len(anchor_rows),
                sum(row.qty for row in anchor_rows),
                approved_by,
                approved_at,
                "APPROVED",
                "Immutable approved physical stock anchor. Do not update or delete; use reversal/new anchor.",
            ),
        )
        anchor_inserted = True

    existing_batch = conn.execute(
        "SELECT * FROM stock_adjustment_batch WHERE batch_id=?",
        (batch_id,),
    ).fetchone()
    if existing_batch is not None:
        if existing_batch["anchor_id"] != anchor_id or existing_batch["status"] != "APPLIED":
            raise RuntimeError(f"Existing stock_adjustment_batch {batch_id} is not an applied match")
        return ApplyEventsResult(
            anchor_inserted=anchor_inserted,
            anchor_events_inserted=0,
            adjustment_events_inserted=0,
            already_applied=True,
        )

    anchor_inserted_count = 0
    for row in anchor_rows:
        idempotency_key = f"{STOCK_ANCHOR_REFERENCE_TYPE}:{anchor_id}:{row.sku_id}"
        cursor = conn.execute(
            """
            INSERT OR IGNORE INTO stock_ledger (
                event_date, event_type, sku_key, sku_id, my_size, store_code,
                qty_change, reference_id, reference_type, notes, input_source,
                created_by, idempotency_key
            ) VALUES (?, 'INITIAL', ?, ?, ?, ?, ?, ?, ?, ?, 'OWNER_APPROVED_ANCHOR', ?, ?)
            """,
            (
                anchor_event_date,
                row.sku_key,
                row.sku_id,
                row.my_size,
                CANONICAL_STORE_CODE,
                row.qty,
                anchor_id,
                STOCK_ANCHOR_REFERENCE_TYPE,
                f"Approved physical stock anchor snapshot_date={anchor_snapshot_date}",
                "agent6_stock_anchor_20pct",
                idempotency_key,
            ),
        )
        anchor_inserted_count += cursor.rowcount

    adjustment_inserted_count = 0
    for row in adjustment_plan.rows:
        if row.reduction_qty <= 0:
            continue
        idempotency_key = f"{BASELINE_REFERENCE_TYPE}:{batch_id}:{row.sku_id}"
        cursor = conn.execute(
            """
            INSERT OR IGNORE INTO stock_ledger (
                event_date, event_type, sku_key, sku_id, my_size, store_code,
                qty_change, reference_id, reference_type, notes, input_source,
                created_by, idempotency_key
            ) VALUES (?, 'ADJUSTMENT', ?, ?, ?, ?, ?, ?, ?, ?, 'OWNER_APPROVED_BASELINE_20PCT', ?, ?)
            """,
            (
                adjustment_event_date,
                row.sku_key,
                row.sku_id,
                row.my_size,
                CANONICAL_STORE_CODE,
                -row.reduction_qty,
                batch_id,
                BASELINE_REFERENCE_TYPE,
                "One-time approved baseline 20 percent decrease after anchor; append-only event.",
                "agent6_stock_anchor_20pct",
                idempotency_key,
            ),
        )
        adjustment_inserted_count += cursor.rowcount

    conn.execute(
        """
        INSERT INTO stock_adjustment_batch (
            batch_id, anchor_id, method, method_version, reduction_rate,
            target_units_delta, generated_units_delta, dry_run_report_path,
            approved_by, approved_at, applied_at, status, notes
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, datetime('now'), 'APPLIED', ?)
        """,
        (
            batch_id,
            anchor_id,
            BASELINE_REFERENCE_TYPE,
            METHOD_VERSION,
            float(Decimal(adjustment_plan.reduction_rate)),
            -adjustment_plan.target_reduction_units,
            -adjustment_plan.generated_reduction_units,
            dry_run_report_path,
            approved_by,
            approved_at,
            "LINE51 and Line61 owner-approved override families excluded to avoid double reduction.",
        ),
    )

    return ApplyEventsResult(
        anchor_inserted=anchor_inserted,
        anchor_events_inserted=anchor_inserted_count,
        adjustment_events_inserted=adjustment_inserted_count,
        already_applied=False,
    )


def _pending_inbound_by_sku(conn: sqlite3.Connection, snapshot_date: str) -> dict[str, int]:
    inbound_by_sku: dict[str, int] = {}
    if _table_exists(conn, "po_line") and _table_exists(conn, "po_header"):
        rows = conn.execute(
            """
            SELECT pl.sku_id, SUM(pl.order_qty - COALESCE(pl.received_qty, 0)) AS inbound_stock
            FROM po_line pl
            JOIN po_header ph ON pl.po_id = ph.po_id
            WHERE pl.status IN ('PENDING', 'PARTIAL', 'IN_TRANSIT')
              AND ph.status NOT IN ('CLOSED', 'CANCELLED')
            GROUP BY pl.sku_id
            """
        ).fetchall()
        inbound_by_sku.update({row["sku_id"]: row["inbound_stock"] or 0 for row in rows})

    if _table_exists(conn, "fact_po_lines"):
        rows = conn.execute(
            """
            SELECT sku_id, SUM(order_quantity - received_qty) AS inbound_stock
            FROM fact_po_lines
            WHERE (order_quantity - received_qty) > 0
              AND (est_arrival_date IS NULL OR est_arrival_date >= ?)
              AND status NOT IN ('ARRIVED', 'CLOSED', 'CANCELLED', 'RECEIVED')
            GROUP BY sku_id
            """,
            (snapshot_date,),
        ).fetchall()
        for row in rows:
            inbound_by_sku[row["sku_id"]] = inbound_by_sku.get(row["sku_id"], 0) + (row["inbound_stock"] or 0)
    return inbound_by_sku


def _active_size_rows(conn: sqlite3.Connection) -> dict[str, tuple[str, str]]:
    if not _table_exists(conn, "dim_sku_size"):
        return {}
    size_cols = _columns(conn, "dim_sku_size")
    if "active_flag" in size_cols and _table_exists(conn, "dim_sku") and "active_flag" in _columns(conn, "dim_sku"):
        rows = conn.execute(
            """
            SELECT ds.sku_id, ds.sku_key, ds.my_size
            FROM dim_sku_size ds
            JOIN dim_sku d ON ds.sku_key = d.sku_key
            WHERE ds.active_flag = 1
              AND d.active_flag = 1
              AND ds.my_size IS NOT NULL
              AND ds.my_size != ''
            """
        ).fetchall()
    else:
        rows = conn.execute(
            """
            SELECT sku_id, sku_key, my_size
            FROM dim_sku_size
            WHERE my_size IS NOT NULL
              AND my_size != ''
            """
        ).fetchall()
    return {row["sku_id"]: (row["sku_key"], row["my_size"]) for row in rows}


def build_anchor_lineage_snapshot(
    conn: sqlite3.Connection,
    *,
    anchor_id: str,
    batch_id: str,
    anchor_date: str,
    snapshot_date: str,
    include_active_sizes: bool = False,
) -> SnapshotBuildResult:
    """Build a physical-stock snapshot from anchor lineage, excluding pre-anchor legacy events."""
    rows = conn.execute(
        """
        SELECT sku_id, sku_key, my_size, SUM(qty_change) AS qty
        FROM stock_ledger
        WHERE event_date < ?
          AND (
            (reference_type = ? AND reference_id = ?)
            OR (reference_type = ? AND reference_id = ?)
            OR (
                event_date > ?
                AND COALESCE(reference_type, '') NOT IN (?, ?, ?)
            )
          )
        GROUP BY sku_id, sku_key, my_size
        """,
        (
            snapshot_date,
            STOCK_ANCHOR_REFERENCE_TYPE,
            anchor_id,
            BASELINE_REFERENCE_TYPE,
            batch_id,
            anchor_date,
            STOCK_ANCHOR_REFERENCE_TYPE,
            BASELINE_REFERENCE_TYPE,
            ROLLBACK_REFERENCE_TYPE,
        ),
    ).fetchall()

    base: dict[str, tuple[str, str, int]] = {
        row["sku_id"]: (row["sku_key"], row["my_size"], row["qty"] or 0) for row in rows
    }
    if include_active_sizes:
        for sku_id, (sku_key, my_size) in _active_size_rows(conn).items():
            base.setdefault(sku_id, (sku_key, my_size, 0))

    inbound_by_sku = _pending_inbound_by_sku(conn, snapshot_date)
    for sku_id in inbound_by_sku:
        base.setdefault(sku_id, (sku_id.rsplit("_", 1)[0], sku_id.rsplit("_", 1)[-1], 0))

    snapshot_rows: list[SnapshotRow] = []
    negative_rows: list[SnapshotRow] = []
    raw_total = 0
    for sku_id, (sku_key, my_size, raw_stock) in sorted(base.items()):
        raw_total += raw_stock
        current_stock = max(0, int(raw_stock))
        row = SnapshotRow(
            sku_key=sku_key,
            sku_id=sku_id,
            my_size=my_size,
            current_stock=current_stock,
            inbound_stock=max(0, int(inbound_by_sku.get(sku_id, 0))),
            raw_current_stock=int(raw_stock),
        )
        snapshot_rows.append(row)
        if raw_stock < 0:
            negative_rows.append(row)

    return SnapshotBuildResult(
        snapshot_date=snapshot_date,
        rows=snapshot_rows,
        negative_rows=negative_rows,
        raw_total_units=raw_total,
        current_stock_total=sum(row.current_stock for row in snapshot_rows),
        inbound_stock_total=sum(row.inbound_stock for row in snapshot_rows),
    )


def build_high_risk_exception_specs(
    anchor_rows: Iterable[AnchorStockRow],
    *,
    negative_rows: Iterable[SnapshotRow],
    run_id: str = "AGENT6_STOCK_REBUILD_20260503",
) -> list[ExceptionSpec]:
    specs: list[ExceptionSpec] = []
    seen: set[str] = set()
    rows = list(anchor_rows)

    def add(reason_code: str, severity: str, reason: str, evidence: dict) -> None:
        key = f"{reason_code}:{evidence.get('sku_id') or evidence.get('sku_key')}"
        if key in seen:
            return
        seen.add(key)
        specs.append(
            ExceptionSpec(
                exception_id=f"{run_id}:{reason_code}:{_slug(str(evidence.get('sku_id') or evidence.get('sku_key')))}",
                domain="STOCK",
                severity=severity,
                reason_code=reason_code,
                reason=reason,
                evidence=evidence,
            )
        )

    for row in rows:
        if (row.sku_key, row.my_size) in OWNER_OOS_ACTIVE_ZERO_KEYS:
            add(
                "OWNER_OOS_ACTIVE_ZERO",
                "HIGH",
                "Owner says this SKU-size is out of stock for active sellable use until reverified.",
                {"sku_key": row.sku_key, "sku_id": row.sku_id, "my_size": row.my_size, "physical_anchor_qty": row.qty},
            )
        if row.sku_key in OWNER_OVERRIDE_NO_DOUBLE_REDUCE_SKU_KEYS:
            add(
                "OWNER_OVERRIDE_NO_DOUBLE_REDUCE",
                "HIGH",
                "Owner-approved LINE51/Line61 override remains active; do not double-reduce by the global batch.",
                {"sku_key": row.sku_key},
            )
        if row.sku_key == LINE61_SKU_KEY and row.my_size == "4XL":
            add(
                "LINE61_4XL_EXCLUDED",
                "HIGH",
                "Line61 4XL remains zero/excluded until explicitly superseded.",
                {"sku_key": row.sku_key, "sku_id": row.sku_id, "my_size": row.my_size, "physical_anchor_qty": row.qty},
            )

    by_lower: dict[str, set[str]] = {}
    for row in rows:
        by_lower.setdefault(row.sku_key.lower(), set()).add(row.sku_key)
    for lower_key, variants in by_lower.items():
        if len(variants) > 1:
            add(
                "MIXED_CASE_ALIAS_QUARANTINE",
                "MEDIUM",
                "Mixed-case SKU aliases require explicit canonicalization before family-level conclusions.",
                {"sku_key": lower_key, "variants": sorted(variants)},
            )

    for row in negative_rows:
        add(
            "NEGATIVE_RAW_LEDGER_BALANCE",
            "HIGH",
            "Anchor-lineage replay produced a raw negative balance; persisted physical stock is clamped to zero and output is provisional.",
            {
                "sku_key": row.sku_key,
                "sku_id": row.sku_id,
                "my_size": row.my_size,
                "raw_current_stock": row.raw_current_stock,
            },
        )

    return specs


def _as_snapshot_tuple(row: SnapshotRow | tuple[str, str, str, int]) -> tuple[str, str, str, int]:
    if isinstance(row, SnapshotRow):
        return row.sku_key, row.sku_id, row.my_size, row.current_stock
    return row


def build_offer_availability_rows(
    snapshot_rows: Iterable[SnapshotRow | tuple[str, str, str, int]],
    exceptions: Iterable[ExceptionSpec],
    *,
    snapshot_date: str,
    source_manifest_id: str,
) -> list[OfferAvailabilityRow]:
    _ = list(exceptions)
    availability: list[OfferAvailabilityRow] = []
    for raw_row in snapshot_rows:
        sku_key, sku_id, my_size, physical_qty = _as_snapshot_tuple(raw_row)
        status = "OFFER_AVAILABILITY_UNKNOWN"
        offer_qty: int | None = None
        if (sku_key, my_size) in OWNER_OOS_ACTIVE_ZERO_KEYS:
            status = "OWNER_OOS_EXCEPTION"
            offer_qty = 0
        elif sku_key == LINE61_SKU_KEY and my_size == "4XL":
            status = "OWNER_OVERRIDE_EXCLUDED"
            offer_qty = 0
        elif sku_key in OWNER_OVERRIDE_NO_DOUBLE_REDUCE_SKU_KEYS:
            status = "OWNER_OVERRIDE_ACTIVE_EVIDENCE"

        availability.append(
            OfferAvailabilityRow(
                snapshot_id=f"OFFER_AVAIL_{snapshot_date}_{_slug(sku_id)}",
                snapshot_date=snapshot_date,
                store_code=CANONICAL_STORE_CODE,
                sku_key=sku_key,
                sku_id=sku_id,
                my_size=my_size,
                offer_available_qty=offer_qty,
                physical_stock_qty=physical_qty,
                status=status,
                source_manifest_id=source_manifest_id,
            )
        )
    return availability


def insert_source_manifest(
    conn: sqlite3.Connection,
    *,
    source_id: str,
    source_type: str,
    source_path: str,
    source_sha256: str,
    as_of_date: str,
    row_count: int,
    notes: str,
) -> None:
    conn.execute(
        """
        INSERT OR REPLACE INTO source_manifest (
            source_id, source_type, source_path, source_sha256, as_of_date,
            freshness_status, row_count, notes
        ) VALUES (?, ?, ?, ?, ?, 'APPROVED', ?, ?)
        """,
        (source_id, source_type, source_path, source_sha256, as_of_date, row_count, notes),
    )


def insert_exception_specs(conn: sqlite3.Connection, specs: Iterable[ExceptionSpec], *, run_id: str) -> int:
    inserted = 0
    for spec in specs:
        cursor = conn.execute(
            """
            INSERT OR IGNORE INTO exception_queue (
                exception_id, run_id, domain, severity, status, reason, evidence_json
            ) VALUES (?, ?, ?, ?, 'OPEN', ?, ?)
            """,
            (
                spec.exception_id,
                run_id,
                spec.domain,
                spec.severity,
                f"{spec.reason_code}: {spec.reason}",
                json.dumps(spec.evidence, sort_keys=True),
            ),
        )
        inserted += cursor.rowcount
    return inserted


def write_inventory_snapshot(conn: sqlite3.Connection, snapshot: SnapshotBuildResult) -> int:
    conn.execute(
        "DELETE FROM fact_inventory_snapshot_size WHERE snapshot_date=?",
        (snapshot.snapshot_date,),
    )
    conn.executemany(
        """
        INSERT INTO fact_inventory_snapshot_size (
            snapshot_date, sku_id, sku_key, my_size, current_stock, inbound_stock
        ) VALUES (?, ?, ?, ?, ?, ?)
        """,
        [
            (
                snapshot.snapshot_date,
                row.sku_id,
                row.sku_key,
                row.my_size,
                row.current_stock,
                row.inbound_stock,
            )
            for row in snapshot.rows
        ],
    )
    return len(snapshot.rows)


def write_offer_availability_snapshot(
    conn: sqlite3.Connection,
    rows: Iterable[OfferAvailabilityRow],
    *,
    snapshot_date: str,
    source_manifest_id: str,
) -> int:
    conn.execute(
        "DELETE FROM offer_availability_snapshot WHERE snapshot_date=? AND source_manifest_id=?",
        (snapshot_date, source_manifest_id),
    )
    payload = list(rows)
    conn.executemany(
        """
        INSERT INTO offer_availability_snapshot (
            snapshot_id, snapshot_date, store_code, sku_key, sku_id, my_size,
            offer_available_qty, physical_stock_qty, status, source_manifest_id
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (
                row.snapshot_id,
                row.snapshot_date,
                row.store_code,
                row.sku_key,
                row.sku_id,
                row.my_size,
                row.offer_available_qty,
                row.physical_stock_qty,
                row.status,
                row.source_manifest_id,
            )
            for row in payload
        ],
    )
    return len(payload)


def count_adjustment_batch_duplicates(conn: sqlite3.Connection, batch_id: str) -> int:
    row = conn.execute(
        """
        SELECT COUNT(*) AS duplicate_groups
        FROM (
            SELECT sku_id, idempotency_key, COUNT(*) AS cnt
            FROM stock_ledger
            WHERE reference_type = ? AND reference_id = ?
            GROUP BY sku_id, idempotency_key
            HAVING COUNT(*) > 1
        )
        """,
        (BASELINE_REFERENCE_TYPE, batch_id),
    ).fetchone()
    return int(row["duplicate_groups"] or 0)


def count_legacy_duplicate_groups(conn: sqlite3.Connection) -> tuple[int, int]:
    row = conn.execute(
        """
        SELECT COUNT(*) AS duplicate_groups, COALESCE(SUM(cnt), 0) AS duplicate_rows
        FROM (
            SELECT event_date, event_type, sku_id, store_code, qty_change,
                   reference_id, reference_type, COUNT(*) AS cnt
            FROM stock_ledger
            WHERE COALESCE(reference_type, '') NOT IN (?, ?)
            GROUP BY event_date, event_type, sku_id, store_code, qty_change, reference_id, reference_type
            HAVING COUNT(*) > 1
        )
        """,
        (STOCK_ANCHOR_REFERENCE_TYPE, BASELINE_REFERENCE_TYPE),
    ).fetchone()
    return int(row["duplicate_groups"] or 0), int(row["duplicate_rows"] or 0)


def write_lineage_json(
    path: Path,
    *,
    anchor_id: str,
    batch_id: str,
    anchor_rows: list[AnchorStockRow],
    adjustment_plan: AdjustmentPlan,
    snapshot: SnapshotBuildResult,
    exceptions: list[ExceptionSpec],
    adjustment_duplicate_groups: int,
    legacy_duplicate_groups: int,
    legacy_duplicate_rows: int,
    apply: bool,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "anchor_id": anchor_id,
        "batch_id": batch_id,
        "apply": apply,
        "anchor_row_count": len(anchor_rows),
        "anchor_total_units": sum(row.qty for row in anchor_rows),
        "adjustment_plan": asdict(adjustment_plan),
        "snapshot": {
            "snapshot_date": snapshot.snapshot_date,
            "rows": len(snapshot.rows),
            "raw_total_units": snapshot.raw_total_units,
            "current_stock_total": snapshot.current_stock_total,
            "inbound_stock_total": snapshot.inbound_stock_total,
            "negative_rows": [asdict(row) for row in snapshot.negative_rows],
        },
        "exceptions": [asdict(spec) for spec in exceptions],
        "validation": {
            "adjustment_duplicate_groups": adjustment_duplicate_groups,
            "legacy_duplicate_groups": legacy_duplicate_groups,
            "legacy_duplicate_rows": legacy_duplicate_rows,
        },
    }
    path.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")


def write_snapshot_csv(path: Path, snapshot: SnapshotBuildResult, availability: list[OfferAvailabilityRow]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    availability_by_sku = {row.sku_id: row for row in availability}
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(
            [
                "snapshot_date",
                "sku_key",
                "sku_id",
                "my_size",
                "physical_stock_qty",
                "raw_physical_stock_qty",
                "inbound_stock_qty",
                "offer_available_qty",
                "availability_status",
            ]
        )
        for row in snapshot.rows:
            offer = availability_by_sku.get(row.sku_id)
            writer.writerow(
                [
                    snapshot.snapshot_date,
                    row.sku_key,
                    row.sku_id,
                    row.my_size,
                    row.current_stock,
                    row.raw_current_stock,
                    row.inbound_stock,
                    "" if offer is None or offer.offer_available_qty is None else offer.offer_available_qty,
                    offer.status if offer else "OFFER_AVAILABILITY_UNKNOWN",
                ]
            )


def write_owner_report(
    path: Path,
    *,
    trust_status: str,
    anchor_id: str,
    batch_id: str,
    snapshot: SnapshotBuildResult,
    exceptions: list[ExceptionSpec],
    adjustment_duplicate_groups: int,
    legacy_duplicate_groups: int,
    legacy_duplicate_rows: int,
    csv_path: Path,
    lineage_json_path: Path,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Owner Stock Output - Agent 6",
        "",
        f"Trust banner: {trust_status}",
        "",
        "Physical warehouse stock is separated from active sellable stock and offer availability.",
        "",
        "## Lineage",
        "",
        f"- Anchor: `{anchor_id}`",
        f"- Adjustment batch: `{batch_id}`",
        f"- Snapshot date: `{snapshot.snapshot_date}`",
        f"- Snapshot rows: `{len(snapshot.rows)}`",
        f"- Physical stock total: `{snapshot.current_stock_total}`",
        f"- Inbound stock total: `{snapshot.inbound_stock_total}`",
        f"- Raw negative rows clamped to zero: `{len(snapshot.negative_rows)}`",
        "",
        "## Validation",
        "",
        f"- Adjustment batch duplicate groups: `{adjustment_duplicate_groups}`",
        f"- Legacy ledger duplicate groups still unresolved: `{legacy_duplicate_groups}`",
        f"- Legacy ledger duplicate rows still unresolved: `{legacy_duplicate_rows}`",
        f"- Exception rows represented: `{len(exceptions)}`",
        "",
        "## Output Files",
        "",
        f"- CSV: `{csv_path}`",
        f"- Lineage JSON: `{lineage_json_path}`",
        "",
        "## Trust Notes",
        "",
    ]
    if trust_status != "TRUSTED":
        lines.extend(
            [
                "- This output is provisional/blocked for owner decisions until open exceptions and downstream strict-gate blockers are resolved.",
                "- Owner OOS exceptions force active sellable quantity to zero for the listed high-risk sizes.",
                "- Offer availability remains unknown except where owner exceptions force zero.",
            ]
        )
    else:
        lines.append("- All Agent 6 stock gates passed without unresolved exception blockers.")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def insert_pipeline_and_report(
    conn: sqlite3.Connection,
    *,
    run_id: str,
    snapshot_date: str,
    trust_status: str,
    report_path: str,
    source_manifest_json: str,
    exception_count: int,
    validation_messages: list[tuple[str, str, str, str]],
) -> None:
    conn.execute(
        """
        INSERT OR REPLACE INTO pipeline_run (
            run_id, run_type, as_of_date, status, source_manifest_json,
            validation_status, exception_count, finished_at, notes
        ) VALUES (?, 'STOCK_ANCHOR_20PCT_REBUILD', ?, ?, ?, ?, ?, datetime('now'), ?)
        """,
        (
            run_id,
            snapshot_date,
            "COMPLETE",
            source_manifest_json,
            trust_status,
            exception_count,
            "Agent 6 approved anchor and one-time baseline adjustment workflow.",
        ),
    )
    conn.execute(
        """
        INSERT OR REPLACE INTO owner_report_snapshot (
            snapshot_id, run_id, report_date, trust_status, report_path,
            source_manifest_json, exception_count
        ) VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            f"OWNER_STOCK_REPORT_{snapshot_date}",
            run_id,
            snapshot_date,
            trust_status,
            report_path,
            source_manifest_json,
            exception_count,
        ),
    )
    for gate_name, status, severity, message in validation_messages:
        conn.execute(
            """
            INSERT INTO validation_result (
                run_id, gate_name, status, severity, message
            ) VALUES (?, ?, ?, ?, ?)
            """,
            (run_id, gate_name, status, severity, message),
        )


def backup_db(db_path: Path, backup_dir: Path, label: str) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = backup_dir / f"app_db_before_{label}_{timestamp}.db"
    shutil.copy2(db_path, backup_path)
    return backup_path
