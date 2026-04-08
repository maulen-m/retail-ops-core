from __future__ import annotations

import csv
from datetime import datetime
import json
from pathlib import Path
import re
import shutil
from typing import Any
import zipfile

from openpyxl import load_workbook
import yaml


class ManifestError(RuntimeError):
    """Raised when an offer manifest is incomplete or invalid."""


ALLOWED_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp"}
BASE_PRODUCT_KEYS = [
    "store",
    "brand",
    "base_model",
    "display_title",
    "family_id",
    "description_path",
    "vendor_id",
    "product_type",
    "gender",
    "gender_rus",
    "season",
    "base_cost_cny",
    "base_cost_kzt",
    "weight_kg",
    "sell_price_kzt",
    "stock_entered_kzt",
]
LEGACY_PRODUCT_KEYS = [
    "color_slug",
    "color_label",
    "color_kaspi",
    "internal_base_sku_key",
    "kaspi_name_core",
]
VARIANT_REQUIRED_KEYS = [
    "color_slug",
    "color_label",
    "color_kaspi",
    "internal_base_sku_key",
    "kaspi_name_core",
    "external_token",
    "images",
    "sizes",
]


def _require_keys(mapping: dict[str, Any], keys: list[str], *, context: str) -> None:
    missing = [key for key in keys if key not in mapping]
    if missing:
        raise ManifestError(f"{context} missing required keys: {', '.join(missing)}")


def _resolve_path(base_dir: Path, value: str | None) -> Path | None:
    if not value:
        return None
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = (base_dir / path).resolve()
    return path


def _validate_sizes(sizes: list[dict[str, Any]], *, context: str) -> None:
    if not sizes:
        raise ManifestError(f"{context} must not be empty")
    for size in sizes:
        _require_keys(size, ["size_label", "numeric_size"], context=context)


def _resolve_images_cfg(images_cfg: dict[str, Any], manifest_dir: Path, *, context: str) -> dict[str, Any]:
    _require_keys(images_cfg, ["source_dir", "selection_mode", "upload_count"], context=context)
    images_cfg["source_dir"] = _resolve_path(manifest_dir, images_cfg["source_dir"])
    if images_cfg["source_dir"] is None or not images_cfg["source_dir"].is_dir():
        raise ManifestError(f"{context}.source_dir must point to an existing directory")
    mode = images_cfg.get("selection_mode", "folder_sequence_first_file")
    if mode == "folder_sequence_first_file":
        if "folder_names" not in images_cfg:
            raise ManifestError(f"{context}.folder_names is required for folder_sequence_first_file mode")
    elif mode == "explicit_filenames":
        if "filenames" not in images_cfg:
            raise ManifestError(f"{context}.filenames is required for explicit_filenames mode")
    else:
        raise ManifestError(f"Unsupported image selection_mode: {mode}")
    return images_cfg


def _now_timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def _midnight_now() -> datetime:
    now = datetime.now()
    return now.replace(hour=0, minute=0, second=0, microsecond=0)


def _render(value: Any, context: dict[str, Any]) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        return value.format_map(context)
    return value


def _sheet_display_map(ws) -> dict[str, int]:
    display_to_col: dict[str, int] = {}
    for col_num in range(1, int(ws.max_column or 0) + 1):
        display = ws.cell(3, col_num).value
        if display is None:
            continue
        display_to_col[str(display).strip()] = col_num
    return display_to_col


def _attribute_image_codes_from_xlsm(xlsm_path: Path) -> list[str]:
    wb = load_workbook(xlsm_path, keep_vba=True, data_only=False)
    try:
        if "attributes" not in wb.sheetnames:
            raise ManifestError(f"Template missing 'attributes' sheet: {xlsm_path}")
        ws = wb["attributes"]
        display_map = _sheet_display_map(ws)
        image_col = display_map.get("Код изображений")
        if image_col is None:
            raise ManifestError(f"Template missing 'Код изображений' column: {xlsm_path}")
        codes: list[str] = []
        seen: set[str] = set()
        for row_num in range(4, int(ws.max_row or 0) + 1):
            value = str(ws.cell(row_num, image_col).value or "").strip()
            if not value:
                continue
            if value not in seen:
                seen.add(value)
                codes.append(value)
        return codes
    finally:
        wb.close()


def validate_offer_package_layout(
    *,
    package_dir: Path,
    zip_path: Path | None = None,
) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    package_dir = Path(package_dir)
    if not package_dir.is_dir():
        raise ManifestError(f"Package directory missing: {package_dir}")

    xlsm_files = sorted(package_dir.glob("*.xlsm"))
    if len(xlsm_files) != 1:
        errors.append(f"Package must contain exactly one root XLSM, found {len(xlsm_files)}.")
        image_codes: list[str] = []
        xlsm_name = None
    else:
        xlsm_name = xlsm_files[0].name
        image_codes = _attribute_image_codes_from_xlsm(xlsm_files[0])
        if not image_codes:
            errors.append("Template contains no image-code rows in attributes sheet.")

    images_dir = package_dir / "images"
    if not images_dir.is_dir():
        errors.append("Package directory must contain lowercase images/ folder.")
    else:
        flat_files = sorted(path.name for path in images_dir.iterdir() if path.is_file())
        if flat_files:
            errors.append("Images must be stored in images/<image_code>/..., not flat files directly under images/.")
        for code in image_codes:
            folder = images_dir / code
            if not folder.is_dir():
                errors.append(f"Missing image folder for code '{code}' under images/<image_code>/.")
                continue
            files = sorted(path for path in folder.iterdir() if path.is_file())
            if not files:
                errors.append(f"Image folder '{code}' is empty.")
                continue
            if len(files) > 5:
                errors.append(f"Image folder '{code}' has {len(files)} files; Kaspi allows at most 5.")
            for file_path in files:
                if file_path.suffix.lower() not in ALLOWED_IMAGE_SUFFIXES:
                    errors.append(f"Unsupported image suffix in '{file_path.name}'.")

    if zip_path is not None:
        zip_path = Path(zip_path)
        if not zip_path.is_file():
            errors.append(f"ZIP file missing: {zip_path}")
        else:
            with zipfile.ZipFile(zip_path) as zf:
                names = [name for name in zf.namelist() if not name.endswith("/")]
            if any(name.startswith("__MACOSX/") for name in names):
                errors.append("ZIP must not contain __MACOSX entries.")
            root_xlsm = [name for name in names if "/" not in name and name.lower().endswith(".xlsm")]
            if len(root_xlsm) != 1:
                errors.append(f"ZIP must contain exactly one root XLSM, found {len(root_xlsm)}.")
            if any(name.count("/") and not name.startswith("images/") for name in names):
                errors.append("ZIP must not contain extra top-level folders; only root XLSM and images/ are allowed.")
            for name in names:
                if name.startswith("Images/"):
                    errors.append("ZIP must use lowercase images/ folder, not Images/.")
                    break
            flat_image_names = [
                name
                for name in names
                if re.fullmatch(r"images/[^/]+\.[A-Za-z0-9]+", name)
            ]
            if flat_image_names:
                errors.append("ZIP images must be nested as images/<image_code>/file, not images/file.")
            for code in image_codes:
                prefix = f"images/{code}/"
                code_files = [name for name in names if name.startswith(prefix)]
                if not code_files:
                    errors.append(f"ZIP missing image folder contents for '{code}' under {prefix}.")
                elif len(code_files) > 5:
                    errors.append(f"ZIP folder {prefix} has {len(code_files)} files; Kaspi allows at most 5.")

    return {
        "ok": not errors,
        "errors": errors,
        "warnings": warnings,
        "package_dir": str(package_dir),
        "zip_path": str(zip_path) if zip_path else None,
        "xlsm_name": xlsm_name,
        "image_codes": image_codes,
    }


def _clear_attribute_rows(ws) -> None:
    if ws.max_row <= 3:
        return
    for row_num in range(4, ws.max_row + 1):
        for col_num in range(1, ws.max_column + 1):
            ws.cell(row_num, col_num).value = None


def _ensure_dir(path: Path, *, apply: bool) -> None:
    if apply:
        path.mkdir(parents=True, exist_ok=True)


def _copy_file(src: Path, dst: Path, *, apply: bool) -> None:
    if apply:
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)


def _write_text(path: Path, text: str, *, apply: bool) -> None:
    if apply:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")


def _write_json(path: Path, payload: Any, *, apply: bool) -> None:
    if apply:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def _select_source_images(images_cfg: dict[str, Any]) -> list[Path]:
    source_dir = images_cfg["source_dir"]
    mode = images_cfg.get("selection_mode", "folder_sequence_first_file")
    selected: list[Path] = []
    if mode == "folder_sequence_first_file":
        for folder_name in images_cfg["folder_names"]:
            folder_path = source_dir / str(folder_name)
            if not folder_path.is_dir():
                raise ManifestError(f"Image source folder missing: {folder_path}")
            candidates = sorted(
                path
                for path in folder_path.iterdir()
                if path.is_file() and path.suffix.lower() in ALLOWED_IMAGE_SUFFIXES
            )
            if not candidates:
                raise ManifestError(f"No image files found in {folder_path}")
            selected.append(candidates[0])
    elif mode == "explicit_filenames":
        for filename in images_cfg["filenames"]:
            file_path = source_dir / str(filename)
            if not file_path.is_file():
                raise ManifestError(f"Image source file missing: {file_path}")
            if file_path.suffix.lower() not in ALLOWED_IMAGE_SUFFIXES:
                raise ManifestError(f"Unsupported image suffix in {file_path}")
            selected.append(file_path)
    else:
        raise ManifestError(f"Unsupported image selection_mode: {mode}")

    upload_count = int(images_cfg["upload_count"])
    if upload_count <= 0:
        raise ManifestError("images.upload_count must be > 0")
    if upload_count > len(selected):
        raise ManifestError("images.upload_count cannot exceed number of selected source images")
    return selected


def _manifest_variants(manifest: dict[str, Any]) -> list[dict[str, Any]]:
    if manifest.get("variants"):
        return list(manifest["variants"])
    product = manifest["product"]
    return [
        {
            "color_slug": product["color_slug"],
            "color_label": product["color_label"],
            "color_title": str(product["color_kaspi"]).replace(", ", " / "),
            "color_kaspi": product["color_kaspi"],
            "internal_base_sku_key": product["internal_base_sku_key"],
            "kaspi_name_core": product["kaspi_name_core"],
            "description_path": product["description_path"],
            "images": manifest["images"],
            "sizes": manifest["sizes"],
            "external_token": product.get("color_slug", "BASE").upper(),
        }
    ]


def load_offer_manifest(manifest_path: Path | str) -> dict[str, Any]:
    path = Path(manifest_path).expanduser().resolve()
    raw = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        raise ManifestError(f"Manifest must be a mapping: {path}")

    manifest_dir = path.parent
    _require_keys(raw, ["schema_version", "product", "categories"], context="manifest")
    has_variants = bool(raw.get("variants"))
    _require_keys(
        raw["product"],
        BASE_PRODUCT_KEYS + ([] if has_variants else LEGACY_PRODUCT_KEYS),
        context="manifest.product",
    )
    if not raw["categories"]:
        raise ManifestError("manifest.categories must not be empty")

    product_root = _resolve_path(manifest_dir, raw.get("product_root"))
    if product_root is None:
        product_root = manifest_dir
    build_root = _resolve_path(manifest_dir, raw.get("build_root"))
    if build_root is None:
        build_root = product_root

    raw["manifest_path"] = path
    raw["manifest_dir"] = manifest_dir
    raw["product_root"] = product_root
    raw["build_root"] = build_root

    raw["product"]["description_path"] = _resolve_path(manifest_dir, raw["product"]["description_path"])
    if raw["product"]["description_path"] is None or not raw["product"]["description_path"].is_file():
        raise ManifestError("product.description_path must point to an existing file")

    if has_variants:
        if not raw["variants"]:
            raise ManifestError("manifest.variants must not be empty")
        for idx, variant in enumerate(raw["variants"]):
            _require_keys(variant, VARIANT_REQUIRED_KEYS, context=f"manifest.variants[{idx}]")
            variant["description_path"] = _resolve_path(
                manifest_dir,
                variant.get("description_path") or str(raw["product"]["description_path"]),
            )
            if variant["description_path"] is None or not variant["description_path"].is_file():
                raise ManifestError(f"manifest.variants[{idx}].description_path must point to an existing file")
            variant["images"] = _resolve_images_cfg(
                variant["images"],
                manifest_dir,
                context=f"manifest.variants[{idx}].images",
            )
            _validate_sizes(variant["sizes"], context=f"manifest.variants[{idx}].sizes[]")
            if not variant.get("color_title"):
                variant["color_title"] = str(variant["color_kaspi"]).replace(", ", " / ")
    else:
        _require_keys(raw, ["images", "sizes"], context="manifest")
        raw["images"] = _resolve_images_cfg(raw["images"], manifest_dir, context="manifest.images")
        _validate_sizes(raw["sizes"], context="manifest.sizes[]")

    for category in raw["categories"]:
        _require_keys(
            category,
            [
                "slug",
                "short_code",
                "template_path",
                "package_slug",
                "zip_slug",
                "group_token",
                "internal_article_pattern",
                "external_sku_pattern",
                "offer_name_pattern",
                "shared_fields",
                "row_fields",
            ],
            context="manifest.categories[]",
        )
        category["template_path"] = _resolve_path(manifest_dir, category["template_path"])
        if category["template_path"] is None or not category["template_path"].is_file():
            raise ManifestError(f"Category template missing: {category.get('slug')}")

    workbook_cfg = raw.get("workbook")
    if workbook_cfg:
        _require_keys(workbook_cfg, ["workbook_path", "agent_session_title"], context="manifest.workbook")
        workbook_cfg["workbook_path"] = _resolve_path(manifest_dir, workbook_cfg["workbook_path"])

    return raw


def _copy_source_assets(
    manifest: dict[str, Any],
    *,
    timestamp: str,
    apply: bool,
) -> dict[str, Any]:
    source_root = manifest["product_root"] / "sources"
    reports_dir = source_root / "reports"
    _ensure_dir(source_root, apply=apply)
    _ensure_dir(reports_dir, apply=apply)

    variants = _manifest_variants(manifest)
    is_multi_variant = len(variants) > 1
    variant_sources: dict[str, Any] = {}
    mapping_rows: list[dict[str, Any]] = []

    for variant in variants:
        selected = _select_source_images(variant["images"])
        selected_count = len(selected)
        upload_count = int(variant["images"]["upload_count"])
        variant_root = source_root if not is_multi_variant else source_root / "by_color" / variant["color_slug"]
        original_dir = variant_root / f"original_{selected_count}"
        renamed_dir = variant_root / f"renamed_1_{selected_count}"
        images_part_2_dir = variant_root / "images_part_2"
        for path in [variant_root, original_dir, renamed_dir, images_part_2_dir]:
            _ensure_dir(path, apply=apply)

        for idx, src in enumerate(selected, start=1):
            original_dst = original_dir / src.name
            renamed_dst = renamed_dir / f"{idx}{src.suffix.lower()}"
            _copy_file(src, original_dst, apply=apply)
            _copy_file(src, renamed_dst, apply=apply)
            if idx > upload_count:
                _copy_file(src, images_part_2_dir / renamed_dst.name, apply=apply)
            mapping_rows.append(
                {
                    "variant_slug": variant["color_slug"],
                    "index": idx,
                    "source_path": str(src),
                    "original_copy": str(original_dst),
                    "renamed_copy": str(renamed_dst),
                    "upload_batch": idx <= upload_count,
                }
            )

        variant_sources[variant["color_slug"]] = {
            "selected_images": selected,
            "selected_count": selected_count,
            "upload_count": upload_count,
            "variant_root": variant_root,
            "original_dir": original_dir,
            "renamed_dir": renamed_dir,
            "images_part_2_dir": images_part_2_dir,
        }

    lines = [
        "# Image Migration Map",
        "",
        f"- Generated: {timestamp}",
        f"- Product root: {manifest['product_root']}",
        "",
        "| Variant | # | Source | Original Copy | Renamed | Upload Batch |",
        "|---------|---|--------|---------------|---------|--------------|",
    ]
    for row in mapping_rows:
        lines.append(
            f"| {row['variant_slug']} | {row['index']} | {row['source_path']} | {row['original_copy']} | {row['renamed_copy']} | {'yes' if row['upload_batch'] else 'no'} |"
        )
    mapping_text = "\n".join(lines) + "\n"
    mapping_path = source_root / "IMAGE_MIGRATION_MAP.md"
    _write_text(mapping_path, mapping_text, apply=apply)
    report_path = reports_dir / f"IMAGE_MIGRATION_MAP_{timestamp}.md"
    _write_text(report_path, mapping_text, apply=apply)

    for variant in variants:
        source_side_note = variant["images"]["source_dir"] / (
            f"{manifest['product']['base_model']}_{variant['color_slug'].upper()}_IMAGE_MIGRATION_TO_PRODUCT_OFFERS_{timestamp}.md"
        )
        _write_text(source_side_note, mapping_text, apply=apply)

    return {
        "source_root": source_root,
        "mapping_rows": mapping_rows,
        "mapping_path": mapping_path,
        "variant_sources": variant_sources,
    }


def _base_context(
    manifest: dict[str, Any],
    category: dict[str, Any],
    variant: dict[str, Any],
    *,
    size_label: str | None = None,
    numeric_size: int | None = None,
) -> dict[str, Any]:
    product = manifest["product"]
    first_size = variant["sizes"][0]
    context: dict[str, Any] = {
        "brand": product["brand"],
        "base_model": product["base_model"],
        "color_slug": variant["color_slug"],
        "color_label": variant["color_label"],
        "color_title": variant.get("color_title", variant["color_kaspi"]),
        "color_kaspi": variant["color_kaspi"],
        "description_text": variant["description_path"].read_text(encoding="utf-8").strip(),
        "display_title": product["display_title"],
        "family_id": product["family_id"],
        "gender": product["gender"],
        "gender_rus": product["gender_rus"],
        "internal_base_sku_key": variant["internal_base_sku_key"],
        "kaspi_name_core": variant["kaspi_name_core"],
        "product_type": product["product_type"],
        "season": product["season"],
        "short_code": category["short_code"],
        "external_token": variant.get("external_token"),
        "size_label": size_label,
        "numeric_size": numeric_size,
        "first_size_label": first_size["size_label"],
        "first_numeric_size": first_size["numeric_size"],
        "variant_slug": variant["color_slug"],
    }
    if size_label is not None:
        context["internal_article"] = category["internal_article_pattern"].format_map(context)
        context["external_sku"] = category["external_sku_pattern"].format_map(context)
        context["offer_name"] = category["offer_name_pattern"].format_map(context)
    shared_code_ctx = dict(context)
    shared_code_ctx["size_label"] = first_size["size_label"]
    shared_code_ctx["numeric_size"] = first_size["numeric_size"]
    context["shared_image_code"] = category["internal_article_pattern"].format_map(shared_code_ctx)
    return context


def _fixed_template_name(template_path: Path) -> str:
    if template_path.stem.endswith("_FIXED"):
        return template_path.name
    return f"{template_path.stem}_FIXED{template_path.suffix}"


def _write_category_template(
    manifest: dict[str, Any],
    category: dict[str, Any],
    *,
    package_template_path: Path,
    apply: bool,
) -> list[dict[str, Any]]:
    row_records: list[dict[str, Any]] = []
    variants = _manifest_variants(manifest)
    if not apply:
        for variant in variants:
            for size in variant["sizes"]:
                ctx = _base_context(
                    manifest,
                    category,
                    variant,
                    size_label=str(size["size_label"]),
                    numeric_size=int(size["numeric_size"]),
                )
                row_records.append(
                    {
                        "variant_slug": ctx["variant_slug"],
                        "color_label": ctx["color_label"],
                        "color_kaspi": ctx["color_kaspi"],
                        "internal_base_sku_key": ctx["internal_base_sku_key"],
                        "kaspi_name_core": ctx["kaspi_name_core"],
                        "shared_image_code": ctx["shared_image_code"],
                        "size_label": ctx["size_label"],
                        "numeric_size": ctx["numeric_size"],
                        "internal_article": ctx["internal_article"],
                        "external_sku": ctx["external_sku"],
                        "offer_name": ctx["offer_name"],
                    }
                )
        return row_records

    wb = load_workbook(package_template_path, keep_vba=True)
    if "attributes" not in wb.sheetnames:
        raise ManifestError(f"Template missing 'attributes' sheet: {package_template_path}")
    ws = wb["attributes"]
    display_map = _sheet_display_map(ws)
    required_displays = set(category["shared_fields"].keys()) | set(category["row_fields"].keys())
    missing = sorted(display for display in required_displays if display not in display_map)
    if missing:
        raise ManifestError(
            f"Template {package_template_path.name} missing required display columns: {', '.join(missing)}"
        )
    _clear_attribute_rows(ws)

    row_num = 4
    for variant in variants:
        for size in variant["sizes"]:
            ctx = _base_context(
                manifest,
                category,
                variant,
                size_label=str(size["size_label"]),
                numeric_size=int(size["numeric_size"]),
            )
            for display_name, raw_value in category["shared_fields"].items():
                ws.cell(row_num, display_map[display_name]).value = _render(raw_value, ctx)
            for display_name, raw_value in category["row_fields"].items():
                ws.cell(row_num, display_map[display_name]).value = _render(raw_value, ctx)
            row_records.append(
                {
                    "variant_slug": ctx["variant_slug"],
                    "color_label": ctx["color_label"],
                    "color_kaspi": ctx["color_kaspi"],
                    "internal_base_sku_key": ctx["internal_base_sku_key"],
                    "kaspi_name_core": ctx["kaspi_name_core"],
                    "shared_image_code": ctx["shared_image_code"],
                    "size_label": ctx["size_label"],
                    "numeric_size": ctx["numeric_size"],
                    "internal_article": ctx["internal_article"],
                    "external_sku": ctx["external_sku"],
                    "offer_name": ctx["offer_name"],
                }
            )
            row_num += 1

    wb.save(package_template_path)
    wb.close()
    return row_records


def _write_row_mapping_csv(path: Path, rows: list[dict[str, Any]], *, apply: bool) -> None:
    if not apply:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "variant_slug",
        "color_label",
        "color_kaspi",
        "internal_base_sku_key",
        "kaspi_name_core",
        "shared_image_code",
        "size_label",
        "numeric_size",
        "internal_article",
        "external_sku",
        "offer_name",
    ]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _write_build_log(
    path: Path,
    *,
    manifest: dict[str, Any],
    category: dict[str, Any],
    timestamp: str,
    rows: list[dict[str, Any]],
    zip_path: Path,
    apply: bool,
) -> None:
    lines = [
        "# Build Log",
        "",
        f"- Timestamp: {timestamp}",
        f"- Manifest: {manifest['manifest_path']}",
        f"- Product root: {manifest['product_root']}",
        f"- Category: {category['slug']}",
        f"- Package slug: {category['package_slug']}",
        f"- ZIP: {zip_path}",
        f"- Rows: {len(rows)}",
        f"- Image groups: {', '.join(sorted({row['shared_image_code'] for row in rows})) if rows else ''}",
        "",
        "| Variant | Size | Numeric | Internal Article | External SKU |",
        "|---------|------|---------|------------------|--------------|",
    ]
    for row in rows:
        lines.append(
            f"| {row['variant_slug']} | {row['size_label']} | {row['numeric_size']} | {row['internal_article']} | {row['external_sku']} |"
        )
    _write_text(path, "\n".join(lines) + "\n", apply=apply)


def _zip_package(package_dir: Path, *, xlsm_name: str, apply: bool) -> Path:
    zip_candidates = list(package_dir.glob("*.zip"))
    if zip_candidates and not apply:
        return zip_candidates[0]
    zip_path = package_dir / f"{package_dir.name}_UPLOAD.zip"
    if apply:
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            zf.write(package_dir / xlsm_name, arcname=xlsm_name)
            for image_path in sorted((package_dir / "images").glob("*")):
                zf.write(image_path, arcname=f"images/{image_path.name}")
    return zip_path


def build_offer_packages(
    manifest: dict[str, Any],
    *,
    timestamp: str | None = None,
    apply: bool = True,
) -> dict[str, Any]:
    ts = timestamp or _now_timestamp()
    source_report = _copy_source_assets(manifest, timestamp=ts, apply=apply)
    product_root = manifest["product_root"]
    build_root = manifest.get("build_root") or product_root
    batch_root = build_root / ts if Path(build_root) != Path(product_root) else build_root

    templates_dir = product_root / "templates"
    _ensure_dir(templates_dir, apply=apply)
    description_copy = product_root / manifest["product"]["description_path"].name
    _copy_file(manifest["product"]["description_path"], description_copy, apply=apply)

    for category in manifest["categories"]:
        _copy_file(category["template_path"], templates_dir / category["template_path"].name, apply=apply)

    packages: list[dict[str, Any]] = []
    for category in manifest["categories"]:
        category_root = batch_root / category["slug"]
        package_dir = category_root / f"{category['package_slug']}_{ts}"
        _ensure_dir(category_root, apply=apply)
        _ensure_dir(package_dir, apply=apply)
        _ensure_dir(package_dir / "images", apply=apply)
        _ensure_dir(package_dir / "backups", apply=apply)

        template_name = _fixed_template_name(category["template_path"])
        package_template_path = package_dir / template_name
        backup_template_path = package_dir / "backups" / f"{Path(template_name).stem}.{ts}{Path(template_name).suffix}"
        _copy_file(category["template_path"], package_template_path, apply=apply)
        _copy_file(category["template_path"], backup_template_path, apply=apply)
        _copy_file(manifest["product"]["description_path"], package_dir / manifest["product"]["description_path"].name, apply=apply)

        row_records = _write_category_template(
            manifest,
            category,
            package_template_path=package_template_path,
            apply=apply,
        )

        image_groups: dict[str, list[dict[str, Any]]] = {}
        for row in row_records:
            image_groups.setdefault(row["shared_image_code"], []).append(row)
        for shared_image_code, rows in image_groups.items():
            image_target_dir = package_dir / "images" / shared_image_code
            _ensure_dir(image_target_dir, apply=apply)
            variant_source = source_report["variant_sources"][rows[0]["variant_slug"]]
            for idx in range(1, variant_source["upload_count"] + 1):
                src = variant_source["renamed_dir"] / f"{idx}{variant_source['selected_images'][idx - 1].suffix.lower()}"
                _copy_file(src, image_target_dir / src.name, apply=apply)

        zip_path = package_dir / f"{category['zip_slug']}_{ts}.zip"
        if apply:
            with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
                zf.write(package_template_path, arcname=package_template_path.name)
                for image_dir in sorted(path for path in (package_dir / "images").iterdir() if path.is_dir()):
                    for image_path in sorted(path for path in image_dir.iterdir() if path.is_file()):
                        zf.write(image_path, arcname=f"images/{image_dir.name}/{image_path.name}")

        if apply:
            package_validation = validate_offer_package_layout(package_dir=package_dir, zip_path=zip_path)
        else:
            package_validation = {
                "ok": True,
                "errors": [],
                "warnings": [],
                "package_dir": str(package_dir),
                "zip_path": None,
                "xlsm_name": template_name,
                "image_codes": [shared_image_code] if shared_image_code else [],
            }
        if apply and not package_validation["ok"]:
            raise ManifestError(
                f"Kaspi package validation failed for {package_dir.name}: " + "; ".join(package_validation["errors"])
            )

        _write_row_mapping_csv(package_dir / "ROW_MAPPING.csv", row_records, apply=apply)
        _write_build_log(
            package_dir / "BUILD_LOG.md",
            manifest=manifest,
            category=category,
            timestamp=ts,
            rows=row_records,
            zip_path=zip_path,
            apply=apply,
        )
        packages.append(
            {
                "category_slug": category["slug"],
                "short_code": category["short_code"],
                "group_token": category["group_token"],
                "package_dir": str(package_dir),
                "template_path": str(package_template_path),
                "zip_path": str(zip_path),
                "rows": row_records,
                "source": f"{category['package_slug']}_{ts}",
                "package_validation": package_validation,
            }
        )

    report = {
        "manifest_path": str(manifest["manifest_path"]),
        "product_root": str(product_root),
        "build_root": str(build_root),
        "batch_root": str(batch_root),
        "timestamp": ts,
        "packages": packages,
        "source_assets": {
            "mapping_path": str(source_report["mapping_path"]),
            "variant_sources": {
                slug: {
                    "selected_count": payload["selected_count"],
                    "upload_count": payload["upload_count"],
                    "original_dir": str(payload["original_dir"]),
                    "renamed_dir": str(payload["renamed_dir"]),
                    "images_part_2_dir": str(payload["images_part_2_dir"]),
                }
                for slug, payload in source_report["variant_sources"].items()
            },
        },
    }
    return report


def build_workbook_ingest_payload(manifest: dict[str, Any], build_report: dict[str, Any]) -> dict[str, Any]:
    product = manifest["product"]
    workbook_cfg = manifest.get("workbook", {})
    entered_at = _midnight_now()

    display_title = product["display_title"]
    candidate_color = str(workbook_cfg.get("candidate_color", "black"))

    m02_rows: list[list[Any]] = []
    sku_map_rows: list[list[Any]] = []
    ingest_candidates: list[list[Any]] = []
    size_engine_rows: list[list[Any]] = []
    offer_map_blocks: list[list[list[Any]]] = []
    agent_journal_rows: list[list[Any]] = [
        [workbook_cfg["agent_session_title"], "Manifest-driven Kaspi offer workflow"],
        ["Manifest", str(manifest["manifest_path"])],
        ["Product root", str(manifest["product_root"])],
        ["Build timestamp", build_report["timestamp"]],
    ]

    sku_map_seen: set[str] = set()
    for package in build_report["packages"]:
        category = next(item for item in manifest["categories"] if item["group_token"] == package["group_token"])
        short_code = category["short_code"]
        grouped_rows: dict[str, list[dict[str, Any]]] = {}
        for row in package["rows"]:
            grouped_rows.setdefault(row["variant_slug"], []).append(row)

        for variant_slug, grouped in grouped_rows.items():
            block_header = (
                f"NEW OFFER GROUP — {package['group_token']}"
                if len(grouped_rows) == 1
                else f"NEW OFFER GROUP — {package['group_token']} — {variant_slug}"
            )
            block_rows: list[list[Any]] = [
                [
                    block_header,
                    f'| "{display_title}" | Status: package_ready | Source: {package["source"]}',
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                ],
                [
                    "#",
                    "Kaspi Offer Name (from package)",
                    "SKU_ID_KSP (prepared)",
                    "Probable Size",
                    "Size RUS",
                    "Variant URL",
                    "SKU_ID (normalized)",
                    "Status",
                    "Kaspi Size #",
                    "Product Code",
                    "SKU_key",
                    "Merchant SKU",
                    None,
                    None,
                    None,
                    None,
                ],
            ]

            for idx, row in enumerate(grouped, start=1):
                size_label = row["size_label"]
                numeric_size = row["numeric_size"]
                base_sku_key = row["internal_base_sku_key"]
                kaspi_name_core = row["kaspi_name_core"]
                color_label = row["color_label"]
                color_kaspi = row["color_kaspi"]
                sku_id = f"{base_sku_key}_{size_label}"
                offer_name = row["offer_name"]

                m02_rows.append(
                    [
                        product["store"],
                        sku_id,
                        row["internal_article"],
                        kaspi_name_core,
                        size_label,
                        size_label,
                        numeric_size,
                        None,
                        None,
                        row["external_sku"],
                        f"{kaspi_name_core}_{product['gender_rus']}_{numeric_size}_({short_code})_({product['base_model']})_({size_label})",
                        product["sell_price_kzt"],
                        product["stock_entered_kzt"],
                        None,
                        entered_at,
                        None,
                        None,
                        base_sku_key,
                        "YES",
                        offer_name,
                        None,
                        product["product_type"],
                        None,
                        product["brand"],
                        product["base_model"],
                        color_label,
                        product["gender"],
                        product["gender_rus"],
                        product["season"],
                        product["base_cost_cny"],
                        None,
                    ]
                )

                ingest_candidates.append(
                    [
                        package["group_token"],
                        candidate_color if len(grouped_rows) == 1 else variant_slug.lower(),
                        None,
                        numeric_size,
                        size_label,
                        f"{numeric_size} RUS",
                        None,
                        display_title,
                        numeric_size,
                        sku_id,
                        base_sku_key,
                        offer_name,
                        sku_id,
                        size_label,
                        row["external_sku"],
                        None,
                        0,
                        None,
                        None,
                        None,
                    ]
                )

                size_engine_rows.append([offer_name] + [None] * 12 + [size_label, 1, None])
                size_engine_rows.append([f"{display_title} {color_kaspi} {numeric_size}"] + [None] * 12 + [size_label, 1, None])

                block_rows.append(
                    [
                        idx,
                        offer_name,
                        row["external_sku"],
                        size_label,
                        numeric_size,
                        None,
                        sku_id,
                        "package_ready",
                        numeric_size,
                        None,
                        base_sku_key,
                        row["internal_article"],
                        None,
                        None,
                        None,
                        None,
                    ]
                )

                if sku_id not in sku_map_seen:
                    sku_map_seen.add(sku_id)
                    sku_map_rows.append(
                        [
                            sku_id,
                            base_sku_key,
                            size_label,
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                            0,
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                            kaspi_name_core,
                            None,
                            None,
                            product["product_type"],
                            None,
                            product["brand"],
                            product["base_model"],
                            color_label,
                            product["gender"],
                            product["season"],
                            product["base_cost_cny"],
                            product["base_cost_kzt"],
                            product["weight_kg"],
                            product["vendor_id"],
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                        ]
                    )

            offer_map_blocks.append(block_rows)

    return {
        "workbook_path": str(workbook_cfg.get("workbook_path") or ""),
        "timestamp": build_report["timestamp"],
        "m02_rows": m02_rows,
        "sku_map_rows": sku_map_rows,
        "ingest_candidates": ingest_candidates,
        "size_engine_rows": size_engine_rows,
        "offer_map_blocks": offer_map_blocks,
        "agent_journal_rows": agent_journal_rows,
    }


def _collect_workbook_state(workbook_path: Path) -> dict[str, Any]:
    wb = load_workbook(workbook_path, read_only=True, data_only=False, keep_vba=True)
    state = {
        "max_rows": {},
        "m02_existing": set(),
        "sku_map_existing": set(),
        "ingest_existing": set(),
        "size_engine_existing": set(),
        "offer_group_headers": set(),
        "agent_journal_existing": set(),
    }
    try:
        ws = wb["M02_SKU_CATALOG_NC"]
        state["max_rows"]["M02_SKU_CATALOG_NC"] = ws.max_row
        for row in ws.iter_rows(min_row=2, values_only=True):
            value = row[9] if len(row) > 9 else None
            if value:
                state["m02_existing"].add(str(value))

        ws = wb["Sku_Map_CRM_01"]
        state["max_rows"]["Sku_Map_CRM_01"] = ws.max_row
        for row in ws.iter_rows(min_row=2, values_only=True):
            value = row[0] if row else None
            if value:
                state["sku_map_existing"].add(str(value))

        ws = wb["ingest_candidates"]
        state["max_rows"]["ingest_candidates"] = ws.max_row
        for row in ws.iter_rows(min_row=2, values_only=True):
            value = row[14] if len(row) > 14 else None
            if value:
                state["ingest_existing"].add(str(value))

        ws = wb["SIZE_engine_basic"]
        state["max_rows"]["SIZE_engine_basic"] = ws.max_row
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = row[0] if len(row) > 0 else None
            probable = row[13] if len(row) > 13 else None
            if name and probable:
                state["size_engine_existing"].add((str(name), str(probable)))

        ws = wb["SKU_Offer_Map_v2"]
        state["max_rows"]["SKU_Offer_Map_v2"] = ws.max_row
        for row in ws.iter_rows(min_row=1, values_only=True):
            head = row[0] if row else None
            if head:
                state["offer_group_headers"].add(str(head))

        ws = wb["AGENT_JOURNAL"]
        state["max_rows"]["AGENT_JOURNAL"] = ws.max_row
        for row in ws.iter_rows(min_row=1, values_only=True):
            head = row[0] if row else None
            if head:
                state["agent_journal_existing"].add(str(head))
    finally:
        wb.close()
    return state


def _append_rows_xlwings(sheet, start_row: int, rows: list[list[Any]]) -> int:
    if not rows:
        return 0
    sheet.range((start_row, 1)).value = rows
    return len(rows)


def apply_workbook_ingest_payload(
    workbook_path: Path,
    payload: dict[str, Any],
    *,
    backup_dir: Path | None = None,
) -> dict[str, Any]:
    try:
        import xlwings as xw  # type: ignore
    except Exception as exc:  # pragma: no cover - environment dependent
        raise RuntimeError("xlwings is required for workbook apply mode") from exc

    workbook_path = workbook_path.expanduser().resolve()
    if not workbook_path.is_file():
        raise ManifestError(f"Workbook not found: {workbook_path}")

    state = _collect_workbook_state(workbook_path)

    filtered_m02 = [row for row in payload["m02_rows"] if str(row[9]) not in state["m02_existing"]]
    filtered_sku_map = [row for row in payload["sku_map_rows"] if str(row[0]) not in state["sku_map_existing"]]
    filtered_ingest = [row for row in payload["ingest_candidates"] if str(row[14]) not in state["ingest_existing"]]
    filtered_size_engine = [
        row
        for row in payload["size_engine_rows"]
        if (str(row[0]), str(row[13])) not in state["size_engine_existing"]
    ]
    filtered_offer_blocks = [
        block for block in payload["offer_map_blocks"] if str(block[0][0]) not in state["offer_group_headers"]
    ]
    filtered_journal = [
        row for row in payload["agent_journal_rows"] if str(row[0]) not in state["agent_journal_existing"]
    ]

    safe_backup_dir = backup_dir or workbook_path.parent / "backups" / "offer_manifest_ingest"
    safe_backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = safe_backup_dir / f"{workbook_path.stem}.{_now_timestamp()}{workbook_path.suffix}"
    shutil.copy2(workbook_path, backup_path)

    app = xw.App(visible=False, add_book=False)
    app.display_alerts = False
    app.screen_updating = False
    write_counts = {
        "m02_rows_added": 0,
        "sku_map_rows_added": 0,
        "ingest_rows_added": 0,
        "size_engine_rows_added": 0,
        "offer_map_rows_added": 0,
        "agent_journal_rows_added": 0,
    }
    try:  # pragma: no cover - Excel runtime path
        wb = app.books.open(str(workbook_path), update_links=False, read_only=False)
        try:
            if filtered_m02:
                start = state["max_rows"]["M02_SKU_CATALOG_NC"] + 1
                write_counts["m02_rows_added"] = _append_rows_xlwings(
                    wb.sheets["M02_SKU_CATALOG_NC"],
                    start,
                    filtered_m02,
                )
            if filtered_sku_map:
                start = state["max_rows"]["Sku_Map_CRM_01"] + 1
                write_counts["sku_map_rows_added"] = _append_rows_xlwings(
                    wb.sheets["Sku_Map_CRM_01"],
                    start,
                    filtered_sku_map,
                )
            if filtered_ingest:
                start = state["max_rows"]["ingest_candidates"] + 1
                write_counts["ingest_rows_added"] = _append_rows_xlwings(
                    wb.sheets["ingest_candidates"],
                    start,
                    filtered_ingest,
                )
            if filtered_size_engine:
                start = state["max_rows"]["SIZE_engine_basic"] + 1
                write_counts["size_engine_rows_added"] = _append_rows_xlwings(
                    wb.sheets["SIZE_engine_basic"],
                    start,
                    filtered_size_engine,
                )
            if filtered_offer_blocks:
                start = state["max_rows"]["SKU_Offer_Map_v2"] + 1
                offer_rows = [row for block in filtered_offer_blocks for row in block]
                write_counts["offer_map_rows_added"] = _append_rows_xlwings(
                    wb.sheets["SKU_Offer_Map_v2"],
                    start,
                    offer_rows,
                )
            if filtered_journal:
                start = state["max_rows"]["AGENT_JOURNAL"] + 1
                write_counts["agent_journal_rows_added"] = _append_rows_xlwings(
                    wb.sheets["AGENT_JOURNAL"],
                    start,
                    filtered_journal,
                )
            wb.save()
        finally:
            wb.close()
    finally:
        app.quit()

    post_state = _collect_workbook_state(workbook_path)
    return {
        "backup_path": str(backup_path),
        "workbook": str(workbook_path),
        **write_counts,
        "post_max_rows": post_state["max_rows"],
    }
